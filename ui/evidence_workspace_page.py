"""사업장 업무 화면: 신입사원도 따라올 수 있게 단계별로 안내합니다.

① 엑셀 작성·올리기 → ② 입력 확인·수정 → ③ 시나리오 후보 → ④ 작성·내려받기

보완 요청은 ②의 한 표에 모으고 ③은 후보 표, ④는 작성 동작과 간단한 진행 상태만 보여 줍니다.
"""
from __future__ import annotations

import re

import streamlit as st

from engine.stage2 import demo_company_workbook, scenario_triage
from engine.stage2 import cap_evidence_links, cap_fact_drafting, cap_fact_status, cap_upload_template
from engine.stage2.intake import build_intake_catalog
from engine.stage2 import company_workbook_intake as company_intake
from engine.stage2 import evidence_workspace as ew
from engine.stage2.storage import delete_all_projects, list_projects, load_project, save_attachment, save_project

STEPS = {1: "① 엑셀 작성·올리기", 2: "② 입력 확인·수정", 3: "③ 시나리오 후보", 4: "④ 작성·내려받기"}
MISSING_FIELDS = ("product", "inventory", "pressure", "temperature", "facility_type", "substance_state")

# ---------------------------------------------------------------- 사람 말로 바꾸기

def _unit(field: str) -> str:
    match = re.search(r"\(([^)]+)\)$", ew.LABELS[field])
    return match.group(1) if match else ""


def _name(field: str) -> str:
    return re.sub(r"\s*\([^)]*\)$", "", ew.LABELS[field])


def _number(value) -> str:
    if isinstance(value, bool) or not isinstance(value, (int, float)):
        return str(value)
    if abs(value) >= 1e6:
        return f"{value:,.0f}"
    return f"{value:,.6g}"


def _value(field: str, value) -> str:
    unit = _unit(field)
    text = _number(value)
    if not unit or not isinstance(value, (int, float)):
        return text
    return f"{text}{unit}" if unit == "%" else f"{text} {unit}"


def _where(fact) -> str:
    """‘CSV 행 2 열 4’ 같은 내부 위치를 ‘02_equipment.csv 2번째 줄’로 바꿉니다."""
    doc, location = fact["document"], str(fact["location"])
    if doc == "담당자 입력":
        return "직접 입력"
    match = re.match(r"^(.*?)\s*행 (\d+)(?: 열 \d+)?$", location)
    if not match:
        return f"{doc} ({location})"
    place, row = match.groups()
    if place == "CSV":
        return f"{doc} {row}번째 줄"
    if doc.lower().endswith((".xlsx", ".xlsm")):
        return f"{doc} ‘{place}’ 시트 {row}번째 줄"
    return f"{doc} {place} {row}번째 줄"


def _raw_note(fact) -> str:
    """자료에 다른 단위로 적힌 값을 바꿨으면 원래 표기를 알려 줍니다 (예: 10,000 L → 10 m³)."""
    raw, value = str(fact.get("raw", "")).strip(), fact["value"]
    try:
        same = float(raw.replace(",", "")) == float(value)
    except (TypeError, ValueError):
        same = raw == str(value)
    if same or not raw:
        return ""
    try:
        raw = _number(float(raw.replace(",", "")))
    except ValueError:
        pass
    header_unit = re.search(r"\(([^)]+)\)\s*$", str(fact.get("header", "")))
    return f" (자료에는 ‘{raw}{' ' + header_unit.group(1) if header_unit else ''}’로 적혀 있음)"


class Names:
    """화면에 ‘가상 저장탱크 (TK-01)’처럼 이름과 번호를 함께 보여 줍니다."""

    def __init__(self, slots):
        self.names = {}
        for slot in slots:
            if slot["field"] == "name":
                fact = slot["chosen"] or (slot["facts"][0] if slot["facts"] else None)
                if fact and str(fact["value"]).strip():
                    self.names[slot["entity"]] = str(fact["value"]).strip()

    def __call__(self, entity: str) -> str:
        if entity == ew.SITE_ENTITY:
            return "사업장"
        if entity.startswith("detector:"):
            return "감지기 " + entity.split(":", 1)[1]
        if not entity.startswith(("facility:", "material:")):
            return entity
        code = entity.split(":", 1)[1]
        name = self.names.get(entity)
        return f"{name} ({code})" if name and name != code else code


# ---------------------------------------------------------------- 질문 모으기

def _collect(project, triage) -> dict:
    """사람에게 물어볼 것을 한곳에 모읍니다. 같은 값을 두 번 묻지 않습니다."""
    all_slots = ew.slots(project)
    queue = ew.review_queue(project)
    asked = {(i["entity"], i["field"]) for i in queue}
    rank_questions = []
    for question in triage["questions"]:
        field = question["field"]
        entity = question["material"] if field == "concentration" else question["facility"]
        if (entity, field) not in asked:
            asked.add((entity, field))
            rank_questions.append({**question, "entity": entity})
    present = {}
    for slot in all_slots:
        present.setdefault(slot["entity"], set()).add(slot["field"])
    missing = {}
    for entity in sorted(e for e in present if e.startswith("facility:")):
        fields = [f for f in MISSING_FIELDS if f not in present[entity] and (entity, f) not in asked]
        if fields:
            missing[entity] = fields
    notes = [t for t in ew.tasks(project) if not t["key"] and not t["task"].startswith("기본 운전조건 미확인")]
    return {"slots": all_slots, "names": Names(all_slots),
            "conflicts": [i for i in queue if i["kind"] == "충돌"],
            "problems": ew.problems(project),
            "held": [i for i in queue if i["kind"] == "확인 필요"],
            "rank": rank_questions, "missing": missing, "notes": notes}


def _question_count(found) -> int:
    return len(found["conflicts"]) + len(found["problems"]) + len(found["held"]) + len(found["rank"]) + len(found["missing"])


# ---------------------------------------------------------------- 공통

def _save(project, note: str | None = None) -> None:
    save_project(project)
    if note:
        st.session_state.setdefault("evidence_notes", []).append(note)
    st.rerun()


def _go(step: int, pid: str) -> None:
    st.session_state["evidence_step_" + pid] = step


def _next_button(step: int, pid: str, label: str) -> None:
    st.button(label, key=f"next_{step}_{pid}", type="primary", on_click=_go, args=(step, pid))


# ---------------------------------------------------------------- 머리

def _project_header():
    projects = list_projects()
    done = st.session_state.pop("evidence_reset_done", None)
    if done is not None:
        st.success(f"사업장 {done}곳을 모두 지웠어요. 작성한 엑셀을 올려 다시 시작하세요.")
    if not projects:
        from engine.stage2.project import Stage2Project
        return Stage2Project(project_id="EXCEL-START", company_name="", cap_required=True,
                             scope_confirmed=True, cap_selected=True)
    _reset_all(len(projects))
    options = {p["project_id"]: p["company_name"] for p in projects}
    current = st.session_state.get("_stage2_active_project_id")
    project_id = st.selectbox("사업장", list(options), index=list(options).index(current) if current in options else 0,
                              format_func=options.get)
    st.session_state["_stage2_active_project_id"] = project_id
    return load_project(project_id)


def _reset_all(count) -> None:
    """연습을 처음부터 다시 할 때: 사업장을 모두 지웁니다(승인한 규정 DB·법령 자료는 그대로)."""
    with st.expander("사업장 모두 지우고 처음부터 하기"):
        st.warning(f"저장된 사업장 {count}곳과 그 안의 올린 파일·답변·작성 내용·버전이 모두 지워지고 되돌릴 수 없어요. "
                   "승인해 둔 규정 DB와 내려받은 법령 자료는 지우지 않아요.")
        sure = st.checkbox("모두 지우는 것을 확인했어요", key="evidence_reset_all_confirm")
        if st.button("사업장 모두 지우기", disabled=not sure, key="evidence_reset_all"):
            removed = delete_all_projects()
            st.session_state.clear()
            st.session_state["evidence_reset_done"] = removed
            st.rerun()


def _status_line(project, found) -> None:
    if company_intake.current(project):
        st.markdown(f"최신 엑셀에서 값 {len(ew.graph_facts(project))}개를 읽었어요 · 입력 보완 {len(company_intake.corrections(project))}개")
        return
    info = ew.summary(project)
    if not info["documents"]:
        return
    count = len(_correction_rows(project, found))
    parts = [f"자료 {info['documents']}개를 읽었어요", f"값 {info['auto'] + info['person']}개를 채웠어요"]
    unread = sum(bool(d.get("needs_reading")) for d in ew.documents(project))
    parts.append(f"보완·확인할 항목 {count}개" if count else "자료에서 확인할 항목이 없어요")
    if unread:
        parts.append(f"읽지 못한 파일 {unread}개는 확인이 필요해요")
    st.markdown(" · ".join(parts))


# ---------------------------------------------------------------- ① 자료 올리기

def _step_upload(project, found) -> None:
    pid = project.project_id
    st.write("엑셀을 내려받아 연두색 입력 칸을 채운 뒤 올려 주세요.")
    st.download_button("회사 작성용 엑셀 내려받기 · 전체 항목", cap_upload_template.build_company(),
                       file_name=cap_upload_template.COMPANY_FILENAME,
                       mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                       key="cap_company_full_template")
    st.download_button("가상 회사 정보가 채워진 엑셀 내려받기", demo_company_workbook.build(),
                       file_name=demo_company_workbook.FILENAME,
                       mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                       key="cap_company_demo_template")
    st.caption("가상 회사 엑셀은 값을 미리 채운 1군 작성 연습용입니다. 내려받은 파일을 그대로 올려 초안을 확인하세요. "
               "실제 법적 판정·현장 확인 자료가 아니며 관련 보완은 남습니다.")
    st.caption("단답형은 ‘사업장 정보’ 시트의 값 열에, 목록은 한 행에 하나씩 적으세요. 예시는 읽지 않습니다. "
               "수정한 엑셀 전체를 다시 올리면 이름이 달라도 최신 내용으로 교체합니다. 도면·SDS·스캔본은 올리지 않습니다.")
    st.caption("엑셀의 사업장명으로 자료를 구분합니다. 처음 올리는 사업장은 파일을 읽은 뒤 자동으로 등록합니다.")
    if project.company_name:
        _latest_workbook(project)
    uploaded = st.file_uploader("작성한 회사 엑셀 선택 (.xlsx, 한 파일)", type=["xlsx"],
                                accept_multiple_files=False, key="company_excel_" + pid)
    if st.button("엑셀 읽고 확인하기", type="primary", disabled=uploaded is None, key="read_" + pid):
        progress = st.progress(0, text="엑셀 양식을 확인하고 있습니다.")
        try:
            candidate = company_intake.project_for_upload(project, uploaded.name, uploaded.getvalue())
            progress.progress(0.35, text="사업장과 입력값을 확인했습니다. 내용을 반영하고 있습니다.")
            changed = company_intake.apply(candidate, uploaded.name, uploaded.getvalue())
            progress.progress(0.75, text="읽은 내용을 저장하고 있습니다.")
            if changed:
                save_attachment(candidate.project_id, uploaded.name, uploaded.getvalue(), note="최신 회사 작성용 엑셀 원본")
                save_project(candidate)
                st.session_state.pop("narrative_progress_" + pid, None)
                st.session_state.pop("evidence_change", None)
            st.session_state["_stage2_active_project_id"] = candidate.project_id
            progress.progress(1.0, text="엑셀 읽기 완료 · ②에서 입력 결과를 확인하세요.")
            st.session_state["company_upload_result_" + candidate.project_id] = "엑셀 읽기 완료 · ②에서 입력 결과를 확인하세요."
            st.session_state.pop("company_upload_error_" + pid, None)
            st.session_state.setdefault("evidence_notes", []).append(
                f"{uploaded.name}: 최신 엑셀로 반영했어요. ②에서 빠진 값과 고칠 칸을 확인하세요." if changed else
                "이미 반영한 엑셀과 내용이 같습니다. ②에서 입력 결과를 확인하세요.")
            st.session_state["evidence_goto"] = 2
            st.rerun()
        except company_intake.WorkbookFormatError as exc:
            st.session_state["company_upload_error_" + pid] = exc.issues
        except Exception:
            st.session_state["company_upload_error_" + pid] = [{"파일·위치": uploaded.name,
                "항목": "엑셀 읽기", "확인된 문제": "엑셀을 반영하지 못했습니다. 이전에 반영한 자료는 유지됩니다.",
                "수정 요청": "회사 작성용 양식인지 확인하고 .xlsx로 저장한 뒤 다시 올려 주세요."}]
    errors = st.session_state.get("company_upload_error_" + pid)
    if errors:
        st.error("이번 파일은 반영하지 못했습니다. 아래 내용을 고쳐 다시 올려 주세요.")
        st.dataframe(_excel_issue_rows(errors, uploaded.getvalue() if uploaded else None), hide_index=True, width="stretch")
    if company_intake.current(project):
        _next_button(2, pid, "다음: 입력 확인·수정 →")


def _latest_workbook(project):
    state = company_intake.current(project)
    if not state:
        st.info("아직 반영한 회사 작성용 엑셀이 없습니다. 기존 자료는 보존되며, 엑셀을 올리면 새 입력 내용으로 전환됩니다.")
        return
    from datetime import datetime
    from zoneinfo import ZoneInfo
    when = datetime.fromisoformat(state['uploaded_at']).astimezone(ZoneInfo('Asia/Seoul')).strftime('%Y-%m-%d %H:%M:%S')
    st.caption(f"현재 반영한 최신 엑셀: {state['name']} · {when} (한국 시간)")


_STEP3_PREFIXES = ('cap.offsite.', 'cap.site.surrounding_environment', 'documents.')


def _validation_location(item) -> dict:
    """문서 검증 결과는 엑셀 칸이 아니라 점검 항목 단위라서, 연결된 입력 위치(엑셀 시트·③ 화면)를 알려 준다."""
    from engine.stage2 import cap_fact_catalog as cat, cap_upload_template as tpl
    if '_keys' not in item:
        return {'시트': '파일 전체', '행': '—', '열': '—'}
    keys = set(item.get('_keys') or ())
    if keys and all(k.startswith(_STEP3_PREFIXES) for k in keys):
        return {'시트': '엑셀 입력 아님 · ③ 시나리오·주변 조사 화면', '행': '—', '열': '—'}
    sheets, headers = [], []
    text = f"{item.get('항목', '')} {item.get('확인된 문제', '')}"
    for fact in cat.facts(docs=['CAP']):
        if not fact.readable or not (fact.field_key in keys or any(t[0] in keys for t in fact.targets)):
            continue
        if fact.kind == 'psm_table':
            sheet = tpl.sheet_title(fact.table)
        elif fact.kind in cat.SITE_KINDS:
            sheet = tpl.SITE_SHEET
        else:
            sheet = tpl.ENTITY_SHEETS.get('scenario' if fact.file == tpl.SCENARIO_FILE else fact.kind)
        if sheet and sheet not in (tpl.ENTITY_SHEETS['scenario'], tpl.ENTITY_SHEETS['target']):
            if sheet not in sheets:
                sheets.append(sheet)
            if fact.template_header and fact.template_header in text and fact.template_header not in headers:
                headers.append(fact.template_header)
    if sheets:
        return {'시트': ' · '.join(sheets) + ' (관련 시트)', '행': '—', '열': ' · '.join(headers[:3]) or '—'}
    return {'시트': '엑셀 입력 아님 · 법령 승인 DB·판정 확인', '행': '—', '열': '—'}


def _excel_issue_rows(rows, content=None):
    """파일명 대신 실제 엑셀의 시트·행·열 제목을 표시한다. 특정 불가한 검증은 추측하지 않는다."""
    from io import BytesIO
    from openpyxl import load_workbook
    from openpyxl.utils import get_column_letter
    book = None
    if content:
        try:
            book = load_workbook(BytesIO(content), read_only=True)
        except Exception:
            pass
    result = []
    for item in rows:
        source = item.get('파일·위치', '')
        match = re.search(r'·\s*(.+?)(?:!([A-Z]+)(\d+)|\s+행\s+(\d+)\s+열\s+(\d+))\b', source)
        if match:
            sheet = match.group(1).strip()
            row = int(match.group(3) or match.group(4))
            col = match.group(2) or get_column_letter(int(match.group(5)))
            title = ''
            if book and sheet in book.sheetnames:
                title = str(book[sheet][f'{col}1'].value or '')
            location = {'시트': sheet, '행': f'{row}행', '열': f'{col}열' + (f' · {title}' if title else '')}
        else:
            location = _validation_location(item)
        result.append({**location, '항목': item.get('항목', ''), '확인된 문제': item.get('확인된 문제', '')})
    return result


# ---------------------------------------------------------------- ② 파일별 보완 요청

def _correction_rows(project, found) -> list[dict]:
    """파일·위치·항목·수정 지시를 한 줄씩 적는다. 추정된 원본은 쓰지 않는다."""
    if company_intake.current(project):
        return company_intake.corrections(project)
    rows = []

    def add(source, item, detail, action):
        rows.append({"파일·위치": source, "항목": item, "확인된 문제": detail, "수정 요청": action})

    for doc in ew.documents(project):
        if doc.get("needs_reading"):
            is_reference = doc.get("kind") == "reference" and not doc.get("needs_ocr")
            add(doc["name"], "서식 값 연결" if is_reference else "문서 읽기",
                doc.get("problem") or "값을 연결하지 못했습니다.",
                "서식 항목 이름(예: 대표자, 사업자 등록번호)과 값을 2열 표로 적어 다시 올리세요." if is_reference else
                "스캔본은 OCR을 적용하세요. 항목명·번호를 확인해 원본을 다시 제출하세요.")
        for link in doc.get("links", []):
            record = project.get_field(link["key"])
            if record and (record.status == "HOLD" or
                           (not cap_evidence_links.linked(record) and str(record.value) != str(link["value"]))):
                add(f"{doc['name']} · {link['location']}", link["label"],
                    "같은 항목의 값이 다른 자료와 다릅니다.",
                    "맞는 값을 원본에서 확인하고 잘못된 파일을 고쳐 다시 올려 주세요.")
        for table in doc.get("tables", []):
            for skipped in table.get("skipped", []):
                add(f"{doc['name']} · {table['location']} {skipped['row']}행", "설비·제품 번호",
                    skipped["reason"], "설비번호 또는 제품코드를 원본 행에 기재한 뒤 다시 올려 주세요.")
        for column in doc.get("unmapped_columns", []):
            add(doc["name"], column + " 열", "알 수 없는 항목 이름", "열 이름과 뜻·단위를 확인해 원본을 수정한 뒤 다시 올려 주세요.")
    for item in found["conflicts"]:
        facts = item["slot"]["facts"]
        add(" / ".join(dict.fromkeys(_where(f) for f in facts)),
            f"{found['names'](item['entity'])} · {_name(item['field'])}",
            "자료마다 값이 다릅니다: " + ", ".join(_value(item["field"], f["value"]) for f in facts[:5]),
            "현재 적용되는 개정본을 확인하고 파일의 값을 바로잡아 다시 올려 주세요. 확인 전에는 자동 확정되지 않습니다.")
    for group in found["problems"]:
        fact = group["facts"][0]
        add(_where(fact), _name(group["field"]), f"‘{group['raw']}’: {group['reason']}",
            "숫자·단위·표기와 적용 대상을 원본에서 수정한 뒤 다시 올려 주세요.")
    for entity, fields in found["missing"].items():
        add("설비대장·운전조건표 (해당 파일 위치 미확인)", found["names"](entity),
            "자료에서 " + ", ".join(_name(field) for field in fields) + " 항목을 찾지 못했습니다.",
            "해당 설비번호의 누락 열을 원본에 기재해 올려 주세요.")
    for question in found["rank"]:
        add("설비대장·물질목록 (해당 파일 위치 미확인)", found["names"](question["entity"]),
            f"시나리오 우선순위에 영향을 줄 {_name(question['field'])} 값을 확인하지 못했습니다.",
            "해당 물질·설비의 값을 단위와 함께 원본에 기재해 올려 주세요.")
    for item in found["held"]:
        fact = item["slot"]["facts"][0]
        add(_where(fact), _name(item["field"]), "출처·단위·법적 판단용 값의 확인이 필요합니다.",
            "원본과 현장 적용 여부를 확인하고 최신본을 다시 올리거나 아래 확인 절차를 이용하세요.")
    narrative_keys = cap_fact_drafting.input_requirement_keys(project.cap_group)
    for req in build_intake_catalog(project):
        if req.requirement_key in narrative_keys:
            continue  # 문장 제출을 요구하지 않음: 단답형 누락·충돌은 사실 표에서 안내
        if req.system not in {"CAP", "COMMON"} or req.coverage_status in {"확인", "해당 없음"}:
            continue
        missing = ", ".join(req.missing_labels or (req.label,))
        add(" / ".join(req.suggested_evidence) or "근거 파일 미지정", req.label,
            f"작성 항목에 {missing} 값이 아직 연결되지 않았습니다. 업로드한 자료에 있어도 자동 연결되지 않았을 수 있습니다.",
            "해당 근거자료와 최신본·적용범위를 확인해 파일에 적어 다시 올려 주세요.")
    return rows


KIND_NAMES = {"material": "물질 목록·SDS", "facility": "설비 목록·운전 자료", "site": "사업장 일반정보·공정 설명",
              "detector": "가스감지기 명세", "target": "주변 보호대상 조사표",
              "mixed": "여러 종류가 든 묶음 파일(시트마다 따로 읽음)"}
DOC_NAMES = {"CAP": "화학사고예방관리계획서", "PSM": "공정안전보고서"}


def _coverage_panel(project, found) -> None:
    """보완 요청을 필요한 사실 표 한 곳에 모읍니다."""
    if company_intake.current(project):
        _latest_workbook(project)
        done = st.session_state.pop("company_upload_result_" + project.project_id, None)
        if done:
            st.progress(1.0, text=done)
        rows = company_intake.corrections(project)
        if project.cap_in_scope:
            from engine.stage2 import report_export
            try:
                state = report_export.evaluate(project, "CAP")
                for issue in (*state.text_issues, *state.manual_issues):
                    if issue.status != "PASS":
                        rows.append({"파일·위치": company_intake.current(project)['name'], "항목": issue.legal_item,
                                     "확인된 문제": issue.message, "_keys": tuple(issue.field_keys),
                                     "수정 요청": "관련 입력값과 적용 기준을 확인해 주세요. 도면·SDS 원본은 올리지 않습니다."})
            except Exception:
                st.error("문서 검증을 완료하지 못했습니다. 입력 검사 통과만으로 문서가 완성된 것은 아닙니다.")
        st.caption("시트·행·열을 찾아 엑셀을 고친 뒤 전체 파일을 다시 올려 주세요. 문서 검증 결과는 칸 단위가 아니라 관련 시트 또는 확인 화면을 표시합니다.")
        if rows:
            from engine.stage2.storage import load_attachment
            current = company_intake.current(project)
            content = load_attachment(project.project_id, current['sha256'])
            st.dataframe(_excel_issue_rows(rows, content), hide_index=True, width="stretch")
        else:
            st.success("엑셀 입력 검사에서 고칠 항목이 없습니다. 법령 기준과 문서 작성 가능 여부는 별도로 검증합니다.")
        st.button("수정한 엑셀 다시 올리기", key="replace_excel_" + project.project_id,
                  on_click=_go, args=(1, project.project_id), type="primary")
        return

def _step_questions(project, found) -> None:
    pid = project.project_id
    if not company_intake.current(project):
        st.info("먼저 ① 엑셀 작성·올리기에서 회사 작성용 엑셀을 올려 주세요.")
        return
    _coverage_panel(project, found)
    _next_button(3, pid, "다음: 시나리오 후보 보기 →")


# ---------------------------------------------------------------- ③ 시나리오 후보

def _step_scenarios(project, found, triage) -> None:
    from engine.stage2 import scenario_survey, scenario_pipeline
    from ui.scenario_results_page import _boundary

    pid, names = project.project_id, found["names"]
    st.caption("지도에서 사업장 경계를 확인하면 주변 지역을 검색하고 사고시나리오를 다시 계산합니다. 검색 결과는 확인 전 후보입니다.")
    _boundary(project)
    candidates = scenario_survey.candidate_rows(project)
    if candidates:
        st.write(f"경계 밖 500m 안에서 찾은 주변 보호대상 후보 {len(candidates)}곳")
        st.dataframe(candidates, hide_index=True, width="stretch")
    linked = [r for r in triage["rows"] if r["linked"]]
    existing = scenario_pipeline.file_inputs(project)
    tags = list(dict.fromkeys(r['facility'].split(':', 1)[-1] for r in linked))
    if tags:
        st.caption("설비별 사고유형은 비워 두면 규칙으로 제안합니다. 설비 위치 좌표가 없으면 경계까지 거리는 지도만으로 알 수 없어 회사에서 확인한 값을 적어야 합니다.")
        review = st.data_editor([
            {'설비번호': tag, '선정할 사고유형': existing.get(tag, {}).get('선정할 사고유형', ''),
             '경계까지 거리(m)': existing.get(tag, {}).get('경계까지 거리(m)', ''),
             '액위(m)': existing.get(tag, {}).get('액위(m)', '')}
            for tag in tags], hide_index=True, num_rows='fixed', disabled=['설비번호'],
            key='scenario_review_' + pid)
        if st.button('시나리오 확인값 반영', key='scenario_review_save_' + pid):
            values = review.to_dict('records') if hasattr(review, 'to_dict') else review
            project.set_field(scenario_pipeline.INPUT_KEY, '설비별 시나리오 확인값', values, 'USER_CONFIRMED')
            scenario_pipeline.refresh(project)
            save_project(project)
            st.rerun()
    st.dataframe([{"순서": r["rank"], "설비": names(r["facility"]), "물질": names(r["material"]),
                   "인식한 파일": " / ".join(dict.fromkeys(source["document"] for source in r["sources"].values() if source)),
                   "비교 점수": round(r["score"], 1)} for r in linked], hide_index=True, width="stretch")
    _next_button(4, pid, "다음: 작성·내려받기 →")


# ---------------------------------------------------------------- ④ 작성·내려받기


def _step_export(project, found) -> None:
    """확인한 값을 작성자료로 보내고, 문서별로 남은 일을 한곳에서 보여 줍니다(여기만 따라가면 되도록)."""
    st.caption("읽은 자료로 초안을 작성합니다. 부족한 자료는 ②에서 확인하세요.")
    if st.button("작성하기", type="primary", key="write_" + project.project_id):
        _auto_publish(project)
        st.success("초안을 작성했어요. 아래에서 내려받으세요.")
    st.caption("서술형 초안 " + str(sum(bool(i.company_text or i.draft) for system in cap_fact_drafting.systems(project) for i in cap_fact_drafting.items(project, system))) + "개 작성됨")
    _narrative_drafts(project)
    from ui.report_export_panel import render_workspace_download
    render_workspace_download(project)


def _auto_publish(project) -> None:
    """회사 파일에서 읽은 값을 작성자료(초안)에 자동으로 반영합니다. 사람의 승인 단계는 없습니다.
    서술형 항목도 근거 사실이 있으면 규칙 초안을 바로 만들어 둡니다(AI 다듬기는 ④에서 따로)."""
    changed = False
    # 파일을 새로 올리면 이전에 반영한 값이 보류로 바뀌므로, 값이 같아도 확정 상태가 아니면 다시 반영합니다.
    if ew.needs_publish(project):
        try:
            ew.publish(project, replace_existing=True)
            changed = True
        except ValueError:
            pass
    if project.cap_in_scope or project.psm_in_scope:
        try:
            changed = cap_fact_drafting.refresh_rule_drafts(project) > 0 or changed
        except Exception:          # 서술형 초안 문제로 화면이 멈추지 않게
            pass
        try:                       # 사고시나리오 결과도 새 파일로 다시 계산(‘사고시나리오 결과’ 화면에서 확인)
            from engine.stage2 import scenario_pipeline

            if scenario_pipeline.needs_refresh(project):
                scenario_pipeline.refresh(project)
                changed = True
        except Exception:
            pass
    if changed:
        save_project(project)


def _narrative_drafts(project) -> None:
    """사전관리방침·비상대응계획 서술형: 회사 파일의 사실로 쓴 초안 현황과 다시 만들기."""
    systems = cap_fact_drafting.systems(project)
    if not systems:
        return
    pid = project.project_id
    tables = [(system, cap_fact_drafting.rows(project, system)) for system in systems]
    if not any(rows for _, rows in tables):
        return
    use_ai = st.checkbox("AI를 이용해서 작성하기 (서술형만)", key="fact_ai_" + pid)
    client = None
    if use_ai:
        from ui.narrative_panel import _local_config

        config, probe, build = _local_config("fact_" + pid)
        if probe is not None and probe.ready:
            client = build(config)
        else:
            st.warning("로컬 AI 연결을 확인해 주세요.")
    progress_key = "narrative_progress_" + pid
    previous = st.session_state.get(progress_key)
    if previous and previous[1]:
        st.progress(previous[0] / previous[1], text=f"서술형 {previous[0]}/{previous[1]}개 완료 (100%)")
    if st.button("AI로 서술형 작성하기" if use_ai else "서술형 작성하기", key="fact_drafts_" + pid, disabled=use_ai and client is None):
        bar = st.progress(0, text="작성할 서술형 항목을 확인하고 있어요.")

        def report_progress(done, total, label):
            percent = round(done / total * 100) if total else 0
            text = f"서술형 {done}/{total}개 완료 ({percent}%)"
            if label:
                text += f" · {done + 1}번째 작성 중: {label}"
            bar.progress(done / total if total else 0, text=text)

        results = cap_fact_drafting.generate(project, client, refresh=True, on_progress=report_progress)
        st.session_state[progress_key] = (len(results), len(results))
        save_project(project)
        checks = sum(r["checks"] for r in results)
        st.session_state.setdefault("evidence_notes", []).append(
            (f"서술형 초안 {len(results)}개를 만들었어요." + (f" ‘[확인 필요]’ {checks}곳을 채워 주세요." if checks else ""))
            if results else "작성할 서술형이 없어요. 완성 문서는 그대로 사용하며, 근거가 없는 항목은 ②에서 필요한 자료를 확인해 주세요.")
        st.rerun()


# ---------------------------------------------------------------- 화면

st.set_page_config(page_title="사업장 업무", page_icon="📂", layout="wide")
st.title("화학사고예방관리계획서 자료 정리")
st.caption("회사 작성용 엑셀을 내려받아 작성하고 올리세요. 빠진 값과 잘못된 칸을 확인해 수정본을 다시 올리면 최신 내용으로 반영합니다.")
_active = _project_header()
if not _active.company_name:
    _step_upload(_active, None)
    st.stop()
_triage = scenario_triage.evaluate(_active)
_found = _collect(_active, _triage)
_auto_publish(_active)
_status_line(_active, _found)
_change = st.session_state.pop("evidence_change", None)
if _change:
    (st.warning if "새 충돌 0개" not in _change and "새 충돌은 없어요" not in _change else st.info)(_change)
for _note in st.session_state.pop("evidence_notes", []):
    (st.error if "읽지 못" in _note else st.warning if "찾지 못" in _note else st.success)(_note)

_step_key = "evidence_step_" + _active.project_id
if "evidence_goto" in st.session_state:
    st.session_state[_step_key] = st.session_state.pop("evidence_goto")
if _step_key not in st.session_state:
    st.session_state[_step_key] = 1 if not company_intake.current(_active) else (2 if _correction_rows(_active, _found) else 3)
# 단계 이름은 고정합니다: 이름이 바뀌면 위젯이 새로 만들어져 보던 단계가 처음으로 돌아갈 수 있습니다.
_step = st.radio("단계", list(STEPS), key=_step_key, horizontal=True, label_visibility="collapsed", format_func=STEPS.get)
st.divider()
if _step == 1:
    _step_upload(_active, _found)
elif _step == 2:
    _step_questions(_active, _found)
elif _step == 3:
    _step_scenarios(_active, _found, _triage)
else:
    _step_export(_active, _found)
