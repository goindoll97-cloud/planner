from __future__ import annotations

from io import BytesIO
import unittest

from docx import Document
from openpyxl import load_workbook

from engine.stage2.integrated_workbook import apply_integrated_authoring_workbook
from engine.stage2.project import Stage2Project
from engine.stage2.statutory_report_v2 import _add_operator_staffing
from engine.stage2.workbook_enhancements import build_enhanced_integrated_authoring_workbook


class CAPOperatorStaffingTableTests(unittest.TestCase):
    def _project(self) -> Stage2Project:
        project = Stage2Project(
            project_id="S2-CAP-STAFFING",
            company_name="테스트화학",
            site_name="제1공장",
            cap_required=True,
            cap_group="1군",
            scope_confirmed=True,
            cap_selected=True,
        )
        project.set_field(
            "business.company_name",
            "회사명",
            "테스트화학",
            "USER_CONFIRMED",
        )
        project.set_field(
            "business.address",
            "사업장 소재지",
            "테스트시 테스트로 1",
            "USER_CONFIRMED",
        )
        return project

    @staticmethod
    def _save(wb) -> bytes:
        out = BytesIO()
        wb.save(out)
        return out.getvalue()

    def test_workbook_exposes_operator_staffing_as_structured_table(self):
        project = self._project()
        data = build_enhanced_integrated_authoring_workbook(project, example=False)
        wb = load_workbook(BytesIO(data), data_only=False)

        self.assertIn("30_운전책임자_작업자현황", wb.sheetnames)
        ws = wb["30_운전책임자_작업자현황"]
        headers = [str(cell.value or "").strip() for cell in ws[4]][:5]
        self.assertEqual(
            headers,
            ["공정·단위공장", "운전책임자", "작업자 수", "교대 형태", "비고"],
        )

        meta = wb["_시스템정보"]
        matches = [
            row
            for row in meta.iter_rows(min_row=8, values_only=True)
            if str(row[2] or "") == "cap.facility.operator_staffing"
        ]
        self.assertEqual(len(matches), 1)
        self.assertEqual(str(matches[0][0]), "TABLE")

    def test_operator_staffing_round_trip_preserves_rows(self):
        project = self._project()
        data = build_enhanced_integrated_authoring_workbook(project, example=False)
        wb = load_workbook(BytesIO(data), data_only=False)
        ws = wb["30_운전책임자_작업자현황"]

        values = [
            "제1공장 염소 저장·공급공정",
            "환경안전팀 홍길동",
            4,
            "2조 2교대",
            "공정별 실제 운영인원 기준",
        ]
        for col, value in enumerate(values, start=1):
            ws.cell(5, col, value)

        result = apply_integrated_authoring_workbook(project, self._save(wb))
        self.assertGreaterEqual(result.table_fields, 1)

        record = project.get_field("cap.facility.operator_staffing")
        self.assertIsNotNone(record)
        self.assertEqual(record.status, "USER_CONFIRMED")
        self.assertIsInstance(record.value, list)
        self.assertEqual(record.value[0]["운전책임자"], "환경안전팀 홍길동")
        self.assertEqual(record.value[0]["작업자 수"], 4)
        self.assertEqual(record.value[0]["교대 형태"], "2조 2교대")

        redownload = build_enhanced_integrated_authoring_workbook(project, example=False)
        wb2 = load_workbook(BytesIO(redownload), data_only=False)
        ws2 = wb2["30_운전책임자_작업자현황"]
        self.assertEqual(ws2.cell(5, 1).value, "제1공장 염소 저장·공급공정")
        self.assertEqual(ws2.cell(5, 3).value, 4)
        self.assertEqual(ws2.cell(5, 4).value, "2조 2교대")

    def test_docx_renderer_outputs_structured_staffing_table(self):
        project = self._project()
        project.set_field(
            "cap.facility.operator_staffing",
            "운전책임자 및 작업자 현황",
            [
                {
                    "공정·단위공장": "제1공장 저장공정",
                    "운전책임자": "홍길동",
                    "작업자 수": 3,
                    "교대 형태": "주간",
                    "비고": "저장공정",
                },
                {
                    "공정·단위공장": "제1공장 공급공정",
                    "운전책임자": "김담당",
                    "작업자 수": 5,
                    "교대 형태": "2조 2교대",
                    "비고": "",
                },
            ],
            "USER_CONFIRMED",
        )

        doc = Document()
        _add_operator_staffing(doc, project)

        self.assertEqual(len(doc.tables), 1)
        table = doc.tables[0]
        header = [cell.text for cell in table.rows[0].cells]
        self.assertEqual(
            header,
            ["공정·단위공장", "운전책임자", "작업자 수", "교대 형태", "비고"],
        )
        text = "\n".join(cell.text for row in table.rows for cell in row.cells)
        self.assertIn("제1공장 저장공정", text)
        self.assertIn("홍길동", text)
        self.assertIn("2조 2교대", text)

    def test_legacy_scalar_staffing_value_still_renders_as_text(self):
        project = self._project()
        project.set_field(
            "cap.facility.operator_staffing",
            "운전책임자 및 작업자 현황",
            "저장공정 운전책임자 1명, 작업자 4명, 2조 2교대",
            "USER_CONFIRMED",
        )

        doc = Document()
        _add_operator_staffing(doc, project)

        self.assertEqual(len(doc.tables), 0)
        text = "\n".join(paragraph.text for paragraph in doc.paragraphs)
        self.assertIn("작업자 4명", text)
        self.assertIn("2조 2교대", text)


if __name__ == "__main__":
    unittest.main()
