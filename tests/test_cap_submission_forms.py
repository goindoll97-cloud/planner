from __future__ import annotations

from io import BytesIO
from pathlib import Path
import tempfile
import unittest

from docx import Document

from engine.stage2.project import Stage2Project
from engine.stage2 import cap_submission_forms as forms
from engine.stage2 import versioning


class CAPSubmissionFormsTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.root = Path(self.tmp.name)
        self.addCleanup(self.tmp.cleanup)

    def _project(self) -> Stage2Project:
        p = Stage2Project(
            project_id="CAP-ADMIN",
            company_name="한빛화학(주)",
            site_name="한빛화학 울산공장",
            cap_required=True,
            cap_group="1군",
            scope_confirmed=True,
            cap_selected=True,
        )
        facts = (
            ("business.registration_no", "사업자등록번호", "123-45-67890"),
            ("business.representative", "대표자", "홍길동"),
            ("business.address", "주소", "울산광역시 남구 산업로 1"),
            ("cap.business.writer_name", "담당자", "김안전"),
            ("cap.business.writer_contact", "연락처", "010-1234-5678"),
            ("cap.business.writer_email", "이메일", "safety@example.com"),
            ("cap.business.joint_emergency_plan", "공동비상대응계획", "N"),
            ("cap.submission.method", "제출방법", "화학물질 종합정보시스템"),
            ("cap.submission.business_permit", "영업허가 대상", "대상"),
        )
        for key, label, value in facts:
            p.set_field(key, label, value, "USER_CONFIRMED")
        return p

    def test_form31_requires_submission_stage_facts_and_does_not_invent_date(self):
        p = self._project()
        p.set_field("cap.business.submission_type", "제출구분", "신규제출", "USER_CONFIRMED")

        result = forms.form31_readiness(p)

        self.assertFalse(result.ready)
        self.assertTrue(any("관할 유역" in blocker for blocker in result.blockers))
        self.assertTrue(any("신청일" in blocker for blocker in result.blockers))

    def test_form31_docx_for_new_submission(self):
        p = self._project()
        for key, label, value in (
            ("cap.business.submission_type", "제출구분", "신규제출"),
            ("cap.submission.office", "관할기관", "낙동강"),
            ("cap.submission.center", "합동방재센터", "울산"),
            ("cap.submission.review_skip", "검토생략", "공정안전보고서"),
            ("cap.submission.form31.application_date", "신청일", "2026-09-21"),
        ):
            p.set_field(key, label, value, "USER_CONFIRMED")

        data = forms.build_form31_docx(p)
        doc = Document(BytesIO(data))
        text = "\n".join(
            [para.text for para in doc.paragraphs]
            + [cell.text for table in doc.tables for row in table.rows for cell in row.cells]
        )

        self.assertIn("화학사고예방관리계획서 검토신청서", text)
        self.assertIn("[✓] 신규제출", text)
        self.assertIn("공정안전보고서", text)
        self.assertIn("2026-09-21", text)
        self.assertEqual(data, forms.build_form31_docx(p))

    def test_form32_requires_change_submission_and_base_version(self):
        p = self._project()
        p.set_field("cap.business.submission_type", "제출구분", "변경제출", "USER_CONFIRMED")

        result = forms.form32_readiness(p)

        self.assertFalse(result.ready)
        self.assertTrue(any("최종" in blocker for blocker in result.blockers))
        self.assertTrue(any("비교 기준 버전" in blocker for blocker in result.blockers))

    def test_form32_uses_version_diff_for_back_page(self):
        p = self._project()
        p.set_field("cap.business.submission_type", "제출구분", "신규제출", "USER_CONFIRMED")
        p.set_field(
            "inventory.chemicals",
            "화학물질",
            [{"물질명": "톨루엔", "CAS 번호": "108-88-3"}],
            "USER_CONFIRMED",
        )
        p.set_field(
            "cap.workspace.facilities",
            "시설",
            [{"설비번호": "TK-101", "설비명": "저장탱크", "최대보유량(kg)": 10000}],
            "USER_CONFIRMED",
        )
        versioning.freeze_version(p, "CAP", "신규제출", root=self.root)

        # The production engine uses the normal runtime root. Mirror the snapshot
        # metadata by monkeypatching only the version loader/diff entry points.
        original_load = versioning.load_version_fields
        original_diff = versioning.diff_versions
        try:
            versioning.load_version_fields = lambda project_id, version_id, root=versioning.DEFAULT_ROOT: original_load(project_id, version_id, self.root)
            versioning.diff_versions = lambda project, base_version_id, target_version_id=None, root=versioning.DEFAULT_ROOT: original_diff(project, base_version_id, target_version_id, self.root)

            p.set_field("cap.business.submission_type", "제출구분", "변경제출", "USER_CONFIRMED")
            p.set_field(
                "inventory.chemicals",
                "화학물질",
                [
                    {"물질명": "톨루엔", "CAS 번호": "108-88-3"},
                    {"물질명": "염소", "CAS 번호": "7782-50-5"},
                ],
                "USER_CONFIRMED",
            )
            p.set_field(
                "cap.workspace.facilities",
                "시설",
                [
                    {"설비번호": "TK-101", "설비명": "저장탱크", "최대보유량(kg)": 20000},
                    {"설비번호": "V-201", "설비명": "염소용기", "최대보유량(kg)": 1000},
                ],
                "USER_CONFIRMED",
            )
            for key, label, value in (
                ("cap.submission.final_approval_no", "최종 적합통보 번호", "A-2026-001"),
                ("cap.submission.final_approval_date", "최종 적합통보 일자", "2026-03-31"),
                ("cap.submission.change_reason_type", "변경사유", "그 밖의 사유"),
                ("cap.submission.change_other_reason", "사유", "염소 취급 및 설비 추가"),
                ("cap.submission.form32.application_date", "신청일", "2026-09-21"),
                ("cap.submission.base_version_id", "기준버전", "CAP-v1.0"),
            ):
                p.set_field(key, label, value, "USER_CONFIRMED")
            p.set_field(
                "cap.submission.oca_details", "장외영향평가",
                {"민원번호": "해당없음", "결과번호": "해당없음", "적합날짜": "해당없음", "위험도": "해당없음"},
                "USER_CONFIRMED",
            )
            p.set_field(
                "cap.submission.rmp_details", "위해관리계획",
                {"민원번호": "해당없음", "결과번호": "해당없음", "적합날짜": "해당없음"},
                "USER_CONFIRMED",
            )

            result = forms.form32_readiness(p)
            self.assertTrue(result.ready, result.blockers)
            details = result.values["변경사항 상세내용"]
            self.assertIn("염소", details["유해화학물질추가"]["변경 후"])
            self.assertIn("V-201", details["시설추가"]["변경 후"])
            self.assertIn("20 ton", details["취급저장량 증가"]["변경 후"])

            data = forms.build_form32_docx(p)
            doc = Document(BytesIO(data))
            text = "\n".join(
                [para.text for para in doc.paragraphs]
                + [cell.text for table in doc.tables for row in table.rows for cell in row.cells]
            )
            self.assertIn("화학사고예방관리계획서 변경 검토신청서", text)
            self.assertIn("염소", text)
            self.assertIn("V-201", text)
            self.assertIn("20 ton", text)
            self.assertEqual(data, forms.build_form32_docx(p))
        finally:
            versioning.load_version_fields = original_load
            versioning.diff_versions = original_diff


if __name__ == "__main__":
    unittest.main()
