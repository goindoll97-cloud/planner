from __future__ import annotations

from io import BytesIO
from pathlib import Path
import unittest

from openpyxl import Workbook, load_workbook

from engine.stage2.direct_template import create_direct_entry_template_project
from engine.stage2.integrated_workbook import build_integrated_authoring_workbook
from engine.stage2.intake import selected_requirement_specs
from engine.stage2.project import Stage2Project
from engine.stage2.standalone_entry import (
    create_project_from_standalone_workbook,
    inspect_standalone_workbook,
    is_standalone_stage2_project,
)
from engine.stage2.workbook_enhancements import build_enhanced_integrated_authoring_workbook


PROJECT_ROOT = Path(__file__).resolve().parents[1]


class Stage2StandaloneEntryTests(unittest.TestCase):
    def _completed_workbook(self) -> bytes:
        project = Stage2Project(
            project_id="S2-DIRECT-TEST",
            company_name="직접시작화학",
            psm_required=True,
            cap_required=True,
            cap_group="2군",
            scope_confirmed=True,
            psm_selected=True,
            cap_selected=True,
        )
        project.set_field("business.company_name", "회사명", "직접시작화학", "VERIFIED")
        project.set_field("business.address", "사업장 소재지", "테스트시 직접로 1", "VERIFIED")
        project.set_field("cap.business.writing_level", "작성수준", "2군 사업장", "USER_CONFIRMED")
        project.set_field(
            "inventory.chemicals",
            "화학물질 목록",
            [{"물질명": "톨루엔", "CAS 번호": "108-88-3", "함량(%)": 99.5, "최대보유량": 15000, "단위": "kg"}],
            "VERIFIED",
        )
        project.set_field(
            "inventory.facilities",
            "시설별 최대보유량 자료",
            [{"설비번호": "TK-101", "설비명": "톨루엔 저장탱크", "설비종류": "저장탱크", "최대보유량": 15000}],
            "VERIFIED",
        )
        return build_integrated_authoring_workbook(project)

    def test_preview_reads_scope_company_and_group_without_stage1(self):
        preview = inspect_standalone_workbook(self._completed_workbook())
        self.assertEqual(preview.project_id, "S2-DIRECT-TEST")
        self.assertEqual(preview.company_name, "직접시작화학")
        self.assertEqual(preview.address, "테스트시 직접로 1")
        self.assertTrue(preview.psm_selected)
        self.assertTrue(preview.cap_selected)
        self.assertEqual(preview.cap_group, "2군")

    def test_direct_project_runs_stage2_scope_without_claiming_stage1_decision(self):
        project, result, preview = create_project_from_standalone_workbook(self._completed_workbook())
        self.assertEqual(preview.cap_group, "2군")
        self.assertTrue(is_standalone_stage2_project(project))
        self.assertTrue(project.psm_in_scope)
        self.assertTrue(project.cap_in_scope)
        self.assertEqual(project.cap_group, "2군")
        self.assertFalse(project.stage1_snapshot["legal_applicability_confirmed"])
        self.assertEqual(project.stage1_snapshot["decision"], {})
        self.assertGreater(result.updated_fields, 0)

        chemicals = project.get_field("inventory.chemicals")
        facilities = project.get_field("inventory.facilities")
        self.assertIsNotNone(chemicals)
        self.assertIsNotNone(facilities)
        self.assertEqual(chemicals.status, "USER_CONFIRMED")
        self.assertEqual(chemicals.value[0]["CAS 번호"], "108-88-3")
        self.assertEqual(facilities.value[0]["설비번호"], "TK-101")

        selected = selected_requirement_specs(project)
        self.assertFalse(any(spec.key.startswith("cap.external.") for spec in selected))
        self.assertFalse(any(key.startswith("cap.external.") for key in project.fields))

    def test_direct_template_can_be_downloaded_before_any_stage1_project_exists(self):
        project = create_direct_entry_template_project(
            psm_selected=True,
            cap_selected=True,
            cap_group="2군",
            project_id="S2-DIRECT-TEMPLATE",
        )
        actual = build_enhanced_integrated_authoring_workbook(project, example=False)
        example = build_enhanced_integrated_authoring_workbook(project, example=True)

        actual_wb = load_workbook(BytesIO(actual), data_only=False)
        example_wb = load_workbook(BytesIO(example), data_only=False)
        self.assertEqual(actual_wb["_시스템정보"]["B1"].value, "stage2-integrated-authoring-v1")
        self.assertEqual(actual_wb["_시스템정보"]["B2"].value, "INPUT")
        self.assertEqual(example_wb["_시스템정보"]["B2"].value, "EXAMPLE")
        self.assertEqual(actual_wb["01_사업장정보"]["B5"].value, None)

        writing_level = ""
        ws = actual_wb["01_사업장정보"]
        for row in range(5, ws.max_row + 1):
            if str(ws.cell(row, 1).value or "").strip() == "작성수준":
                writing_level = str(ws.cell(row, 2).value or "")
                break
        self.assertEqual(writing_level, "2군 사업장")

    def test_non_program_workbook_is_rejected(self):
        wb = Workbook()
        wb.active["A1"] = "not a stage2 workbook"
        out = BytesIO()
        wb.save(out)
        with self.assertRaisesRegex(ValueError, "프로그램 통합 작성자료가 아닙니다"):
            inspect_standalone_workbook(out.getvalue())

    def test_direct_start_requires_cap_writing_level(self):
        data = self._completed_workbook()
        wb = load_workbook(BytesIO(data))
        ws = wb["01_사업장정보"]
        for row in range(5, ws.max_row + 1):
            if str(ws.cell(row, 1).value or "").strip() == "작성수준":
                ws.cell(row, 2, "")
                break
        out = BytesIO()
        wb.save(out)
        with self.assertRaisesRegex(ValueError, "작성수준"):
            inspect_standalone_workbook(out.getvalue())

    def test_scope_page_offers_stage2_direct_start_and_both_downloads(self):
        text = (PROJECT_ROOT / "ui/stage2_scope_page.py").read_text(encoding="utf-8")
        self.assertIn("Stage 2 통합 작성자료로 직접 시작", text)
        self.assertIn("create_project_from_standalone_workbook", text)
        self.assertIn("법적 대상 여부를 새로 판정하지 않고", text)
        self.assertIn("통합 작성자료.xlsx 다운로드", text)
        self.assertIn("통합 작성자료_작성예시.xlsx 다운로드", text)
        self.assertIn("create_direct_entry_template_project", text)


if __name__ == "__main__":
    unittest.main()
