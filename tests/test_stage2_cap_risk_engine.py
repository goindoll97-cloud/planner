from __future__ import annotations

from io import BytesIO
from unittest.mock import patch
import unittest

from openpyxl import load_workbook

from engine.stage2.cap_risk_engine import (
    APPENDIX3_THRESHOLDS,
    build_cap_form14_data,
    build_cap_form15_data,
)
from engine.stage2.project import Stage2Project
from engine.stage2.workbook_enhancements import build_enhanced_integrated_authoring_workbook


class CAPRiskEngineTests(unittest.TestCase):
    def _project(self, *, industrial_complex="해당 없음") -> Stage2Project:
        p = Stage2Project(
            project_id="S2-RISK",
            company_name="테스트화학",
            cap_required=True,
            cap_group="1군",
            scope_confirmed=True,
            cap_selected=True,
        )
        p.set_field(
            "cap.business.industrial_complex",
            "산업단지",
            industrial_complex,
            "USER_CONFIRMED",
        )
        p.set_field(
            "cap.offsite.scenario_frequency",
            "시설빈도",
            [{
                "사고시나리오명": "염소 독성누출-1",
                "고압용기파열": 1,
                "배관파열": 0,
                "배관누출": 2,
                "상압 탱크 파열 및 누출": 0,
                "플랜지 등의 가스켓 파손": 0,
                "펌프/컴프레서 누출": 0,
                "안전밸브 오작동 및 조기개방": 0,
                "냉각수 손실": 0,
                "입/출하 시설 누출 사고": 0,
                "외부화재": 0,
                "개수 산정근거": "PID-201 및 설비목록",
                "수동적 완화장치": "방류벽",
                "능동적 완화장치": "",
                "안전성확보설비 증빙": "GA-201",
            }],
            "USER_CONFIRMED",
        )
        p.set_field(
            "cap.offsite.scenario_impact_table",
            "영향평가",
            [{
                "사고시나리오명": "염소 독성누출-1",
                "장외거리(m)": 180,
                "거주민수": 25,
                "근로자수": 10,
                "갑종 보호대상 수": 1,
                "을종 보호대상 수": 0,
                "환경수용체 수": 1,
                "KORA/GIS 근거": "KORA-01",
            }],
            "USER_CONFIRMED",
        )
        return p

    @patch("engine.stage2.cap_risk_engine.approved_source_is_current", return_value=True)
    def test_form14_multiplies_reference_frequency_by_count(self, _current):
        result = build_cap_form14_data(self._project())
        self.assertEqual(result.blockers, ())
        self.assertAlmostEqual(float(result.scenario_rows[0]["시설빈도(/연)"]), 0.002001)
        piping_leak = next(row for row in result.event_rows if row["개시사건"] == "배관누출")
        self.assertEqual(piping_leak["개수"], 2)
        self.assertAlmostEqual(float(piping_leak["사고빈도(/연)"]), 0.002)

    @patch("engine.stage2.cap_risk_engine.approved_source_is_current", return_value=False)
    def test_form14_does_not_auto_apply_frequency_when_law_source_is_not_current(self, _current):
        result = build_cap_form14_data(self._project())
        self.assertTrue(any("CURRENT" in blocker for blocker in result.blockers))
        self.assertEqual(result.scenario_rows[0]["시설빈도(/연)"], "")

    @patch("engine.stage2.cap_risk_engine.approved_source_is_current", return_value=True)
    def test_form15_calculates_a_b_c_d_and_pre_adjustment_score(self, _current):
        result = build_cap_form15_data(self._project())
        self.assertEqual(result.blockers, ())
        self.assertEqual(result.totals["사고시나리오 총 개수(A)"], 1)
        self.assertAlmostEqual(result.totals["사고시나리오 시설빈도의 합(B)"], 0.002001)
        self.assertEqual(result.totals["사고시나리오 거리의 합(C)"], 180)
        self.assertEqual(result.totals["주민수 합(D)"], 35)
        self.assertEqual(result.scores["사고시나리오 개수 구간점수"], 0)
        self.assertEqual(result.scores["시설빈도 구간점수"], 0)
        self.assertEqual(result.scores["거리 구간점수"], 2)
        self.assertEqual(result.scores["주민수 구간점수"], 1)
        self.assertEqual(result.scores["사고빈도점수(A+B)"], 0)
        self.assertEqual(result.scores["사고영향점수(C+D)"], 3)
        self.assertEqual(result.scores["위험도 판정표 점수(증감 전)"], 3)
        self.assertEqual(result.scores["증감 전 위험도"], "다")
        self.assertEqual(result.scores["최종 위험도"], "안전원 최종결정 전")

    @patch("engine.stage2.cap_risk_engine.approved_source_is_current", return_value=True)
    def test_industrial_park_business_excludes_workers_from_risk_population(self, _current):
        result = build_cap_form15_data(self._project(industrial_complex="울산미포국가산업단지"))
        self.assertEqual(result.blockers, ())
        self.assertEqual(result.totals["주민수 합(D)"], 25)

    @patch("engine.stage2.cap_risk_engine.approved_source_is_current", return_value=True)
    def test_explicit_no_offsite_scenario_applies_risk_d_without_form14(self, _current):
        p = self._project()
        p.fields.pop("cap.offsite.scenario_frequency", None)
        p.fields.pop("cap.offsite.scenario_impact_table", None)
        p.set_field(
            "cap.offsite.risk_control",
            "장외 사고시나리오 존재 여부",
            [{"장외 사고시나리오 없음 여부": "예", "확인근거": "KORA-00"}],
            "USER_CONFIRMED",
        )
        result = build_cap_form15_data(p)
        self.assertEqual(result.blockers, ())
        self.assertTrue(result.no_offsite_scenario)
        self.assertEqual(result.scores["최종 위험도"], "다")

    def test_appendix3_thresholds_keep_statutory_boundaries(self):
        self.assertEqual(APPENDIX3_THRESHOLDS["scenario_count"], (4.0, 16.0, 64.0))
        self.assertEqual(APPENDIX3_THRESHOLDS["facility_frequency"], (0.1, 1.0, 10.0))
        self.assertEqual(APPENDIX3_THRESHOLDS["offsite_distance_m"], (10.0, 100.0, 1000.0))
        self.assertEqual(APPENDIX3_THRESHOLDS["population"], (10.0, 100.0, 1000.0))

    def test_workbook_exposes_structured_risk_inputs(self):
        data = build_enhanced_integrated_authoring_workbook(self._project(), example=False)
        wb = load_workbook(BytesIO(data), data_only=False)
        for sheet in ("22_사고시나리오_영향평가", "23_사고시나리오_시설빈도", "24_장외시나리오_확인"):
            self.assertIn(sheet, wb.sheetnames)
        headers = [str(cell.value or "") for cell in wb["23_사고시나리오_시설빈도"][4]]
        self.assertIn("고압용기파열", headers)
        self.assertIn("개수 산정근거", headers)
        impact_headers = [str(cell.value or "") for cell in wb["22_사고시나리오_영향평가"][4]]
        self.assertIn("장외거리(m)", impact_headers)
        self.assertIn("KORA/GIS 근거", impact_headers)


if __name__ == "__main__":
    unittest.main()
