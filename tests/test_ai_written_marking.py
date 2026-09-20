from __future__ import annotations

from io import BytesIO
import json
from pathlib import Path
from unittest.mock import patch
import unittest

from docx import Document

from engine.stage2 import ai_drafting as drafting
from engine.stage2 import psm_narrative_workspace as nw
from engine.stage2 import report_export as export
from engine.stage2.output_provenance import build_output_provenance, output_provenance_json_bytes
from tests.test_psm_narrative_workspace import FakeClient, _with_facts
from tests.test_psm_workspace_page import _project
from ui.narrative_panel import state_label


def _generated():
    project = _with_facts(_project())
    with patch("engine.stage2.storage.save_project"):
        result = nw.generate(project, FakeClient())
    return project, result.generated[0].requirement_key


class MarkingTests(unittest.TestCase):
    def test_unapproved_and_approved_ai_text_are_both_listed_with_their_state(self):
        project, key = _generated()
        rows = {r["requirement_key"]: r for r in drafting.ai_written_items(project, "PSM")}
        self.assertFalse(rows[key]["approved"])
        self.assertIn("문서에 반영되지 않음", rows[key]["state"])
        nw.adopt(project, key)
        rows = {r["requirement_key"]: r for r in drafting.ai_written_items(project, "PSM")}
        self.assertTrue(rows[key]["approved"])
        self.assertEqual(rows[key]["state"], "담당자 확인 완료")
        self.assertFalse(rows[key]["edited_by_reviewer"])

    def test_reviewer_edits_are_recorded(self):
        project, key = _generated()
        nw.adopt(project, key, "확인된 설비와 물질을 기준으로 운전자가 따라야 할 절차를 담당자가 고쳐 기술합니다.")
        row = next(r for r in drafting.ai_written_items(project, "PSM") if r["requirement_key"] == key)
        self.assertTrue(row["edited_by_reviewer"])

    def test_provenance_json_carries_the_ai_written_items(self):
        project, key = _generated()
        nw.adopt(project, key)
        provenance = build_output_provenance(project, "PSM", b"docx", file_name="x_검토용.docx", final_ready=False)
        data = json.loads(output_provenance_json_bytes(provenance))
        self.assertIn(key, [row["requirement_key"] for row in data["ai_written_items"]])

    def test_review_sheet_lists_items_and_has_recheck_boxes(self):
        project, key = _generated()
        nw.adopt(project, key)
        doc = Document(BytesIO(export.build_ai_review_sheet(project, "PSM")))
        text = "\n".join(cell.text for table in doc.tables for row in table.rows for cell in row.cells)
        self.assertIn("담당자 확인 완료", text)
        self.assertIn("☐", text)
        self.assertIn("AI 작성 항목 점검표", "\n".join(p.text for p in doc.paragraphs))

    def test_empty_sheet_says_nothing_was_written_by_ai(self):
        doc = Document(BytesIO(export.build_ai_review_sheet(_project(), "PSM")))
        self.assertIn("AI가 작성한 항목이 없습니다", "\n".join(p.text for p in doc.paragraphs))

    def test_narrative_screen_labels_ai_items(self):
        self.assertEqual(state_label({"state": "초안 있음", "text": "글"}), "AI 초안 · 확인 전")
        self.assertEqual(state_label({"state": "확인 완료", "text": "글"}), "AI 작성 · 담당자 확인 완료")
        self.assertEqual(state_label({"state": "만들 수 없음", "text": ""}), "만들 수 없음")

    def test_export_panel_shows_the_review_list(self):
        text = (Path(__file__).resolve().parents[1] / "ui/report_export_panel.py").read_text(encoding="utf-8")
        self.assertIn("제출 전 다시 점검", text)
        self.assertIn("build_ai_review_sheet", text)


if __name__ == "__main__":
    unittest.main()


class PanelRenderTests(unittest.TestCase):
    def test_export_panel_shows_the_ai_review_list_and_sheet_download(self):
        from streamlit.testing.v1 import AppTest

        code = (
            "import sys; sys.path.insert(0, %r)\n"
            "from unittest.mock import patch\n"
            "from engine.stage2 import psm_narrative_workspace as nw, report_export as export\n"
            "from tests.test_psm_narrative_workspace import FakeClient, _with_facts\n"
            "from tests.test_psm_workspace_page import _project\n"
            "from ui import report_export_panel\n"
            "project = _with_facts(_project())\n"
            "with patch('engine.stage2.storage.save_project'):\n"
            "    result = nw.generate(project, FakeClient())\n"
            "nw.adopt(project, result.generated[0].requirement_key)\n"
            "export.acknowledge_holds(project)\n"
            "with patch('ui.report_export_panel.save_project'):\n"
            "    report_export_panel.render(project, 'PSM')\n"
        ) % str(Path(__file__).resolve().parents[1])
        at = AppTest.from_string(code, default_timeout=120).run()
        self.assertFalse(at.exception)
        self.assertTrue(any("제출 전 다시 점검" in e.label for e in at.expander))
        self.assertTrue(any("AI가 초안을 쓴 것" in w.value for w in at.warning))
