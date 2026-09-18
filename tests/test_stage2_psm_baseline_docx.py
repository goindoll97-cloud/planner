from __future__ import annotations

import base64
from hashlib import sha256
from io import BytesIO
from pathlib import Path
import unittest

from docx import Document

from engine.stage2.project import Stage2Project
from engine.stage2.psm_baseline_docx import (
    build_psm_baseline_draft,
    load_psm_baseline_bytes,
    psm_baseline_filename,
)


PROJECT_ROOT = Path(__file__).resolve().parents[1]
EXPECTED_SHA256 = "82457f0775850d95abd166fa6e63aae67da8fdc2ae27349d7efdb8bbc9501335"


def _project() -> Stage2Project:
    project = Stage2Project(
        project_id="PSM-BASELINE-TEST",
        company_name="세계화학(주) 충주공장",
        psm_required=True,
        cap_required=False,
    )
    project.set_authoring_scope(psm_selected=True, cap_selected=False)
    return project


def _set(project: Stage2Project, key: str, value, label: str | None = None) -> None:
    project.set_field(key, label or key, value, "USER_CONFIRMED")


def _unique_cells(row):
    seen = set()
    cells = []
    for cell in row.cells:
        marker = id(cell._tc)
        if marker in seen:
            continue
        seen.add(marker)
        cells.append(cell)
    return cells


class PSMBaselineDocxTests(unittest.TestCase):
    def test_baseline_bytes_match_uploaded_layout_snapshot(self):
        template_dir = PROJECT_ROOT / "data" / "templates" / "psm"
        parts = sorted(template_dir.glob("psm_statutory_forms_baseline.docx.b64.*"))
        raw = b"".join(
            base64.b64decode("".join(path.read_text(encoding="ascii").split()), validate=True)
            for path in parts
        )
        self.assertEqual(
            sha256(raw).hexdigest(),
            EXPECTED_SHA256,
            msg=f"stored_parts={len(parts)} decoded_bytes={len(raw)}",
        )

        data = load_psm_baseline_bytes()
        self.assertEqual(sha256(data).hexdigest(), EXPECTED_SHA256)
        doc = Document(BytesIO(data))
        self.assertEqual(len(doc.tables), 15)
        text = "\n".join(cell.text for table in doc.tables for row in table.rows for cell in row.cells)
        for marker in (
            "별지 제12호서식",
            "별지 제13호서식",
            "별지 제17호의2",
            "별지 제19호의2",
            "별지 제21호서식",
        ):
            self.assertIn(marker, text)

    def test_confirmed_business_and_chemical_values_are_written_into_existing_forms(self):
        project = _project()
        _set(project, "psm.business.project_type", "변경")
        _set(project, "business.registration_no", "123-45-67890")
        _set(project, "business.representative", "홍길동")
        _set(project, "psm.business.target_facility", "MIC 반응공정")
        _set(project, "inventory.chemicals", [
            {
                "물질명": "Methyl isocyanate",
                "CAS 번호": "624-83-9",
                "분자식": "C2H3NO",
                "폭발한계 하한": "5.3",
                "폭발한계 상한": "26",
                "인화점": "-7",
                "일일사용량": "100 kg/day",
                "저장량": "500 kg",
            }
        ])

        out = build_psm_baseline_draft(project)
        doc = Document(BytesIO(out))
        self.assertEqual(len(doc.tables), 15)

        form12 = "\n".join(cell.text for row in doc.tables[0].rows for cell in row.cells)
        self.assertIn("세계화학(주) 충주공장", form12)
        self.assertIn("123-45-67890", form12)
        self.assertIn("홍길동", form12)
        self.assertIn("MIC 반응공정", form12)
        self.assertIn("☒ 변경", form12)
        self.assertIn("☐ 설치·이전", form12)

        form13 = "\n".join(cell.text for row in doc.tables[1].rows for cell in row.cells)
        self.assertIn("Methyl isocyanate", form13)
        self.assertIn("624-83-9", form13)
        self.assertIn("C2H3NO", form13)
        self.assertIn("500 kg", form13)

    def test_form12_uses_uploaded_location_building_and_schedule_rows(self):
        project = _project()
        _set(project, "business.address", "충청북도 충주시 산업로 100")
        _set(project, "psm.business.site_building", "반응동 2동 / 연면적 1,200㎡")
        _set(project, "psm.business.schedule", "2026-10-01 ~ 2027-03-31")

        doc = Document(BytesIO(build_psm_baseline_draft(project)))
        table = doc.tables[0]
        self.assertIn("충청북도 충주시 산업로 100", _unique_cells(table.rows[14])[2].text)
        self.assertIn("반응동 2동", _unique_cells(table.rows[16])[2].text)
        self.assertIn("2026-10-01", _unique_cells(table.rows[17])[2].text)
        self.assertNotIn("반응동 2동", _unique_cells(table.rows[18])[2].text)

    def test_form12_prefers_structured_details_but_keeps_stage1_identity(self):
        project = _project()
        _set(project, "business.address", "충청북도 충주시 산업로 100")
        _set(project, "psm.business.form12_details", [{
            "사업장명": "다른회사명",
            "제출구분": "변경",
            "사업자등록번호": "123-45-67890",
            "대표자": "홍길동",
            "대상 유해·위험설비": "MIC 반응공정",
            "한국표준산업분류": "C20119",
            "근로자수": "80",
            "계약전력(kW)": "1200",
            "작성자 성명": "김작성",
            "작성자 자격": "산업안전기사",
            "주요 원료": "Methyl isocyanate",
            "주요 생산품": "제품 A",
            "사업개요": "MIC 반응·정제공정 변경",
            "사업장 소재지": "다른 주소",
            "전화번호": "043-000-0000",
            "전송번호": "해당 없음",
            "부지면적": "10,000㎡",
            "주요 건물": "반응동 2동",
            "총 사업기간": "2026-10-01 ~ 2027-03-31",
            "착공예정일": "2026-10-01",
            "시운전기간": "2027-03-01 ~ 2027-03-31",
        }])

        doc = Document(BytesIO(build_psm_baseline_draft(project)))
        form12 = "\n".join(cell.text for row in doc.tables[0].rows for cell in row.cells)

        self.assertIn("세계화학(주) 충주공장", form12)
        self.assertNotIn("다른회사명", form12)
        self.assertIn("충청북도 충주시 산업로 100", form12)
        self.assertNotIn("다른 주소", form12)
        self.assertIn("MIC 반응·정제공정 변경", form12)
        self.assertIn("김작성", form12)
        self.assertIn("산업안전기사", form12)
        self.assertIn("☒ 변경", form12)

    def test_form19_2_uses_structured_worst_and_alternative_rows(self):
        project = _project()
        base = {
            "대기온도(℃)": "25",
            "습도(%)": "50",
            "표면거칠기": "도시",
            "물질명": "염소",
            "물질의 상태": "기체",
            "설비명(또는 배관부위)": "V-201",
            "운전압력(MPa)": "0.7",
            "운전온도(℃)": "25",
            "웅덩이 크기(m2)": "해당 없음",
            "누출결과": "연속누출",
            "웅덩이(kg/s)": "해당 없음",
            "화재-4 kW/m2": "해당 없음",
            "화재-12.5 kW/m2": "해당 없음",
            "화재-37.5 kW/m2": "해당 없음",
            "폭발-7 kPa": "해당 없음",
            "폭발-21 kPa": "해당 없음",
            "폭발-70 kPa": "해당 없음",
            "인화성-25% LEL": "해당 없음",
            "인화성-LEL": "해당 없음",
            "인화성-UEL": "해당 없음",
            "계산모델·결과 근거": "KORA 결과파일",
        }
        worst = dict(base, **{
            "시나리오 구분": "최악의 사고 시나리오",
            "풍속(m/s)": "1.5",
            "대기안정도(A~F)": "F",
            "누출구의 크기(mm2)": "25",
            "직접계산(kg/s or kg)": "0.25",
            "설비/배관(kg/s)": "0.25",
            "독성-ERPG 1": "650",
            "독성-ERPG 2": "300",
            "독성-ERPG 3": "180",
        })
        alternative = dict(base, **{
            "시나리오 구분": "대안의 사고 시나리오",
            "풍속(m/s)": "3.0",
            "대기안정도(A~F)": "D",
            "누출구의 크기(mm2)": "10",
            "직접계산(kg/s or kg)": "0.10",
            "설비/배관(kg/s)": "0.10",
            "독성-ERPG 1": "320",
            "독성-ERPG 2": "150",
            "독성-ERPG 3": "90",
        })
        _set(project, "psm.risk.consequence_table", [worst, alternative])

        doc = Document(BytesIO(build_psm_baseline_draft(project)))
        form19_2 = "\n".join(cell.text for row in doc.tables[12].rows for cell in row.cells)

        self.assertIn("1.5", form19_2)
        self.assertIn("3.0", form19_2)
        self.assertIn("염소", form19_2)
        self.assertIn("650", form19_2)
        self.assertIn("300", form19_2)
        self.assertIn("90", form19_2)

    def test_form19_matches_uploaded_explosion_proof_then_exhaust_sequence_order(self):
        project = _project()
        _set(project, "psm.psi.local_exhaust", [
            {
                "공정 또는 작업장명": "혼합공정",
                "전동기용량": "7.5 kW",
                "방폭형식": "Ex d IIB T4",
                "배기 및 처리순서": "후드 → 세정기 → 배기구",
            }
        ])

        doc = Document(BytesIO(build_psm_baseline_draft(project)))
        row = _unique_cells(doc.tables[11].rows[5])
        self.assertEqual(row[8].text, "7.5 kW")
        self.assertEqual(row[9].text, "Ex d IIB T4")
        self.assertEqual(row[10].text, "후드 → 세정기 → 배기구")

    def test_unconfirmed_values_stay_blank_in_regulation_forms(self):
        project = _project()
        out = build_psm_baseline_draft(project)
        doc = Document(BytesIO(out))
        text = "\n".join(cell.text for table in doc.tables for row in table.rows for cell in row.cells)
        self.assertNotIn("[확인 필요]", text)
        self.assertIn("사업장명:", text)
        self.assertIn("세계화학(주) 충주공장", text)

    def test_psm_regulation_filename_is_distinct_from_internal_review(self):
        name = psm_baseline_filename(_project())
        self.assertTrue(name.endswith("_공정안전보고서_규정서식_작성본.docx"))
        self.assertNotIn("내부", name)

    def test_stage5_renders_regulation_form_directly_before_psm_review_heading(self):
        # Stage 5 used to detect the (already relabeled) PSM review heading via
        # a chained st.markdown monkeypatch, which never actually fired because
        # cap_official_docx's own relabeling wrapper called its own frozen
        # pre-patch st.markdown reference and bypassed the outer patch. Calling
        # _render_psm_regulation_form directly, in source order before the
        # heading, removes that failure mode entirely.
        text = (PROJECT_ROOT / "ui/stage2_review_page.py").read_text(encoding="utf-8")
        self.assertIn("공정안전보고서 · 규정서식 작성본", text)
        self.assertIn("공정안전보고서 규정서식 작성본 DOCX 다운로드", text)
        # The last occurrence is the call site in the page body; the function
        # definition itself (which also contains this substring) comes first.
        call_site = text.rindex("_render_psm_regulation_form(\n        project,")
        heading = text.index('st.markdown("### 공정안전보고서 · 내부 검토용")')
        self.assertLess(call_site, heading)
        self.assertGreater(
            call_site,
            text.index("def _render_psm_regulation_form(project, *, final_ready: bool)")
        )
        self.assertIn('final_ready=download_readiness.get("PSM", False)', text[call_site:heading])


if __name__ == "__main__":
    unittest.main()
