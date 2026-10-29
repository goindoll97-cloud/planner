from copy import deepcopy
from io import BytesIO

import pytest
from openpyxl import load_workbook

from engine.stage2 import cap_upload_template as tpl, company_workbook_intake as intake
from engine.stage2 import evidence_workspace as ew
from engine.stage2.direct_template import create_direct_cap_project


def workbook(rep='김대표', pressure=0.2, include_tank=True):
    book = load_workbook(BytesIO(tpl.build_company()))
    ws = book[tpl.SITE_SHEET]
    for row in range(2, ws.max_row + 1):
        if ws.cell(row, 1).value == '사업장명':
            ws.cell(row, 2).value = '시험'
        if ws.cell(row, 1).value == '대표자':
            ws.cell(row, 2).value = rep
    if include_tank:
        ws = book['설비']
        for c in ws[1]:
            values = {'설비번호':'TK-01', '설비명':'탱크', '용량(m3)':10, '운전압력(MPa)':pressure}
            if c.value in values:
                ws.cell(2, c.column).value = values[c.value]
    out = BytesIO(); book.save(out)
    return out.getvalue()


def project():
    return create_direct_cap_project('시험', address='주소', cap_group='')


def test_latest_workbook_replaces_renamed_file_removed_rows_and_cleared_values():
    p = project()
    intake.apply(p, '초본.xlsx', workbook())
    assert p.get_field('cap.business.representative').value == '김대표'
    assert any(f['entity'] == 'facility:TK-01' for f in ew.graph_facts(p))
    intake.apply(p, '수정본.xlsx', workbook('이대표', include_tank=False))
    assert intake.current(p)['name'] == '수정본.xlsx'
    assert len(ew.documents(p)) == 1
    assert p.get_field('cap.business.representative').value == '이대표'
    assert not any(f['entity'] == 'facility:TK-01' for f in ew.graph_facts(p))
    assert not p.get_field('inventory.facilities') or not p.get_field('inventory.facilities').value
    intake.apply(p, '최종.xlsx', workbook(None, include_tank=False))
    assert p.get_field('cap.business.representative') is None
    assert '대표자' in [r['항목'] for r in intake.corrections(p)]
    assert p.get_field(intake.HISTORY_KEY).value


def test_bad_cells_have_exact_locations_and_do_not_reach_authoring_fields():
    p = project()
    intake.apply(p, '작성.xlsx', workbook(pressure='잘못된 압력'))
    issues = intake.corrections(p)
    assert any('설비!' in r['파일·위치'] and '압력' in r['항목'] for r in issues)
    assert not any(f['field'] == 'pressure' for f in ew.graph_facts(p))
    assert p.get_field('cap.business.representative').value == '김대표'
    assert any('사업장 정보!B' in r['파일·위치'] for r in issues)


def test_unreadable_or_wrong_schema_preserves_latest_accepted_workbook():
    p = project()
    intake.apply(p, '작성.xlsx', workbook())
    before = deepcopy(p.to_dict())
    for name, data in [('파일.pdf', b'pdf'), ('깨진.xlsx', b'broken')]:
        with pytest.raises(intake.WorkbookFormatError):
            intake.apply(p, name, data)
        assert p.to_dict() == before
    book = load_workbook(BytesIO(workbook()))
    del book['설비']
    out = BytesIO(); book.save(out)
    with pytest.raises(intake.WorkbookFormatError) as error:
        intake.apply(p, '다른양식.xlsx', out.getvalue())
    assert any('설비' in r['파일·위치'] for r in error.value.issues)
    assert p.to_dict() == before


def test_migration_keeps_old_graph_in_history_and_uses_only_current_excel():
    p = project()
    ew.ingest(p, '옛자료.csv', '설비번호,설비명,용량(m3)\nOLD-01,구형탱크,5\n'.encode(), 'facility')
    ew.publish(p, replace_existing=True)
    intake.apply(p, '새양식.xlsx', workbook())
    assert not any(f['entity'] == 'facility:OLD-01' for f in ew.graph_facts(p))
    history = p.get_field(intake.HISTORY_KEY).value
    assert any('옛자료.csv' in str(h) for h in history)
    assert not any(r.get('설비번호') == 'OLD-01' for r in p.get_field('inventory.facilities').value)


def test_identical_file_is_not_reapplied_or_added_to_history():
    p = project(); data = workbook()
    assert intake.apply(p, '작성.xlsx', data)
    before = deepcopy(p.to_dict())
    assert not intake.apply(p, '작성.xlsx', data)
    assert p.to_dict() == before


def test_formula_bad_cas_and_choice_are_reported_and_corrected_in_latest_file():
    p = project()
    book = load_workbook(BytesIO(workbook()))
    def set_value(sheet, header, value):
        ws = book[sheet]
        col = next(c.column for c in ws[1] if c.value == header)
        ws.cell(2, col).value = value
    set_value('설비', '용량(m3)', '=10+20')
    set_value('설비', '시설유형', '아무 시설')
    set_value('물질', '제품코드', 'P-01')
    set_value('물질', '제품명', '톨루엔')
    set_value('물질', 'CAS 번호', '108-88-4')
    out = BytesIO(); book.save(out)
    intake.apply(p, '작성.xlsx', out.getvalue())
    messages = str(intake.corrections(p))
    assert '수식 대신' in messages
    assert '선택 가능한 값' in messages
    assert 'CAS 번호의 형식' in messages
    assert not any(f['field'] == 'capacity' for f in ew.graph_facts(p))
    set_value('설비', '용량(m3)', 30)
    set_value('설비', '시설유형', '저장시설')
    set_value('물질', 'CAS 번호', '108-88-3')
    out = BytesIO(); book.save(out)
    intake.apply(p, '고친작성.xlsx', out.getvalue())
    messages = str(intake.corrections(p))
    assert '수식 대신' not in messages
    assert 'CAS 번호의 형식' not in messages
    assert any(f['field'] == 'capacity' and f['value'] == 30 for f in ew.graph_facts(p))


def test_blank_new_excel_clears_previous_address_and_does_not_read_examples():
    p = project()
    assert p.get_field('business.address').value == '주소'
    intake.apply(p, '빈양식.xlsx', tpl.build_company())
    assert not ew.graph_facts(p)
    assert p.get_field('business.address') is None
    assert any(r['항목'] == '사업장 주소' for r in intake.corrections(p))


def test_component_percentages_and_references_are_checked():
    p = project(); book = load_workbook(BytesIO(workbook()))
    ws = book['물질']
    ws.cell(2, 1, 'M-01'); ws.cell(2, 2, '혼합제품')
    parts = book['혼합제품 구성성분']
    for row, values in enumerate([['M-01', '혼합제품', '108-88-3', '톨루엔', 80],
                                 ['M-01', '혼합제품', '67-64-1', '아세톤', 50]], 2):
        for col, value in enumerate(values, 1): parts.cell(row, col, value)
    out = BytesIO(); book.save(out)
    intake.apply(p, '작성.xlsx', out.getvalue())
    assert any('100%를 넘습니다' in r['확인된 문제'] for r in intake.corrections(p))
    parts.cell(3, 5, 20)
    parts.cell(4, 1, '없는제품'); parts.cell(4, 3, '108-88-3'); parts.cell(4, 5, 10)
    out = BytesIO(); book.save(out)
    intake.apply(p, '수정.xlsx', out.getvalue())
    assert any('물질 시트에 없는 제품코드' in r['확인된 문제'] for r in intake.corrections(p))
    assert not any('100%를 넘습니다' in r['확인된 문제'] for r in intake.corrections(p))
