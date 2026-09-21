from __future__ import annotations

from io import BytesIO
import os
from pathlib import Path
import tempfile
import unittest
import zipfile

from docx import Document

from engine.stage2.project import Stage2Project
from engine.stage2 import cap_implementation_self_check as impl
from engine.stage2 import versioning


class CAPImplementationSelfCheckTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.cwd = os.getcwd()
        os.chdir(self.tmp.name)
        self.addCleanup(os.chdir, self.cwd)
        self.addCleanup(self.tmp.cleanup)

    def _project(self, group="1군"):
        p = Stage2Project(
            project_id=f"CAP-IMPL-{group}",
            company_name="한빛화학",
            site_name="제1공장",
            cap_required=True,
            cap_group=group,
            scope_confirmed=True,
            cap_selected=True,
        )
        facts = (
            ("business.representative", "대표자", "홍길동"),
            ("business.address", "주소", "울산광역시 남구 산업로 1"),
            ("business.ksic", "KSIC", "20111"),
            ("cap.business.unit_plant_name", "단위공장명", "제1공장"),
        )
        for key, label, value in facts:
            p.set_field(key, label, value, "USER_CONFIRMED")
        versioning.freeze_version(p, "CAP", "신규제출")
        impl.save_header(p, {
            impl.BASE_VERSION_KEY: "CAP-v1.0",
            "cap.implementation.target_process": "제1공장 반응공정",
            "cap.implementation.workplace_registration_no": "CHEM-001",
            "cap.implementation.business_permit_type": "유해화학물질 영업허가",
            "cap.implementation.period_start": "2026-09-01",
            "cap.implementation.period_end": "2026-09-05",
            "cap.implementation.report_date": "2026-09-21",
        })
        impl.save_people(
            p, impl.TEAM_KEY, "자체점검반",
            [{"소속": "환경안전팀", "직급": "팀장", "성명": "김안전", "서명": ""}],
        )
        impl.save_people(
            p, impl.CONFIRMER_KEY, "사업장 확인자",
            [{"소속": "공장", "직급": "공장장", "성명": "이공장", "서명": ""}],
        )
        return p

    def _mark_all(self, p, status="확인"):
        rows = impl.checklist_rows(p)
        for row in rows:
            row["status"] = status
        impl.save_checklist(p, rows)

    def test_current_checklist_has_all_77_official_check_items(self):
        p = self._project("2군")
        rows = impl.checklist_rows(p)
        self.assertEqual(len(rows), 77)
        self.assertEqual(rows[0]["id"], "S01")
        self.assertEqual(rows[-1]["id"], "E13")
        self.assertIn("지역사회 고지 실시 여부", rows[-1]["check"])

    def test_missing_check_items_fail_closed(self):
        p = self._project("2군")
        state = impl.readiness(p)
        self.assertFalse(state.ready)
        self.assertEqual(state.total, 77)
        self.assertEqual(state.checked, 0)
        self.assertTrue(any("미확인 항목이 77건" in item for item in state.blockers))

    def test_improvement_needed_creates_form3_row_and_blocks_until_completed(self):
        p = self._project("2군")
        rows = impl.checklist_rows(p)
        for row in rows:
            row["status"] = "확인"
        rows[10]["status"] = "개선필요"
        rows[11]["status"] = "해당없음"
        impl.save_checklist(p, rows)

        proposed = impl.proposed_improvements(p)
        self.assertEqual(len(proposed), 1)
        self.assertEqual(proposed[0]["_check_id"], rows[10]["id"])
        self.assertEqual(proposed[0]["자체점검결과 개선사항"], rows[10]["check"])

        state = impl.readiness(p)
        self.assertFalse(state.ready)
        self.assertEqual(state.improvements_required, 1)
        self.assertEqual(state.improvements_completed, 0)

        proposed[0].update({
            "조치결과": "개인임무카드 재교육 및 배포 완료",
            "조치일자": "2026-09-10",
            "책임부서 (담당자)": "환경안전팀 김안전",
            "확인자": "이공장",
        })
        impl.save_improvements(p, proposed)
        state = impl.readiness(p)
        self.assertTrue(state.ready, state.blockers)
        self.assertEqual(state.improvements_completed, 1)

    def test_major_facility_requires_existing_authoring_regulation_change_log(self):
        p = self._project("1군")
        self._mark_all(p)
        state = impl.readiness(p)
        self.assertFalse(state.ready)
        self.assertTrue(any("변경내역 관리대장" in item for item in state.blockers))

        p.set_field(
            "cap.prevention.change_log",
            "변경내역 관리대장",
            [{
                "일자": "2026-08-01",
                "변경항목": "시설정보",
                "변경의 종류": "㈎ 취급시설 변경",
                "변경 내용(변경전 → 변경후)": "TK-101 → TK-101A",
                "후속조치": "㈎ 변경제출",
                "담당자": "김안전",
            }],
            "USER_CONFIRMED",
        )
        state = impl.readiness(p)
        self.assertTrue(state.ready, state.blockers)

    def test_second_group_does_not_require_change_log_for_package(self):
        p = self._project("2군")
        self._mark_all(p)
        state = impl.readiness(p)
        self.assertTrue(state.ready, state.blockers)
        package = impl.build_submission_package(p)
        with zipfile.ZipFile(BytesIO(package)) as zf:
            names = zf.namelist()
        self.assertEqual(len(names), 3)
        self.assertTrue(any("별지제1호" in name for name in names))
        self.assertTrue(any("별지제2호" in name for name in names))
        self.assertTrue(any("별지제3호" in name for name in names))

    def test_major_facility_package_contains_existing_change_log(self):
        p = self._project("1군")
        self._mark_all(p)
        p.set_field(
            "cap.prevention.change_log",
            "변경내역 관리대장",
            [{
                "일자": "2026-08-01",
                "변경항목": "시설정보",
                "변경의 종류": "㈎ 취급시설 변경",
                "변경 내용(변경전 → 변경후)": "TK-101 → TK-101A",
                "후속조치": "㈎ 변경제출",
                "담당자": "김안전",
            }],
            "USER_CONFIRMED",
        )
        package = impl.build_submission_package(p)
        with zipfile.ZipFile(BytesIO(package)) as zf:
            names = zf.namelist()
            self.assertEqual(len(names), 4)
            change_name = next(name for name in names if "변경내역_관리대장" in name)
            change_doc = Document(BytesIO(zf.read(change_name)))
            text = "\n".join(cell.text for t in change_doc.tables for row in t.rows for cell in row.cells)
            self.assertIn("TK-101", text)

    def test_form1_and_form2_output_confirmed_values(self):
        p = self._project("2군")
        self._mark_all(p)

        doc1 = Document(BytesIO(impl.build_form1_docx(p)))
        text1 = "\n".join(
            [para.text for para in doc1.paragraphs]
            + [cell.text for table in doc1.tables for row in table.rows for cell in row.cells]
        )
        self.assertIn("화학사고예방관리계획서 자체점검 결과서", text1)
        self.assertIn("CHEM-001", text1)
        self.assertIn("2026-09-01", text1)

        doc2 = Document(BytesIO(impl.build_form2_docx(p)))
        text2 = "\n".join(
            [para.text for para in doc2.paragraphs]
            + [cell.text for table in doc2.tables for row in table.rows for cell in row.cells]
        )
        self.assertIn("화학사고예방관리계획서 자체점검표", text2)
        self.assertIn("지역사회 고지 실시 여부", text2)
        self.assertIn("확인", text2)


if __name__ == "__main__":
    unittest.main()
