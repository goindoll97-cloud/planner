from __future__ import annotations

from io import BytesIO
from pathlib import Path
from unittest.mock import patch
import unittest

from docx import Document

from engine.stage2 import cap_form2_workspace as f2
from engine.stage2 import cap_form4_workspace as f4
from engine.stage2 import cap_guideline as guide
from engine.stage2 import cap_workspace as ws
from engine.stage2.cap_baseline_docx import FORM_TABLE_INDEX, build_cap_baseline_draft
from tests import test_stage2_cap_form1_engine as form1_tests


def _project():
    project = form1_tests.CAPForm1EngineTests()._project()
    rows = ws.facility_editor_rows(project)
    rows[0].update({"설비번호": "TK-1", "설비명": "염소 저장탱크", "시설유형": "저장탱크",
                    "물질성상": "액체", "용량": 2.5, "용량단위": "m3", "비중": 1.4})
    rows.append({"설비번호": "R-1", "설비명": "반응기", "취급물질": "염소", "시설유형": "제조·사용시설",
                 "물질성상": "액체", "용량": 1, "용량단위": "m3", "비중": 1.4})
    rows.append({"설비번호": "LR-1", "설비명": "탱크로리", "취급물질": "염소", "시설유형": "탱크로리·운송차량"})
    ws.save_facility_rows(project, rows)
    return project


def _docx_tables(project):
    with patch("engine.stage2.cap_form1_engine.screen_facility_stage", return_value=form1_tests.CAPForm1EngineTests._screen()):
        return Document(BytesIO(build_cap_baseline_draft(project))).tables


class CAPForm4WorkspaceTests(unittest.TestCase):
    def test_schema_follows_guideline(self):
        self.assertTrue(ws.load_form_schema(4)["title"].endswith(guide.form_guidelines()[4].title))
        header = guide.form_guidelines()[4].tables[0]
        self.assertIn("공정개요", {row[0] for row in header})

    def test_counts_and_lorries_come_from_form1_grid(self):
        project = _project()
        self.assertEqual(dict(f4.equipment_counts(project)), {"저장탱크": 1, "반응시설": 1})
        self.assertEqual(f4.lorry_count(project), 1)

    def test_overview_draft_is_built_from_entered_facts(self):
        project = _project()
        f2.save_submission(project, "신규제출", "최초", "염소 공장")
        with patch("engine.stage2.cap_form1_engine.screen_facility_stage", return_value=form1_tests.CAPForm1EngineTests._screen()):
            draft = f4.draft_overview(project)
        self.assertIn("염소 공장", draft)
        self.assertIn("저장탱크 1기", draft)
        self.assertIn("반응시설 1기", draft)

    def test_saved_inputs_reach_form4_and_form5(self):
        project = _project()
        f4.save_inputs(project, "단위공장 구성 문장", "원료 → 반응 → 출하", "2")
        self.assertEqual(f4.inputs(project).loading_units, "2")
        tables = _docx_tables(project)
        form4 = "\n".join(c.text for r in tables[FORM_TABLE_INDEX["4"][0]].rows for c in r.cells)
        self.assertIn("단위공장 구성 문장", form4)
        self.assertIn("원료 → 반응 → 출하", form4)
        self.assertIn("☒ 입·출하 시설 (2)기", form4)
        self.assertIn("☒ 보유 탱크로리 (1)기", form4)
        self.assertIn("☒ 저장탱크 (1)기", form4)
        form5 = "\n".join(c.text for r in tables[FORM_TABLE_INDEX["5"][0]].rows for c in r.cells)
        self.assertIn("원료 → 반응 → 출하", form5)  # 공정개요는 한 번만 입력

    def test_missing_items_are_listed(self):
        project = _project()
        self.assertEqual(f4.missing(project), ["단위공장 구성", "공정개요"])

    def test_page_offers_form4(self):
        page = (Path(__file__).resolve().parents[1] / "ui/cap_workspace_page.py").read_text(encoding="utf-8")
        self.assertIn("cap_form4_view", page)


if __name__ == "__main__":
    unittest.main()
