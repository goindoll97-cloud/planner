from __future__ import annotations

from io import BytesIO
import unittest

from openpyxl import load_workbook

from engine.stage2.cap_impact_engine import build_cap_form12_data, build_cap_form13_data
from engine.stage2.project import Stage2Project
from engine.stage2.workbook_enhancements import build_enhanced_integrated_authoring_workbook


class CAPImpactEngineTests(unittest.TestCase):
    def _project(self) -> Stage2Project:
        project = Stage2Project(
            project_id="S2-IMPACT",
            company_name="영향평가테스트화학",
            cap_required=True,
            cap_group="1군",
            scope_confirmed=True,
            cap_selected=True,
        )
        project.set_field(
            "inventory.chemicals",
            "화학물질 목록",
            [{"물질명": "염소", "CAS 번호": "7782-50-5", "함량(%)": 99.9, "물리적 상태": "기체"}],
            "USER_CONFIRMED",
        )
        project.set_field(
            "inventory.facilities",
            "설비 목록",
            [{"설비번호": "V-201", "설비명": "염소 용기군", "설비종류": "압력용기"}],
            "USER_CONFIRMED",
        )
        project.set_field(
            "cap.offsite.scenario_impact_table",
            "사고시나리오 영향평가",
            [{
                "사고시나리오명": "염소 독성누출-1",
                "유해화학물질명": "염소",
                "대상 설비번호": "V-201",
                "사고유형": "독성누출",
                "장외거리(m)": 180,
                "거주민수": 25,
                "근로자수": 10,
                "갑종 보호대상 수": 1,
                "을종 보호대상 수": 0,
                "환경수용체 수": 1,
                "사고원점 좌표": "35.0000,129.0000",
                "KORA/GIS 근거": "KORA-01",
            }],
            "USER_CONFIRMED",
        )
        project.set_field(
            "cap.offsite.overall_impact_summary",
            "총괄영향범위 요약",
            [{
                "총괄영향범위 산출방법": "KORA 사고시나리오 결과와 GIS 공간중첩",
                "총괄영향범위 결과 요약": "총괄영향범위는 승인 GIS 결과도면 기준",
                "GIS/KORA 근거": "KORA-ALL-01",
                "총괄영향범위 내 거주민수": 65,
                "총괄영향범위 내 근로자수": 25,
                "보호대상 없음 여부": "아니오",
            }],
            "USER_CONFIRMED",
        )
        project.set_field(
            "cap.offsite.population_and_protected_targets",
            "보호대상 명세",
            [{
                "보호대상 명칭": "○○초등학교",
                "보호대상 구분": "갑종",
                "세부유형": "교육·연구시설",
                "주소·위치": "○○시 ○○로 10",
                "좌표": "35.0010,129.0010",
                "사업장 경계와 거리(m)": 420,
                "인원수": 350,
                "GIS 근거": "GIS-PT-01",
            }],
            "USER_CONFIRMED",
        )
        project.set_field(
            "documents.kora_impact_result",
            "KORA/GIS 결과파일",
            {"file_name": "KORA_impact_result.pdf", "stored_path": "uploads/KORA_impact_result.pdf"},
            "USER_CONFIRMED",
        )
        return project

    def test_form12_cross_checks_and_prepares_confirmed_scenario_impact(self):
        result = build_cap_form12_data(self._project())

        self.assertTrue(result.ready)
        self.assertEqual(result.blockers, ())
        row = result.rows[0]
        self.assertEqual(row["사고시나리오명"], "염소 독성누출-1")
        self.assertEqual(row["유해화학물질명"], "염소")
        self.assertEqual(row["대상 설비번호"], "V-201")
        self.assertEqual(row["장외거리(m)"], 180)
        self.assertEqual(row["사고원점 좌표"], "35.0000,129.0000")

    def test_form12_unknown_facility_fails_closed(self):
        project = self._project()
        rec = project.get_field("cap.offsite.scenario_impact_table")
        row = dict(rec.value[0])
        row["대상 설비번호"] = "V-999"
        project.set_field(
            "cap.offsite.scenario_impact_table",
            rec.label,
            [row],
            "USER_CONFIRMED",
        )

        result = build_cap_form12_data(project)

        self.assertTrue(any("V-999" in blocker for blocker in result.blockers))

    def test_form13_requires_confirmed_kora_or_gis_result_file(self):
        project = self._project()
        project.fields.pop("documents.kora_impact_result", None)

        result = build_cap_form13_data(project)

        self.assertTrue(any("결과파일" in blocker for blocker in result.blockers))
        self.assertFalse(result.ready)

    def test_form13_does_not_infer_overall_geometry_from_scenario_distance(self):
        project = self._project()
        project.fields.pop("cap.offsite.overall_impact_summary", None)

        result = build_cap_form13_data(project)

        self.assertTrue(any("총괄영향범위" in blocker for blocker in result.blockers))
        self.assertEqual(result.summary, {})

    def test_explicit_no_protected_targets_allows_empty_target_table(self):
        project = self._project()
        project.fields.pop("cap.offsite.population_and_protected_targets", None)
        project.set_field(
            "cap.offsite.overall_impact_summary",
            "총괄영향범위 요약",
            [{
                "총괄영향범위 산출방법": "KORA + GIS",
                "총괄영향범위 결과 요약": "보호대상 없음",
                "GIS/KORA 근거": "KORA-ALL-00",
                "총괄영향범위 내 거주민수": 0,
                "총괄영향범위 내 근로자수": 0,
                "보호대상 없음 여부": "예",
            }],
            "USER_CONFIRMED",
        )

        result = build_cap_form13_data(project)

        self.assertTrue(result.ready)
        self.assertTrue(result.no_protected_targets)
        self.assertEqual(result.protected_targets, ())

    def test_workbook_exposes_overall_impact_and_protected_target_sheets(self):
        data = build_enhanced_integrated_authoring_workbook(self._project(), example=False)
        wb = load_workbook(BytesIO(data), data_only=False)

        self.assertIn("25_총괄영향범위_요약", wb.sheetnames)
        self.assertIn("26_총괄영향범위_보호대상", wb.sheetnames)
        summary_headers = [str(cell.value or "") for cell in wb["25_총괄영향범위_요약"][4]]
        target_headers = [str(cell.value or "") for cell in wb["26_총괄영향범위_보호대상"][4]]
        self.assertIn("보호대상 없음 여부", summary_headers)
        self.assertIn("보호대상 구분", target_headers)
        self.assertIn("사업장 경계와 거리(m)", target_headers)


if __name__ == "__main__":
    unittest.main()
