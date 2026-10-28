"""‘계획서 끝까지 연습 자료’(demo_full_company)로 안내서대로 처음부터 내려받기까지 가 봅니다.

사람이 문서를 보고 화면에 적는 일을 엔진 함수로 흉내 냅니다. 이 환경에는 법령 기준표(별표 1~4) 승인 DB가 없으므로
- 사고시나리오는 07 조사표·08 검토자료 파일을 올리면 프로그램이 규정수량(내장 별표 2 수량)으로 대상 설비를 정하고 계산합니다.
- 끝에 남는 것은 승인 DB가 필요한 항목과 검토 안내뿐이어야 합니다.
"""
import json
import unittest
from hashlib import sha256
from io import BytesIO
from unittest.mock import patch
from zipfile import ZipFile

from docx import Document
from openpyxl import load_workbook
from pyproj import Transformer
from shapely.geometry import box, mapping
from shapely.ops import transform

from engine.stage2 import (cap_form2_workspace as f2, cap_form3_workspace as f3, cap_form4_workspace as f4,
                           cap_form5_workspace as f5, cap_form6_workspace as f6, cap_form7_workspace as f7,
                           cap_form8_workspace as f8, cap_form9_workspace as f9, cap_form10_workspace as f10,
                           cap_form11_workspace as f11, cap_form14_workspace as f14, cap_form16_workspace as f16,
                           cap_impact_workspace as iw, cap_narrative_workspace as nw, cap_scenario_workspace as sc,
                           cap_workspace as ws, demo_full_company as d, evidence_workspace as ew,
                           psm_narrative_workspace as pn, report_export as rx, cap_surroundings as survey,
                           cap_site_lookup as kakao, cap_vworld_lookup as vworld, scenario_pipeline as sp)
from engine.stage2.cap_baseline_docx import build_cap_baseline_draft
from engine.stage2.completeness import evaluate_project_completeness
from engine.stage2.direct_template import create_direct_cap_project
from engine.stage2.project import EvidenceRef
from engine.stage2.cap_site_lookup import Candidate

D, C = d.DATA, d.DATA["company"]
# 승인 DB(법령 기준표)가 있어야 풀리는 것과, 최종 제출 전 사람이 확인하라는 안내
ALLOWED = ("승인 DB", "법령 DB", "별표 2", "별표 3", "Stage 1 판정", "적용하지 않습니다")


REAL = "울산광역시 남구 연습로 1"           # 실제 주소 판 시험용(조회는 가짜 응답으로 대신함)
CENTER = (129.3589, 35.4779)
TO_M = Transformer.from_crs(4326, 5179, always_xy=True).transform
TO_LL = Transformer.from_crs(5179, 4326, always_xy=True).transform


def _ref(name, content=None):
    content = d.files()[name] if content is None else content
    return EvidenceRef(source_type="COMPANY_EVIDENCE", source_name=name, sha256=sha256(content).hexdigest(), location=name)


def _box(dx1, dy1, dx2, dy2):
    x, y = TO_M(*CENTER)
    return mapping(transform(TO_LL, box(x + dx1, y + dy1, x + dx2, y + dy2)))


# 브이월드가 돌려줄 법한 주변 객체(주소점에서 동서남북)
MAP_OBJECTS = [Candidate("연습 초등학교", "", "", "주소점 남동쪽", None, "브이월드 도로명주소 건물", "b1", _box(250, -300, 330, -240)),
               Candidate("연습 아파트", "", "", "주소점 남쪽", None, "브이월드 도로명주소 건물", "b2", _box(-40, -260, 60, -200)),
               Candidate("연습천", "", "", "주소점 서쪽", None, "브이월드 하천망", "r1", _box(-260, -400, -240, 400))]


def form8_from_map(p):
    """실제 주소 판 안내서의 별지 제8호 ①~⑤를 그대로 따라 합니다."""
    with patch.object(vworld, "api_key", return_value="key"), patch.object(vworld, "geocode", return_value=CENTER), \
            patch.object(vworld, "search", return_value=(MAP_OBJECTS, [])), patch.object(kakao, "api_key", return_value=""):
        survey.collect(p)                                                                           # ①
    survey.save_boundary(p, _box(-75, -50, 75, 50), "연습용 가상 부지 경계, 2026-10-27", confirmed=True)   # ②
    for group in survey.group_suggestions(p, 60):                                                   # ③
        saved = survey.save_group(p, group["name"], group["members"], confirmed=True)
        survey.import_group(p, saved["group_id"])
    rows, _, _ = f8.classify_rows(f8.saved_rows(p))                                                 # ④
    distance = {r["name"]: r["boundary_distance_m"] for r in survey.annotate(p)}
    kinds = {"연습 초등학교": ("갑종", "교육·연구시설"), "연습 아파트": ("갑종", "주택"), "연습천": ("환경수용체", "하천")}
    for row in rows:
        row["보호대상 구분"] = row["보호대상 구분"] or kinds[row["보호대상 명칭"]][0]
        row["세부유형"] = row["세부유형"] or kinds[row["보호대상 명칭"]][1]
        row["사업장 경계와 거리(m)"] = round(distance[row["보호대상 명칭"]])
        row["GIS/현장 근거"] = "브이월드 공간자료 2026-10-27"
    f8.save(p, rows, False, "", scope_reviewed=True)
    shot = b"\x89PNG map screenshot"
    f8.save_review(p, {key: True for key, _ in f8.REVIEW_ITEMS}, "브이월드 공간자료·항공사진, 2026-10-27",   # ⑤
                   "그린 부지 경계에서 500m 안을 지도로 확인", ",".join(str(i) for i in range(1, len(rows) + 1)),
                   [_ref("지도_캡처.png", shot)])


def walk(address=""):
    C = d.company(address)
    p = create_direct_cap_project(d.COMPANY, address=C["주소"], cap_group="")
    for name, kind in d.UPLOAD.items():                                              # ① 자료 올리기
        ew.ingest(p, name, d.files(address)[name], kind)
    questions = [i for i in ew.review_queue(p) if i["kind"] != "확인 필요"] + ew.problems(p)
    ew.approve_bulk(p)                                                                # ② 확인
    ew.confirm_auto(p, [s["key"] for s in ew.slots(p) if s["method"] == "AUTO"])     # ④ 마지막 확인
    ew.publish(p, replace_existing=True)                                              # ④ 보내기
    p.cap_group = C["작성수준"]                                                        # 작성수준 확인
    p.set_field("cap.business.writing_level", "작성수준", f"{p.cap_group} 사업장", "USER_CONFIRMED", note="회사 확인값")
    rows = ws.facility_editor_rows(p)                                                 # 별지 제1호
    extra = {f["id"]: f["form1"] for f in D["facilities"]}
    for row in rows:
        row.update(extra.get(row.get("설비번호"), {}))
    ws.save_facility_rows(p, rows)
    f2.save_submission(p, C["제출구분"], C["제출 사유"], C["단위공장명"])                 # 별지 제2호
    f3.save_fields(p, {f["id"]: C.get(f["id"], C["주소"] if f["id"] == "address" else "") for f in f3.fields()})
    f4.save_inputs(p, D["process"]["총괄 개요"], D["process"]["공정개요"], D["process"]["입출하 시설 수"])
    f5.save_overview(p, D["process"]["단위공장 개요"])
    f6.save_properties(p, [{"CAS 번호": pr["cas"], "물질상태": d._sds_state(pr["state"]), "비중": str(pr["density"]), **pr["props"]}
                           for pr in D["products"]])
    reps = [pr for pr in D["products"] if pr["code"] in ("P-01", "P-03")]
    f7.save_rows(p, [{**f7.draft_row(p, pr["cas"], pr["name"], pr["hazard"]["선정 사유"]), **pr["hazard"],
                      "SDS 파일명": pr["props"]["SDS 파일명"], "SDS 개정일": pr["props"]["SDS 개정일"]} for pr in reps])
    if address:                                                                       # 별지 제8호(주소판: 지도 조회)
        form8_from_map(p)                                                             # 가상판은 07 조사표 파일에서 읽음
    specs = {f["id"]: f["spec"] for f in D["facilities"]}                             # 별지 제9·10·11호
    f9.save_specs(p, [{**r, **specs.get(r["구분기호"], {})} for r in f9.rows(p)])
    dikes = {f["id"]: f["dike"] for f in D["facilities"]}
    f10.save(p, [{**r, **dikes.get(r.get("구분기호"), {})} for r in f10.edit_rows(p)], D["dike_rule"]["비율(%)"], D["dike_rule"]["근거"])
    f11.save(p, D["detectors"])
    f11.attach_evidence(p, [_ref("06_가스감지기_배치도.png")])
    for fact in nw.DECISION_FACTS:                                                    # 서술형 결정사항
        values = D["narrative"].get(fact.key.split(".")[-1])
        if values:
            pn.save_fact(p, fact, values)
    for item in f16.items("internal"):                                                # 별지 제16호
        if item["id"] in D["emergency"]:
            f16.save_text(p, item, D["emergency"][item["id"]])
    return p, questions


def scenarios(p):
    """화면이 하는 일: 올린 파일로 시나리오를 다시 계산합니다(사람은 고르지 않음)."""
    outcome = sp.refresh(p)
    rows = p.get_field(iw.IMPACT_TABLE_KEY).value
    return outcome.impacts, {"scenarios": len(rows), "problems": [(x.where, x.what) for x in outcome.problems]}


class FullDemoTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.p, cls.questions = walk()
        cls.impacts, cls.applied = scenarios(cls.p)

    def test_package_has_every_document_the_guide_names(self):
        with ZipFile(BytesIO(d.full_demo_zip())) as z:
            names = set(z.namelist())
            guide = z.read("README_테스트_안내서.txt").decode("utf-8")
        self.assertEqual(names, set(d.files()) | {"README_테스트_안내서.txt", "answer_key.json"})
        with ZipFile(BytesIO(d.full_demo_zip())) as z:
            key = json.loads(z.read("answer_key.json"))
        self.assertEqual(len(key["facilities"]), 4)
        self.assertEqual(len(key["products"]), 3)
        for name in d.files():
            self.assertIn(name, guide, name)
        for name, kind in d.UPLOAD.items():
            self.assertEqual(ew.classify_document(name, d.files()[name]), kind)

    def test_clean_company_data_raises_no_questions(self):
        self.assertEqual(self.questions, [])

    def test_every_company_input_form_is_complete(self):
        p = self.p
        self.assertEqual(f3.missing_labels(p), [])
        self.assertEqual(f4.missing(p), [])
        self.assertEqual(f5.missing(p), [])
        self.assertEqual(f8.needs(p), [])
        self.assertEqual(f9.needs(p), [])
        self.assertEqual(f10.needs(p), [])
        self.assertEqual(f11.needs(p), [])
        self.assertEqual(f14.needs(p), [])
        self.assertEqual(f16.missing(p), [])
        self.assertTrue(all(any(word in need for word in ALLOWED) for need in f7.needs(p)), f7.needs(p))

    def test_scenarios_compute_from_the_package_values(self):
        self.assertEqual(len(self.impacts), 4)
        self.assertTrue(all(i.radius_m and not i.problems for i in self.impacts), [(i.name, i.problems) for i in self.impacts])
        self.assertGreaterEqual(self.applied["scenarios"], 1)
        self.assertEqual(self.applied["problems"], [])
        kinds = {s["대상 설비번호"]: s["사고유형"] for s in sc.saved_scenarios(self.p)}
        self.assertEqual(kinds, {"TK-01": sc.FIRE, "TK-02": sc.TOXIC, "RX-01": sc.FIRE, "CY-01": sc.TOXIC})   # 08 파일의 선정

    def test_only_legal_db_items_remain(self):
        state = rx.evaluate(self.p, "CAP")
        left = [i.message for i in list(state.text_issues) + list(state.manual_issues) if i.status != "PASS"]
        self.assertTrue(left)
        self.assertEqual([m for m in left if not any(word in m for word in ALLOWED)], [])

    def test_docx_contains_the_company_values(self):
        doc = Document(BytesIO(build_cap_baseline_draft(self.p)))
        text = "\n".join(c.text for t in doc.tables for row in t.rows for c in row.cells)
        for value in (C["representative"], C["registration_no"], "GD-01", "가상 아파트 단지", "TK-01", "CY-01 가상 액화암모니아 독성 누출"):
            self.assertIn(value, text)

    def test_evidence_survives_resave(self):
        f8.save(self.p, f8.saved_rows(self.p), no_target=False, scope_reviewed=True)
        f11.save(self.p, f11.saved_rows(self.p))
        self.assertTrue(self.p.get_field(f8.SITE_KEY).evidence)
        self.assertTrue(self.p.get_field(f11.GAS_KEY).evidence)


class RealAddressTests(unittest.TestCase):
    """주소를 넣어 내려받은 판: 그 주소로 지도 조회해 별지 제8호를 채우고 끝까지 갑니다."""

    @classmethod
    def setUpClass(cls):
        cls.p, cls.questions = walk(REAL)
        cls.impacts, cls.applied = scenarios(cls.p)

    def test_package_uses_the_address_and_drops_invented_surroundings(self):
        with ZipFile(BytesIO(d.full_demo_zip(REAL, "연습 국가산업단지"))) as z:
            names = set(z.namelist())
            guide = z.read("README_테스트_안내서.txt").decode("utf-8")
            general = Document(BytesIO(z.read("03_사업장_일반정보.docx")))
            survey_sheet = load_workbook(BytesIO(z.read("07_주변보호대상_조사표.xlsx")))["보호대상"]
        cells = [c.text for t in general.tables for row in t.rows for c in row.cells]
        self.assertIn(REAL, cells)
        self.assertIn("연습 국가산업단지", cells)
        self.assertNotIn("07_주변보호대상_지도.png", names)
        self.assertEqual(survey_sheet.max_row, 1)                       # 빈 양식(머리글만)
        with ZipFile(BytesIO(d.full_demo_zip(REAL))) as z:
            self.assertIsNone(json.loads(z.read("answer_key.json"))["surroundings"])
        for words in (REAL, "주소 필지 경계를 초안으로 불러오기", "VWORLD_API_KEY", "그린 경계 저장", "주변 보호대상 조사표 내려받기"):
            self.assertIn(words, guide)
        self.assertNotIn("가상 초등학교", guide)
        self.assertEqual(d.company(REAL)["industrial_complex"], "해당 없음")
        with ZipFile(BytesIO(d.full_demo_zip())) as z:                 # 주소를 안 넣으면 예전 그대로
            self.assertIn("07_주변보호대상_지도.png", z.namelist())

    def test_map_lookup_fills_form8_and_the_plan_reaches_the_end(self):
        self.assertEqual(self.questions, [])
        self.assertEqual(f8.needs(self.p), [])
        self.assertEqual(f8.address(self.p), REAL)
        names = {r["보호대상 명칭"] for r in f8.saved_rows(self.p)}
        self.assertEqual(names, {"연습 초등학교", "연습 아파트", "연습천"})
        self.assertTrue(all(i.radius_m and not i.problems for i in self.impacts))
        state = rx.evaluate(self.p, "CAP")
        left = [i.message for i in list(state.text_issues) + list(state.manual_issues) if i.status != "PASS"]
        self.assertEqual([m for m in left if not any(word in m for word in ALLOWED)], [])


class NonDrawingAuditTests(unittest.TestCase):
    """도면을 제외해도 남는 작업을 별도 공개해 '끝까지'라는 이름으로 가리지 않습니다."""

    def test_full_walk_keeps_non_drawing_gaps_visible(self):
        p, questions = walk()
        scenarios(p)
        self.assertEqual(questions, [])
        state = rx.evaluate(p, "CAP")
        self.assertFalse(state.final_ready)
        drawing_keys = {"documents.pfd", "documents.pid", "documents.site_plan",
                        "cap.site.unit_plant_layout", "cap.site.equipment_layout", "cap.safety.dike_layout"}
        remaining = []
        for item in evaluate_project_completeness(p)["requirements"]:
            if item["system"] not in {"CAP", "COMMON"} or item["state"] in {"READY", "NOT_REQUIRED"}:
                continue
            keys = (set(item["missing_fields"]) | set(item["hold_fields"]) | set(item["draft_fields"])) - drawing_keys
            if keys:
                remaining.append((item["label"], keys))
        self.assertGreater(len(remaining), 0, "도면 외 항목이 실제로 연결되면 이 시험과 안내를 갱신하세요")
        self.assertTrue(any("cap.facility.equipment_specs" in keys for _, keys in remaining))
        self.assertTrue(any("cap.offsite.impact_range_result" in keys for _, keys in remaining))

    def test_stress_pack_exposes_merges_instead_of_claiming_success(self):
        files = d.stress_files()
        with ZipFile(BytesIO(d.stress_demo_zip())) as z:
            self.assertEqual(set(z.namelist()), set(files) | {"answer_key.json", "README.txt"})
            truth = json.loads(z.read("answer_key.json"))
        self.assertEqual(len(truth["facilities"]), 3)
        self.assertEqual(len(truth["components"]), 2)
        p = create_direct_cap_project("가상 압박 시험", address="가상 주소", cap_group="")
        for name, content in files.items():
            if name.endswith(".docx"):
                ew.ingest(p, name, content, "facility", "T-02")
            else:
                kind = "material" if "혼합물" in name else "facility"
                self.assertEqual(ew.classify_document(name, content), kind)
                ew.ingest(p, name, content, kind)
        facilities = {s["entity"] for s in ew.slots(p) if s["entity"].startswith("facility:")}
        self.assertEqual(facilities, {"facility:T-01", "facility:T-02"})  # 실제 정답은 3대: 알려진 결함
        cas = next(s for s in ew.slots(p) if s["key"] == "material:MIX-01|cas")
        self.assertEqual(cas["status"], "값 충돌")                  # 성분 2개를 하나로 묶음
        inventory = next(s for s in ew.slots(p) if s["key"] == "facility:T-02|inventory")
        self.assertEqual(inventory["status"], "값 충돌")          # 개정 자료를 자동으로 채택하지 않음


if __name__ == "__main__":
    unittest.main()
