from __future__ import annotations

from io import BytesIO
import unittest

from openpyxl import load_workbook

from engine.stage2.completeness import evaluate_project_completeness
from engine.stage2.integrated_workbook import apply_integrated_authoring_workbook
from engine.stage2.project import create_project_from_stage1_snapshot
from engine.stage2.workbook_enhancements import build_enhanced_integrated_authoring_workbook


class PSMLaterFormWorkbookTests(unittest.TestCase):
    def _project(self):
        snapshot = {
            "source_fingerprint": "d" * 64,
            "business": {
                "회사명": "PSM후속서식테스트",
                "사업장명": "제1공장",
                "사업장 소재지": "테스트시 산업로 1",
            },
            "documents": {},
            "chemicals": [{
                "물질명": "톨루엔",
                "CAS 번호": "108-88-3",
                "함량(%)": 99.5,
                "최대보유량": 15000,
                "단위": "kg",
            }],
            "facilities": [{
                "설비번호": "TK-101",
                "설비명": "톨루엔 저장탱크",
                "설비종류": "저장탱크",
            }],
            "decision": {
                "psm_status": "공정안전보고서 제출 대상",
                "cap_status": "비대상",
            },
        }
        project = create_project_from_stage1_snapshot(snapshot)
        project.set_authoring_scope(psm_selected=True, cap_selected=False)
        return project

    @staticmethod
    def _headers(ws):
        return {
            str(cell.value or "").strip(): cell.column
            for cell in ws[4]
            if str(cell.value or "").strip()
        }

    def test_input_workbook_exposes_applicability_and_all_remaining_structured_forms(self):
        wb = load_workbook(BytesIO(
            build_enhanced_integrated_authoring_workbook(self._project(), example=False)
        ))

        expected_sheets = {
            "09_PSM_조건부서식_적용여부",
            "12_PSM_인터록",
            "13_PSM_소화설비",
            "14_PSM_화재탐지",
            "15_PSM_내화구조",
            "16_PSM_국소배기",
            "17_PSM_방폭기기",
            "18_PSM_위험성평가자",
        }
        self.assertTrue(expected_sheets.issubset(set(wb.sheetnames)))

        app = wb["09_PSM_조건부서식_적용여부"]
        form_numbers = [app.cell(row=row, column=1).value for row in range(5, 12)]
        self.assertEqual(form_numbers, ["17-2", "17-3", "17-4", "17-5", "18", "19", "20"])
        self.assertTrue(all(app.cell(row=row, column=3).value in (None, "") for row in range(5, 12)))
        self.assertTrue(any(
            "_선택목록" in str(dv.formula1)
            for dv in app.data_validations.dataValidation
        ))

        gas = wb["05_가스누출감지_경보장치"]
        self.assertIn("경보 위치", self._headers(gas))
        self.assertIn("연동 설비·조치", self._headers(gas))

    def test_incomplete_applicability_table_import_stays_hold(self):
        project = self._project()
        raw = build_enhanced_integrated_authoring_workbook(project, example=False)
        wb = load_workbook(BytesIO(raw))
        ws = wb["09_PSM_조건부서식_적용여부"]
        ws.cell(5, 3, "적용")
        ws.cell(5, 4, "P&ID 확인")

        out = BytesIO()
        wb.save(out)
        apply_integrated_authoring_workbook(project, out.getvalue())

        record = project.get_field("psm.psi.form_applicability")
        self.assertIsNotNone(record)
        self.assertEqual(record.status, "HOLD")
        self.assertIn("일부 비어", record.note)

    def test_complete_applicability_table_import_is_user_confirmed(self):
        project = self._project()
        raw = build_enhanced_integrated_authoring_workbook(project, example=False)
        wb = load_workbook(BytesIO(raw))
        ws = wb["09_PSM_조건부서식_적용여부"]
        values = {
            "17-2": ("적용", "P&ID 및 인터록 목록"),
            "17-3": ("적용", "소화설비 배치도"),
            "17-4": ("해당 없음", "화재탐지 적용대상 없음 검토"),
            "17-5": ("적용", "가스감지기 목록"),
            "18": ("해당 없음", "내화 적용대상 없음 검토"),
            "19": ("해당 없음", "국소배기 적용대상 없음 검토"),
            "20": ("적용", "폭발위험장소 구분도"),
        }
        for row in range(5, 12):
            form_no = str(ws.cell(row, 1).value)
            applicable, basis = values[form_no]
            ws.cell(row, 3, applicable)
            ws.cell(row, 4, basis)

        out = BytesIO()
        wb.save(out)
        apply_integrated_authoring_workbook(project, out.getvalue())

        record = project.get_field("psm.psi.form_applicability")
        self.assertIsNotNone(record)
        self.assertEqual(record.status, "USER_CONFIRMED")
        self.assertEqual(len(record.value), 7)

    def test_structured_table_import_uses_table_field_not_attachment_field(self):
        project = self._project()
        raw = build_enhanced_integrated_authoring_workbook(project, example=False)
        wb = load_workbook(BytesIO(raw))
        ws = wb["15_PSM_내화구조"]
        headers = self._headers(ws)
        ws.cell(5, headers["내화설비 또는 지역"], "R-101 지지철골")
        ws.cell(5, headers["내화부위"], "주기둥 및 보")
        ws.cell(5, headers["내화시험기준 및 시간"], "2시간")

        out = BytesIO()
        wb.save(out)
        apply_integrated_authoring_workbook(project, out.getvalue())

        table = project.get_field("psm.psi.fireproofing_table")
        attachment = project.get_field("psm.psi.fireproofing")
        self.assertIsNotNone(table)
        self.assertEqual(table.status, "USER_CONFIRMED")
        self.assertIsNone(attachment)

    def test_attachment_requirements_remain_separate_from_structured_tables(self):
        project = self._project()
        wb = load_workbook(BytesIO(
            build_enhanced_integrated_authoring_workbook(project, example=False)
        ))
        ws = wb["07_도면_첨부자료목록"]
        labels = {
            str(ws.cell(row, 1).value or "").strip()
            for row in range(5, ws.max_row + 1)
        }

        self.assertIn("내화구조 명세·도면", labels)
        self.assertIn("소화설비 설치계획·시방·용량계산·도면", labels)
        self.assertIn("화재탐지 및 경보설비 설치계획", labels)
        self.assertIn("가스누출감지 및 경보장치 설치계획", labels)
        self.assertIn("국소배기장치 설치계획", labels)

    def test_structured_table_alone_does_not_make_fireproofing_requirement_ready(self):
        project = self._project()
        project.set_field(
            "psm.psi.fireproofing_table",
            "내화구조 명세 표",
            [{
                "내화설비 또는 지역": "R-101 지지철골",
                "내화부위": "주기둥 및 보",
                "내화시험기준 및 시간": "2시간",
            }],
            "USER_CONFIRMED",
        )

        completeness = evaluate_project_completeness(project)
        fireproof = next(
            item for item in completeness["requirements"]
            if item["key"] == "psm.psi.fireproofing"
        )

        self.assertEqual(fireproof["state"], "HOLD")
        self.assertIn("psm.psi.fireproofing", fireproof["missing_fields"])


if __name__ == "__main__":
    unittest.main()
