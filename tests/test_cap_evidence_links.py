import unittest
from io import BytesIO

from docx import Document

from engine.stage2 import cap_evidence_links as links
from engine.stage2 import demo_full_company as demo
from engine.stage2 import evidence_workspace as ew
from engine.stage2 import cap_form2_workspace as form2
from engine.stage2 import cap_form16_engine as form16
from engine.stage2 import statutory_report as report
from engine.stage2.direct_template import create_direct_cap_project


def document(rows):
    doc = Document()
    table = doc.add_table(rows=0, cols=2)
    for label, value in rows:
        cells = table.add_row().cells
        cells[0].text, cells[1].text = label, value
    out = BytesIO()
    doc.save(out)
    return out.getvalue()


class EvidenceLinksTests(unittest.TestCase):
    def setUp(self):
        self.project = create_direct_cap_project("가상 사업장", cap_group="2군")

    def test_synthetic_company_documents_populate_real_form_keys(self):
        files = demo._files()
        for name in ("03_사업장_일반정보.docx", "04_공정설명서.docx", "09_안전관리_비상대응계획.docx"):
            self.assertTrue(ew.register_reference(self.project, name, files[name]))
        registration = self.project.get_field("cap.business.registration_no")
        self.assertEqual(registration.status, "FILE_READ")                     # 승인 단계 없이 바로 초안
        self.assertEqual(registration.evidence[0].source_name, "03_사업장_일반정보.docx")
        self.assertIn("행", registration.evidence[0].location)
        self.assertEqual(self.project.get_field("process.description").status, "FILE_READ")
        self.assertEqual(self.project.get_field("cap.basic.unit_facility_overview").status, "FILE_READ")
        self.assertEqual(self.project.get_field("cap.prevention.emergency_contact_system").status, "FILE_READ")
        self.assertTrue(links.linked(self.project.get_field("cap.prevention.emergency_contact_system")))

    def test_one_company_file_reuses_facts_in_forms_2_3_16(self):
        files = demo._files()
        ew.register_reference(self.project, "03_사업장_일반정보.docx", files["03_사업장_일반정보.docx"])
        self.assertEqual(form2.submission(self.project)["unit_plant"], demo.DATA["company"]["단위공장명"])
        self.assertEqual(self.project.get_field(form2.SUBMISSION_TYPE_KEY).status, "FILE_READ")
        self.assertEqual(form2.submission(self.project)["type"], demo.DATA["company"]["제출구분"])
        self.assertEqual(form16.build_cap_form16_data(self.project).business["대표자"],
                         demo.DATA["company"]["representative"])
        doc = Document()
        report._cap_form3(doc, self.project)
        values = [cell.text for table in doc.tables for row in table.rows for cell in row.cells]
        self.assertIn(demo.DATA["company"]["단위공장명"], values)
        self.assertIn(demo.DATA["company"]["writer_name"], " ".join(values))

    def test_one_process_file_reuses_description_in_forms_4_5(self):
        files = demo._files()
        ew.register_reference(self.project, "04_공정설명서.docx", files["04_공정설명서.docx"])
        doc = Document()
        report._cap_facility_overview(doc, self.project, detailed=False)
        report._cap_facility_overview(doc, self.project, detailed=True)
        rows = [[cell.text for cell in row.cells] for table in doc.tables for row in table.rows]
        descriptions = [row[1] for row in rows if row[0] == "공정개요"]
        self.assertEqual(len(descriptions), 2)
        self.assertEqual(descriptions, [demo.DATA["process"]["공정개요"]] * 2)
        overviews = [row[1] for row in rows if row[0] == "단위공장 구성"]
        self.assertEqual(overviews, [demo.DATA["process"]["총괄 개요"], demo.DATA["process"]["단위공장 개요"]])

    def test_hwpx_and_binary_hwp_link_identically(self):
        files = demo._files()
        for suffix in ("hwpx", "hwp"):
            project = create_direct_cap_project("시험", cap_group="2군")
            name = f"03_사업장_일반정보_한글판.{suffix}"
            ew.register_reference(project, name, files[name])
            self.assertEqual(project.get_field("cap.business.registration_no").value,
                             demo.DATA["company"]["registration_no"])

    def test_conflict_holds_both_sources_and_preserves_user_edit(self):
        ew.register_reference(self.project, "old.docx", document([("대표전화", "052-111-1111")]))
        ew.register_reference(self.project, "new.docx", document([("대표전화", "052-222-2222")]))
        record = self.project.get_field("cap.business.contact")
        self.assertEqual(record.status, "HOLD")
        self.assertEqual(len(record.evidence), 2)
        self.assertEqual(record.value, "")
        self.project.set_field(record.key, record.label, "052-333-3333", "USER_CONFIRMED")
        ew.register_reference(self.project, "third.docx", document([("대표전화", "052-444-4444")]))
        self.assertEqual(self.project.get_field(record.key).value, "052-333-3333")

    def test_removing_the_file_removes_the_values_it_alone_supplied(self):
        ew.register_reference(self.project, "a.docx", document([("대표전화", "052-111-1111")]))
        sha = ew.documents(self.project)[0]["sha256"]
        ew.remove_document(self.project, sha)
        self.assertIsNone(self.project.get_field("cap.business.contact"))

    def test_unstructured_and_invalid_choice_never_confirm(self):
        ew.register_reference(self.project, "other.docx", document([
            ("최근 3년간 화학사고 발생 여부", "조사 중"), ("임의 항목", "값"),
            ("입·출하 시설 수(기)", "수량 미확인")]))
        self.assertIsNone(self.project.get_field("cap.business.recent_accident"))
        self.assertIsNone(self.project.get_field("cap.basic.loading_transport"))
        self.assertTrue(ew.documents(self.project)[0]["needs_reading"])


if __name__ == "__main__":
    unittest.main()
