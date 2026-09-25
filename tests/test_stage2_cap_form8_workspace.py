from __future__ import annotations

from io import BytesIO
from pathlib import Path
import os
import unittest
from unittest.mock import patch

from docx import Document

from engine.stage2 import cap_form8_workspace as f8
from engine.stage2 import cap_guideline as guide
from engine.stage2 import cap_site_lookup as lookup
from engine.stage2 import cap_workspace as ws
from engine.stage2.cap_baseline_docx import FORM_TABLE_INDEX, build_cap_baseline_draft
from engine.stage2.project import Stage2Project

ROOT = Path(__file__).resolve().parents[1]


def _project() -> Stage2Project:
    project = Stage2Project(project_id="S2-F8", company_name="테스트화학", site_name="공장", cap_required=True,
                            scope_confirmed=True, cap_selected=True, cap_group="2군")
    project.set_field("business.address", "사업장 소재지", "울산광역시 남구 산업로 1", "VERIFIED")
    return project


def _fake_get(url, params, headers):
    if url.endswith("address.json"):
        return {"documents": [{"x": "129.3", "y": "35.5"}]}
    code = params.get("category_group_code")
    if code == "SC4":
        return {"documents": [{"place_name": "한빛초등학교", "distance": "310", "road_address_name": "산업로 9"}]}
    if params.get("query") == "하천":
        return {"documents": [{"place_name": "태화강", "distance": "450", "address_name": "울산 남구"}]}
    return {"documents": []}


class CAPForm8WorkspaceTests(unittest.TestCase):
    def test_schema_follows_guideline_and_annex4(self):
        self.assertTrue(ws.load_form_schema(8)["title"].endswith(guide.form_guidelines()[8].title))
        rules = guide.protected_target_rules()
        self.assertEqual(set(rules), {"갑종", "을종", "환경수용체"})
        self.assertIn("300명", f8.type_hint("갑종", "종교시설"))
        self.assertIn("20명", f8.type_hint("갑종", "노유자시설"))

    def test_printed_options_match_the_form_checkboxes(self):
        header = " ".join(cell for row in guide.form_guidelines()[8].tables[0] for cell in row)
        for options in f8.SUBTYPES.values():
            for option in options:
                self.assertIn(option.replace(" ", ""), header.replace(" ", ""))

    def test_no_key_means_no_lookup(self):
        with patch.dict(os.environ, {lookup.ENV_KEY: ""}):
            found, message = lookup.find_candidates("울산", get=_fake_get)
        self.assertEqual(found, [])
        self.assertIn(lookup.ENV_KEY, message)

    def test_candidates_are_proposed_from_address_and_sorted_by_distance(self):
        with patch.dict(os.environ, {lookup.ENV_KEY: "k"}):
            found, _ = lookup.find_candidates("울산광역시 남구 산업로 1", get=_fake_get)
        self.assertEqual([c.name for c in found], ["한빛초등학교", "태화강"])
        self.assertEqual((found[0].category, found[0].subtype), ("갑종", "교육·연구시설"))
        self.assertEqual(found[1].category, "환경수용체")

    def test_candidate_distance_is_reference_only_until_boundary_distance_is_entered(self):
        with patch.dict(os.environ, {lookup.ENV_KEY: "k"}):
            found, _ = lookup.find_candidates("울산광역시 남구 산업로 1", get=_fake_get)
        row = f8.candidate_row(found[0])
        self.assertEqual(row["사업장 경계와 거리(m)"], "")
        self.assertEqual(row["검색결과 거리(주소점 기준, 참고)"], 310)
        self.assertEqual(row["GIS/현장 근거"], "")
        self.assertFalse(row["500m 범위 전체 확인"])

    def test_added_search_candidate_is_proposed_not_confirmed(self):
        project = _project()
        with patch.dict(os.environ, {lookup.ENV_KEY: "k"}):
            found, _ = lookup.find_candidates("울산광역시 남구 산업로 1", get=_fake_get)
        f8.save(project, [f8.candidate_row(found[0])], no_target=False, status="PROPOSED")
        record = project.get_field(f8.SITE_KEY)
        self.assertEqual(record.status, "PROPOSED")
        self.assertTrue(f8.needs(project))

    def test_lookup_failure_falls_back_to_manual_entry(self):
        def boom(url, params, headers):
            raise lookup.requests.ConnectionError("down")

        with patch.dict(os.environ, {lookup.ENV_KEY: "k"}):
            found, message = lookup.find_candidates("울산", get=boom)
        self.assertEqual(found, [])
        self.assertIn("직접 입력", message)

    def test_confirmed_list_derives_checkboxes_and_reaches_the_docx(self):
        project = _project()
        with patch.dict(os.environ, {lookup.ENV_KEY: "k"}):
            found, _ = lookup.find_candidates("주소", get=_fake_get)
        rows = [f8.candidate_row(c) for c in found]
        rows[0]["사업장 경계와 거리(m)"] = 320
        rows[0]["GIS/현장 근거"] = "공식 지도에서 경계부터 측정, 2026-09-25"
        rows[1]["사업장 경계와 거리(m)"] = 430
        rows[1]["GIS/현장 근거"] = "현장 확인 및 지도 측정, 2026-09-25"
        self.assertEqual(f8.save(project, rows, no_target=False, scope_reviewed=True), 2)
        chosen = f8.selected_options(project)
        self.assertEqual(chosen["갑종"], {"교육·연구시설"})
        self.assertEqual(chosen["환경수용체"], {"하천"})
        self.assertEqual(f8.needs(project), [])
        tables = Document(BytesIO(build_cap_baseline_draft(project))).tables
        context = "\n".join(c.text for r in tables[FORM_TABLE_INDEX["8"][0]].rows for c in r.cells)
        listing = "\n".join(c.text for r in tables[FORM_TABLE_INDEX["8"][1]].rows for c in r.cells)
        self.assertIn("☒ 교육·연구시설", context)
        self.assertIn("☐ 의료시설", context)
        self.assertIn("☒ 하천", context)
        self.assertNotIn("☒ 주택", context)
        self.assertIn("한빛초등학교", listing)

    def test_no_target_declaration_needs_evidence(self):
        project = _project()
        f8.save(project, [], no_target=True, evidence="", scope_reviewed=True)
        self.assertTrue(f8.declared_no_target(project))
        self.assertTrue(f8.needs(project))
        f8.save(project, [], no_target=True, evidence="지도 확인 2026-09-19", scope_reviewed=True)
        self.assertEqual(f8.needs(project), [])

    def test_page_is_wired(self):
        self.assertIn(8, __import__("ui.cap_forms_registry", fromlist=["x"]).FORM_NUMBERS)


if __name__ == "__main__":
    unittest.main()
