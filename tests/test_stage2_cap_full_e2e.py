from __future__ import annotations

from io import BytesIO
from types import SimpleNamespace
import unittest
from unittest.mock import patch

from docx import Document
import pandas as pd

from engine.stage2.cap_baseline_docx import build_cap_baseline_draft
from engine.stage2.project import Stage2Project
from engine.stage2.scope_validation import validate_selected_scope


class CAPFullStatutoryFormE2ETests(unittest.TestCase):
    def _project(self) -> Stage2Project:
        p = Stage2Project(
            project_id="S2-CAP-FULL-E2E",
            company_name="한빛정밀화학(주)",
            site_name="울산제1공장",
            cap_required=True,
            cap_group="1군",
            scope_confirmed=True,
            cap_selected=True,
            stage1_source_fingerprint="stage1-e2e-fixture",
        )

        for key, value in (
            ("business.address", "울산광역시 테스트로 1"),
            ("cap.business.representative", "홍길동"),
            ("cap.business.registration_no", "123-45-67890"),
            ("cap.business.contact", "052-260-4000"),
            ("cap.business.unit_plant_name", "울산제1공장"),
            ("cap.business.industrial_complex", "해당 없음"),
            ("cap.business.submission_type", "신규제출"),
            ("cap.business.submission_reason", "최초"),
            ("cap.business.writing_level", "1군"),
            ("cap.business.joint_emergency_plan", "단독제출"),
            ("cap.business.other_system_review", "미해당"),
            ("cap.business.residents_in_overall_range", "예"),
            ("cap.business.recent_accident", "아니오"),
            ("cap.business.writer_name", "최유진"),
            ("cap.business.writer_department", "환경안전팀"),
            ("cap.business.writer_contact", "052-260-4123"),
            ("cap.business.writer_email", "safety@example.com"),
            ("cap.business.report_date", "2026-09-18"),
            ("cap.basic.total_facility_overview", "염소 저장·공급 계통으로 구성"),
            ("cap.basic.unit_facility_overview", "염소 저장탱크와 공급배관으로 구성"),
            ("cap.basic.loading_transport", "입·출하 시설 1기 / 보유 탱크로리 1기"),
            ("process.description", "염소를 저장탱크에서 보관한 뒤 차단밸브와 공급배관을 통해 공정에 공급한다."),
        ):
            p.set_field(key, key, value, "USER_CONFIRMED")

        chemical = {
            "물질명": "염소",
            "CAS 번호": "7782-50-5",
            "함량(%)": 99.9,
            "물리적 상태": "기체",
            "최대보유량": 800,
            "단위": "kg",
            "비중": "2.49 (공기=1)",
            "폭발한계 하한": "해당 없음",
            "폭발한계 상한": "해당 없음",
            "독성구분 항목": "급성독성(흡입)",
            "독성구분": "구분 2",
            "위험노출수준": "ERPG-2 3 ppm",
            "허용농도값": "TWA 0.5 ppm",
            "증기압": "기체",
            "부식성": "예",
            "SDS 파일명": "chlorine_company_SDS.pdf",
            "SDS 개정일": "2026-05-10",
        }
        p.set_field("inventory.chemicals", "화학물질 목록", [chemical], "USER_CONFIRMED")
        p.set_field("cap.chemical.details", "유해화학물질 상세", [chemical], "USER_CONFIRMED")

        facility = {
            "설비번호": "TK-301",
            "설비명": "염소 저장탱크",
            "설비종류": "저장탱크",
            "단위공장·공정": "울산제1공장",
            "취급물질": "염소",
            "CAS 번호": "7782-50-5",
            "용량": 1.2,
            "용량단위": "m3",
            "설계압력": "1.0 MPa",
            "운전압력": "0.5 MPa",
            "설계온도": "40 ℃",
            "운전온도": "25 ℃",
            "재질": "Carbon Steel",
            "최대보유량(kg)": 800,
            "P&ID 번호": "PID-301",
            "최대 연결구 크기(mm)": 25,
        }
        p.set_field("inventory.facilities", "시설·설비 목록", [facility], "USER_CONFIRMED")
        p.set_field("cap.facility.equipment_specs", "장치·설비 상세 명세", [facility], "USER_CONFIRMED")

        p.set_field(
            "cap.chemical.hazard_information",
            "유해화학물질 유해성 정보",
            [{
                "물질명": "염소",
                "CAS 번호": "7782-50-5",
                "인체유해성": "흡입 시 급성 독성 및 호흡기 자극 우려",
                "물리적 위험성": "가압가스 및 산화성 관련 위험",
                "환경유해성": "수생생물에 매우 유독",
                "출처": "회사 제품 SDS 제2·11·12항",
                "선정 사유": "사고시나리오의 독성영향을 대표하는 물질로 선정",
                "SDS 파일명": "chlorine_company_SDS.pdf",
                "SDS 개정일": "2026-05-10",
            }],
            "USER_CONFIRMED",
        )

        p.set_field(
            "cap.site.surrounding_environment",
            "사업장 주변 환경정보",
            [{
                "보호대상 없음 여부": "아니오",
                "보호대상 명칭": "○○초등학교",
                "보호대상 구분": "갑종",
                "세부유형": "교육·연구시설",
                "주소·위치": "울산광역시 테스트로 100",
                "좌표": "35.0010,129.0010",
                "사업장 경계와 거리(m)": 420,
                "GIS/현장 근거": "GIS-SITE-01",
            }],
            "USER_CONFIRMED",
        )

        p.set_field(
            "cap.safety.dike_layout",
            "확산방지설비 현황 및 배치도",
            [{"대상 설비번호": "TK-301", "확산방지설비 종류": "방류벽", "도면번호": "GA-301"}],
            "USER_CONFIRMED",
        )
        p.set_field(
            "cap.safety.dike_calculation",
            "확산방지설비 계산자료",
            [{
                "적용여부": "예",
                "대상 설비번호": "TK-301",
                "설비형태": "저장탱크",
                "확산방지설비 종류": "방류벽",
                "필요용량(m3)": 1.0,
                "필요용량 기준·근거": "회사 확산방지설비 검토자료 DB-301",
                "내부 길이(m)": 2.0,
                "내부 폭(m)": 1.5,
                "유효높이(m)": 0.5,
                "내부 차감용적(m3)": 0.1,
                "직접확인 유효용량(m3)": "",
                "비고": "",
            }],
            "USER_CONFIRMED",
        )

        p.set_field(
            "cap.safety.gas_detection",
            "고정식 유해감지시설 명세",
            [{
                "감지기 번호": "GD-301",
                "설치형태": "고정식",
                "설치위치": "TK-301 방유제 내",
                "검출대상 물질": "염소",
                "작동시간": "30초 이내",
                "측정방식": "전기화학식",
                "경보 설정값": "0.5 ppm",
                "경보 위치": "중앙제어실",
                "연동여부": "예",
                "연동 설비·조치": "HH 경보 시 TK-301 공급 차단",
                "정밀도": "±3% F.S.",
                "유지관리": "월 1회 기능점검·연 1회 교정",
                "비상전원 여부": "예",
                "관련 도면번호": "GA-301",
                "비고": "",
            }],
            "USER_CONFIRMED",
        )

        p.set_field(
            "cap.offsite.scenario_impact_table",
            "사고시나리오 영향평가",
            [{
                "사고시나리오명": "염소 독성누출-1",
                "유해화학물질명": "염소",
                "대상 설비번호": "TK-301",
                "사고유형": "독성누출",
                "장외거리(m)": 180,
                "거주민수": 25,
                "근로자수": 10,
                "갑종 보호대상 수": 1,
                "을종 보호대상 수": 0,
                "환경수용체 수": 1,
                "사고원점 좌표": "35.0000,129.0000",
                "KORA/GIS 근거": "KORA-01",
            }],
            "USER_CONFIRMED",
        )
        p.set_field(
            "cap.offsite.overall_impact_summary",
            "총괄영향범위 요약",
            [{
                "총괄영향범위 산출방법": "KORA 사고시나리오 결과와 GIS 공간중첩",
                "총괄영향범위 결과 요약": "총괄영향범위는 승인 GIS 결과도면 기준",
                "GIS/KORA 근거": "KORA-ALL-01",
                "총괄영향범위 내 거주민수": 65,
                "총괄영향범위 내 근로자수": 25,
                "보호대상 없음 여부": "아니오",
            }],
            "USER_CONFIRMED",
        )
        p.set_field(
            "cap.offsite.population_and_protected_targets",
            "총괄영향범위 보호대상",
            [{
                "보호대상 명칭": "○○초등학교",
                "보호대상 구분": "갑종",
                "세부유형": "교육·연구시설",
                "주소·위치": "울산광역시 테스트로 100",
                "좌표": "35.0010,129.0010",
                "사업장 경계와 거리(m)": 420,
                "인원수": 350,
                "GIS 근거": "GIS-PT-01",
            }],
            "USER_CONFIRMED",
        )
        p.set_field(
            "documents.kora_impact_result",
            "KORA/GIS 결과파일",
            {"file_name": "KORA_impact_result.pdf", "stored_path": "uploads/KORA_impact_result.pdf"},
            "USER_CONFIRMED",
        )
        p.set_field(
            "cap.offsite.scenario_frequency",
            "사고시나리오 시설빈도",
            [{
                "사고시나리오명": "염소 독성누출-1",
                "고압용기파열": 1,
                "배관파열": 0,
                "배관누출": 2,
                "상압 탱크 파열 및 누출": 0,
                "플랜지 등의 가스켓 파손": 1,
                "펌프/컴프레서 누출": 0,
                "안전밸브 오작동 및 조기개방": 0,
                "냉각수 손실": 0,
                "입/출하 시설 누출 사고": 0,
                "외부화재": 0,
                "개수 산정근거": "PID-301 및 설비목록",
                "수동적 완화장치": "방류벽",
                "능동적 완화장치": "",
                "안전성확보설비 증빙": "GA-301",
            }],
            "USER_CONFIRMED",
        )

        for key, value in (
            ("cap.prevention.emergency_contact_system", "비상연락망으로 사고상황을 전파한다."),
            ("cap.prevention.emergency_org_chart", "공장장을 비상대응 책임자로 하는 조직을 운영한다."),
            ("cap.prevention.emergency_roles", "상황전파·초기대응·대피 임무를 분담한다."),
            ("cap.internal.shutdown_authority", "공정반장"),
            ("cap.internal.shutdown_procedure", "원료차단 후 관련 설비를 정지한다."),
            ("cap.internal.response_personnel", "방재반 8명"),
            ("cap.internal.response_equipment", "보호복·공기호흡기·흡착재"),
            ("cap.internal.response_operations", "누출원 차단과 확산방지 조치를 수행한다."),
            ("cap.internal.communication_system", "비상방송과 무전기로 내부에 전파한다."),
            ("cap.internal.facility_response_plan", "염소 누출 시 원격차단 후 방재반이 접근한다."),
            ("cap.internal.investigation_plan", "사고조사팀이 원인을 조사한다."),
            ("cap.internal.recurrence_prevention", "원인별 개선조치를 추적관리한다."),
            ("cap.internal.recovery_plan", "오염제거와 설비복구 후 재가동을 승인한다."),
            ("cap.external.communication_plan", "지자체·소방·주민에게 사고정보를 제공한다."),
            ("cap.external.stakeholders", "관할 소방서·지자체·인근 사업장"),
            ("cap.external.communication_schedule", "평상시 연 1회, 사고 시 즉시"),
            ("cap.external.mutual_aid_contacts", "관할 소방서와 인근 사업장 연락망"),
            ("cap.external.resource_support", "소방차·방재인력 상호지원"),
            ("cap.external.joint_drill_plan", "연 1회 합동훈련"),
            ("cap.external.warning_system", "재난문자·방송으로 주민경보"),
            ("cap.external.evacuation_routes", "남문 방향 대피로"),
            ("cap.external.shelters", "○○체육관"),
            ("cap.external.medical_contacts", "○○병원 응급실"),
            ("cap.external.notice_targets", "총괄영향범위 내 주민"),
            ("cap.external.notice_method", "서면·홈페이지 고지"),
            ("cap.external.notice_content", "취급물질·사고위험·행동요령"),
        ):
            p.set_field(key, key, value, "USER_CONFIRMED")
        return p

    @staticmethod
    def _screen():
        return SimpleNamespace(
            ready=True,
            blockers=(),
            messages=("fixture legal screen",),
            legal_hits=[{
                "row_no": 1,
                "source_key": "CAP_QTY_APP3",
                "lower_quantity_ton": 0.1,
                "upper_quantity_ton": 1.0,
                "cas": "7782-50-5",
                "legal_substance": "염소",
                "item_no": "fixture-1",
            }],
        )

    @staticmethod
    def _law_tables():
        return {
            "CAP_QTY_APP2": pd.DataFrame([{
                "direct_cas": "7782-50-5",
                "hazard_category": "인체급성유해성",
                "designation_id": "97-1-1",
                "source_key": "CAP_QTY_APP2",
                "active": True,
                "content_threshold_pct": 0,
            }])
        }

    @staticmethod
    def _table_text(doc: Document, *indices: int) -> str:
        return "\n".join(
            cell.text
            for index in indices
            for row in doc.tables[index].rows
            for cell in row.cells
        )

    @patch("engine.stage2.cap_risk_engine.approved_source_is_current", return_value=True)
    @patch("engine.stage2.cap_chemical_legal.screen_cap_scope_candidates_from_tables", return_value=pd.DataFrame())
    @patch("engine.stage2.cap_chemical_legal.load_approved_scope_tables")
    @patch("engine.stage2.cap_form1_engine.screen_facility_stage")
    def test_fully_confirmed_project_populates_core_forms_without_program_blanks(
        self, screen_mock, law_tables_mock, _scope_mock, _current
    ):
        screen_mock.return_value = self._screen()
        law_tables_mock.return_value = (self._law_tables(), ())
        p = self._project()

        out = build_cap_baseline_draft(p)
        doc = Document(BytesIO(out))
        self.assertEqual(len(doc.tables), 44)

        form1 = self._table_text(doc, 6, 7, 8)
        self.assertIn("염소", form1)
        self.assertIn("0.8", form1)
        self.assertIn("0.1", form1)
        self.assertIn("1", form1)
        self.assertIn("1군", form1)
        self.assertIn("사고대비물질", form1)

        form3 = self._table_text(doc, 12)
        self.assertIn("☒ 신규제출", form3)
        self.assertIn("☒ 최초", form3)
        self.assertIn("☒ 1군", form3)
        self.assertIn("☒ 단독제출", form3)
        self.assertIn("환경안전팀 최유진", form3)
        self.assertIn("052-260-4123", form3)
        self.assertIn("safety@example.com", form3)
        self.assertEqual(form3.count("☒ 신규제출"), 1)
        self.assertEqual(form3.count("☒ 최초"), 1)

        for index in (13, 14):
            overview = self._table_text(doc, index)
            self.assertIn("☒ 저장탱크 (1)기", overview)
            self.assertEqual(overview.count("☒ 저장탱크 (1)기"), 1)
            self.assertNotIn("반응기 1기 / 저장탱크", overview)

        form6 = self._table_text(doc, 15)
        self.assertIn("인체급성유해성물질", form6)
        self.assertIn("사고대비물질", form6)
        self.assertIn("97-1-1", form6)
        self.assertIn("ERPG-2 3 ppm", form6)
        self.assertIn("TWA 0.5 ppm", form6)
        self.assertIn("2.49 (공기=1)", form6)

        form7 = self._table_text(doc, 16)
        self.assertIn("97-1-1", form7)
        self.assertIn("0.8 ton", form7)
        self.assertIn("흡입 시 급성 독성", form7)
        self.assertIn("chlorine_company_SDS.pdf", form7)
        self.assertIn("사고시나리오의 독성영향", form7)

        form8 = self._table_text(doc, 17, 18)
        self.assertIn("○○초등학교", form8)
        self.assertIn("교육·연구시설", form8)
        self.assertIn("420", form8)

        form9 = self._table_text(doc, 19)
        self.assertIn("TK-301", form9)
        self.assertIn("25", form9)
        self.assertIn("1.2", form9)
        self.assertIn("0.8", form9)

        form10 = self._table_text(doc, 20)
        self.assertIn("TK-301", form10)
        self.assertIn("1.4", form10)
        self.assertIn("적정", form10)

        form11 = self._table_text(doc, 21)
        self.assertIn("GD-301", form11)
        self.assertIn("0.5 ppm", form11)
        self.assertIn("전기화학식", form11)

        form12 = self._table_text(doc, 22, 23)
        self.assertIn("염소 독성누출-1", form12)
        self.assertIn("독성누출", form12)
        self.assertIn("180", form12)
        self.assertIn("35.0000,129.0000", form12)

        form13_14 = self._table_text(doc, 24, 25, 26, 27, 28)
        self.assertIn("KORA-ALL-01", form13_14)
        self.assertIn("○○초등학교", form13_14)
        self.assertIn("배관누출", form13_14)
        self.assertIn("2", form13_14)

        form15 = self._table_text(doc, 29, 30, 31)
        self.assertIn("염소 독성누출-1", form15)
        self.assertIn("180", form15)
        self.assertIn("35", form15)

        form16 = self._table_text(doc, *range(32, 44))
        self.assertIn("2026-09-18", form16)
        self.assertIn("환경안전팀 최유진", form16)
        self.assertIn("독성누출", form16)
        self.assertIn("비상연락망으로 사고상황을 전파한다.", form16)
        self.assertIn("○○병원 응급실", form16)

    @patch("engine.stage2.cap_risk_engine.approved_source_is_current", return_value=True)
    @patch("engine.stage2.cap_chemical_legal.screen_cap_scope_candidates_from_tables", return_value=pd.DataFrame())
    @patch("engine.stage2.cap_chemical_legal.load_approved_scope_tables")
    @patch("engine.stage2.cap_form1_engine.screen_facility_stage")
    def test_core_statutory_form_validation_has_no_hold_when_all_facts_are_confirmed(
        self, screen_mock, law_tables_mock, _scope_mock, _current
    ):
        screen_mock.return_value = self._screen()
        law_tables_mock.return_value = (self._law_tables(), ())
        report = validate_selected_scope(self._project())

        form_issues = [
            issue
            for issue in report.issues
            if issue.code.startswith("CAP-FORM")
        ]
        holds = [issue for issue in form_issues if issue.status == "HOLD"]

        self.assertTrue(form_issues)
        self.assertEqual(
            holds,
            [],
            msg="\n".join(f"{issue.code}: {issue.message}" for issue in holds),
        )


if __name__ == "__main__":
    unittest.main()
