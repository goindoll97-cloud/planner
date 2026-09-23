from __future__ import annotations

from io import BytesIO
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

from openpyxl import Workbook

from engine.stage2 import cap_chemical_upload as up
from engine.stage2 import cap_chemical_workspace as chem
from tests.test_stage2_cap_form1_engine import CAPForm1EngineTests


def _xlsx(rows, before=()):
    workbook = Workbook()
    sheet = workbook.active
    for line in before:
        sheet.append(line)
    for line in rows:
        sheet.append(line)
    out = BytesIO()
    workbook.save(out)
    return out.getvalue()


HEADER = ["No", "품명", "CAS번호", "농도(%)", "최대보유량(kg)", "비고"]
ROWS = [
    [1, "톨루엔", "108-88-3", "99.5", "12500", "저장탱크"],
    [2, "염소", "7782-50-5", "100", "3000", ""],
    [3, "아세톤", "67-64-1", "95~99", "800", ""],
    [4, "오타물질", "108-88-4", "100", "10", ""],  # 검산 숫자 오류
    [5, "혼합제품", "67-64-1, 108-88-3", "30", "", ""],  # CAS 여러 개
    [6, "톨루엔", "108-88-3", "99.5", "1", ""],  # 파일 안 중복
]


class CasTests(unittest.TestCase):
    def test_check_digit(self):
        for cas in ("108-88-3", "7782-50-5", "67-64-1", "64-17-5", "50-00-0"):
            self.assertTrue(up.cas_valid(cas), cas)
        for cas in ("108-88-4", "7782-50-6", "1234-56-7", "abc", "108-88"):
            self.assertFalse(up.cas_valid(cas), cas)


class ParseTests(unittest.TestCase):
    def test_header_aliases_are_matched_and_units_come_from_the_header(self):
        parsed = up.parse(_xlsx([HEADER, *ROWS], before=[["회사 물질 목록 2026"], []]), "list.xlsx")
        self.assertEqual(parsed.mapping["제품명"], "품명")
        self.assertEqual(parsed.mapping["CAS No."], "CAS번호")
        self.assertEqual(parsed.mapping["함량(%)"], "농도(%)")
        self.assertEqual(parsed.mapping["최대 동시보유량(ton)"], "최대보유량(kg)")
        self.assertEqual(len(parsed.sha256), 64)
        rows = up.normalize(parsed, parsed.mapping)
        self.assertEqual(rows[0]["최대 동시보유량(ton)"], "12.5")  # 12500 kg -> 12.5 ton
        self.assertEqual(rows[2]["함량(%)"], "99")  # 범위는 상한값
        self.assertIn("범위", rows[2]["메모"])

    def test_annual_amount_columns_are_not_mistaken_for_max_holding(self):
        mapping = up.guess_mapping(["품명", "CAS", "연간취급량(ton)", "최대보유량(ton)"])
        self.assertEqual(mapping["최대 동시보유량(ton)"], "최대보유량(ton)")

    def test_csv_in_cp949_and_utf8_are_both_read(self):
        text = "품명,CAS No.,함량\n톨루엔,108-88-3,99.5\n염소,7782-50-5,100\n"
        for encoding in ("utf-8-sig", "cp949"):
            parsed = up.parse(text.encode(encoding), "list.csv")
            self.assertEqual(len(parsed.frame), 2, encoding)
            self.assertEqual(parsed.mapping["제품명"], "품명")

    def test_unreadable_and_old_excel_files_give_a_clear_error(self):
        with self.assertRaises(ValueError):
            up.parse(b"", "a.xlsx")
        with self.assertRaises(ValueError):
            up.parse(b"x", "a.xls")

    def test_unknown_or_unit_less_amounts_are_left_blank_with_a_reason(self):
        value, note = up._to_ton("5", "", "보유량")
        self.assertEqual(value, "")
        self.assertIn("단위", note)
        value, note = up._to_ton("500", "L", "")
        self.assertEqual(value, "")
        self.assertIn("비중", note)
        self.assertEqual(up._to_ton("2", "톤", "")[0], "2")


class CheckTests(unittest.TestCase):
    def test_errors_warnings_and_duplicates_are_reported_per_row(self):
        parsed = up.parse(_xlsx([HEADER, *ROWS]), "list.xlsx")
        checked = up.check_rows(up.normalize(parsed, parsed.mapping))
        rows = checked.rows
        self.assertTrue(rows[0]["_ok"] and rows[1]["_ok"] and rows[2]["_ok"])
        self.assertFalse(rows[3]["_ok"])
        self.assertIn("검산", rows[3]["확인"])
        self.assertTrue(rows[4]["_ok"])  # 혼합물은 경고만
        self.assertIn("혼합물", rows[4]["확인"])
        self.assertFalse(rows[5]["_ok"])
        self.assertIn("중복", rows[5]["확인"])
        self.assertEqual(checked.errors, 2)

    def test_explicit_single_substance_requires_cas(self):
        checked = up.check_rows([{
            "제품명": "CAS 없는 단일물질", "혼합물 여부": "N", "CAS No.": "",
            "최대 제조·사용량": "1", "최대 저장량": "2", "단위": "ton",
        }])
        self.assertFalse(checked.rows[0]["_ok"])
        self.assertIn("CAS No.", checked.rows[0]["확인"])

    def test_explicit_mixture_can_be_added_without_product_cas(self):
        checked = up.check_rows([{
            "제품명": "세척제 A", "혼합물 여부": "Y", "CAS No.": "",
            "최대 제조·사용량": "1", "최대 저장량": "2", "단위": "ton",
        }])
        self.assertTrue(checked.rows[0]["_ok"])

    def test_blank_name_after_single_is_not_silently_attached_as_a_mixture_component(self):
        products = up.group_products([
            {"제품명": "톨루엔", "혼합물 여부": "N", "CAS No.": "108-88-3", "함량(%)": ""},
            {"제품명": "", "혼합물 여부": "", "CAS No.": "67-64-1", "함량(%)": "100"},
        ])
        self.assertEqual(len(products), 2)
        self.assertEqual(products[0]["혼합물 여부"], "N")
        self.assertTrue(products[1]["_component_problems"])

    def test_rows_already_in_the_list_are_skipped_not_overwritten(self):
        checked = up.check_rows([{"제품명": "톨루엔", "CAS No.": "108-88-3", "함량(%)": "99", "최대 동시보유량(ton)": ""}],
                                existing_cas={"108-88-3"})
        self.assertFalse(checked.rows[0]["_ok"])
        self.assertIn("이미 목록에", checked.rows[0]["확인"])

    def test_non_numeric_content_is_an_error(self):
        checked = up.check_rows([{"제품명": "물", "CAS No.": "7732-18-5", "함량(%)": "많이", "최대 동시보유량(ton)": ""}])
        self.assertFalse(checked.rows[0]["_ok"])


class TemplateTests(unittest.TestCase):
    def test_blank_template_round_trips_and_examples_are_only_on_the_guide_sheet(self):
        from openpyxl import load_workbook

        book = load_workbook(BytesIO(up.blank_template()))
        first = book["물질 목록"]
        self.assertEqual([c.value for c in first[1]], [
            "제품명", "단일물질/혼합물", "CAS No.", "최대 제조·사용량", "최대 저장량", "단위"])
        self.assertEqual(first.max_row, 1)  # 빈 양식: 예시 행이 실수로 올라가지 않는다
        self.assertTrue(any("(예시)" in str(c.value) for row in book["작성 안내"].iter_rows() for c in row))
        parsed = up.parse(up.blank_template(), "template.xlsx")
        self.assertEqual(parsed.mapping["제품명"], "제품명")
        self.assertEqual(parsed.mapping["최대 제조·사용량"], "최대 제조·사용량")
        self.assertEqual(parsed.mapping["최대 저장량"], "최대 저장량")
        self.assertEqual(parsed.mapping["단위"], "단위")


class MergeTests(unittest.TestCase):
    def test_only_valid_rows_are_added_and_existing_rows_are_kept(self):
        project = CAPForm1EngineTests()._project()
        before = [dict(r) for r in chem._rows(project)[1]]
        parsed = up.parse(_xlsx([HEADER, *ROWS]), "list.xlsx")
        with tempfile.TemporaryDirectory() as tmp, patch("engine.stage2.storage.DEFAULT_ROOT", Path(tmp)):
            added, skipped = up.add_to_project(project, up.normalize(parsed, parsed.mapping), file_name="list.xlsx",
                                               sha256=parsed.sha256)
            self.assertEqual(added, 3)  # 톨루엔, 아세톤, 혼합제품 (염소는 이미 목록에 있어 건너뜀)
            self.assertEqual(skipped, 3)  # 이미 있음(염소), 오타물질, 파일 안 중복
            rows = chem._rows(project)[1]
            self.assertEqual(rows[:len(before)], before)
            added_rows = {r["물질명"]: r for r in rows[len(before):]}
            self.assertEqual(set(added_rows), {"톨루엔", "아세톤", "혼합제품"})
            self.assertEqual(added_rows["톨루엔"]["최대 동시보유량(알면 입력)"], 12.5)
            self.assertEqual(added_rows["톨루엔"]["수량 단위"], "ton")
            self.assertTrue(any("list.xlsx" in n for n in project.notes))
            again = up.add_to_project(project, up.normalize(parsed, parsed.mapping), file_name="list.xlsx",
                                      sha256=parsed.sha256)
            self.assertEqual(again[0], 0)  # 같은 파일을 다시 올려도 중복 추가되지 않는다

    def test_details_key_is_kept_in_sync_so_new_rows_are_not_hidden(self):
        project = CAPForm1EngineTests()._project()
        rows = chem._rows(project)[1]
        project.set_field(chem.DETAILS_KEY, "상세", rows, "USER_CONFIRMED")
        self.assertEqual(chem._rows(project)[0], chem.DETAILS_KEY)
        parsed = up.parse(_xlsx([HEADER, ROWS[0]]), "l.xlsx")
        up.add_to_project(project, up.normalize(parsed, parsed.mapping), file_name="l.xlsx", sha256="a" * 64)
        self.assertIn("톨루엔", [r.get("물질명") for r in chem._rows(project)[1]])
        self.assertIn("톨루엔", [r.get("물질명") for r in project.get_field(chem.INVENTORY_KEY).value])


class ReplaceTests(unittest.TestCase):
    def test_old_rows_are_gone_and_only_the_new_file_remains(self):
        from engine.stage2.cap_judgement import CHEM_INPUTS_KEY, MIXTURE_COMPONENTS_KEY

        project = CAPForm1EngineTests()._project()  # 염소를 이미 물질 목록에 갖고 있다
        project.set_field(MIXTURE_COMPONENTS_KEY, "혼합물 구성성분",
                          [{"제품목록행번호": 1, "CAS No.": "7782-50-5", "함량(%)": "100"}], "USER_CONFIRMED")
        project.set_field(CHEM_INPUTS_KEY, "물질별 확인값", [{"함량(%)": "100"}], "USER_CONFIRMED")

        parsed = up.parse(_xlsx([HEADER, *ROWS]), "list.xlsx")
        added, skipped = up.replace_project(project, up.normalize(parsed, parsed.mapping), file_name="list.xlsx",
                                            sha256=parsed.sha256)
        self.assertEqual(added, 4)  # 톨루엔, 염소, 아세톤, 혼합제품(오타물질·파일 안 중복은 제외)
        names = [r["물질명"] for r in chem._rows(project)[1]]
        self.assertEqual(sorted(names), sorted(["톨루엔", "염소", "아세톤", "혼합제품"]))
        # 옛 염소 행 번호를 가리키던 자료는 새 목록 기준으로 다시 만들어지지 않는 한 남겨 두지 않는다.
        self.assertEqual(project.get_field(MIXTURE_COMPONENTS_KEY).value, [])
        self.assertEqual(project.get_field(CHEM_INPUTS_KEY).value, [])

    def test_details_and_inventory_keys_both_end_up_with_the_same_replaced_list(self):
        project = CAPForm1EngineTests()._project()
        project.set_field(chem.DETAILS_KEY, "상세", chem._rows(project)[1], "USER_CONFIRMED")
        parsed = up.parse(_xlsx([HEADER, ROWS[0]]), "l.xlsx")
        up.replace_project(project, up.normalize(parsed, parsed.mapping), file_name="l.xlsx", sha256="a" * 64)
        self.assertEqual([r["물질명"] for r in project.get_field(chem.DETAILS_KEY).value], ["톨루엔"])
        self.assertEqual([r["물질명"] for r in project.get_field(chem.INVENTORY_KEY).value], ["톨루엔"])

    def test_uploading_the_same_file_twice_still_leaves_exactly_one_copy(self):
        project = CAPForm1EngineTests()._project()
        parsed = up.parse(_xlsx([HEADER, ROWS[0]]), "l.xlsx")
        rows = up.normalize(parsed, parsed.mapping)
        up.replace_project(project, rows, file_name="l.xlsx", sha256="a" * 64)
        up.replace_project(project, rows, file_name="l.xlsx", sha256="a" * 64)
        self.assertEqual([r["물질명"] for r in chem._rows(project)[1]], ["톨루엔"])


if __name__ == "__main__":
    unittest.main()




if __name__ == "__main__":
    unittest.main()
