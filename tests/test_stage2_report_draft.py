from __future__ import annotations

from io import BytesIO
from pathlib import Path
import unittest
from zipfile import ZipFile

from docx import Document

from engine.stage2.project import EvidenceRef, Stage2Project
from engine.stage2.report_draft import (
    build_draft_bundle,
    build_report_draft,
    draft_filename,
    report_generation_status,
)


class Stage2ReportDraftTests(unittest.TestCase):
    def _project(self, *, psm=True, cap=True, cap_group="1군") -> Stage2Project:
        project = Stage2Project(
            project_id="S2-TEST-REPORT",
            company_name="테스트정밀화학",
            site_name="울산공장",
            psm_required=psm,
            cap_required=cap,
            cap_group=cap_group if cap else "",
            scope_confirmed=True,
            psm_selected=psm,
            cap_selected=cap,
        )
        project.set_field("business.company_name", "회사명", "테스트정밀화학", "VERIFIED")
        project.set_field("business.site_name", "사업장명", "울산공장", "VERIFIED")
        project.set_field("business.address", "사업장 소재지", "울산광역시 테스트로 1", "VERIFIED")
        project.set_field(
            "inventory.chemicals",
            "화학물질 목록",
            [
                {
                    "물질명": "톨루엔",
                    "CAS 번호": "108-88-3",
                    "물질구분": "사고대비",
                    "물리적 상태": "액체",
                    "함량(%)": 99.5,
                    "비중": 0.87,
                    "폭발한계 하한": 1.2,
                    "폭발한계 상한": 7.1,
                    "허용농도값": "50 ppm",
                    "최대보유량": 15000,
                    "단위": "kg",
                }
            ],
            "VERIFIED",
        )
        project.set_field(
            "inventory.facilities",
            "시설·설비 목록",
            [
                {
                    "설비번호": "TK-101",
                    "설비명": "톨루엔 저장탱크",
                    "설비종류": "저장탱크",
                    "취급물질": "톨루엔",
                    "용량": 20,
                    "설계압력": "0.49 MPa",
                    "운전압력": "0.20 MPa",
                    "설계온도": "60 ℃",
                    "운전온도": "25 ℃",
                    "최대보유량(kg)": 15000,
                    "P&ID 번호": "PID-101",
                }
            ],
            "VERIFIED",
        )
        project.set_field(
            "process.description",
            "공정설명서",
            "원료를 저장탱크에서 혼합공정으로 이송한 뒤 제품을 생산한다.",
            "USER_CONFIRMED",
        )
        project.set_field(
            "psm.psi.machinery_list",
            "동력기계 목록",
            [{"기계번호": "P-101", "기계명": "이송펌프", "명세": "10 m3/h", "주요재질": "SUS304"}],
            "USER_CONFIRMED",
        )
        project.set_field(
            "psm.psi.equipment_specs",
            "장치 및 설비명세",
            [{"설비번호": "TK-101", "설비명": "톨루엔 저장탱크", "내용물": "톨루엔", "설계압력": "0.49 MPa"}],
            "USER_CONFIRMED",
        )
        project.set_field(
            "psm.psi.relief_device_specs",
            "안전밸브 및 파열판 명세",
            [{"안전밸브·파열판 번호": "PSV-101", "보호대상 설비번호": "TK-101", "설정압력": "0.45 MPa"}],
            "USER_CONFIRMED",
        )
        project.set_field(
            "cap.facility.equipment_specs",
            "장치·설비 상세 명세",
            [{"설비번호": "TK-101", "설비명": "톨루엔 저장탱크", "취급물질": "톨루엔", "설계압력": "0.49 MPa"}],
            "USER_CONFIRMED",
        )
        project.set_field(
            "cap.safety.gas_detection",
            "고정식 유해감지시설 명세 및 배치도",
            [{"감지기 번호": "GD-101", "검출대상 물질": "톨루엔", "설치위치": "TK-101 주변", "경보 설정값": "20 ppm"}],
            "USER_CONFIRMED",
        )
        project.set_field(
            "documents.pfd",
            "공정흐름도(PFD)",
            {"file_name": "PFD_Rev3.pdf", "reference_no": "PFD-001", "revision": "Rev.3"},
            "HOLD",
            evidence=[EvidenceRef(source_type="COMPANY_EVIDENCE", source_name="PFD_Rev3.pdf", sha256="a" * 64)],
            note="도면은 접수되었으나 내용 확인 전",
        )
        if cap:
            project.set_field("cap.business.representative", "대표자 성명", "김테스트", "USER_CONFIRMED")
            project.set_field("cap.business.writing_level", "작성수준", f"{cap_group} 사업장", "USER_CONFIRMED")
            project.set_field("cap.business.submission_type", "제출구분", "신규제출", "USER_CONFIRMED")
            project.set_field("cap.internal.shutdown_authority", "가동중지 권한", "비상상황 시 당직 책임자가 가동중지를 지시한다.", "USER_CONFIRMED")
            if cap_group == "1군":
                project.set_field("cap.external.notice_method", "지역사회 고지 방법", "화학물질종합정보시스템 등록", "USER_CONFIRMED")
        return project

    @staticmethod
    def _doc_text(data: bytes) -> str:
        doc = Document(BytesIO(data))
        parts = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                parts.extend(cell.text for cell in row.cells)
        return "\n".join(parts)

    def test_psm_draft_uses_statutory_annex_forms_and_fixed_columns(self):
        project = self._project(psm=True, cap=False)
        data = build_report_draft(project, "PSM")
        text = self._doc_text(data)

        self.assertIn("법정서식 기반 검토용 작성본", text)
        self.assertIn("별지 제12호서식", text)
        self.assertIn("별지 제13호서식", text)
        self.assertIn("별지 제14호서식", text)
        self.assertIn("별지 제15호서식", text)
        self.assertIn("별지 제17호서식", text)
        self.assertIn("유해·위험물질 목록", text)
        self.assertIn("동력기계 번호", text)
        self.assertIn("장치번호", text)
        self.assertIn("보호기기 번호", text)
        self.assertIn("톨루엔", text)
        self.assertIn("108-88-3", text)
        self.assertIn("TK-101", text)
        self.assertIn("P-101", text)
        self.assertIn("PSV-101", text)
        self.assertIn("PFD_Rev3.pdf", text)
        self.assertIn("담당자 확인 필요", text)
        self.assertIn("[확인 필요]", text)
        self.assertIn("검토 참고사항 (법정서식 외)", text)

    def test_cap_draft_uses_statutory_forms_and_fixed_columns(self):
        project = self._project(psm=False, cap=True, cap_group="1군")
        data = build_report_draft(project, "CAP")
        text = self._doc_text(data)

        self.assertIn("별지 제1호서식", text)
        self.assertIn("사업장의 작성수준 구분", text)
        self.assertIn("별지 제3호서식", text)
        self.assertIn("사업장 일반정보", text)
        self.assertIn("별지 제6호서식", text)
        self.assertIn("유해화학물질 목록 및 명세", text)
        self.assertIn("화학물질식별번호(CAS 번호)", text)
        self.assertIn("별지 제9호서식", text)
        self.assertIn("연결구 크기(mm)", text)
        self.assertIn("별지 제11호서식", text)
        self.assertIn("경보설정값", text)
        self.assertIn("GD-101", text)
        self.assertIn("외부 비상대응계획", text)
        self.assertIn("지역사회 고지계획", text)

    def test_cap_group2_draft_does_not_add_external_emergency_section(self):
        project = self._project(psm=False, cap=True, cap_group="2군")
        data = build_report_draft(project, "CAP")
        doc = Document(BytesIO(data))
        headings = [p.text for p in doc.paragraphs if p.style.name.startswith("Heading")]
        text = self._doc_text(data)

        self.assertIn("기본정보", headings)
        self.assertIn("시설정보", headings)
        self.assertIn("내부 비상대응계획", headings)
        self.assertNotIn("외부 비상대응계획", headings)
        self.assertNotIn("지역사회 고지계획", headings)
        self.assertIn("2군 사업장은 외부 비상대응계획 내용을 생략할 수 있으므로", text)

    def test_report_status_is_fail_closed_but_statutory_draft_is_still_buildable(self):
        project = self._project(psm=True, cap=False)
        status = report_generation_status(project, "PSM")
        self.assertFalse(status.final_ready)
        self.assertEqual(status.state, "HOLD")
        self.assertGreater(status.unresolved_n, 0)
        self.assertGreater(len(build_report_draft(project, "PSM")), 1000)

    def test_unselected_report_cannot_be_generated(self):
        project = self._project(psm=True, cap=False)
        with self.assertRaisesRegex(ValueError, "현재 작성범위"):
            build_report_draft(project, "CAP")

    def test_filename_and_bundle_identify_statutory_form_output(self):
        project = self._project(psm=True, cap=True)
        self.assertIn("법정서식_검토용_초안", draft_filename(project, "PSM"))
        data = build_draft_bundle(project)
        with ZipFile(BytesIO(data)) as archive:
            names = archive.namelist()
            self.assertIn(draft_filename(project, "PSM"), names)
            self.assertIn(draft_filename(project, "CAP"), names)
            self.assertIn("생성상태.txt", names)
            manifest = archive.read("생성상태.txt").decode("utf-8")
            self.assertIn("법정서식 기반 검토용 작성본", manifest)
            self.assertIn("공정안전보고서", manifest)
            self.assertIn("화학사고예방관리계획서", manifest)

    def test_report_page_describes_statutory_form_based_output(self):
        source = Path("ui/stage2_review_page.py").read_text(encoding="utf-8")
        self.assertIn('st.title("📝 5. 보고서 작성")', source)
        self.assertIn("보고서 초안 내려받기", source)
        self.assertIn("build_report_draft", source)
        self.assertIn("기본 초안 다운로드", source)
        self.assertIn("AI 문장 다듬기 포함 초안", source)
        self.assertIn("법제처 원본서식 HWPX", source)
        self.assertIn("최종 제출자료와 대조한 뒤 제출본으로 확정", source)
        self.assertNotIn('"감사·검토자료"', source)


if __name__ == "__main__":
    unittest.main()
