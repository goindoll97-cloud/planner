from __future__ import annotations

from io import BytesIO
import unittest

from openpyxl import load_workbook

from engine.stage2.integrated_workbook import apply_integrated_authoring_workbook
from engine.stage2.project import create_project_from_stage1_snapshot
from engine.stage2.workbook_enhancements import build_enhanced_integrated_authoring_workbook


class PSMCoreWorkbookTests(unittest.TestCase):
    def _project(self):
        snapshot = {
            "source_fingerprint": "c" * 64,
            "business": {
                "회사명": "PSM테스트",
                "사업장명": "제1공장",
                "사업장 소재지": "테스트시 산업로 1",
            },
            "documents": {},
            "chemicals": [{
                "물질명": "톨루엔",
                "CAS 번호": "108-88-3",
                "함량(%)": 99.5,
                "최대보유량": 15000,
                "단위": "kg",
            }],
            "facilities": [{
                "설비번호": "TK-101",
                "설비명": "톨루엔 저장탱크",
                "설비종류": "저장탱크",
                "취급물질": "톨루엔",
            }],
            "decision": {
                "psm_status": "공정안전보고서 제출 대상",
                "cap_status": "비대상",
            },
        }
        project = create_project_from_stage1_snapshot(snapshot)
        project.set_authoring_scope(psm_selected=True, cap_selected=False)
        return project

    @staticmethod
    def _headers(ws):
        return {
            str(cell.value or "").strip(): cell.column
            for cell in ws[4]
            if str(cell.value or "").strip()
        }

    def test_psm_only_workbook_exposes_core_statutory_columns_and_prefills_stage1_identity(self):
        project = self._project()
        wb = load_workbook(BytesIO(build_enhanced_integrated_authoring_workbook(project, example=False)))

        chemical = wb["02_화학물질정보"]
        headers = self._headers(chemical)
        for name in (
            "분자식", "일일사용량", "노출기준", "독성치", "인화점",
            "발화점", "이상반응 유무",
        ):
            self.assertIn(name, headers)
        self.assertEqual(chemical.cell(5, headers["물질명"]).value, "톨루엔")
        self.assertEqual(chemical.cell(5, headers["CAS 번호"]).value, "108-88-3")

        equipment = wb["03_설비정보"]
        equipment_headers = self._headers(equipment)
        for name in (
            "부속품재질", "개스킷재질", "용접효율", "계산두께",
            "부식여유", "사용두께", "후열처리 여부", "비파괴검사율",
        ):
            self.assertIn(name, equipment_headers)
        self.assertEqual(equipment.cell(5, equipment_headers["설비번호"]).value, "TK-101")

        relief = wb["04_안전밸브_파열판"]
        relief_headers = self._headers(relief)
        for name in (
            "정격용량", "노즐크기 입구", "노즐크기 출구",
            "보호기기 운전압력", "보호기기 설계압력",
            "몸체재질", "TRIM 재질", "정밀도", "배출원인",
        ):
            self.assertIn(name, relief_headers)

    def test_psm_chemical_enrichment_import_does_not_overwrite_stage1_inventory(self):
        project = self._project()
        before = [dict(row) for row in project.get_field("inventory.chemicals").value]

        raw = build_enhanced_integrated_authoring_workbook(project, example=False)
        wb = load_workbook(BytesIO(raw))
        ws = wb["02_화학물질정보"]
        headers = self._headers(ws)
        values = {
            "분자식": "C7H8",
            "일일사용량": "2500 kg/day",
            "폭발한계 하한": "1.2 vol%",
            "폭발한계 상한": "7.1 vol%",
            "노출기준": "TWA 50 ppm",
            "독성치": "회사 SDS 확인값",
            "인화점": "4 ℃",
            "발화점": "480 ℃",
            "증기압": "28.4 mmHg (25℃)",
            "부식성": "해당 없음",
            "이상반응 유무": "해당 없음",
        }
        for header, value in values.items():
            ws.cell(5, headers[header], value)

        out = BytesIO()
        wb.save(out)
        result = apply_integrated_authoring_workbook(project, out.getvalue())

        self.assertGreaterEqual(result.table_fields, 1)
        self.assertEqual(project.get_field("inventory.chemicals").value, before)
        detail = project.get_field("psm.psi.chemical_details")
        self.assertIsNotNone(detail)
        self.assertEqual(detail.status, "USER_CONFIRMED")
        self.assertEqual(detail.value[0]["분자식"], "C7H8")
        self.assertEqual(detail.value[0]["독성치"], "회사 SDS 확인값")
        self.assertIsNone(project.get_field("cap.chemical.details"))

    def test_re_download_prefers_psm_enrichment_over_minimal_stage1_inventory(self):
        project = self._project()
        project.set_field(
            "psm.psi.chemical_details",
            "공정안전보고서 유해·위험물질 상세명세",
            [{
                "물질명": "톨루엔",
                "CAS 번호": "108-88-3",
                "분자식": "C7H8",
                "독성치": "회사 SDS 확인값",
            }],
            "USER_CONFIRMED",
        )

        wb = load_workbook(BytesIO(build_enhanced_integrated_authoring_workbook(project, example=False)))
        ws = wb["02_화학물질정보"]
        headers = self._headers(ws)

        self.assertEqual(ws.cell(5, headers["분자식"]).value, "C7H8")
        self.assertEqual(ws.cell(5, headers["독성치"]).value, "회사 SDS 확인값")


if __name__ == "__main__":
    unittest.main()
