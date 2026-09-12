from __future__ import annotations

import unittest
from io import BytesIO
from pathlib import Path

from openpyxl import load_workbook

from engine.template import build_legal_reference_workbook


class DecisionExplanationContractTests(unittest.TestCase):
    def test_stage1_has_company_fact_linked_explanations(self) -> None:
        text = Path('engine/stage1_workbook.py').read_text(encoding='utf-8')
        for expected in [
            '「산업안전보건법」 제44조제1항',
            '한국표준산업분류 코드',
            '합산한 값(R)',
            '「화학물질관리법」 제23조제1항',
            '최대보유량',
            '하위 규정수량',
            '상위 규정수량',
            '주요취급시설',
        ]:
            self.assertIn(expected, text)

    def test_result_card_labels_explanation(self) -> None:
        text = Path('ui/diagnosis_page.py').read_text(encoding='utf-8')
        self.assertIn('판정 근거 설명', text)

    def test_reference_workbook_uses_full_legal_names(self) -> None:
        wb = load_workbook(BytesIO(build_legal_reference_workbook()), data_only=False)
        self.assertEqual(wb.sheetnames, [
            '00_사용안내',
            '01_공정안전보고서_법령참고',
            '02_화학사고예방관리계획서_법령참고',
        ])
        values = [str(cell.value or '') for ws in wb.worksheets for row in ws.iter_rows() for cell in row]
        joined = '\n'.join(values)
        self.assertNotIn('화사계', joined)
        self.assertNotIn('PSM', joined)
        self.assertNotIn('화사계', joined)

    def test_company_workbook_guide_uses_full_names(self) -> None:
        from engine.template import build_minimal_input_workbook
        wb = load_workbook(BytesIO(build_minimal_input_workbook()), data_only=False)
        values = [str(cell.value or '') for ws in wb.worksheets for row in ws.iter_rows() for cell in row]
        joined = '\n'.join(values)
        self.assertNotIn('PSM', joined)
        self.assertNotIn('화관법', joined)
        self.assertNotIn('화사계', joined)
        self.assertIn('공정안전보고서', joined)
        self.assertIn('화학사고예방관리계획서', joined)

    def test_download_filenames_do_not_use_abbreviations(self) -> None:
        text = Path('engine/__init__.py').read_text(encoding='utf-8')
        self.assertIn('공정안전보고서_화학사고예방관리계획서_법령작성참고_v1.0.xlsx', text)
        self.assertIn('공정안전보고서_화학사고예방관리계획서_회사입력_작성예시_가이드_v1.0.xlsx', text)
        self.assertNotIn('PSM_CAP_법령작성참고', text)


if __name__ == '__main__':
    unittest.main()
