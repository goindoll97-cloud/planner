from __future__ import annotations

from io import BytesIO
import unittest

from docx import Document

from engine.stage2 import cap_form8_workspace as f8
from engine.stage2 import cap_impact_engine as impact_engine
from engine.stage2 import cap_impact_workspace as iw
from engine.stage2 import cap_release_workspace as rw
from engine.stage2 import cap_scenario_workspace as sc
from tests.test_stage2_cap_dispersion import _toxic_project


def _project_with_targets():
    project = _toxic_project()
    project.set_field("business.address", "사업장 소재지", "울산광역시 남구 산업로 1", "VERIFIED")
    f8.save(project, [
        {"보호대상 명칭": "한빛초등학교", "보호대상 구분": "갑종", "세부유형": "교육·연구시설", "주소·위치": "산업로 9",
         "사업장 경계와 거리(m)": 120, "GIS/현장 근거": "지도 캡처 2026-09-19", "거주민수": 0, "근로자수": 30},
        {"보호대상 명칭": "태화강", "보호대상 구분": "환경수용체", "세부유형": "하천", "주소·위치": "울산 남구",
         "사업장 경계와 거리(m)": 450, "GIS/현장 근거": "지도 캡처 2026-09-19", "거주민수": "", "근로자수": ""},
        {"보호대상 명칭": "먼 마을", "보호대상 구분": "을종", "세부유형": "주택·업무시설", "주소·위치": "먼 곳",
         "사업장 경계와 거리(m)": 5000, "GIS/현장 근거": "지도", "거주민수": 200, "근로자수": 0},
    ], no_target=False)
    sc.save_scenarios(project, [{
        "사고시나리오명": "E-1 염소 독성 누출", "대상 설비번호": "E-1", "유해화학물질명": "염소", "사고유형": sc.TOXIC,
        rw.BOUNDARY_COLUMN: "30"}])
    return project


class ImpactWorkspaceTests(unittest.TestCase):
    def test_off_site_scenario_counts_only_targets_inside_the_range(self):
        [impact] = iw.evaluate(_project_with_targets())
        self.assertTrue(impact.is_off_site)
        inside = {t["보호대상 명칭"] for t in impact.targets}
        self.assertIn("한빛초등학교", inside)
        self.assertNotIn("먼 마을", inside)
        expected_range = impact.off_site_m
        self.assertEqual(inside, {t["보호대상 명칭"] for t in f8.saved_rows(_project_with_targets())
                                  if t["사업장 경계와 거리(m)"] <= expected_range})
        self.assertEqual(impact.count("갑종"), 1)
        self.assertEqual(impact.people("근로자수"), 30)

    def test_range_beyond_500m_warns_that_far_receptors_are_not_listed(self):
        [impact] = iw.evaluate(_project_with_targets(), worst_case=True)
        if impact.off_site_m > iw.LISTED_RANGE_M:
            self.assertTrue(any("500m" in note for note in impact.notes))

    def test_facts_feed_the_existing_form12_and_form13_engines(self):
        project = _project_with_targets()
        impacts = iw.evaluate(project)
        counts = iw.apply_to_forms(project, impacts)
        self.assertEqual(counts["scenarios"], 1)
        form12 = impact_engine.build_cap_form12_data(project)
        self.assertEqual(form12.blockers, ())
        [row] = form12.rows
        self.assertEqual(row["대상 설비번호"], "E-1")
        self.assertEqual(row["갑종 보호대상 수"], 1)
        self.assertIn("자체 영향범위 분석 근거서", row["KORA/GIS 근거"])
        form13 = impact_engine.build_cap_form13_data(project)
        self.assertEqual(form13.blockers, ())
        self.assertGreaterEqual(len(form13.protected_targets), 1)
        self.assertEqual(form13.summary["보호대상 없음 여부"], "아니오")

    def test_scenario_that_stays_inside_the_site_is_not_an_accident_scenario(self):
        project = _project_with_targets()
        scenarios = sc.saved_scenarios(project)
        scenarios[0][rw.BOUNDARY_COLUMN] = "1000000"
        sc.save_scenarios(project, scenarios)
        impacts = iw.evaluate(project)
        self.assertFalse(impacts[0].is_off_site)
        iw.apply_to_forms(project, impacts)
        self.assertEqual(project.get_field(iw.IMPACT_TABLE_KEY).value, [])
        self.assertEqual(project.get_field(iw.NO_SCENARIO_KEY).value, "예")
        self.assertIn("없습니다", iw.summary_text(impacts))

    def test_unfinished_inputs_are_reported_not_counted(self):
        project = _project_with_targets()
        scenarios = sc.saved_scenarios(project)
        scenarios[0].pop(rw.BOUNDARY_COLUMN)
        sc.save_scenarios(project, scenarios)
        [impact] = iw.evaluate(project)
        self.assertFalse(impact.is_off_site)
        self.assertTrue(impact.problems)

    def test_report_documents_inputs_models_and_limits(self):
        project = _project_with_targets()
        impacts = iw.evaluate(project)
        doc = Document(BytesIO(iw.report_docx(project, impacts)))
        text = "\n".join(p.text for p in doc.paragraphs) + "\n" + "\n".join(
            c.text for t in doc.tables for r in t.rows for c in r.cells)
        for expected in ("영향범위 분석 근거서", "KOSHA GUIDE P-92-2023", "ERPG-2", "E-1 염소 독성 누출", "풍향·지형은 반영하지 않았다"):
            self.assertIn(expected, text)


if __name__ == "__main__":
    unittest.main()
