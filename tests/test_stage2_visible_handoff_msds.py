from __future__ import annotations

from io import BytesIO
from pathlib import Path
from tempfile import TemporaryDirectory
import unittest

from openpyxl import load_workbook

from engine.stage2.integrated_workbook import apply_integrated_authoring_workbook
from engine.stage2.msds_reference import compare_supplier_sds, inventory_chemicals
from engine.stage2.project import EvidenceRef, create_project_from_stage1_snapshot
from engine.stage2.workbook_enhancements import build_enhanced_integrated_authoring_workbook


class Stage2VisibleHandoffMSDSTests(unittest.TestCase):
    def _snapshot(self):
        return {
            "source_fingerprint": "a" * 64,
            "business": {
                "사업장명": "테스트화학 울산공장",
                "사업장 주소": "울산광역시 남구 테스트로 1",
                "대표자 성명": "홍길동",
                "사업자등록번호": "123-45-67890",
                "대표전화": "052-000-0000",
            },
            "documents": {},
            "chemicals": [
                {
                    "No.": 1,
                    "제품명": "세정제 ABC",
                    "CAS No.": "108-88-3",
                    "물질명(알면 입력)": "톨루엔",
                    "함량(%)": 60,
                    "최대 동시보유량(알면 입력)": 10000,
                    "수량 단위": "kg",
                    "혼합물 여부": "Y",
                }
            ],
            "mixture_components": [
                {
                    "적용여부": "해당",
                    "제품목록행번호": 1,
                    "제품명(확인용)": "세정제 ABC",
                    "구성성분명": "톨루엔",
                    "CAS No.": "108-88-3",
                    "함량(%)": 60,
                    "함량 최저(%)": "",
                    "함량 최고(%)": "",
                    "SDS 제3항 근거": "회사 제품 SDS 제3항",
                    "비고": "",
                },
                {
                    "적용여부": "해당",
                    "제품목록행번호": 1,
                    "제품명(확인용)": "세정제 ABC",
                    "구성성분명": "메탄올",
                    "CAS No.": "67-56-1",
                    "함량(%)": 30,
                    "함량 최저(%)": "",
                    "함량 최고(%)": "",
                    "SDS 제3항 근거": "회사 제품 SDS 제3항",
                    "비고": "",
                },
            ],
            "facilities": [],
            "decision": {
                "psm_status": "현재 확인 범위에서 공정안전보고서 제출 대상 기준 미해당",
                "cap_status": "화학사고예방관리계획서 작성수준 2군 사업장",
            },
        }

    def _project(self):
        project = create_project_from_stage1_snapshot(self._snapshot())
        project.set_authoring_scope(psm_selected=False, cap_selected=True)
        return project

    def test_stage1_mixture_components_are_carried_as_verified_field(self):
        project = self._project()
        record = project.get_field("inventory.mixture_components")
        self.assertIsNotNone(record)
        self.assertEqual(record.status, "VERIFIED")
        self.assertEqual(len(record.value), 2)
        self.assertEqual(project.stage1_snapshot["mixture_components"][1]["CAS No."], "67-56-1")

    def test_stage2_workbook_shows_handoff_source_and_mixture_sheet(self):
        project = self._project()
        data = build_enhanced_integrated_authoring_workbook(project, example=False)
        wb = load_workbook(BytesIO(data), data_only=False)
        self.assertIn("00A_Stage1승계정보", wb.sheetnames)
        self.assertIn("02A_혼합물구성성분", wb.sheetnames)
        handoff_text = "\n".join(
            str(cell.value)
            for row in wb["00A_Stage1승계정보"].iter_rows()
            for cell in row
            if cell.value is not None
        )
        self.assertIn("Stage 1 회사확정 승계값", handoff_text)
        self.assertIn("Stage 1 회사 입력 Excel", handoff_text)
        self.assertIn("혼합물 구성성분", handoff_text)
        component = wb["02A_혼합물구성성분"]
        self.assertEqual(component["E5"].value, "108-88-3")
        self.assertEqual(component["E6"].value, "67-56-1")
        self.assertEqual(component["K5"].value, "Stage 1 승계값")

    def test_stage2_does_not_overwrite_stage1_business_or_chemical_identity(self):
        project = self._project()
        data = build_enhanced_integrated_authoring_workbook(project, example=False)
        wb = load_workbook(BytesIO(data), data_only=False)

        business = wb["01_사업장정보"]
        representative_row = None
        for row_no in range(5, business.max_row + 1):
            if str(business.cell(row_no, 1).value or "") == "대표자 성명":
                representative_row = row_no
                break
        self.assertIsNotNone(representative_row)
        business.cell(representative_row, 2, "다른대표")

        chemical = wb["02_화학물질정보"]
        chemical["B5"] = "999-99-9"

        out = BytesIO()
        wb.save(out)
        result = apply_integrated_authoring_workbook(project, out.getvalue())

        self.assertTrue(any("Stage 1 승계값" in warning for warning in result.warnings))
        self.assertEqual(project.get_field("cap.business.representative").value, "홍길동")
        cap_details = project.get_field("cap.chemical.details")
        self.assertIsNotNone(cap_details)
        self.assertEqual(cap_details.value[0]["CAS 번호"], "108-88-3")

    def test_msds_inventory_includes_all_stage1_mixture_component_cas(self):
        rows = inventory_chemicals(self._project())
        self.assertEqual({row.cas for row in rows}, {"108-88-3", "67-56-1"})
        methanol = next(row for row in rows if row.cas == "67-56-1")
        self.assertEqual(methanol.chemical_name, "메탄올")
        self.assertEqual(methanol.product_name, "세정제 ABC")

    def _attach_supplier_msds(self, project, text: str):
        tmp = TemporaryDirectory()
        path = Path(tmp.name) / "세정제 ABC_MSDS.txt"
        path.write_text(text, encoding="utf-8")
        project.set_field(
            "psm.psi.msds",
            "물질안전보건자료",
            {"file_name": path.name},
            "HOLD",
            evidence=[
                EvidenceRef(
                    source_type="ATTACHMENT",
                    source_name=path.name,
                    location=str(path),
                    sha256="b" * 64,
                )
            ],
        )
        return tmp

    def test_supplier_msds_section3_composition_match_can_pass_auxiliary_check(self):
        project = self._project()
        tmp = self._attach_supplier_msds(
            project,
            "제품명: 세정제 ABC\n3. 구성성분의 명칭 및 함유량\n톨루엔 108-88-3 60%\n메탄올 67-56-1 30%\n4. 응급조치 요령\n",
        )
        try:
            result = compare_supplier_sds(project)
        finally:
            tmp.cleanup()
        self.assertEqual(result.status, "CAS_COVERED")
        self.assertIn("구성성분·함량 2건", result.message)

    def test_supplier_msds_section3_concentration_mismatch_stays_review_required(self):
        project = self._project()
        tmp = self._attach_supplier_msds(
            project,
            "제품명: 세정제 ABC\n3. 구성성분의 명칭 및 함유량\n톨루엔 108-88-3 60%\n메탄올 67-56-1 20%\n4. 응급조치 요령\n",
        )
        try:
            result = compare_supplier_sds(project)
        finally:
            tmp.cleanup()
        self.assertEqual(result.status, "PARTIAL")
        self.assertIn("일치하지 않는", result.message)
        self.assertIn("메탄올", result.message)


if __name__ == "__main__":
    unittest.main()
