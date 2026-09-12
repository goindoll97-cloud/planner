from __future__ import annotations

import unittest
from io import BytesIO

from openpyxl import load_workbook

from engine.template import build_legal_reference_workbook, build_minimal_input_workbook


class CompanyTemplateContractTests(unittest.TestCase):
    def setUp(self) -> None:
        self.workbook = load_workbook(BytesIO(build_minimal_input_workbook()), data_only=False)

    def test_guide_download_uses_current_company_input_structure(self) -> None:
        self.assertEqual(
            self.workbook.sheetnames,
            [
                "00_작성가이드",
                "01_사업장기본정보",
                "02_화학물질목록",
                "03_기존문서보유여부",
                "04_시설별최대보유량",
                "05_최종판정조건",
                "06_공정안전보고서_비고8제외수량",
            ],
        )

    def test_chemical_sheet_contains_current_screening_columns(self) -> None:
        ws = self.workbook["02_화학물질목록"]
        headers = [ws.cell(3, col).value for col in range(1, 16)]
        self.assertEqual(
            headers,
            [
                "No.",
                "제품명",
                "CAS No.",
                "물질명(알면 입력)",
                "함량(%)",
                "취급형태",
                "최대 제조·사용량",
                "최대 저장량",
                "수량 단위",
                "최대 동시보유량(알면 입력)",
                "비고",
                "상온·상압 액체 여부(해당 시)",
                "최대보유량 법정 산정 여부",
                "회사/제품 SDS 제2항 보유·확인 여부",
                "SDS 제2항 유해성·위험성 분류(선택 입력)",
            ],
        )

    def test_facility_final_and_psm_note8_sheets_are_present(self) -> None:
        facility = self.workbook["04_시설별최대보유량"]
        final_conditions = self.workbook["05_최종판정조건"]
        note8 = self.workbook["06_공정안전보고서_비고8제외수량"]
        self.assertEqual(facility["A3"].value, "적용여부")
        self.assertEqual(facility["V3"].value, "비고")
        self.assertEqual(final_conditions["A2"].value, "제도")
        self.assertEqual(final_conditions["F2"].value, "판정에 미치는 영향")
        self.assertEqual(note8["A3"].value, "적용여부")
        self.assertEqual(note8["B3"].value, "별표13 호수")
        self.assertEqual(note8["D3"].value, "저장 제외량(kg)")

    def test_final_conditions_include_psm_followup_facts(self) -> None:
        ws = self.workbook["05_최종판정조건"]
        questions = [ws.cell(row, 2).value for row in range(3, 20)]
        for expected in [
            "시행령 제43조제2항 제외설비 해당 여부",
            "시행령 제43조제2항 제외설비 유형",
            "별표 13 제1호 인화성 가스 해당 여부",
            "별표 13 제1호 하루 최대 제조·취급량(kg)",
            "별표 13 제2호 인화성 액체 해당 여부",
            "별표 13 제23호 발연황산 삼산화황(SO3) 중량%",
            "별표 13 제42호 니트로셀룰로오스 질소 함유량%",
            "가스를 전문으로 저장·판매하는 시설 내 가스 여부",
            "법적 예외 적용 유형",
        ]:
            self.assertIn(expected, questions)

    def test_guide_points_to_04_and_06_conditional_sheets(self) -> None:
        guide = self.workbook["00_작성가이드"]
        values = [cell.value for row in guide.iter_rows() for cell in row if isinstance(cell.value, str)]
        self.assertTrue(any("04_시설별최대보유량" in value for value in values))
        self.assertTrue(any("06_공정안전보고서_비고8제외수량" in value for value in values))
        self.assertFalse(any("03_시설별최대보유량" in value for value in values))

    def test_separate_legal_reference_workbook(self) -> None:
        legal = load_workbook(BytesIO(build_legal_reference_workbook()), data_only=False)
        self.assertEqual(legal.sheetnames, ["00_사용안내", "01_공정안전보고서_법령참고", "02_화학사고예방관리계획서_법령참고"])
        psm_values = "\n".join(str(c.value) for row in legal["01_공정안전보고서_법령참고"].iter_rows() for c in row if c.value is not None)
        cap_values = "\n".join(str(c.value) for row in legal["02_화학사고예방관리계획서_법령참고"].iter_rows() for c in row if c.value is not None)
        self.assertIn("「산업안전보건법 시행령」 제43조제1항", psm_values)
        self.assertIn("별표 13 비고 제7호·제8호", psm_values)
        self.assertIn("「화학물질관리법」 제23조제1항", cap_values)
        self.assertIn("「유해화학물질의 규정수량에 관한 규정」", cap_values)
        self.assertIn("「화학사고예방관리계획서 작성 등에 관한 규정」 제9조", cap_values)
        combined = psm_values + "\n" + cap_values
        for forbidden in ["사업자등록증", "생산계획", "사내", "회사 내부자료"]:
            self.assertNotIn(forbidden, combined)



if __name__ == "__main__":
    unittest.main()
