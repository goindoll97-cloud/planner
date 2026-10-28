"""정상 회사자료는 자동 분류·계산·서술형·DOCX까지 이어져야 합니다."""
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile
import os
import tempfile
import unittest
from docx import Document
from engine.stage2 import demo_authoring_company as demo, evidence_workspace as ew
from engine.stage2 import cap_fact_drafting as fd, scenario_pipeline as sp, report_export as rx
from engine.stage2.direct_template import create_direct_cap_project

class AuthoringDemoTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        cwd = os.getcwd()
        os.chdir(self.tmp.name)
        self.addCleanup(self.tmp.cleanup)
        self.addCleanup(os.chdir, cwd)

    def test_all_formats_ingest_without_conflict_and_produce_a_draft(self):
        p = create_direct_cap_project(demo.COMPANY, cap_group='1군')
        for name, content in demo.files().items():
            kind = ew.classify_document(name, content)
            self.assertIsNotNone(kind, name)
            self.assertGreater(ew.ingest(p, name, content, kind), 0, name)
        self.assertEqual(ew.problems(p), [])
        self.assertFalse(any(s['status'] == '값 충돌' for s in ew.slots(p)))
        self.assertFalse(any(d.get('needs_reading') or d.get('unmapped_columns') for d in ew.documents(p)))
        ew.publish(p, replace_existing=True)
        self.assertIsNone(p.get_field('business.address'))
        self.assertIsNone(p.get_field('cap.workspace.facilities'))  # 수동 작성대 방문 없이 계산
        results = fd.generate(p)
        self.assertEqual(len(results), 14)
        self.assertEqual(len([r for r in results if r['key'].startswith('cap.external.')]), 4)
        outcome = sp.refresh(p)
        self.assertEqual(len(outcome.plan.scenarios), 4)
        self.assertEqual(len(outcome.impacts), 4)
        self.assertEqual(outcome.problems, [])
        self.assertIsNone(p.get_field('cap.workspace.facilities'))  # FILE_READ 사실을 수동 승인으로 바꾸지 않음
        draft = rx.build_cap_workspace_draft(p)
        doc = Document(BytesIO(draft.docx))
        self.assertGreaterEqual(len(doc.tables), 16)
        self.assertIn('AI 초안', '\n'.join(p.text for p in doc.paragraphs))
        self.assertFalse(draft.state.final_ready)  # 제외한 주소·도면·현장조사 등은 보완으로 남음

    def test_archive_contains_only_supported_uploads_without_addresses_or_images(self):
        self.assertEqual({Path(n).suffix for n in demo.files()}, {'.xlsx', '.csv', '.docx', '.hwpx', '.hwp', '.pdf', '.txt'})
        with ZipFile(BytesIO(demo.archive())) as archive:
            company = [n for n in archive.namelist() if n.startswith('회사자료/')]
            self.assertEqual(len(company), 15)
            self.assertFalse(any(n.lower().endswith(('.png', '.jpg', '.jpeg', '.svg', '.dxf', '.dwg')) for n in archive.namelist()))
        for name, content in demo.files().items():
            if name.endswith(('.docx', '.xlsx', '.hwpx')):
                with ZipFile(BytesIO(content)) as archive:
                    self.assertFalse(any('/media/' in n or n.startswith('BinData/') for n in archive.namelist()), name)
        for name, content in demo.files().items():
            tables, passages = ew.read_document(name, content)
            text = str([t.rows for t in tables]) + str(passages)
            self.assertNotIn('우편번호/주소', text, name)
            self.assertNotIn('가상산업로', text, name)
            self.assertNotIn('사고원점 좌표', text, name)
            self.assertNotIn('보관계획도 최대량', text, name)
        self.assertTrue(demo.files()['10_지역사회_공조대피_기초정보.hwp'].startswith(bytes.fromhex('d0cf11e0a1b11ae1')))
