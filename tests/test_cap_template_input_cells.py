"""사용자가 제공한 규정서식: 음영 제목 보존, 흰색 입력 칸과 병합 구조 확인."""
from collections import Counter
from io import BytesIO
import unittest
from docx import Document
from engine.stage2 import cap_baseline_docx as renderer
from engine.stage2 import cap_catalog_docx as catalog_renderer
from engine.stage2 import demo_authoring_company as demo, evidence_workspace as ew, scenario_pipeline as sp, cap_fact_drafting as fd
from engine.stage2.direct_template import create_direct_cap_project


def project():
    p = create_direct_cap_project(demo.COMPANY, cap_group='1군')
    for name, content in demo.files().items():
        ew.ingest(p, name, content, ew.classify_document(name, content))
    ew.publish(p, replace_existing=True)
    fd.generate(p)
    sp.refresh(p)
    return p


def shaded(doc):
    return Counter((cell.text, str(cell._tc.tcPr.xml)) for table in doc.tables for row in table.rows
                   for cell in renderer._unique_cells(row) if renderer._is_label_cell(cell))


class TemplateInputTests(unittest.TestCase):
    def test_labels_are_immutable_in_every_form_and_both_renderers(self):
        source = Document(BytesIO(renderer.load_cap_baseline_bytes()))
        expected = shaded(source)
        for build in (renderer.build_cap_baseline_draft, catalog_renderer.build_cap_catalog_draft):
            doc = Document(BytesIO(build(project())))
            actual = shaded(doc)
            self.assertFalse(expected - actual)
            self.assertFalse(set(actual) - set(expected))

    def test_colored_cell_write_fails_without_modifying_the_template(self):
        doc = Document(BytesIO(renderer.load_cap_baseline_bytes()))
        cell = doc.tables[16].cell(6, 0)
        before = cell._tc.xml
        with self.assertRaises(ValueError):
            renderer._write_cell(cell, '흡인 유해성 구분 1')
        self.assertEqual(before, cell._tc.xml)

    def test_form7_examples_are_cleared_and_values_go_below_titles(self):
        blank = Document(BytesIO(renderer.load_cap_baseline_bytes())).tables[16]
        self.assertEqual(blank.cell(1, 0).text, '가.')
        self.assertEqual(blank.cell(3, 0).text, '다.')
        self.assertEqual(blank.cell(7, 0).text, '')
        doc = Document(BytesIO(catalog_renderer.build_cap_catalog_draft(project())))
        tables = [t for t in doc.tables if t.cell(0, 0).text == '취급물질의 일반정보']
        self.assertEqual(len(tables), 2)
        self.assertTrue(tables[0].cell(1, 0).text.startswith('가. 가상 톨루엔'))
        self.assertNotIn('물질명', tables[0].cell(1, 0).text)
        self.assertEqual(tables[0].cell(4, 0).text, '라. 99.5 %')
        self.assertEqual(tables[0].cell(6, 0).text, blank.cell(6, 0).text)
        self.assertEqual(tables[0].cell(7, 0).text, '흡인 유해성 구분 1, 생식독성 구분 2')
        self.assertFalse(renderer._is_label_cell(tables[0].cell(7, 0)))
        self.assertNotIn('254-55-3', '\n'.join(c.text for t in tables for r in t.rows for c in r.cells))

    def test_form13_checkboxes_frequency_headers_and_notes_survive(self):
        original = Document(BytesIO(renderer.load_cap_baseline_bytes()))
        doc = Document(BytesIO(catalog_renderer.build_cap_catalog_draft(project())))
        checklist = next(t for t in doc.tables if t.cell(0, 0).text.startswith('갑종 보호대상'))
        self.assertEqual(checklist.cell(0, 0).text, original.tables[24].cell(0, 0).text)
        self.assertIn('☒ 공업시설', ' '.join(c.text for r in checklist.rows for c in r.cells))
        self.assertEqual(checklist.rows[11].cells[0].text, '')  # 지도 칸에 요약문을 써넣지 않음
        frequency = [t for t in doc.tables if len(t.rows) == 13 and len(t.columns) == 5 and
                     t.cell(1, 1).text == '개시사건']
        self.assertEqual(len(frequency), 3)
        for table in frequency:
            self.assertEqual([c.text for c in table.rows[1].cells], ['연번', '개시사건', '빈도', '개수', '사고빈도'])
            self.assertEqual(table.cell(2, 1).text, '고압용기파열')
            self.assertEqual(table.cell(12, 0).text, '합')
            self.assertTrue(table.cell(12, 4).text)
        summary = next(t for t in doc.tables if t.cell(0, 0).text == '보호대상 포함 여부')
        self.assertIn('☒ 을종 보호대상', summary.cell(1, 0).text)
        self.assertEqual([c.text for c in renderer._unique_cells(summary.rows[2])],
                         ['보호대상 종류', '보호대상 명칭', '장외거리(m)', '주민수(개수)'])
        self.assertTrue(summary.rows[-1].cells[0].text.lstrip().startswith('주)'))
        # 원본 병합 폭은 모든 행에서 Word 표 그리드와 일치해야 합니다.
        for table in doc.tables:
            width = len(table._tbl.tblGrid.gridCol_lst)
            for row in table.rows:
                offsets = sum(int(n) for n in row._tr.xpath('./w:trPr/w:gridBefore/@w:val | ./w:trPr/w:gridAfter/@w:val'))
                self.assertEqual(offsets + sum(int(tc.tcPr.gridSpan.val) if tc.tcPr is not None and tc.tcPr.gridSpan is not None else 1
                                               for tc in row._tr.tc_lst), width)
