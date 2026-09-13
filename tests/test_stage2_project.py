from __future__ import annotations

from io import BytesIO
from pathlib import Path
import tempfile
import unittest

from openpyxl import load_workbook

from engine.stage2.cap_requests import build_cap_data_requests
from engine.stage2.completeness import evaluate_project_completeness
from engine.stage2.export import build_progress_workbook
from engine.stage2.intake import (
    build_intake_catalog,
    build_program_input_workbook,
    extract_form_references,
    selected_requirement_specs,
)
from engine.stage2.project import create_project_from_stage1_snapshot
from engine.stage2.psm_requests import build_psm_data_requests
from engine.stage2.requirements import (
    cap_manual_source,
    cap_requirement_specs,
    psm_example_source,
    psm_requirement_specs,
    requirement_specs_for_project,
)
from engine.stage2.storage import load_project, save_attachment, save_project


class Stage2ProjectTests(unittest.TestCase):
    def _snapshot(self, *, psm_status="공정안전보고서 제출 대상", cap_status="작성수준 — 2군 사업장"):
        return {
            "source_fingerprint": "a" * 64,
            "business": {
                "회사명": "테스트회사",
                "사업장명": "테스트공장",
                "사업장 소재지": "테스트시 테스트구 1",
            },
            "documents": {},
            "chemicals": [{"제품명": "물질A", "CAS No.": "50-00-0"}],
            "facilities": [{"시설명": "TK-101", "시설유형": "저장탱크"}],
            "decision": {
                "psm_status": psm_status,
                "cap_status": cap_status,
                "psm_legal_basis": ["근거1"],
                "cap_legal_basis": ["근거2"],
            },
        }

    def _selected_project(self, *, psm=True, cap=True, cap_status="작성수준 — 2군 사업장"):
        project = create_project_from_stage1_snapshot(self._snapshot(cap_status=cap_status))
        project.set_authoring_scope(psm_selected=psm, cap_selected=cap)
        return project

    def test_stage1_status_is_carried_without_redecision(self):
        project = create_project_from_stage1_snapshot(self._snapshot())
        self.assertTrue(project.psm_required)
        self.assertTrue(project.cap_required)
        self.assertEqual(project.cap_group, "2군")
        self.assertEqual(project.get_field("inventory.chemicals").status, "VERIFIED")
        self.assertFalse(project.scope_confirmed)
        self.assertFalse(project.psm_selected)
        self.assertFalse(project.cap_selected)

    def test_non_subject_labels_are_false(self):
        project = create_project_from_stage1_snapshot(
            self._snapshot(
                psm_status="현재 확인 범위에서 공정안전보고서 제출 대상 기준 미해당",
                cap_status="화학사고예방관리계획서 작성·제출 의무 없음 — 하위 규정수량 미만",
            )
        )
        self.assertFalse(project.psm_required)
        self.assertFalse(project.cap_required)

    def test_authoring_scope_does_not_change_legal_applicability(self):
        project = create_project_from_stage1_snapshot(self._snapshot())
        project.set_authoring_scope(psm_selected=False, cap_selected=True)
        self.assertTrue(project.psm_required)
        self.assertTrue(project.cap_required)
        self.assertFalse(project.psm_in_scope)
        self.assertTrue(project.cap_in_scope)

    def test_non_subject_document_cannot_be_selected(self):
        project = create_project_from_stage1_snapshot(
            self._snapshot(psm_status="현재 확인 범위에서 공정안전보고서 제출 대상 기준 미해당")
        )
        with self.assertRaises(ValueError):
            project.set_authoring_scope(psm_selected=True, cap_selected=False)

    def test_manual_source_is_versioned_by_hash(self):
        source = cap_manual_source()
        self.assertEqual(source["document_code"], "NICS-GP2026-8")
        self.assertEqual(source["pdf_pages"], 156)
        self.assertEqual(len(source["sha256"]), 64)
        self.assertEqual(source["source_role"], "작성 실무지침")

    def test_psm_sources_are_versioned_and_separated_from_current_law(self):
        sources = psm_example_source()
        example = sources["source"]
        outline = sources["outline_source"]
        legal_ref = sources["legal_reference"]
        self.assertEqual(example["pdf_pages"], 155)
        self.assertEqual(len(example["sha256"]), 64)
        self.assertEqual(len(outline["sha256"]), 64)
        self.assertEqual(example["source_role"], "과거 작성예시·형식 참고")
        self.assertEqual(legal_ref["checked_as_of"], "2026-09-13")
        self.assertTrue(legal_ref["final_export_must_revalidate"])

    def test_psm_registry_is_detailed_and_contains_key_drawings(self):
        specs = psm_requirement_specs()
        keys = {spec.key for spec in specs}
        self.assertGreaterEqual(len(specs), 25)
        self.assertIn("psm.psi.pfd", keys)
        self.assertIn("psm.psi.pid", keys)
        self.assertIn("psm.psi.hazardous_area", keys)
        self.assertIn("psm.operation.moc", keys)
        self.assertIn("psm.emergency.core", keys)

    def test_psm_unchecked_emergency_is_still_required_by_current_law(self):
        specs = {spec.key: spec for spec in psm_requirement_specs()}
        emergency = specs["psm.emergency.core"]
        self.assertTrue(emergency.required)
        self.assertEqual(emergency.outline_checkbox_status, "UNCHECKED")
        self.assertEqual(emergency.legal_status, "STATUTORY_REQUIRED")
        self.assertIn("시행규칙 제50조", emergency.legal_basis)

    def test_psm_unchecked_risk_procedure_is_nonblocking_until_verified(self):
        specs = {spec.key: spec for spec in psm_requirement_specs()}
        procedure = specs["psm.risk.procedure"]
        self.assertFalse(procedure.required)
        self.assertEqual(procedure.legal_status, "VERIFY_CURRENT")
        self.assertEqual(procedure.outline_checkbox_status, "UNCHECKED")

    def test_psm_request_engine_waits_for_scope_selection(self):
        project = create_project_from_stage1_snapshot(self._snapshot())
        self.assertEqual(build_psm_data_requests(project), [])
        project.set_authoring_scope(psm_selected=True, cap_selected=False)
        self.assertTrue(build_psm_data_requests(project))

    def test_psm_request_engine_does_not_ask_confirmed_stage1_inventory_again(self):
        project = self._selected_project()
        rows = build_psm_data_requests(project)
        keys = {row.requirement_key for row in rows}
        self.assertNotIn("psm.psi.chemical_inventory", keys)
        equipment = next(row for row in rows if row.requirement_key == "psm.psi.equipment_specs")
        self.assertNotIn("inventory.facilities", equipment.missing_fields)
        self.assertIn("psm.psi.equipment_specs", equipment.missing_fields)

    def test_psm_pid_request_exposes_cross_checks(self):
        project = self._selected_project()
        row = next(row for row in build_psm_data_requests(project) if row.requirement_key == "psm.psi.pid")
        self.assertIn("psm.psi.equipment_specs", row.cross_checks)
        self.assertIn("psm.psi.relief_device_specs", row.cross_checks)
        self.assertEqual(row.priority, "HIGH")

    def test_group2_excludes_external_emergency_requirements(self):
        specs = cap_requirement_specs("2군")
        self.assertEqual(len(specs), 33)
        self.assertTrue(all(spec.section != "3.6 외부 비상대응계획" for spec in specs))
        self.assertTrue(any(spec.key == "cap.internal.facility_response" for spec in specs))

    def test_group1_includes_external_emergency_requirements(self):
        specs = cap_requirement_specs("1군")
        self.assertEqual(len(specs), 37)
        external = [spec for spec in specs if spec.section == "3.6 외부 비상대응계획"]
        self.assertEqual(len(external), 4)
        self.assertTrue(any(spec.key == "cap.external.public_notice" for spec in external))

    def test_project_requirements_use_detailed_registries(self):
        project = create_project_from_stage1_snapshot(self._snapshot(cap_status="작성수준 — 2군 사업장"))
        keys = {spec.key for spec in requirement_specs_for_project(project)}
        self.assertIn("psm.psi.pid", keys)
        self.assertIn("psm.emergency.core", keys)
        self.assertIn("cap.facility.pid", keys)
        self.assertIn("cap.offsite.risk", keys)
        self.assertIn("cap.prevention.self_inspection", keys)
        self.assertNotIn("cap.external.public_notice", keys)

    def test_selected_requirement_specs_exclude_unselected_legal_target(self):
        project = self._selected_project(psm=False, cap=True)
        specs = selected_requirement_specs(project)
        systems = {spec.system for spec in specs}
        self.assertNotIn("PSM", systems)
        self.assertIn("CAP", systems)
        self.assertIn("COMMON", systems)

    def test_cap_request_engine_does_not_ask_confirmed_stage1_fields_again(self):
        project = self._selected_project(cap_status="작성수준 — 2군 사업장")
        rows = build_cap_data_requests(project)
        chemical_row = next(row for row in rows if row.requirement_key == "cap.basic.chemical_inventory")
        self.assertNotIn("inventory.chemicals", chemical_row.missing_fields)
        self.assertIn("cap.chemical.details", chemical_row.missing_fields)
        site_row = next(row for row in rows if row.requirement_key == "cap.basic.site_location")
        self.assertIn("documents.site_plan", site_row.missing_fields)

    def test_intake_catalog_marks_stage1_reuse_and_missing_materials(self):
        project = self._selected_project(psm=False, cap=True)
        rows = build_intake_catalog(project)
        chemical = next(row for row in rows if row.requirement_key == "cap.basic.chemical_inventory")
        self.assertIn("inventory.chemicals", chemical.confirmed_fields)
        self.assertIn("cap.chemical.details", chemical.missing_fields)
        self.assertEqual(chemical.coverage_status, "일부 확인")

    def test_form_reference_parser_handles_multiple_annex_forms(self):
        refs = extract_form_references("작성규정 별지 제4호·제5호서식")
        self.assertIn("별지 제4호서식", refs)
        self.assertIn("별지 제5호서식", refs)

    def test_ai_draft_does_not_count_as_ready(self):
        project = self._selected_project()
        project.set_field("process.description", "공정설명서", "AI가 작성한 설명", "AI_DRAFT")
        result = evaluate_project_completeness(project)
        row = next(r for r in result["requirements"] if r["key"] == "common.process_description")
        self.assertEqual(row["state"], "REVIEW_REQUIRED")
        self.assertEqual(row["completion_pct"], 0.0)

    def test_hold_has_priority_over_review_required(self):
        project = self._selected_project()
        project.set_field("process.description", "공정설명서", "AI가 작성한 설명", "AI_DRAFT")
        result = evaluate_project_completeness(project)
        self.assertEqual(result["overall"]["state"], "HOLD")
        self.assertEqual(result["psm"]["state"], "HOLD")
        self.assertEqual(result["cap"]["state"], "HOLD")

    def test_project_and_attachment_persist_with_scope_and_hash(self):
        project = self._selected_project(psm=False, cap=True)
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            save_project(project, root=root)
            loaded = load_project(project.project_id, root=root)
            self.assertEqual(loaded.company_name, "테스트회사")
            self.assertTrue(loaded.scope_confirmed)
            self.assertFalse(loaded.psm_in_scope)
            self.assertTrue(loaded.cap_in_scope)

            ref = save_attachment(
                project.project_id,
                "PFD test.pdf",
                b"fake-pdf-bytes",
                root=root,
            )
            self.assertEqual(len(ref.sha256), 64)
            self.assertTrue(Path(ref.location).exists())

    def test_program_input_workbook_is_clearly_not_statutory_form(self):
        project = self._selected_project(psm=False, cap=True)
        workbook = load_workbook(BytesIO(build_program_input_workbook(project)))
        self.assertIn("안내", workbook.sheetnames)
        self.assertIn("요구자료목록", workbook.sheetnames)
        guide_values = [workbook["안내"].cell(row=r, column=2).value for r in range(1, workbook["안내"].max_row + 1)]
        self.assertTrue(any("법정 별지서식이 아님" in str(value) for value in guide_values))

    def test_review_workbook_uses_selected_scope_only(self):
        project = self._selected_project(psm=False, cap=True)
        workbook = load_workbook(BytesIO(build_progress_workbook(project)))
        sheet = workbook["작성현황"]
        visible_system_values = {sheet.cell(row=row, column=1).value for row in range(2, sheet.max_row + 1)}
        self.assertNotIn("공정안전보고서", visible_system_values)
        self.assertIn("화학사고예방관리계획서", visible_system_values)
        self.assertIn("자료준비·접수현황", workbook.sheetnames)
        self.assertNotIn("공정안전보고서 요청자료", workbook.sheetnames)
        self.assertIn("화학사고예방관리계획서 요청자료", workbook.sheetnames)


if __name__ == "__main__":
    unittest.main()
