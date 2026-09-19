from __future__ import annotations

from io import BytesIO
from pathlib import Path
import unittest

from docx import Document

from engine.stage2 import cap_form10_workspace as f10
from engine.stage2 import cap_form9_workspace as f9
from engine.stage2 import cap_guideline as guide
from engine.stage2 import cap_workspace as ws
from engine.stage2.cap_baseline_docx import FORM_TABLE_INDEX, build_cap_baseline_draft
from tests import test_stage2_cap_form1_engine as form1_tests

ROOT = Path(__file__).resolve().parents[1]


def _project():
    project = form1_tests.CAPForm1EngineTests()._project()
    rows = ws.facility_editor_rows(project)
    rows[0].update({"설비번호": "TK-1", "설비명": "염소 저장탱크", "시설유형": "저장탱크",
                    "물질성상": "액체", "용량": 2.5, "용량단위": "m3", "비중": 1.4})
    ws.save_facility_rows(project, rows)
    [spec] = f9.rows(project)
    spec.update({"최대 연결구 크기(mm)": "50", "설계압력": "0.3", "운전압력": "0.2", "설계온도": "60", "운전온도": "25"})
    f9.save_specs(project, [spec])
    return project


def _filled(project, **overrides):
    [row] = f10.edit_rows(project)
    row.update({"확산방지설비 종류": "방류벽", "내부 길이(m)": "2", "내부 폭(m)": "2", "유효높이(m)": "1",
                "내부 차감용적(m3)": "0"})
    row.update(overrides)
    return row


class CAPForm10WorkspaceTests(unittest.TestCase):
    def test_schema_follows_guideline(self):
        self.assertTrue(ws.load_form_schema(10)["title"].endswith(guide.form_guidelines()[10].title))
        header = " ".join(guide.form_guidelines()[10].table_header(0))
        for label in ("설비형태", "필요용량", "유효용량", "검토결과"):
            self.assertIn(label, header)
        self.assertTrue(all(help_text for _, _, help_text in f10.EDIT_COLUMNS))

    def test_liquid_storage_tank_is_proposed_as_applicable_with_form_type(self):
        [row] = f10.edit_rows(_project())
        self.assertEqual((row["적용여부"], row["설비형태"]), ("예", "저장"))
        self.assertEqual(row["설계용량(m3)"], "2.5")

    def test_required_capacity_is_computed_from_the_entered_ratio_never_assumed(self):
        project = _project()
        f10.save(project, [_filled(project)], ratio="", basis="")
        self.assertTrue(any("필요용량" in item for item in f10.needs(project)))
        f10.save(project, [_filled(project)], ratio="110", basis="시행규칙 별표 5 확인값")
        self.assertEqual(f10.needs(project), [])
        self.assertEqual(f10.rule(project)["비율(%)"], "110")

    def test_result_is_computed_and_reaches_the_docx(self):
        project = _project()
        f10.save(project, [_filled(project)], ratio="110", basis="확인 기준")  # 필요 2.75, 유효 4
        table = Document(BytesIO(build_cap_baseline_draft(project))).tables[FORM_TABLE_INDEX["10"][0]]
        text = "\n".join(c.text for r in table.rows for c in r.cells)
        for expected in ("TK-1", "방류벽", "2.75", "4", "적정"):
            self.assertIn(expected, text)
        f10.save(project, [_filled(project, **{"내부 길이(m)": "1"})], ratio="110", basis="확인 기준")  # 유효 2 < 2.75
        table = Document(BytesIO(build_cap_baseline_draft(project))).tables[FORM_TABLE_INDEX["10"][0]]
        self.assertIn("부족", "\n".join(c.text for r in table.rows for c in r.cells))

    def test_direct_required_capacity_overrides_the_ratio(self):
        project = _project()
        f10.save(project, [_filled(project, **{"필요용량 직접입력(m3)": "3"})], ratio="110", basis="검토자료")
        [row] = f10.edit_rows(project)
        self.assertEqual(row["필요용량 직접입력(m3)"], "3")
        self.assertEqual(f10.needs(project), [])

    def test_not_applicable_facilities_are_recorded(self):
        project = _project()
        [row] = f10.edit_rows(project)
        row["적용여부"] = f10.NOT_APPLICABLE
        f10.save(project, [row], ratio="", basis="")
        self.assertEqual(f10.needs(project), [])

    def test_page_is_wired(self):
        self.assertIn('10: "별지 제10호"', (ROOT / "ui/cap_workspace_page.py").read_text(encoding="utf-8"))


if __name__ == "__main__":
    unittest.main()
