from __future__ import annotations

import base64
from hashlib import sha256
from io import BytesIO
from pathlib import Path
import unittest

from docx import Document

from engine.stage2.project import Stage2Project
from engine.stage2.cap_baseline_docx import (
    build_cap_baseline_draft,
    cap_baseline_filename,
    load_cap_baseline_bytes,
)


PROJECT_ROOT = Path(__file__).resolve().parents[1]
EXPECTED_SHA256 = "f41c3b26c72fb3e0186d5d25b99004ed7a8d3644527c12844a6a7f2c36112d2d"


def _project() -> Stage2Project:
    project = Stage2Project(
        project_id="CAP-BASELINE-TEST",
        company_name="세계화학(주) 충주공장",
        psm_required=False,
        cap_required=True,
        cap_group="1군",
    )
    project.set_authoring_scope(psm_selected=False, cap_selected=True)
    return project


def _set(project: Stage2Project, key: str, value, label: str | None = None) -> None:
    project.set_field(key, label or key, value, "USER_CONFIRMED")


class CAPBaselineDocxTests(unittest.TestCase):
    def test_baseline_bytes_match_uploaded_layout_snapshot(self):
        template_dir = PROJECT_ROOT / "data" / "templates" / "cap"
        parts = sorted(template_dir.glob("cap_statutory_forms_baseline.docx.b64.*"))
        raw = b"".join(
            base64.b64decode("".join(path.read_text(encoding="ascii").split()), validate=True)
            for path in parts
        )
        self.assertEqual(
            sha256(raw).hexdigest(),
            EXPECTED_SHA256,
            msg=f"stored_parts={len(parts)} decoded_bytes={len(raw)}",
        )

        data = load_cap_baseline_bytes()
        self.assertEqual(sha256(data).hexdigest(), EXPECTED_SHA256)
        doc = Document(BytesIO(data))
        self.assertEqual(len(doc.tables), 44)
        text = "\n".join(p.text for p in doc.paragraphs)
        for marker in ("별지 제1호서식", "별지 제6호서식", "별지 제9호서식", "별지 제16호서식"):
            self.assertIn(marker, text)

    def test_confirmed_business_and_chemical_values_are_written_into_existing_forms(self):
        project = _project()
        _set(project, "cap.business.representative", "홍길동")
        _set(project, "cap.business.registration_no", "123-45-67890")
        _set(project, "business.address", "충청북도 충주시 테스트로 1")
        _set(project, "inventory.chemicals", [
            {
                "물질명": "Methyl isocyanate",
                "CAS 번호": "624-83-9",
                "물질구분": "사고대비물질",
                "물리적 상태": "액체",
                "함량(%)": "99",
                "최대보유량": "500",
            }
        ])
        _set(project, "psm.psi.equipment_specs", [
            {
                "설비번호": "TK-101",
                "설비명": "MIC 저장탱크",
                "취급물질": "Methyl isocyanate",
                "설계용량": "20",
            }
        ])

        out = build_cap_baseline_draft(project)
        doc = Document(BytesIO(out))
        self.assertEqual(len(doc.tables), 44)
        text = "\n".join(cell.text for table in doc.tables for row in table.rows for cell in row.cells)

        self.assertIn("홍길동", text)
        self.assertIn("123-45-67890", text)
        self.assertIn("Methyl isocyanate", text)
        self.assertIn("624-83-9", text)
        self.assertIn("TK-101", text)
        self.assertIn("MIC 저장탱크", text)

    def test_unconfirmed_values_stay_blank_and_static_tables_are_untouched(self):
        project = _project()
        out = build_cap_baseline_draft(project)
        doc = Document(BytesIO(out))
        text = "\n".join(cell.text for table in doc.tables for row in table.rows for cell in row.cells)

        self.assertNotIn("[확인 필요]", text)
        self.assertIn("사업장명", text)

        # 별표 1~4 (indexes 0-5) are the static regulation-threshold reference
        # tables, never company data. Confirm the very first one is untouched.
        first_table = doc.tables[0]
        self.assertIn("예비시나리오 규정 수량", first_table.rows[0].cells[0].text)

        # 별지 제2호서식's worked example table (index 11) must stay exactly
        # as the government supplied it, never overwritten with live data.
        example_table = doc.tables[11]
        example_text = "\n".join(cell.text for row in example_table.rows for cell in row.cells)
        self.assertIn("홍길동", example_text)
        self.assertIn("장치 설비 목록 및 명세", example_text)

    def test_form6_fills_kosha_reference_only_where_company_left_it_unconfirmed(self):
        project = _project()
        _set(project, "inventory.chemicals", [
            {
                "물질명": "톨루엔",
                "CAS 번호": "108-88-3",
                "물질구분": "사고대비물질",
                # 물질상태 deliberately left unconfirmed so the KOSHA reference
                # should fill it in; 비중 is company-confirmed and must win.
                "비중": "0.87 (회사 확인값)",
            }
        ])
        project.set_field(
            "reference.kosha_msds.108883",
            "KOSHA MSDS 참고자료 · 108-88-3",
            {
                "cas": "108-88-3",
                "chemical_name": "톨루엔",
                "sections": {
                    "9": {"items": [["성상", "액체"], ["비중", "1.0 (KOSHA 값 · 회사값과 다름)"]]},
                },
            },
            "HOLD",
        )

        out = build_cap_baseline_draft(project)
        doc = Document(BytesIO(out))
        form6_text = "\n".join(cell.text for row in doc.tables[15].rows for cell in row.cells)

        self.assertIn("액체 (KOSHA 참고값·확인필요)", form6_text)
        self.assertIn("0.87 (회사 확인값)", form6_text)
        self.assertNotIn("1.0 (KOSHA 값", form6_text)

    def test_cap_regulation_filename_is_distinct_from_internal_review(self):
        name = cap_baseline_filename(_project())
        self.assertTrue(name.endswith("_화학사고예방관리계획서_규정서식_작성본.docx"))
        self.assertNotIn("내부", name)


if __name__ == "__main__":
    unittest.main()
