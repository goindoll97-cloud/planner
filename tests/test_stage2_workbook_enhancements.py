from __future__ import annotations

from io import BytesIO
import unittest

from openpyxl import load_workbook

from engine.stage2.integrated_workbook import apply_integrated_authoring_workbook
from engine.stage2.project import create_project_from_stage1_snapshot
from engine.stage2.workbook_enhancements import build_enhanced_integrated_authoring_workbook


class Stage2WorkbookEnhancementTests(unittest.TestCase):
    def _project(self):
        snapshot = {
            "source_fingerprint": "b" * 64,
            "business": {
                "회사명": "예시화학",
                "사업장명": "예시공장",
                "사업장 소재지": "테스트시 테스트구 1",
            },
            "documents": {},
            "chemicals": [
                {
                    "제품명": "톨루엔",
                    "CAS No.": "108-88-3",
                    "함량": 99.5,
                    "최대보유량": 15000,
                    "단위": "kg",
                }
            ],
            "facilities": [
                {
                    "설비번호": "TK-101",
                    "설비명": "원료 저장탱크",
                    "시설유형": "저장탱크",
                }
            ],
            "decision": {
                "psm_status": "공정안전보고서 제출 대상",
                "cap_status": "작성수준 — 1군 사업장",
                "psm_legal_basis": ["근거1"],
                "cap_legal_basis": ["근거2"],
            },
        }
        project = create_project_from_stage1_snapshot(snapshot)
        project.set_authoring_scope(psm_selected=True, cap_selected=True)
        return project

    def test_input_workbook_uses_dropdowns_and_keeps_custom_entry_open(self):
        wb = load_workbook(BytesIO(build_enhanced_integrated_authoring_workbook(self._project(), example=False)))
        chemical = wb["02_화학물질정보"]
        self.assertGreaterEqual(len(chemical.data_validations.dataValidation), 4)
        formulas = [str(dv.formula1) for dv in chemical.data_validations.dataValidation]
        self.assertTrue(any("_선택목록" in formula for formula in formulas))
        self.assertTrue(all(dv.showErrorMessage is False for dv in chemical.data_validations.dataValidation))

        relief = wb["04_안전밸브_파열판"]
        relief_formulas = [str(dv.formula1) for dv in relief.data_validations.dataValidation]
        self.assertTrue(any("03_설비정보" in formula for formula in relief_formulas))

        equipment = wb["03_설비정보"]
        self.assertIsNotNone(equipment["A4"].comment)
        self.assertIn("Tag No.", equipment["A4"].comment.text)

    def test_input_workbook_minimizes_full_example_answers(self):
        wb = load_workbook(BytesIO(build_enhanced_integrated_authoring_workbook(self._project(), example=False)))
        process = wb["06_공정정보"]
        self.assertEqual(process["C4"].value, "입력 도움말")
        self.assertIn("사실 위주", str(process["C5"].value))
        self.assertNotIn("원료는 저장탱크에서", str(process["C5"].value))

    def test_example_workbook_contains_multiple_table_scenarios(self):
        wb = load_workbook(BytesIO(build_enhanced_integrated_authoring_workbook(self._project(), example=True)))
        chemical = wb["02_화학물질정보"]
        names = [chemical.cell(row=row, column=1).value for row in range(5, 9)]
        self.assertEqual(names, ["톨루엔", "염소", "황산", "아세톤"])

        equipment = wb["03_설비정보"]
        tags = [equipment.cell(row=row, column=1).value for row in range(5, 10)]
        self.assertEqual(tags, ["TK-101", "R-201", "P-301", "E-401", "V-501"])

    def test_example_workbook_has_multiple_narrative_variants(self):
        wb = load_workbook(BytesIO(build_enhanced_integrated_authoring_workbook(self._project(), example=True)))
        process = wb["06_공정정보"]
        self.assertEqual(process["B4"].value, "예시 A")
        self.assertEqual(process["C4"].value, "예시 B")
        self.assertEqual(process["D4"].value, "예시 C")
        self.assertEqual(process["E4"].value, "작성 포인트")
        variants = [process.cell(row=5, column=col).value for col in (2, 3, 4)]
        self.assertEqual(len(set(variants)), 3)
        self.assertIn("저장·이송형", variants[0])
        self.assertIn("반응공정형", variants[1])
        self.assertIn("혼합·충전형", variants[2])

    def test_enhanced_input_remains_importable(self):
        project = self._project()
        raw = build_enhanced_integrated_authoring_workbook(project, example=False)
        wb = load_workbook(BytesIO(raw))
        wb["06_공정정보"]["B5"] = "원료 저장 후 펌프로 반응공정에 이송한다."
        output = BytesIO()
        wb.save(output)

        result = apply_integrated_authoring_workbook(project, output.getvalue())
        self.assertGreaterEqual(result.updated_fields, 1)
        record = project.get_field("process.description")
        self.assertIsNotNone(record)
        self.assertIn("반응공정", str(record.value))

    def test_example_workbook_is_still_rejected_as_company_input(self):
        project = self._project()
        raw = build_enhanced_integrated_authoring_workbook(project, example=True)
        with self.assertRaises(ValueError):
            apply_integrated_authoring_workbook(project, raw)


if __name__ == "__main__":
    unittest.main()
