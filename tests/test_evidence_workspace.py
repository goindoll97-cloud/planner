import re
import unittest
from io import BytesIO
from zipfile import ZipFile

from openpyxl import Workbook

from engine.stage2 import evidence_workspace as e
from engine.stage2.project import Stage2Project


def workbook_bytes(build, cached=None):
    """openpyxl은 수식의 계산 결과를 저장하지 못하므로, 시험용으로 캐시 값을 직접 써 넣습니다."""
    wb = Workbook()
    build(wb.active)
    buffer = BytesIO()
    wb.save(buffer)
    if not cached:
        return buffer.getvalue()
    out = BytesIO()
    with ZipFile(BytesIO(buffer.getvalue())) as src, ZipFile(out, "w") as dst:
        for item in src.infolist():
            data = src.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml":
                text = data.decode()
                for ref, value in cached.items():
                    text = re.sub(rf'(<c r="{ref}"[^>]*>)(<f>[^<]*</f>)<v></v>', rf"\1\2<v>{value}</v>", text)
                data = text.encode()
            dst.writestr(item, data)
    return out.getvalue()


HWPX_SECTION = (
    '<hs:sec xmlns:hs="http://www.hancom.co.kr/hwpml/2011/section" xmlns:hp="http://www.hancom.co.kr/hwpml/2011/paragraph">'
    '<hp:p><hp:run><hp:t>가상 설비 대장</hp:t></hp:run></hp:p>'
    '<hp:p><hp:run><hp:tbl>'
    '<hp:tr><hp:tc><hp:subList><hp:p><hp:run><hp:t>설비번호</hp:t></hp:run></hp:p></hp:subList></hp:tc>'
    '<hp:tc><hp:subList><hp:p><hp:run><hp:t>최대보유량(kg)</hp:t></hp:run></hp:p></hp:subList></hp:tc></hp:tr>'
    '<hp:tr><hp:tc><hp:subList><hp:p><hp:run><hp:t>TK-7</hp:t></hp:run></hp:p></hp:subList></hp:tc>'
    '<hp:tc><hp:subList><hp:p><hp:run><hp:t>4,500</hp:t></hp:run></hp:p></hp:subList></hp:tc></hp:tr>'
    '</hp:tbl></hp:run></hp:p></hs:sec>'
)


def hwpx_bytes(section=HWPX_SECTION):
    out = BytesIO()
    with ZipFile(out, "w") as z:
        z.writestr("mimetype", "application/hwp+zip")
        z.writestr("Contents/section0.xml", section)
    return out.getvalue()


class EvidenceTests(unittest.TestCase):
    def setUp(self):
        self.p = Stage2Project('evidence-test', '가상 회사')

    def upload(self, name, text, kind='facility'):
        return e.ingest(self.p, name, text.encode('utf-8'), kind)

    def slot(self, field):
        return next(s for s in e.slots(self.p) if s['field'] == field)

    def approve_all(self):
        for s in e.slots(self.p):
            e.approve(self.p, s['key'], s['facts'][0]['id'])

    # ---- 출처·충돌·단위 ----
    def test_conflicts_units_and_provenance(self):
        self.upload('a.csv', '설비번호,용량(L)\nTK-1,10000\n')
        self.upload('b.csv', '설비번호,용량(m3)\nTK-1,12\n')
        slot = e.slots(self.p)[0]
        self.assertEqual(slot['status'], '값 충돌')
        self.assertEqual([f['value'] for f in slot['facts']], [10, 12])
        e.approve(self.p, slot['key'], slot['facts'][0]['id'])
        self.assertEqual(e.confirmed_entities(self.p)['facility:TK-1']['capacity'], 10)
        self.assertIn('행 2 열 2', slot['facts'][0]['location'])
        self.upload('c.csv', '설비번호,용량(m3)\nTK-1,15\n')
        self.assertEqual(e.slots(self.p)[0]['status'], '값 충돌')

    def test_identical_values_from_two_sources_are_corroborated_and_auto_filled(self):
        self.upload('a.csv', '설비번호,용량(L)\nTK,10000\n')
        self.upload('b.csv', '설비번호,용량(m3)\nTK,10\n')
        slot = e.slots(self.p)[0]
        self.assertEqual(slot['status'], '자동 기입')
        self.assertEqual(len(slot['facts']), 2)

    def test_ranges_and_unspecified_units_not_silently_accepted(self):
        for text in ['設', '10~20', '10']:
            self.p = Stage2Project('test', 'test')
            self.upload('a.csv', f'설비번호,운전압력\nTK,{text}\n')
            slot = e.slots(self.p)[0]
            self.assertFalse(slot['chosen'])
            with self.assertRaises(ValueError):
                e.approve(self.p, slot['key'], slot['facts'][0]['id'])

    def test_implausible_values_are_held_and_cannot_be_approved(self):
        self.upload('a.csv', '설비번호,운전압력(MPa),함량(%)\nTK,5000,250\n')
        for field in ('pressure', 'concentration'):
            slot = self.slot(field)
            self.assertFalse(slot['chosen'])
            self.assertIn('범위', e.review_gate(slot)[0])
            with self.assertRaises(ValueError):
                e.approve(self.p, slot['key'], slot['facts'][0]['id'])

    def test_publish_is_atomic_and_invalidated_on_source_change(self):
        self.upload('a.csv', '설비번호,설비명,용량(m3)\nTK,가상 탱크,10\n')
        self.approve_all()
        e.publish(self.p)
        self.assertEqual(self.p.get_field('inventory.facilities').status, 'USER_CONFIRMED')
        self.upload('b.csv', '설비번호,용량(m3)\nTK,12\n')
        self.assertEqual(self.p.get_field('inventory.facilities').status, 'HOLD')
        with self.assertRaises(ValueError):
            e.publish(self.p)

    def test_same_file_dedup_and_material_identity(self):
        text = '제품코드,제품명,CAS 번호\nP1,제품 A,64-17-5\nP2,제품 B,64-17-5\n'
        self.upload('a.csv', text, 'material')
        self.approve_all()
        self.assertEqual(self.upload('copy.csv', text, 'material'), 0)
        self.assertEqual(len(e.confirmed_entities(self.p)), 2)

    def test_unparsed_document_and_missing_scenario_inputs_are_tasks(self):
        e.ingest(self.p, 'sds.txt', b'Unstructured safety document', 'material', 'P1')
        self.upload('a.csv', '설비번호,설비명\nTK,가상 탱크\n')
        self.assertTrue(any('연결하지 못' in t['task'] for t in e.tasks(self.p)))
        self.assertTrue(any('운전압력' in t['task'] for t in e.tasks(self.p)))

    def test_manual_corrections_cas_validation_and_report(self):
        e.add_manual(self.p, 'material:P1', 'cas', '64-17-5', '시험 SDS p.3')
        self.approve_all()
        self.assertIn('시험 SDS p.3', e.report(self.p))
        e.add_manual(self.p, 'material:P1', 'cas', '64-17-6', '오류 시험')
        slot = e.slots(self.p)[0]
        with self.assertRaises(ValueError):
            e.approve(self.p, slot['key'], slot['facts'][-1]['id'])

    def test_manual_confirm_is_one_step_and_rejects_bad_values_without_saving(self):
        e.add_manual(self.p, 'facility:T1', 'inventory', '3,000', '현장 확인 2026-10-26', confirm=True)
        self.assertEqual(e.confirmed_entities(self.p)['facility:T1']['inventory'], 3000)
        with self.assertRaises(ValueError):
            e.add_manual(self.p, 'facility:T1', 'pressure', '10~20', '범위 입력', confirm=True)
        self.assertEqual([s['field'] for s in e.slots(self.p)], ['inventory'])

    def test_demo_end_to_end(self):
        with ZipFile(BytesIO(e.demo_zip())) as z:
            for name in z.namelist():
                if name.endswith(('.csv', '.xlsx')):
                    e.ingest(self.p, name, z.read(name), 'material' if 'materials' in name else 'facility')
        queue = e.review_queue(self.p)
        self.assertEqual(sorted((i['kind'], i['field']) for i in queue), [('값 오류', 'inventory'), ('충돌', 'capacity')])
        info = e.summary(self.p)
        self.assertGreater(info['auto'], info['pending'] * 5)
        conflict = next(i for i in queue if i['kind'] == '충돌')['slot']
        e.approve(self.p, conflict['key'], conflict['facts'][0]['id'])
        e.add_manual(self.p, 'facility:RX-01', 'inventory', '2000', '현장 확인(수식 결과 확인)', confirm=True)
        self.assertEqual(e.review_queue(self.p), [])
        self.assertTrue(e.unreviewed_legal(self.p))
        with self.assertRaises(ValueError):
            e.publish(self.p)
        e.confirm_auto(self.p, [s['key'] for s in e.unreviewed_legal(self.p)])
        self.assertEqual(set(e.publish(self.p)), {'inventory.chemicals', 'inventory.facilities'})

    def test_mixed_batch_files_are_classified_before_ingestion(self):
        with ZipFile(BytesIO(e.demo_zip())) as archive:
            files = [name for name in archive.namelist() if name.endswith(('.csv', '.xlsx'))]
            kinds = {name: e.classify_document(name, archive.read(name)) for name in files}
        self.assertEqual(kinds['01_materials.csv'], 'material')
        self.assertEqual({v for k, v in kinds.items() if k != '01_materials.csv'}, {'facility'})
        with self.assertRaisesRegex(ValueError, '자료 종류'):
            e.classify_document('unclear.txt', b'No recognizable headings')

    def test_classifier_uses_identifier_column_over_incidental_headings(self):
        def kind(text):
            return e.classify_document('x.csv', text.encode('utf-8-sig'))
        # 실제 설비대장에는 취급 물질 열이 흔합니다
        self.assertEqual(kind('설비번호,설비명,물질명,CAS번호,용량(m3),최대보유량(kg)\nTK-01,탱크,톨루엔,108-88-3,10,8000\n'), 'facility')
        self.assertEqual(kind('제품코드,제품명,CAS번호,함량(%),용량(L)\nP-01,신나,108-88-3,60,18\n'), 'material')
        # 식별자가 없어도 항목명이 한쪽으로 뚜렷하면 분류합니다
        self.assertEqual(kind('물질명,CAS번호,함량(%)\n톨루엔,108-88-3,60\n'), 'material')
        self.assertEqual(kind('설비명,용량(m3),운전압력(MPa)\n탱크,10,0.1\n'), 'facility')
        # 두 대상을 잇는 대응표나 근거가 엇갈리면 묻습니다
        for text in ('설비번호,제품코드\nTK-01,P-01\n', '물질명,용량(m3)\n톨루엔,10\n'):
            with self.assertRaisesRegex(ValueError, '자료 종류'):
                kind(text)

    def test_classifier_reads_each_file_once_even_when_screen_redraws(self):
        from unittest import mock
        good = '설비번호,설비명,용량(m3)\nTK-99,탱크,10\n'.encode('utf-8-sig')
        unclear = '시설과 무관한 머리글,값\n가,1\n'.encode('utf-8-sig')
        with mock.patch.object(e, 'read_document', wraps=e.read_document) as reader:
            for _ in range(3):                       # 화면이 세 번 다시 그려짐
                self.assertEqual(e.classify_document('cache_a.csv', good), 'facility')
                with self.assertRaisesRegex(ValueError, '자료 종류'):
                    e.classify_document('cache_b.csv', unclear)
            self.assertEqual(reader.call_count, 2)   # 판별 실패도 다시 읽지 않음
            with self.assertRaises(ValueError):       # 확장자가 다르면 읽는 방식이 달라 새로 읽음 (TXT는 표가 없음)
                e.classify_document('cache_a.txt', good)
            self.assertEqual(reader.call_count, 3)

    def test_invalid_file_leaves_graph_unchanged(self):
        before = e.graph(self.p)
        with self.assertRaises(ValueError):
            e.ingest(self.p, 'bad.hwp', b'data')
        self.assertEqual(e.graph(self.p), before)

    # ---- 자동 기입 규칙 ----
    def test_clean_table_values_are_auto_filled_but_name_linked_rows_are_held(self):
        self.upload('f.csv', '설비번호,설비명,최대보유량(kg),운전압력(MPa),운전온도(C)\nT1,탱크,5000,0.2,25\n')
        for slot in e.slots(self.p):
            self.assertEqual(slot['status'], '자동 기입', slot['field'])
        self.upload('g.csv', '설비명,최대보유량(kg)\n이름만 있는 탱크,100\n')
        held = next(s for s in e.slots(self.p) if s['entity'] == 'facility:이름만 있는 탱크' and s['field'] == 'inventory')
        self.assertFalse(held['chosen'])
        self.assertIn('이름으로 연결', e.review_gate(held)[0])

    def test_unit_missing_from_header_is_held(self):
        self.upload('f.csv', '설비번호,운전온도\nT1,25\n')
        slot = self.slot('temperature')
        self.assertFalse(slot['chosen'])
        self.assertIn('단위', e.review_gate(slot)[0])

    def test_confidence_threshold_depends_on_impact(self):
        # 병합 셀(0.90)은 영향이 작은 항목은 통과하지만, 법적 판단에 쓰이는 항목이나 핵심값은 막습니다.
        def build(ws):
            ws.append(['설비번호', '설비명', '운전압력(MPa)', '최대보유량(kg)'])
            ws.append(['T1', '탱크', 0.2, 5000])
            ws.append(['T2', '탱크2', None, None])
            ws.merge_cells('C2:C3')
            ws.merge_cells('D2:D3')
        e.ingest(self.p, 'm.xlsx', workbook_bytes(build), 'facility')
        pressure = next(s for s in e.slots(self.p) if s['entity'] == 'facility:T2' and s['field'] == 'pressure')
        inventory = next(s for s in e.slots(self.p) if s['entity'] == 'facility:T2' and s['field'] == 'inventory')
        self.assertEqual(pressure['facts'][0]['flags'], ['merged'])
        self.assertEqual(pressure['status'], '자동 기입')
        self.assertFalse(inventory['chosen'])
        self.assertIn('병합 셀', e.review_gate(inventory)[0])
        self.assertIn('엄격', e.review_gate(inventory)[0])
        self.assertEqual(e.review_gate(pressure, critical=frozenset({pressure['key']}))[0][:4], '예비 순')

    def test_critical_auto_filled_values_are_listed_and_can_be_confirmed(self):
        self.upload('m.csv', '제품코드,제품명,CAS 번호,함량(%)\nP,제품,64-17-5,100\n', 'material')
        self.upload('f.csv', '설비번호,취급제품,최대보유량(kg),운전압력(MPa),운전온도(C)\nT1,제품,9000,0.5,20\nT2,제품,8000,0.5,20\n')
        keys = {s['key'] for s in e.critical_auto(self.p)}
        self.assertIn('facility:T1|inventory', keys)
        self.assertEqual(e.summary(self.p)['critical_auto'], len(keys))
        self.assertEqual(e.confirm_auto(self.p, keys), len(keys))
        self.assertEqual(e.critical_auto(self.p), [])
        self.assertIn('사람이 아직 보지 않은 핵심값 0개', e.report(self.p))

    def test_bulk_approve_skips_conflicts_errors_and_ai_values(self):
        self.upload('a.csv', '설비번호,용량(m3),운전압력\nT,10,0.5\n')
        self.upload('b.csv', '설비번호,용량(m3)\nT,12\n')
        result = e.approve_bulk(self.p)
        self.assertEqual(result['approved'], 0)
        self.assertEqual({reason for _, reason in result['skipped']}, {'값 충돌', '범위·단위·수식을 확인해야 하는 값'})
        self.upload('c.csv', '설비번호,용량(L)\nU,5000\n')
        self.upload('d.csv', '설비명,용량(L)\n이름만,7000\n')
        self.assertEqual(e.approve_bulk(self.p)['approved'], 2)  # 이름만으로 연결한 명칭·용량
        self.assertEqual(self.slot_by_entity('facility:이름만', 'capacity')['method'], 'BULK')

    def slot_by_entity(self, entity, field):
        return next(s for s in e.slots(self.p) if s['entity'] == entity and s['field'] == field)

    # ---- 실제 자료 읽기 ----
    def test_header_row_is_found_below_title_lines(self):
        text = '○○공장 설비대장 (2026.03 기준),,,\n,,,\n설비번호,설비명,최대보유량(kg)\nTK-01,저장탱크,7000\n'
        self.assertEqual(self.upload('ledger.csv', text), 2)  # 식별번호 열은 사실이 아니라 대상 식별자
        doc = e.graph(self.p)['documents'][0]
        self.assertEqual(doc['tables'][0]['header_row'], 3)
        self.assertIn('행 4 열 3', self.slot('inventory')['facts'][0]['location'])

    def test_unrecognized_table_lists_its_columns_instead_of_failing_silently(self):
        self.assertEqual(self.upload('x.csv', '번호,이름\n1,가\n'), 0)
        doc = e.graph(self.p)['documents'][0]
        self.assertEqual(doc['unmapped_columns'], ['번호', '이름'])
        self.assertTrue(doc['needs_reading'])

    def test_rows_without_any_identifier_are_counted(self):
        self.upload('x.csv', '설비번호,설비명,최대보유량(kg)\nTK-1,탱크,100\n,,500\n')
        self.assertEqual(e.graph(self.p)['documents'][0]['rows_without_id'], 1)
        self.assertTrue(any('건너뛴 행 1개' in t['task'] for t in e.tasks(self.p)))

    def test_excel_formulas_use_cached_value_but_are_not_blindly_trusted(self):
        def build(ws):
            ws.append(['설비번호', '운전압력(MPa)', '최대보유량(kg)'])
            ws.append(['T1', '=0.1+0.1', '=5000+2000'])
        data = workbook_bytes(build, cached={'B2': '0.2'})
        e.ingest(self.p, 'f.xlsx', data, 'facility')
        pressure, inventory = self.slot('pressure'), self.slot('inventory')
        self.assertEqual(pressure['facts'][0]['value'], 0.2)
        self.assertEqual(pressure['facts'][0]['flags'], ['formula'])
        self.assertEqual(pressure['status'], '자동 기입')
        self.assertEqual(inventory['facts'][0]['flags'], ['formula_uncached'])
        self.assertIn('수식', e.review_gate(inventory)[0])
        self.assertFalse(inventory['chosen'])

    def test_hwpx_tables_are_read_with_location(self):
        self.assertEqual(e.ingest(self.p, 'ledger.hwpx', hwpx_bytes(), 'facility'), 1)
        fact = self.slot('inventory')['facts'][0]
        self.assertEqual(fact['value'], 4500)
        self.assertEqual(fact['entity'], 'facility:TK-7')
        self.assertIn('구역 1 표 1 행 2 열 2', fact['location'])
        self.assertTrue(any('문단' in p['location'] for p in e.graph(self.p)['documents'][0]['passages']))

    def test_hwpx_rejects_entity_declarations_and_legacy_hwp(self):
        with self.assertRaises(ValueError):
            e.ingest(self.p, 'x.hwpx', hwpx_bytes('<!DOCTYPE x [<!ENTITY a "b">]><x/>'), 'facility')
        with self.assertRaises(ValueError):
            e.ingest(self.p, 'old.hwpx', b'not a zip', 'facility')

    # ---- 별지 제1호로 이어지는 선택형 항목 ----
    def test_choice_fields_match_form1_vocabulary_and_normalize_common_wording(self):
        from engine.stage2 import cap_workspace as ws
        self.assertEqual(list(e.CHOICES['facility_type']), ws.FACILITY_TYPES)
        self.assertEqual(list(e.CHOICES['substance_state']), ws.SUBSTANCE_STATES)
        self.upload('f.csv', '설비번호,시설유형,물질상태\nT1,저장 탱크,기체\nT2,탱크,액체\n')
        slots = {(s['entity'], s['field']): s for s in e.slots(self.p)}
        t1_type, t1_state = slots[('facility:T1', 'facility_type')], slots[('facility:T1', 'substance_state')]
        self.assertEqual((t1_type['chosen']['value'], t1_state['chosen']['value']), ('저장탱크', '기체·고압가스'))
        bad = slots[('facility:T2', 'facility_type')]
        self.assertFalse(bad['chosen'])
        self.assertIn('허용 값', e.review_gate(bad)[0])

    def test_published_rows_prefill_the_legal_holding_calculation_inputs(self):
        # 별표4 산정(설계용량×비중)은 별지 제1호 시설 입력에서 저장할 때 일어납니다.
        # 그래프는 그 화면이 묻는 값(시설유형·성상·용량·비중)을 미리 채워, 다시 묻지 않게 합니다.
        from engine.stage2 import cap_workspace as ws
        self.upload('m.csv', '제품코드,제품명,CAS 번호,함량(%)\nP,제품 A,64-17-5,99\n', 'material')
        self.upload('f.csv', '설비번호,설비명,취급제품,용량(m3),비중,시설유형,물질성상\nT1,탱크1,제품 A,10,0.79,저장탱크,액체\nT2,탱크2,제품 A,5,0.79,저장탱크,액체\n')
        self.assertEqual(e.review_queue(self.p), [])
        e.confirm_auto(self.p, [s['key'] for s in e.unreviewed_legal(self.p)])
        e.publish(self.p)
        rows = ws.facility_editor_rows(self.p)
        self.assertEqual([(r['시설유형'], r['물질성상'], r['비중']) for r in rows], [('저장탱크', '액체', 0.79)] * 2)
        results = ws.compute_holdings(self.p, rows)
        self.assertEqual([r.problem for r in results], ['', ''])
        self.assertAlmostEqual(results[1].ton, 3.95)
        self.assertEqual(e.form1_needs(self.p), ())

    def test_missing_calculation_inputs_are_reported_from_the_engine_not_guessed(self):  # noqa
        self.upload('m.csv', '제품코드,제품명,CAS 번호,함량(%)\nP,제품 A,64-17-5,99\n', 'material')
        self.upload('f.csv', '설비번호,설비명,취급제품,용량(m3),시설유형,물질성상\nT1,탱크,제품 A,10,저장탱크,액체\n')
        e.confirm_auto(self.p, [s['key'] for s in e.unreviewed_legal(self.p)])
        e.publish(self.p)
        self.assertTrue(any('비중' in t['task'] and t['entity'] == '별지 제1호' for t in e.tasks(self.p)))

    def test_missing_facility_type_is_a_task(self):
        self.upload('f.csv', '설비번호,설비명\nT1,탱크\n')
        tasks = [t['task'] for t in e.tasks(self.p)]
        self.assertIn('기본 운전조건 미확인: 시설유형', tasks)
        self.assertIn('기본 운전조건 미확인: 물질성상', tasks)


if __name__ == '__main__':
    unittest.main()
