from io import BytesIO

from openpyxl import load_workbook

from engine.stage2 import cap_upload_template as tpl, cap_fact_catalog as cat
from engine.stage2 import evidence_workspace as ew, cap_fact_drafting as drafting
from engine.stage2.direct_template import create_direct_cap_project


def save(book):
    out = BytesIO()
    book.save(out)
    return out.getvalue()


def test_full_company_workbook_has_all_input_facts_and_no_results():
    book = load_workbook(BytesIO(tpl.build_company()))
    assert len(book.sheetnames) == 9
    assert '주변 보호대상' not in book.sheetnames
    scenario_headers = [c.value for c in book['사고시나리오 검토'][1]]
    assert {'액위(m)', '비열비', '배관누출', '개수 산정근거', '수동적 완화장치'} <= set(scenario_headers)
    assert not {'경계까지 거리(m)', '사고원점 좌표', '선정할 사고유형', '시나리오 선정 근거'} & set(scenario_headers)      # 지도·프로그램이 채우는 값
    assert [c.value for c in book[tpl.SITE_SHEET][1]][:2] == ['항목', '값']
    assert book[tpl.SITE_SHEET].max_column == 3
    assert all(book[tpl.SITE_SHEET].cell(r, 2).value is None for r in range(2, book[tpl.SITE_SHEET].max_row + 1))
    generated = drafting.generated_fact_ids()
    for item in cat.facts(docs=['CAP']):
        if item.id in generated or item.source not in {'company', 'project', 'survey'} or not item.readable:
            continue
        if item.kind in cat.SITE_KINDS:
            assert item.template_header in {r[0].value for r in book[tpl.SITE_SHEET].iter_rows(min_row=2)}, item.id
        else:
            title = tpl.sheet_title(item.table) if item.table else (
                tpl.ENTITY_SHEETS['scenario'] if item.file == tpl.SCENARIO_FILE else tpl.ENTITY_SHEETS[item.kind])
            if item.id in tpl.MAP_SCENARIO_IDS or title == '주변 보호대상':
                continue  # 지도로 확인하는 값은 엑셀에 적지 않는다
            assert item.template_header in [c.value for c in book[title][1]], item.id
    assert '안내' not in book.sheetnames
    assert not any('첨부' in s for s in book.sheetnames)


def test_blank_company_workbook_examples_are_not_ingested():
    p = create_direct_cap_project('시험', address='주소', cap_group='')
    data = tpl.build_company()
    ew.ingest(p, tpl.COMPANY_FILENAME, data, ew.classify_document(tpl.COMPANY_FILENAME, data))
    assert not ew.graph_facts(p)
    assert not ew.documents(p)[0]['unmapped_columns']
    assert ew.documents(p)[0]['rows_without_id'] == 0


def test_company_workbook_reads_every_filled_sheet_and_replaces_same_name():
    p = create_direct_cap_project('시험', address='주소', cap_group='')
    book = load_workbook(BytesIO(tpl.build_company(include_scenario_sheets=True)))
    sheet = book[tpl.SITE_SHEET]
    rep_row = next(r for r in range(2, sheet.max_row + 1) if sheet.cell(r, 1).value == '대표자')
    sheet.cell(rep_row, 2, '김시험')
    for title, values in {
        '설비': {'설비번호': 'TK-01', '설비명': '저장탱크', '용량(m3)': 10},
        '물질': {'제품코드': 'P-01', '제품명': '시험제품', 'CAS 번호': '108-88-3'},
        '가스감지기': {'감지기 번호': 'GD-01', '설치위치': '저장구역'},
        '사고시나리오 검토': {'설비번호': 'TK-01', '경계까지 거리(m)': 30},
        '주변 보호대상': {'보호대상 명칭': '시험학교', '주소·위치': '학교로 1'},
        '처리시설': {'처리시설 번호': 'SC-01', '시설명': '스크러버'},
        '안전밸브': {'안전밸브 식별번호': 'PSV-01', '보호대상 설비번호': 'TK-01'},
        '변경내역': {'변경기록 일자': '2026-01-01', '변경항목': '탱크', '담당자': '시설팀'},
        '혼합제품 구성성분': {'혼합제품 코드': 'M-01', '혼합제품명': '혼합제품', '구성성분 CAS 번호': '108-88-3', '구성성분 함량(%)': '40'},
    }.items():
        ws = book[title]
        for cell in ws[1]:
            if cell.value in values:
                ws.cell(2, cell.column, values[cell.value])
    ws = book['혼합제품 구성성분']
    ws.append(['M-01', '혼합제품', '67-64-1', '아세톤', '60'])
    def upload():
        data = save(book)
        ew.replace_same_name(p, tpl.COMPANY_FILENAME, data)
        ew.ingest(p, tpl.COMPANY_FILENAME, data, ew.classify_document(tpl.COMPANY_FILENAME, data))
    upload()
    assert not ew.documents(p)[0]['unmapped_columns']
    assert ew.documents(p)[0]['rows_without_id'] == 0
    preview = ew.catalog_preview(p)
    assert preview['cap.business.representative'][0] == '김시험'
    assert preview['cap.safety.waste_treatment'][0][0]['시설명'] == '스크러버'
    assert len(preview['cap.company.mixture_components'][0]) == 2
    assert preview['cap.safety.relief_device_specs'][0][0]['보호대상 설비번호'] == 'TK-01'
    sheet.cell(rep_row, 2, '이시험')
    upload()
    assert len(ew.documents(p)) == 1
    assert ew.catalog_preview(p)['cap.business.representative'][0] == '이시험'
    assert not any(f['value'] == '김시험' for f in ew.graph_facts(p))
    sheet.cell(rep_row, 2).value = None
    upload()
    assert not any(f['value'] == '이시험' for f in ew.graph_facts(p))


def test_every_full_workbook_input_survives_reading_with_source_location():
    book = load_workbook(BytesIO(tpl.build_company(include_scenario_sheets=True)))
    wanted = tpl.company_columns()
    for r, item in enumerate(wanted['site'], 2):
        book[tpl.SITE_SHEET].cell(r, 2, item.options[0] if item.options else ('10' if item.format == 'number' else f'확인값 {r}'))
    expected = []
    for kind, title in tpl.ENTITY_SHEETS.items():
        ws = book[title]
        ws.cell(2, 1, {'facility':'TK-01','scenario':'TK-01','material':'P-01','detector':'GD-01','target':'학교'}[kind])
        for item in wanted[kind]:
            col = next(c.column for c in ws[1] if c.value == item.template_header)
            value = item.options[0] if item.options else ('108-88-3' if item.evidence_field == 'cas' else '10')
            ws.cell(2, col, value)
            expected.append(item.evidence_field)
    for table, items in wanted['tables'].items():
        for col, item in enumerate(items, 1):
            book[tpl.sheet_title(table)].cell(2, col, '108-88-3' if 'CAS' in item.template_header else str(col))
            expected.append(item.evidence_field)
    p = create_direct_cap_project('시험', address='주소', cap_group='')
    data = save(book)
    ew.ingest(p, tpl.COMPANY_FILENAME, data, ew.classify_document(tpl.COMPANY_FILENAME, data))
    assert not ew.documents(p)[0]['unmapped_columns']
    assert not ew.documents(p)[0]['rows_without_id']
    facts = ew.graph_facts(p)
    fields = {f['field'] for f in facts}
    assert set(expected) <= fields
    assert {i.evidence_field for i in wanted['site']} <= fields
