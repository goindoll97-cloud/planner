from __future__ import annotations

from io import BytesIO
import unittest

from openpyxl import load_workbook

from engine.stage2.cap_form11_engine import build_cap_form11_data
from engine.stage2.project import Stage2Project
from engine.stage2.workbook_enhancements import build_enhanced_integrated_authoring_workbook


class CAPForm11EngineTests(unittest.TestCase):
    def _project(self) -> Stage2Project:
        project = Stage2Project(
            project_id="S2-FORM11",
            company_name="테스트화학",
            cap_required=True,
            cap_group="1군",
            scope_confirmed=True,
            cap_selected=True,
        )
        project.set_field(
            "inventory.chemicals",
            "화학물질 목록",
            [
                {"물질명": "톨루엔", "CAS 번호": "108-88-3", "함량(%)": 99.5, "물리적 상태": "액체"},
                {"물질명": "염소", "CAS 번호": "7782-50-5", "함량(%)": 99.9, "물리적 상태": "기체"},
            ],
            "USER_CONFIRMED",
        )
        project.set_field(
            "inventory.facilities",
            "시설·설비 목록",
            [
                {"설비번호": "TK-101", "설비명": "톨루엔 저장탱크", "설비종류": "저장탱크"},
                {"설비번호": "V-201", "설비명": "염소 공급용기군", "설비종류": "압력용기"},
            ],
            "USER_CONFIRMED",
        )
        return project

    @staticmethod
    def _fixed_row(**overrides):
        row = {
            "감지기 번호": "GD-101",
            "설치형태": "고정식",
            "설치위치": "TK-101 방유제 내",
            "검출대상 물질": "톨루엔",
            "작동시간": "30초 이내",
            "측정방식": "접촉연소식",
            "경보 설정값": "10% LEL",
            "경보 위치": "중앙제어실",
            "연동여부": "예",
            "연동 설비·조치": "HH 경보 시 P-101 원료이송펌프 정지",
            "정밀도": "±3% F.S.",
            "유지관리": "월 1회 기능점검·연 1회 교정",
            "비상전원 여부": "예",
            "관련 도면번호": "GA-101",
            "비고": "",
        }
        row.update(overrides)
        return row

    def test_complete_fixed_detector_is_ready(self):
        project = self._project()
        project.set_field(
            "cap.safety.gas_detection",
            "가스감지기",
            [self._fixed_row()],
            "USER_CONFIRMED",
        )

        result = build_cap_form11_data(project)

        self.assertEqual(result.blockers, ())
        self.assertTrue(result.ready)
        row = result.rows[0]
        self.assertEqual(row["구분기호"], "GD-101")
        self.assertEqual(row["측정방식"], "접촉연소식")
        self.assertEqual(row["연동여부"], "예")
        self.assertIn("연동조치:", row["비고"])
        self.assertIn("도면: GA-101", row["비고"])

    def test_portable_detector_is_excluded_from_form11(self):
        project = self._project()
        portable = self._fixed_row(
            **{
                "감지기 번호": "PD-301",
                "설치형태": "휴대식",
                "설치위치": "정비팀 비치",
                "검출대상 물질": "복합가스",
            }
        )
        project.set_field("cap.safety.gas_detection", "가스감지기", [portable], "USER_CONFIRMED")

        result = build_cap_form11_data(project)

        self.assertEqual(result.excluded_portable_count, 1)
        self.assertEqual(result.rows, ())
        self.assertTrue(any("고정식 유해감지시설" in blocker for blocker in result.blockers))

    def test_mixed_installation_row_is_blocking(self):
        project = self._project()
        project.set_field(
            "cap.safety.gas_detection",
            "가스감지기",
            [self._fixed_row(**{"설치형태": "고정식+휴대식"})],
            "USER_CONFIRMED",
        )

        result = build_cap_form11_data(project)

        self.assertTrue(any("별도 행으로 분리" in blocker for blocker in result.blockers))

    def test_legacy_fixed_value_is_not_reused_as_measurement_method(self):
        project = self._project()
        row = self._fixed_row()
        row.pop("설치형태")
        row.pop("측정방식")
        row["감지방식"] = "고정식"
        project.set_field("cap.safety.gas_detection", "가스감지기", [row], "USER_CONFIRMED")

        result = build_cap_form11_data(project)

        self.assertEqual(result.rows[0]["측정방식"], "")
        self.assertTrue(any("측정방식이 비어 있습니다" in blocker for blocker in result.blockers))

    def test_interlock_yes_requires_linked_action(self):
        project = self._project()
        project.set_field(
            "cap.safety.gas_detection",
            "가스감지기",
            [self._fixed_row(**{"연동 설비·조치": ""})],
            "USER_CONFIRMED",
        )

        result = build_cap_form11_data(project)

        self.assertTrue(any("연동 설비·조치내용" in blocker for blocker in result.blockers))

    def test_alarm_setting_must_have_number_and_unit(self):
        project = self._project()
        project.set_field(
            "cap.safety.gas_detection",
            "가스감지기",
            [self._fixed_row(**{"경보 설정값": "사업장 기준"})],
            "USER_CONFIRMED",
        )

        result = build_cap_form11_data(project)

        self.assertTrue(any("수치와 단위" in blocker for blocker in result.blockers))

    def test_unknown_equipment_tag_in_location_is_blocking(self):
        project = self._project()
        project.set_field(
            "cap.safety.gas_detection",
            "가스감지기",
            [self._fixed_row(**{"설치위치": "TK-999 방유제 내"})],
            "USER_CONFIRMED",
        )

        result = build_cap_form11_data(project)

        self.assertTrue(any("TK-999" in blocker for blocker in result.blockers))

    def test_workbook_has_separate_installation_and_measurement_columns(self):
        project = self._project()
        legacy = self._fixed_row()
        legacy.pop("설치형태")
        legacy.pop("측정방식")
        legacy["감지방식"] = "고정식"
        project.set_field("cap.safety.gas_detection", "가스감지기", [legacy], "USER_CONFIRMED")

        data = build_enhanced_integrated_authoring_workbook(project, example=False)
        wb = load_workbook(BytesIO(data), data_only=False)
        ws = wb["05_가스누출감지_경보장치"]
        headers = [str(cell.value or "") for cell in ws[4]]
        self.assertIn("설치형태", headers)
        self.assertIn("측정방식", headers)
        self.assertNotIn("감지방식", headers)

        header_map = {str(cell.value or ""): cell.column for cell in ws[4]}
        self.assertEqual(ws.cell(5, header_map["설치형태"]).value, "고정식")
        self.assertIn(ws.cell(5, header_map["측정방식"]).value, (None, ""))


if __name__ == "__main__":
    unittest.main()
