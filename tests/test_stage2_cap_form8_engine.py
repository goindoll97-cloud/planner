from __future__ import annotations

from io import BytesIO
import unittest
from types import SimpleNamespace
from unittest.mock import patch

from docx import Document
from openpyxl import load_workbook

from engine.stage2.cap_chemical_legal import CAPChemicalLegalData, build_cap_chemical_legal_data
from engine.stage2.cap_sds_engine import CAPForm6SDSData
from engine.stage2.cap_form8_engine import build_cap_form8_data
from engine.stage2.integrated_workbook import build_integrated_authoring_workbook
from engine.stage2.project import Stage2Project
from engine.stage2.cap_baseline_docx import build_cap_baseline_draft


class CAPForm8EngineTests(unittest.TestCase):
    def _project(self) -> Stage2Project:
        project = Stage2Project(
            project_id="S2-FORM8",
            company_name="입지테스트화학",
            cap_required=True,
            cap_group="1군",
            scope_confirmed=True,
            cap_selected=True,
        )
        return project

    def test_structured_500m_target_is_ready(self):
        project = self._project()
        project.set_field(
            "cap.site.surrounding_environment",
            "사업장 주변 환경정보",
            [{
                "보호대상 없음 여부": "아니오",
                "보호대상 명칭": "○○초등학교",
                "보호대상 구분": "갑종",
                "세부유형": "교육·연구시설",
                "주소·위치": "○○시 ○○로 10",
                "좌표": "35.0,129.0",
                "사업장 경계와 거리(m)": 420,
                "GIS/현장 근거": "GIS-SITE-01",
            }],
            "USER_CONFIRMED",
        )

        result = build_cap_form8_data(project)

        self.assertTrue(result.ready, result.blockers)
        self.assertFalse(result.no_protected_targets)
        self.assertEqual(result.rows[0]["500m 이내 여부"], "예")
        self.assertEqual(result.rows[0]["사업장 경계와 거리(m)"], 420)

    def test_distance_over_500m_fails_closed(self):
        project = self._project()
        project.set_field(
            "cap.site.surrounding_environment",
            "사업장 주변 환경정보",
            [{
                "보호대상 없음 여부": "아니오",
                "보호대상 명칭": "원거리 시설",
                "보호대상 구분": "을종",
                "세부유형": "근린생활시설",
                "주소·위치": "○○시 원거리로 1",
                "사업장 경계와 거리(m)": 620,
                "GIS/현장 근거": "GIS-SITE-02",
            }],
            "USER_CONFIRMED",
        )

        result = build_cap_form8_data(project)

        self.assertFalse(result.ready)
        self.assertTrue(any("500m를 초과" in item for item in result.blockers))

    def test_explicit_no_targets_requires_evidence_and_then_is_ready(self):
        project = self._project()
        project.set_field(
            "cap.site.surrounding_environment",
            "사업장 주변 환경정보",
            [{"보호대상 없음 여부": "예", "GIS/현장 근거": "GIS-NONE-01"}],
            "USER_CONFIRMED",
        )

        result = build_cap_form8_data(project)

        self.assertTrue(result.ready, result.blockers)
        self.assertTrue(result.no_protected_targets)
        self.assertEqual(result.rows, ())

    def test_no_target_without_gis_or_field_evidence_is_hold(self):
        project = self._project()
        project.set_field(
            "cap.site.surrounding_environment",
            "사업장 주변 환경정보",
            [{"보호대상 없음 여부": "예"}],
            "USER_CONFIRMED",
        )

        result = build_cap_form8_data(project)

        self.assertFalse(result.ready)
        self.assertTrue(any("GIS/현장 근거" in item for item in result.blockers))

    def test_integrated_workbook_contains_structured_form8_sheet(self):
        project = self._project()
        data = build_integrated_authoring_workbook(project, example=False)
        wb = load_workbook(BytesIO(data), data_only=False)

        self.assertIn("27_사업장주변_500m_보호대상", wb.sheetnames)
        headers = [
            str(cell.value or "")
            for cell in wb["27_사업장주변_500m_보호대상"][4]
        ]
        self.assertIn("보호대상 없음 여부", headers)
        self.assertIn("보호대상 구분", headers)
        self.assertIn("사업장 경계와 거리(m)", headers)
        self.assertIn("GIS/현장 근거", headers)

    @patch("engine.stage2.cap_chemical_legal.build_cap_form1_data")
    @patch("engine.stage2.cap_chemical_legal.load_approved_scope_tables")
    def test_stale_company_legal_identity_does_not_survive_failed_current_law_lookup(
        self, tables_mock, form1_mock
    ):
        project = self._project()
        project.set_field(
            "inventory.chemicals",
            "화학물질 목록",
            [{
                "물질명": "가상물질",
                "CAS 번호": "123-45-6",
                "함량(%)": 50,
                "물질구분": "과거 분류",
                "고유번호": "OLD-001",
            }],
            "USER_CONFIRMED",
        )
        tables_mock.return_value = ({}, ("CAP_QTY_APP2",))
        form1_mock.return_value = SimpleNamespace(chemical_rows=())

        result = build_cap_chemical_legal_data(project)

        self.assertFalse(result.ready)
        self.assertNotEqual(result.rows[0].get("물질구분"), "과거 분류")
        self.assertNotEqual(result.rows[0].get("고유번호"), "OLD-001")
        self.assertTrue(result.blockers)

    @patch("engine.stage2.cap_baseline_docx.build_cap_form6_sds_data")
    def test_baseline_docx_uses_validated_form6_legal_identity_and_form8_rows(
        self, form6_mock
    ):
        project = self._project()
        project.set_field(
            "inventory.chemicals",
            "화학물질 목록",
            [{
                "물질명": "염소",
                "CAS 번호": "7782-50-5",
                "함량(%)": 99.9,
                "물리적 상태": "기체",
                "물질구분": "과거 회사 입력값",
                "고유번호": "OLD-999",
            }],
            "USER_CONFIRMED",
        )
        project.set_field(
            "cap.site.surrounding_environment",
            "사업장 주변 환경정보",
            [{
                "보호대상 없음 여부": "아니오",
                "보호대상 명칭": "○○초등학교",
                "보호대상 구분": "갑종",
                "세부유형": "교육·연구시설",
                "주소·위치": "○○시 ○○로 10",
                "사업장 경계와 거리(m)": 420,
                "GIS/현장 근거": "GIS-SITE-01",
            }],
            "USER_CONFIRMED",
        )
        form6_mock.return_value = CAPForm6SDSData(
            rows=({
                "물질명": "염소",
                "CAS 번호": "7782-50-5",
                "함량(%)": 99.9,
                "물리적 상태": "기체",
                "물질구분": "인체급성유해성물질 / 사고대비물질",
                "고유번호": "97-1-1",
                "법적분류 근거": "CAP_QTY_APP2 / CAP_QTY_APP3",
            },),
            blockers=(),
            messages=(),
        )

        data = build_cap_baseline_draft(project)
        doc = Document(BytesIO(data))
        form6_text = "\n".join(
            cell.text for row in doc.tables[15].rows for cell in row.cells
        )
        form8_text = "\n".join(
            cell.text
            for table_index in (17, 18)
            for row in doc.tables[table_index].rows
            for cell in row.cells
        )

        self.assertIn("인체급성유해성물질 / 사고대비물질", form6_text)
        self.assertIn("97-1-1", form6_text)
        self.assertNotIn("과거 회사 입력값", form6_text)
        self.assertNotIn("OLD-999", form6_text)
        self.assertIn("○○초등학교", form8_text)
        self.assertIn("교육·연구시설", form8_text)
        self.assertIn("420", form8_text)
        self.assertIn("500m 내 보호대상 1건", form8_text)


if __name__ == "__main__":
    unittest.main()
