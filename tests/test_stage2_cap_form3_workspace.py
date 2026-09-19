from __future__ import annotations

from io import BytesIO
from pathlib import Path
import unittest

from docx import Document

from engine.stage2 import cap_form2_workspace as f2
from engine.stage2 import cap_form3_workspace as f3
from engine.stage2 import cap_guideline as guide
from engine.stage2 import cap_workspace as ws
from engine.stage2.cap_baseline_docx import FORM_TABLE_INDEX, build_cap_baseline_draft
from tests import test_stage2_cap_form1_engine as form1_tests


def _project():
    return form1_tests.CAPForm1EngineTests()._project()


def _form3_text(project):
    table = Document(BytesIO(build_cap_baseline_draft(project))).tables[FORM_TABLE_INDEX["3"][0]]
    return "\n".join(cell.text for row in table.rows for cell in row.cells)


class CAPForm3WorkspaceTests(unittest.TestCase):
    def test_schema_matches_guideline_and_all_fields_have_help(self):
        schema = ws.load_form_schema(3)
        self.assertTrue(schema["title"].endswith(guide.form_guidelines()[3].title))
        self.assertTrue(all(item["help"] for item in f3.fields()))
        labels = {" ".join(row[0].split()) for row in guide.form_guidelines()[3].tables[0]}
        for asked in ("사업자 등록번호", "대표자", "우편번호/주소", "산업단지", "대표전화",
                      "공동비상대응계획 수립 여부", "유사제도 심사결과 활용", "총괄영향범위내 주민여부",
                      "최근 3년간 화학사고 발생 여부"):
            self.assertIn(asked, labels)

    def test_earlier_answers_are_auto_rows_not_questions(self):
        project = _project()
        f2.save_submission(project, "변경제출", "부적합", "염소 공장")
        auto = {row.label: row for row in f3.auto_rows(project)}
        self.assertEqual(auto["단위공장명"].value, "염소 공장")
        self.assertIn("변경제출", auto["제출구분"].value)
        asked = {item["label"] for item in f3.fields()}
        self.assertFalse(asked & {"사업장명", "단위공장명", "제출구분", "작성수준"})

    def test_answers_flow_into_the_docx(self):
        project = _project()
        f3.save_fields(project, {
            "registration_no": "123-45-67890", "representative": "김대표", "address": "(12345) 서울 어딘가 1",
            "joint": "단독제출", "other_review": "해당(공정안전보고서)", "recent_accident": "없음",
            "writer_department": "안전팀", "writer_name": "홍길동", "blank": "",
        })
        text = _form3_text(project)
        for expected in ("123-45-67890", "김대표", "서울 어딘가 1", "☒ 단독제출", "안전팀 홍길동", "☒ 없음"):
            self.assertIn(expected, text)
        self.assertIn("☒ 공정안전보고서", text)

    def test_missing_cells_are_listed_and_residents_is_optional(self):
        project = _project()
        self.assertIn("대표자", f3.missing_labels(project))
        self.assertNotIn("총괄영향범위내 주민여부", f3.missing_labels(project))

    def test_page_offers_form3(self):
        page = (Path(__file__).resolve().parents[1] / "ui/cap_workspace_page.py").read_text(encoding="utf-8")
        self.assertIn("cap_form3_view", page)


if __name__ == "__main__":
    unittest.main()
