from __future__ import annotations

from io import BytesIO
import unittest

from openpyxl import load_workbook

from engine.stage2.cap_form10_engine import build_cap_form10_data
from engine.stage2.project import Stage2Project
from engine.stage2.workbook_enhancements import build_enhanced_integrated_authoring_workbook


class CAPForm10EngineTests(unittest.TestCase):
    def _project(self) -> Stage2Project:
        project = Stage2Project(
            project_id="S2-FORM10",
            company_name="테스트화학",
            cap_required=True,
            cap_group="1군",
            scope_confirmed=True,
            cap_selected=True,
        )
        project.set_field(
            "inventory.chemicals",
            "화학물질 목록",
            [{
                "물질명": "톨루엔",
                "CAS 번호": "108-88-3",
                "함량(%)": 99.5,
                "물리적 상태": "액체",
                "최대보유량": 15000,
                "단위": "kg",
            }],
            "USER_CONFIRMED",
        )
        project.set_field(
            "inventory.facilities",
            "시설·설비 목록",
            [{
                "설비번호": "TK-101",
                "설비명": "톨루엔 저장탱크",
                "설비종류": "저장탱크",
                "취급물질": "톨루엔",
                "용량": 20,
                "용량단위": "m3",
                "설계압력": "0.49 MPa",
                "운전압력": "0.15 MPa",
                "설계온도": "80 ℃",
                "운전온도": "30 ℃",
                "최대보유량(kg)": 15000,
                "최대 연결구 크기(mm)": 50,
                "P&ID 번호": "PID-101",
            }],
            "USER_CONFIRMED",
        )
        return project

    def _set_calculation(self, project: Stage2Project, **overrides) -> None:
        row = {
            "적용여부": "예",
            "대상 설비번호": "TK-101",
            "설비형태": "저장탱크",
            "확산방지설비 종류": "방류벽",
            "필요용량(m3)": 22,
            "필요용량 기준·근거": "적용 시설기준 검토자료",
            "내부 길이(m)": 6,
            "내부 폭(m)": 5,
            "유효높이(m)": 0.8,
            "내부 차감용적(m3)": 1.2,
            "직접확인 유효용량(m3)": "",
            "비고": "",
        }
        row.update(overrides)
        project.set_field(
            "cap.safety.dike_calculation",
            "확산방지설비 계산자료",
            [row],
            "USER_CONFIRMED",
        )

    def test_geometric_capacity_and_review_result_are_calculated(self):
        project = self._project()
        self._set_calculation(project)

        result = build_cap_form10_data(project)

        self.assertEqual(result.blockers, ())
        row = result.rows[0]
        self.assertEqual(row["구분기호"], "TK-101")
        self.assertEqual(row["설비형태"], "저장탱크")
        self.assertEqual(row["설계용량"], "20")
        self.assertEqual(row["필요용량"], "22")
        self.assertEqual(row["유효용량"], "22.8")
        self.assertEqual(row["검토결과"], "적정")
        self.assertIn("길이×폭×유효높이", row["유효용량 산정근거"])

    def test_insufficient_effective_capacity_is_marked_without_guessing_required_capacity(self):
        project = self._project()
        self._set_calculation(project, **{"필요용량(m3)": 25})

        result = build_cap_form10_data(project)

        self.assertEqual(result.rows[0]["유효용량"], "22.8")
        self.assertEqual(result.rows[0]["검토결과"], "부족")

    def test_required_capacity_basis_is_mandatory(self):
        project = self._project()
        self._set_calculation(project, **{"필요용량 기준·근거": ""})

        result = build_cap_form10_data(project)

        self.assertTrue(any("필요용량의 적용 기준·근거" in blocker for blocker in result.blockers))

    def test_blank_deduction_is_not_assumed_zero(self):
        project = self._project()
        self._set_calculation(project, **{"내부 차감용적(m3)": ""})

        result = build_cap_form10_data(project)

        self.assertTrue(any("차감용적" in blocker for blocker in result.blockers))
        self.assertEqual(result.rows[0]["검토결과"], "")

    def test_direct_and_geometric_capacity_mismatch_is_blocking(self):
        project = self._project()
        self._set_calculation(project, **{"직접확인 유효용량(m3)": 30})

        result = build_cap_form10_data(project)

        self.assertTrue(any("서로 다릅니다" in blocker for blocker in result.blockers))
        self.assertEqual(result.rows[0]["유효용량"], "30")

    def test_facility_type_mismatch_is_blocking(self):
        project = self._project()
        self._set_calculation(project, **{"설비형태": "반응기"})

        result = build_cap_form10_data(project)

        self.assertTrue(any("03_설비정보의 설비종류" in blocker for blocker in result.blockers))
        self.assertEqual(result.rows[0]["설비형태"], "저장탱크")

    def test_explicit_global_not_applicable_is_accepted(self):
        project = self._project()
        project.set_field(
            "cap.safety.dike_calculation",
            "확산방지설비 계산자료",
            [{"적용여부": "해당 없음", "비고": "확산방지설비 적용대상 없음"}],
            "USER_CONFIRMED",
        )

        result = build_cap_form10_data(project)

        self.assertEqual(result.blockers, ())
        self.assertEqual(result.rows[0]["검토결과"], "해당 없음")

    def test_workbook_separates_calculation_table_from_layout_attachment(self):
        project = self._project()
        data = build_enhanced_integrated_authoring_workbook(project, example=False)
        wb = load_workbook(BytesIO(data), data_only=False)

        self.assertIn("21_확산방지설비", wb.sheetnames)
        headers = [str(cell.value or "") for cell in wb["21_확산방지설비"][4]]
        self.assertIn("필요용량 기준·근거", headers)
        self.assertIn("내부 차감용적(m3)", headers)

        meta = wb["_시스템정보"]
        table_targets = []
        attachment_targets = []
        for row in meta.iter_rows(min_row=8, values_only=True):
            if row[0] == "TABLE":
                table_targets.append(str(row[2] or ""))
            if row[0] == "ATTACHMENT":
                attachment_targets.append(str(row[2] or ""))
        self.assertTrue(any("cap.safety.dike_calculation" in value for value in table_targets))
        self.assertFalse(any("cap.safety.dike_layout" in value for value in table_targets))
        self.assertTrue(any("cap.safety.dike_layout" in value for value in attachment_targets))


if __name__ == "__main__":
    unittest.main()
