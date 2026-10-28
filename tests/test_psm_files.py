"""공정안전보고서도 회사 파일로: 표 서식·사업개요·행정서식·서술형 사실을 파일에서 읽어 초안에 채웁니다."""
from __future__ import annotations

import unittest
from io import BytesIO

from docx import Document
from openpyxl import Workbook, load_workbook

from engine.stage2 import cap_fact_catalog as catalog
from engine.stage2 import cap_fact_status as fs
from engine.stage2 import cap_upload_template as tpl
from engine.stage2 import evidence_workspace as ew
from engine.stage2 import psm_admin_forms as admin
from engine.stage2 import psm_form12_workspace as f12
from engine.stage2 import psm_table_workspace as pt
from engine.stage2 import statutory_report as report
from engine.stage2.project import Stage2Project


def _psm_project(cap=False):
    project = Stage2Project(project_id="PSM-FILES", company_name="가상 공정화학", psm_required=True,
                            cap_required=cap or None, cap_group="2군" if cap else "")
    project.set_authoring_scope(psm_selected=True, cap_selected=cap)
    project.set_field("business.address", "주소", "가상시 가상로 1", "USER_CONFIRMED")
    return project


def _book(sheets):
    book = Workbook()
    book.remove(book.active)
    for title, rows in sheets.items():
        sheet = book.create_sheet(title)
        for row in rows:
            sheet.append(list(row))
    out = BytesIO()
    book.save(out)
    return out.getvalue()


def _sample(item, row):
    if item.options:
        return item.options[0]
    if item.format == "number":
        return str(row + 1)
    return f"가상 {item.name} {row}"


def _fill(content):
    """양식의 빈칸을 형식에 맞는 값으로 모두 채우고, 표 시트에는 두 줄을 적습니다."""
    book = load_workbook(BytesIO(content))
    site = {item.template_header: item for item in catalog.facts(kinds=["site", "narrative"])}
    tables = {tpl.sheet_title(t): items for t, items in catalog.psm_tables().items()}
    for title in book.sheetnames:
        sheet = book[title]
        if title == tpl.SITE_SHEET:
            for r in range(2, sheet.max_row + 1):
                item = site[sheet.cell(row=r, column=1).value]
                sheet.cell(row=r, column=2, value="해당 없음" if item.column.endswith("#근거") else _sample(item, r))
        elif title in tables:
            items = {i.template_header: i for i in tables[title]}
            headers = [c.value for c in sheet[1] if c.value != "예시 (입력하지 마세요)"]
            for r in (2, 3):
                for c, header in enumerate(headers, start=1):
                    sheet.cell(row=r, column=c, value=_sample(items[header], r))
    out = BytesIO()
    book.save(out)
    return out.getvalue()


class TemplateRoundTripTests(unittest.TestCase):
    def test_filled_psm_template_fills_every_table_and_form12(self):
        project = _psm_project()
        content = tpl.build(project)
        names = load_workbook(BytesIO(content)).sheetnames
        self.assertIn(tpl.SITE_SHEET, names)
        self.assertIn("동력기계 목록", names)
        self.assertEqual(tpl.filename(project), tpl.PSM_FILENAME)
        filled = _fill(content)
        self.assertEqual(ew.classify_document(tpl.PSM_FILENAME, filled), "mixed")
        ew.ingest(project, tpl.PSM_FILENAME, filled, "mixed")
        doc = ew.documents(project)[0]
        self.assertEqual(doc["unmapped_columns"], [])
        self.assertEqual(doc["rows_without_id"], 0)
        ew.publish(project, replace_existing=True)
        for table, spec in pt.SPECS.items():
            if table.startswith("cap-") or table == "17-5":
                continue
            self.assertEqual(pt.needs(project, table), [], table)
            self.assertEqual(len(pt.rows(project, table)), 2, table)
            self.assertEqual(project.get_field(spec.key).status, "FILE_READ", table)
        self.assertTrue(f12.current(project)["대상 유해·위험설비"].startswith("가상 대상 유해·위험설비"))
        self.assertEqual(admin.form1_values(project)["사업장관리번호"][:2], "가상")
        rows = {r["_id"]: r for r in fs.fact_rows(project)}
        self.assertTrue(all(rows[i.id]["_state"] == fs.O for items in catalog.psm_tables().values() for i in items),
                        [(i, r["고칠 점"]) for i, r in rows.items() if r["_kind"] == "psm_table" and r["_state"] != fs.O])

    def test_empty_project_shows_one_line_per_table(self):
        project = _psm_project()
        rows = [r for r in fs.fact_rows(project) if r["_kind"] == "psm_table"]
        self.assertEqual(len(rows), len(catalog.psm_tables()))
        self.assertTrue(all(r["_state"] == fs.BLANK and "표를 올려 주세요" in r["고칠 점"] for r in rows))
        self.assertTrue(all("공정안전보고서" in r["쓰이는 곳"] for r in fs.fact_rows(project)))


class ReadingTests(unittest.TestCase):
    def test_tables_are_recognised_by_their_headers(self):
        content = _book({
            "A": [("설치지역", "소화기", "옥내소화전"), ("A동", "10", "2")],
            "B": [("설치지역", "자동화재탐지설비", "비상방송설비"), ("A동", "1", "1")],
            "C": [("분류코드", "유체의 명칭 또는 구분", "배관재질", "비고"), ("A1A", "염소", "SUS316L", "P&ID-1")],
        })
        project = _psm_project()
        ew.ingest(project, "psm.xlsx", content, ew.classify_document("psm.xlsx", content))
        kinds = [t["kind"] for t in ew.documents(project)[0]["tables"]]
        self.assertEqual(kinds, ["psm:17-3", "psm:17-4", "psm:16"])
        ew.publish(project, replace_existing=True)
        self.assertEqual(pt.rows(project, "16")[0]["배관재질"], "SUS316L")
        self.assertEqual(pt.applicability(project, "17-3")[0], pt.APPLICABLE)       # 표를 올리면 ‘적용’

    def test_not_applicable_is_read_from_the_site_file(self):
        project = _psm_project()
        content = _book({"개요": [("항목", "값"), ("내화구조 명세 적용여부", "해당 없음"),
                                 ("내화구조 명세 해당 없음 근거", "내화 대상 구조물 없음(현장 확인)")]})
        ew.ingest(project, "개요.xlsx", content, ew.classify_document("개요.xlsx", content))
        ew.publish(project, replace_existing=True)
        self.assertEqual(pt.applicability(project, "18"), (pt.NOT_APPLICABLE, "내화 대상 구조물 없음(현장 확인)"))
        self.assertEqual(pt.needs(project, "18"), [])
        row = next(r for r in fs.fact_rows(project) if r["_kind"] == "psm_table" and r["_id"].startswith("psm.t18."))
        self.assertEqual((row["_state"], row["값"]), (fs.O, "해당 없음"))

    def test_rows_without_the_key_column_are_skipped_and_reported(self):
        project = _psm_project()
        content = _book({"기계": [("기계번호", "기계명", "주요재질"), ("P-1", "펌프", "SUS"), ("", "합계", "")]})
        ew.ingest(project, "m.xlsx", content, ew.classify_document("m.xlsx", content))
        self.assertEqual(ew.documents(project)[0]["rows_without_id"], 1)

    def test_same_machine_in_two_files_with_different_values_is_a_conflict(self):
        project = _psm_project()
        for name, material in (("a.xlsx", "SUS316"), ("b.xlsx", "SS400")):
            content = _book({"기계": [("기계번호", "기계명", "주요재질"), ("P-1", "펌프", material)]})
            ew.ingest(project, name, content, ew.classify_document(name, content))
        row = next(r for r in fs.fact_rows(project) if r["_id"] == "psm.t14.주요재질")
        self.assertEqual(row["_state"], fs.CONFLICT)
        ew.publish(project, replace_existing=True)
        self.assertEqual(pt.rows(project, "14")[0]["주요재질"], "")                  # 충돌한 값은 넣지 않음

    def test_detectors_and_psm_only_equipment_columns(self):
        project = _psm_project(cap=True)
        content = _book({
            "설비": [("설비번호", "설비명", "용량(m3)", "본체재질", "적용 법령·검사 여부"), ("TK-1", "탱크", "10", "SUS304", "해당 없음")],
            "감지기": [("감지기 번호", "검출대상 물질", "설치위치", "경보 설정값", "연동 설비·조치"),
                    ("GD-1", "염소", "TK-1", "0.5 ppm", "차단밸브 닫힘")],
        })
        ew.ingest(project, "f.xlsx", content, ew.classify_document("f.xlsx", content))
        ew.publish(project, replace_existing=True)
        alarm = pt.rows(project, "17-5")[0]
        self.assertEqual((alarm["감지기번호"], alarm["경보설정값"], alarm["경보시 조치내용"]), ("GD-1", "0.5 ppm", "차단밸브 닫힘"))
        self.assertEqual(pt.applicability(project, "17-5")[0], pt.APPLICABLE)
        form15 = report._psm_form15_rows(project)
        self.assertIn("SUS304", form15[0])
        self.assertIn("해당 없음", form15[0])


class NarrativeTests(unittest.TestCase):
    def test_psm_narrative_uses_company_text_or_file_facts_only(self):
        from engine.stage2 import cap_fact_drafting as fd
        from engine.stage2 import report_export as rx

        project = _psm_project()
        content = _book({"계획": [("항목", "값"), ("안전운전지침서", "별첨 SOP-01~12 안전운전지침서 참조"),
                                 ("근로자 교육 대상", "공정 운전원 전원"), ("비상 총괄", "공장장")]})
        ew.ingest(project, "plan.xlsx", content, ew.classify_document("plan.xlsx", content))
        ew.publish(project, replace_existing=True)
        self.assertEqual(project.get_field("psm.facts.emergency_org").value, {"총괄": "공장장"})   # 두 문서가 같은 사실
        items = {i.key: i for i in fd.items(project, "PSM")}
        self.assertEqual(items["psm.operation.sop"].state, fd.FROM_FILE)
        self.assertEqual(items["psm.operation.training"].state, fd.CAN_DRAFT)
        self.assertGreater(fd.refresh_rule_drafts(project), 0)
        training = next(i for i in fd.items(project, "PSM") if i.key == "psm.operation.training")
        self.assertIn("근로자 교육 대상: 공정 운전원 전원", training.draft)
        self.assertIn("[확인 필요: 근로자 교육 주기]", training.draft)
        review = Document(BytesIO(rx.build_outputs(project, "PSM").review_docx))
        text = "\n".join(par.text for par in review.paragraphs)
        self.assertIn("[AI 초안", text)
        self.assertIn("근로자 교육 대상: 공정 운전원 전원", text)
        self.assertIn("별첨 SOP-01~12", text)


class DocxTests(unittest.TestCase):
    def test_psm_draft_contains_file_values(self):
        from engine.stage2.psm_baseline_docx import build_psm_baseline_draft

        project = _psm_project()
        content = _book({
            "기계": [("기계번호", "기계명", "주요재질", "전동기용량", "방호·보호장치 종류"), ("P-101", "원료펌프", "SUS316", "7.5", "EOCR")],
            "개요": [("항목", "값"), ("대상 유해·위험설비", "염소 저장·투입 설비"), ("한국표준산업분류", "20119 기타 기초무기화학물질 제조업")],
        })
        ew.ingest(project, "psm.xlsx", content, ew.classify_document("psm.xlsx", content))
        ew.publish(project, replace_existing=True)
        doc = Document(BytesIO(build_psm_baseline_draft(project)))
        text = "\n".join(c.text for t in doc.tables for row in t.rows for c in row.cells)
        for value in ("P-101", "원료펌프", "EOCR", "염소 저장·투입 설비", "20119"):
            self.assertIn(value, text)


if __name__ == "__main__":
    unittest.main()
