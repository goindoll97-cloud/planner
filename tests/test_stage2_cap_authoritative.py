from __future__ import annotations

from io import BytesIO
import unittest

from docx import Document

from engine.stage2.cap_authoritative import (
    NEVER_KOSHA_FORM6_FIELDS,
    cap_form6_msds_candidates,
    cap_form7_reference_candidates,
    load_cap_source_lock,
)
from engine.stage2.msds_reference import reference_field_key
from engine.stage2.project import Stage2Project
from engine.stage2.report_draft import build_report_draft
from engine.stage2.statutory_report import CAP_FORMS


class Stage2CAPAuthoritativeSourceTests(unittest.TestCase):
    def _project(self, *, cap_group: str = "1군") -> Stage2Project:
        project = Stage2Project(
            project_id="S2-CAP-SOURCE-TEST",
            company_name="테스트화학",
            site_name="테스트공장",
            psm_required=False,
            cap_required=True,
            cap_group=cap_group,
            scope_confirmed=True,
            psm_selected=False,
            cap_selected=True,
        )
        project.set_field("business.company_name", "회사명", "테스트화학", "VERIFIED")
        project.set_field("business.address", "사업장 소재지", "울산광역시 테스트로 1", "VERIFIED")
        project.set_field(
            "inventory.chemicals",
            "화학물질 목록",
            [{
                "물질명": "시험물질",
                "CAS 번호": "108-88-3",
                "물질구분": "사고대비물질",
                "고유번호": "97-1-001",
                "함량(%)": "99.5",
                "비중": "0.866",
            }],
            "VERIFIED",
        )
        project.set_field("cap.business.writing_level", "작성수준", f"{cap_group} 사업장", "VERIFIED")
        payload = {
            "reference_status": "REFERENCE_READY",
            "cas": "108-88-3",
            "chemical_name": "톨루엔",
            "chem_id": "TEST-001",
            "source": "한국산업안전보건공단 물질안전보건자료 조회 서비스",
            "source_dataset_id": "15157612",
            "checked_at_utc": "2026-09-14T00:00:00+00:00",
            "message": "test fixture",
            "sections": {
                "2": {"items": [["유해성·위험성 분류", "급성 독성(흡입) : 구분 2 | 금속부식성 물질 : 구분 1"]]},
                "5": {"items": [["화재 위험", "증기는 공기와 폭발성 혼합물을 형성할 수 있음"]]},
                "8": {"items": [["위험노출수준", "AEGL-2 60 ppm; ERPG-2 50 ppm"], ["노출기준", "TWA : 10 ppm"]]},
                "9": {"items": [
                    ["물리적 상태", "액체"],
                    ["비중", "0.87"],
                    ["폭발한계 하한", "1.2 %"],
                    ["폭발한계 상한", "7.1 %"],
                    ["증기압", "28.4 mmHg (20℃)"],
                ]},
                "10": {"items": [["반응성", "강산화제와 반응 가능"]]},
                "11": {"items": [["흡입", "증기 흡입 시 중추신경계 영향 가능"]]},
                "12": {"items": [["수생환경유해성", "수생생물에 유해할 수 있음"]]},
                "15": {"items": [["법적규제 현황", "사고대비물질 / 고유번호 99-9-999"]]},
            },
        }
        project.set_field(
            reference_field_key("108-88-3"),
            "KOSHA MSDS 참고자료 · 108-88-3",
            payload,
            "HOLD",
        )
        return project

    def test_source_lock_matches_user_supplied_current_sources(self):
        source = load_cap_source_lock()
        legal = source["legal_structure"]
        manual = source["writing_guidance"]
        self.assertEqual(legal["effective_date"], "2026-04-22")
        self.assertEqual(legal["notice"], "화학물질안전원고시 제2026-07호")
        self.assertEqual(
            legal["national_law_center_text_sha256"],
            "312117944d2c93eb8e8a538d7e090662754beb2dbe820b2f042351de38c62466",
        )
        self.assertEqual(
            legal["national_law_center_annex_forms_sha256"],
            "f41c3b26c72fb3e0186d5d25b99004ed7a8d3644527c12844a6a7f2c36112d2d",
        )
        self.assertEqual(manual["document_code"], "NICS-GP2026-8")
        self.assertEqual(
            manual["sha256"],
            "e6c55a0a1d87e97ab85afe7afc608abfc773540390e01f29eeac38f364ba2565",
        )

    def test_form6_headers_follow_current_annex_form(self):
        self.assertEqual(
            CAP_FORMS["6"].headers,
            (
                "연번", "유해화학물질명", "물질구분", "화학물질식별번호(CAS 번호)", "고유번호", "물질상태", "함량(%)", "비중",
                "폭발한계 하한(%)", "폭발한계 상한(%)", "독성구분-항목", "독성구분-구분", "위험노출수준", "허용농도값",
                "증기압(20℃, mmHg)", "부식성(유, 무)",
            ),
        )

    def test_kosha_candidates_fill_only_missing_review_fields(self):
        project = self._project()
        candidates = cap_form6_msds_candidates(project)
        by_field = {item.field: item for item in candidates}

        # Company-entered density must win over the KOSHA 0.87 reference.
        self.assertNotIn("비중", by_field)
        self.assertEqual(by_field["물질상태"].value, "액체")
        self.assertEqual(by_field["폭발한계 하한(%)"].value, "1.2 %")
        self.assertEqual(by_field["폭발한계 상한(%)"].value, "7.1 %")
        self.assertEqual(by_field["증기압(20℃, mmHg)"].value, "28.4 mmHg (20℃)")
        self.assertEqual(by_field["독성구분-항목"].value, "급성 독성(흡입)")
        self.assertEqual(by_field["독성구분-구분"].value, "구분 2")
        self.assertIn("ERPG", by_field["위험노출수준"].value)
        self.assertEqual(by_field["부식성(유, 무)"].value, "유")

        self.assertTrue(set(NEVER_KOSHA_FORM6_FIELDS).isdisjoint(by_field))
        self.assertEqual(project.get_field(reference_field_key("108-88-3")).status, "HOLD")

    def test_absence_of_corrosivity_never_becomes_no(self):
        project = self._project()
        payload = dict(project.get_field(reference_field_key("108-88-3")).value)
        sections = dict(payload["sections"])
        sections["2"] = {"items": [["유해성·위험성 분류", "급성 독성(흡입) : 구분 2"]]}
        sections["10"] = {"items": [["안정성", "통상적인 조건에서 안정함"]]}
        payload["sections"] = sections
        project.set_field(reference_field_key("108-88-3"), "KOSHA MSDS 참고자료", payload, "HOLD")
        by_field = {item.field: item for item in cap_form6_msds_candidates(project)}
        self.assertNotIn("부식성(유, 무)", by_field)

    def test_form7_reference_does_not_invent_representative_selection_reason(self):
        rows = cap_form7_reference_candidates(self._project())
        self.assertEqual(len(rows), 1)
        self.assertIn("확인 필요", rows[0]["선정 사유"])
        self.assertIn("증기 흡입", rows[0]["인체유해성 후보"])
        self.assertIn("강산화제", rows[0]["물리적 위험성 후보"])
        self.assertIn("수생생물", rows[0]["환경유해성 후보"])

    def test_cap_report_starts_with_form1_and_separates_candidates(self):
        data = build_report_draft(self._project(), "CAP")
        doc = Document(BytesIO(data))
        text = "\n".join(
            [p.text for p in doc.paragraphs]
            + [cell.text for table in doc.tables for row in table.rows for cell in row.cells]
        )
        first_visible = next(p.text.strip() for p in doc.paragraphs if p.text.strip())

        # The draft opens the real regulation-form baseline as-is (its own
        # title page, not a synthetic cover this program generated).
        self.assertNotIn("화학물질안전원고시 제2026-07호", first_visible)
        self.assertNotIn("국가법령정보센터 별표·별지 서식", first_visible)
        self.assertIn("NICS-GP2026-8", text)
        self.assertIn("별지 제6호 자동입력 후보 검토", text)
        self.assertIn("법정서식 외 검토자료", text)
        self.assertIn("제품 SDS/증빙 확인 후 회사자료에 확정 입력", text)
        self.assertIn("별지 제7호 유해성정보 후보 검토", text)


if __name__ == "__main__":
    unittest.main()
