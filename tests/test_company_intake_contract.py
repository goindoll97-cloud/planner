from __future__ import annotations

from io import BytesIO
import unittest

from openpyxl import load_workbook

from engine.company_intake_contract import AI_DRAFT_EXCLUDED_LABELS
from engine.template import build_minimal_input_workbook
from engine.inventory import read_intake_workbook
from engine.stage2.project import create_project_from_stage1_snapshot


class CompanyIntakeContractTests(unittest.TestCase):
    def test_initial_workbook_contains_company_only_final_report_facts(self):
        workbook = load_workbook(BytesIO(build_minimal_input_workbook()), data_only=False)
        ws = workbook["01_사업장기본정보"]
        labels = [ws.cell(row, 1).value for row in range(4, ws.max_row + 1)]

        for expected in [
            "사업자등록번호",
            "대표자 성명",
            "대표전화",
            "예상근무 근로자수",
            "전기계약용량(kW)",
            "PSM 사업 구분",
            "PSM 심사대상 설비명",
            "PSM 부지면적(㎡)",
            "PSM 보고서 작성자 성명",
            "PSM 보고서 작성자 자격",
            "PSM 총사업기간",
            "PSM 착공예정일",
            "PSM 시운전기간",
            "CAP 단위공장명",
            "CAP 산업단지명",
            "CAP 제출구분",
            "CAP 공동비상대응계획 수립 여부",
            "CAP 최근 3년간 화학사고 발생 여부",
            "CAP 작성자 성명",
            "CAP 담당자 연락처",
            "CAP 담당자 이메일",
        ]:
            self.assertIn(expected, labels)

        for excluded in AI_DRAFT_EXCLUDED_LABELS:
            self.assertNotIn(excluded, labels)

        self.assertIn("AI/시스템이 작성·계산", str(ws["A2"].value))

    def test_workbook_parser_keeps_new_business_facts(self):
        intake = read_intake_workbook(build_minimal_input_workbook())
        self.assertIn("사업자등록번호", intake.business)
        self.assertIn("PSM 심사대상 설비명", intake.business)
        self.assertIn("CAP 담당자 이메일", intake.business)

    def test_stage2_carries_company_facts_without_inference(self):
        snapshot = {
            "source_fingerprint": "a" * 64,
            "business": {
                "사업장명": "테스트공장",
                "사업장 주소": "울산광역시 테스트구 1",
                "사업자등록번호": "111-22-33333",
                "대표자 성명": "홍길동",
                "업종 또는 주요 생산품": "정밀화학 중간체",
                "한국표준산업분류(KSIC) 코드": "20111",
                "대표전화": "052-111-2222",
                "팩스번호": "052-111-2223",
                "예상근무 근로자수": 88,
                "전기계약용량(kW)": 2100,
                "PSM 사업 구분": "변경",
                "PSM 심사대상 설비명": "R-101 반응공정",
                "PSM 부지면적(㎡)": 12345,
                "PSM 주요건물(동/층/연면적)": "2동/3층/5,000㎡",
                "PSM 보고서 작성자 성명": "김안전",
                "PSM 보고서 작성자 자격": "산업안전기사",
                "PSM 총사업기간": "2026.10~2027.06",
                "PSM 착공예정일": "2026-10-15",
                "PSM 시운전기간": "2027.05~2027.06",
                "CAP 단위공장명": "제1생산공장",
                "CAP 산업단지명": "울산미포국가산업단지",
                "CAP 제출구분": "신규",
                "CAP 공동비상대응계획 수립 여부": "N",
                "CAP 유사제도 심사결과 활용 여부": "Y",
                "CAP 최근 3년간 화학사고 발생 여부": "N",
                "CAP 작성자 성명": "이환경",
                "CAP 담당자 연락처": "010-1111-2222",
                "CAP 담당자 이메일": "safe@example.com",
            },
            "documents": {},
            "chemicals": [{"제품명": "물질A", "CAS No.": "50-00-0"}],
            "facilities": [],
            "decision": {
                "psm_status": "공정안전보고서 제출 대상",
                "cap_status": "작성수준 — 2군 사업장",
            },
        }
        project = create_project_from_stage1_snapshot(snapshot)

        self.assertEqual(project.get_field("business.registration_no").value, "111-22-33333")
        self.assertEqual(project.get_field("business.representative").value, "홍길동")
        self.assertEqual(project.get_field("business.main_products").value, "정밀화학 중간체")
        self.assertEqual(project.get_field("business.employee_count").value, 88)
        self.assertEqual(project.get_field("psm.business.project_type").value, "변경")
        self.assertEqual(project.get_field("psm.business.target_facility").value, "R-101 반응공정")
        self.assertEqual(project.get_field("psm.business.site_area").value, 12345)
        self.assertEqual(project.get_field("psm.business.start_date").value, "2026-10-15")
        self.assertEqual(
            project.get_field("psm.business.writer_info").value,
            {"작성자": "김안전", "작성자 자격": "산업안전기사"},
        )
        self.assertEqual(project.get_field("cap.business.registration_no").value, "111-22-33333")
        self.assertEqual(project.get_field("cap.business.unit_plant_name").value, "제1생산공장")
        self.assertEqual(project.get_field("cap.business.recent_accident").value, "N")
        self.assertEqual(project.get_field("cap.business.writer_email").value, "safe@example.com")
        self.assertEqual(project.get_field("cap.business.writer_email").status, "VERIFIED")

    def test_unknown_company_fact_stays_hold(self):
        snapshot = {
            "source_fingerprint": "b" * 64,
            "business": {
                "사업장명": "테스트공장",
                "사업장 주소": "테스트시 1",
                "CAP 산업단지명": "모름",
            },
            "documents": {},
            "chemicals": [],
            "facilities": [],
            "decision": {"psm_status": "", "cap_status": ""},
        }
        project = create_project_from_stage1_snapshot(snapshot)
        record = project.get_field("cap.business.industrial_complex")
        self.assertIsNotNone(record)
        self.assertEqual(record.status, "HOLD")
        self.assertEqual(record.value, "모름")


if __name__ == "__main__":
    unittest.main()
