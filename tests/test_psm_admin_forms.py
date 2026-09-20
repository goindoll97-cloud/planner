from __future__ import annotations

from io import BytesIO
import unittest

from docx import Document

from engine.stage2.project import Stage2Project
from engine.stage2 import psm_admin_forms as forms


class PSMAdminFormsTests(unittest.TestCase):
    def _project(self) -> Stage2Project:
        p = Stage2Project(
            project_id="S2-ADMIN",
            company_name="한빛화학(주)",
            site_name="한빛화학 울산공장",
            psm_required=True,
            scope_confirmed=True,
            psm_selected=True,
        )
        for key, label, value in (
            ("business.registration_no", "사업자등록번호", "123-45-67890"),
            ("business.workplace_management_no", "사업장관리번호", "12345678901"),
            ("business.phone", "대표전화", "052-123-4567"),
            ("business.address", "사업장 소재지", "울산광역시 남구 산업로 1"),
            ("business.representative", "대표자 성명", "홍길동"),
            ("psm.admin.form1.application_date", "심사신청일", "2026-09-20"),
        ):
            p.set_field(key, label, value, "USER_CONFIRMED")
        return p

    def test_form1_reuses_shared_company_facts(self):
        p = self._project()

        ready = forms.form1_readiness(p)

        self.assertTrue(ready.ready)
        self.assertEqual(ready.values["사업자명"], "한빛화학(주)")
        self.assertEqual(ready.values["사업장관리번호"], "12345678901")
        self.assertEqual(ready.values["소재지"], "울산광역시 남구 산업로 1")

    def test_form1_does_not_invent_application_date(self):
        p = self._project()
        p.fields.pop("psm.admin.form1.application_date")

        ready = forms.form1_readiness(p)

        self.assertFalse(ready.ready)
        self.assertTrue(any("신청일" in blocker for blocker in ready.blockers))
        with self.assertRaises(ValueError):
            forms.build_form1_docx(p)

    def test_form1_docx_contains_confirmed_values(self):
        p = self._project()

        doc = Document(BytesIO(forms.build_form1_docx(p)))
        text = "\n".join(
            [para.text for para in doc.paragraphs]
            + [cell.text for table in doc.tables for row in table.rows for cell in row.cells]
        )

        self.assertIn("공정안전보고서 심사신청서", text)
        self.assertIn("한빛화학(주)", text)
        self.assertIn("12345678901", text)
        self.assertIn("2026-09-20", text)

    def test_form9_is_independent_and_requires_post_review_facts(self):
        p = self._project()

        ready = forms.form9_readiness(p)

        self.assertFalse(ready.ready)
        self.assertTrue(any("심사완료일" in blocker for blocker in ready.blockers))
        self.assertTrue(any("확인요청일" in blocker for blocker in ready.blockers))

    def test_form9_can_be_completed_without_changing_form1(self):
        p = self._project()
        values = {
            "psm.admin.contact_name": "김안전",
            "psm.admin.contact_mobile": "010-1234-5678",
            "psm.admin.contact_email": "safety@example.com",
            "psm.admin.confirmation_target": "반응·정제 공정",
            "psm.admin.review_completion_date": "2026-10-01",
            "psm.admin.construction_period": "2026-10-10~2027-03-31",
            "psm.admin.confirmation_request_date": "2027-04-01",
            "psm.admin.confirmation_period_start": "2027-04-10",
            "psm.admin.confirmation_period_end": "2027-04-20",
            "psm.admin.form9.application_date": "2027-04-01",
        }
        forms.save_admin_values(p, values)

        self.assertTrue(forms.form1_readiness(p).ready)
        ready = forms.form9_readiness(p)
        self.assertTrue(ready.ready)

        doc = Document(BytesIO(forms.build_form9_docx(p)))
        text = "\n".join(
            [para.text for para in doc.paragraphs]
            + [cell.text for table in doc.tables for row in table.rows for cell in row.cells]
        )
        self.assertIn("공정안전보고서확인요청서", text)
        self.assertIn("김안전", text)
        self.assertIn("2027-04-10 ~ 2027-04-20", text)

    def test_admin_docx_is_deterministic(self):
        p = self._project()
        self.assertEqual(forms.build_form1_docx(p), forms.build_form1_docx(p))


if __name__ == "__main__":
    unittest.main()
