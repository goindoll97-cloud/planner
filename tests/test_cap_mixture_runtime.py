from __future__ import annotations

from io import BytesIO
from types import SimpleNamespace
import unittest

import pandas as pd
from openpyxl import load_workbook

from engine.cap_mixture_runtime import (
    COMPONENT_SHEET,
    FACILITY_COMPONENT_CAS_COLUMN,
    MIXTURE_FLAG_COLUMN,
    _prepare_mixture_facilities,
    _screen_component_direct,
)
from engine.inventory import IntakeData, _mixture_validation_issues
from engine.template import build_minimal_input_workbook


class CAPMixtureRuntimeTests(unittest.TestCase):
    def _intake(self) -> IntakeData:
        chemicals = pd.DataFrame([
            {
                "No.": 1,
                "제품명": "혼합제품 A",
                "CAS No.": "108-88-3",
                "물질명(알면 입력)": "톨루엔(주성분)",
                "함량(%)": 60,
                "취급형태": "사용+저장",
                "최대 제조·사용량": 1000,
                "최대 저장량": 5000,
                "수량 단위": "kg",
                "최대 동시보유량(알면 입력)": 5000,
                MIXTURE_FLAG_COLUMN: "Y",
            }
        ])
        data = IntakeData(business={}, chemicals=chemicals, documents={})
        data.mixture_components = pd.DataFrame([
            {
                "적용여부": "해당", "제품목록행번호": 1, "제품명(확인용)": "혼합제품 A",
                "구성성분명": "톨루엔", "CAS No.": "108-88-3", "함량(%)": 60,
                "함량 최저(%)": None, "함량 최고(%)": None, "SDS 제3항 근거": "제품 SDS 제3항", "비고": "",
            },
            {
                "적용여부": "해당", "제품목록행번호": 1, "제품명(확인용)": "혼합제품 A",
                "구성성분명": "메탄올", "CAS No.": "67-56-1", "함량(%)": 10,
                "함량 최저(%)": None, "함량 최고(%)": None, "SDS 제3항 근거": "제품 SDS 제3항", "비고": "",
            },
        ])
        return data

    def test_generated_template_has_component_contract(self) -> None:
        wb = load_workbook(BytesIO(build_minimal_input_workbook()), data_only=False)
        self.assertIn(COMPONENT_SHEET, wb.sheetnames)
        chem = wb["02_화학물질목록"]
        headers = [chem.cell(3, c).value for c in range(1, chem.max_column + 1)]
        self.assertIn(MIXTURE_FLAG_COLUMN, headers)
        facilities = wb["04_시설별최대보유량"]
        facility_headers = [facilities.cell(3, c).value for c in range(1, facilities.max_column + 1)]
        self.assertIn(FACILITY_COMPONENT_CAS_COLUMN, facility_headers)

    def test_marked_mixture_requires_component_rows(self) -> None:
        data = self._intake()
        data.mixture_components = pd.DataFrame()
        issues = _mixture_validation_issues(data)
        self.assertTrue(any("구성성분을 한 줄 이상" in issue for issue in issues))

    def test_storage_uses_component_pct_without_reducing_facility_mass(self) -> None:
        data = self._intake()
        data.facilities = pd.DataFrame([
            {
                "목록행번호": 1, "시설명": "T-101", "시설유형": "저장탱크",
                "제외시설여부": "N", "제외사유": "해당없음", "물질성상": "액체",
                "공정유형": "변화없음", "설계용량": 5, "용량단위": "m3",
                "비중 또는 밀도(kg/L=ton/m3)": 0.9,
            }
        ])
        component = data.mixture_components.iloc[1]
        prepared, blockers = _prepare_mixture_facilities(data, 1, component, [{"content_threshold_pct": 5}])
        self.assertEqual(blockers, [])
        self.assertEqual(len(prepared), 1)
        self.assertEqual(float(prepared.iloc[0]["별표4 기준함량(%)"]), 10.0)
        self.assertEqual(float(prepared.iloc[0]["설계용량"]), 5.0)

    def test_multicomponent_reaction_requires_component_cas(self) -> None:
        data = self._intake()
        data.facilities = pd.DataFrame([
            {
                "목록행번호": 1, "시설명": "R-101", "시설유형": "제조·사용시설",
                "제외시설여부": "N", "제외사유": "해당없음", "물질성상": "액체",
                "공정유형": "반응", "별표4 기준함량(%)": 10, "함량근거": "공정배합표",
                "설계용량": 3, "용량단위": "m3", "비중 또는 밀도(kg/L=ton/m3)": 1.0,
                FACILITY_COMPONENT_CAS_COLUMN: "",
            }
        ])
        component = data.mixture_components.iloc[1]
        prepared, blockers = _prepare_mixture_facilities(data, 1, component, [{"content_threshold_pct": 5}])
        self.assertTrue(prepared.empty)
        self.assertTrue(any(FACILITY_COMPONENT_CAS_COLUMN in blocker for blocker in blockers))

    def test_range_crossing_threshold_is_held(self) -> None:
        data = self._intake()
        component = data.mixture_components.iloc[0].copy()
        component["함량(%)"] = None
        component["함량 최저(%)"] = 20
        component["함량 최고(%)"] = 30

        def fake_screen(single):
            pct = float(single.chemicals.iloc[0]["함량(%)"])
            hit = []
            if pct >= 25:
                hit = [{
                    "source_key": "CAP_QTY_APP2", "row_no": 1, "cas": "108-88-3",
                    "item_no": "X", "content_threshold_pct": 25,
                    "lower_quantity_ton": 1, "upper_quantity_ton": 10,
                }]
            return SimpleNamespace(ready=True, legal_hits=hit, blockers=[], row_numbers=[1] if hit else [])

        hits, blockers, claimed, ready = _screen_component_direct(fake_screen, data, 1, component)
        self.assertTrue(ready)
        self.assertTrue(claimed)
        self.assertEqual(hits, [])
        self.assertTrue(any("법정 함량기준" in blocker and "20~30%" in blocker for blocker in blockers))


if __name__ == "__main__":
    unittest.main()
