from __future__ import annotations

from io import BytesIO
import unittest
from unittest.mock import patch

from docx import Document

from engine.stage2.cross_validation import CrossValidationReport
from engine.stage2.intake import field_label, selected_requirement_specs
from engine.stage2.project import Stage2Project
from engine.stage2.psm_baseline_docx import build_psm_baseline_draft
from engine.stage2.report_draft import report_generation_status
from engine.stage2.scope_validation import validate_selected_scope
from engine.stage2.system_final_gate import PASS, evaluate_system_final_gate


class PSMFullStatutoryE2ETests(unittest.TestCase):
    @staticmethod
    def _scenario(kind: str, wind: str, erpg1: str, erpg2: str, erpg3: str):
        return {
            "시나리오 구분": kind,
            "풍속(m/s)": wind,
            "대기안정도(A~F)": "F" if "최악" in kind else "D",
            "대기온도(℃)": "25",
            "습도(%)": "50",
            "표면거칠기": "도시",
            "물질명": "염소",
            "물질의 상태": "기체",
            "설비명(또는 배관부위)": "V-201",
            "운전압력(MPa)": "0.7",
            "운전온도(℃)": "25",
            "누출구의 크기(mm2)": "25" if "최악" in kind else "10",
            "웅덩이 크기(m2)": "해당 없음",
            "누출결과": "연속누출",
            "직접계산(kg/s or kg)": "0.25" if "최악" in kind else "0.10",
            "웅덩이(kg/s)": "해당 없음",
            "설비/배관(kg/s)": "0.25" if "최악" in kind else "0.10",
            "화재-4 kW/m2": "해당 없음",
            "화재-12.5 kW/m2": "해당 없음",
            "화재-37.5 kW/m2": "해당 없음",
            "폭발-7 kPa": "해당 없음",
            "폭발-21 kPa": "해당 없음",
            "폭발-70 kPa": "해당 없음",
            "인화성-25% LEL": "해당 없음",
            "인화성-LEL": "해당 없음",
            "인화성-UEL": "해당 없음",
            "독성-ERPG 1": erpg1,
            "독성-ERPG 2": erpg2,
            "독성-ERPG 3": erpg3,
            "계산모델·결과 근거": "KORA/ALOHA 결과파일",
        }

    def _project(self) -> Stage2Project:
        p = Stage2Project(
            project_id="S2-PSM-FULL-E2E",
            company_name="세계화학(주) 제1공장",
            psm_required=True,
            cap_required=False,
            scope_confirmed=True,
            psm_selected=True,
            cap_selected=False,
        )

        def setf(key: str, value, status: str = "USER_CONFIRMED"):
            p.set_field(key, field_label(key), value, status)

        setf("business.company_name", "세계화학(주) 제1공장", "VERIFIED")
        setf("business.address", "충청북도 충주시 산업로 100", "VERIFIED")
        setf("inventory.chemicals", [{
            "물질명": "염소",
            "CAS 번호": "7782-50-5",
            "함량(%)": 99.9,
            "최대보유량": 2500,
            "단위": "kg",
        }], "VERIFIED")
        setf("inventory.facilities", [{
            "설비번호": "V-201",
            "설비명": "염소 저장용기",
            "설비종류": "압력용기",
        }], "VERIFIED")

        setf("psm.business.form12_details", [{
            "사업장명": "세계화학(주) 제1공장",
            "제출구분": "변경",
            "사업자등록번호": "123-45-67890",
            "대표자": "홍길동",
            "대상 유해·위험설비": "염소 저장·공급공정",
            "한국표준산업분류": "C20119",
            "근로자수": "85",
            "계약전력(kW)": "1500",
            "작성자 성명": "김작성",
            "작성자 자격": "산업안전기사",
            "주요 원료": "염소",
            "주요 생산품": "제품 A",
            "사업개요": "염소 저장 및 공급공정 변경",
            "사업장 소재지": "충청북도 충주시 산업로 100",
            "전화번호": "043-000-0000",
            "전송번호": "해당 없음",
            "부지면적": "12,500㎡",
            "주요 건물": "생산동 2동 / 연면적 3,200㎡",
            "총 사업기간": "2026-10-01 ~ 2027-03-31",
            "착공예정일": "2026-10-01",
            "시운전기간": "2027-03-01 ~ 2027-03-31",
        }])

        setf("psm.psi.chemical_details", [{
            "물질명": "염소",
            "CAS 번호": "7782-50-5",
            "분자식": "Cl2",
            "폭발한계 하한": "해당 없음",
            "폭발한계 상한": "해당 없음",
            "노출기준": "TWA 0.5 ppm",
            "독성치": "LC50 293 ppm (rat, 1h)",
            "인화점": "해당 없음",
            "발화점": "해당 없음",
            "증기압": "기체",
            "부식성": "예",
            "이상반응 유무": "예",
            "일일사용량": "0.4 ton/day",
            "저장량": "2.5 ton",
            "비고": "회사 SDS 확인",
        }])

        setf("psm.psi.machinery_list", [{
            "기계번호": "P-101",
            "기계명": "염소 이송펌프",
            "형식": "원심펌프",
            "용량": "20 m3/h",
            "재질": "SUS316",
            "동력": "7.5 kW",
            "방호·보호장치 종류": "커플링 가드·모터 과부하 보호",
            "비고": "회사 설비명세 확인",
        }])

        setf("psm.psi.equipment_specs", [{
            "설비번호": "V-201",
            "설비명": "염소 저장용기",
            "취급물질": "염소",
            "용량": "3 m3",
            "운전압력": "0.7 MPa",
            "설계압력": "1.5 MPa",
            "운전온도": "25 ℃",
            "설계온도": "80 ℃",
            "재질": "SUS316L",
            "부속품재질": "SUS316L",
            "개스킷재질": "PTFE",
            "용접효율": "1.0",
            "계산두께": "8 mm",
            "부식여유": "1 mm",
            "사용두께": "10 mm",
            "후열처리 여부": "해당 없음",
            "비파괴검사율": "20%",
            "비고": "PID-201",
        }])

        setf("psm.psi.piping_gasket_specs", [{
            "배관번호·Class": "PCL-201",
            "유체명": "염소",
            "설계온도": "80 ℃",
            "설계압력": "1.5 MPa",
            "배관재질": "SUS316L",
            "개스킷 재질": "PTFE",
            "비파괴검사율": "20%",
            "후열처리여부": "해당 없음",
            "비고": "PID-201",
        }])

        setf("psm.psi.relief_device_specs", [{
            "안전밸브·파열판 번호": "PSV-201",
            "배출물질": "염소",
            "배출상태": "기체",
            "배출용량": "850 kg/h",
            "정격용량": "900 kg/h",
            "노즐크기 입구": "20A",
            "노즐크기 출구": "32A",
            "보호대상 설비번호": "V-201",
            "보호기기 운전압력": "0.7 MPa",
            "보호기기 설계압력": "1.5 MPa",
            "설정압력": "1.3 MPa",
            "몸체재질": "WCB",
            "TRIM 재질": "SUS316",
            "정밀도": "±3%",
            "최종 배출·처리 지점": "스크러버",
            "배출원인": "과압",
            "형식": "안전밸브",
        }])

        setf("psm.psi.form_applicability", [
            {"서식번호": "17-2", "적용여부": "적용", "확인근거": "P&ID 및 인터록 목록"},
            {"서식번호": "17-3", "적용여부": "적용", "확인근거": "소화설비 배치도"},
            {"서식번호": "17-4", "적용여부": "적용", "확인근거": "화재감지기 목록"},
            {"서식번호": "17-5", "적용여부": "적용", "확인근거": "가스감지기 목록"},
            {"서식번호": "18", "적용여부": "적용", "확인근거": "내화도면"},
            {"서식번호": "19", "적용여부": "적용", "확인근거": "국소배기 명세"},
            {"서식번호": "20", "적용여부": "적용", "확인근거": "폭발위험장소 구분도"},
            {"서식번호": "19-2", "적용여부": "적용", "확인근거": "위험성평가 및 피해예측 결과"},
        ])

        setf("psm.psi.interlock_conditions", [{
            "인터록번호": "IL-201",
            "대상설비번호": "V-201",
            "설정값-온도(℃)": "50 ℃",
            "설정값-압력(MPa)": "1.0 MPa",
            "설정값-액위(m)": "해당 없음",
            "설정값-기타": "해당 없음",
            "감지기번호": "TI-201/PI-201",
            "최종 작동설비번호": "XV-201",
            "가동중지범위": "염소 공급 차단",
            "점검주기": "월 1회",
            "비고": "SIS-201",
        }])

        setf("psm.psi.fire_protection_table", [{
            "설치지역": "염소 저장실",
            "소화기": "분말 4기",
            "자동확산소화기": "해당 없음",
            "자동소화장치": "해당 없음",
            "옥내소화전": "2개소",
            "스프링클러": "전면 설치",
            "물분무소화설비": "해당 없음",
            "포소화설비": "해당 없음",
            "CO2 소화설비": "해당 없음",
            "할로겐화합물 소화설비": "해당 없음",
            "청정소화약제 소화설비": "해당 없음",
            "옥외소화전": "2개소",
        }])

        setf("psm.psi.fire_detection_table", [{
            "설치지역": "염소 저장실",
            "단독경보형 감지기": "해당 없음",
            "비상경보설비": "1식",
            "시각경보기": "2개소",
            "자동화재탐지설비": "열감지기 6개",
            "비상방송설비": "1식",
            "자동화재속보설비": "1식",
            "통합감시시설": "중앙제어실 연동",
            "누전경보기": "1식",
        }])

        setf("psm.psi.gas_detection_table", [{
            "감지기 번호": "GD-201",
            "검출대상 물질": "염소",
            "설치위치": "염소 저장실",
            "작동시간": "30초 이내",
            "측정방식": "전기화학식",
            "경보 설정값": "0.5 ppm",
            "경보 위치": "중앙제어실·현장 경광등",
            "정밀도": "제조사 사양",
            "경보시 조치내용": "염소 공급 차단",
            "유지관리": "월 1회 기능점검",
            "비고": "GA-201",
        }])

        setf("psm.psi.fireproofing_table", [{
            "내화설비 또는 지역": "V-201 지지철골",
            "내화부위": "주기둥 및 보",
            "내화시험기준 및 시간": "2시간 내화성능",
            "비고": "내화피복 적용",
        }])

        setf("psm.psi.local_exhaust_table", [{
            "공정 또는 작업장명": "염소 취급실",
            "실내외 구분": "실내",
            "발생원": "용기 연결부",
            "유해물질 종류": "염소",
            "후드형식": "포위식",
            "후드 제어풍속(m/s)": "0.5",
            "덕트내 반송속도(m/s)": "12",
            "배풍량(m3/min)": "80",
            "전동기용량(kW)": "7.5",
            "배기 및 처리순서": "후드 → 스크러버 → 배기구",
            "방폭형식": "Ex d IIB T4",
        }])

        setf("psm.psi.ex_equipment", [{
            "설치장소 또는 공정": "염소 저장실",
            "전기/계장 기계·기구명": "모터·현장계기",
            "0종장소 선정기준(방폭형식)": "해당 없음",
            "1종장소 선정기준(방폭형식)": "Ex d IIB T4",
            "2종장소 선정기준(방폭형식)": "Ex e IIB T4",
        }])

        setf("psm.risk.team", [{
            "책임분야": "공정",
            "성명": "홍길동",
            "소속회사": "세계화학(주)",
            "직책": "공정팀장",
            "주요경력": "공정설계 및 운전 15년",
        }])

        setf("psm.risk.consequence_table", [
            self._scenario("최악의 사고 시나리오", "1.5", "650", "300", "180"),
            self._scenario("대안의 사고 시나리오", "3.0", "320", "150", "90"),
        ])
        setf("psm.risk.consequence", "KORA/ALOHA 계산 입력·결과 및 검토자료")

        # Fill every remaining required PSM/common field with an explicit
        # company-confirmed fact/evidence placeholder. Form-specific structured
        # fields above remain authoritative for the statutory annex checks.
        for spec in selected_requirement_specs(p):
            for key in spec.field_keys:
                if p.get_field(key) is None:
                    setf(key, f"{field_label(key)} 회사 확인자료")

        return p

    @patch(
        "engine.stage2.scope_validation.validate_stage2_project",
        return_value=CrossValidationReport(issues=(), checked_rules=0),
    )
    def test_full_psm_statutory_docx_stage4_and_final_gate_are_ready(self, _raw_validation):
        project = self._project()

        report = validate_selected_scope(project)
        form_issues = [
            issue for issue in report.issues
            if issue.system == "PSM" and issue.code.startswith("PSM-FORM")
        ]

        self.assertEqual(len(form_issues), 15)
        self.assertTrue(all(issue.status == "PASS" for issue in form_issues))
        self.assertTrue(any(issue.code == "PSM-FORM12-READY" for issue in form_issues))
        self.assertTrue(any(issue.code == "PSM-FORM19-2-READY" for issue in form_issues))
        self.assertTrue(any(issue.code == "PSM-FORM21-READY" for issue in form_issues))

        status = report_generation_status(project, "PSM")
        self.assertTrue(status.final_ready)
        self.assertEqual(status.state, "READY")
        self.assertEqual(status.completion_pct, 100.0)

        gate = evaluate_system_final_gate(project, "PSM", report)
        self.assertTrue(gate.ready)
        self.assertEqual(
            [item.key for item in gate.checkpoints],
            [
                "psm.final.completeness",
                "psm.final.statutory_forms",
                "psm.final.validation",
            ],
        )
        self.assertTrue(all(item.status == PASS for item in gate.checkpoints))

        doc = Document(BytesIO(build_psm_baseline_draft(project)))
        self.assertEqual(len(doc.tables), 15)

        expected_by_table = (
            ("세계화학(주) 제1공장", "염소 저장 및 공급공정 변경"),
            ("염소", "7782-50-5", "Cl2"),
            ("P-101", "염소 이송펌프"),
            ("V-201", "SUS316L"),
            ("PCL-201", "PTFE"),
            ("PSV-201", "스크러버"),
            ("IL-201", "염소 공급 차단"),
            ("염소 저장실", "분말 4기"),
            ("염소 저장실", "열감지기 6개"),
            ("GD-201", "0.5 ppm"),
            ("V-201 지지철골", "2시간 내화성능"),
            ("염소 취급실", "스크러버"),
            ("1.5", "3.0", "650", "90"),
            ("염소 저장실", "Ex d IIB T4"),
            ("홍길동", "공정설계 및 운전 15년"),
        )
        for table, expected in zip(doc.tables, expected_by_table):
            text = "\n".join(cell.text for row in table.rows for cell in row.cells)
            for value in expected:
                self.assertIn(value, text)

        full_text = "\n".join(
            cell.text
            for table in doc.tables
            for row in table.rows
            for cell in row.cells
        )
        self.assertNotIn("[확인 필요]", full_text)


    def test_psm_statutory_writer_is_byte_deterministic(self):
        project = self._project()

        first = build_psm_baseline_draft(project)
        second = build_psm_baseline_draft(project)

        self.assertEqual(first, second)
        self.assertGreater(len(first), 0)



if __name__ == "__main__":
    unittest.main()
