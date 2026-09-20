from __future__ import annotations

from io import BytesIO
import unittest

from openpyxl import load_workbook

from engine.stage2.completeness import evaluate_project_completeness
from engine.stage2.intake import build_intake_catalog, COVERAGE_NOT_APPLICABLE
from engine.stage2.integrated_workbook import apply_integrated_authoring_workbook
from engine.stage2.project import create_project_from_stage1_snapshot
from engine.stage2.workbook_enhancements import build_enhanced_integrated_authoring_workbook


class PSMPhase24WorkbookTests(unittest.TestCase):
    def _project(self):
        snapshot = {
            "source_fingerprint": "e" * 64,
            "business": {
                "회사명": "PSM24테스트",
                "사업장명": "제1공장",
                "사업장 소재지": "테스트시 산업로 1",
            },
            "documents": {},
            "chemicals": [{
                "물질명": "염소",
                "CAS 번호": "7782-50-5",
                "함량(%)": 99.9,
                "최대보유량": 2500,
                "단위": "kg",
            }],
            "facilities": [{
                "설비번호": "V-201",
                "설비명": "염소 저장용기",
                "설비종류": "압력용기",
            }],
            "decision": {
                "psm_status": "공정안전보고서 제출 대상",
                "cap_status": "비대상",
            },
        }
        project = create_project_from_stage1_snapshot(snapshot)
        project.set_authoring_scope(psm_selected=True, cap_selected=False)
        return project

    @staticmethod
    def _headers(ws):
        return {
            str(cell.value or "").strip(): cell.column
            for cell in ws[4]
            if str(cell.value or "").strip()
        }

    @staticmethod
    def _scenario(kind: str, wind: str, erpg1: str, erpg2: str, erpg3: str):
        return {
            "단위공장": "제1공장",
            "사고유형": "독성 누출",
            "시나리오명": "염소-" + ("최악" if "최악" in kind else "대안") + "-" + wind,
            "대상 설비번호": "V-201",
            "시나리오 구분": kind,
            "풍속(m/s)": wind,
            "대기안정도(A~F)": "F" if "최악" in kind else "D",
            "대기온도(℃)": "25",
            "습도(%)": "50",
            "표면거칠기": "도시",
            "물질명": "염소",
            "물질의 상태": "기체",
            "설비명(또는 배관부위)": "V-201",
            "운전압력(MPa)": "0.7",
            "운전온도(℃)": "25",
            "누출구의 크기(mm2)": "25" if "최악" in kind else "10",
            "웅덩이 크기(m2)": "해당 없음",
            "누출결과": "연속누출",
            "직접계산(kg/s or kg)": "0.25" if "최악" in kind else "0.10",
            "웅덩이(kg/s)": "해당 없음",
            "설비/배관(kg/s)": "0.25" if "최악" in kind else "0.10",
            "화재-4 kW/m2": "해당 없음",
            "화재-12.5 kW/m2": "해당 없음",
            "화재-37.5 kW/m2": "해당 없음",
            "폭발-7 kPa": "해당 없음",
            "폭발-21 kPa": "해당 없음",
            "폭발-70 kPa": "해당 없음",
            "인화성-25% LEL": "해당 없음",
            "인화성-LEL": "해당 없음",
            "인화성-UEL": "해당 없음",
            "독성-ERPG 1": erpg1,
            "독성-ERPG 2": erpg2,
            "독성-ERPG 3": erpg3,
            "계산모델·결과 근거": "KORA/ALOHA 결과파일",
        }

    def test_workbook_exposes_form12_and_form19_2_with_prefill_and_dropdowns(self):
        project = self._project()
        wb = load_workbook(BytesIO(
            build_enhanced_integrated_authoring_workbook(project, example=False)
        ))

        self.assertIn("08_PSM_별지12_사업개요", wb.sheetnames)
        self.assertIn("19_PSM_사고피해예측", wb.sheetnames)

        form12 = wb["08_PSM_별지12_사업개요"]
        h12 = self._headers(form12)
        self.assertEqual(form12.cell(5, h12["사업장명"]).value, "PSM24테스트")
        self.assertEqual(form12.cell(5, h12["사업장 소재지"]).value, "테스트시 산업로 1")
        self.assertEqual(form12.cell(5, h12["주요 원료"]).value, "염소")
        self.assertTrue(any(
            "_선택목록" in str(dv.formula1)
            for dv in form12.data_validations.dataValidation
        ))

        consequence = wb["19_PSM_사고피해예측"]
        h19 = self._headers(consequence)
        self.assertIn("단위공장", h19)
        self.assertIn("사고유형", h19)
        self.assertIn("시나리오명", h19)
        self.assertIn("시나리오 구분", h19)
        self.assertIn("독성-ERPG 3", h19)
        self.assertTrue(any(
            "_선택목록" in str(dv.formula1)
            for dv in consequence.data_validations.dataValidation
        ))

        app = wb["09_PSM_조건부서식_적용여부"]
        form_numbers = [str(app.cell(row, 1).value or "") for row in range(5, 13)]
        self.assertIn("19-2", form_numbers)

    def test_form12_roundtrip_saves_structured_row(self):
        project = self._project()
        raw = build_enhanced_integrated_authoring_workbook(project, example=False)
        wb = load_workbook(BytesIO(raw))
        ws = wb["08_PSM_별지12_사업개요"]
        h = self._headers(ws)

        values = {
            "제출구분": "변경",
            "사업자등록번호": "123-45-67890",
            "대표자": "홍길동",
            "대상 유해·위험설비": "염소 저장·공급공정",
            "한국표준산업분류": "C20119",
            "근로자수": "85",
            "계약전력(kW)": "1500",
            "작성자 성명": "김작성",
            "작성자 자격": "산업안전기사",
            "주요 생산품": "제품 A",
            "사업개요": "염소 저장 및 공급공정",
            "전화번호": "053-000-0000",
            "전송번호": "해당 없음",
            "부지면적": "12,500㎡",
            "주요 건물": "생산동 2동",
            "총 사업기간": "2026-10-01 ~ 2027-03-31",
            "착공예정일": "2026-10-01",
            "시운전기간": "2027-03-01 ~ 2027-03-31",
        }
        for header, value in values.items():
            ws.cell(5, h[header], value)

        out = BytesIO()
        wb.save(out)
        apply_integrated_authoring_workbook(project, out.getvalue())

        record = project.get_field("psm.business.form12_details")
        self.assertIsNotNone(record)
        self.assertEqual(record.status, "USER_CONFIRMED")
        self.assertEqual(record.value[0]["사업장명"], "PSM24테스트")
        self.assertEqual(record.value[0]["사업장 소재지"], "테스트시 산업로 1")
        self.assertEqual(record.value[0]["제출구분"], "변경")

    def test_consequence_table_roundtrip_saves_both_scenarios(self):
        project = self._project()
        raw = build_enhanced_integrated_authoring_workbook(project, example=False)
        wb = load_workbook(BytesIO(raw))
        ws = wb["19_PSM_사고피해예측"]
        h = self._headers(ws)

        for row_no, scenario in enumerate((
            self._scenario("최악의 사고 시나리오", "1.5", "650", "300", "180"),
            self._scenario("대안의 사고 시나리오", "3.0", "320", "150", "90"),
        ), start=5):
            for header, value in scenario.items():
                ws.cell(row_no, h[header], value)

        out = BytesIO()
        wb.save(out)
        apply_integrated_authoring_workbook(project, out.getvalue())

        record = project.get_field("psm.risk.consequence_table")
        self.assertIsNotNone(record)
        self.assertEqual(record.status, "USER_CONFIRMED")
        self.assertEqual(len(record.value), 2)
        self.assertEqual(record.value[0]["시나리오 구분"], "최악의 사고 시나리오")

    def test_consequence_table_does_not_replace_model_evidence_requirement(self):
        project = self._project()
        project.set_field(
            "psm.risk.consequence_table",
            "별지 제19호의2서식 사고피해예측 수치표",
            [
                self._scenario("최악의 사고 시나리오", "1.5", "650", "300", "180"),
                self._scenario("대안의 사고 시나리오", "3.0", "320", "150", "90"),
            ],
            "USER_CONFIRMED",
        )
        project.set_field(
            "psm.psi.form_applicability",
            "PSM 조건부 별지서식 적용여부",
            [{"서식번호": "19-2", "적용여부": "적용", "확인근거": "위험성평가 결과"}],
            "USER_CONFIRMED",
        )

        completeness = evaluate_project_completeness(project)
        item = next(
            row for row in completeness["requirements"]
            if row["key"] == "psm.risk.consequence"
        )
        self.assertEqual(item["state"], "HOLD")
        self.assertIn("psm.risk.consequence", item["missing_fields"])

    def test_explicit_form19_2_non_applicability_removes_model_attachment_request(self):
        project = self._project()
        project.set_field(
            "psm.psi.form_applicability",
            "PSM 조건부 별지서식 적용여부",
            [{
                "서식번호": "19-2",
                "적용여부": "해당 없음",
                "확인근거": "위험성평가 결과 별지 제19호의2 미적용 확인",
            }],
            "USER_CONFIRMED",
        )

        wb = load_workbook(BytesIO(
            build_enhanced_integrated_authoring_workbook(project, example=False)
        ))
        ws = wb["07_도면_첨부자료목록"]
        labels = {
            str(ws.cell(row, 1).value or "").strip()
            for row in range(5, ws.max_row + 1)
        }

        self.assertNotIn("사고피해예측 결과", labels)

    def test_explicit_form19_2_non_applicability_aligns_completeness_and_intake(self):
        project = self._project()
        project.set_field(
            "psm.psi.form_applicability",
            "PSM 조건부 별지서식 적용여부",
            [{
                "서식번호": "19-2",
                "적용여부": "해당 없음",
                "확인근거": "위험성평가 결과 별지 제19호의2 미적용 확인",
            }],
            "USER_CONFIRMED",
        )

        completeness = evaluate_project_completeness(project)
        item = next(
            row for row in completeness["requirements"]
            if row["key"] == "psm.risk.consequence"
        )
        self.assertEqual(item["state"], "NOT_REQUIRED")

        intake = build_intake_catalog(project)
        intake_item = next(row for row in intake if row.requirement_key == "psm.risk.consequence")
        self.assertEqual(intake_item.coverage_status, COVERAGE_NOT_APPLICABLE)


if __name__ == "__main__":
    unittest.main()
