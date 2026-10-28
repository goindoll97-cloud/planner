"""④ 단일 다운로드는 검증을 유지하고 규정서식과 서술형을 함께 제공합니다."""
from io import BytesIO
import unittest
from unittest.mock import patch
from docx import Document
from engine.stage2 import report_export as export, cap_fact_drafting as fd
from tests.test_cap_fact_drafting import _project


class WorkspaceDownloadTests(unittest.TestCase):
    def test_current_narratives_are_included_and_validation_is_preserved(self):
        project = _project()
        fd.refresh_rule_drafts(project)
        with patch.object(export, 'evaluate', wraps=export.evaluate) as check:
            draft = export.build_cap_workspace_draft(project)
        check.assert_called_once_with(project, 'CAP')
        doc = Document(BytesIO(draft.docx))
        self.assertGreaterEqual(len(doc.tables), 16)
        text = '\n'.join(p.text for p in doc.paragraphs)
        self.assertIn('안전관리 책임자: 공장장', text)
        self.assertIn('AI 초안', text)
        self.assertIn('초안', draft.name)
        self.assertGreater(draft.state.hold + draft.state.review, 0)
        self.assertIn(b'REVIEW_ONLY', draft.provenance_json)

    def test_validation_failure_never_produces_downloadable_bytes(self):
        with patch.object(export, 'evaluate', side_effect=RuntimeError('validation unavailable')):
            with self.assertRaises(RuntimeError):
                export.build_cap_workspace_draft(_project())
