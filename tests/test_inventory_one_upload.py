from __future__ import annotations

import unittest
from io import BytesIO

from openpyxl import Workbook

from engine.inventory import read_intake_workbook


class InventoryOneUploadTests(unittest.TestCase):
    def _workbook_bytes(self, facility_sheet_name: str = "04_시설별최대보유량") -> bytes:
        wb = Workbook()
        ws = wb.active
        ws.title = "01_사업장기본정보"
        ws.append(["title"])
        ws.append([])
        ws.append(["항목", "입력값"])
        ws.append(["사업장명", "테스트사업장"])
        ws.append(["사업장 주소", "테스트주소"])
        ws.append(["업종 또는 주요 생산품", "테스트업종"])

        chem = wb.create_sheet("02_화학물질목록")
        chem.append(["title"])
        chem.append([])
        chem.append([
            "No.", "제품명", "CAS No.", "함량(%)", "취급형태",
            "최대 제조·사용량", "최대 저장량", "수량 단위",
            "최대 동시보유량(알면 입력)", "최대보유량 법정 산정 여부",
        ])
        chem.append([1, "포스겐", "75-44-5", 100, "사용", 100, 1600, "kg", 1600, "Y"])

        docs = wb.create_sheet("03_기존문서보유여부")
        docs.append(["title"])
        docs.append([])
        docs.append(["자료", "보유 여부"])
        docs.append(["공정안전보고서(PSM)", "N"])

        fac = wb.create_sheet(facility_sheet_name)
        fac.append(["title"])
        fac.append([])
        fac.append([
            "적용여부", "목록행번호", "제품명", "CAS No.", "시설명", "시설유형",
            "제외시설여부", "제외사유", "물질성상", "공정유형", "별표4 기준함량(%)",
            "함량근거", "설계용량", "용량단위", "비중 또는 밀도(kg/L=ton/m3)",
            "보관계획도 최대량", "일일최대보관량", "질량단위", "직접확인 최대보유량",
            "직접확인 근거", "복수성상 증빙", "비고",
        ])
        fac.append([
            "해당", 1, "포스겐", "75-44-5", "V-201", "기타", "N", "해당없음",
            "기체·고압가스", "해당없음", 100, "SDS", None, None, None,
            None, None, "kg", 1600, "용기 충전질량 합계", "N", "test",
        ])
        fac.append(["해당없음"])

        final = wb.create_sheet("05_최종판정조건")
        final.append(["title"])
        final.append(["제도", "확인항목", "입력값"])
        final.append(["화학사고예방관리계획서", "법정 작성 면제시설 해당 여부", "해당없음"])
        final.append(["화학사고예방관리계획서", "상위 규정수량 이상을 취급하는 개별 주요취급시설 존재 여부", "Y"])

        out = BytesIO()
        wb.save(out)
        return out.getvalue()

    def test_reads_canonical_04_facility_and_final_condition_sheets(self):
        intake = read_intake_workbook(self._workbook_bytes())
        self.assertEqual(len(intake.facilities), 1)
        self.assertEqual(intake.facilities.iloc[0]["시설명"], "V-201")
        self.assertEqual(intake.final_conditions["법정 작성 면제시설 해당 여부"], "해당없음")
        self.assertTrue(bool(intake.source_fingerprint))

    def test_reads_legacy_03_facility_sheet_for_backward_compatibility(self):
        intake = read_intake_workbook(self._workbook_bytes("03_시설별최대보유량"))
        self.assertEqual(len(intake.facilities), 1)
        self.assertEqual(intake.facilities.iloc[0]["시설명"], "V-201")


if __name__ == "__main__":
    unittest.main()
