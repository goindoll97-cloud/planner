from __future__ import annotations

from io import BytesIO
from types import SimpleNamespace
import unittest
from unittest.mock import patch

from docx import Document
from openpyxl import load_workbook

from engine.stage2.cap_chemical_legal import CAPChemicalLegalData
from engine.stage2.cap_sds_engine import (
    CAPForm6SDSData,
    CAPForm7Data,
    build_cap_form6_sds_data,
    build_cap_form7_data,
)
from engine.stage2.cap_baseline_docx import build_cap_baseline_draft
from engine.stage2.integrated_workbook import (
    apply_integrated_authoring_workbook,
    attach_company_file,
    build_integrated_authoring_workbook,
)
from engine.stage2.project import EvidenceRef, Stage2Project


class CAPCompanySDSTests(unittest.TestCase):
    def _project(self) -> Stage2Project:
        return Stage2Project(
            project_id="S2-SDS",
            company_name="회사SDS테스트",
            cap_required=True,
            cap_group="1군",
            scope_confirmed=True,
            cap_selected=True,
        )

    def _chemical_row(self) -> dict:
        return {
            "물질명": "염소",
            "CAS 번호": "7782-50-5",
            "함량(%)": 99.9,
            "물리적 상태": "기체",
            "최대보유량": 800,
            "단위": "kg",
            "비중": "2.49 (공기=1)",
            "폭발한계 하한": "해당 없음",
            "폭발한계 상한": "해당 없음",
            "독성구분 항목": "급성독성(흡입)",
            "독성구분": "구분 2",
            "위험노출수준": "ERPG-2 3 ppm",
            "허용농도값": "TWA 0.5 ppm",
            "증기압": "기체",
            "부식성": "예",
            "SDS 파일명": "chlorine_company_SDS.pdf",
            "SDS 개정일": "2026-05-10",
        }

    @patch("engine.stage2.cap_sds_engine.build_cap_chemical_legal_data")
    def test_form6_accepts_explicit_not_applicable_sds_values(self, legal_mock):
        project = self._project()
        row = self._chemical_row()
        legal_mock.return_value = CAPChemicalLegalData(
            rows=({**row, "물질구분": "인체급성유해성물질 / 사고대비물질", "고유번호": "97-1-1"},),
            blockers=(),
            messages=("법령확인",),
            app2_ready=True,
        )

        result = build_cap_form6_sds_data(project)

        self.assertTrue(result.ready, result.blockers)
        self.assertEqual(result.rows[0]["폭발한계 하한"], "해당 없음")
        self.assertEqual(result.rows[0]["SDS 파일명"], "chlorine_company_SDS.pdf")

    @patch("engine.stage2.cap_sds_engine.build_cap_chemical_legal_data")
    def test_form6_missing_company_sds_property_is_hold(self, legal_mock):
        project = self._project()
        row = self._chemical_row()
        row["증기압"] = ""
        legal_mock.return_value = CAPChemicalLegalData(
            rows=({**row, "물질구분": "사고대비물질", "고유번호": "97-1-1"},),
            blockers=(),
            messages=(),
            app2_ready=True,
        )

        result = build_cap_form6_sds_data(project)

        self.assertFalse(result.ready)
        self.assertTrue(any("증기압" in item for item in result.blockers))

    @patch("engine.stage2.cap_sds_engine.build_cap_form1_data")
    @patch("engine.stage2.cap_sds_engine.build_cap_chemical_legal_data")
    def test_form7_reuses_legal_unique_number_and_form1_holding(
        self, legal_mock, form1_mock
    ):
        project = self._project()
        project.set_field(
            "cap.chemical.hazard_information",
            "유해성 정보",
            [{
                "물질명": "염소",
                "CAS 번호": "7782-50-5",
                "인체유해성": "흡입 급성독성",
                "물리적 위험성": "가압가스",
                "환경유해성": "수생생물 유해성",
                "출처": "회사 제품 SDS 제2·11·12항",
                "선정 사유": "사고시나리오 독성영향 대표물질",
                "SDS 파일명": "chlorine_company_SDS.pdf",
                "SDS 개정일": "2026-05-10",
            }],
            "USER_CONFIRMED",
        )
        legal_mock.return_value = CAPChemicalLegalData(
            rows=({
                "물질명": "염소",
                "CAS 번호": "7782-50-5",
                "고유번호": "97-1-1",
                "물질구분": "인체급성유해성물질 / 사고대비물질",
            },),
            blockers=(),
            messages=(),
            app2_ready=True,
        )
        form1_mock.return_value = SimpleNamespace(
            chemical_rows=({
                "물질명": "염소",
                "CAS No.": "7782-50-5",
                "함량(%)": "99.9",
                "사업장 내 최대보유량(ton)": "0.8",
            },),
            blockers=(),
        )

        result = build_cap_form7_data(project)

        self.assertTrue(result.ready, result.blockers)
        self.assertEqual(result.row["고유번호"], "97-1-1")
        self.assertEqual(result.row["최대보유량(ton)"], "0.8")
        self.assertEqual(result.row["함량(%)"], "99.9")

    def test_workbook_contains_company_sds_columns_and_form7_sheet(self):
        project = self._project()
        data = build_integrated_authoring_workbook(project, example=False)
        wb = load_workbook(BytesIO(data), data_only=False)

        self.assertIn("28_대표물질_유해성정보", wb.sheetnames)
        chemical_headers = [str(cell.value or "") for cell in wb["02_화학물질정보"][4]]
        for header in (
            "비중", "폭발한계 하한", "독성구분", "위험노출수준",
            "허용농도값", "증기압", "부식성", "SDS 파일명", "SDS 개정일",
        ):
            self.assertIn(header, chemical_headers)

    def test_uploaded_sds_links_to_structured_chemical_table_without_downgrading_company_fact(self):
        project = self._project()
        data = build_integrated_authoring_workbook(project, example=False)
        wb = load_workbook(BytesIO(data))
        ws = wb["02_화학물질정보"]
        headers = {str(cell.value or ""): cell.column for cell in ws[4]}
        row = self._chemical_row()
        for header, value in row.items():
            if header in headers:
                ws.cell(5, headers[header], value)

        output = BytesIO()
        wb.save(output)
        apply_integrated_authoring_workbook(project, output.getvalue())

        before = project.get_field("cap.chemical.details")
        self.assertIsNotNone(before)
        self.assertEqual(before.status, "USER_CONFIRMED")

        ref = EvidenceRef(
            source_type="ATTACHMENT",
            source_name="chlorine_company_SDS.pdf",
            sha256="b" * 64,
        )
        matched = attach_company_file(project, ref)

        self.assertIn("cap.chemical.details", matched)
        after = project.get_field("cap.chemical.details")
        self.assertEqual(after.status, "USER_CONFIRMED")
        self.assertTrue(any(item.source_name == "chlorine_company_SDS.pdf" for item in after.evidence))

    @patch("engine.stage2.cap_baseline_docx.build_cap_form7_data")
    @patch("engine.stage2.cap_baseline_docx.build_cap_form6_sds_data")
    def test_final_baseline_docx_writes_company_sds_form6_and_linked_form7(
        self, form6_mock, form7_mock
    ):
        project = self._project()
        form6_mock.return_value = CAPForm6SDSData(
            rows=({
                **self._chemical_row(),
                "물질구분": "인체급성유해성물질 / 사고대비물질",
                "고유번호": "97-1-1",
            },),
            blockers=(),
            messages=(),
        )
        form7_mock.return_value = CAPForm7Data(
            row={
                "물질명": "염소",
                "CAS 번호": "7782-50-5",
                "고유번호": "97-1-1",
                "함량(%)": "99.9",
                "최대보유량(ton)": "0.8",
                "인체유해성": "흡입 급성독성",
                "물리적 위험성": "가압가스",
                "환경유해성": "수생생물 유해성",
                "출처": "회사 제품 SDS 제2·11·12항",
                "선정 사유": "사고시나리오 독성영향 대표물질",
                "SDS 파일명": "chlorine_company_SDS.pdf",
                "SDS 개정일": "2026-05-10",
            },
            blockers=(),
            messages=(),
        )

        out = build_cap_baseline_draft(project)
        doc = Document(BytesIO(out))
        form6_text = "\n".join(cell.text for row in doc.tables[15].rows for cell in row.cells)
        form7_text = "\n".join(cell.text for row in doc.tables[16].rows for cell in row.cells)

        self.assertIn("ERPG-2 3 ppm", form6_text)
        self.assertIn("TWA 0.5 ppm", form6_text)
        self.assertIn("97-1-1", form7_text)
        self.assertIn("0.8 ton", form7_text)
        self.assertIn("chlorine_company_SDS.pdf", form7_text)
        self.assertIn("사고시나리오 독성영향 대표물질", form7_text)


if __name__ == "__main__":
    unittest.main()
