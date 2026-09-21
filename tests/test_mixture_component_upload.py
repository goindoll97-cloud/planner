from __future__ import annotations

from io import BytesIO
import unittest

from openpyxl import load_workbook, Workbook

from engine.stage2 import cap_judgement as jd
from engine.stage2 import mixture_component_upload as mix
from engine.stage2.project import create_project_from_stage1_snapshot


def _xlsx(rows):
    wb = Workbook()
    ws = wb.active
    for row in rows:
        ws.append(row)
    out = BytesIO()
    wb.save(out)
    return out.getvalue()


def _project():
    snapshot = {
        "business": {"사업장명": "가상공장", "사업장 주소": "울산", "업종 또는 주요 생산품": "화학제품 제조"},
        "chemicals": [{
            "제품명": "세척제 A", "물질명": "세척제 A", "CAS No.": "", "함량(%)": "",
            "혼합물 여부": "Y", "최대 제조·사용량": 1.0, "최대 저장량": 2.0, "수량 단위": "ton",
        }],
        "mixture_components": [], "facilities": [], "documents": {}, "decision": {},
    }
    return create_project_from_stage1_snapshot(snapshot)


class MixtureComponentFileTests(unittest.TestCase):
    def test_template_contains_only_three_user_columns_and_prefills_product_names(self):
        book = load_workbook(BytesIO(mix.blank_template(["세척제 A"])))
        ws = book["혼합물 구성성분"]
        self.assertEqual([cell.value for cell in ws[1]], ["혼합제품명", "CAS No.", "함량(%)"])
        self.assertEqual(ws["A2"].value, "세척제 A")

    def test_valid_component_file_is_checked_by_cas_and_percentage(self):
        checked = mix.check(_xlsx([
            ["혼합제품명", "CAS No.", "함량(%)"],
            ["세척제 A", "67-64-1", 60],
            ["세척제 A", "108-88-3", 40],
        ]), "mix.xlsx", ["세척제 A"])
        self.assertTrue(checked.ok)
        self.assertEqual(len(checked.rows), 2)
        self.assertEqual(checked.warnings, [])

    def test_unknown_product_and_bad_cas_are_errors(self):
        checked = mix.check(_xlsx([
            ["혼합제품명", "CAS No.", "함량(%)"],
            ["없는제품", "67-64-1", 50],
            ["세척제 A", "108-88-4", 50],
        ]), "mix.xlsx", ["세척제 A"])
        self.assertFalse(checked.ok)
        self.assertTrue(any("찾을 수 없습니다" in e for e in checked.errors))
        self.assertTrue(any("CAS No." in e for e in checked.errors))

    def test_saving_second_file_resolves_the_existing_mixture_without_reasking_type(self):
        project = _project()
        self.assertEqual(jd.composition_rows(project), [1])
        mix.save_to_project(project, [
            {"혼합제품명": "세척제 A", "CAS No.": "67-64-1", "함량(%)": 60},
            {"혼합제품명": "세척제 A", "CAS No.": "108-88-3", "함량(%)": 40},
        ], sds_confirmed=True)
        self.assertEqual(jd.composition_rows(project), [])
        rows = jd.mixture_components(project)
        self.assertEqual({r["CAS No."] for r in rows}, {"67-64-1", "108-88-3"})


if __name__ == "__main__":
    unittest.main()
