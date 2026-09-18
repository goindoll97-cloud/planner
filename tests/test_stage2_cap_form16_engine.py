from __future__ import annotations

from io import BytesIO
import unittest
from unittest.mock import patch

from docx import Document

from engine.stage2.cap_form16_engine import build_cap_form16_data
from engine.stage2.project import Stage2Project
from engine.stage2.report_draft import build_report_draft


class CAPForm16EngineTests(unittest.TestCase):
    def _project(self) -> Stage2Project:
        project = Stage2Project(
            project_id="S2-FORM16",
            company_name="가상화학",
            site_name="울산공장",
            cap_required=True,
            cap_group="1군",
            scope_confirmed=True,
            cap_selected=True,
            stage1_source_fingerprint="stage1-test",
        )

        for key, label, value in (
            ("business.address", "주소", "울산광역시 테스트로 1"),
            ("cap.business.representative", "대표자", "홍길동"),
            ("cap.business.registration_no", "사업자등록번호", "123-45-67890"),
            ("cap.business.writer_name", "작성자", "최유진"),
            ("cap.business.writer_department", "부서", "환경안전팀"),
            ("cap.business.writer_contact", "연락처", "052-000-0000"),
            ("cap.business.writer_email", "메일", "safety@example.com"),
            ("cap.business.industrial_complex", "산업단지", "해당 없음"),
        ):
            project.set_field(key, label, value, "USER_CONFIRMED")

        project.set_field(
            "inventory.chemicals",
            "화학물질 목록",
            [{
                "물질명": "염소",
                "CAS 번호": "7782-50-5",
                "함량(%)": 99.9,
                "물리적 상태": "기체",
                "최대보유량": 800,
                "단위": "kg",
            }],
            "USER_CONFIRMED",
        )
        project.set_field(
            "inventory.facilities",
            "설비 목록",
            [{
                "설비번호": "V-301",
                "설비명": "염소 공급용기군",
                "설비종류": "압력용기",
                "취급물질": "염소",
                "CAS 번호": "7782-50-5",
                "최대보유량": 800,
                "단위": "kg",
            }],
            "USER_CONFIRMED",
        )
        project.set_field(
            "cap.offsite.scenario_impact_table",
            "사고시나리오 영향평가",
            [{
                "사고시나리오명": "염소 독성누출-1",
                "유해화학물질명": "염소",
                "대상 설비번호": "V-301",
                "사고유형": "독성누출",
                "장외거리(m)": 180,
                "거주민수": 25,
                "근로자수": 10,
                "갑종 보호대상 수": 1,
                "을종 보호대상 수": 0,
                "환경수용체 수": 1,
                "사고원점 좌표": "35.0000,129.0000",
                "KORA/GIS 근거": "KORA-01",
            }],
            "USER_CONFIRMED",
        )
        project.set_field(
            "cap.offsite.scenario_frequency",
            "시설빈도",
            [{
                "사고시나리오명": "염소 독성누출-1",
                "고압용기파열": 1,
                "배관파열": 1,
                "배관누출": 2,
                "상압 탱크 파열 및 누출": 0,
                "플랜지 등의 가스켓 파손": 1,
                "펌프/컴프레서 누출": 1,
                "안전밸브 오작동 및 조기개방": 0,
                "냉각수 손실": 0,
                "입/출하 시설 누출 사고": 0,
                "외부화재": 0,
                "개수 산정근거": "PID-301",
            }],
            "USER_CONFIRMED",
        )

        for key, value in (
            ("cap.prevention.emergency_contact_system", "비상연락망으로 사고상황을 전파한다."),
            ("cap.prevention.emergency_org_chart", "공장장을 비상대응 책임자로 하는 조직을 운영한다."),
            ("cap.prevention.emergency_roles", "상황전파·초기대응·대피 임무를 분담한다."),
            ("cap.internal.shutdown_authority", "공정반장"),
            ("cap.internal.shutdown_procedure", "원료차단 후 관련 설비를 정지한다."),
            ("cap.internal.response_personnel", "방재반 8명"),
            ("cap.internal.response_equipment", "보호복·공기호흡기·흡착재"),
            ("cap.internal.response_operations", "누출원 차단과 확산방지 조치를 수행한다."),
            ("cap.internal.communication_system", "비상방송과 무전기로 내부에 전파한다."),
            ("cap.internal.facility_response_plan", "염소 누출 시 원격차단 후 방재반이 접근한다."),
            ("cap.internal.investigation_plan", "사고조사팀이 원인을 조사한다."),
            ("cap.internal.recurrence_prevention", "원인별 개선조치를 추적관리한다."),
            ("cap.internal.recovery_plan", "오염제거와 설비복구 후 재가동을 승인한다."),
            ("cap.external.communication_plan", "지자체·소방·주민에게 사고정보를 제공한다."),
            ("cap.external.stakeholders", "관할 소방서·지자체·인근 사업장"),
            ("cap.external.communication_schedule", "평상시 연 1회, 사고 시 즉시"),
            ("cap.external.mutual_aid_contacts", "관할 소방서와 인근 사업장 연락망"),
            ("cap.external.resource_support", "소방차·방재인력 상호지원"),
            ("cap.external.joint_drill_plan", "연 1회 합동훈련"),
            ("cap.external.warning_system", "재난문자·방송으로 주민경보"),
            ("cap.external.evacuation_routes", "남문 방향 대피로"),
            ("cap.external.shelters", "○○체육관"),
            ("cap.external.medical_contacts", "○○병원 응급실"),
            ("cap.external.notice_targets", "총괄영향범위 내 주민"),
            ("cap.external.notice_method", "서면·홈페이지 고지"),
            ("cap.external.notice_content", "취급물질·사고위험·행동요령"),
        ):
            project.set_field(key, key, value, "USER_CONFIRMED")
        return project

    @patch("engine.stage2.cap_risk_engine.approved_source_is_current", return_value=True)
    @patch("engine.stage2.cap_form1_engine.screen_facility_stage")
    def test_form16_reuses_scenario_accident_type_and_fills_report_date(
        self, screen_mock, _current
    ):
        class Screen:
            ready = True
            blockers = ()
            messages = ()
            legal_hits = [{
                "row_no": 1,
                "source_key": "CAP_QTY_APP3",
                "lower_quantity_ton": 0.01,
                "upper_quantity_ton": 0.1,
                "cas": "7782-50-5",
                "legal_substance": "염소",
                "item_no": "1",
            }]
        screen_mock.return_value = Screen()

        result = build_cap_form16_data(self._project(), report_date="2026-09-18")

        self.assertTrue(result.ready, result.blockers)
        self.assertEqual(result.business["작성일"], "2026-09-18")
        self.assertEqual(result.business["담당자"], "환경안전팀 최유진")
        self.assertEqual(result.chemical_rows[0]["사고유형"], "독성누출")
        self.assertEqual(result.chemical_rows[0]["화학물질식별번호(CAS 번호)"], "7782-50-5")
        self.assertEqual(result.chemical_rows[0]["최대함량(%)"], "99.9")
        self.assertEqual(result.chemical_rows[0]["최대보유량(ton)"], "0.8")
        self.assertTrue(result.scenario_rows[0]["시설빈도(/연)"])
        self.assertEqual(result.scenario_rows[0]["장외거리(m)"], 180)
        self.assertEqual(len(result.external_summary), 5)

    @patch("engine.stage2.cap_risk_engine.approved_source_is_current", return_value=True)
    @patch("engine.stage2.cap_form1_engine.screen_facility_stage")
    def test_form16_fails_closed_when_company_emergency_fact_is_missing(
        self, screen_mock, _current
    ):
        class Screen:
            ready = True
            blockers = ()
            messages = ()
            legal_hits = [{
                "row_no": 1,
                "source_key": "CAP_QTY_APP3",
                "lower_quantity_ton": 0.01,
                "upper_quantity_ton": 0.1,
                "cas": "7782-50-5",
                "legal_substance": "염소",
                "item_no": "1",
            }]
        screen_mock.return_value = Screen()
        project = self._project()
        project.fields.pop("cap.external.medical_contacts")

        result = build_cap_form16_data(project, report_date="2026-09-18")

        self.assertFalse(result.ready)
        self.assertTrue(any("응급의료" in item for item in result.company_blockers))

    @patch("engine.stage2.cap_risk_engine.approved_source_is_current", return_value=True)
    @patch("engine.stage2.cap_form1_engine.screen_facility_stage")
    def test_generated_cap_docx_contains_completed_form16_sections(
        self, screen_mock, _current
    ):
        class Screen:
            ready = True
            blockers = ()
            messages = ()
            legal_hits = [{
                "row_no": 1,
                "source_key": "CAP_QTY_APP3",
                "lower_quantity_ton": 0.01,
                "upper_quantity_ton": 0.1,
                "cas": "7782-50-5",
                "legal_substance": "염소",
                "item_no": "1",
            }]
        screen_mock.return_value = Screen()
        project = self._project()

        data = build_report_draft(project, "CAP")
        doc = Document(BytesIO(data))
        text = "\n".join(paragraph.text for paragraph in doc.paragraphs)
        table_text = "\n".join(
            cell.text
            for table in doc.tables
            for row in table.rows
            for cell in row.cells
        )
        combined = text + "\n" + table_text

        self.assertIn("화학사고예방관리계획서 비상대응분야 요약서", combined)
        self.assertIn("독성누출", combined)
        self.assertIn("염소 독성누출-1", combined)
        self.assertIn("비상연락체계", combined)
        self.assertIn("응급의료", combined)
        self.assertIn("safety@example.com", combined)


if __name__ == "__main__":
    unittest.main()
