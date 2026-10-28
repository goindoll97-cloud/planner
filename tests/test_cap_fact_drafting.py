"""서술형 초안: 회사 파일에서 읽은 사실만 쓰고, 근거가 없으면 [확인 필요]로 남깁니다."""
from __future__ import annotations

import unittest
from io import BytesIO

from docx import Document

from engine.stage2 import ai_drafting as drafting
from engine.stage2 import cap_fact_drafting as fd
from engine.stage2 import demo_full_company as d
from engine.stage2 import evidence_workspace as ew
from engine.stage2.direct_template import create_direct_cap_project


def _doc(rows):
    doc = Document()
    table = doc.add_table(rows=0, cols=2)
    for label, value in rows:
        cells = table.add_row().cells
        cells[0].text, cells[1].text = label, value
    out = BytesIO()
    doc.save(out)
    return out.getvalue()


def _project(group="2군", files=("02_설비대장.xlsx", "09_안전관리_비상대응계획.docx")):
    demo = d.files()
    p = create_direct_cap_project(d.COMPANY, address=d.DATA["company"]["주소"], cap_group=group)
    for name in files:
        ew.ingest(p, name, demo[name], ew.classify_document(name, demo[name]))
    ew.publish(p, replace_existing=True)
    return p


class FakeClient:
    model = "fake-local"

    def __init__(self, text):
        self.text = text

    def generate_json(self, *, instructions, prompt):
        assert "draft_text" in instructions
        return {"draft_text": self.text}


class BrokenClient:
    model = "broken"

    def generate_json(self, *, instructions, prompt):
        raise TimeoutError("응답 없음")


def _item(project, key):
    return next(i for i in fd.items(project) if i.key == key)


class ItemTests(unittest.TestCase):
    def test_company_text_is_used_as_is_and_decisions_become_basis(self):
        p = _project()
        shutdown = _item(p, "cap.internal.shutdown")
        self.assertEqual(shutdown.state, fd.FROM_FILE)
        self.assertEqual(shutdown.company_text, d.DATA["emergency"]["shutdown"])
        safety = _item(p, "cap.prevention.safety_management")
        self.assertEqual(safety.state, fd.CAN_DRAFT)
        self.assertEqual([b.label for b in safety.basis], ["안전관리 책임자", "안전관리 조직", "안전 점검 회의"])
        self.assertIn("09_안전관리_비상대응계획.docx", safety.basis[0].source)

    def test_group_2_has_no_external_items(self):
        self.assertFalse(any(i.key.startswith("cap.external.") for i in fd.items(_project("2군"))))
        self.assertTrue(any(i.key.startswith("cap.external.") for i in fd.items(_project("1군"))))

    def test_item_without_any_basis_is_not_drafted(self):
        p = _project(files=("02_설비대장.xlsx",))
        investigation = _item(p, "cap.internal.investigation")
        self.assertEqual(investigation.state, fd.NO_BASIS)
        self.assertEqual(fd.generate(p), [])


class RuleDraftTests(unittest.TestCase):
    def test_rule_draft_uses_only_facts_and_marks_missing(self):
        p = create_direct_cap_project("가상 사업장", address="가상 주소", cap_group="2군")
        ew.ingest(p, "plan.docx", _doc([("교육 대상", "전 직원"), ("공정개요", "혼합 후 충전")]), "site")
        ew.publish(p, replace_existing=True)
        results = fd.generate(p)
        self.assertEqual([r["key"] for r in results], ["cap.prevention.training"])
        text = _item(p, "cap.prevention.training").draft
        self.assertIn("교육 대상: 전 직원", text)
        self.assertIn("[확인 필요: 교육 주기]", text)
        self.assertIn("[확인 필요: 교육 방법]", text)
        self.assertEqual(results[0]["checks"], 2)

    def test_refresh_after_new_upload(self):
        p = _project()
        self.assertGreater(fd.refresh_rule_drafts(p), 0)
        self.assertEqual(fd.refresh_rule_drafts(p), 0)          # 이미 있으면 다시 만들지 않음
        ew.ingest(p, "more.docx", _doc([("대표자", "김가상"), ("대표전화", "02-1")]), "site")
        stale = _item(p, "cap.prevention.safety_management")
        self.assertTrue(stale.stale)
        self.assertIn("다시 만들어야", next(r["상태"] for r in fd.rows(p) if r["_key"] == stale.key))
        self.assertGreater(fd.refresh_rule_drafts(p), 0)
        self.assertEqual(_item(p, "cap.prevention.safety_management").state, fd.HAS_DRAFT)


class AIDraftTests(unittest.TestCase):
    def test_values_not_in_the_files_become_check_marks(self):
        p = _project()
        client = FakeClient("공장장이 안전관리를 총괄하고 TK-99 설비와 108-88-3 물질을 매주 3회 점검한다. 월 1회 안전점검 회의를 연다.")
        result = fd.generate(p, client, keys=["cap.prevention.safety_management"])[0]
        self.assertEqual(sorted(result["replaced"]), sorted(["TK-99", "108-88-3", "3회"]))
        text = _item(p, "cap.prevention.safety_management").draft
        self.assertNotIn("TK-99", text.replace("‘TK-99’", ""))
        self.assertIn("[확인 필요: 근거 자료에 없는 값 ‘TK-99’]", text)
        self.assertIn("월 1회", text)                               # 파일에 있는 값은 그대로
        self.assertEqual(result["model"], "fake-local")

    def test_missing_items_are_added_when_ai_leaves_them_out(self):
        p = create_direct_cap_project("가상 사업장", address="가상 주소", cap_group="2군")
        ew.ingest(p, "plan.docx", _doc([("교육 대상", "전 직원"), ("공정개요", "혼합 후 충전")]), "site")
        ew.publish(p, replace_existing=True)
        fd.generate(p, FakeClient("전 직원을 대상으로 교육한다."))
        text = _item(p, "cap.prevention.training").draft
        self.assertIn("[확인 필요: 교육 주기]", text)

    def test_ai_failure_falls_back_to_rule_draft(self):
        p = _project()
        result = fd.generate(p, BrokenClient(), keys=["cap.prevention.training"])[0]
        self.assertEqual(result["model"], fd.RULE_MODEL)
        self.assertIn("규칙 초안", result["note"])
        self.assertIn("교육 대상:", _item(p, "cap.prevention.training").draft)


class ProgressTests(unittest.TestCase):
    def test_progress_counts_only_eligible_items_and_reports_before_ai_call(self):
        p = _project()
        events = []
        class Client(BrokenClient):
            def generate_json(self, **kwargs):
                self_test.assertTrue(events[-1][2])
                return super().generate_json(**kwargs)
        self_test = self
        expected = [i for i in fd.items(p) if i.state != fd.FROM_FILE and i.basis]
        results = fd.generate(p, Client(), refresh=True, on_progress=lambda *args: events.append(args))
        self.assertGreater(len(expected), 0)
        self.assertEqual(events[0], (0, len(expected), ""))
        self.assertEqual(events[-1], (len(expected), len(expected), ""))
        self.assertEqual([e[0] for e in events if e[2]], list(range(len(expected))))
        self.assertEqual(len(results), len(expected))
        self.assertTrue(all(r["note"] for r in results))  # AI 실패 시 규칙 초안으로 완료

    def test_empty_run_does_not_claim_completed_work(self):
        p = _project(files=("02_설비대장.xlsx",))
        events = []
        self.assertEqual(fd.generate(p, on_progress=lambda *args: events.append(args)), [])
        self.assertEqual(events, [(0, 0, "")])

    def test_group_one_external_plans_can_use_ai_with_company_basis(self):
        p = create_direct_cap_project("가상", address="주소", cap_group="1군")
        ew.ingest(p, "plan.docx", _doc([("지역사회 소통 대상", "인근 주민"),
                  ("협력 기관", "관할 소방서"), ("대피 장소", "지정 집결지")]), "site")
        ew.publish(p, replace_existing=True)
        candidates = [i for i in fd.items(p) if i.key.startswith("cap.external.") and i.basis]
        self.assertEqual(len(candidates), 3)
        results = fd.generate(p, FakeClient("회사 자료에 따라 계획을 운영한다."), keys=[i.key for i in candidates])
        self.assertEqual(len(results), 3)
        self.assertTrue(all(r["model"] == "fake-local" for r in results))


class OutputTests(unittest.TestCase):
    def test_drafts_are_listed_and_put_into_the_internal_draft(self):
        from engine.stage2 import report_export as rx

        p = _project()
        fd.refresh_rule_drafts(p)
        rows = {r["requirement_key"]: r for r in drafting.ai_written_items(p, "CAP")}
        self.assertIn("cap.prevention.safety_management", rows)
        self.assertIn("AI 초안", rows["cap.prevention.safety_management"]["state"])
        outputs = rx.build_outputs(p, "CAP")
        text = "\n".join(par.text for par in Document(BytesIO(outputs.review_docx)).paragraphs)
        self.assertIn("[AI 초안", text)
        self.assertIn("안전관리 책임자: 공장장", text)
        self.assertIn(d.DATA["emergency"]["shutdown"], text)        # 회사 문서의 글은 그대로


if __name__ == "__main__":
    unittest.main()
