from __future__ import annotations

from io import BytesIO
import unittest

from openpyxl import load_workbook

from engine.stage2.integrated_workbook import (
    META_SHEET,
    apply_integrated_authoring_workbook,
    build_integrated_authoring_workbook,
)
from engine.stage2.project import EvidenceRef, Stage2Project


class IntegratedStage2WorkbookTests(unittest.TestCase):
    def _project(self, *, psm=True, cap=True) -> Stage2Project:
        project = Stage2Project(
            project_id="S2-TEST-INTEGRATED",
            company_name="테스트화학",
            psm_required=psm,
            cap_required=cap,
            cap_group="1군" if cap else "",
            scope_confirmed=True,
            psm_selected=psm,
            cap_selected=cap,
        )
        project.set_field("business.company_name", "회사명", "테스트화학", "VERIFIED")
        project.set_field("business.address", "사업장 소재지", "대구광역시 테스트로 1", "VERIFIED")
        project.set_field(
            "inventory.chemicals",
            "화학물질 목록",
            [{"물질명": "톨루엔", "CAS 번호": "108-88-3", "함량(%)": 99.5, "최대보유량": 15000, "단위": "kg"}],
            "VERIFIED",
        )
        project.set_field(
            "inventory.facilities",
            "시설별 최대보유량 자료",
            [{"설비번호": "TK-101", "설비명": "톨루엔 저장탱크", "최대보유량": 15000}],
            "VERIFIED",
        )
        return project

    def test_both_reports_use_one_common_workbook_with_scope_specific_sheets(self):
        project = self._project(psm=True, cap=True)
        data = build_integrated_authoring_workbook(project)
        wb = load_workbook(BytesIO(data), data_only=False)

        self.assertEqual(wb.sheetnames.count("02_화학물질정보"), 1)
        self.assertEqual(wb.sheetnames.count("03_설비정보"), 1)
        self.assertIn("04_안전밸브_파열판", wb.sheetnames)
        self.assertIn("10_동력기계", wb.sheetnames)
        self.assertIn("11_배관_개스킷", wb.sheetnames)
        self.assertIn("20_배출물질_처리시설", wb.sheetnames)
        self.assertEqual(wb["01_사업장정보"]["B5"].value, "테스트화학")
        self.assertEqual(wb["02_화학물질정보"]["B5"].value, "108-88-3")
        self.assertEqual(wb["03_설비정보"]["A5"].value, "TK-101")
        self.assertEqual(wb[META_SHEET]["B2"].value, "INPUT")
        self.assertEqual(wb[META_SHEET].sheet_state, "hidden")

    def test_single_scope_hides_other_report_specific_tables(self):
        psm_project = self._project(psm=True, cap=False)
        psm_wb = load_workbook(BytesIO(build_integrated_authoring_workbook(psm_project)))
        self.assertIn("10_동력기계", psm_wb.sheetnames)
        self.assertNotIn("20_배출물질_처리시설", psm_wb.sheetnames)

        cap_project = self._project(psm=False, cap=True)
        cap_wb = load_workbook(BytesIO(build_integrated_authoring_workbook(cap_project)))
        self.assertNotIn("10_동력기계", cap_wb.sheetnames)
        self.assertIn("20_배출물질_처리시설", cap_wb.sheetnames)

    def test_example_workbook_cannot_be_submitted(self):
        project = self._project()
        example = build_integrated_authoring_workbook(project, example=True)
        with self.assertRaisesRegex(ValueError, "작성예시 파일은 제출할 수 없습니다"):
            apply_integrated_authoring_workbook(project, example)

    def test_common_equipment_and_relief_tables_feed_both_reports(self):
        project = self._project()
        data = build_integrated_authoring_workbook(project)
        wb = load_workbook(BytesIO(data))

        equipment = wb["03_설비정보"]
        equipment["A5"] = "TK-101"
        equipment["B5"] = "톨루엔 저장탱크"
        equipment["C5"] = "저장탱크"
        equipment["D5"] = "원료저장"
        equipment["E5"] = "톨루엔"
        equipment["H5"] = "0.49 MPa"
        equipment["I5"] = "80 ℃"
        equipment["J5"] = "0.15 MPa"
        equipment["K5"] = "30 ℃"
        equipment["L5"] = "SUS304"

        relief = wb["04_안전밸브_파열판"]
        relief["A5"] = "PSV-101"
        relief["B5"] = "TK-101"
        relief["C5"] = "안전밸브"
        relief["D5"] = "0.45 MPa"
        relief["E5"] = "1200 kg/h"
        relief["F5"] = "톨루엔 증기"

        process = wb["06_공정정보"]
        process["B5"] = "원료 저장 후 혼합공정으로 이송하여 제품을 생산한다."

        output = BytesIO()
        wb.save(output)
        evidence = EvidenceRef(
            source_type="STAGE2_INTEGRATED_WORKBOOK",
            source_name="통합작성자료.xlsx",
            sha256="a" * 64,
        )
        result = apply_integrated_authoring_workbook(project, output.getvalue(), workbook_evidence=evidence)

        self.assertGreater(result.updated_fields, 0)
        psm_equipment = project.get_field("psm.psi.equipment_specs")
        cap_equipment = project.get_field("cap.facility.equipment_specs")
        self.assertIsNotNone(psm_equipment)
        self.assertIsNotNone(cap_equipment)
        self.assertEqual(psm_equipment.value, cap_equipment.value)
        self.assertEqual(psm_equipment.value[0]["설비번호"], "TK-101")

        psm_relief = project.get_field("psm.psi.relief_device_specs")
        cap_relief = project.get_field("cap.safety.relief_device_specs")
        self.assertEqual(psm_relief.value, cap_relief.value)
        self.assertEqual(psm_relief.value[0]["보호대상 설비번호"], "TK-101")
        self.assertEqual(project.get_field("process.description").status, "USER_CONFIRMED")

    def test_stage1_protected_value_is_not_overwritten(self):
        project = self._project()
        data = build_integrated_authoring_workbook(project)
        wb = load_workbook(BytesIO(data))
        wb["01_사업장정보"]["B5"] = "다른회사"
        output = BytesIO()
        wb.save(output)

        result = apply_integrated_authoring_workbook(project, output.getvalue())
        self.assertEqual(project.get_field("business.company_name").value, "테스트화학")
        self.assertTrue(any("판정진단을 다시 수행" in warning for warning in result.warnings))

    def test_attachment_declaration_stays_hold_until_file_is_uploaded(self):
        project = self._project()
        data = build_integrated_authoring_workbook(project)
        wb = load_workbook(BytesIO(data))
        ws = wb["07_도면_첨부자료목록"]
        target_row = None
        for row in range(5, ws.max_row + 1):
            if "공정흐름도" in str(ws.cell(row, 1).value or ""):
                target_row = row
                break
        self.assertIsNotNone(target_row)
        ws.cell(target_row, 2, "PFD_Rev3.pdf")
        ws.cell(target_row, 3, "PFD-001")
        output = BytesIO()
        wb.save(output)

        result = apply_integrated_authoring_workbook(project, output.getvalue())
        self.assertGreater(result.attachment_declarations, 0)
        record = project.get_field("documents.pfd")
        self.assertIsNotNone(record)
        self.assertEqual(record.status, "HOLD")
        self.assertEqual(record.value["file_name"], "PFD_Rev3.pdf")


if __name__ == "__main__":
    unittest.main()
