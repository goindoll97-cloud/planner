from __future__ import annotations

from io import BytesIO
import unittest

from openpyxl import load_workbook

from engine.stage2.cap_form2_engine import PASS, build_cap_form2_readiness
from engine.stage2.integrated_workbook import apply_integrated_authoring_workbook
from engine.stage2.project import Stage2Project
from engine.stage2.workbook_enhancements import build_enhanced_integrated_authoring_workbook


class CAPForm2ChangeLogWorkbookTests(unittest.TestCase):
    def _project(self) -> Stage2Project:
        project = Stage2Project(
            project_id="S2-CAP-FORM2-WORKBOOK",
            company_name="테스트화학",
            site_name="제1공장",
            cap_required=True,
            cap_group="1군",
            scope_confirmed=True,
            cap_selected=True,
        )
        project.set_field(
            "business.company_name",
            "회사명",
            "테스트화학",
            "USER_CONFIRMED",
        )
        project.set_field(
            "business.address",
            "사업장 소재지",
            "테스트시 테스트로 1",
            "USER_CONFIRMED",
        )
        project.set_field(
            "cap.business.submission_type",
            "제출구분",
            "변경제출",
            "USER_CONFIRMED",
        )
        project.set_field(
            "cap.business.unit_plant_name",
            "단위공장명",
            "제1공장",
            "USER_CONFIRMED",
        )
        return project

    @staticmethod
    def _save(wb) -> bytes:
        out = BytesIO()
        wb.save(out)
        return out.getvalue()

    def test_input_workbook_exposes_change_log_as_structured_table(self):
        project = self._project()
        data = build_enhanced_integrated_authoring_workbook(project, example=False)
        wb = load_workbook(BytesIO(data), data_only=False)

        self.assertIn("29_변경내역_관리대장", wb.sheetnames)
        ws = wb["29_변경내역_관리대장"]
        headers = [str(cell.value or "").strip() for cell in ws[4]][:6]
        self.assertEqual(
            headers,
            [
                "일자",
                "변경항목",
                "변경의 종류",
                "변경 내용(변경전 → 변경후)",
                "후속조치",
                "담당자",
            ],
        )

        meta = wb["_시스템정보"]
        change_log_records = [
            row
            for row in meta.iter_rows(min_row=8, values_only=True)
            if str(row[2] or "") == "cap.prevention.change_log"
        ]
        self.assertEqual(len(change_log_records), 1)
        self.assertEqual(str(change_log_records[0][0]), "TABLE")

    def test_change_log_round_trip_satisfies_form2_readiness(self):
        project = self._project()
        data = build_enhanced_integrated_authoring_workbook(project, example=False)
        wb = load_workbook(BytesIO(data), data_only=False)
        ws = wb["29_변경내역_관리대장"]

        values = [
            "2026-09-18",
            "장치·설비 목록 및 명세",
            "설비변경",
            "TK-101 저장탱크 → TK-101A 저장탱크",
            "관련 도면 및 설비명세 갱신",
            "환경안전팀 홍길동",
        ]
        for col, value in enumerate(values, start=1):
            ws.cell(5, col, value)

        result = apply_integrated_authoring_workbook(project, self._save(wb))
        self.assertGreaterEqual(result.table_fields, 1)

        record = project.get_field("cap.prevention.change_log")
        self.assertIsNotNone(record)
        self.assertEqual(record.status, "USER_CONFIRMED")
        self.assertIsInstance(record.value, list)
        self.assertEqual(record.value[0]["변경항목"], "장치·설비 목록 및 명세")
        self.assertEqual(
            record.value[0]["변경 내용(변경전 → 변경후)"],
            "TK-101 저장탱크 → TK-101A 저장탱크",
        )

        readiness = build_cap_form2_readiness(project)
        self.assertEqual(readiness.status, PASS)
        self.assertTrue(readiness.ready)

    def test_redownload_preserves_confirmed_change_log_rows(self):
        project = self._project()
        project.set_field(
            "cap.prevention.change_log",
            "변경내역 관리대장",
            [{
                "일자": "2026-09-18",
                "변경항목": "공정배관계장도",
                "변경의 종류": "도면변경",
                "변경 내용(변경전 → 변경후)": "PID-101 Rev.1 → Rev.2",
                "후속조치": "관련 문서 개정",
                "담당자": "김담당",
            }],
            "USER_CONFIRMED",
        )

        data = build_enhanced_integrated_authoring_workbook(project, example=False)
        wb = load_workbook(BytesIO(data), data_only=False)
        ws = wb["29_변경내역_관리대장"]

        self.assertEqual(ws.cell(5, 1).value, "2026-09-18")
        self.assertEqual(ws.cell(5, 2).value, "공정배관계장도")
        self.assertEqual(ws.cell(5, 4).value, "PID-101 Rev.1 → Rev.2")
        self.assertEqual(ws.cell(5, 6).value, "김담당")

    def test_example_workbook_contains_only_example_change_log(self):
        project = self._project()
        data = build_enhanced_integrated_authoring_workbook(project, example=True)
        wb = load_workbook(BytesIO(data), data_only=False)
        ws = wb["29_변경내역_관리대장"]

        self.assertEqual(ws.cell(5, 2).value, "장치·설비 목록 및 명세")
        self.assertIn("TK-101", str(ws.cell(5, 4).value))


if __name__ == "__main__":
    unittest.main()
