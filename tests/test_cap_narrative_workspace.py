from __future__ import annotations

from pathlib import Path
import unittest
from unittest.mock import patch

from streamlit.testing.v1 import AppTest

from engine.stage2 import ai_drafting as drafting
from engine.stage2 import cap_narrative_workspace as cn
from engine.stage2 import narrative_examples as ex
from engine.stage2 import psm_narrative_workspace as nw
from engine.stage2 import psm_table_workspace as tables
from tests.test_stage2_cap_form1_engine import CAPForm1EngineTests

ROOT = Path(__file__).resolve().parents[1]
PROFILE = cn.CAP_PROFILE


def _project(group="1군"):
    project = CAPForm1EngineTests()._project()
    project.cap_group = group
    for fact in nw.BASIC_FACTS:
        nw.save_fact(project, fact, "염소를 탱크에 받아 이송한다." if fact.long else "연속 운전")
    return project


def _used():
    keys = ["inventory.chemicals"]
    if _current["project"].get_field("cap.internal.response_equipment") is not None:
        keys.append("cap.internal.response_equipment")  # 항목 고유 사실이 있으면 그 키를 연결해야 검증을 통과한다
    return keys


class FakeClient:
    model = "fake-local"

    def generate_json(self, *, instructions, prompt):
        return {"profile_summary": "염소를 저장하고 이송하는 사업장입니다.",
                "drafts": [{"requirement_key": key,
                            "draft_text": "확인된 설비와 물질을 기준으로 사고 예방과 대응 절차를 단계별로 기술합니다.",
                            "used_fact_keys": _used(),
                            "suggested_additions": ["회사 고유 절차가 있으면 추가해 주세요."]}
                           for key in cn.items(_current["project"])]}


_current: dict = {}


class ItemTests(unittest.TestCase):
    def test_first_group_writes_all_three_chapters_second_group_skips_external(self):
        first, second = _project("1군"), _project("2군")
        first_items, second_items = cn.items(first), cn.items(second)
        self.assertTrue(any(k.startswith("cap.external") for k in first_items))
        self.assertFalse(any(k.startswith("cap.external") for k in second_items))
        self.assertTrue(any(k.startswith("cap.internal") for k in second_items))
        self.assertNotIn(cn.CONTACT_ITEM, first_items)  # 연락처는 표로 받는다
        self.assertTrue(cn.includes_external(first) and not cn.includes_external(second))
        self.assertEqual({f.key for f in cn.visible_facts(second)} & cn.EXTERNAL_FACT_KEYS, set())
        self.assertEqual(len(cn.visible_facts(first)), len(cn.DECISION_FACTS))

    def test_items_wait_for_the_basic_facts(self):
        project = CAPForm1EngineTests()._project()
        project.cap_group = "1군"
        states = {i["key"]: i for i in nw.item_status(project, PROFILE)}
        self.assertTrue(all(i["state"] == "만들 수 없음" for i in states.values()))
        self.assertTrue(all("공정 개요" in i["reason"] for i in states.values()))

    def test_drafts_are_stored_unconfirmed_then_adopted_into_the_report_fields(self):
        project = _project("1군")
        _current["project"] = project
        with patch("engine.stage2.storage.save_project"):
            result = nw.generate(project, FakeClient(), profile=PROFILE)
        self.assertTrue(result.generated, [r.validation_warnings for r in result.rejected])
        item = result.generated[0].requirement_key
        record = project.get_field(drafting.ai_draft_field_key("CAP", item))
        self.assertEqual(record.status, "AI_DRAFT")
        spec = nw._specs(project, PROFILE)[item]
        self.assertIsNone(project.get_field(spec.field_keys[0]))
        nw.adopt(project, item, profile=PROFILE)
        self.assertEqual(project.get_field(spec.field_keys[0]).status, "USER_CONFIRMED")
        self.assertEqual({i["key"]: i["state"] for i in nw.item_status(project, PROFILE)}[item], "확인 완료")

    def test_adopting_keeps_tables_that_share_a_requirement(self):
        project = _project("1군")
        tables.save(project, "cap-resources", [{"구분": "장비", "명칭": "공기호흡기", "수량": "4", "보관위치": "제어실"}])
        _current["project"] = project
        with patch("engine.stage2.storage.save_project"):
            result = nw.generate(project, FakeClient(), profile=PROFILE)
        self.assertIn("cap.internal.resources", [g.requirement_key for g in result.generated])
        nw.adopt(project, "cap.internal.resources", profile=PROFILE)
        rows = project.get_field("cap.internal.response_equipment").value
        self.assertIsInstance(rows, list)  # 방재 장비 표는 글로 덮어쓰지 않는다
        self.assertEqual(rows[0]["명칭"], "공기호흡기")

    def test_an_item_of_the_other_group_cannot_be_adopted(self):
        project = _project("2군")
        with self.assertRaises(ValueError):
            nw.adopt(project, "cap.external.evacuation", "글", profile=PROFILE)


class ExampleCoverageTests(unittest.TestCase):
    def test_every_cap_decision_field_has_examples_without_duplicates(self):
        for fact in cn.DECISION_FACTS:
            for name, _label, _help in fact.fields:
                options = ex.choices(fact.key, name)
                self.assertGreaterEqual(len(options), 3, (fact.key, name))
                self.assertEqual(len(options), len(set(options)), (fact.key, name))

    def test_saved_example_verbatim_is_flagged(self):
        project = _project()
        fact = cn.DECISION_FACTS[0]
        nw.save_fact(project, fact, {"책임자": ex.choices(fact.key, "책임자")[0], "조직": "우리 회사 방식", "회의": ""})
        self.assertIn(fact.label, nw.chosen_as_is(project, PROFILE))
        self.assertNotIn("안전관리 책임자", nw.chosen_as_is(project, PROFILE))


class ScreenTests(unittest.TestCase):
    def _run(self, project, option):
        from engine.stage2.local_llm import LocalLLMProbe

        rows = [{"project_id": project.project_id, "company_name": project.company_name}]
        down = LocalLLMProbe(False, "ollama", "http://127.0.0.1:11434", message="연결되지 않았습니다.")
        with patch("engine.stage2.storage.list_projects", return_value=rows), \
             patch("engine.stage2.storage.load_project", return_value=project), \
             patch("engine.stage2.storage.save_project"), \
             patch("engine.stage2.local_llm.probe_local_llm_runtime", return_value=down):
            at = AppTest.from_file(str(ROOT / "ui/cap_workspace_page.py"), default_timeout=90)
            at.session_state["_stage2_active_project_id"] = project.project_id
            at.run()
            at.selectbox(key="cap_form_no").select(option).run()
            return at

    def test_narrative_and_export_options_render_for_a_cap_project(self):
        project = _project("1군")
        for option in ("narrative", "export"):
            at = self._run(project, option)
            self.assertFalse(at.exception, option)
        at = self._run(project, "narrative")
        labels = [e.label for e in at.expander]
        self.assertTrue(any("안전관리 운영" in label for label in labels))
        self.assertTrue(any("주민 보호" in label for label in labels))  # 1군은 외부 비상대응 사실도 받는다

    def test_second_group_does_not_ask_for_external_plan_facts(self):
        at = self._run(_project("2군"), "narrative")
        self.assertFalse(at.exception)
        labels = " ".join(e.label for e in at.expander)
        self.assertIn("안전관리 운영", labels)
        self.assertNotIn("주민 보호", labels)


if __name__ == "__main__":
    unittest.main()


class FinalDocumentTests(unittest.TestCase):
    def test_adopted_cap_text_appears_in_the_narrative_chapters(self):
        from docx import Document

        from engine.stage2 import statutory_report_v2 as v2

        project = _project("1군")
        _current["project"] = project
        with patch("engine.stage2.storage.save_project"):
            result = nw.generate(project, FakeClient(), profile=PROFILE)
        for item in result.generated:
            nw.adopt(project, item.requirement_key, profile=PROFILE)
        doc = Document()
        v2._render_cap_narrative(doc, project)
        text = "\n".join(p.text for p in doc.paragraphs)
        self.assertIn("사고 예방과 대응 절차를 단계별로 기술합니다", text)
        from io import BytesIO

        from engine.stage2.cap_baseline_docx import build_cap_baseline_draft

        baseline = Document(BytesIO(build_cap_baseline_draft(project)))
        cells = chr(10).join(c.text for t in baseline.tables for r in t.rows for c in r.cells)
        self.assertIn("염소를 탱크에 받아 이송한다.", cells)  # 공정 개요(사실)는 규정서식 표에 들어간다
