from __future__ import annotations

import unittest
from io import BytesIO
from pathlib import Path

from openpyxl import load_workbook

from engine.template import build_legal_reference_workbook, build_minimal_input_workbook


class UserFacingAbbreviationContractTests(unittest.TestCase):
    def _joined_workbook_text(self, payload: bytes) -> str:
        wb = load_workbook(BytesIO(payload), data_only=False)
        return '\n'.join(
            str(cell.value)
            for ws in wb.worksheets
            for row in ws.iter_rows()
            for cell in row
            if cell.value is not None
        )

    def test_generated_workbooks_do_not_expose_psm_or_hwagye_abbreviations(self) -> None:
        for payload in (build_minimal_input_workbook(), build_legal_reference_workbook()):
            text = self._joined_workbook_text(payload)
            self.assertNotIn('PSM', text)
            self.assertNotIn('화사계', text)
            self.assertNotIn('화관법', text)

    def test_main_diagnosis_page_uses_full_report_name(self) -> None:
        text = Path('ui/diagnosis_page.py').read_text(encoding='utf-8')
        self.assertIn('PSM_FULL = "공정안전보고서"', text)
        self.assertNotIn('공정안전보고서(PSM)', text)


if __name__ == '__main__':
    unittest.main()
