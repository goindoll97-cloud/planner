"""사실 카탈로그로 넓힌 파일 읽기: 사업장 일반정보·공정 설명, 설비 사양·확산방지설비, SDS 물성·유해성, 가스감지기."""
from __future__ import annotations

import unittest
from io import BytesIO

from docx import Document
from openpyxl import Workbook

from engine.stage2 import demo_full_company as d
from engine.stage2 import evidence_workspace as ew
from engine.stage2.direct_template import create_direct_cap_project

DEMO_UPLOADS = ("01_제품목록.xlsx", "02_설비대장.xlsx", "03_사업장_일반정보.docx", "04_공정설명서.docx",
                "05_SDS_물성표.xlsx", "06_설비사양_안전설비.xlsx")


def _docx(rows=(), lines=()):
    doc = Document()
    for line in lines:
        doc.add_paragraph(line)
    if rows:
        table = doc.add_table(rows=0, cols=2)
        for label, value in rows:
            cells = table.add_row().cells
            cells[0].text, cells[1].text = label, value
    out = BytesIO()
    doc.save(out)
    return out.getvalue()


def _xlsx(sheets):
    book = Workbook()
    book.remove(book.active)
    for title, rows in sheets.items():
        sheet = book.create_sheet(title)
        for row in rows:
            sheet.append(list(row))
    out = BytesIO()
    book.save(out)
    return out.getvalue()


def _value(project, entity, field):
    slot = next((s for s in ew.slots(project) if s["entity"] == entity and s["field"] == field), None)
    return None if slot is None or not slot["chosen"] else slot["chosen"]["value"]


class SiteFactTests(unittest.TestCase):
    def setUp(self):
        self.p = create_direct_cap_project("가상 사업장", address="가상 주소 1", cap_group="")

    def test_label_value_table_is_read_as_site_facts_and_published(self):
        content = _docx(rows=[("사업자 등록번호", "000-81-11111"), ("대표자", "홍가상"), ("대표전화", "02-000-0000")])
        self.assertEqual(ew.classify_document("일반정보.docx", content), "site")
        self.assertEqual(ew.ingest(self.p, "일반정보.docx", content, "site"), 3)
        self.assertEqual(_value(self.p, ew.SITE_ENTITY, "representative"), "홍가상")
        ew.publish(self.p, replace_existing=True)
        record = self.p.get_field("cap.business.representative")
        self.assertEqual(record.value, "홍가상")
        self.assertEqual(record.status, "FILE_READ")             # 파일에서 읽은 값(사람이 확인한 값과 구분)
        self.assertTrue(record.confirmed)
        self.assertTrue(record.note.startswith("공통정보 연결"))
        self.assertEqual(record.evidence[0].source_name, "일반정보.docx")

    def test_label_lines_in_a_site_document_are_read(self):
        content = "사업장 일반정보\n대표자: 홍가상\n사업자등록번호 : 000-81-22222\n".encode("utf-8")
        self.assertEqual(ew.classify_document("info.txt", content), "site")
        ew.ingest(self.p, "info.txt", content, "site")
        self.assertEqual(_value(self.p, ew.SITE_ENTITY, "registration_no"), "000-81-22222")

    def test_single_label_line_in_another_document_is_not_taken_as_site_fact(self):
        # SDS 같은 문서의 공급자 정보를 사업장 정보로 착각하지 않도록: 사업장 문서가 아니면 한 줄만으로는 읽지 않습니다
        content = _docx(lines=["공급자 정보", "대표자: 공급사 사장", "긴급전화: 119"])
        ew.ingest(self.p, "sds.docx", content, "material")
        self.assertIsNone(_value(self.p, ew.SITE_ENTITY, "representative"))

    def test_two_files_disagree_so_the_value_is_not_sent(self):
        ew.ingest(self.p, "a.docx", _docx(rows=[("대표자", "홍가상"), ("대표전화", "02-1")]), "site")
        ew.ingest(self.p, "b.docx", _docx(rows=[("대표자", "김가상"), ("대표전화", "02-1")]), "site")
        slot = next(s for s in ew.slots(self.p) if s["key"] == ew.SITE_ENTITY + "|representative")
        self.assertEqual(slot["status"], "값 충돌")
        ew.publish(self.p, replace_existing=True)
        self.assertIsNone(self.p.get_field("cap.business.representative"))
        self.assertEqual(self.p.get_field("cap.business.contact").value, "02-1")

    def test_value_entered_elsewhere_is_not_overwritten(self):
        self.p.set_field("cap.business.representative", "대표자", "직접 적은 값", "USER_CONFIRMED", note="사람이 입력")
        ew.ingest(self.p, "a.docx", _docx(rows=[("대표자", "홍가상"), ("대표전화", "02-1")]), "site")
        ew.publish(self.p, replace_existing=True)
        self.assertEqual(self.p.get_field("cap.business.representative").value, "직접 적은 값")

    def test_loading_count_becomes_the_form_text(self):
        ew.ingest(self.p, "공정.docx", _docx(rows=[("공정개요", "저장 후 반응"), ("입·출하 시설 수(기)", "2")]), "site")
        ew.publish(self.p, replace_existing=True)
        self.assertEqual(self.p.get_field("cap.basic.loading_transport").value, "입·출하 시설 2기")
        self.assertEqual(self.p.get_field("process.description").value, "저장 후 반응")

    def test_reupload_needs_publish_again_even_when_values_are_unchanged(self):
        ew.ingest(self.p, "a.docx", _docx(rows=[("대표자", "홍가상"), ("대표전화", "02-1")]), "site")
        ew.publish(self.p, replace_existing=True)
        self.assertFalse(ew.needs_publish(self.p))
        ew.ingest(self.p, "c.csv", "설비번호,설비명\nTK-1,탱크\n".encode("utf-8"), "facility")
        self.assertEqual(self.p.get_field("cap.business.representative").status, "HOLD")   # 새 파일을 읽으면 다시 확인
        self.assertTrue(ew.needs_publish(self.p))
        ew.publish(self.p, replace_existing=True)
        self.assertEqual(self.p.get_field("cap.business.representative").status, "FILE_READ")


class EntityFactTests(unittest.TestCase):
    def setUp(self):
        self.p = create_direct_cap_project("가상 사업장", address="가상 주소 1", cap_group="")

    def test_detector_sheet_in_a_facility_workbook(self):
        content = _xlsx({"설비": [("설비번호", "설비명", "용량(m3)"), ("TK-1", "탱크", 10)],
                         "감지기": [("감지기 번호", "검출대상 물질", "설치위치", "경보 설정값", "연동여부"),
                                  ("GD-1", "톨루엔", "TK-1 옆", "LEL 20%", "예")]})
        self.assertEqual(ew.classify_document("설비.xlsx", content), "facility")
        ew.ingest(self.p, "설비.xlsx", content, "facility")
        self.assertEqual(_value(self.p, "detector:GD-1", "detector_alarm_setting"), "LEL 20%")
        self.assertEqual(_value(self.p, "facility:TK-1", "capacity"), 10)
        ew.publish(self.p, replace_existing=True)
        rows = self.p.get_field("cap.safety.gas_detection").value
        self.assertEqual(rows[0]["감지기 번호"], "GD-1")
        self.assertEqual(rows[0]["연동여부"], "예")
        self.assertEqual(rows[0]["설치형태"], "고정식")

    def test_sds_table_without_product_code_joins_the_product_by_cas_in_either_order(self):
        products = "제품코드,제품명,CAS 번호,함량(%)\nP-1,가상 톨루엔,108-88-3,99\n".encode("utf-8")
        sds = "제품명,CAS 번호,폭발한계 하한,증기압\n톨루엔,108-88-3,1.1 %,22\n".encode("utf-8")
        for order in ((("p.csv", products), ("s.csv", sds)), (("s.csv", sds), ("p.csv", products))):
            project = create_direct_cap_project("가상", address="주소", cap_group="")
            for name, content in order:
                ew.ingest(project, name, content, "material")
            self.assertEqual(_value(project, "material:P-1", "lel"), "1.1 %")
            self.assertFalse(any(s["entity"] == "material:톨루엔" for s in ew.slots(project)))

    def test_cas_shared_by_two_products_is_not_guessed(self):
        products = "제품코드,제품명,CAS 번호\nP-1,신나 A,108-88-3\nP-2,신나 B,108-88-3\n".encode("utf-8")
        sds = "제품명,CAS 번호,폭발한계 하한\n톨루엔,108-88-3,1.1 %\n".encode("utf-8")
        ew.ingest(self.p, "p.csv", products, "material")
        ew.ingest(self.p, "s.csv", sds, "material")
        self.assertTrue(any(s["entity"] == "material:톨루엔" for s in ew.slots(self.p)))
        self.assertIsNone(_value(self.p, "material:P-1", "lel"))

    def test_catalog_choice_spelling_is_matched(self):
        content = "설비번호,확산방지설비 적용여부\nTK-1,해당없음\n".encode("utf-8")
        ew.ingest(self.p, "dike.csv", content, "facility")
        self.assertEqual(_value(self.p, "facility:TK-1", "dike_applicable"), "해당 없음")

    def test_generic_columns_are_not_reported_but_whole_unread_tables_are(self):
        ew.ingest(self.p, "x.csv", "설비번호,설비명,비고\nTK-1,탱크,메모\n".encode("utf-8"), "facility")
        self.assertEqual(ew.documents(self.p)[0]["unmapped_columns"], [])

    def test_mixed_workbook_reads_each_sheet_by_its_own_kind(self):
        content = _xlsx({"사업장": [("항목", "값"), ("대표자", "홍가상"), ("대표전화", "02-1")],
                         "설비": [("설비번호", "설비명"), ("TK-1", "탱크")],
                         "물질": [("제품코드", "제품명", "CAS 번호"), ("P-1", "톨루엔", "108-88-3")]})
        self.assertEqual(ew.classify_document("묶음.xlsx", content), "mixed")
        ew.ingest(self.p, "묶음.xlsx", content, "mixed")
        self.assertEqual(_value(self.p, ew.SITE_ENTITY, "representative"), "홍가상")
        self.assertEqual(_value(self.p, "facility:TK-1", "name"), "탱크")
        self.assertEqual(_value(self.p, "material:P-1", "cas"), "108-88-3")


class ReuploadTests(unittest.TestCase):
    def test_fixed_file_with_the_same_name_replaces_the_old_one(self):
        p = create_direct_cap_project("가상", address="주소", cap_group="")
        ew.ingest(p, "eq.csv", "설비번호,용량(m3)\nTK-1,10\n".encode("utf-8"), "facility")
        ew.ingest(p, "other.csv", "설비번호,용량(m3)\nTK-1,12\n".encode("utf-8"), "facility")
        self.assertEqual(next(s for s in ew.slots(p) if s["key"] == "facility:TK-1|capacity")["status"], "값 충돌")
        fixed = "설비번호,용량(m3)\nTK-1,10\nTK-2,5\n".encode("utf-8")
        self.assertEqual(ew.replace_same_name(p, "other.csv", fixed), 1)
        ew.ingest(p, "other.csv", fixed, "facility")
        self.assertEqual([d_["name"] for d_ in ew.documents(p)], ["eq.csv", "other.csv"])
        self.assertEqual(_value(p, "facility:TK-1", "capacity"), 10)
        self.assertEqual(ew.replace_same_name(p, "other.csv", fixed), 0)       # 같은 내용이면 그대로


class FilesOnlyDemoTests(unittest.TestCase):
    """연습 회사 파일만 올리고(화면에 적는 일 없이) 서식 초안이 채워지는지 봅니다."""

    @classmethod
    def setUpClass(cls):
        files = d.files()
        cls.p = create_direct_cap_project(d.COMPANY, address=d.DATA["company"]["주소"], cap_group="2군")
        for name in DEMO_UPLOADS:
            ew.ingest(cls.p, name, files[name], ew.classify_document(name, files[name]))
        ew.publish(cls.p, replace_existing=True)

    def test_clean_files_leave_nothing_to_ask(self):
        self.assertEqual([s["key"] for s in ew.slots(self.p) if not s["chosen"]], [])
        self.assertEqual([d_["name"] for d_ in ew.documents(self.p) if d_["unmapped_columns"]], [])

    def test_form_engines_have_no_company_gaps(self):
        from engine.stage2 import (cap_form3_workspace as f3, cap_form4_workspace as f4, cap_form5_workspace as f5,
                                   cap_form9_workspace as f9, cap_form10_workspace as f10, cap_form11_workspace as f11)
        self.assertEqual(f3.missing_labels(self.p), [])
        self.assertEqual(f4.missing(self.p), [])
        self.assertEqual(f5.missing(self.p), [])
        self.assertEqual(f9.needs(self.p), [])
        self.assertEqual(f10.needs(self.p), [])
        self.assertEqual(f11.needs(self.p), [])

    def test_draft_docx_carries_values_from_every_file(self):
        from engine.stage2.cap_baseline_docx import build_cap_baseline_draft

        doc = Document(BytesIO(build_cap_baseline_draft(self.p)))
        text = "\n".join(c.text for t in doc.tables for row in t.rows for c in row.cells)
        C = d.DATA["company"]
        for value in (C["representative"], C["registration_no"], d.DATA["process"]["공정개요"][:20], "GD-03", "LEL 20%",
                      "흡인 유해성 구분 1", "방류벽", "TK-01", "가상 톨루엔"):
            self.assertIn(value, text)


if __name__ == "__main__":
    unittest.main()
