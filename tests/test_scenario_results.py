"""사고시나리오 결과: 회사 파일로 선정·계산하고, 사람은 지도에 경계만 그립니다."""
from __future__ import annotations

from io import BytesIO
from pathlib import Path
import unittest
from unittest.mock import patch

from openpyxl import Workbook, load_workbook
from streamlit.testing.v1 import AppTest

from engine.stage2 import cap_form8_workspace as f8
from engine.stage2 import cap_impact_workspace as iw
from engine.stage2 import cap_scenario_workspace as sc
from engine.stage2 import cap_surroundings as survey
from engine.stage2 import cap_vworld_lookup as vworld
from engine.stage2 import demo_full_company as d
from engine.stage2 import evidence_workspace as ew
from engine.stage2 import psm_form19_2_workspace as f19
from engine.stage2 import scenario_pipeline as sp
from engine.stage2 import scenario_survey as ss
from tests.test_demo_full_company import CENTER, MAP_OBJECTS, _box, walk

ROOT = Path(__file__).resolve().parents[1]
SCENARIO_FILE = "08_사고시나리오_검토자료.xlsx"


def _book(sheets):
    book = Workbook()
    book.remove(book.active)
    for title, rows in sheets.items():
        sheet = book.create_sheet(title)
        for row in rows:
            sheet.append(list(row))
    out = BytesIO()
    book.save(out)
    return out.getvalue()


def _upload(project, name, content):
    ew.replace_same_name(project, name, content)
    ew.ingest(project, name, content, ew.classify_document(name, content))
    ew.publish(project, replace_existing=True)


def _scenario_file(changes):
    """08 검토자료의 ‘별지12_추가값’ 시트를 설비별로 고친 판."""
    book = load_workbook(BytesIO(d.files()[SCENARIO_FILE]))
    sheet = book["별지12_추가값"]
    headers = [c.value for c in sheet[1]]
    for column in {c for values in changes.values() for c in values}:
        if column not in headers:
            headers.append(column)
            sheet.cell(row=1, column=len(headers), value=column)
    for row in range(2, sheet.max_row + 1):
        tag = sheet.cell(row=row, column=1).value
        for column, value in changes.get(tag, {}).items():
            sheet.cell(row=row, column=headers.index(column) + 1).value = value     # None이면 칸을 비움
    out = BytesIO()
    book.save(out)
    return out.getvalue()


class FakeClient:
    model = "가짜 로컬 AI"

    def __init__(self, answer):
        self.answer = answer
        self.calls = 0

    def generate_json(self, *, instructions, prompt):
        self.calls += 1
        return self.answer(prompt) if callable(self.answer) else self.answer


class ReadingTests(unittest.TestCase):
    def test_tables_are_told_apart(self):
        self.assertEqual(ew._classify_table({"대상설비번호", "선정할사고유형", "액위m"}), "facility")
        self.assertEqual(ew._classify_table({"인터록번호", "대상설비번호", "감지기번호", "가동중지범위"}), "psm:17-2")
        self.assertEqual(ew._classify_table({"보호대상명칭", "보호대상구분", "세부유형"}), "target")

    def test_survey_file_fills_form8_and_its_review(self):
        p, _ = walk()
        self.assertEqual(p.get_field(f8.SITE_KEY).status, ew.FILE_STATUS)
        self.assertEqual(len(f8.saved_rows(p)), 4)
        self.assertTrue(f8.scope_reviewed(p))
        self.assertEqual(f8.needs(p), [])
        self.assertTrue(any(ref.source_name == "07_주변보호대상_조사표.xlsx" for ref in p.get_field(f8.SITE_KEY).evidence))

    def test_no_target_declaration(self):
        p, _ = walk()
        _upload(p, "07_주변보호대상_조사표.xlsx", _book({"보호대상": [("보호대상 명칭", "보호대상 구분")],
                                                       "조사 방법": [("항목", "값"), ("500m 안 보호대상 없음", "예"),
                                                                 ("500m 범위 전체 확인", "예"),
                                                                 ("확인한 지도·도면과 확인일", "브이월드 2026-10-27")]}))
        self.assertTrue(f8.declared_no_target(p))

    def test_scenario_only_numbers_do_not_become_new_facilities(self):
        p, _ = walk()
        _upload(p, "추가.xlsx", _book({"추가": [("대상 설비번호", "액위(m)"), ("XX-99", "2")]}))
        tags = {str(r.get("설비번호")) for r in p.get_field("inventory.facilities").value}
        self.assertNotIn("XX-99", tags)


class PipelineTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.p, _ = walk()
        cls.outcome = sp.refresh(cls.p)

    def test_file_choice_decides_the_accident_kind_and_numbers_come_from_the_engines(self):
        kinds = {s["대상 설비번호"]: s["사고유형"] for s in sc.saved_scenarios(self.p)}
        self.assertEqual(kinds, {"TK-01": sc.FIRE, "TK-02": sc.TOXIC, "RX-01": sc.FIRE, "CY-01": sc.TOXIC})
        self.assertEqual(self.outcome.problems, [])
        self.assertTrue(all(i.radius_m for i in self.outcome.impacts))
        reason = next(s for s in sc.saved_scenarios(self.p) if s["대상 설비번호"] == "TK-01")["선정 근거"]
        self.assertIn("규정수량", reason)
        self.assertIn("회사 검토자료", reason)
        self.assertTrue(self.p.get_field(iw.IMPACT_TABLE_KEY).note.startswith("공통정보 연결"))
        self.assertFalse(sp.needs_refresh(self.p))

    def test_reupload_marks_results_stale_and_not_selected_removes_a_scenario(self):
        p, _ = walk()
        sp.refresh(p)
        _upload(p, SCENARIO_FILE, _scenario_file({"RX-01": {"선정할 사고유형": "선정 안 함"},
                                                   "TK-02": {"선정할 사고유형": "아무거나"}}))
        self.assertTrue(sp.needs_refresh(p))
        outcome = sp.refresh(p)
        tags = {s["대상 설비번호"] for s in sc.saved_scenarios(p)}
        self.assertEqual(tags, {"TK-01", "CY-01"})
        self.assertIn("RX-01", outcome.plan.skipped[0])
        problem = next(x for x in outcome.problems if x.where == "TK-02")
        self.assertIn("‘선정할 사고유형’", problem.fix)
        self.assertIn("사고시나리오 검토자료", problem.fix)

    def test_blank_choice_uses_the_program_proposal(self):
        p, _ = walk()
        _upload(p, SCENARIO_FILE, _scenario_file({"TK-01": {"선정할 사고유형": None}}))
        sp.refresh(p)
        kinds = sorted(s["사고유형"] for s in sc.saved_scenarios(p) if s["대상 설비번호"] == "TK-01")
        self.assertEqual(kinds, sorted(k for k in sc.accident_types(p, "가상 톨루엔")))

    def test_missing_boundary_distance_says_which_file_and_column(self):
        p, _ = walk()
        _upload(p, SCENARIO_FILE, _scenario_file({"CY-01": {"경계까지 거리(m)": None}}))
        outcome = sp.refresh(p)
        fixes = [x.fix for x in outcome.problems if x.where.startswith("CY-01")]
        self.assertTrue(fixes)
        self.assertTrue(all("‘경계까지 거리(m)’" in fix and "CY-01" in fix for fix in fixes), fixes)

    def test_psm_19_2_uses_worst_per_kind_unless_the_file_says_otherwise(self):
        p, _ = walk()
        p.psm_required = True
        p.set_authoring_scope(psm_selected=True, cap_selected=True)
        p.set_field(sp.WEATHER_INPUT_KEY, "기상", {"대기온도(℃)": "33", "습도(%)": "70"}, "FILE_READ")
        outcome = sp.refresh(p)
        marks = {r.row["대상 설비번호"]: r.row["시나리오 구분"] for r in outcome.psm}
        self.assertEqual(marks, {"TK-01": f19.WORST, "TK-02": f19.WORST})      # 사고유형마다 피해반경이 가장 큰 것
        self.assertEqual({r.row["대기온도(℃)"] for r in outcome.psm}, {"33"})
        _upload(p, SCENARIO_FILE, _scenario_file({"CY-01": {"공정안전보고서 시나리오 구분": "대안"}}))
        outcome = sp.refresh(p)
        marks = {r.row["대상 설비번호"]: r.row["시나리오 구분"] for r in outcome.psm}
        self.assertEqual(marks["CY-01"], f19.ALTERNATIVE)
        self.assertTrue(p.get_field(f19.TABLE_KEY).value)

    def test_ai_polishes_reasons_but_cannot_add_numbers(self):
        p, _ = walk()
        sp.refresh(p)
        client = FakeClient({"text": "TK-01은 규정수량 400 kg 이상으로 대상이며 누출량은 999 kg입니다."})
        results = sp.polish_reasons(p, client)
        self.assertTrue(any(r["결과"] == "다듬음" for r in results))
        sp.refresh(p)
        scenario = next(s for s in sc.saved_scenarios(p) if s["대상 설비번호"] == "TK-01")
        self.assertIn("[확인 필요: 근거 자료에 없는 값 ‘999", scenario["선정 근거"])
        self.assertIn("AI", scenario["선정 근거 출처"])
        failing = FakeClient(lambda prompt: (_ for _ in ()).throw(RuntimeError("down")))
        self.assertTrue(all("규칙 문장 유지" in r["결과"] for r in sp.polish_reasons(p, failing)))


class SurveyTests(unittest.TestCase):
    def test_boundary_from_vertex_csv(self):
        geometry = ss.boundary_from_csv("경도,위도\n129.35,35.47\n129.36,35.47\n129.36,35.48\n".encode("utf-8"))
        self.assertEqual(geometry["type"], "Polygon")
        self.assertEqual(geometry["coordinates"][0][0], geometry["coordinates"][0][-1])
        with self.assertRaises(ValueError):
            ss.boundary_from_csv("x\n1\n".encode("utf-8"))

    def test_parcel_boundary_picks_the_parcel_holding_the_address_point(self):
        far, here = _box(400, 400, 500, 500), _box(-50, -50, 50, 50)
        payload = {"response": {"status": "OK", "result": {"featureCollection": {"features": [
            {"geometry": far, "properties": {"addr": "먼 필지"}}, {"geometry": here, "properties": {"addr": "연습로 1"}}]}}}}
        with patch.object(vworld, "api_key", return_value="key"):
            geometry, note = vworld.parcel_boundary(CENTER[1], CENTER[0], get=lambda params: payload)
            self.assertEqual(geometry, here)
            self.assertIn("연습로 1", note)
            missing, why = vworld.parcel_boundary(CENTER[1], CENTER[0], get=lambda params: {"response": {"status": "NOT_FOUND"}})
            self.assertIsNone(missing)
            self.assertIn("직접 그려", why)
        with patch.object(vworld, "api_key", return_value=""):
            self.assertIsNone(vworld.parcel_boundary(CENTER[1], CENTER[0])[0])

    def test_map_candidates_become_a_survey_sheet_that_reads_back_into_form8(self):
        p, _ = walk("울산광역시 남구 연습로 1")
        with patch.object(vworld, "api_key", return_value="key"), patch.object(vworld, "geocode", return_value=CENTER), \
                patch.object(vworld, "search", return_value=(MAP_OBJECTS, [])):
            survey.collect(p)
        ss.save_drawn_boundary(p, [{"type": "Feature", "geometry": _box(-75, -50, 75, 50), "properties": {}}], "")
        self.assertIn("화면 지도에서 그린 경계", survey.boundary_record(p)["source"])
        names = {"연습 초등학교": ("갑종", "교육·연구시설"), "연습 아파트": ("갑종", "주택"), "연습천": ("환경수용체", "하천")}
        client = FakeClient({"items": [{"name": "연습 아파트", "category": "갑종", "subtype": "주택"},
                                       {"name": "연습천", "category": "환경수용체", "subtype": "없는 분류"}]})
        rows = {r["보호대상 명칭"]: r for r in ss.candidate_rows(p, client)}
        self.assertEqual(set(rows), set(names))
        self.assertEqual((rows["연습 초등학교"]["보호대상 구분"], rows["연습 아파트"]["보호대상 구분"]), ("갑종", "갑종"))
        self.assertEqual(rows["연습천"]["보호대상 구분"], "")            # 법정 분류 목록에 없는 AI 답은 버림
        self.assertTrue(all(0 < r["사업장 경계와 거리(m)"] <= 500 for r in rows.values()))
        book = load_workbook(BytesIO(ss.workbook(p, client)))
        sheet, method = book[ss.TARGET_SHEET], book[ss.METHOD_SHEET]
        for row in range(2, sheet.max_row + 1):                        # 사람이 조사표를 고침
            name = sheet.cell(row=row, column=1).value
            sheet.cell(row=row, column=2, value=names[name][0])
            sheet.cell(row=row, column=3, value=names[name][1])
        for row in range(2, method.max_row + 1):
            if method.cell(row=row, column=1).value != "500m 안 보호대상 없음":
                method.cell(row=row, column=2, value=method.cell(row=row, column=2).value or "예")
        out = BytesIO()
        book.save(out)
        _upload(p, ss.FILENAME, out.getvalue())
        self.assertEqual({r["보호대상 명칭"] for r in f8.saved_rows(p)}, set(names))
        self.assertEqual(f8.needs(p), [])


class BoundaryPickingTests(unittest.TestCase):
    """주소를 찾으면 필지 경계가 표시되고, 필지를 눌러 더하거나 빼고, ‘이 경계로 확인’으로 이어서 계산합니다."""

    def _feature(self, parcels):
        def feature(lat, lon):
            for parcel in parcels:
                from shapely.geometry import Point, shape
                if shape(parcel["geometry"]).covers(Point(lon, lat)):
                    return parcel, "브이월드 연속지적도 필지"
            return None, "이 위치에서 필지 경계를 찾지 못했습니다. 지도에 직접 그려 주세요."
        return feature

    def test_locate_tries_vworld_then_kakao(self):
        class Provider:
            def __init__(self, key, point):
                self.key, self.point = key, point

            def api_key(self):
                return self.key

            def geocode(self, address):
                return self.point

        point, where = ss.locate("연습로 1", geocoders=(("브이월드", Provider("k", None)), ("카카오", Provider("k", CENTER))))
        self.assertEqual((point, where), (CENTER, "카카오"))
        self.assertIsNone(ss.locate("연습로 1", geocoders=(("브이월드", Provider("", CENTER)),))[0])
        self.assertIn("찾지 못했습니다", ss.locate("연습로 1", geocoders=(("브이월드", Provider("k", None)),))[1])

    def test_click_adds_and_removes_parcels_and_neighbours_merge(self):
        west = {"id": "A", "label": "연습동 1", "geometry": _box(-80, -50, 0, 50)}
        east = {"id": "B", "label": "연습동 2", "geometry": _box(0, -50, 80, 50)}
        feature = self._feature([west, east])
        lon, lat = CENTER
        parcels, note = ss.toggle_parcel([], lat, lon - 0.0003, feature=feature)
        self.assertEqual([p["id"] for p in parcels], ["A"])
        parcels, _ = ss.toggle_parcel(parcels, lat, lon + 0.0003, feature=feature)
        self.assertEqual([p["id"] for p in parcels], ["A", "B"])
        self.assertEqual(ss.parcels_geometry(parcels)["type"], "Polygon")          # 맞닿은 두 필지는 한 면
        parcels, note = ss.toggle_parcel(parcels, lat, lon - 0.0003, feature=feature)
        self.assertEqual([p["id"] for p in parcels], ["B"])
        self.assertIn("뺐습니다", note)
        same, note = ss.toggle_parcel(parcels, lat + 0.01, lon, feature=feature)
        self.assertEqual(same, parcels)
        self.assertIn("직접 그려", note)

    def test_confirm_saves_the_boundary_then_surveys_and_recalculates(self):
        p, _ = walk("울산광역시 남구 연습로 1")
        with patch.object(vworld, "api_key", return_value="key"), patch.object(vworld, "geocode", return_value=CENTER), \
                patch.object(vworld, "search", return_value=(MAP_OBJECTS, [])):
            summary = ss.confirm(p, _box(-75, -50, 75, 50), "브이월드 연속지적도 필지를 고른 경계")
        self.assertIsNotNone(survey.boundary(p))
        self.assertIn("브이월드 연속지적도 필지를 고른 경계", survey.boundary_record(p)["source"])
        self.assertEqual(summary["candidates"], 3)
        self.assertTrue(sc.saved_scenarios(p))
        self.assertFalse(sp.needs_refresh(p))

    def test_confirm_without_map_keys_still_saves_the_boundary(self):
        p, _ = walk()
        with patch.object(vworld, "api_key", return_value=""), patch("engine.stage2.cap_site_lookup.api_key", return_value=""):
            summary = ss.confirm(p, _box(-75, -50, 75, 50), "")
        self.assertIsNotNone(survey.boundary(p))
        self.assertEqual(summary["candidates"], 0)
        self.assertIn("주변 지도 조회를 하지 못했습니다", summary["notices"][0])

    def test_tiles_need_a_key(self):
        self.assertEqual(vworld.tile_layers(""), {})
        self.assertIn("/KEY/Satellite/", vworld.tile_layers("KEY")["satellite"])


class PageTests(unittest.TestCase):
    def _run(self, project):
        rows = [{"project_id": project.project_id, "company_name": project.company_name}]
        with patch("engine.stage2.storage.list_projects", return_value=rows), \
                patch("engine.stage2.storage.load_project", return_value=project), \
                patch("engine.stage2.storage.save_project"):
            at = AppTest.from_file(str(ROOT / "ui/scenario_results_page.py"), default_timeout=120)
            at.session_state["_stage2_active_project_id"] = project.project_id
            at.run()
            return at

    def test_page_shows_results_without_input_widgets(self):
        p, _ = walk()
        at = self._run(p)
        self.assertFalse(at.exception, at.exception)
        self.assertEqual([t.label for t in at.text_input], ["주소로 찾기"])   # 지도를 옮길 주소만. 시나리오·숫자는 적지 않음
        self.assertEqual(len(at.number_input), 0)
        texts = " ".join(m.value for m in at.markdown)
        self.assertIn("피해반경", texts)
        self.assertTrue(sc.saved_scenarios(p))               # 처음 열 때 파일로 계산


    def test_opening_finds_the_address_and_shows_its_parcel(self):
        p, _ = walk("울산광역시 남구 연습로 1")
        survey.clear_boundary(p)                          # 경계를 아직 정하지 않은 사업장
        parcel = {"id": "P1", "label": "연습동 1-1", "geometry": _box(-60, -40, 60, 40)}
        with patch.object(vworld, "api_key", return_value="key"), patch.object(vworld, "geocode", return_value=CENTER), \
                patch.object(vworld, "parcel_feature", return_value=(parcel, "브이월드 연속지적도 필지 연습동 1-1")), \
                patch("engine.stage2.cap_site_lookup.api_key", return_value=""):
            at = self._run(p)
        self.assertFalse(at.exception, at.exception)
        self.assertTrue(any("고른 필지: 연습동 1-1" in m.value for m in at.markdown))
        self.assertTrue(any("브이월드에서 주소를 찾았습니다" in m.value for m in at.success))
        self.assertTrue(any(b.label == "이 경계로 확인" and not b.disabled for b in at.button))


if __name__ == "__main__":
    unittest.main()
