from io import BytesIO

from docx import Document
from openpyxl import load_workbook

from engine.stage2 import demo_company_workbook as demo, cap_upload_template as template
from engine.stage2 import company_workbook_intake as intake, cap_fact_drafting, scenario_pipeline, report_export
from engine.stage2.project import Stage2Project


def test_filled_workbook_uses_same_schema_and_keeps_examples_separate():
    empty = load_workbook(BytesIO(template.build_company()))
    filled = load_workbook(BytesIO(demo.build()))
    assert filled.sheetnames == empty.sheetnames
    for ws in filled:
        assert [c.value for c in ws[1]] == [c.value for c in empty[ws.title][1]]
    assert filled.properties.identifier == demo.IDENTIFIER
    assert '가상' in filled.properties.description
    assert not any(ws._images for ws in filled)
    assert [r[2].value for r in filled[template.SITE_SHEET].iter_rows(min_row=2)] == [
        r[2].value for r in empty[template.SITE_SHEET].iter_rows(min_row=2)]
    assert len([r for r in filled['설비'].iter_rows(min_row=2) if r[0].value]) == 4
    _, issues = intake.inspect(demo.build(), demo.FILENAME, '1군')
    assert issues == []


def test_uploaded_demo_generates_narratives_scenarios_and_docx_without_forging_approval(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    selected = Stage2Project(project_id='EXCEL-START', company_name='')
    project = intake.project_for_upload(selected, demo.FILENAME, demo.build())
    assert project.company_name == demo.COMPANY
    assert project.cap_group == '1군'
    assert not project.stage1_snapshot['legal_applicability_confirmed']
    intake.apply(project, demo.FILENAME, demo.build())
    assert intake.corrections(project) == []
    assert project.get_field('cap.business.representative').value == '김가상'
    assert project.get_field('cap.workspace.dike_rule').value['비율(%)'] == '110'
    assert len(cap_fact_drafting.generate(project)) == 14
    scenarios = scenario_pipeline.refresh(project)
    assert len(scenarios.plan.scenarios) >= 4  # ③에서 경계 확인 후 상세 계산을 진행한다.
    assert any('경계까지 거리' in issue.what for issue in scenarios.problems)
    draft = report_export.build_cap_workspace_draft(project)
    doc = Document(BytesIO(draft.docx))
    text = '\n'.join([p.text for p in doc.paragraphs] + [c.text for t in doc.tables for r in t.rows for c in r.cells])
    assert demo.COMPANY in text
    assert '김가상' in text
    assert 'TK-01' in text
    assert 'AI 초안' in text
    assert not draft.state.final_ready


def test_reuploading_the_workbook_keeps_map_survey_results(tmp_path, monkeypatch):
    from openpyxl import load_workbook
    from engine.stage2 import evidence_workspace as ew
    monkeypatch.chdir(tmp_path)
    project = intake.project_for_upload(Stage2Project(project_id='EXCEL-START', company_name=''), demo.FILENAME, demo.build())
    intake.apply(project, demo.FILENAME, demo.build())
    project.set_field(ew.SURVEY_KEY, '보호대상', [{'보호대상 명칭': '가상학교'}], 'HOLD')
    rows = project.get_field(ew.SCENARIO_INPUT_KEY).value
    next(r for r in rows if r['설비번호'] == 'TK-01').update({'경계까지 거리(m)': 120, '사고원점 좌표': '35,127'})
    project.set_field(ew.SCENARIO_INPUT_KEY, '시나리오 입력', rows, 'HOLD')
    book = load_workbook(BytesIO(demo.build()))
    book['설비']['B2'] = '가상 톨루엔 저장탱크(수정)'
    out = BytesIO()
    book.save(out)
    assert intake.apply(project, demo.FILENAME, out.getvalue())
    assert [r['보호대상 명칭'] for r in project.get_field(ew.SURVEY_KEY).value] == ['가상학교']
    tk = next(r for r in project.get_field(ew.SCENARIO_INPUT_KEY).value if r['설비번호'] == 'TK-01')
    assert tk['경계까지 거리(m)'] == 120 and tk['사고원점 좌표'] == '35,127'
    assert str(tk['액위(m)']) == '6'       # 회사 엑셀의 시나리오 입력값은 그대로 읽힌다


def test_workbook_emergency_facts_fill_the_form16_summary(tmp_path, monkeypatch):
    from engine.stage2.cap_form16_engine import build_cap_form16_data
    monkeypatch.chdir(tmp_path)
    project = intake.project_for_upload(Stage2Project(project_id='EXCEL-START', company_name=''), demo.FILENAME, demo.build())
    intake.apply(project, demo.FILENAME, demo.build())
    data = build_cap_form16_data(project)
    assert not [b for b in data.company_blockers if '회사 확정자료가 없습니다' in b]
    values = dict(data.internal_summary + data.external_summary)
    assert '박가상' in values['가동중지 권한·절차']
    assert '가상 집결지 A' in values['주민 경보·대피']
