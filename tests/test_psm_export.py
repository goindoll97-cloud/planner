from __future__ import annotations

from io import BytesIO
import unittest
from unittest.mock import patch
import zipfile

from docx import Document

from engine.stage2 import psm_export as export
from tests.test_psm_workspace_page import _project, _run


class ExportTests(unittest.TestCase):
    def test_evaluate_splits_issues(self):
        project = _project()
        state = export.evaluate(project)
        self.assertGreater(state.hold + state.review, 0)  # 비어 있는 별지가 많다
        self.assertFalse(state.can_confirm)

    def test_outputs_are_always_drafts_without_any_approval(self):
        project = _project()
        outputs = export.build_outputs(project)
        self.assertIn("초안", outputs.regulation_name)
        self.assertNotIn("작성본", outputs.regulation_name)
        Document(BytesIO(outputs.regulation_docx))
        review = Document(BytesIO(outputs.review_docx))
        self.assertIn("프로그램 보완 안내", " ".join(p.text for p in review.paragraphs))  # 문서 끝에 보완 안내
        with zipfile.ZipFile(BytesIO(outputs.bundle)) as archive:
            self.assertTrue(any(n.endswith(".docx") for n in archive.namelist()))
        self.assertIn(b"REVIEW_ONLY", outputs.provenance_json)

    def test_export_requires_psm_in_scope(self):
        project = _project()
        project.psm_selected = False
        with self.assertRaises(ValueError):
            export.evaluate(project)


class ScreenTests(unittest.TestCase):
    def test_screen_offers_downloads_immediately_as_drafts(self):
        project = _project()
        at = _run(project, "export")
        self.assertFalse(at.exception)
        self.assertFalse(any(b.key in ("export_psm_ack", "export_psm_confirm") for b in at.button))
        self.assertTrue(any("초안입니다" in i.value for i in at.info))


if __name__ == "__main__":
    unittest.main()


class CapExportTests(unittest.TestCase):
    def _cap_project(self):
        from tests.test_stage2_cap_form1_engine import CAPForm1EngineTests

        project = CAPForm1EngineTests()._project()
        project.psm_required = False
        return project

    def test_cap_download_is_labelled_as_draft(self):
        from engine.stage2 import report_export as shared

        project = self._cap_project()
        self.assertTrue(project.cap_in_scope)
        state = shared.evaluate(project, "CAP")
        self.assertIsNotNone(state.readiness.cap_gate)  # 제출 전 체크포인트도 함께 판정
        outputs = shared.build_outputs(project, system="CAP")
        self.assertIn("규정서식_초안", outputs.regulation_name)
        self.assertNotIn("작성본", outputs.regulation_name)
        self.assertIn(b"REVIEW_ONLY", outputs.provenance_json)
        Document(BytesIO(outputs.regulation_docx))

    def test_scope_is_checked_per_system(self):
        from engine.stage2 import report_export as shared

        project = self._cap_project()
        with self.assertRaises(ValueError):
            shared.evaluate(project, "PSM")
        with self.assertRaises(ValueError):
            shared.evaluate(project, "XYZ")


class CapScreenTests(unittest.TestCase):
    def test_cap_workspace_export_step_uses_the_shared_panel(self):
        from pathlib import Path
        source = (Path(__file__).resolve().parents[1] / "ui/evidence_workspace_page.py").read_text(encoding="utf-8")
        self.assertIn('render_workspace_download(project)', source)
        self.assertNotIn("cap_form01_download_", source)  # 항상 '작성본'으로 이름 붙이던 예전 다운로드 제거


class CapPanelRenderTests(unittest.TestCase):
    def test_panel_renders_for_a_cap_project_with_immediate_downloads(self):
        from streamlit.testing.v1 import AppTest

        code = (
            "import sys; sys.path.insert(0, %r)\n"
            "from unittest.mock import patch\n"
            "from tests.test_stage2_cap_form1_engine import CAPForm1EngineTests\n"
            "from ui import report_export_panel\n"
            "project = CAPForm1EngineTests()._project()\n"
            "report_export_panel.render(project, 'CAP')\n"
        ) % str(__import__("pathlib").Path(__file__).resolve().parents[1])
        at = AppTest.from_string(code, default_timeout=90).run()
        self.assertFalse(at.exception)
        self.assertFalse(any(b.key == "export_cap_ack" for b in at.button))   # 승인·확인 단추가 없다
        self.assertTrue(any("초안입니다" in i.value for i in at.caption))
        self.assertEqual([d.proto.label for d in at.get("download_button")], ["규정서식 초안 DOCX 내려받기"])
