"""엑셀 업로드 양식: 빈 사실의 열 이름으로 만들고, 채워 올리면 모두 읽히는지(왕복) 확인합니다."""
from __future__ import annotations

import unittest
from io import BytesIO

from openpyxl import load_workbook

from engine.stage2 import cap_fact_catalog as catalog
from engine.stage2 import cap_fact_status as fs
from engine.stage2 import cap_upload_template as tpl
from engine.stage2 import evidence_workspace as ew
from engine.stage2.direct_template import create_direct_cap_project

VALID_CAS = "108-88-3"


def _sample(item: catalog.Fact, row: int) -> str:
    if item.options:
        return item.options[0]
    if item.evidence_field == "cas":
        return VALID_CAS
    if item.format == "number":
        return str(row + 1)
    return f"가상 {item.name} {row}"


def _fill(content: bytes, new_ids: dict[str, list[str]]) -> bytes:
    """모든 빈칸을 형식에 맞는 값으로 채우고, 대상 시트에는 새 번호 줄을 더합니다."""
    book = load_workbook(BytesIO(content))
    by_header = {}
    for item in catalog.facts(company=True):
        for header in item.headers[:1]:
            by_header[(item.kind if item.kind not in catalog.SITE_KINDS else "site", header)] = item
    if tpl.SITE_SHEET in book.sheetnames:
        sheet = book[tpl.SITE_SHEET]
        for r in range(2, sheet.max_row + 1):
            item = by_header[("site", sheet.cell(row=r, column=1).value)]
            sheet.cell(row=r, column=2, value=_sample(item, r))
    for kind, title in tpl.ENTITY_SHEETS.items():
        if title not in book.sheetnames:
            continue
        sheet = book[title]
        headers = [c.value for c in sheet[1] if c.value != "예시 (입력하지 마세요)"]
        for code in new_ids.get(kind, []):
            sheet.append([code] + [""] * (len(headers) - 1))
        for r in range(2, sheet.max_row + 1):
            if not sheet.cell(row=r, column=1).value:
                continue
            for c, header in enumerate(headers[1:], start=2):
                if sheet.cell(row=r, column=c).value in (None, ""):
                    lookup = "facility" if kind == "scenario" else kind      # 사고시나리오 검토 시트는 설비 사실
                    sheet.cell(row=r, column=c, value=_sample(by_header[(lookup, header)], r))
    for table, items in catalog.input_tables().items():
        title = tpl.sheet_title(table)
        if title in book.sheetnames:
            for c, item in enumerate(items, 1):
                book[title].cell(2, c, _sample(item, c))
    out = BytesIO()
    book.save(out)
    return out.getvalue()


class TemplateTests(unittest.TestCase):
    def test_empty_project_template_has_every_readable_company_fact(self):
        p = create_direct_cap_project("가상", address="가상 주소", cap_group="")
        wanted = tpl.columns(p)
        site_headers = {i.template_header for i in wanted["site"]}
        self.assertIn("대표자", site_headers)
        self.assertIn("공정개요", site_headers)
        self.assertNotIn("사업장명", site_headers)           # 사업장을 만들 때 이미 입력
        self.assertNotIn("우편번호/주소", site_headers)        # 직접 입력한 주소가 있음
        book = load_workbook(BytesIO(tpl.build(p)))
        self.assertNotIn(tpl.GUIDE_SHEET, book.sheetnames)
        self.assertEqual(book[tpl.SITE_SHEET]["C1"].value, "예시 (입력하지 마세요)")
        self.assertEqual([c.value for c in book["설비"][1]][:3], ["설비번호", "설비명", "단위공장·공정"])
        self.assertEqual(book["가스감지기"]["A1"].value, "감지기 번호")

    def test_round_trip_fills_every_template_column(self):
        p = create_direct_cap_project("가상", address="가상 주소", cap_group="")
        content = _fill(tpl.build(p), {"facility": ["TK-01", "TK-02"], "material": ["P-01"], "detector": ["GD-01"],
                                       "scenario": ["TK-01", "TK-02"]})
        self.assertEqual(ew.classify_document(tpl.FILENAME, content), "mixed")
        ew.ingest(p, tpl.FILENAME, content, "mixed")
        doc = ew.documents(p)[0]
        self.assertEqual(doc["unmapped_columns"], [])
        self.assertEqual(doc["rows_without_id"], 0)
        rows = {r["_id"]: r for r in fs.fact_rows(p)}
        wanted = tpl.columns(create_direct_cap_project("빈", address="주소", cap_group=""))
        for kind, items in wanted.items():
            if kind == "tables":
                items = [item for table in items.values() for item in table]
            for item in items:
                self.assertEqual(rows[item.id]["_state"], fs.O, f"{item.id}: {rows[item.id]['고칠 점']}")
        self.assertEqual(sorted(e for e in fs.entities(p, "facility")), ["facility:TK-01", "facility:TK-02"])

    def test_unfilled_template_adds_nothing_and_reports_nothing_unknown(self):
        p = create_direct_cap_project("가상", address="가상 주소", cap_group="")
        content = tpl.build(p)
        ew.ingest(p, tpl.FILENAME, content, ew.classify_document(tpl.FILENAME, content))
        doc = ew.documents(p)[0]
        self.assertEqual(doc["fact_count"], 0)
        self.assertEqual(doc["unmapped_columns"], [])
        self.assertEqual([s for s in ew.slots(p)], [])

    def test_only_missing_facts_are_asked_after_some_files_were_read(self):
        p = create_direct_cap_project("가상", address="가상 주소", cap_group="")
        ew.ingest(p, "eq.csv", "설비번호,설비명,용량(m3)\nTK-01,탱크,10\n".encode("utf-8"), "facility")
        wanted = tpl.columns(p)
        facility_headers = [i.template_header for i in wanted["facility"]]
        self.assertNotIn("용량(m3)", facility_headers)
        self.assertIn("운전압력(MPa)", facility_headers)
        book = load_workbook(BytesIO(tpl.build(p)))
        sheet = book["설비"]
        self.assertEqual([sheet["A2"].value, sheet["B2"].value], ["TK-01", "탱크"])   # 아는 설비는 미리 적어 둠
        # 채운 값이 원래 파일과 같은 설비에 붙습니다
        content = _fill(tpl.build(p), {})
        ew.ingest(p, tpl.FILENAME, content, ew.classify_document(tpl.FILENAME, content))
        self.assertEqual(fs.entities(p, "facility"), ["facility:TK-01"])
        row = next(r for r in fs.fact_rows(p) if r["_id"] == "facility.pressure")
        self.assertEqual(row["_state"], fs.O)

    def test_examples_are_never_company_facts_and_same_filename_replaces(self):
        p = create_direct_cap_project("가상", address="주소", cap_group="")
        book = load_workbook(BytesIO(tpl.build(p, docs=("CAP",))))
        sheet = book[tpl.SITE_SHEET]
        row = next(r for r in range(2, sheet.max_row + 1) if sheet.cell(r, 1).value == "대표자")
        self.assertTrue(sheet.cell(row, 3).value)
        def upload(value):
            sheet.cell(row, 2, value)
            out = BytesIO()
            book.save(out)
            content = out.getvalue()
            ew.replace_same_name(p, tpl.FILENAME, content)
            ew.ingest(p, tpl.FILENAME, content, ew.classify_document(tpl.FILENAME, content))
        upload("김대표")
        upload("이대표")
        self.assertEqual(len(ew.documents(p)), 1)
        facts = ew.graph_facts(p)
        self.assertTrue(any(f["value"] == "이대표" for f in facts))
        self.assertFalse(any(f["value"] in ("김대표", "홍길동") for f in facts))

    def test_plan_template_requests_facts_and_never_complete_narratives(self):
        p = create_direct_cap_project("가상", address="주소", cap_group="1군")
        book = load_workbook(BytesIO(tpl.build(p, docs=("CAP",))))
        headers = {row[0].value for row in book[tpl.SITE_SHEET].iter_rows(min_row=2)}
        self.assertTrue({"지역사회 소통 대상", "협력 기관", "대피 장소"} <= headers)
        self.assertFalse({"지역사회 소통", "지역비상대응기관·인근사업장 공조", "주민 경보·대피"} & headers)

    def test_full_template_lists_all_columns_even_when_filled(self):
        p = create_direct_cap_project("가상", address="가상 주소", cap_group="")
        ew.ingest(p, "eq.csv", "설비번호,설비명,용량(m3)\nTK-01,탱크,10\n".encode("utf-8"), "facility")
        everything = tpl.columns(p, only_blank=False)
        self.assertIn("용량(m3)", [i.template_header for i in everything["facility"]])
        self.assertGreater(tpl.blank_count(p), 0)


if __name__ == "__main__":
    unittest.main()
