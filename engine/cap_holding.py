from __future__ import annotations

"""Second-stage CAP maximum-holding workflow based on approved Appendix 4.

The company-facing first workbook intentionally stays small.  This module is
used only after a chemical has a confirmed Appendix 2/3 direct-rule hit.  It
creates a focused facility workbook and calculates maximum holding from the
facility facts supplied by the user.

Fail-closed rules:
- Appendix 4 must be present in the approved regulatory DB.
- tanker/transport, outside-pipeline and formally stopped-facility exclusions
  are accepted only when the user explicitly selects them.
- gas/high-pressure and multi-phase equipment are not inferred; a verified
  direct mass with its evidence is required until dedicated rules are added.
- mixture concentration is used to test whether a legal concentration
  threshold applies.  It is NOT multiplied into the material mass here.
"""

from dataclasses import dataclass, field
from io import BytesIO
from pathlib import Path
from typing import Any, Iterable

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from .inventory import IntakeData


PROJECT_ROOT = Path(__file__).resolve().parents[1]
APPROVED_APP4_DB = PROJECT_ROOT / "data" / "regulatory" / "approved" / "cap_qty_app4.csv"
EXPECTED_APP4_RULE_IDS = {"1", "2-가", "2-나", "3-가", "3-나", "3-다", "비고"}
FACILITY_SHEET = "01_CAP시설정보"

TITLE_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FILL = PatternFill("solid", fgColor="5B9BD5")
INPUT_FILL = PatternFill("solid", fgColor="FFF2CC")
NOTE_FILL = PatternFill("solid", fgColor="F3F6F9")
WHITE_FONT = Font(color="FFFFFF", bold=True)

FACILITY_COLUMNS = [
    "목록행번호",
    "제품명",
    "CAS No.",
    "시설명",
    "시설유형",
    "제외시설여부",
    "제외사유",
    "물질성상",
    "공정유형",
    "별표4 기준함량(%)",
    "함량근거",
    "설계용량",
    "용량단위",
    "비중 또는 밀도(kg/L=ton/m3)",
    "보관계획도 최대량",
    "일일최대보관량",
    "질량단위",
    "직접확인 최대보유량",
    "직접확인 근거",
    "복수성상 증빙",
    "비고",
]

RECOGNIZED_EXCLUSIONS = {
    "탱크로리·운송차량",
    "사외배관",
    "취급중단 신고시설",
}


@dataclass
class FacilityCalculation:
    inventory_row_no: int
    product_name: str
    cas: str
    facility_name: str
    facility_type: str
    content_pct: float | None
    max_holding_ton: float | None
    excluded: bool
    calculation_basis: str
    blocker: str = ""


@dataclass
class CAPHoldingResult:
    status: str
    label: str
    app4_ready: bool
    facility_rows: list[FacilityCalculation] = field(default_factory=list)
    comparison_rows: list[dict[str, Any]] = field(default_factory=list)
    blockers: list[str] = field(default_factory=list)
    messages: list[str] = field(default_factory=list)


def _clean(value: Any) -> str:
    if value is None:
        return ""
    try:
        if pd.isna(value):
            return ""
    except (TypeError, ValueError):
        pass
    return str(value).strip()


def _num(value: Any) -> float | None:
    text = _clean(value).replace(",", "")
    if not text:
        return None
    try:
        return float(text)
    except ValueError:
        return None


def _int(value: Any) -> int | None:
    number = _num(value)
    if number is None or int(number) != number:
        return None
    return int(number)


def _mass_ton(value: Any, unit: Any) -> float | None:
    amount = _num(value)
    if amount is None or amount < 0:
        return None
    norm = _clean(unit).lower().replace(" ", "")
    if norm == "kg":
        return amount / 1000.0
    if norm in {"ton", "t", "톤"}:
        return amount
    return None


def _volume_mass_ton(volume: Any, unit: Any, density: Any) -> float | None:
    amount = _num(volume)
    rho = _num(density)
    if amount is None or rho is None or amount < 0 or rho <= 0:
        return None
    norm = _clean(unit).lower().replace(" ", "")
    # 1 kg/L == 1 ton/m3
    if norm in {"l", "liter", "litre", "리터"}:
        return amount * rho / 1000.0
    if norm in {"m3", "㎥", "m³"}:
        return amount * rho
    return None


def app4_db_ready() -> bool:
    if not APPROVED_APP4_DB.exists():
        return False
    try:
        frame = pd.read_csv(APPROVED_APP4_DB, dtype=str, keep_default_na=False)
    except Exception:
        return False
    if "rule_id" not in frame.columns:
        return False
    return EXPECTED_APP4_RULE_IDS.issubset(set(frame["rule_id"].astype(str)))


def _candidate_rows(intake: IntakeData, row_numbers: Iterable[int]) -> list[tuple[int, pd.Series]]:
    rows: list[tuple[int, pd.Series]] = []
    for row_no in sorted(set(int(v) for v in row_numbers if int(v) > 0)):
        idx = row_no - 1
        if 0 <= idx < len(intake.chemicals):
            rows.append((row_no, intake.chemicals.iloc[idx]))
    return rows


def build_cap_facility_workbook(intake: IntakeData, row_numbers: Iterable[int]) -> bytes:
    """Build a focused second-stage workbook only for CAP direct-match rows."""
    wb = Workbook()
    guide = wb.active
    guide.title = "00_작성가이드"
    guide.merge_cells("A1:H1")
    guide["A1"] = "화사계 별표 4 최대보유량 시설정보 — 2차 입력"
    guide["A1"].fill = TITLE_FILL
    guide["A1"].font = Font(color="FFFFFF", bold=True, size=14)
    guide["A3"] = "작성 원칙"
    guide["A3"].font = Font(bold=True)
    guide_rows = [
        ["1", "한 시설당 한 줄", "같은 물질이 반응기·탱크·창고에 있으면 시설별로 줄을 나눠 입력합니다."],
        ["2", "제조·사용시설/저장탱크", "설계용량과 비중(또는 밀도)을 입력합니다. 부피단위는 L 또는 m3를 사용합니다."],
        ["3", "보관시설", "보관계획도 기준 최대량과 일일최대보관량을 모두 입력합니다. 프로그램은 둘 중 큰 값을 사용합니다."],
        ["4", "혼합물", "규제 함량 이상인지 판단할 기준함량을 입력합니다. 최대보유량 계산에서 함량%를 다시 곱해 순물질량으로 줄이지 않습니다."],
        ["5", "단순혼합/반응", "단순혼합은 투입완료 후, 반응공정은 반응 전의 법정 기준함량을 확인해 입력합니다."],
        ["6", "기체·고압가스/복수성상", "자동 추정하지 않습니다. 검증된 직접 최대보유량과 근거가 있으면 입력하고, 없으면 판정보류됩니다."],
        ["7", "제외시설", "탱크로리·운송차량, 사외배관, 취급중단 신고시설만 명확한 근거가 있을 때 Y로 표시합니다."],
        ["8", "직접확인 최대보유량", "공식 산정표 등으로 이미 확인된 질량을 사용할 때만 값과 근거를 함께 입력합니다."],
    ]
    for r, values in enumerate(guide_rows, 5):
        for c, value in enumerate(values, 1):
            guide.cell(r, c, value)
            guide.cell(r, c).alignment = Alignment(vertical="top", wrap_text=True)
    guide.column_dimensions["A"].width = 8
    guide.column_dimensions["B"].width = 26
    guide.column_dimensions["C"].width = 100

    ws = wb.create_sheet(FACILITY_SHEET)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(FACILITY_COLUMNS))
    ws.cell(1, 1, "화사계 별표 4 시설정보 — 노란색 칸만 입력, 시설이 여러 개면 행을 복사해 추가")
    ws.cell(1, 1).fill = TITLE_FILL
    ws.cell(1, 1).font = WHITE_FONT
    for col, header in enumerate(FACILITY_COLUMNS, 1):
        cell = ws.cell(3, col, header)
        cell.fill = HEADER_FILL
        cell.font = WHITE_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    candidates = _candidate_rows(intake, row_numbers)
    start = 4
    for offset, (row_no, item) in enumerate(candidates):
        r = start + offset
        ws.cell(r, 1, row_no)
        ws.cell(r, 2, _clean(item.get("제품명")))
        ws.cell(r, 3, _clean(item.get("CAS No.")))
        pct = _num(item.get("함량(%)"))
        if pct is not None:
            ws.cell(r, 10, pct)
            ws.cell(r, 11, "1차 화학물질 목록 함량 — 공정 중 변하면 수정")
        for c in range(4, len(FACILITY_COLUMNS) + 1):
            ws.cell(r, c).fill = INPUT_FILL
        for c in (1, 2, 3):
            ws.cell(r, c).fill = NOTE_FILL

    # Reserve extra rows so users can duplicate one material across facilities.
    for r in range(start + len(candidates), start + len(candidates) + 60):
        for c in range(1, len(FACILITY_COLUMNS) + 1):
            ws.cell(r, c).fill = INPUT_FILL

    dv_type = DataValidation(type="list", formula1='"제조·사용시설,저장탱크,보관시설,기타"', allow_blank=True)
    dv_yn = DataValidation(type="list", formula1='"N,Y,모름"', allow_blank=True)
    dv_ex = DataValidation(type="list", formula1='"해당없음,탱크로리·운송차량,사외배관,취급중단 신고시설,기타"', allow_blank=True)
    dv_state = DataValidation(type="list", formula1='"액체,고체,기체·고압가스,복수성상,모름"', allow_blank=True)
    dv_process = DataValidation(type="list", formula1='"변화없음,단순혼합,반응,해당없음,모름"', allow_blank=True)
    dv_vol = DataValidation(type="list", formula1='"L,m3"', allow_blank=True)
    dv_mass = DataValidation(type="list", formula1='"kg,ton"', allow_blank=True)
    dv_evidence = DataValidation(type="list", formula1='"Y,N,모름"', allow_blank=True)
    for dv in (dv_type, dv_yn, dv_ex, dv_state, dv_process, dv_vol, dv_mass, dv_evidence):
        ws.add_data_validation(dv)
    last = start + len(candidates) + 59
    dv_type.add(f"E4:E{last}")
    dv_yn.add(f"F4:F{last}")
    dv_ex.add(f"G4:G{last}")
    dv_state.add(f"H4:H{last}")
    dv_process.add(f"I4:I{last}")
    dv_vol.add(f"M4:M{last}")
    dv_mass.add(f"Q4:Q{last}")
    dv_evidence.add(f"T4:T{last}")

    widths = [10, 24, 16, 24, 18, 14, 24, 18, 16, 18, 34, 14, 12, 24, 20, 20, 12, 22, 40, 16, 38]
    for i, width in enumerate(widths, 1):
        ws.column_dimensions[chr(64 + i) if i <= 26 else "A"].width = width
    ws.freeze_panes = "D4"

    out = BytesIO()
    wb.save(out)
    return out.getvalue()


def read_cap_facility_workbook(file_bytes: bytes) -> pd.DataFrame:
    xls = pd.ExcelFile(BytesIO(file_bytes), engine="openpyxl")
    if FACILITY_SHEET not in xls.sheet_names:
        raise ValueError(f"'{FACILITY_SHEET}' 시트가 없습니다.")
    frame = pd.read_excel(xls, sheet_name=FACILITY_SHEET, header=2)
    missing = [col for col in FACILITY_COLUMNS if col not in frame.columns]
    if missing:
        raise ValueError("시설정보 필수 열이 없습니다: " + ", ".join(missing))
    frame = frame.dropna(how="all").copy()
    value_cols = ["시설명", "시설유형", "직접확인 최대보유량", "설계용량", "보관계획도 최대량", "일일최대보관량"]
    frame = frame[frame[value_cols].notna().any(axis=1)].copy()
    frame.reset_index(drop=True, inplace=True)
    return frame


def _inventory_item(intake: IntakeData, row_no: int) -> pd.Series | None:
    idx = row_no - 1
    if 0 <= idx < len(intake.chemicals):
        return intake.chemicals.iloc[idx]
    return None


def calculate_facility_rows(intake: IntakeData, facilities: pd.DataFrame) -> tuple[list[FacilityCalculation], list[str]]:
    calculations: list[FacilityCalculation] = []
    blockers: list[str] = []

    for idx, row in facilities.iterrows():
        excel_row = idx + 4
        row_no = _int(row.get("목록행번호"))
        if row_no is None:
            blockers.append(f"시설정보 {excel_row}행: 목록행번호 확인 필요")
            continue
        source_item = _inventory_item(intake, row_no)
        if source_item is None:
            blockers.append(f"시설정보 {excel_row}행: 목록행번호 {row_no}가 1차 화학물질 목록에 없습니다.")
            continue

        product = _clean(source_item.get("제품명"))
        cas = _clean(source_item.get("CAS No."))
        facility_name = _clean(row.get("시설명"))
        facility_type = _clean(row.get("시설유형"))
        if not facility_name or facility_type not in {"제조·사용시설", "저장탱크", "보관시설", "기타"}:
            reason = f"시설정보 {excel_row}행({product or cas}): 시설명과 시설유형을 확인해 주세요."
            blockers.append(reason)
            calculations.append(FacilityCalculation(row_no, product, cas, facility_name, facility_type, None, None, False, "", reason))
            continue

        excluded_flag = _clean(row.get("제외시설여부"))
        exclusion_reason = _clean(row.get("제외사유"))
        if excluded_flag not in {"Y", "N"}:
            reason = f"시설정보 {excel_row}행({facility_name}): 제외시설 여부를 Y 또는 N으로 확정해 주세요."
            blockers.append(reason)
            calculations.append(FacilityCalculation(row_no, product, cas, facility_name, facility_type, None, None, False, "", reason))
            continue
        if excluded_flag == "Y":
            if exclusion_reason not in RECOGNIZED_EXCLUSIONS:
                reason = f"시설정보 {excel_row}행({facility_name}): 별표4 산정 제외 사유를 확인할 수 없습니다."
                blockers.append(reason)
                calculations.append(FacilityCalculation(row_no, product, cas, facility_name, facility_type, None, None, False, "", reason))
            else:
                calculations.append(
                    FacilityCalculation(row_no, product, cas, facility_name, facility_type, None, 0.0, True, f"별표4 산정 제외: {exclusion_reason}")
                )
            continue

        state = _clean(row.get("물질성상"))
        process_type = _clean(row.get("공정유형"))
        source_pct = _num(source_item.get("함량(%)"))
        content_pct = _num(row.get("별표4 기준함량(%)"))
        if content_pct is None:
            content_pct = source_pct
        if content_pct is not None and not (0 <= content_pct <= 100):
            content_pct = None

        direct_mass = _mass_ton(row.get("직접확인 최대보유량"), row.get("질량단위"))
        direct_basis = _clean(row.get("직접확인 근거"))

        if state in {"기체·고압가스", "복수성상"}:
            if direct_mass is None or not direct_basis:
                reason = (
                    f"시설정보 {excel_row}행({facility_name}): {state}은 자동 추정하지 않습니다. "
                    "검증된 직접 최대보유량과 근거를 입력해 주세요."
                )
                blockers.append(reason)
                calculations.append(FacilityCalculation(row_no, product, cas, facility_name, facility_type, content_pct, None, False, "", reason))
                continue
            calculations.append(
                FacilityCalculation(row_no, product, cas, facility_name, facility_type, content_pct, direct_mass, False, f"직접확인 질량 / 근거: {direct_basis}")
            )
            continue

        if facility_type == "제조·사용시설" and process_type in {"단순혼합", "반응"}:
            if _num(row.get("별표4 기준함량(%)")) is None or not _clean(row.get("함량근거")):
                reason = (
                    f"시설정보 {excel_row}행({facility_name}): {process_type} 공정은 별표4 기준시점의 함량과 근거가 필요합니다."
                )
                blockers.append(reason)
                calculations.append(FacilityCalculation(row_no, product, cas, facility_name, facility_type, content_pct, None, False, "", reason))
                continue

        if direct_mass is not None:
            if not direct_basis:
                reason = f"시설정보 {excel_row}행({facility_name}): 직접확인 최대보유량의 근거를 입력해 주세요."
                blockers.append(reason)
                calculations.append(FacilityCalculation(row_no, product, cas, facility_name, facility_type, content_pct, None, False, "", reason))
                continue
            mass_ton = direct_mass
            basis = f"직접확인 질량 / 근거: {direct_basis}"
        elif facility_type in {"제조·사용시설", "저장탱크"}:
            mass_ton = _volume_mass_ton(
                row.get("설계용량"), row.get("용량단위"), row.get("비중 또는 밀도(kg/L=ton/m3)")
            )
            if mass_ton is None:
                reason = f"시설정보 {excel_row}행({facility_name}): 설계용량·용량단위·비중(밀도)이 필요합니다."
                blockers.append(reason)
                calculations.append(FacilityCalculation(row_no, product, cas, facility_name, facility_type, content_pct, None, False, "", reason))
                continue
            basis = "별표4 설계용량 × 비중(밀도)"
        elif facility_type == "보관시설":
            plan_ton = _mass_ton(row.get("보관계획도 최대량"), row.get("질량단위"))
            daily_ton = _mass_ton(row.get("일일최대보관량"), row.get("질량단위"))
            if plan_ton is None or daily_ton is None:
                reason = f"시설정보 {excel_row}행({facility_name}): 보관계획도 최대량과 일일최대보관량을 모두 입력해 주세요."
                blockers.append(reason)
                calculations.append(FacilityCalculation(row_no, product, cas, facility_name, facility_type, content_pct, None, False, "", reason))
                continue
            mass_ton = max(plan_ton, daily_ton)
            basis = "별표4 보관계획도 최대량과 일일최대보관량 중 큰 값"
        else:
            reason = f"시설정보 {excel_row}행({facility_name}): 기타 시설유형은 자동 산정하지 않습니다."
            blockers.append(reason)
            calculations.append(FacilityCalculation(row_no, product, cas, facility_name, facility_type, content_pct, None, False, "", reason))
            continue

        calculations.append(
            FacilityCalculation(
                inventory_row_no=row_no,
                product_name=product,
                cas=cas,
                facility_name=facility_name,
                facility_type=facility_type,
                content_pct=content_pct,
                max_holding_ton=round(float(mass_ton), 8),
                excluded=False,
                calculation_basis=basis,
            )
        )

    return calculations, list(dict.fromkeys(blockers))


def _legal_row(hit: dict[str, Any]) -> dict[str, Any]:
    source_key = _clean(hit.get("source_key")) or "CAP_QTY_APP3"
    item_no = _clean(hit.get("item_no")) or _clean(hit.get("legal_item_no"))
    variant = _clean(hit.get("hazard_category")) or _clean(hit.get("variant_type")) or "BASE"
    return {
        "source_key": source_key,
        "row_no": _int(hit.get("row_no")),
        "cas": _clean(hit.get("cas")),
        "item_no": item_no,
        "variant": variant,
        "legal_substance": _clean(hit.get("legal_substance")),
        "content_threshold_pct": _num(hit.get("content_threshold_pct")),
        "lowest_quantity_ton": _num(hit.get("lowest_quantity_ton")),
        "lower_quantity_ton": _num(hit.get("lower_quantity_ton")),
        "upper_quantity_ton": _num(hit.get("upper_quantity_ton")),
    }


def _band(value: float, lowest: float | None, lower: float | None, upper: float | None) -> str:
    if upper is not None and value >= upper:
        return "상위 규정수량 이상"
    if lower is not None and value >= lower:
        return "하위 이상·상위 미만"
    if lowest is not None and value >= lowest:
        return "최하위 이상·하위 미만"
    return "하위 규정수량 미만"


def compare_with_legal_rules(
    calculations: list[FacilityCalculation],
    legal_hits: Iterable[dict[str, Any]],
) -> tuple[list[dict[str, Any]], list[str]]:
    """Aggregate facility mass for each same legal rule and compare quantities."""
    legal = [_legal_row(dict(hit)) for hit in legal_hits]
    legal = [row for row in legal if row["row_no"] is not None and row["cas"]]
    groups: dict[tuple[Any, ...], dict[str, Any]] = {}
    for row in legal:
        key = (
            row["source_key"], row["cas"], row["item_no"], row["variant"],
            row["content_threshold_pct"], row["lowest_quantity_ton"],
            row["lower_quantity_ton"], row["upper_quantity_ton"],
        )
        group = groups.setdefault(key, {**row, "row_nos": set()})
        group["row_nos"].add(int(row["row_no"]))

    comparison: list[dict[str, Any]] = []
    blockers: list[str] = []
    for group in groups.values():
        threshold = group["content_threshold_pct"]
        eligible: list[FacilityCalculation] = []
        for calc in calculations:
            if calc.inventory_row_no not in group["row_nos"] or calc.cas != group["cas"]:
                continue
            if calc.excluded or calc.max_holding_ton is None:
                continue
            if threshold is not None:
                if calc.content_pct is None:
                    blockers.append(
                        f"CAS {group['cas']} / {group['source_key']} {group['item_no']}: 함량기준 {threshold:g}% 적용을 위한 시설별 기준함량 미확인"
                    )
                    continue
                if calc.content_pct < threshold:
                    continue
            eligible.append(calc)

        total = round(sum(float(calc.max_holding_ton or 0.0) for calc in eligible), 8)
        comparison.append(
            {
                "source_key": group["source_key"],
                "cas": group["cas"],
                "item_no": group["item_no"],
                "variant": group["variant"],
                "legal_substance": group["legal_substance"],
                "content_threshold_pct": threshold,
                "inventory_rows": ",".join(str(v) for v in sorted(group["row_nos"])),
                "included_facilities": len(eligible),
                "calculated_max_holding_ton": total,
                "lowest_quantity_ton": group["lowest_quantity_ton"],
                "lower_quantity_ton": group["lower_quantity_ton"],
                "upper_quantity_ton": group["upper_quantity_ton"],
                "quantity_band": _band(total, group["lowest_quantity_ton"], group["lower_quantity_ton"], group["upper_quantity_ton"]),
                "basis": "별표4 시설별 최대보유량 합산",
            }
        )
    return comparison, list(dict.fromkeys(blockers))


def assess_cap_holding(
    intake: IntakeData,
    facilities: pd.DataFrame,
    legal_hits: Iterable[dict[str, Any]],
    required_row_numbers: Iterable[int],
) -> CAPHoldingResult:
    if not app4_db_ready():
        return CAPHoldingResult(
            status="DB_NOT_READY",
            label="별표 4 승인 DB 필요",
            app4_ready=False,
            blockers=["화사계 별표 4 최대보유량 산정방법 승인 DB가 필요합니다."],
        )

    calculations, blockers = calculate_facility_rows(intake, facilities)
    required = sorted(set(int(v) for v in required_row_numbers if int(v) > 0))
    represented = {row.inventory_row_no for row in calculations}
    for row_no in required:
        if row_no not in represented:
            item = _inventory_item(intake, row_no)
            label = _clean(item.get("제품명")) if item is not None else ""
            blockers.append(f"화학물질 목록 {row_no}행({label or '물질'}): 별표4 시설정보가 한 줄 이상 필요합니다.")

    comparison, compare_blockers = compare_with_legal_rules(calculations, legal_hits)
    blockers.extend(compare_blockers)
    blockers = list(dict.fromkeys(blockers))

    if blockers:
        return CAPHoldingResult(
            status="HOLD",
            label="화사계 최대보유량 판정보류",
            app4_ready=True,
            facility_rows=calculations,
            comparison_rows=comparison,
            blockers=blockers,
            messages=["별표4 시설별 계산은 수행했지만 미확인 정보가 있어 규정수량 비교를 확정하지 않습니다."],
        )

    upper = any(row.get("quantity_band") == "상위 규정수량 이상" for row in comparison)
    lower = any(row.get("quantity_band") == "하위 이상·상위 미만" for row in comparison)
    if upper:
        label = "최대보유량이 상위 규정수량 이상"
        status = "UPPER_CANDIDATE"
    elif lower:
        label = "최대보유량이 하위 규정수량 이상·상위 규정수량 미만"
        status = "LOWER_CANDIDATE"
    else:
        label = "확인된 유해화학물질의 최대보유량이 하위 규정수량 미만"
        status = "BELOW_LOWER"

    return CAPHoldingResult(
        status=status,
        label=label,
        app4_ready=True,
        facility_rows=calculations,
        comparison_rows=comparison,
        messages=[
            "승인된 별표 4 기준으로 시설별 최대보유량을 산정하고 동일 법적 규칙별로 합산했습니다.",
            "이 결과는 규정수량 비교 단계이며, 화사계 1군·2군 최종확정에는 면제조건과 군 분류 규칙 확인이 추가로 필요합니다.",
        ],
    )
