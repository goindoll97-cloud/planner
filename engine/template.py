from __future__ import annotations

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation


TITLE_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FILL = PatternFill("solid", fgColor="5B9BD5")
INPUT_FILL = PatternFill("solid", fgColor="FFF2CC")
NOTE_FILL = PatternFill("solid", fgColor="F3F6F9")
EXAMPLE_FILL = PatternFill("solid", fgColor="E2F0D9")
CAP_FILL = PatternFill("solid", fgColor="FFF2CC")
PSM_FILL = PatternFill("solid", fgColor="FCE4D6")
WHITE_FONT = Font(color="FFFFFF", bold=True)


def _write_row(ws, row_index: int, values: list[object]) -> None:
    for col_index, value in enumerate(values, start=1):
        ws.cell(row=row_index, column=col_index, value=value)


def _style_header(ws, cell_range: str) -> None:
    for row in ws[cell_range]:
        for cell in row:
            cell.fill = HEADER_FILL
            cell.font = WHITE_FONT
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _fill_range(ws, cell_range: str, fill: PatternFill) -> None:
    for row in ws[cell_range]:
        for cell in row:
            cell.fill = fill


def _wrap_range(ws, cell_range: str) -> None:
    for row in ws[cell_range]:
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)


def _add_list_validation(ws, cell_range: str, values: list[str]) -> None:
    formula = '"' + ",".join(values) + '"'
    dv = DataValidation(type="list", formula1=formula, allow_blank=True)
    dv.error = "목록에서 값을 선택해 주세요."
    dv.errorTitle = "입력값 확인"
    dv.prompt = "해당하지 않으면 '해당없음', 확인하지 못했으면 '모름'을 선택하세요."
    dv.promptTitle = "작성 도움말"
    dv.showInputMessage = True
    dv.showErrorMessage = True
    ws.add_data_validation(dv)
    dv.add(cell_range)


def _build_guide_sheet(wb: Workbook) -> None:
    ws = wb.active
    ws.title = "00_작성가이드"

    ws.merge_cells("A1:H1")
    ws["A1"] = "PSM·화학사고예방관리계획서 회사 입력파일 작성가이드"
    ws["A1"].fill = TITLE_FILL
    ws["A1"].font = Font(color="FFFFFF", bold=True, size=14)
    ws["A1"].alignment = Alignment(vertical="center")
    ws.row_dimensions[1].height = 28

    ws.merge_cells("A3:H3")
    ws["A3"] = (
        "이 파일은 누가 미리 작성해 놓은 예시입니다. 초록색 예시값을 귀사 정보로 바꾸고, "
        "선택형 항목은 드롭다운에서 고르세요."
    )
    ws["A3"].fill = EXAMPLE_FILL
    ws["A3"].font = Font(bold=True)
    ws["A3"].alignment = Alignment(wrap_text=True)

    rows = [
        ["구분", "무엇을 입력하나", "필수 수준", "쉽게 확인하는 곳", "해당하지 않을 때", "모를 때", "프로그램에서 쓰는 이유", "예시"],
        ["사업장", "사업장명·주소·업종·KSIC", "기본 필수", "사업자등록증·회사 기본정보", "-", "KSIC는 모름 가능", "시행령 제43조제1항 각 호의 사업 종류 및 사업장 식별", "KSIC 20111"],
        ["물질", "제품명·CAS·함량(%)", "기본 필수", "제품 SDS 제3항", "-", "추측하지 말고 모름", "PSM/화학사고예방관리계획서 물질기준", "포스겐 75-44-5, 100%"],
        ["수량", "최대 제조·사용량·최대 저장량", "기본 필수", "생산계획·탱크/창고 자료", "0 또는 실제 미취급값", "모름", "별표 13 유해·위험물질 규정량", "kg 또는 ton 권장"],
        ["화학사고예방관리계획서", "법정 사업장 최대보유량", "최종판정 핵심", "시설별 최대체류량·설계자료", "실제 0/미취급", "모름", "하위·상위 규정수량 비교", "04_시설별최대보유량 참고"],
        ["물질상태", "상온·상압 액체 여부", "조건부 필수", "SDS 제9항·물성자료", "해당없음", "모름", "상태별 규정수량 선택", "암모니아 N / 염산용액 Y"],
        ["SDS", "제품 SDS 제2항 유해성·위험성 분류", "조건부 필수", "현재 공급자/제조자 SDS 제2항", "해당없음", "모름", "별표 1 유해·위험성 그룹 연결", "인화성 액체 구분 2"],
        ["법정 예외·특수조건", "PSM 시행령 제43조제2항 제외설비·화관법 제23조제1항 단서·고시 제9조 작성 면제 시설 등", "조건부 필수", "설비용도·인허가·운영자료", "해당없음", "모름", "제출·작성 의무 및 작성수준 결정", "05_최종판정조건·06_PSM_비고8제외수량 참고"],
    ]
    for r_idx, row in enumerate(rows, start=5):
        _write_row(ws, r_idx, row)
    _style_header(ws, "A5:H5")
    _wrap_range(ws, "A5:H12")

    ws.merge_cells("A14:H14")
    ws["A14"] = (
        "입력 원칙: '해당없음' = 확인 결과 조건 자체가 적용되지 않음 / "
        "'모름' = 아직 확인하지 못함 → 판정보류 가능"
    )
    ws["A14"].fill = NOTE_FILL
    ws["A14"].font = Font(bold=True)
    ws["A14"].alignment = Alignment(wrap_text=True)

    widths = [24, 34, 18, 34, 28, 18, 42, 30]
    for i, width in enumerate(widths, 1):
        ws.column_dimensions[chr(64 + i)].width = width
    ws.freeze_panes = "A5"


def _build_business_sheet(wb: Workbook) -> None:
    ws = wb.create_sheet("01_사업장기본정보")
    ws.merge_cells("A1:E1")
    ws["A1"] = "1. 사업장 기본정보 — 초록색 예시값을 귀사 정보로 수정"
    ws["A1"].fill = TITLE_FILL
    ws["A1"].font = Font(color="FFFFFF", bold=True, size=13)

    _write_row(ws, 3, ["항목", "입력값", "필수", "프로그램 활용", "작성예시/설명"])
    _style_header(ws, "A3:E3")
    rows = [
        ["사업장명", "한빛정밀화학(주) 울산공장 (가상)", "Y", "공통", "귀사 사업장명을 입력"],
        ["사업장 주소", "울산광역시 남구 산업로 000 (가상)", "Y", "화학사고예방관리계획서", "실제 주소 입력"],
        ["업종 또는 주요 생산품", "석유화학계 기초화학물질 및 정밀화학 중간체 제조", "Y", "PSM", "주요 생산품까지 적으면 판정에 도움"],
        ["한국표준산업분류(KSIC) 코드", "20111", "N", "PSM", "모르면 '모름' 입력 가능"],
        ["기존 PSM 보유 여부", "Y", "Y", "공통", "Y / N / 해당없음 / 모름"],
        ["기존 장외영향평가서 보유 여부", "Y", "Y", "화학사고예방관리계획서", "Y / N / 해당없음 / 모름"],
        ["기존 화학사고예방관리계획서 보유 여부", "N", "Y", "화학사고예방관리계획서", "Y / N / 해당없음 / 모름"],
        ["현재 목적", "신규 사전진단", "Y", "공통", "신규 사전진단 / 변경검토 / 재제출검토"],
    ]
    for r_idx, row in enumerate(rows, start=4):
        _write_row(ws, r_idx, row)
        ws.cell(r_idx, 2).fill = EXAMPLE_FILL
    _wrap_range(ws, "A3:E11")
    _add_list_validation(ws, "B8:B10", ["Y", "N", "해당없음", "모름"])
    _add_list_validation(ws, "B11", ["신규 사전진단", "변경검토", "재제출검토"])

    widths = [34, 46, 10, 26, 52]
    for i, width in enumerate(widths, 1):
        ws.column_dimensions[chr(64 + i)].width = width
    ws.freeze_panes = "A4"


def _build_chemical_sheet(wb: Workbook) -> None:
    ws = wb.create_sheet("02_화학물질목록")
    ws.merge_cells("A1:O1")
    ws["A1"] = "2. 화학물질 목록 — 예시값을 귀사 물질로 수정"
    ws["A1"].fill = TITLE_FILL
    ws["A1"].font = Font(color="FFFFFF", bold=True, size=13)

    ws.merge_cells("A2:O2")
    ws["A2"] = "※ 선택형 정보가 귀사에 적용되지 않으면 '해당없음', 아직 확인하지 못했으면 '모름'을 선택하세요."
    ws["A2"].fill = NOTE_FILL
    ws["A2"].alignment = Alignment(wrap_text=True)

    headers = [
        "No.", "제품명", "CAS No.", "물질명(알면 입력)", "함량(%)", "취급형태",
        "최대 제조·사용량", "최대 저장량", "수량 단위", "최대 동시보유량(알면 입력)", "비고",
        "상온·상압 액체 여부(해당 시)", "최대보유량 법정 산정 여부",
        "회사/제품 SDS 제2항 보유·확인 여부", "SDS 제2항 유해성·위험성 분류(선택 입력)",
    ]
    _write_row(ws, 3, headers)
    _style_header(ws, "A3:O3")

    examples = [
        [1, "메틸 이소시아네이트", "624-83-9", "메틸 이소시아네이트", 100, "사용", 600, 0, "kg", 600, "100% 단일물질 예시", "해당없음", "Y", "Y", "인화성 액체 : 구분 2|급성 독성(경구) : 구분 3|급성 독성(경피) : 구분 3"],
        [2, "포스겐", "75-44-5", "포스겐", 100, "사용", 250, 1600, "kg", 1600, "1군 판정 테스트용 가상값", "해당없음", "Y", "해당없음", "해당없음"],
        [3, "염소", "7782-50-5", "염소", 100, "저장", 0, 800, "kg", 800, "사고대비물질 예시", "해당없음", "Y", "해당없음", "해당없음"],
        [4, "암모니아", "7664-41-7", "암모니아", 100, "사용", 1200, 4500, "kg", 5000, "상온·상압 액체 여부 확인 예시", "N", "Y", "해당없음", "해당없음"],
        [5, "염산 35%", "7647-01-0", "염화수소", 35, "사용", 2000, 6000, "kg", 6500, "농도조건 예시", "Y", "Y", "해당없음", "해당없음"],
        [6, "과산화수소 35%", "7722-84-1", "과산화수소", 35, "사용", 1000, 3000, "kg", 2500, "농도조건 예시", "해당없음", "Y", "해당없음", "해당없음"],
        [7, "톨루엔", "108-88-3", "톨루엔", 100, "사용", 1500, 4500, "kg", 4000, "별표 1/SDS 예시", "해당없음", "Y", "Y", "인화성 액체 : 구분 2"],
        [8, "이소프로필알코올 수용액 70%", "67-63-0", "2-프로판올", 70, "사용", 800, 1200, "kg", 1500, "혼합제품: 회사 제품 SDS 제2항 확인", "해당없음", "Y", "Y", "인화성 액체 : 구분 2"],
        [9, "메탄올", "67-56-1", "메탄올", 100, "저장", 1000, 7000, "kg", 6500, "SDS 다중분류 예시", "해당없음", "Y", "Y", "인화성 액체 : 구분 2|특정표적장기 독성(1회 노출) : 구분 1"],
        [10, "질산 68%", "7697-37-2", "질산", 68, "사용", 1500, 3000, "kg", 2800, "농도조건 예시", "해당없음", "Y", "해당없음", "해당없음"],
        [11, "수산화나트륨 수용액 30%", "1310-73-2", "수산화나트륨", 30, "사용", 600, 2500, "kg", 2000, "별표 1 해당없음 예시", "해당없음", "Y", "Y", "별표1 해당없음"],
        [12, "아세톤", "67-64-1", "아세톤", 100, "사용", 900, 1800, "kg", 1600, "SDS 예시", "해당없음", "Y", "Y", "인화성 액체 : 구분 2"],
    ]
    for r_idx, row in enumerate(examples, start=4):
        _write_row(ws, r_idx, row)
    _fill_range(ws, "B4:O15", EXAMPLE_FILL)

    for r in range(16, 104):
        ws.cell(r, 1, r - 3)
        for c in range(2, 16):
            ws.cell(r, c).fill = INPUT_FILL

    _add_list_validation(ws, "F4:F103", ["제조", "사용", "저장", "보관", "제조+저장", "사용+저장", "기타", "해당없음", "모름"])
    _add_list_validation(ws, "I4:I103", ["kg", "ton", "L", "m3", "기타"])
    for col in ["L", "M", "N"]:
        _add_list_validation(ws, f"{col}4:{col}103", ["Y", "N", "해당없음", "모름"])

    _wrap_range(ws, "A1:O103")
    widths = [7, 29, 16, 24, 11, 15, 18, 16, 12, 25, 42, 24, 24, 29, 58]
    for i, width in enumerate(widths, 1):
        ws.column_dimensions[chr(64 + i)].width = width
    ws.freeze_panes = "D4"


def _build_documents_sheet(wb: Workbook) -> None:
    ws = wb.create_sheet("03_기존문서보유여부")
    ws.merge_cells("A1:E1")
    ws["A1"] = "3. 기존 작성자료 보유 여부"
    ws["A1"].fill = TITLE_FILL
    ws["A1"].font = Font(color="FFFFFF", bold=True, size=13)
    _write_row(ws, 3, ["자료", "보유 여부", "프로그램 활용 시점", "왜 확인하는가", "예시/비고"])
    _style_header(ws, "A3:E3")
    rows = [
        ["공정안전보고서(PSM)", "Y", "PSM 대상 판정 후", "기존 공정안전자료 재활용", "2024년 작성본 보유 가정"],
        ["장외영향평가서", "Y", "화학사고예방관리계획서 대상 판정 후", "기존 시설·사고영향 자료 재활용", "과거 작성본 보유 가정"],
        ["기존 화학사고예방관리계획서", "N", "변경/재제출 판정 시", "신규/변경 구분", "현재 없음 가정"],
        ["위해관리계획서", "Y", "화학사고예방관리계획서 대상 판정 후", "기존 비상대응 정보 재활용", "기존 자료 보유 가정"],
    ]
    for r_idx, row in enumerate(rows, start=4):
        _write_row(ws, r_idx, row)
        ws.cell(r_idx, 2).fill = EXAMPLE_FILL
    _add_list_validation(ws, "B4:B7", ["Y", "N", "해당없음", "모름"])
    _wrap_range(ws, "A1:E7")
    widths = [34, 14, 34, 46, 38]
    for i, width in enumerate(widths, 1):
        ws.column_dimensions[chr(64 + i)].width = width


def _build_facility_sheet(wb: Workbook) -> None:
    ws = wb.create_sheet("04_시설별최대보유량")
    ws.merge_cells("A1:V1")
    ws["A1"] = "화학사고예방관리계획서 사업장 최대보유량 산정용 시설정보 — 시설 하나당 한 줄"
    ws["A1"].fill = CAP_FILL
    ws["A1"].font = Font(bold=True, size=13)

    ws.merge_cells("A2:V2")
    ws["A2"] = (
        "※ 법정 사업장 최대보유량을 이미 정확히 알고 있으면 02 시트에 값을 입력하고 '최대보유량 법정 산정 여부=Y'로 표시하세요. "
        "그렇지 않으면 이 시트에 시설별 정보를 작성합니다."
    )
    ws["A2"].fill = NOTE_FILL
    ws["A2"].alignment = Alignment(wrap_text=True)

    headers = [
        "적용여부", "목록행번호", "제품명", "CAS No.", "시설명", "시설유형", "제외시설여부", "제외사유",
        "물질성상", "공정유형", "별표4 기준함량(%)", "함량근거", "설계용량", "용량단위",
        "비중 또는 밀도(kg/L=ton/m3)", "보관계획도 최대량", "일일최대보관량", "질량단위",
        "직접확인 최대보유량", "직접확인 근거", "복수성상 증빙", "비고",
    ]
    _write_row(ws, 3, headers)
    _style_header(ws, "A3:V3")

    examples = [
        ["해당", 1, "메틸 이소시아네이트", "624-83-9", "R-101 반응기", "제조·사용시설", "N", "해당없음", "액체", "반응", 100, "제품 SDS 및 공정배합표", 0.35, "m3", 0.96, None, None, "kg", None, None, "N", "가상 예시"],
        ["해당", 1, "메틸 이소시아네이트", "624-83-9", "T-101 원료탱크", "저장탱크", "N", "해당없음", "액체", "해당없음", 100, "제품 SDS", 0.30, "m3", 0.96, None, None, "kg", None, None, "N", "가상 예시"],
        ["해당", 2, "포스겐", "75-44-5", "V-201 공급용기군", "기타", "N", "해당없음", "기체·고압가스", "해당없음", 100, "제품 SDS", None, "해당없음", None, None, None, "kg", 1600, "1군 테스트용: 법정 최대보유량 합계 1.6 ton", "N", "가상값"],
        ["해당", 3, "염소", "7782-50-5", "V-301 염소용기군", "기타", "N", "해당없음", "기체·고압가스", "해당없음", 100, "제품 SDS", None, "해당없음", None, None, None, "kg", 800, "최대 동시 연결·보관 용기 질량 합계", "N", "가상값"],
        ["해당", 4, "암모니아", "7664-41-7", "T-401 암모니아 저장조", "기타", "N", "해당없음", "기체·고압가스", "해당없음", 100, "제품 SDS", None, "해당없음", None, None, None, "kg", 5000, "설비 운영자료상 최대보유 질량", "N", "가상값"],
        ["해당", 5, "염산 35%", "7647-01-0", "T-501 염산탱크", "저장탱크", "N", "해당없음", "액체", "변화없음", 35, "제품 SDS", 5.5, "m3", 1.18, None, None, "kg", None, None, "N", "가상 예시"],
        ["해당", 7, "톨루엔", "108-88-3", "T-701 톨루엔탱크", "저장탱크", "N", "해당없음", "액체", "변화없음", 100, "제품 SDS", 4.6, "m3", 0.87, None, None, "kg", None, None, "N", "가상 예시"],
        ["해당", 8, "이소프로필알코올 수용액 70%", "67-63-0", "T-801 IPA 혼합액탱크", "저장탱크", "N", "해당없음", "액체", "변화없음", 70, "회사 제품 SDS", 1.9, "m3", 0.79, None, None, "kg", None, None, "N", "혼합제품 예시"],
    ]
    for r_idx, row in enumerate(examples, start=4):
        _write_row(ws, r_idx, row)
    _fill_range(ws, "A4:V11", EXAMPLE_FILL)

    for r in range(12, 31):
        ws.cell(r, 1, "해당없음")
        for c in range(1, 23):
            ws.cell(r, c).fill = INPUT_FILL

    _add_list_validation(ws, "A4:A30", ["해당", "해당없음", "모름"])
    _add_list_validation(ws, "F4:F30", ["제조·사용시설", "저장탱크", "보관시설", "기타", "해당없음", "모름"])
    _add_list_validation(ws, "G4:G30", ["Y", "N", "해당없음", "모름"])
    _add_list_validation(ws, "H4:H30", ["해당없음", "탱크로리·운송차량", "사외배관", "취급중단 신고시설", "기타", "모름"])
    _add_list_validation(ws, "I4:I30", ["액체", "고체", "기체·고압가스", "복수성상", "해당없음", "모름"])
    _add_list_validation(ws, "J4:J30", ["변화없음", "단순혼합", "반응", "해당없음", "모름"])
    _add_list_validation(ws, "N4:N30", ["L", "m3", "해당없음"])
    _add_list_validation(ws, "R4:R30", ["kg", "ton", "해당없음"])
    _add_list_validation(ws, "U4:U30", ["Y", "N", "해당없음", "모름"])

    _wrap_range(ws, "A1:V30")
    widths = [12, 10, 28, 16, 24, 18, 15, 22, 18, 16, 18, 28, 14, 12, 24, 20, 20, 12, 22, 36, 16, 28]
    for i, width in enumerate(widths, 1):
        ws.column_dimensions[chr(64 + i)].width = width
    ws.freeze_panes = "A4"


def _build_final_conditions_sheet(wb: Workbook) -> None:
    ws = wb.create_sheet("05_최종판정조건")
    ws.merge_cells("A1:F1")
    ws["A1"] = "5. 최종 판정 조건 — 회사에서 확인한 사실을 Excel에 직접 작성"
    ws["A1"].fill = PSM_FILL
    ws["A1"].font = Font(bold=True, size=13)
    _write_row(ws, 2, ["제도", "확인항목", "입력값", "필수수준", "확인자료 예시", "판정에 미치는 영향"])
    _style_header(ws, "A2:F2")

    rows = [
        ["공정안전보고서", "시행령 제43조제2항 제외설비 해당 여부", "N", "조건부 필수", "설비 용도·인허가·도면", "공정안전보고서 제출 대상 여부"],
        ["공정안전보고서", "시행령 제43조제2항 제외설비 유형", "해당없음", "해당 시 필수", "법령상 제외설비 유형을 그대로 작성", "제외설비 법적 유형 확인"],
        ["공정안전보고서", "별표 13 제1호 인화성 가스 해당 여부", "N", "조건부 필수", "회사 SDS·공정안전자료", "별표 13 제1호 적용 여부"],
        ["공정안전보고서", "별표 13 제1호 하루 최대 제조·취급량(kg)", 0, "제1호 해당 시 필수", "하루 최대 제조·취급량", "비고 제7호 합산한 값(R)"],
        ["공정안전보고서", "별표 13 제1호 최대 저장량(kg)", 0, "제1호 해당 시 필수", "최대 저장량", "비고 제7호 합산한 값(R)"],
        ["공정안전보고서", "별표 13 제2호 인화성 액체 해당 여부", "N", "조건부 필수", "회사 SDS·공정안전자료", "별표 13 제2호 적용 여부"],
        ["공정안전보고서", "별표 13 제2호 하루 최대 제조·취급량(kg)", 0, "제2호 해당 시 필수", "하루 최대 제조·취급량", "비고 제7호 합산한 값(R)"],
        ["공정안전보고서", "별표 13 제2호 최대 저장량(kg)", 0, "제2호 해당 시 필수", "최대 저장량", "비고 제7호 합산한 값(R)"],
        ["공정안전보고서", "별표 13 제23호 발연황산 삼산화황(SO3) 중량%", "해당없음", "해당 물질 보유 시 필수", "제품 SDS·성분분석자료", "별표 13 제23호 성분조건"],
        ["공정안전보고서", "별표 13 제42호 니트로셀룰로오스 질소 함유량%", "해당없음", "해당 물질 보유 시 필수", "제품 SDS·성분분석자료", "별표 13 제42호 성분조건"],
        ["공정안전보고서", "가스를 전문으로 저장·판매하는 시설 내 가스 여부", "N", "조건부 필수", "시설 용도·사업형태", "별표 13 비고 제8호"],
        ["화학사고예방관리계획서", "법 제23조제1항 단서 해당 여부", "N", "규정수량 이상 시 필수", "시설 용도·인허가·운영현황", "작성·제출 의무 여부"],
        ["화학사고예방관리계획서", "법적 예외 적용 유형", "해당없음", "해당 시 필수", "법 제23조제1항 단서·시행규칙 제19조제2항·고시 제9조의 유형", "법적 예외 유형 확인"],
        ["화학사고예방관리계획서", "법적 예외가 관련 취급시설 전체에 적용되는지", "Y", "해당 시 필수", "시설별 적용범위", "전체 적용 여부"],
        ["화학사고예방관리계획서", "상위 규정수량 이상을 취급하는 개별 주요취급시설 존재 여부", "Y", "상위 규정수량 해당 시 필수", "시설별 최대보유량·설계자료", "작성수준 1군/2군"],
        ["화학사고예방관리계획서", "회사/제품 SDS 제2항이 KOSHA 자동조회 분류와 일치하는지", "Y", "자동조회 사용 시", "현재 공급자 SDS 제2항", "별표 1 규정수량 확정"],
        ["공통", "미확인 결정조건 존재 여부", "N", "필수", "사내 검토", "Y면 판정보류"],
    ]
    for r_idx, row in enumerate(rows, start=3):
        _write_row(ws, r_idx, row)
        ws.cell(r_idx, 3).fill = EXAMPLE_FILL

    for cell in ["C3", "C5", "C8", "C13", "C14", "C16", "C17", "C18", "C19"]:
        _add_list_validation(ws, cell, ["Y", "N", "해당없음", "모름"])
    _wrap_range(ws, "A1:F19")
    widths = [26, 58, 20, 24, 40, 32]
    for i, width in enumerate(widths, 1):
        ws.column_dimensions[chr(64 + i)].width = width
    ws.freeze_panes = "A3"


def _build_psm_note8_sheet(wb: Workbook) -> None:
    ws = wb.create_sheet("06_PSM_비고8제외수량")
    ws.merge_cells("A1:G1")
    ws["A1"] = "6. PSM 별표 13 비고 제8호 제외수량 — 전문 가스 저장·판매시설에 해당할 때만 작성"
    ws["A1"].fill = PSM_FILL
    ws["A1"].font = Font(bold=True, size=13)
    ws.merge_cells("A2:G2")
    ws["A2"] = "※ 05_최종판정조건에서 '가스를 전문으로 저장·판매하는 시설 내 가스 여부=Y'인 경우에만 작성합니다. 시설 하나당 한 줄로 작성하세요."
    ws["A2"].fill = NOTE_FILL
    ws["A2"].alignment = Alignment(wrap_text=True)
    headers = ["적용여부", "별표13 호수", "제조·취급 제외량(kg)", "저장 제외량(kg)", "시설명", "근거", "비고"]
    _write_row(ws, 3, headers)
    _style_header(ws, "A3:G3")
    _write_row(ws, 4, ["해당없음", None, None, None, None, None, "비고 제8호 미해당이면 이 시트는 작성하지 않아도 됩니다."])
    _fill_range(ws, "A4:G4", EXAMPLE_FILL)
    for r in range(5, 31):
        for c in range(1, 8):
            ws.cell(r, c).fill = INPUT_FILL
    _add_list_validation(ws, "A4:A30", ["해당", "해당없음", "모름"])
    _wrap_range(ws, "A1:G30")
    widths = [12, 14, 24, 20, 26, 38, 38]
    for i, width in enumerate(widths, 1):
        ws.column_dimensions[chr(64 + i)].width = width
    ws.freeze_panes = "A4"



def build_legal_reference_workbook() -> bytes:
    """Return a separate law-only reference workbook for company preparers."""
    wb = Workbook()

    guide = wb.active
    guide.title = "00_사용안내"
    guide.merge_cells("A1:F1")
    guide["A1"] = "PSM·화학사고예방관리계획서 법령 작성 참고"
    guide["A1"].fill = TITLE_FILL
    guide["A1"].font = Font(color="FFFFFF", bold=True, size=14)
    guide["A1"].alignment = Alignment(vertical="center")
    guide.row_dimensions[1].height = 28
    guide.merge_cells("A3:F3")
    guide["A3"] = (
        "회사 입력파일을 작성할 때 법령 확인이 필요한 항목만 정리한 별도 참고파일입니다. "
        "회사 내부자료의 위치·문서명·확인방법은 수록하지 않습니다."
    )
    guide["A3"].fill = NOTE_FILL
    guide["A3"].alignment = Alignment(wrap_text=True)
    guide.merge_cells("A5:F5")
    guide["A5"] = "법령은 개정될 수 있으므로 실제 판정 시 프로그램의 법령 동기화 상태와 국가법령정보센터 최신 현행본을 함께 확인하세요."
    guide["A5"].fill = NOTE_FILL
    guide["A5"].font = Font(bold=True)
    guide["A5"].alignment = Alignment(wrap_text=True)
    _write_row(guide, 7, ["구분", "현행 기준", "핵심 법령", "핵심 조문·별표", "비고", "공식 링크"])
    _style_header(guide, "A7:F7")
    guide_rows = [
        ["공정안전보고서(PSM)", "2026-09-12 확인", "「산업안전보건법」 / 「산업안전보건법 시행령」", "법 제44조제1항 / 시행령 제43조 / 별표 13", "별표 13: 유해·위험물질 규정량", "https://www.law.go.kr/법령/산업안전보건법시행령/제43조"],
        ["화학사고예방관리계획서", "2026-09-12 확인", "「화학물질관리법」 / 같은 법 시행규칙", "법 제23조 / 시행규칙 제19조", "작성·제출 의무 및 법정 예외", "https://www.law.go.kr/법령/화학물질관리법/제23조"],
        ["화학사고예방관리계획서", "시행 2026.5.6", "「유해화학물질의 규정수량에 관한 규정」", "제2조~제4조 / 별표 1~4", "화학물질안전원고시 제2026-4호", "https://www.law.go.kr/admRulLsInfoP.do?admRulId=93578&efYd=0"],
        ["화학사고예방관리계획서", "시행 2026.4.22", "「화학사고예방관리계획서 작성 등에 관한 규정」", "제2조 / 제4조 / 제6조 / 제9조", "화학물질안전원고시 제2026-7호", "https://www.law.go.kr/admRulInfoP.do?admRulSeq=2100000278102"],
    ]
    for row_idx, row in enumerate(guide_rows, start=8):
        _write_row(guide, row_idx, row)
        link = guide.cell(row_idx, 6)
        link.hyperlink = row[5]
        link.style = "Hyperlink"
    _wrap_range(guide, "A1:F11")
    for idx, width in enumerate([28, 22, 48, 46, 44, 50], 1):
        guide.column_dimensions[chr(64 + idx)].width = width
    guide.freeze_panes = "A8"

    def build_reference_sheet(title: str, rows: list[list[str]]) -> None:
        ws = wb.create_sheet(title)
        ws.merge_cells("A1:F1")
        ws["A1"] = f"{title.replace('_', ' ')} — 회사 입력파일 작성항목별 법령 참고"
        ws["A1"].fill = TITLE_FILL
        ws["A1"].font = Font(color="FFFFFF", bold=True, size=13)
        _write_row(ws, 3, ["입력파일 시트", "작성 항목", "법령에서 확인할 내용", "관련 법령·조문·별표", "현행 기준", "국가법령정보센터"])
        _style_header(ws, "A3:F3")
        for row_idx, row in enumerate(rows, start=4):
            _write_row(ws, row_idx, row)
            link = ws.cell(row_idx, 6)
            link.hyperlink = row[5]
            link.style = "Hyperlink"
        _wrap_range(ws, f"A1:F{3 + len(rows)}")
        for idx, width in enumerate([25, 43, 58, 70, 32, 50], 1):
            ws.column_dimensions[chr(64 + idx)].width = width
        ws.freeze_panes = "A4"

    psm_url = "https://www.law.go.kr/법령/산업안전보건법시행령/제43조"
    psm_rows = [
        ["01_사업장기본정보", "업종 또는 주요 생산품 / KSIC 코드", "시행령 제43조제1항 각 호의 사업 종류 해당 여부 확인", "「산업안전보건법」 제44조제1항; 「산업안전보건법 시행령」 제43조제1항", "시행령 [시행 2026.8.1.] 대통령령 제36540호", psm_url],
        ["02_화학물질목록", "CAS No. / 물질명 / 함량(%)", "별표 13 개별 유해·위험물질 및 농도·성분조건 해당 여부 확인", "「산업안전보건법 시행령」 제43조제1항 및 별표 13 「유해·위험물질 규정량」", "별표 13 현행본", psm_url],
        ["02_화학물질목록", "최대 제조·사용량 / 최대 저장량 / 수량 단위", "각 별표 13 항목의 제조·취급 규정량과 저장 규정량 비교", "「산업안전보건법 시행령」 별표 13 및 비고 제1호·제7호", "별표 13 현행본", psm_url],
        ["05_최종판정조건", "시행령 제43조제2항 제외설비 해당 여부 / 유형", "공정안전보고서 제출 대상에서 제외되는 설비인지 확인", "「산업안전보건법 시행령」 제43조제2항", "시행령 현행본", psm_url],
        ["05_최종판정조건", "별표 13 제1호 인화성 가스 해당 여부 및 수량", "인화성 가스에 해당하는지와 제조·취급 5,000kg / 저장 200,000kg 규정량 적용", "「산업안전보건법 시행령」 별표 13 제1호", "제조·취급 5,000kg / 저장 200,000kg", psm_url],
        ["05_최종판정조건", "별표 13 제2호 인화성 액체 해당 여부 및 수량", "인화성 액체에 해당하는지와 제조·취급 5,000kg / 저장 200,000kg 규정량 적용", "「산업안전보건법 시행령」 별표 13 제2호", "제조·취급 5,000kg / 저장 200,000kg", psm_url],
        ["05_최종판정조건", "별표 13 제23호 발연황산 SO3 중량%", "삼산화황 중량 65% 이상 80% 미만 조건 해당 여부 확인", "「산업안전보건법 시행령」 별표 13 제23호", "SO3 65% 이상 80% 미만", psm_url],
        ["05_최종판정조건", "별표 13 제42호 니트로셀룰로오스 질소 함유량%", "질소 함유량 12.6% 이상 조건 해당 여부 확인", "「산업안전보건법 시행령」 별표 13 제42호", "질소 12.6% 이상", psm_url],
        ["05_최종판정조건", "가스를 전문으로 저장·판매하는 시설 내 가스 여부", "별표 13 비고 제8호 적용 여부 확인", "「산업안전보건법 시행령」 별표 13 비고 제8호", "별표 13 현행본", psm_url],
        ["06_PSM_비고8제외수량", "별표13 호수 / 제조·취급 제외량 / 저장 제외량", "비고 제8호에 따라 제외할 수량을 항목별로 구분하고 비고 제7호 합산한 값(R)에 반영", "「산업안전보건법 시행령」 별표 13 비고 제7호·제8호", "별표 13 현행본", psm_url],
    ]
    build_reference_sheet("01_PSM_법령참고", psm_rows)

    cap_law_url = "https://www.law.go.kr/법령/화학물질관리법/제23조"
    cap_rule_url = "https://www.law.go.kr/법령/화학물질관리법시행규칙/제19조"
    cap_qty_url = "https://www.law.go.kr/admRulLsInfoP.do?admRulId=93578&efYd=0"
    cap_write_url = "https://www.law.go.kr/admRulInfoP.do?admRulSeq=2100000278102"
    cap_rows = [
        ["02_화학물질목록", "CAS No. / 물질명 / 함량(%)", "유해화학물질별 상위·하위 규정수량 및 유해·위험성 그룹·물질범위 확인", "「유해화학물질의 규정수량에 관한 규정」 제3조 및 별표 1~3", "화학물질안전원고시 제2026-4호", cap_qty_url],
        ["02_화학물질목록", "상온·상압 액체 여부", "물질 성상에 따라 적용되는 규정수량 및 최대보유량 산정조건 확인", "「유해화학물질의 규정수량에 관한 규정」 제3조·제4조 및 별표 1~4", "화학물질안전원고시 제2026-4호", cap_qty_url],
        ["02_화학물질목록", "최대 동시보유량 / 최대보유량 법정 산정 여부", "사업장 최대보유량의 정의 및 물질별 규정수량 비교방법 확인", "「유해화학물질의 규정수량에 관한 규정」 제2조제2호·제4조; 「화학사고예방관리계획서 작성 등에 관한 규정」 제4조", "고시 제2026-4호 / 제2026-7호", cap_qty_url],
        ["02_화학물질목록", "SDS 제2항 유해성·위험성 분류", "별표 1 유해·위험성 그룹별 규정수량 적용 여부 확인", "「유해화학물질의 규정수량에 관한 규정」 제3조제1호 및 별표 1", "화학물질안전원고시 제2026-4호", cap_qty_url],
        ["04_시설별최대보유량", "시설유형 / 제외시설여부 / 제외사유", "최대보유량에 포함되는 취급시설과 제외되는 범위 확인", "「유해화학물질의 규정수량에 관한 규정」 제4조 및 별표 4; 「화학사고예방관리계획서 작성 등에 관한 규정」 제2조제11호·제4조", "고시 제2026-4호 / 제2026-7호", cap_write_url],
        ["04_시설별최대보유량", "물질성상 / 공정유형 / 별표4 기준함량(%)", "제조·사용시설의 함량기준과 공정유형별 적용조건 확인", "「유해화학물질의 규정수량에 관한 규정」 별표 4", "화학물질안전원고시 제2026-4호", cap_qty_url],
        ["04_시설별최대보유량", "설계용량 / 용량단위 / 비중 또는 밀도", "저장탱크 및 제조·사용시설 최대보유량의 질량 환산기준 확인", "「유해화학물질의 규정수량에 관한 규정」 제4조 및 별표 4", "화학물질안전원고시 제2026-4호", cap_qty_url],
        ["04_시설별최대보유량", "보관계획도 최대량 / 일일최대보관량", "보관시설 최대보유량 산정기준 확인", "「유해화학물질의 규정수량에 관한 규정」 별표 4", "화학물질안전원고시 제2026-4호", cap_qty_url],
        ["04_시설별최대보유량", "직접확인 최대보유량 / 복수성상 증빙", "기체·고압가스·복수성상 등 별표 4의 특수 산정조건 확인", "「유해화학물질의 규정수량에 관한 규정」 별표 4 및 비고", "화학물질안전원고시 제2026-4호", cap_qty_url],
        ["05_최종판정조건", "법 제23조제1항 단서 해당 여부", "화학사고예방관리계획서 작성·제출 의무의 법정 예외 해당 여부 확인", "「화학물질관리법」 제23조제1항 단서; 「화학물질관리법 시행규칙」 제19조제2항; 「화학사고예방관리계획서 작성 등에 관한 규정」 제9조", "법률·시행규칙·고시 현행본", cap_law_url],
        ["05_최종판정조건", "법적 예외 적용 유형 / 관련 취급시설 전체 적용 여부", "법정 예외의 정확한 유형과 적용범위 확인", "「화학물질관리법」 제23조제1항; 같은 법 시행규칙 제19조제2항; 「화학사고예방관리계획서 작성 등에 관한 규정」 제9조", "고시 제2026-7호", cap_rule_url],
        ["05_최종판정조건", "상위 규정수량 이상을 취급하는 개별 주요취급시설 존재 여부", "1군 사업장 작성수준 판단에 필요한 주요취급시설 존재 여부 확인", "「화학물질관리법 시행규칙」 제19조제8항; 「화학사고예방관리계획서 작성 등에 관한 규정」 제2조·제4조", "화학물질안전원고시 제2026-7호", cap_write_url],
        ["05_최종판정조건", "작성수준 1군 / 2군 판단", "물질별 최대보유량과 상위·하위 규정수량을 비교하여 작성수준 결정", "「유해화학물질의 규정수량에 관한 규정」 제3조·제4조 및 별표 1~4; 「화학사고예방관리계획서 작성 등에 관한 규정」 제2조·제4조·제6조", "고시 제2026-4호 / 제2026-7호", cap_write_url],
    ]
    build_reference_sheet("02_화사계_법령참고", cap_rows)

    output = BytesIO()
    wb.save(output)
    return output.getvalue()

def build_minimal_input_workbook() -> bytes:
    """Return the guided company intake workbook as valid XLSX bytes.

    The historical function name is kept for compatibility. The workbook now
    includes example values, dropdowns, facility-level maximum-holding inputs
    and final-decision conditions. The old final-preparation checklist sheet is
    intentionally not created.
    """
    wb = Workbook()
    _build_guide_sheet(wb)
    _build_business_sheet(wb)
    _build_chemical_sheet(wb)
    _build_documents_sheet(wb)
    _build_facility_sheet(wb)
    _build_final_conditions_sheet(wb)
    _build_psm_note8_sheet(wb)

    output = BytesIO()
    wb.save(output)
    return output.getvalue()
