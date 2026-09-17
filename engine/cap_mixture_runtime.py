from __future__ import annotations

"""Backwards-compatible mixture/component handling for CAP Stage 1.

A commercial mixture stays one product row in ``02_화학물질목록``.  Its SDS
Section-3 components are listed in ``02A_혼합물구성성분``.  CAP Appendix 3 and
Appendix 2 screening is then performed per component, while Appendix-4 maximum
holding remains the whole mixture mass (the component percentage is a legal
threshold test, not a mass multiplier).

The hook is fail-closed.  Old workbooks without the new sheet keep their current
behaviour; explicit mixtures require component CAS/concentration/evidence; a
concentration range that crosses a legal threshold is held; and a multi-
component reaction/simple-mixing facility must identify which component CAS its
process concentration belongs to.
"""

from io import BytesIO
from typing import Any, Callable, Iterable

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation


COMPONENT_SHEET = "02A_혼합물구성성분"
MIXTURE_FLAG_COLUMN = "혼합물 여부"
FACILITY_COMPONENT_CAS_COLUMN = "규제성분 CAS(혼합물)"
_SHORT_NAME = "화사계"
_FULL_NAME = "화학사고예방관리계획서"
COMPONENT_COLUMNS = [
    "적용여부", "제품목록행번호", "제품명(확인용)", "구성성분명", "CAS No.",
    "함량(%)", "함량 최저(%)", "함량 최고(%)", "SDS 제3항 근거", "비고",
]

TITLE_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FILL = PatternFill("solid", fgColor="5B9BD5")
INPUT_FILL = PatternFill("solid", fgColor="FFF2CC")
EXAMPLE_FILL = PatternFill("solid", fgColor="E2F0D9")
NOTE_FILL = PatternFill("solid", fgColor="F3F6F9")
WHITE_FONT = Font(color="FFFFFF", bold=True)


def _clean(value: Any) -> str:
    if value is None:
        return ""
    try:
        if pd.isna(value):
            return ""
    except (TypeError, ValueError):
        pass
    text = str(value).strip()
    return "" if text.lower() in {"nan", "none", "null", "<na>"} else text


def _num(value: Any) -> float | None:
    text = _clean(value).replace(",", "")
    if not text:
        return None
    try:
        return float(text)
    except ValueError:
        return None


def _int(value: Any) -> int | None:
    value = _num(value)
    if value is None or int(value) != value:
        return None
    return int(value)


def _yes(value: Any) -> bool:
    return _clean(value).lower().replace(" ", "") in {"y", "yes", "예", "해당", "혼합물", "1", "true"}


def _no(value: Any) -> bool:
    return _clean(value).lower().replace(" ", "") in {"n", "no", "아니오", "아님", "단일물질", "0", "false", "해당없음"}


def _component_frame(intake: Any) -> pd.DataFrame:
    frame = getattr(intake, "mixture_components", None)
    return frame if isinstance(frame, pd.DataFrame) else pd.DataFrame()


def _mixture_rows(intake: Any) -> set[int]:
    return set(intake.mixture_parent_rows())


def _components_for_parent(intake: Any, parent: int) -> pd.DataFrame:
    frame = _component_frame(intake)
    if frame.empty or "제품목록행번호" not in frame.columns:
        return pd.DataFrame(columns=COMPONENT_COLUMNS)
    return frame[frame["제품목록행번호"].map(_int).eq(parent)].copy().reset_index(drop=True)


def _concentration_values(component: pd.Series) -> tuple[list[float], str]:
    exact = _num(component.get("함량(%)"))
    if exact is not None:
        return ([exact] if 0 <= exact <= 100 else []), "exact"
    low = _num(component.get("함량 최저(%)"))
    high = _num(component.get("함량 최고(%)"))
    if low is None or high is None or low < 0 or high > 100 or low > high:
        return [], "invalid"
    if low == high:
        return [low], "exact"
    return [low, high], "range"


def _clone_intake(intake: Any, *, chemicals: pd.DataFrame | None = None, facilities: pd.DataFrame | None = None):
    from .inventory import IntakeData

    return IntakeData(
        business=dict(getattr(intake, "business", {}) or {}),
        chemicals=(chemicals.copy() if chemicals is not None else intake.chemicals.copy()),
        documents=dict(getattr(intake, "documents", {}) or {}),
        facilities=(facilities.copy() if facilities is not None else getattr(intake, "facilities", pd.DataFrame()).copy()),
        final_conditions=dict(getattr(intake, "final_conditions", {}) or {}),
        psm_note8_exclusions=getattr(intake, "psm_note8_exclusions", pd.DataFrame()).copy(),
        mixture_components=_component_frame(intake).copy(),
        source_fingerprint=str(getattr(intake, "source_fingerprint", "") or ""),
    )


def _single_component_intake(intake: Any, parent: int, component: pd.Series, pct: float):
    row = intake.chemicals.iloc[parent - 1].copy()
    row["CAS No."] = _clean(component.get("CAS No."))
    row["물질명(알면 입력)"] = _clean(component.get("구성성분명")) or row["CAS No."]
    row["함량(%)"] = pct
    row[MIXTURE_FLAG_COLUMN] = "Y"
    return _clone_intake(intake, chemicals=pd.DataFrame([row]))


def _map_records(records: Iterable[dict[str, Any]], parent: int, product: str, component: pd.Series) -> list[dict[str, Any]]:
    output: list[dict[str, Any]] = []
    for raw in records:
        row = dict(raw)
        row["row_no"] = parent
        row["product_name"] = product
        row["mixture_product_name"] = product
        row["component_name"] = _clean(component.get("구성성분명"))
        row["component_cas"] = _clean(component.get("CAS No."))
        output.append(row)
    return output


def _signature(rows: Iterable[dict[str, Any]]) -> tuple[tuple[Any, ...], ...]:
    values = []
    for row in rows:
        values.append((
            _clean(row.get("source_key")), _clean(row.get("cas")), _clean(row.get("item_no")),
            _clean(row.get("designation_id")), _clean(row.get("variant_type")), _clean(row.get("hazard_category")),
            _num(row.get("content_threshold_pct")), _num(row.get("lowest_quantity_ton")),
            _num(row.get("lower_quantity_ton")), _num(row.get("upper_quantity_ton")),
        ))
    return tuple(sorted(values, key=str))


def _screen_component_direct(original_screen: Callable[[Any], Any], intake: Any, parent: int, component: pd.Series):
    values, mode = _concentration_values(component)
    product = _clean(intake.chemicals.iloc[parent - 1].get("제품명"))
    label = _clean(component.get("구성성분명")) or _clean(component.get("CAS No."))
    if not values:
        return [], [f"{product} / {label}: 혼합물 구성성분의 함량 또는 함량범위를 확인해 주세요."], True, True

    runs = [original_screen(_single_component_intake(intake, parent, component, pct)) for pct in values]
    if any(not getattr(run, "ready", False) for run in runs):
        blockers: list[str] = []
        for run in runs:
            blockers.extend(getattr(run, "blockers", []) or [])
        return [], list(dict.fromkeys(blockers)), False, False

    mapped = [_map_records(getattr(run, "legal_hits", []) or [], parent, product, component) for run in runs]
    signatures = [_signature(rows) for rows in mapped]
    claimed = any(bool(getattr(run, "row_numbers", []) or getattr(run, "blockers", [])) for run in runs)
    if mode == "range" and len(set(signatures)) > 1:
        low, high = values
        return [], [
            f"{product} / {label}: SDS 제3항 함량범위 {low:g}~{high:g}%가 법정 함량기준 또는 특수 규정수량 경계를 가로질러 적용 여부를 확정할 수 없습니다. "
            "법정 판정에 사용할 실제 최대함량 또는 공정 기준함량과 근거를 확인해 주세요."
        ], True, True
    blockers: list[str] = []
    for run in runs:
        blockers.extend(str(v) for v in (getattr(run, "blockers", []) or []))
    return (mapped[-1] if mapped else []), list(dict.fromkeys(blockers)), claimed, True


def _replace_short_name(value: Any) -> Any:
    if isinstance(value, str) and _SHORT_NAME in value:
        return value.replace(_SHORT_NAME, _FULL_NAME)
    return value


def _cleanup_company_workbook(raw: bytes) -> bytes:
    """Final user-facing pass: full legal names only, and blank future input
    rows are not left flagged as non-mixtures by the new mixture column."""
    wb = load_workbook(BytesIO(raw), data_only=False)

    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                replacement = _replace_short_name(cell.value)
                if replacement != cell.value:
                    cell.value = replacement

    if "02_화학물질목록" in wb.sheetnames:
        ws = wb["02_화학물질목록"]
        headers = {
            str(ws.cell(3, col).value or "").strip(): col
            for col in range(1, ws.max_column + 1)
        }
        mix_col = headers.get(MIXTURE_FLAG_COLUMN)
        if mix_col:
            for row_no in range(4, ws.max_row + 1):
                has_product_data = any(
                    ws.cell(row_no, col).value not in (None, "")
                    for col in (2, 3, 4, 5)  # 제품명, CAS, 물질명, 함량
                )
                if not has_product_data:
                    ws.cell(row_no, mix_col).value = None

    output = BytesIO()
    wb.save(output)
    return output.getvalue()


def _patch_template(original: Callable[[], bytes]) -> Callable[[], bytes]:
    def wrapped() -> bytes:
        raw = original()
        wb = load_workbook(BytesIO(raw))
        if COMPONENT_SHEET in wb.sheetnames:
            return _cleanup_company_workbook(raw)

        chem = wb["02_화학물질목록"]
        mix_col = chem.max_column + 1
        chem.cell(3, mix_col, MIXTURE_FLAG_COLUMN)
        chem.cell(3, mix_col).fill = HEADER_FILL
        chem.cell(3, mix_col).font = WHITE_FONT
        chem.cell(3, mix_col).alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for row in range(4, 104):
            if chem.cell(row, 1).value not in (None, ""):
                chem.cell(row, mix_col, "N")
            chem.cell(row, mix_col).fill = EXAMPLE_FILL if row <= 15 else INPUT_FILL
        # Existing aqueous examples: product row numbers 5,6,8,10,11.
        for worksheet_row in (8, 9, 11, 13, 14):
            chem.cell(worksheet_row, mix_col, "Y")
        dv = DataValidation(type="list", formula1='"N,Y,모름"', allow_blank=True)
        chem.add_data_validation(dv)
        dv.add(f"{chem.cell(3, mix_col).column_letter}4:{chem.cell(3, mix_col).column_letter}103")
        chem.column_dimensions[chem.cell(3, mix_col).column_letter].width = 16
        for merged in list(chem.merged_cells.ranges):
            if str(merged) in {"A1:O1", "A2:O2"}:
                chem.unmerge_cells(str(merged))
        chem.merge_cells(start_row=1, start_column=1, end_row=1, end_column=mix_col)
        chem.merge_cells(start_row=2, start_column=1, end_row=2, end_column=mix_col)
        chem["A2"] = (
            "※ 단일물질은 기존 방식대로 작성합니다. 혼합제품은 '혼합물 여부=Y'로 표시하고 02A_혼합물구성성분에 SDS 제3항의 규제 가능 성분을 모두 입력하세요. "
            "화사계 최대보유량은 성분량으로 줄이지 않고 혼합물 전체량을 사용합니다."
        )
        chem["A2"].fill = NOTE_FILL
        chem["A2"].alignment = Alignment(wrap_text=True)

        ws = wb.create_sheet(COMPONENT_SHEET, 3)
        ws.merge_cells("A1:J1")
        ws["A1"] = "2A. 혼합물 구성성분 — 혼합제품 하나에 성분이 여러 개면 성분별로 한 줄"
        ws["A1"].fill = TITLE_FILL
        ws["A1"].font = Font(color="FFFFFF", bold=True, size=13)
        ws.merge_cells("A2:J2")
        ws["A2"] = (
            "※ 제품목록행번호로 02_화학물질목록의 부모 제품과 연결합니다. 성분별 법정 함량기준을 판단하되 적용되는 성분의 최대보유량은 '성분함량×제품량'이 아니라 혼합물 전체 총량입니다. "
            "SDS가 함량범위로 제시되면 최저·최고를 입력하며 법정기준을 가로지르면 판정보류합니다."
        )
        ws["A2"].fill = NOTE_FILL
        ws["A2"].alignment = Alignment(wrap_text=True)
        for col, header in enumerate(COMPONENT_COLUMNS, start=1):
            ws.cell(3, col, header)
            ws.cell(3, col).fill = HEADER_FILL
            ws.cell(3, col).font = WHITE_FONT
            ws.cell(3, col).alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        examples = [
            ["해당", 5, "염산 35%", "염화수소", "7647-01-0", 35, None, None, "회사 제품 SDS 제3항", ""],
            ["해당", 6, "과산화수소 35%", "과산화수소", "7722-84-1", 35, None, None, "회사 제품 SDS 제3항", ""],
            ["해당", 8, "이소프로필알코올 수용액 70%", "2-프로판올", "67-63-0", 70, None, None, "회사 제품 SDS 제3항", ""],
            ["해당", 10, "질산 68%", "질산", "7697-37-2", 68, None, None, "회사 제품 SDS 제3항", ""],
            ["해당", 11, "수산화나트륨 수용액 30%", "수산화나트륨", "1310-73-2", 30, None, None, "회사 제품 SDS 제3항", ""],
            ["해당없음", 999, "다성분 혼합제품 형식예시(업로드 제외)", "톨루엔", "108-88-3", 90, None, None, "회사 제품 SDS 제3항", "같은 제품번호에 성분별로 여러 줄 작성"],
            ["해당없음", 999, "다성분 혼합제품 형식예시(업로드 제외)", "메탄올", "67-56-1", 8, None, None, "회사 제품 SDS 제3항", "같은 제품번호에 성분별로 여러 줄 작성"],
        ]
        for row_no, values in enumerate(examples, start=4):
            for col, value in enumerate(values, start=1):
                ws.cell(row_no, col, value)
                ws.cell(row_no, col).fill = EXAMPLE_FILL
                ws.cell(row_no, col).alignment = Alignment(vertical="top", wrap_text=True)
        for row in range(11, 104):
            ws.cell(row, 1, "해당없음")
            for col in range(1, len(COMPONENT_COLUMNS) + 1):
                ws.cell(row, col).fill = INPUT_FILL
        dv_active = DataValidation(type="list", formula1='"해당,해당없음,모름"', allow_blank=True)
        ws.add_data_validation(dv_active)
        dv_active.add("A4:A103")
        for idx, width in enumerate([12, 16, 28, 24, 16, 12, 14, 14, 38, 40], start=1):
            ws.column_dimensions[ws.cell(3, idx).column_letter].width = width
        ws.freeze_panes = "A4"

        if "04_시설별최대보유량" in wb.sheetnames:
            facilities = wb["04_시설별최대보유량"]
            component_col = facilities.max_column + 1
            facilities.cell(3, component_col, FACILITY_COMPONENT_CAS_COLUMN)
            facilities.cell(3, component_col).fill = HEADER_FILL
            facilities.cell(3, component_col).font = WHITE_FONT
            facilities.cell(3, component_col).alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            for row in range(4, 31):
                facilities.cell(row, component_col).fill = EXAMPLE_FILL if row <= 11 else INPUT_FILL
            facilities.column_dimensions[facilities.cell(3, component_col).column_letter].width = 24
            for merged in list(facilities.merged_cells.ranges):
                if str(merged) in {"A1:V1", "A2:V2"}:
                    facilities.unmerge_cells(str(merged))
            facilities.merge_cells(start_row=1, start_column=1, end_row=1, end_column=component_col)
            facilities.merge_cells(start_row=2, start_column=1, end_row=2, end_column=component_col)
            facilities["A2"] = (
                "※ 혼합물 저장·변화없음 공정은 02A 성분함량을 자동 연결할 수 있습니다. 하나의 혼합제품에 규제성분이 2개 이상이고 단순혼합·반응으로 성분함량이 달라지면 시설행을 성분별로 복사하여 "
                f"'{FACILITY_COMPONENT_CAS_COLUMN}'와 별표4 기준함량·근거를 각각 작성하세요."
            )
            facilities["A2"].fill = NOTE_FILL
            facilities["A2"].alignment = Alignment(wrap_text=True)

        out = BytesIO()
        wb.save(out)
        return _cleanup_company_workbook(out.getvalue())

    wrapped._planner_cap_mixture_template = True
    return wrapped


def _resolved_pct(component: pd.Series, hits: Iterable[dict[str, Any]]) -> float | None:
    values, mode = _concentration_values(component)
    if not values:
        return None
    if mode == "exact":
        return values[0]
    low, high = values
    thresholds = [_num(hit.get("content_threshold_pct")) for hit in hits]
    thresholds = [value for value in thresholds if value is not None]
    if not thresholds or all(low >= threshold for threshold in thresholds):
        return low
    return high


def _prepare_mixture_facilities(intake: Any, parent: int, component: pd.Series, hits: list[dict[str, Any]]):
    facilities = getattr(intake, "facilities", pd.DataFrame())
    if facilities is None or facilities.empty or "목록행번호" not in facilities.columns:
        return pd.DataFrame(), []
    frame = facilities[facilities["목록행번호"].map(_int).eq(parent)].copy()
    component_cas = _clean(component.get("CAS No."))
    component_count = len(_components_for_parent(intake, parent))
    pct = _resolved_pct(component, hits)
    evidence = _clean(component.get("SDS 제3항 근거"))
    selected: list[pd.Series] = []
    blockers: list[str] = []
    for _, row in frame.iterrows():
        selector = _clean(row.get(FACILITY_COMPONENT_CAS_COLUMN)) if FACILITY_COMPONENT_CAS_COLUMN in row.index else ""
        process = _clean(row.get("공정유형"))
        if selector and selector != component_cas:
            continue
        if not selector and component_count > 1 and process in {"단순혼합", "반응"}:
            facility = _clean(row.get("시설명")) or "시설"
            blockers.append(
                f"혼합물 제품 {parent}행 / {facility}: 규제성분이 여러 개인 {process} 공정은 시설행을 성분별로 나누고 '{FACILITY_COMPONENT_CAS_COLUMN}'에 {component_cas}를 지정한 뒤 별표4 기준함량과 근거를 작성해 주세요."
            )
            continue
        copy = row.copy()
        if process not in {"단순혼합", "반응"}:
            if pct is not None:
                copy["별표4 기준함량(%)"] = pct
            if not _clean(copy.get("함량근거")) and evidence:
                copy["함량근거"] = f"02A 혼합물 구성성분 / {evidence}"
        selected.append(copy)
    return (pd.DataFrame(selected).reset_index(drop=True) if selected else frame.iloc[0:0].copy()), list(dict.fromkeys(blockers))


def install_cap_mixture_runtime() -> None:
    from . import template as template_module
    from . import inventory as inventory_module
    from . import cap_engine as cap_engine_module
    from . import cap_scope_engine as cap_scope_module
    from . import cap_holding_screen as screen_module
    from . import cap_holding as holding_module
    from . import cap_quick_holding as quick_module

    if getattr(inventory_module, "_cap_mixture_runtime_installed", False):
        return

    template_module.build_minimal_input_workbook = _patch_template(template_module.build_minimal_input_workbook)

    # inventory.read_intake_workbook/validate_intake already read and validate
    # mixture_components directly (it is a real IntakeData field); only the
    # downstream CAP screening/scope/holding functions still need patching.

    original_screen = screen_module.screen_facility_stage
    original_scope = cap_scope_module.assess_cap_scope
    original_assess_cap = cap_engine_module.assess_cap
    original_holding = holding_module.assess_cap_holding
    original_quick = quick_module.compare_confirmed_declared_holding

    def screen_with_components(intake):
        mixtures = _mixture_rows(intake)
        if not mixtures:
            return original_screen(intake)
        chemicals = intake.chemicals.copy()
        for parent in mixtures:
            if 0 < parent <= len(chemicals):
                chemicals.at[parent - 1, "CAS No."] = ""
        base = original_screen(_clone_intake(intake, chemicals=chemicals))
        if not base.ready:
            return base
        hits = list(base.legal_hits)
        blockers = list(base.blockers)
        row_numbers = set(base.row_numbers)
        count = 0
        for parent in sorted(mixtures):
            if parent <= 0 or parent > len(intake.chemicals):
                continue
            for _, component in _components_for_parent(intake, parent).iterrows():
                count += 1
                comp_hits, comp_blockers, claimed, ready = _screen_component_direct(original_screen, intake, parent, component)
                if not ready:
                    return screen_module.FacilityStageScreen(legal_hits=hits, row_numbers=sorted(row_numbers), blockers=list(dict.fromkeys(blockers + comp_blockers)), messages=list(base.messages), ready=False)
                hits.extend(comp_hits)
                blockers.extend(comp_blockers)
                if claimed:
                    row_numbers.add(parent)
        unique_hits: list[dict[str, Any]] = []
        seen = set()
        for hit in hits:
            key = (hit.get("source_key"), hit.get("row_no"), hit.get("cas"), hit.get("item_no"), hit.get("designation_id"), hit.get("variant_type"), hit.get("hazard_category"), hit.get("content_threshold_pct"), hit.get("lower_quantity_ton"), hit.get("upper_quantity_ton"))
            if key not in seen:
                seen.add(key)
                unique_hits.append(hit)
        messages = list(base.messages)
        if count:
            messages.append(f"혼합제품의 SDS 제3항 구성성분 {count}건을 성분별로 별표 3→별표 2 순서로 확인했습니다.")
        return screen_module.FacilityStageScreen(legal_hits=unique_hits, row_numbers=sorted(row_numbers), blockers=list(dict.fromkeys(str(v) for v in blockers if str(v).strip())), messages=messages, ready=True)

    screen_module.screen_facility_stage = screen_with_components

    def component_scope(intake):
        direct: list[dict[str, Any]] = []
        candidates: list[dict[str, Any]] = []
        for parent in sorted(_mixture_rows(intake)):
            if parent <= 0 or parent > len(intake.chemicals):
                continue
            product = _clean(intake.chemicals.iloc[parent - 1].get("제품명"))
            for _, component in _components_for_parent(intake, parent).iterrows():
                values, _ = _concentration_values(component)
                if not values:
                    continue
                result = original_scope(_single_component_intake(intake, parent, component, values[-1]))
                direct.extend(_map_records(result.direct_hits, parent, product, component))
                candidates.extend(_map_records(result.candidate_rows, parent, product, component))
        return direct, candidates

    def assess_cap_with_components(intake):
        mixtures = _mixture_rows(intake)
        if not mixtures:
            return original_assess_cap(intake)
        chemicals = intake.chemicals.copy()
        for parent in mixtures:
            if 0 < parent <= len(chemicals):
                chemicals.at[parent - 1, "CAS No."] = ""
                chemicals.at[parent - 1, "물질명(알면 입력)"] = ""
        base = original_assess_cap(_clone_intake(intake, chemicals=chemicals))
        base.blockers = [str(b) for b in base.blockers if not any(f"{parent}행" in str(b) and "CAS" in str(b) for parent in mixtures)]
        base.scope_direct_hits = [r for r in base.scope_direct_hits if _int(r.get("row_no")) not in mixtures]
        base.scope_candidates = [r for r in base.scope_candidates if _int(r.get("row_no")) not in mixtures]
        base.app1_required_rows = [r for r in base.app1_required_rows if _int(r.get("row_no")) not in mixtures]
        direct_scope, candidates = component_scope(intake)
        base.scope_direct_hits.extend(direct_scope)
        base.scope_candidates.extend(candidates)
        base.scope_review_required = bool(base.scope_candidates)
        direct_screen = screen_with_components(intake)
        claimed_rows = {int(v) for v in direct_screen.row_numbers if int(v) in mixtures}
        candidate_rows = {_int(v.get("row_no")) for v in candidates}
        for parent in sorted(mixtures):
            if parent in claimed_rows or parent in candidate_rows or parent > len(intake.chemicals):
                continue
            row = intake.chemicals.iloc[parent - 1]
            base.app1_required_rows.append({"row_no": parent, "product_name": _clean(row.get("제품명")), "cas": _clean(row.get("CAS No."))})
        base.messages.append("혼합제품은 02A 구성성분으로 별표 3·2를 먼저 확인하고, 해당 성분이 없을 때 부모 제품의 SDS 제2항을 별표 1 fallback으로 확인합니다.")
        return base

    cap_engine_module.assess_cap = assess_cap_with_components

    def holding_with_components(intake, facilities, legal_hits, required_row_numbers):
        mixtures = _mixture_rows(intake)
        legal_hits = list(legal_hits)
        required = sorted(set(int(v) for v in required_row_numbers if int(v) > 0))
        if not any(row in mixtures for row in required):
            return original_holding(intake, facilities, legal_hits, required)
        comparisons: list[dict[str, Any]] = []
        facility_rows: list[Any] = []
        blockers: list[str] = []
        messages: list[str] = []

        nonmix_hits = [hit for hit in legal_hits if _int(hit.get("row_no")) not in mixtures]
        nonmix_required = [row for row in required if row not in mixtures]
        if nonmix_hits or nonmix_required:
            nonmix_facilities = facilities
            if isinstance(facilities, pd.DataFrame) and not facilities.empty and "목록행번호" in facilities.columns:
                nonmix_facilities = facilities[~facilities["목록행번호"].map(_int).isin(mixtures)].copy().reset_index(drop=True)
            result = original_holding(intake, nonmix_facilities, nonmix_hits, nonmix_required)
            if result.status == "DB_NOT_READY":
                return result
            comparisons.extend(result.comparison_rows)
            facility_rows.extend(result.facility_rows)
            blockers.extend(result.blockers)
            messages.extend(result.messages)

        groups: dict[tuple[int, str], list[dict[str, Any]]] = {}
        for hit in legal_hits:
            parent = _int(hit.get("row_no"))
            cas = _clean(hit.get("cas"))
            if parent in mixtures and cas:
                groups.setdefault((parent, cas), []).append(dict(hit))
        for (parent, cas), hits in groups.items():
            comps = _components_for_parent(intake, parent)
            matched = comps[comps["CAS No."].map(_clean).eq(cas)] if not comps.empty else pd.DataFrame()
            if matched.empty:
                blockers.append(f"혼합물 제품 {parent}행: 법적 규칙 CAS {cas}와 02A 구성성분을 연결하지 못했습니다.")
                continue
            component = matched.iloc[0]
            pct = _resolved_pct(component, hits)
            if pct is None:
                blockers.append(f"혼합물 제품 {parent}행 / {cas}: 법정 판정용 구성성분 함량을 확정하지 못했습니다.")
                continue
            synthetic = intake.chemicals.copy()
            synthetic.at[parent - 1, "CAS No."] = cas
            synthetic.at[parent - 1, "물질명(알면 입력)"] = _clean(component.get("구성성분명"))
            synthetic.at[parent - 1, "함량(%)"] = pct
            comp_facilities, prep_blockers = _prepare_mixture_facilities(intake, parent, component, hits)
            blockers.extend(prep_blockers)
            result = original_holding(_clone_intake(intake, chemicals=synthetic, facilities=comp_facilities), comp_facilities, hits, [parent])
            if result.status == "DB_NOT_READY":
                return result
            comparisons.extend(result.comparison_rows)
            facility_rows.extend(result.facility_rows)
            blockers.extend(result.blockers)
            messages.extend(result.messages)

        blockers = list(dict.fromkeys(str(v) for v in blockers if str(v).strip()))
        if blockers:
            return holding_module.CAPHoldingResult(status="HOLD", label="화사계 최대보유량 판정보류", app4_ready=True, facility_rows=facility_rows, comparison_rows=comparisons, blockers=blockers, messages=list(dict.fromkeys(messages + ["혼합물은 규제성분별 적용 여부를 판단하되 최대보유량에는 혼합물 전체 질량을 사용합니다."])))
        if any(row.get("quantity_band") == "상위 규정수량 이상" for row in comparisons):
            status, label = "UPPER_CANDIDATE", "최대보유량이 상위 규정수량 이상"
        elif any(row.get("quantity_band") == "하위 이상·상위 미만" for row in comparisons):
            status, label = "LOWER_CANDIDATE", "최대보유량이 하위 규정수량 이상·상위 규정수량 미만"
        else:
            status, label = "BELOW_LOWER", "확인된 유해화학물질의 최대보유량이 하위 규정수량 미만"
        return holding_module.CAPHoldingResult(status=status, label=label, app4_ready=True, facility_rows=facility_rows, comparison_rows=comparisons, blockers=[], messages=list(dict.fromkeys(messages + ["혼합물 규제성분별로 동일한 혼합물 전체 최대보유량을 각각 규정수량과 비교했습니다."])))

    holding_module.assess_cap_holding = holding_with_components

    def quick_with_components(intake, legal_hits):
        result = original_quick(intake, legal_hits)
        mixtures = _mixture_rows(intake)
        if not mixtures:
            return result
        # The quantity calculation is already the parent's whole mixture mass.
        # Restore the component identity in reviewer-facing rows.
        by_key: dict[tuple[int, str, str], list[dict[str, Any]]] = {}
        for hit in legal_hits:
            parent = _int(hit.get("row_no"))
            if parent in mixtures:
                by_key.setdefault((parent or 0, _clean(hit.get("source_key")), _clean(hit.get("item_no"))), []).append(hit)
        used: dict[tuple[int, str, str], int] = {}
        for row in result.comparison_rows:
            parent = _int(row.get("row_no"))
            if parent not in mixtures:
                continue
            key = (parent or 0, _clean(row.get("source_key")), _clean(row.get("item_no")))
            options = by_key.get(key, [])
            index = used.get(key, 0)
            if options:
                hit = options[min(index, len(options) - 1)]
                row["cas"] = _clean(hit.get("cas"))
                row["legal_substance"] = _clean(hit.get("legal_substance"))
                row["component_name"] = _clean(hit.get("component_name"))
                row["basis"] = "회사 확인 혼합물 전체 최대동시보유량 / 규제성분별 규정수량 비교"
                used[key] = index + 1
        return result

    quick_module.compare_confirmed_declared_holding = quick_with_components
    inventory_module._cap_mixture_runtime_installed = True
