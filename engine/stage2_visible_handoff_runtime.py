from __future__ import annotations

"""Visible Stage-1 handoff and mixture-aware Stage-2 MSDS validation.

Goals
-----
1. Keep Stage 1 and Stage 2 workbooks separate, but make every carried company
   fact visible in the Stage 2 workbook with an explicit source and correction
   path.
2. Never silently overwrite a Stage-1-confirmed fact from a Stage 2 workbook.
3. Carry Stage-1 mixture component rows into Stage 2 so every component CAS can
   participate in KOSHA reference lookup and supplier-MSDS comparison.
4. For mixtures, do not treat global CAS presence as sufficient.  Where the
   supplier MSDS can be associated with the product, inspect Section 3 and
   compare component CAS/concentration.  Ambiguous/unreadable composition stays
   review-required instead of being promoted to a false-safe success.

The module is installed as a compatibility runtime so existing saved projects
and Stage-2 workbook schema remain readable.
"""

from io import BytesIO
from pathlib import Path
import re
from typing import Any, Mapping

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


HANDOFF_SHEET = "00A_Stage1승계정보"
COMPONENT_SHEET = "02A_혼합물구성성분"
META_SHEET = "_시스템정보"
STAGE1_FILL = "D9EAF7"
STAGE2_FILL = "FFF2CC"
HEADER_FILL = "5B9BD5"
NOTE_FILL = "F2F2F2"
PROVENANCE_HEADERS = ("정보구분", "출처/수정방법")


def _text(value: object) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _norm(value: object) -> str:
    return re.sub(r"\s+", "", _text(value)).lower()


def _is_stage1_record(record: Any) -> bool:
    if record is None:
        return False
    return any(_text(getattr(ref, "source_type", "")) == "STAGE1_WORKBOOK" for ref in getattr(record, "evidence", []) or [])


def _records_equal(left: Any, right: Any) -> bool:
    if isinstance(left, (int, float)) and isinstance(right, (int, float)):
        try:
            return abs(float(left) - float(right)) <= 1e-9
        except Exception:
            pass
    return _norm(left) == _norm(right)


def _stage1_evidence(project: Any) -> list[Any]:
    source = _text(getattr(project, "stage1_source_fingerprint", ""))
    if not source:
        return []
    from engine.stage2.project import EvidenceRef

    return [
        EvidenceRef(
            source_type="STAGE1_WORKBOOK",
            source_name="회사 입력 Excel",
            sha256=source,
            note="Stage 1 판정진단에서 회사가 입력·확인한 사실",
        )
    ]


def _value_summary(key: str, value: Any) -> str:
    if isinstance(value, list):
        if key == "inventory.chemicals":
            return f"{len(value)}개 행 — 02_화학물질정보에서 상세 확인"
        if key == "inventory.facilities":
            return f"{len(value)}개 행 — 03_설비정보에서 상세 확인"
        if key == "inventory.mixture_components":
            return f"{len(value)}개 행 — 02A_혼합물구성성분에서 상세 확인"
        return f"{len(value)}개 항목"
    if isinstance(value, Mapping):
        return f"{len(value)}개 항목"
    return _text(value)


def _style_header(ws, row: int, max_col: int) -> None:
    for col in range(1, max_col + 1):
        cell = ws.cell(row, col)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=HEADER_FILL)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _write_handoff_sheet(wb: Any, project: Any) -> None:
    if HANDOFF_SHEET in wb.sheetnames:
        del wb[HANDOFF_SHEET]
    ws = wb.create_sheet(HANDOFF_SHEET, 1)
    ws["A1"] = "Stage 1 → Stage 2 승계정보"
    ws["A1"].font = Font(size=14, bold=True, color="FFFFFF")
    ws["A1"].fill = PatternFill("solid", fgColor="1F4E78")
    ws.merge_cells("A1:E1")
    ws["A2"] = (
        "이 시트는 Stage 1에서 Stage 2로 전달된 내용을 숨기지 않고 보여주는 확인용 시트입니다. "
        "'Stage 1 회사확정 승계값'을 바꾸려면 이 파일에서 수정하지 말고 Stage 1 회사 입력파일을 수정하여 판정진단을 다시 수행하세요. "
        "Stage 1 시스템 판정값은 참고용이며 회사 입력값으로 취급하지 않습니다."
    )
    ws["A2"].fill = PatternFill("solid", fgColor=NOTE_FILL)
    ws["A2"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells("A2:E2")
    headers = ["정보구분", "항목", "승계값", "출처", "수정방법"]
    for col, value in enumerate(headers, 1):
        ws.cell(4, col, value)
    _style_header(ws, 4, 5)

    row_no = 5
    decision = dict(getattr(project, "stage1_snapshot", {}).get("decision") or {})
    for label, key in (("공정안전보고서 판정", "psm_status"), ("화학사고예방관리계획서 판정", "cap_status")):
        value = _text(decision.get(key))
        if not value:
            continue
        ws.cell(row_no, 1, "Stage 1 시스템 판정(참고)")
        ws.cell(row_no, 2, label)
        ws.cell(row_no, 3, value)
        ws.cell(row_no, 4, "Stage 1 규칙엔진")
        ws.cell(row_no, 5, "이 파일에서 수정 불가 · 판정 변경이 필요하면 Stage 1 재수행")
        row_no += 1

    rows = []
    for key, record in sorted(getattr(project, "fields", {}).items()):
        if not _is_stage1_record(record):
            continue
        rows.append((key, record))
    for key, record in rows:
        ws.cell(row_no, 1, "Stage 1 회사확정 승계값")
        ws.cell(row_no, 2, _text(getattr(record, "label", "")) or key)
        ws.cell(row_no, 3, _value_summary(key, getattr(record, "value", None)))
        ws.cell(row_no, 4, "Stage 1 회사 입력 Excel")
        ws.cell(row_no, 5, "변경 시 Stage 1 회사 입력파일 수정 후 판정진단 재수행")
        row_no += 1

    if row_no == 5:
        ws.cell(5, 1, "승계값 없음")
        ws.cell(5, 2, "Stage 2 직접 시작 프로젝트")
        ws.cell(5, 3, "Stage 1에서 승계된 회사확정값이 없습니다.")
        ws.cell(5, 4, "-")
        ws.cell(5, 5, "Stage 2 통합 작성자료에서 직접 입력")
        row_no = 6

    for row in ws.iter_rows(min_row=5, max_row=row_no - 1, min_col=1, max_col=5):
        for cell in row:
            cell.fill = PatternFill("solid", fgColor=STAGE1_FILL)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    widths = {"A": 26, "B": 34, "C": 62, "D": 26, "E": 56}
    for col, width in widths.items():
        ws.column_dimensions[col].width = width
    ws.freeze_panes = "A5"


def _mixture_rows(project: Any) -> list[dict[str, Any]]:
    record = project.get_field("inventory.mixture_components")
    if record and isinstance(record.value, list):
        return [dict(row) for row in record.value if isinstance(row, Mapping)]
    raw = list(getattr(project, "stage1_snapshot", {}).get("mixture_components") or [])
    return [dict(row) for row in raw if isinstance(row, Mapping)]


def _write_mixture_sheet(wb: Any, project: Any) -> None:
    components = _mixture_rows(project)
    if not components:
        return
    if COMPONENT_SHEET in wb.sheetnames:
        del wb[COMPONENT_SHEET]
    insert_at = wb.sheetnames.index("02_화학물질정보") + 1 if "02_화학물질정보" in wb.sheetnames else 2
    ws = wb.create_sheet(COMPONENT_SHEET, insert_at)
    ws["A1"] = "혼합물 구성성분 — Stage 1 승계정보"
    ws["A1"].font = Font(size=14, bold=True, color="FFFFFF")
    ws["A1"].fill = PatternFill("solid", fgColor="1F4E78")
    ws.merge_cells("A1:L1")
    ws["A2"] = (
        "Stage 1에서 회사 제품 SDS 제3항을 근거로 입력한 혼합물 구성성분입니다. "
        "Stage 2에서는 이 값을 다시 입력하지 않습니다. 변경이 필요하면 Stage 1을 다시 수행하세요. "
        "이 구성성분 CAS는 KOSHA 참고조회와 공급자 MSDS 제3항 대조에도 사용됩니다."
    )
    ws["A2"].fill = PatternFill("solid", fgColor=NOTE_FILL)
    ws["A2"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells("A2:L2")
    headers = [
        "적용여부", "제품목록행번호", "제품명(확인용)", "구성성분명", "CAS No.",
        "함량(%)", "함량 최저(%)", "함량 최고(%)", "SDS 제3항 근거", "비고",
        "정보구분", "출처/수정방법",
    ]
    for col, value in enumerate(headers, 1):
        ws.cell(4, col, value)
    _style_header(ws, 4, len(headers))
    for r_idx, item in enumerate(components, 5):
        values = [
            item.get("적용여부", "해당"), item.get("제품목록행번호", ""), item.get("제품명(확인용)", ""),
            item.get("구성성분명", ""), item.get("CAS No.", ""), item.get("함량(%)", ""),
            item.get("함량 최저(%)", ""), item.get("함량 최고(%)", ""), item.get("SDS 제3항 근거", ""),
            item.get("비고", ""), "Stage 1 승계값", "회사 입력 Excel / 변경 시 Stage 1 재판정",
        ]
        for c_idx, value in enumerate(values, 1):
            ws.cell(r_idx, c_idx, value)
            ws.cell(r_idx, c_idx).fill = PatternFill("solid", fgColor=STAGE1_FILL)
            ws.cell(r_idx, c_idx).alignment = Alignment(vertical="top", wrap_text=True)
    widths = [12, 16, 28, 24, 16, 12, 14, 14, 30, 30, 20, 46]
    for idx, width in enumerate(widths, 1):
        ws.column_dimensions[chr(64 + idx)].width = width
    ws.freeze_panes = "A5"

    if META_SHEET in wb.sheetnames:
        meta = wb[META_SHEET]
        meta.append(["PROTECTED_TABLE", COMPONENT_SHEET, "inventory.mixture_components", 4, "", "Stage 1 판정자료"])


def _append_source_columns(ws: Any, *, field_rows: dict[int, bool]) -> None:
    if not field_rows:
        return
    # FORM sheets use four visible columns in the base workbook.
    info_col = max(ws.max_column, 4) + 1
    source_col = info_col + 1
    ws.cell(4, info_col, PROVENANCE_HEADERS[0])
    ws.cell(4, source_col, PROVENANCE_HEADERS[1])
    _style_header(ws, 4, source_col)
    for row_no, is_stage1 in field_rows.items():
        if is_stage1:
            ws.cell(row_no, info_col, "Stage 1 승계값")
            ws.cell(row_no, source_col, "회사 입력 Excel / 변경 시 Stage 1 재판정")
            ws.cell(row_no, 2).fill = PatternFill("solid", fgColor=STAGE1_FILL)
        else:
            ws.cell(row_no, info_col, "Stage 2 신규입력")
            ws.cell(row_no, source_col, "이 통합 작성자료에서 작성·수정")
            ws.cell(row_no, 2).fill = PatternFill("solid", fgColor=STAGE2_FILL)
        for col in (info_col, source_col):
            ws.cell(row_no, col).alignment = Alignment(vertical="top", wrap_text=True)
    ws.column_dimensions[chr(64 + info_col)].width = 20
    ws.column_dimensions[chr(64 + source_col)].width = 46


def _annotate_form_provenance(wb: Any, project: Any) -> None:
    if META_SHEET not in wb.sheetnames:
        return
    meta = wb[META_SHEET]
    by_sheet: dict[str, dict[int, bool]] = {}
    for row_no in range(8, meta.max_row + 1):
        kind = _text(meta.cell(row_no, 1).value)
        if kind not in {"FORM", "PROTECTED"}:
            continue
        sheet = _text(meta.cell(row_no, 2).value)
        field_key = _text(meta.cell(row_no, 3).value).split("|")[0]
        target_row = int(meta.cell(row_no, 4).value or 0)
        record = project.get_field(field_key) if field_key else None
        stage1 = _is_stage1_record(record)
        by_sheet.setdefault(sheet, {})[target_row] = stage1
        if stage1:
            meta.cell(row_no, 1, "PROTECTED")
            meta.cell(row_no, 6, "Stage 1 판정자료 · 변경 시 Stage 1 재판정")
    for sheet, rows in by_sheet.items():
        if sheet in wb.sheetnames:
            _append_source_columns(wb[sheet], field_rows=rows)


def _table_prefill_rows(project: Any, sheet: str) -> list[list[Any]]:
    from engine.stage2 import integrated_workbook as integrated
    if sheet == "02_화학물질정보":
        return integrated._prefill_chemicals(project)
    if sheet == "03_설비정보":
        return integrated._prefill_facilities(project)
    return []


def _annotate_table_provenance(wb: Any, project: Any) -> None:
    for sheet in ("02_화학물질정보", "03_설비정보"):
        if sheet not in wb.sheetnames:
            continue
        expected = _table_prefill_rows(project, sheet)
        if not expected:
            continue
        ws = wb[sheet]
        last_header = max((cell.column for cell in ws[4] if _text(cell.value)), default=ws.max_column)
        info_col = last_header + 1
        source_col = last_header + 2
        ws.cell(4, info_col, PROVENANCE_HEADERS[0])
        ws.cell(4, source_col, PROVENANCE_HEADERS[1])
        _style_header(ws, 4, source_col)
        for idx in range(len(expected)):
            row_no = 5 + idx
            ws.cell(row_no, info_col, "Stage 1 승계값")
            ws.cell(row_no, source_col, "회사 입력 Excel / 기존값 변경 시 Stage 1 재판정")
            for col in range(1, last_header + 1):
                if ws.cell(row_no, col).value not in (None, ""):
                    ws.cell(row_no, col).fill = PatternFill("solid", fgColor=STAGE1_FILL)
            ws.cell(row_no, info_col).fill = PatternFill("solid", fgColor=STAGE1_FILL)
            ws.cell(row_no, source_col).fill = PatternFill("solid", fgColor=STAGE1_FILL)
        ws.column_dimensions[get_column_letter(info_col)].width = 20
        ws.column_dimensions[get_column_letter(source_col)].width = 48


def _decorate_workbook(project: Any, raw: bytes, *, example: bool) -> bytes:
    if example:
        return raw
    wb = load_workbook(BytesIO(raw))
    _write_handoff_sheet(wb, project)
    _write_mixture_sheet(wb, project)
    _annotate_form_provenance(wb, project)
    _annotate_table_provenance(wb, project)

    if "00_작성안내" in wb.sheetnames:
        ws = wb["00_작성안내"]
        row = ws.max_row + 1
        ws.cell(row, 1, "Stage 1 승계표시")
        ws.cell(row, 2, "하늘색=Stage 1 승계값, 노란색=Stage 2 신규입력. 승계값은 00A_Stage1승계정보에서 출처를 확인할 수 있습니다.")
        row += 1
        ws.cell(row, 1, "승계값 수정")
        ws.cell(row, 2, "Stage 1 승계값과 다른 내용을 Stage 2 파일에 입력해도 자동 덮어쓰지 않습니다. 변경이 필요하면 Stage 1 판정진단을 다시 수행합니다.")
        for r in range(ws.max_row - 1, ws.max_row + 1):
            ws.cell(r, 2).alignment = Alignment(wrap_text=True, vertical="top")

    out = BytesIO()
    wb.save(out)
    return out.getvalue()


def _canonical_rows(rows: list[dict[str, Any]], headers: list[str]) -> list[tuple[str, ...]]:
    output: list[tuple[str, ...]] = []
    for row in rows:
        output.append(tuple(_norm(row.get(header)) for header in headers))
    return output


def _parse_sheet_rows(ws: Any, header_row: int, headers: list[str]) -> list[dict[str, Any]]:
    names = [_text(cell.value) for cell in ws[header_row]]
    index = {name: idx + 1 for idx, name in enumerate(names) if name}
    rows: list[dict[str, Any]] = []
    for row_no in range(header_row + 1, ws.max_row + 1):
        item = {header: ws.cell(row_no, index[header]).value if header in index else "" for header in headers}
        if any(_text(value) for value in item.values()):
            rows.append(item)
    return rows


def _sanitize_stage2_workbook(project: Any, raw: bytes) -> tuple[bytes, list[str]]:
    from engine.stage2 import integrated_workbook as integrated

    wb = load_workbook(BytesIO(raw), data_only=False)
    warnings: list[str] = []

    # Any field that came from Stage 1 remains immutable in Stage 2, even when
    # an older downloaded workbook still marked it as FORM.
    if META_SHEET in wb.sheetnames:
        meta = wb[META_SHEET]
        for row_no in range(8, meta.max_row + 1):
            kind = _text(meta.cell(row_no, 1).value)
            if kind not in {"FORM", "PROTECTED"}:
                continue
            sheet = _text(meta.cell(row_no, 2).value)
            key = _text(meta.cell(row_no, 3).value).split("|")[0]
            target_row = int(meta.cell(row_no, 4).value or 0)
            target_col = int(meta.cell(row_no, 5).value or 0)
            record = project.get_field(key) if key else None
            if not _is_stage1_record(record) or sheet not in wb.sheetnames or not target_row or not target_col:
                continue
            cell = wb[sheet].cell(target_row, target_col)
            if cell.value not in (None, "") and not _records_equal(cell.value, record.value):
                warnings.append(f"{getattr(record, 'label', key)}은(는) Stage 1 승계값과 달라 덮어쓰지 않았습니다. 변경하려면 Stage 1 판정진단을 다시 수행해 주세요.")
            cell.value = record.value
            meta.cell(row_no, 1, "PROTECTED")

    # Preserve non-empty Stage-1-prefilled table facts while still allowing the
    # company to fill previously blank Stage-2-only columns in those rows.
    for sheet in ("02_화학물질정보", "03_설비정보"):
        if sheet not in wb.sheetnames:
            continue
        expected = _table_prefill_rows(project, sheet)
        if not expected:
            continue
        ws = wb[sheet]
        headers = [_text(cell.value) for cell in ws[4]]
        protected_cols = [idx + 1 for idx, name in enumerate(headers) if name and name not in PROVENANCE_HEADERS]
        for idx, expected_row in enumerate(expected):
            excel_row = 5 + idx
            changed = False
            for col, expected_value in enumerate(expected_row, 1):
                if col not in protected_cols or expected_value in (None, ""):
                    continue
                current = ws.cell(excel_row, col).value
                if current not in (None, "") and not _records_equal(current, expected_value):
                    changed = True
                ws.cell(excel_row, col, expected_value)
            if changed:
                warnings.append(f"{sheet} {excel_row}행의 Stage 1 승계값 변경은 반영하지 않았습니다. 판정값을 바꾸려면 Stage 1을 다시 수행해 주세요.")

        # Workbook-only provenance columns must not become statutory data fields.
        header_names = [_text(cell.value) for cell in ws[4]]
        delete_cols = [idx + 1 for idx, name in enumerate(header_names) if name in PROVENANCE_HEADERS]
        for col in sorted(delete_cols, reverse=True):
            ws.delete_cols(col, 1)

    # The mixture table is display/provenance only in Stage 2.  Detect edits but
    # never promote them over the Stage-1-confirmed component list.
    expected_components = _mixture_rows(project)
    if expected_components:
        headers = ["적용여부", "제품목록행번호", "제품명(확인용)", "구성성분명", "CAS No.", "함량(%)", "함량 최저(%)", "함량 최고(%)", "SDS 제3항 근거", "비고"]
        if COMPONENT_SHEET not in wb.sheetnames:
            warnings.append("Stage 1 혼합물 구성성분 확인시트가 통합 작성자료에서 누락되었습니다. 내부 승계값은 유지하며 최신 Stage 2 파일을 다시 내려받는 것을 권장합니다.")
        else:
            actual = _parse_sheet_rows(wb[COMPONENT_SHEET], 4, headers)
            if _canonical_rows(actual, headers) != _canonical_rows(expected_components, headers):
                warnings.append("02A_혼합물구성성분의 Stage 1 승계값 변경은 반영하지 않았습니다. 혼합물 조성이 바뀌었다면 Stage 1 판정진단을 다시 수행해 주세요.")

    out = BytesIO()
    wb.save(out)
    return out.getvalue(), list(dict.fromkeys(warnings))


def _install_workbook_handoff() -> None:
    from engine.stage2 import integrated_workbook as integrated
    from engine.stage2 import workbook_enhancements as enhancements

    original_build = integrated.build_integrated_authoring_workbook
    if not getattr(original_build, "_planner_visible_handoff_wrapped", False):
        def build_with_visible_handoff(project, *, example=False):
            return _decorate_workbook(project, original_build(project, example=example), example=example)

        build_with_visible_handoff._planner_visible_handoff_wrapped = True
        integrated.build_integrated_authoring_workbook = build_with_visible_handoff
        enhancements._base_build = build_with_visible_handoff

    original_apply = integrated.apply_integrated_authoring_workbook
    if not getattr(original_apply, "_planner_visible_handoff_wrapped", False):
        def apply_with_visible_handoff(project, workbook_bytes, *, workbook_evidence=None):
            sanitized, extra = _sanitize_stage2_workbook(project, workbook_bytes)
            result = original_apply(project, sanitized, workbook_evidence=workbook_evidence)
            return integrated.IntegratedImportResult(
                updated_fields=result.updated_fields,
                table_fields=result.table_fields,
                attachment_declarations=result.attachment_declarations,
                warnings=tuple(extra) + tuple(result.warnings),
            )

        apply_with_visible_handoff._planner_visible_handoff_wrapped = True
        integrated.apply_integrated_authoring_workbook = apply_with_visible_handoff


def _pick_mapping(row: Mapping[str, Any], *aliases: str) -> str:
    normalized = {re.sub(r"[^0-9A-Za-z가-힣]", "", _text(k)).lower(): v for k, v in row.items()}
    for alias in aliases:
        key = re.sub(r"[^0-9A-Za-z가-힣]", "", alias).lower()
        value = normalized.get(key)
        if value not in (None, ""):
            return _text(value)
    return ""


def _parent_product_names(project: Any) -> dict[int, str]:
    record = project.get_field("inventory.chemicals")
    rows = record.value if record and isinstance(record.value, list) else []
    result: dict[int, str] = {}
    for idx, row in enumerate(rows, 1):
        if isinstance(row, Mapping):
            result[idx] = _pick_mapping(row, "제품명", "상품명", "물질명", "화학물질명")
    return result


def _component_expectations(project: Any) -> list[dict[str, Any]]:
    products = _parent_product_names(project)
    output: list[dict[str, Any]] = []
    for row in _mixture_rows(project):
        parent_text = _pick_mapping(row, "제품목록행번호")
        try:
            parent = int(float(parent_text)) if parent_text else 0
        except ValueError:
            parent = 0
        output.append({
            "parent": parent,
            "product": _pick_mapping(row, "제품명(확인용)", "제품명") or products.get(parent, ""),
            "component": _pick_mapping(row, "구성성분명", "성분명", "물질명"),
            "cas": _pick_mapping(row, "CAS No.", "CAS 번호", "CAS"),
            "exact": _pick_mapping(row, "함량(%)", "함량"),
            "low": _pick_mapping(row, "함량 최저(%)", "함량최저"),
            "high": _pick_mapping(row, "함량 최고(%)", "함량최고"),
        })
    return output


def _section3_text(text: str) -> str:
    lines = text.splitlines()
    start = None
    for idx, line in enumerate(lines):
        compact = re.sub(r"\s+", "", line).lower()
        if re.search(r"(^|[^0-9])3[\.\)\-:]?", compact) and any(token in compact for token in ("구성성분", "함유량", "성분의명칭")):
            start = idx
            break
    if start is None:
        return ""
    end = len(lines)
    for idx in range(start + 1, len(lines)):
        compact = re.sub(r"\s+", "", lines[idx]).lower()
        if re.match(r"^4[\.\)\-:]", compact) or (compact.startswith("4.") and "응급" in compact):
            end = idx
            break
    return "\n".join(lines[start:end])


def _pct_on_cas_line(section: str, cas: str) -> tuple[float, float] | None:
    for line in section.splitlines():
        if cas not in line:
            continue
        cleaned = line.replace(cas, " ")
        range_match = re.search(r"(\d+(?:\.\d+)?)\s*(?:~|～|–|—|-)\s*(\d+(?:\.\d+)?)\s*%", cleaned)
        if range_match:
            low = float(range_match.group(1))
            high = float(range_match.group(2))
            return (min(low, high), max(low, high))
        exact = re.search(r"(?<![\d.-])(\d+(?:\.\d+)?)\s*%", cleaned)
        if exact:
            value = float(exact.group(1))
            return (value, value)
    return None


def _expected_pct(item: Mapping[str, Any]) -> tuple[float, float] | None:
    try:
        if _text(item.get("exact")):
            value = float(_text(item.get("exact")).replace(",", ""))
            return value, value
        if _text(item.get("low")) and _text(item.get("high")):
            low = float(_text(item.get("low")).replace(",", ""))
            high = float(_text(item.get("high")).replace(",", ""))
            return min(low, high), max(low, high)
    except ValueError:
        return None
    return None


def _same_pct(expected: tuple[float, float], actual: tuple[float, float]) -> bool:
    return abs(expected[0] - actual[0]) <= 0.05 and abs(expected[1] - actual[1]) <= 0.05


def _comparison_with(base: Any, *, status: str, message: str):
    cls = type(base)
    return cls(
        status=status,
        expected_cas=base.expected_cas,
        found_cas=base.found_cas,
        matched_cas=base.matched_cas,
        missing_cas=base.missing_cas,
        files_checked=base.files_checked,
        unreadable_files=base.unreadable_files,
        message=message,
    )


def _install_msds_mixture_handoff() -> None:
    from engine.stage2 import msds_reference as msds

    original_inventory = msds.inventory_chemicals
    if not getattr(original_inventory, "_planner_mixture_handoff_wrapped", False):
        def inventory_with_components(project):
            base = list(original_inventory(project))
            seen = {item.cas for item in base}
            products = _parent_product_names(project)
            for row in _mixture_rows(project):
                cas_value = _pick_mapping(row, "CAS No.", "CAS 번호", "CAS")
                for cas in msds.extract_cas_numbers(cas_value):
                    if cas in seen:
                        continue
                    parent_text = _pick_mapping(row, "제품목록행번호")
                    try:
                        parent = int(float(parent_text)) if parent_text else 0
                    except ValueError:
                        parent = 0
                    base.append(
                        msds.ChemicalCAS(
                            cas=cas,
                            chemical_name=_pick_mapping(row, "구성성분명", "성분명", "물질명"),
                            product_name=_pick_mapping(row, "제품명(확인용)", "제품명") or products.get(parent, ""),
                        )
                    )
                    seen.add(cas)
            return base

        inventory_with_components._planner_mixture_handoff_wrapped = True
        msds.inventory_chemicals = inventory_with_components

    original_compare = msds.compare_supplier_sds
    if not getattr(original_compare, "_planner_mixture_handoff_wrapped", False):
        def compare_with_composition(project):
            base = original_compare(project)
            expectations = [row for row in _component_expectations(project) if row.get("cas")]
            if not expectations:
                return base

            refs = msds._candidate_sds_evidence(project)
            docs: list[tuple[str, str]] = []
            for ref in refs:
                path = Path(ref.location) if ref.location else None
                if path is None or not path.exists():
                    continue
                try:
                    docs.append((ref.source_name, msds._read_local_document_text(path)))
                except Exception:
                    continue

            if not docs:
                return _comparison_with(
                    base,
                    status="UNREADABLE" if base.status != "NO_FILES" else base.status,
                    message=base.message + " 혼합물 구성성분은 Stage 1에서 승계되었으나 공급자 MSDS 제3항을 읽지 못해 성분·함량 대조를 완료하지 못했습니다.",
                )

            by_product: dict[tuple[int, str], list[dict[str, Any]]] = {}
            for item in expectations:
                by_product.setdefault((int(item.get("parent") or 0), _text(item.get("product"))), []).append(item)

            mismatches: list[str] = []
            unresolved: list[str] = []
            verified = 0
            for (_, product), items in by_product.items():
                token = _norm(product)
                candidates = []
                if token:
                    for name, text in docs:
                        if token in _norm(name) or token in _norm(text):
                            candidates.append((name, text))
                if len(candidates) != 1:
                    unresolved.append(f"{product or '혼합제품'}: 제품명으로 공급자 MSDS를 하나로 특정하지 못함")
                    continue
                name, text = candidates[0]
                section = _section3_text(text)
                if not section:
                    unresolved.append(f"{product}: {name}에서 MSDS 제3항 영역을 자동 식별하지 못함")
                    continue
                for item in items:
                    cas = _text(item.get("cas"))
                    label = _text(item.get("component")) or cas
                    if cas not in section:
                        mismatches.append(f"{product} / {label}: 제3항에서 CAS {cas} 미확인")
                        continue
                    expected_pct = _expected_pct(item)
                    actual_pct = _pct_on_cas_line(section, cas)
                    if expected_pct is None:
                        unresolved.append(f"{product} / {label}: Stage 1 함량값 확인 필요")
                        continue
                    if actual_pct is None:
                        unresolved.append(f"{product} / {label}: 제3항 함량을 자동 해석하지 못함")
                        continue
                    if not _same_pct(expected_pct, actual_pct):
                        mismatches.append(
                            f"{product} / {label}: Stage 1 {expected_pct[0]:g}~{expected_pct[1]:g}% vs MSDS 제3항 {actual_pct[0]:g}~{actual_pct[1]:g}%"
                        )
                        continue
                    verified += 1

            if mismatches:
                details = "; ".join(mismatches[:4])
                return _comparison_with(
                    base,
                    status="PARTIAL",
                    message=(
                        "혼합물 Stage 1 승계정보와 공급자 MSDS 제3항이 일치하지 않는 항목이 있습니다. 자동으로 확정하지 않고 담당자 확인이 필요합니다. "
                        + details
                    ),
                )
            if unresolved:
                details = "; ".join(unresolved[:4])
                return _comparison_with(
                    base,
                    status="PARTIAL",
                    message=(
                        "CAS 존재 여부는 비교했지만 혼합물의 제품별 제3항 성분·함량을 충분히 확정하지 못했습니다. 담당자 확인 전까지 완료로 보지 않습니다. "
                        + details
                    ),
                )
            if verified and base.status == "CAS_COVERED":
                return _comparison_with(
                    base,
                    status="CAS_COVERED",
                    message=(
                        f"연결된 공급자 MSDS에서 대상 CAS를 확인했고, 제품명으로 연결된 혼합물 MSDS 제3항의 구성성분·함량 {verified}건도 Stage 1 승계정보와 일치했습니다. "
                        "이는 자동 보조검토 결과이며 MSDS 전체의 법적 적정성을 확정하는 것은 아닙니다."
                    ),
                )
            return base

        compare_with_composition._planner_mixture_handoff_wrapped = True
        msds.compare_supplier_sds = compare_with_composition


def install_stage2_visible_handoff_runtime() -> None:
    """Install visible provenance and immutable Stage-1 handoff for workbooks
    and MSDS linkage. Mixture-component handoff into Stage 2 project fields is
    already native to engine.stage2.project.create_project_from_stage1_snapshot."""
    import engine.stage2.project as project_module
    if getattr(project_module, "_visible_stage1_handoff_runtime_installed", False):
        return
    _install_workbook_handoff()
    _install_msds_mixture_handoff()
    project_module._visible_stage1_handoff_runtime_installed = True
