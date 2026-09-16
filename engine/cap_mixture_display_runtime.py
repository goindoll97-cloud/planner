from __future__ import annotations

"""Final user-facing cleanup for the CAP mixture workbook extension.

The mixture calculation runtime intentionally focuses on regulatory behavior.
This small wrapper keeps the generated company workbook aligned with the
existing UI contract: full Korean report names only, and blank future input rows
remain blank instead of being prefilled as non-mixtures.
"""

from io import BytesIO
from typing import Callable

from openpyxl import load_workbook

from .cap_mixture_runtime import MIXTURE_FLAG_COLUMN


_SHORT_NAME = "화사계"
_FULL_NAME = "화학사고예방관리계획서"


def _replace_short_name(value):
    if isinstance(value, str) and _SHORT_NAME in value:
        return value.replace(_SHORT_NAME, _FULL_NAME)
    return value


def _cleanup_company_workbook(raw: bytes) -> bytes:
    wb = load_workbook(BytesIO(raw), data_only=False)

    # Existing public workbook contract uses full legal/report names and never
    # exposes the internal shorthand '화사계'.
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                replacement = _replace_short_name(cell.value)
                if replacement != cell.value:
                    cell.value = replacement

    # The base template reserves numbered blank rows.  Do not let the new
    # mixture flag turn those otherwise empty rows into visible records.
    if "02_화학물질목록" in wb.sheetnames:
        ws = wb["02_화학물질목록"]
        headers = {
            str(ws.cell(3, col).value or "").strip(): col
            for col in range(1, ws.max_column + 1)
        }
        mix_col = headers.get(MIXTURE_FLAG_COLUMN)
        if mix_col:
            for row_no in range(4, ws.max_row + 1):
                has_product_data = any(
                    ws.cell(row_no, col).value not in (None, "")
                    for col in (2, 3, 4, 5)  # 제품명, CAS, 물질명, 함량
                )
                if not has_product_data:
                    ws.cell(row_no, mix_col).value = None

    output = BytesIO()
    wb.save(output)
    return output.getvalue()


def _wrap_builder(builder: Callable[[], bytes]) -> Callable[[], bytes]:
    def wrapped() -> bytes:
        return _cleanup_company_workbook(builder())

    wrapped._planner_cap_mixture_display = True
    return wrapped


def install_cap_mixture_display_runtime() -> None:
    from . import template as template_module

    current = template_module.build_minimal_input_workbook
    if getattr(current, "_planner_cap_mixture_display", False):
        return
    template_module.build_minimal_input_workbook = _wrap_builder(current)
