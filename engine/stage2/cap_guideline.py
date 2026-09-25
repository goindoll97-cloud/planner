from __future__ import annotations

"""Reads the two regulation source documents kept under data/stage2/cap_sources.

- cap_forms_guideline.docx: 별지 제1~16호 only. It is the guideline for what the
  screens show (form titles, table headers, form notes).
- cap_annex_rules_260409.docx: contains 별표 1~4. These are rules the program
  learns (계산·판정 로직과 도움말 근거), not screen content.
"""

from dataclasses import dataclass
from functools import lru_cache
from pathlib import Path
import re

from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph

SOURCE_DIR = Path(__file__).resolve().parents[2] / "data" / "stage2" / "cap_sources"
FORMS_GUIDELINE = SOURCE_DIR / "cap_forms_guideline.docx"
ANNEX_RULES = SOURCE_DIR / "cap_annex_rules_260409.docx"

_FORM_HEAD = re.compile(r"^■.*\[별지 제(\d+)호서식\]")
_ANNEX_HEAD = re.compile(r"^\[별표 (\d+)\]$")


@dataclass(frozen=True)
class FormGuideline:
    no: int
    title: str
    headings: tuple[str, ...]
    tables: tuple[tuple[tuple[str, ...], ...], ...]
    notes: tuple[str, ...]

    def table_header(self, index: int) -> tuple[str, ...]:
        """Distinct header cells of table `index` within this form."""
        return tuple(dict.fromkeys(c for row in self.tables[index][:2] for c in row if c))


def _blocks(doc):
    for el in doc.element.body.iterchildren():
        if el.tag.endswith("}p"):
            yield Paragraph(el, doc)
        elif el.tag.endswith("}tbl"):
            yield Table(el, doc)


def _cells(table: Table) -> tuple[tuple[str, ...], ...]:
    return tuple(tuple(c.text.strip().replace("\n", " ") for c in row.cells) for row in table.rows)


@lru_cache(maxsize=1)
def form_guidelines() -> dict[int, FormGuideline]:
    doc = Document(str(FORMS_GUIDELINE))
    raw: dict[int, dict] = {}
    current = None
    for block in _blocks(doc):
        if isinstance(block, Table):
            if current:
                raw[current]["tables"].append(_cells(block))
            continue
        text = block.text.strip()
        match = _FORM_HEAD.match(text)
        if match:
            current = int(match.group(1))
            if current in raw:  # the source labels 별지 제14호 as 제13호 a second time
                current = max(raw) + 1
            raw[current] = {"title": "", "headings": [], "tables": [], "notes": []}
        elif current and text:
            item = raw[current]
            if not item["title"]:
                item["title"] = text
            elif re.match(r"^\d+(-\d+)?\. ", text):
                item["headings"].append(text)
            elif text.startswith("주)") or text[:1] in "①②③④⑤⑥⑦⑧⑨⑩":
                item["notes"].append(text)
    return {
        no: FormGuideline(no, v["title"], tuple(v["headings"]), tuple(v["tables"]), tuple(v["notes"]))
        for no, v in raw.items()
    }


@lru_cache(maxsize=1)
def annex_rules() -> dict[int, tuple[str, ...]]:
    """별표 number -> its numbered rule paragraphs (별표 1 = 최대보유량 산정 방법)."""
    doc = Document(str(ANNEX_RULES))
    rules: dict[int, list[str]] = {}
    current = None
    for block in _blocks(doc):
        if isinstance(block, Table):
            continue
        text = block.text.strip()
        match = _ANNEX_HEAD.match(text)
        if match:
            current = int(match.group(1))
            rules[current] = []
        elif text.startswith("■") or _FORM_HEAD.match(text):
            current = None
        elif current and re.match(r"^\d+\. ", text):
            rules[current].append(text)
    return {no: tuple(items) for no, items in rules.items()}


@lru_cache(maxsize=1)
def protected_target_rules() -> dict[str, dict[str, str]]:
    """작성 규정 별표 4 「보호대상」: 구분별 세부유형 및 규모 조건."""
    doc = Document(str(ANNEX_RULES))
    in_annex4 = False
    tables: list[tuple[tuple[str, ...], ...]] = []
    for block in _blocks(doc):
        if isinstance(block, Paragraph):
            text = block.text.strip()
            if _ANNEX_HEAD.match(text):
                in_annex4 = text == "[별표 4]"
            elif text.startswith("■"):
                in_annex4 = False
        elif in_annex4 and len(tables) < 3:
            tables.append(_cells(block))
    result: dict[str, dict[str, str]] = {}
    for label, table in zip(("갑종", "을종", "환경수용체"), tables):
        entries: dict[str, str] = {}
        for row in table[1:]:
            cells = list(dict.fromkeys(c for c in row if c))
            if len(cells) >= 2:
                entries[" ".join(cells[-2].split())] = cells[-1]
        result[label] = entries
    return result


@lru_cache(maxsize=1)
def preliminary_scenario_quantities() -> dict[tuple[str, str], float]:
    """별표 2: (성상, 유해성 분류) -> 예비시나리오 규정수량(kg)."""
    doc = Document(str(ANNEX_RULES))
    result: dict[tuple[str, str], float] = {}
    for block in _blocks(doc):
        if not isinstance(block, Table):
            continue
        rows = _cells(block)
        if rows and "예비시나리오" in rows[0][0]:
            for row in rows[2:]:
                match = re.search(r"([\d,]+)\s*kg", row[2] if len(row) > 2 else "")
                if match:
                    result[(row[0].strip(), row[1].strip())] = float(match.group(1).replace(",", ""))
            break
    return result


@lru_cache(maxsize=1)
def initiating_event_frequencies() -> dict[str, float]:
    """별지 제14호 서식 표에 인쇄된 개시사건 -> 기준빈도(/연)."""
    result: dict[str, float] = {}
    for row in form_guidelines()[14].tables[0]:
        cells = [c.strip() for c in row]
        if len(cells) >= 3 and cells[0].isdigit():
            found = re.match(r"^(\d+(?:\.\d+)?)\s*[×x]\s*10\s*-\s*(\d+)$", cells[2])
            if found:
                result[cells[1]] = float(found.group(1)) * 10 ** -int(found.group(2))
    return result


@lru_cache(maxsize=1)
def risk_score_thresholds() -> dict[str, tuple[float, float, float]]:
    """별표 3 제1호 구간별 점수표: 0·1·2점 구간의 '미만' 경계값(3점은 마지막 경계 이상)."""
    doc = Document(str(ANNEX_RULES))
    keys = ("scenario_count", "facility_frequency", "offsite_distance_m", "population")
    for block in _blocks(doc):
        if not isinstance(block, Table):
            continue
        rows = _cells(block)
        if rows and "구간" in rows[0][0] and "시설빈도" in " ".join(rows[0]):
            columns: dict[str, list[float]] = {key: [] for key in keys}
            for row in rows[1:4]:
                for key, cell in zip(keys, row[1:5]):
                    found = re.match(r"^\s*([\d.]+)\s*미만", cell)
                    if found:
                        columns[key].append(float(found.group(1)))
            if all(len(values) == 3 for values in columns.values()):
                return {key: tuple(values) for key, values in columns.items()}
    return {}


@lru_cache(maxsize=1)
def risk_matrix_legible_cells() -> dict[tuple[int, int], str]:
    """별표 3 제2호 위험도 판정표에서 글자로 남아 있는 칸: (사고영향점수, 사고빈도점수) -> 등급."""
    doc = Document(str(ANNEX_RULES))
    for block in _blocks(doc):
        if not isinstance(block, Table):
            continue
        rows = _cells(block)
        if rows and rows[0][0].startswith("위험도 판정표"):
            found: dict[tuple[int, int], str] = {}
            for row in rows[3:]:
                impact = row[1].strip() if len(row) > 1 else ""
                if not impact.isdigit():
                    continue
                for frequency, cell in enumerate(row[2:9]):
                    if cell.strip() in ("가", "나", "다"):
                        found[(int(impact), frequency)] = cell.strip()
            return found
    return {}
