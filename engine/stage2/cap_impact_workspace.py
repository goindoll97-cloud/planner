from __future__ import annotations

"""시나리오별 영향범위 → 장외 사고시나리오 판정 → 영향범위 내 주민·보호대상 집계 → 별지 제12·13호 사실 생성.

피해반경은 cap_release_workspace(독성: 끝점농도 + 가우시안 플룸, 화재·폭발: 1 psi / 5 kW/m2)에서 오고,
보호대상은 별지 제8호 목록(사업장 경계와의 거리)에서 온다. 방향(풍향)과 지형은 반영하지 않고 사업장
경계 기준 원형 범위로 보수적으로 집계한다. 사업장 경계에서 500m 밖은 별지 제8호에 없으므로 그 밖의
보호대상은 집계되지 않는다(경고). 결과 사실은 CALCULATED 상태로 기존 별지 제12·13호 엔진이 읽는 키에 쓴다.
"""

from dataclasses import dataclass, field
from datetime import date
import io
from typing import Any, Mapping

from . import cap_dispersion as disp
from . import cap_form8_workspace as f8
from . import cap_release_workspace as rw
from . import cap_scenario_workspace as sc
from .project import Stage2Project

IMPACT_TABLE_KEY = "cap.offsite.scenario_impact_table"
SUMMARY_KEY = "cap.offsite.overall_impact_summary"
TARGETS_KEY = "cap.offsite.population_and_protected_targets"
NO_SCENARIO_KEY = "cap.offsite.no_offsite_scenario"
DOCUMENT_KEY = "documents.kora_impact_result"
LISTED_RANGE_M = 500.0  # 별지 제8호가 다루는 사업장 경계 밖 범위


@dataclass(frozen=True)
class ScenarioImpact:
    name: str
    material: str
    tag: str
    kind: str
    radius_m: float | None
    off_site_m: float | None
    basis: str
    problems: tuple[str, ...]
    notes: tuple[str, ...]
    targets: tuple[dict[str, Any], ...] = field(default_factory=tuple)

    @property
    def is_off_site(self) -> bool:
        return bool(self.off_site_m and self.off_site_m > 0)

    def count(self, category: str) -> int:
        return sum(1 for t in self.targets if t.get("보호대상 구분") == category)

    def people(self, column: str) -> int:
        return sum(int(_number(t.get(column)) or 0) for t in self.targets)


def _number(value: object) -> float | None:
    try:
        text = str(value).strip().replace(",", "")
        return float(text) if text else None
    except ValueError:
        return None


def _targets_within(project: Stage2Project, distance_m: float) -> list[dict[str, Any]]:
    out = []
    for row in f8.saved_rows(project):
        distance = _number(row.get("사업장 경계와 거리(m)"))
        if distance is not None and distance <= distance_m:
            out.append(dict(row))
    return out


def evaluate(project: Stage2Project, *, worst_case: bool = False, detection: str = "C",
             isolation: str = "C") -> list[ScenarioImpact]:
    weather = rw.default_weather(project, worst_case)
    results = []
    for scenario in sc.saved_scenarios(project):
        kind = str(scenario.get("사고유형") or "")
        if kind == sc.TOXIC:
            effect = rw.toxic_effect_for_scenario(project, scenario, weather=weather, detection=detection, isolation=isolation)
            radius, off_site, basis = effect.radius_m, effect.off_site_m, effect.endpoint_basis
            problems, notes = effect.problems, effect.notes
        elif kind == sc.FIRE:
            effect = rw.fire_effect_for_scenario(project, scenario, weather=weather, detection=detection, isolation=isolation)
            radius, off_site = effect.radius_m, effect.off_site_m
            basis = "; ".join(part for part in (effect.explosion_basis, effect.fire_basis) if part)
            problems, notes = effect.problems, effect.notes
        else:
            radius = off_site = None
            basis, problems, notes = "", ("사고유형(독성 누출/화재·폭발)을 정해야 합니다.",), ()
        targets = tuple(_targets_within(project, off_site)) if off_site else ()
        extra_notes = list(notes)
        if off_site and off_site > LISTED_RANGE_M:
            extra_notes.append(f"장외거리 {off_site:.0f} m가 별지 제8호가 다루는 500m를 넘어 500m 밖 보호대상·인구는 집계되지 않았습니다.")
        results.append(ScenarioImpact(
            str(scenario.get("사고시나리오명") or ""), str(scenario.get("유해화학물질명") or ""),
            str(scenario.get("대상 설비번호") or ""), kind, radius, off_site, basis, tuple(problems), tuple(extra_notes), targets))
    return results


def summary_text(impacts: list[ScenarioImpact]) -> str:
    off_site = [i for i in impacts if i.is_off_site]
    if not off_site:
        return "영향범위가 사업장 경계를 벗어나는 사고시나리오가 없습니다."
    toxic = max((i.off_site_m for i in off_site if i.kind == sc.TOXIC), default=0.0)
    fire = max((i.off_site_m for i in off_site if i.kind == sc.FIRE), default=0.0)
    return f"장외 사고시나리오 {len(off_site)}건. 최대 장외거리: 독성 {toxic:.0f} m, 화재·폭발 {fire:.0f} m."


def apply_to_forms(project: Stage2Project, impacts: list[ScenarioImpact], *, worst_case: bool = False) -> dict[str, int]:
    """Write CALCULATED facts read by the 별지 제12·13호 engines. Returns counts."""
    today = date.today().isoformat()
    weather = rw.default_weather(project, worst_case)
    evidence = f"자체 영향범위 분석 근거서 {today}"
    address = f8.address(project)
    scenarios = {str(s.get("사고시나리오명")): s for s in sc.saved_scenarios(project)}
    off_site = [i for i in impacts if i.is_off_site]

    rows, union = [], {}
    for impact in off_site:
        scenario = scenarios.get(impact.name, {})
        for target in impact.targets:
            union[(str(target.get("보호대상 명칭") or ""), str(target.get("보호대상 구분") or ""))] = target
        rows.append({
            "사고시나리오명": impact.name, "유해화학물질명": impact.material, "대상 설비번호": impact.tag,
            "사고유형": impact.kind, "장외거리(m)": round(impact.off_site_m, 1),
            "거주민수": impact.people("거주민수"), "근로자수": impact.people("근로자수"),
            "갑종 보호대상 수": impact.count("갑종"), "을종 보호대상 수": impact.count("을종"),
            "환경수용체 수": impact.count("환경수용체"),
            "사고원점 좌표": str(scenario.get("사고원점 좌표") or "").strip() or (f"주소 기준: {address}" if address else ""),
            "KORA/GIS 근거": f"{evidence} ({impact.basis or '기술지침 2021-3, KOSHA GUIDE P-92-2023'}; {weather.label})",
        })
    project.set_field(IMPACT_TABLE_KEY, "사고시나리오별 영향평가(자체 계산)", rows, "CALCULATED",
                      note="cap_impact_workspace: 끝점 도달거리 − 경계 거리, 방향 미반영 원형 범위")

    listed = []
    for target in union.values():
        people = int(_number(target.get("거주민수")) or 0) + int(_number(target.get("근로자수")) or 0)
        listed.append({
            "보호대상 명칭": target.get("보호대상 명칭"), "보호대상 구분": target.get("보호대상 구분"),
            "세부유형": target.get("세부유형"), "주소·위치": target.get("주소·위치"),
            "사업장 경계와 거리(m)": target.get("사업장 경계와 거리(m)"), "인원수": people,
            "GIS 근거": target.get("GIS/현장 근거") or evidence,
        })
    project.set_field(TARGETS_KEY, "총괄영향범위 내 보호대상(자체 계산)", listed, "CALCULATED", note="별지 제8호 목록에서 집계")
    project.set_field(SUMMARY_KEY, "총괄영향범위 요약(자체 계산)", [{
        "총괄영향범위 산출방법": "시나리오별 영향범위(끝점 도달거리)를 사업장 경계 기준 원형으로 보고 최외곽을 총괄영향범위로 함(풍향·지형 미반영, 보수적)",
        "총괄영향범위 결과 요약": summary_text(impacts),
        "GIS/KORA 근거": evidence,
        "보호대상 없음 여부": "예" if not listed else "아니오",
        "총괄영향범위 내 거주민수": sum(int(_number(t.get("거주민수")) or 0) for t in union.values()),
        "총괄영향범위 내 근로자수": sum(int(_number(t.get("근로자수")) or 0) for t in union.values()),
    }], "CALCULATED", note="cap_impact_workspace")
    project.set_field(DOCUMENT_KEY, "영향범위 분석 근거서", f"프로그램 생성 자체 분석 근거서({today})", "CALCULATED",
                      note="내려받은 근거서를 화학사고예방관리계획서에 첨부한다(규정 제23조 ⑥)")
    if not off_site and impacts:
        project.set_field(NO_SCENARIO_KEY, "장외 사고시나리오 없음", "예", "CALCULATED",
                          note="모든 시나리오의 영향범위가 사업장 경계 안(규정 제23조 ⑩)")
    return {"scenarios": len(rows), "targets": len(listed)}


def report_docx(project: Stage2Project, impacts: list[ScenarioImpact], *, worst_case: bool = False) -> bytes:
    """영향범위 분석 근거서: 입력, 적용 기준·모델, 시나리오별 결과, 가정과 한계."""
    from docx import Document

    weather = rw.default_weather(project, worst_case)
    doc = Document()
    doc.add_heading("사고시나리오 영향범위 분석 근거서", level=1)
    doc.add_paragraph(f"{project.company_name} · 작성일 {date.today().isoformat()} · 프로그램 자체 계산(제안값)")
    doc.add_heading("적용 기준", level=2)
    for line in (
        "화학물질안전원지침 제2021-3호 사고시나리오 선정 및 위험도 분석에 관한 기술지침(끝점, 기상, 누출공, 증발속도)",
        "KOSHA GUIDE P-92-2023 누출원 모델링에 관한 기술지침(누출률 식 2·3·4·6)",
        f"기상: {weather.label}, {weather.terrain}지형",
        "독성 끝점: ERPG-2 → AEGL-2 → PAC-2 → IDLH×0.1(기술지침 붙임 1)",
        "폭발 끝점 1 psi(EPA RMP TNT 당량식, 효율 10%), 화재 끝점 5 kW/m²(점광원 복사열 모델)",
        "확산: 지표 연속 누출 가우시안 플룸(Briggs 확산계수). 중가스 효과, TNO 멀티에너지, BLEVE 화구는 반영하지 않음",
    ):
        doc.add_paragraph(line, style="List Bullet")
    doc.add_heading("시나리오별 결과", level=2)
    table = doc.add_table(rows=1, cols=7)
    table.style = "Table Grid"
    for cell, label in zip(table.rows[0].cells, ("사고시나리오", "유형", "피해반경(m)", "장외거리(m)", "갑종/을종/환경", "거주민/근로자", "근거·비고")):
        cell.text = label
    for impact in impacts:
        cells = table.add_row().cells
        cells[0].text = impact.name
        cells[1].text = impact.kind
        cells[2].text = "" if impact.radius_m is None else f"{impact.radius_m:.0f}"
        cells[3].text = "" if impact.off_site_m is None else f"{impact.off_site_m:.0f}"
        cells[4].text = f"{impact.count('갑종')}/{impact.count('을종')}/{impact.count('환경수용체')}"
        cells[5].text = f"{impact.people('거주민수')}/{impact.people('근로자수')}"
        cells[6].text = "; ".join(impact.problems) or impact.basis
    doc.add_paragraph(summary_text(impacts))
    doc.add_heading("가정과 한계", level=2)
    for line in (
        "영향범위는 사업장 경계 기준 원형 범위로 집계했으며 풍향·지형은 반영하지 않았다(보수적).",
        "보호대상·인구는 별지 제8호 목록(경계 500m 이내)에 적은 값만 집계했다.",
        "누출시간은 API 581 감지·차단 등급표 참고값이며 수동 차단은 인정하지 않았다(기술지침 3-3 ②).",
        "결과는 제안값이며 KORA 등 동등 프로그램 결과와 대조해 확정한다(작성 규정 제23조 ⑤).",
    ):
        doc.add_paragraph(line, style="List Bullet")
    for impact in impacts:
        for note in impact.notes:
            doc.add_paragraph(f"{impact.name}: {note}", style="List Bullet")
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()
