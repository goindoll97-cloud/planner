from __future__ import annotations

"""CAP report blocks governed by current regulation/forms and NICS manual."""

from typing import Sequence

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT

from . import statutory_report as base
from .cap_authoritative import (
    cap_form6_msds_candidates,
    cap_form7_reference_candidates,
    cap_source_declaration,
    load_cap_source_lock,
)
from .project import Stage2Project


def add_cap_source_declaration(doc: Document) -> None:
    legal, manual = cap_source_declaration()
    source = load_cap_source_lock()

    heading = doc.add_heading("작성 기준", level=2)
    base._set_keep_with_next(heading)
    p = doc.add_paragraph()
    r = p.add_run("법적 작성구조·별표·별지 서식: ")
    r.bold = True
    p.add_run(legal)
    p = doc.add_paragraph()
    r = p.add_run("항목별 작성방법: ")
    r.bold = True
    p.add_run(manual)
    p = doc.add_paragraph()
    r = p.add_run("회사 사실 원칙: ")
    r.bold = True
    p.add_run(source["company_fact_rule"])
    p = doc.add_paragraph()
    r = p.add_run("KOSHA 참고자료 원칙: ")
    r.bold = True
    p.add_run(source["kosha_msds_rule"])


def _add_review_table(doc: Document, title: str, headers: Sequence[str], rows: Sequence[Sequence[object]]) -> None:
    heading = base._add_heading_safe(doc, title, level=2)
    base._set_keep_with_next(heading)
    if not rows:
        doc.add_paragraph("[자동입력 후보 없음 · 회사자료 및 제품 SDS 확인 필요]")
        return
    table = doc.add_table(rows=len(rows) + 1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    base._set_table_borders(table)
    for idx, header in enumerate(headers):
        base._set_cell_text(table.rows[0].cells[idx], header, bold=True, size=6.2, center=True)
    base._set_repeat_header(table.rows[0])
    for ridx, row in enumerate(rows, 1):
        for cidx, header in enumerate(headers):
            value = row[cidx] if cidx < len(row) else base.MISSING
            base._set_cell_text(table.rows[ridx].cells[cidx], value, size=6.0)


def add_cap_form6_msds_candidate_review(doc: Document, project: Stage2Project) -> None:
    """Show missing-field candidates separately from the statutory Form 6.

    This deliberately never writes KOSHA values into the official Form 6 cells.
    The user must confirm them against product SDS/evidence and enter them as
    company data before they become final statutory content.
    """
    candidates = cap_form6_msds_candidates(project)
    rows = [
        (
            item.cas,
            item.chemical_name or base.MISSING,
            item.field,
            item.value,
            item.source_text,
            "제품 SDS/증빙 확인 후 회사자료에 확정 입력",
        )
        for item in candidates
    ]
    _add_review_table(
        doc,
        "별지 제6호 자동입력 후보 검토 (법정서식 외 검토자료)",
        ("CAS 번호", "유해화학물질명", "항목", "후보값", "출처", "검토상태"),
        rows,
    )
    p = doc.add_paragraph(
        "※ NICS-GP2026-8 p.41~42 및 현행 별지 제6호서식 기준. 물질구분과 고유번호는 KOSHA MSDS 제15항에서 자동 확정하지 않으며, "
        "현행 지정·분류 근거와 회사 제품자료를 확인한다. 부식성은 명시적인 금속부식성 근거가 있을 때만 '유' 후보를 제시하고, 자료 부재를 '무'로 추정하지 않는다."
    )
    if p.runs:
        p.runs[0].font.size = base.Pt(7)


def add_cap_form7_msds_candidate_review(doc: Document, project: Stage2Project) -> None:
    rows_data = cap_form7_reference_candidates(project)
    rows = [
        (
            row.get("CAS 번호", ""),
            row.get("유해화학물질명", ""),
            row.get("인체유해성 후보", ""),
            row.get("물리적 위험성 후보", ""),
            row.get("환경유해성 후보", ""),
            row.get("출처", ""),
            row.get("선정 사유", ""),
        )
        for row in rows_data
    ]
    _add_review_table(
        doc,
        "별지 제7호 유해성정보 후보 검토 (법정서식 외 검토자료)",
        ("CAS 번호", "유해화학물질명", "인체유해성 후보", "물리적 위험성 후보", "환경유해성 후보", "출처", "선정 사유"),
        rows,
    )
    p = doc.add_paragraph(
        "※ NICS-GP2026-8 p.43~44 기준. 대표물질 선정과 선정 사유는 사업장 취급현황·물리화학적 특성·독성·영향범위 등을 함께 검토해야 하므로 자동으로 결정하지 않는다."
    )
    if p.runs:
        p.runs[0].font.size = base.Pt(7)
