from __future__ import annotations

from collections import Counter
from dataclasses import dataclass
from io import BytesIO
from typing import Any, Iterable, Mapping, Sequence
import re

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

from .cap_chemical_legal import build_cap_chemical_legal_data
from .cap_form1_engine import build_cap_form1_data
from .cap_form9_engine import build_cap_form9_data
from .cap_form10_engine import build_cap_form10_data
from .cap_form11_engine import build_cap_form11_data
from .cap_impact_engine import build_cap_form12_data, build_cap_form13_data
from .cap_risk_engine import build_cap_form14_data, build_cap_form15_data
from .cap_form16_engine import build_cap_form16_data
from .intake import selected_requirement_specs
from .project import CONFIRMED_STATUSES, Stage2Project


MISSING = "[확인 필요]"
PSM_SOURCE = (
    "공정안전보고서의 제출·심사·확인 및 이행상태평가 등에 관한 규정 "
    "[시행 2025.5.30., 고용노동부고시 제2025-30호]"
)
CAP_SOURCE = (
    "화학사고예방관리계획서 작성 등에 관한 규정 "
    "[시행 2026.4.22., 화학물질안전원고시 제2026-07호]"
)


@dataclass(frozen=True)
class FormSpec:
    reference: str
    title: str
    headers: tuple[str, ...]
    note: str = ""


PSM_FORMS = {
    "13": FormSpec(
        "별지 제13호서식",
        "유해·위험물질 목록",
        (
            "화학물질", "CAS No.", "분자식", "폭발한계 하한", "폭발한계 상한", "노출기준", "독성치",
            "인화점", "발화점", "증기압(20℃, mmHg)", "부식성 유무", "이상반응 유무", "일일사용량", "저장량", "비고",
        ),
        "저장량은 설비의 최대 저장량, 취급량은 하루 동안 취급할 수 있는 최대량을 기준으로 확인한다.",
    ),
    "14": FormSpec(
        "별지 제14호서식",
        "동력기계 목록",
        ("동력기계 번호", "동력기계명", "명세", "주요재질", "전동기용량(kW)", "방호·보호장치 종류", "비고"),
    ),
    "15": FormSpec(
        "별지 제15호서식",
        "장치 및 설비 명세",
        (
            "장치번호", "장치명", "내용물", "용량", "압력(MPa)-운전", "압력(MPa)-설계",
            "온도(℃)-운전", "온도(℃)-설계", "사용재질-본체", "사용재질-부속품", "사용재질-개스킷",
            "용접효율", "계산두께(mm)", "부식여유(mm)", "사용두께(mm)", "후열처리 여부", "비파괴검사율(%)", "비고",
        ),
    ),
    "16": FormSpec(
        "별지 제16호서식",
        "배관 및 개스킷 명세",
        ("분류코드", "유체의 명칭 또는 구분", "설계온도", "설계압력", "배관재질", "개스킷 재질 및 형태", "비파괴검사율", "후열처리여부", "비고"),
    ),
    "17": FormSpec(
        "별지 제17호서식",
        "안전밸브 및 파열판 명세",
        (
            "계기번호", "내용물", "물상태", "배출용량(kg/hr)", "정격용량(kg/hr)", "노즐크기-입구", "노즐크기-출구",
            "보호기기 번호", "보호기기 운전압력(MPa)", "보호기기 설계압력(MPa)", "설정압력(MPa)",
            "몸체재질", "TRIM 재질", "정밀도(오차범위)", "배출 연결부위", "배출원인", "형식",
        ),
    ),
    "17-2": FormSpec(
        "별지 제17호의2서식",
        "이상발생시 인터록 작동조건 및 가동중지 범위",
        ("인터록번호", "대상설비번호", "설정값-온도(℃)", "설정값-압력(MPa)", "설정값-액위(m)", "설정값-기타", "감지기번호", "최종 작동설비번호", "가동중지범위", "점검주기", "비고"),
    ),
    "17-3": FormSpec(
        "별지 제17호의3서식",
        "소화설비 설치계획",
        ("설치지역", "소화기", "자동확산소화기", "자동소화장치", "옥내소화전", "스프링클러", "물분무소화설비", "포소화설비", "CO2 소화설비", "할로겐화합물 소화설비", "청정소화약제 소화설비", "옥외소화전"),
    ),
    "17-4": FormSpec(
        "별지 제17호의4서식",
        "화재탐지경보설비 설치계획",
        ("설치지역", "단독경보형 감지기", "비상경보설비", "시각경보기", "자동화재탐지설비", "비상방송설비", "자동화재속보설비", "통합감시시설", "누전경보기"),
    ),
    "17-5": FormSpec(
        "별지 제17호의5서식",
        "가스누출감지경보기 설치계획",
        ("감지기번호", "감지대상", "설치장소", "작동시간", "측정방식", "경보설정값", "경보기 위치", "정밀도", "경보시 조치내용", "유지관리", "비고"),
    ),
    "18": FormSpec(
        "별지 제18호서식",
        "내화구조 명세",
        ("내화설비 또는 지역", "내화부위", "내화시험기준 및 시간", "비고"),
    ),
    "19": FormSpec(
        "별지 제19호서식",
        "국소배기장치 개요",
        ("공정 또는 작업장명", "실내외 구분", "발생원", "유해물질 종류", "후드형식", "후드 제어풍속(m/s)", "덕트내 반송속도(m/s)", "배풍량(m3/min)", "전동기용량(kW)", "배기 및 처리순서", "방폭형식"),
    ),
    "20": FormSpec(
        "별지 제20호서식",
        "방폭전기/계장 기계·기구 선정기준",
        ("설치장소 또는 공정", "전기/계장 기계·기구명", "0종장소 선정기준(방폭형식)", "1종장소 선정기준(방폭형식)", "2종장소 선정기준(방폭형식)"),
    ),
    "21": FormSpec(
        "별지 제21호서식",
        "위험성평가 참여 전문가 명단",
        ("책임분야", "성명", "소속회사", "직책", "주요경력"),
    ),
}


CAP_FORMS = {
    "6": FormSpec(
        "별지 제6호서식",
        "유해화학물질 목록 및 명세",
        ("연번", "유해화학물질명", "물질구분", "화학물질식별번호(CAS 번호)", "고유번호", "물질상태", "함량(%)", "비중", "폭발한계 하한(%)", "폭발한계 상한(%)", "독성구분-항목", "독성구분-구분", "위험노출수준", "허용농도값", "증기압(20℃, mmHg)", "부식성(유, 무)"),
    ),
    "9": FormSpec(
        "별지 제9호서식",
        "장치·설비 목록 및 명세",
        ("연번", "구분기호", "장치·설비명", "취급물질", "CAS No.", "물질상태", "함량(%)", "연결구 크기(mm)", "압력(MPa)-설계", "압력(MPa)-운전", "온도(℃)-설계", "온도(℃)-운전", "설계용량(m3)", "취급량(ton)", "비고"),
    ),
    "10": FormSpec(
        "별지 제10호서식",
        "확산방지설비 현황",
        ("연번", "취급시설-설비형태", "취급시설-구분기호", "취급시설-장치·설비명", "취급시설-설계용량", "확산방지설비-설비종류", "확산방지설비-필요용량", "확산방지설비-유효용량", "확산방지설비-검토결과", "비고"),
        "확산방지설비 배치도는 담당자가 별도 작성·첨부할 수 있다.",
    ),
    "11": FormSpec(
        "별지 제11호서식",
        "고정식 유해감지시설 명세",
        ("연번", "구분기호", "감지대상", "설치위치", "작동시간", "측정방식", "경보설정값", "경보기 설치장소", "연동여부", "정밀도", "유지관리", "비고"),
    ),
    "14": FormSpec(
        "별지 제14호서식",
        "사고시나리오별 시설빈도",
        ("연번", "개시사건", "빈도", "개수", "사고빈도"),
    ),
    "15": FormSpec(
        "별지 제15호서식",
        "위험도 분석 - 사업장 내 위험도 판단 요소 선정",
        ("연번", "사고시나리오 명", "사고시나리오 시설빈도", "사고시나리오 거리(장외)", "주민수"),
    ),
}


PSM_STRUCTURED_REQUIREMENTS = {
    "psm.business.overview",
    "psm.psi.chemical_inventory",
    "psm.psi.machinery",
    "psm.psi.equipment_specs",
    "psm.psi.piping_gasket",
    "psm.psi.relief_devices",
    "psm.psi.fireproofing",
    "psm.psi.fire_protection",
    "psm.psi.fire_detection",
    "psm.psi.gas_detection",
    "psm.psi.local_exhaust",
    "psm.psi.ex_equipment",
    "psm.risk.team",
    "psm.risk.consequence",
}

CAP_STRUCTURED_REQUIREMENTS = {
    "cap.basic.business",
    "cap.basic.facility_overview",
    "cap.basic.chemical_inventory",
    "cap.basic.hazard_info",
    "cap.basic.site_location",
    "cap.facility.equipment_specs",
    "cap.safety.dike_layout",
    "cap.safety.gas_detection",
    "cap.offsite.surrounding_impact",
    "cap.offsite.scenario_frequency",
    "cap.offsite.risk_analysis",
    "cap.prevention.change_log",
}


PSM_ARTICLE_ITEMS = {
    "psm.operation.sop": (
        "최초의 시운전", "정상운전", "비상시 운전", "정상적인 운전 정지", "비상정지", "정비 후 운전 개시",
        "운전범위를 벗어났을 경우 조치 절차", "화학물질의 물성과 유해·위험성", "위험물질 누출 예방 조치",
        "개인보호구 착용방법", "위험물질에 폭로시의 조치요령과 절차", "안전설비 계통의 기능·운전방법 및 절차",
    ),
    "psm.operation.maintenance": (
        "목적", "적용범위", "구성 기기의 우선순위 등급", "기기의 점검", "기기의 결함관리", "기기의 정비",
        "기기 및 기자재의 품질관리", "외주업체 관리", "설비의 유지관리",
    ),
    "psm.operation.work_permit": (
        "목적", "적용범위", "안전작업허가의 일반사항", "안전작업 준비", "화기작업 허가", "일반위험작업 허가",
        "밀폐공간 출입작업 허가", "정전작업 허가", "굴착작업 허가", "방사선 사용작업 허가 등",
    ),
    "psm.operation.contractor": (
        "목적", "적용범위", "적용대상", "사업주의 의무", "도급업체 사업주의 의무", "계획서 작성 및 승인 등",
    ),
    "psm.operation.training": ("목적", "적용범위", "교육대상", "교육의 종류", "교육계획의 수립", "교육의 실시", "교육의 평가 및 사후관리"),
    "psm.operation.prestartup": ("목적", "적용범위", "점검팀의 구성", "점검시기", "점검표의 작성", "점검보고서", "점검결과의 처리"),
    "psm.operation.moc": (
        "목적", "적용범위", "변경요소 관리의 원칙", "정상변경 관리절차", "비상변경 관리절차", "변경관리위원회의 구성",
        "변경시의 검토항목", "변경업무분담", "변경에 대한 기술적 근거", "변경요구서 서식 등",
    ),
    "psm.operation.audit": ("목적", "적용범위", "감사계획", "감사팀의 구성", "감사 시행", "평가 및 시정", "문서화 등"),
    "psm.operation.incident_investigation": ("목적", "적용범위", "공정사고 조사팀의 구성", "공정사고 조사 보고서의 작성", "공정사고 조사 결과의 처리"),
    "psm.emergency.core": (
        "목적", "비상사태의 구분", "위험성 및 재해의 파악·분석", "유해·위험물질의 성질·상태 조사", "비상조치계획의 수립",
        "비상조치 계획의 검토", "비상대피 계획", "비상사태의 발령", "비상경보의 사업장 내·외부 전파",
        "비상사태의 종결", "사고조사", "비상조치 위원회의 구성", "비상통제 조직의 기능 및 책무",
        "장비보유현황 및 비상통제소의 설치", "운전정지 절차", "비상훈련의 실시 및 조정", "주민 홍보계획 등",
    ),
}


CAP_ARTICLE_ITEMS = {
    "cap.prevention.safety_management": (
        "사업장의 종합적인 화학사고 안전관리 방향 및 목표", "목표 달성을 위한 구체적인 실행과제",
        "관리적 대책", "기술적 대책",
    ),
    "cap.prevention.training": (
        "연간 교육·훈련 계획", "화학사고예방관리계획서 전문교육 및 비상대응조직 역할별 교육·훈련",
        "교육·훈련 평가방법 및 평가결과에 따른 보완계획",
    ),
    "cap.prevention.self_inspection": (
        "자체점검반 구성·점검시기·점검항목", "자체점검 실시 및 내부 보고체계", "결과 환류계획",
        "화학물질안전원 서면 제출계획(1군 사업장 해당)",
    ),
    "cap.prevention.change_management": ("변경내역 확인 담당자", "변경 확인 주기", "변경 확인 방법 및 후속조치"),
    "cap.prevention.emergency_system": (
        "사업장 내·외부 사고신고 체계", "유관기관 목록 및 사고신고 체계", "인근 사업장 공조 연락체계",
        "비상대응조직별 편성인원 및 임무", "협력업체 비상대응조직 및 임무(해당 시)", "비상통제실 지정 및 운영",
    ),
    "cap.internal.accident_response": (
        "가동중지 권한 및 절차", "자체 방재 인력", "방재장비·물품 및 개인보호장구", "방재장비·물품 관리·유지·확충",
        "사업장 내부 경보전달체계", "취급시설 유형별 응급조치계획",
    ),
    "cap.internal.post_accident": (
        "사고조사팀 구성 및 역할", "사고조사보고서 작성", "개선대책 및 이행방법", "사고복구 조직 및 역할",
        "책임보험 가입계획(해당 시)", "환경복원 전문업체 활용계획",
    ),
    "cap.external.community_coordination": (
        "대외소통 담당 조직 및 임무", "정보 제공 방법 또는 절차", "이해당사자 목록 및 제공정보",
        "평상시 지역사회 소통계획", "지역비상대응기관·인근 사업장과의 공조계획",
    ),
    "cap.external.evacuation": (
        "사고유형별 대피경보 방법", "대상별 경보전달 방법", "기초지방자치단체 담당부서 및 연락처",
        "주민행동 요령", "응급의료계획", "주민대피장소 및 방법",
    ),
    "cap.external.community_notice": ("고지대상", "고지방법", "고지정보"),
}


def _norm(value: object) -> str:
    return re.sub(r"[\s·ㆍ․/()_%℃㎥³㎡².-]+", "", str(value or "")).lower()


def _row_value(row: Mapping[str, object], *aliases: str) -> str:
    normalized = {_norm(key): value for key, value in row.items()}
    for alias in aliases:
        value = normalized.get(_norm(alias))
        if value not in (None, ""):
            return str(value).strip()
    return MISSING


def _record(project: Stage2Project, key: str):
    rec = project.get_field(key)
    if rec is None or rec.status not in CONFIRMED_STATUSES:
        return None
    return rec


def _value(project: Stage2Project, *keys: str, default: str = MISSING) -> Any:
    for key in keys:
        rec = _record(project, key)
        if rec is not None and rec.value not in (None, "", [], {}):
            return rec.value
    return default


def _text(project: Stage2Project, *keys: str, default: str = MISSING) -> str:
    value = _value(project, *keys, default=default)
    if value == default:
        return default
    if isinstance(value, Mapping):
        return " / ".join(f"{k}: {v}" for k, v in value.items() if v not in (None, "")) or default
    if isinstance(value, (list, tuple)):
        if all(isinstance(item, Mapping) for item in value):
            return f"구조화 자료 {len(value)}건 확인"
        return ", ".join(str(v) for v in value if v not in (None, "")) or default
    return str(value).strip() or default


def _rows(project: Stage2Project, *keys: str) -> list[Mapping[str, object]]:
    for key in keys:
        rec = _record(project, key)
        if rec is None:
            continue
        if isinstance(rec.value, list):
            rows = [dict(item) for item in rec.value if isinstance(item, Mapping)]
            if rows:
                return rows
        if isinstance(rec.value, Mapping):
            return [dict(rec.value)]
    return []


def _doc_value(project: Stage2Project, key: str) -> str:
    rec = project.get_field(key)
    if rec is None:
        return MISSING
    if isinstance(rec.value, Mapping):
        name = str(rec.value.get("file_name") or rec.value.get("filename") or "").strip()
        ref = str(rec.value.get("reference_no") or rec.value.get("drawing_no") or "").strip()
        rev = str(rec.value.get("revision") or "").strip()
        pieces = [v for v in (name, ref, rev) if v]
        if pieces:
            status = "확인완료" if rec.status in CONFIRMED_STATUSES else "담당자 확인 필요"
            return " / ".join(pieces) + f" ({status})"
    if rec.value not in (None, ""):
        return str(rec.value)
    return MISSING


def _checkbox(selected: bool, text: str) -> str:
    return ("☒ " if selected else "☐ ") + text


def _set_cell_text(cell, value: object, *, bold: bool = False, size: float = 7.2, center: bool = False) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.LEFT
    run = paragraph.add_run(str(value if value not in (None, "") else MISSING))
    run.font.name = "Malgun Gothic"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Malgun Gothic")
    run.font.size = Pt(size)
    run.bold = bold
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), "D9D9D9" if bold else "FFFFFF")


def _set_table_borders(table) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        elem = borders.find(qn(tag))
        if elem is None:
            elem = OxmlElement(tag)
            borders.append(elem)
        elem.set(qn("w:val"), "single")
        elem.set(qn("w:sz"), "4")
        elem.set(qn("w:color"), "000000")


def _set_repeat_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def _set_keep_with_next(paragraph) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    keep = OxmlElement("w:keepNext")
    p_pr.append(keep)


def _configure_doc(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Cm(1.3)
    section.bottom_margin = Cm(1.3)
    section.left_margin = Cm(1.2)
    section.right_margin = Cm(1.2)
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Malgun Gothic"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Malgun Gothic")
    normal.font.size = Pt(9)
    for style_name, size in (("Title", 19), ("Heading 1", 14), ("Heading 2", 11), ("Heading 3", 10)):
        style = styles[style_name]
        style.font.name = "Malgun Gothic"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Malgun Gothic")
        style.font.size = Pt(size)


def _add_title(doc: Document, title: str, source: str, project: Stage2Project, *, cap_group: str = "") -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(title)
    run.bold = True
    run.font.name = "Malgun Gothic"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Malgun Gothic")
    run.font.size = Pt(18)
    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r2 = p2.add_run("법정서식 기반 검토용 작성본")
    r2.bold = True
    r2.font.size = Pt(10)
    doc.add_paragraph(f"사업장: {project.company_name}{' / ' + project.site_name if project.site_name else ''}")
    if cap_group:
        doc.add_paragraph(f"작성수준: {cap_group}")
    doc.add_paragraph(f"서식 기준: {source}")
    notice = doc.add_paragraph()
    nr = notice.add_run(
        "※ 이 문서는 현행 고시의 작성서식과 작성순서를 기준으로 회사 확인자료를 배치한 검토용 문서입니다. "
        "확인되지 않은 사실·수치·설비사양은 임의로 채우지 않고 [확인 필요]로 표시합니다. 도면·이미지는 담당자가 별도 작성·첨부할 수 있습니다."
    )
    nr.bold = True
    nr.font.size = Pt(8)
    doc.add_paragraph()


def _add_form_heading(doc: Document, reference: str, title: str) -> None:
    p = doc.add_paragraph()
    _set_keep_with_next(p)
    run = p.add_run(f"■ {reference}  {title}")
    run.bold = True
    run.font.name = "Malgun Gothic"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Malgun Gothic")
    run.font.size = Pt(11)


def _add_form_table(doc: Document, spec: FormSpec, rows: Sequence[Sequence[object]], *, min_rows: int = 1) -> None:
    _add_form_heading(doc, spec.reference, spec.title)
    count = max(min_rows, len(rows), 1)
    table = doc.add_table(rows=count + 1, cols=len(spec.headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    _set_table_borders(table)
    for idx, header in enumerate(spec.headers):
        _set_cell_text(table.rows[0].cells[idx], header, bold=True, size=6.5, center=True)
    _set_repeat_header(table.rows[0])
    for ridx in range(count):
        source = rows[ridx] if ridx < len(rows) else [MISSING] * len(spec.headers)
        for cidx in range(len(spec.headers)):
            value = source[cidx] if cidx < len(source) else MISSING
            _set_cell_text(table.rows[ridx + 1].cells[cidx], value, size=6.5, center=False)
    if spec.note:
        p = doc.add_paragraph("주) " + spec.note)
        p.runs[0].font.size = Pt(7)


def _add_key_value_form(doc: Document, reference: str, title: str, rows: Sequence[tuple[str, object]]) -> None:
    _add_form_heading(doc, reference, title)
    table = doc.add_table(rows=len(rows), cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    _set_table_borders(table)
    for idx, (label, value) in enumerate(rows):
        _set_cell_text(table.rows[idx].cells[0], label, bold=True, size=8)
        _set_cell_text(table.rows[idx].cells[1], value, size=8)


def _add_attachment_line(doc: Document, label: str, value: str) -> None:
    p = doc.add_paragraph()
    p.add_run(label + ": ").bold = True
    p.add_run(value)


_HEADING_SIZES = {1: 14, 2: 11, 3: 10}


def _add_heading_safe(doc: Document, text: str, level: int = 1):
    """Add a heading even when the doc has no Word "Heading N" styles.

    A statutory-form baseline docx loaded from disk (rather than built via
    ``_configure_doc``) may not define "Heading N" at all, so plain
    ``doc.add_heading`` raises ``KeyError``. Register the missing style
    (matching python-docx's own default) instead of falling back to a bold
    paragraph, so callers that check ``paragraph.style.name`` still see a
    real "Heading N" style.
    """
    style_name = f"Heading {level}"
    try:
        doc.styles[style_name]
    except KeyError:
        from docx.enum.style import WD_STYLE_TYPE

        style = doc.styles.add_style(style_name, WD_STYLE_TYPE.PARAGRAPH)
        style.base_style = doc.styles["Normal"]
        style.font.bold = True
        style.font.size = Pt(_HEADING_SIZES.get(level, 10))
    return doc.add_heading(text, level=level)


def _add_narrative_requirement(doc: Document, spec, project: Stage2Project, *, article_items: Sequence[str] = ()) -> None:
    heading = _add_heading_safe(doc, spec.label, level=2)
    _set_keep_with_next(heading)
    if article_items:
        p = doc.add_paragraph()
        p.add_run("법정 작성항목: ").bold = True
        p.add_run(" / ".join(article_items))
    values = []
    for key in spec.field_keys:
        rec = project.get_field(key)
        if rec is None:
            continue
        if rec.status in CONFIRMED_STATUSES and rec.value not in (None, "", [], {}):
            values.append((rec.label or key, rec.value))
        elif rec.status == "HOLD" and rec.value not in (None, "", [], {}):
            values.append((rec.label or key, "[담당자 확인 필요]"))
    if not values:
        doc.add_paragraph(MISSING)
        return
    for label, value in values:
        if isinstance(value, list) and all(isinstance(item, Mapping) for item in value):
            doc.add_paragraph(f"{label}: 구조화 자료 {len(value)}건 확인")
        elif isinstance(value, Mapping):
            text = " / ".join(f"{k}: {v}" for k, v in value.items() if v not in (None, ""))
            doc.add_paragraph(f"{label}: {text or MISSING}")
        else:
            doc.add_paragraph(str(value))


def _chemical_rows(project: Stage2Project) -> list[Mapping[str, object]]:
    return _rows(project, "cap.chemical.details", "inventory.chemicals")


def _psm_chemical_rows(project: Stage2Project) -> list[Mapping[str, object]]:
    return _rows(project, "psm.psi.chemical_details", "inventory.chemicals", "cap.chemical.details")


def _facility_rows(project: Stage2Project) -> list[Mapping[str, object]]:
    return _rows(project, "cap.facility.equipment_specs", "psm.psi.equipment_specs", "inventory.facilities")


def _psm_facility_rows(project: Stage2Project) -> list[Mapping[str, object]]:
    return _rows(project, "psm.psi.equipment_specs", "inventory.facilities", "cap.facility.equipment_specs")


def _facility_counts(rows: Sequence[Mapping[str, object]]) -> str:
    labels = []
    for row in rows:
        value = _row_value(row, "설비종류", "설비형태", "장치·설비 종류")
        if value != MISSING:
            labels.append(value)
    if not labels:
        return MISSING
    counter = Counter(labels)
    return " / ".join(f"{name} {count}기" for name, count in counter.items())


def _psm_form12(doc: Document, project: Stage2Project) -> None:
    chemicals = _psm_chemical_rows(project)
    raw_names = ", ".join(_row_value(r, "물질명", "화학물질", "유해화학물질명") for r in chemicals[:8]) if chemicals else MISSING
    rows = (
        ("사업장명", project.company_name or MISSING),
        ("사업의 구분", _text(project, "psm.business.project_type", default=MISSING)),
        ("사업자등록번호", _text(project, "business.registration_no", "cap.business.registration_no", default=MISSING)),
        ("대표자 성명", _text(project, "business.representative", "cap.business.representative", default=MISSING)),
        ("소재지", _text(project, "business.address", default=MISSING)),
        ("심사대상 설비명", _text(project, "psm.business.target_facility", default=MISSING)),
        ("표준산업분류", _text(project, "business.ksic", default=MISSING)),
        ("예상근무 근로자수", _text(project, "business.employee_count", default=MISSING)),
        ("전기계약용량", _text(project, "business.electric_contract_capacity", default=MISSING)),
        ("보고서 작성자 / 작성자 자격", _text(project, "psm.business.writer_info", default=MISSING)),
        ("주요 사업 내용 또는 변경내용", _text(project, "psm.business.overview", default=MISSING)),
        ("주원료 또는 재료", raw_names),
        ("주생산품", _text(project, "business.main_products", default=MISSING)),
        ("사업장 위치·부지·건물", _text(project, "psm.business.site_building", default=MISSING)),
        ("공사·가동 일정", _text(project, "psm.business.schedule", default=MISSING)),
    )
    _add_key_value_form(doc, "별지 제12호서식", "사업개요", rows)


def _psm_form13_rows(project: Stage2Project) -> list[list[str]]:
    rows = []
    for row in _psm_chemical_rows(project):
        rows.append([
            _row_value(row, "물질명", "화학물질", "유해화학물질명"),
            _row_value(row, "CAS 번호", "CAS No.", "화학물질식별번호"),
            _row_value(row, "분자식"),
            _row_value(row, "폭발한계 하한", "폭발하한", "LEL"),
            _row_value(row, "폭발한계 상한", "폭발상한", "UEL"),
            _row_value(row, "노출기준", "허용농도값", "TWA"),
            _row_value(row, "독성치"),
            _row_value(row, "인화점"),
            _row_value(row, "발화점"),
            _row_value(row, "증기압", "증기압(20℃,mmHg)"),
            _row_value(row, "부식성", "부식성 유무"),
            _row_value(row, "이상반응 유무", "이상반응"),
            _row_value(row, "일일사용량", "취급량", "사용량"),
            _row_value(row, "저장량", "최대보유량", "최대보유량(kg)"),
            _row_value(row, "비고"),
        ])
    return rows


def _psm_form14_rows(project: Stage2Project) -> list[list[str]]:
    out = []
    for row in _rows(project, "psm.psi.machinery_list"):
        out.append([
            _row_value(row, "기계번호", "동력기계 번호", "설비번호"),
            _row_value(row, "기계명", "동력기계명", "설비명"),
            _row_value(row, "명세", "형식", "용량"),
            _row_value(row, "주요재질", "재질"),
            _row_value(row, "전동기용량", "동력", "전동기용량(kW)"),
            _row_value(row, "방호·보호장치 종류", "방호장치의 종류", "보호장치"),
            _row_value(row, "비고"),
        ])
    return out


def _psm_form15_rows(project: Stage2Project) -> list[list[str]]:
    out = []
    for row in _psm_facility_rows(project):
        material = _row_value(row, "취급물질", "내용물")
        body_mat = _row_value(row, "본체재질", "재질", "사용재질")
        out.append([
            _row_value(row, "설비번호", "장치번호", "구분기호"),
            _row_value(row, "설비명", "장치명", "장치·설비명"),
            material,
            _row_value(row, "용량", "설계용량", "설계용량(m3)"),
            _row_value(row, "운전압력", "압력-운전"),
            _row_value(row, "설계압력", "압력-설계"),
            _row_value(row, "운전온도", "온도-운전"),
            _row_value(row, "설계온도", "온도-설계"),
            body_mat,
            _row_value(row, "부속품재질", "부속품"),
            _row_value(row, "개스킷재질", "개스킷 재질"),
            _row_value(row, "용접효율"),
            _row_value(row, "계산두께"),
            _row_value(row, "부식여유"),
            _row_value(row, "사용두께"),
            _row_value(row, "후열처리 여부", "후열처리여부"),
            _row_value(row, "비파괴검사율", "비파괴율검사"),
            _row_value(row, "비고", "P&ID 번호"),
        ])
    return out


def _psm_form16_rows(project: Stage2Project) -> list[list[str]]:
    out = []
    for row in _rows(project, "psm.psi.piping_gasket_specs"):
        out.append([
            _row_value(row, "분류코드", "배관번호·Class", "배관번호", "Class"),
            _row_value(row, "유체의 명칭 또는 구분", "유체명"),
            _row_value(row, "설계온도"),
            _row_value(row, "설계압력"),
            _row_value(row, "배관재질"),
            _row_value(row, "개스킷 재질 및 형태", "개스킷 재질"),
            _row_value(row, "비파괴검사율"),
            _row_value(row, "후열처리여부", "후열처리 여부"),
            _row_value(row, "비고", "관련 P&ID 번호"),
        ])
    return out


def _psm_form17_rows(project: Stage2Project) -> list[list[str]]:
    out = []
    for row in _rows(project, "psm.psi.relief_device_specs", "cap.safety.relief_device_specs"):
        out.append([
            _row_value(row, "안전밸브·파열판 번호", "계기번호", "안전밸브번호"),
            _row_value(row, "배출물질", "내용물"),
            _row_value(row, "배출상태", "물상태"),
            _row_value(row, "배출용량"),
            _row_value(row, "정격용량"),
            _row_value(row, "노즐크기 입구", "입구크기"),
            _row_value(row, "노즐크기 출구", "출구크기"),
            _row_value(row, "보호대상 설비번호", "보호기기 번호"),
            _row_value(row, "보호기기 운전압력", "운전압력"),
            _row_value(row, "보호기기 설계압력", "설계압력"),
            _row_value(row, "설정압력"),
            _row_value(row, "몸체재질"),
            _row_value(row, "TRIM 재질", "트림재질"),
            _row_value(row, "정밀도", "오차범위"),
            _row_value(row, "최종 배출·처리 지점", "배출 연결부위"),
            _row_value(row, "배출원인"),
            _row_value(row, "형식"),
        ])
    return out


def _generic_form_rows(project: Stage2Project, field_key: str | Sequence[str], spec: FormSpec, alias_groups: Sequence[Sequence[str]]) -> list[list[str]]:
    keys = (field_key,) if isinstance(field_key, str) else tuple(field_key)
    rows = _rows(project, *keys)
    if not rows:
        return []
    return [[_row_value(row, *aliases) for aliases in alias_groups] for row in rows]


def _psm_form19_2(doc: Document, project: Stage2Project) -> None:
    _add_form_heading(doc, "별지 제19호의2서식", "시나리오 및 피해예측 결과")
    fields = (
        "풍속(m/s)", "대기안정도(A~F)", "대기온도(℃)", "습도(%)", "표면거칠기(m)",
        "물질명", "물질의 상태", "설비명(또는 배관부위)", "운전압력(MPa)", "운전온도(℃)",
        "누출구의 크기(mm2)", "웅덩이 크기(m2)", "누출결과", "화재-복사열이 미치는 거리", "폭발-과압이 미치는 거리",
        "확산결과-인화성", "확산결과-독성",
    )
    source = _value(project, "psm.risk.consequence", default={})
    worst: Mapping[str, object] = {}
    alternative: Mapping[str, object] = {}
    if isinstance(source, Mapping):
        worst = source.get("최악의 사고 시나리오", source.get("worst_case", {})) if isinstance(source.get("최악의 사고 시나리오", source.get("worst_case", {})), Mapping) else {}
        alternative = source.get("대안의 사고 시나리오", source.get("alternative_case", {})) if isinstance(source.get("대안의 사고 시나리오", source.get("alternative_case", {})), Mapping) else {}
    table = doc.add_table(rows=len(fields) + 1, cols=3)
    _set_table_borders(table)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for idx, value in enumerate(("구분", "최악의 사고 시나리오", "대안의 사고 시나리오")):
        _set_cell_text(table.rows[0].cells[idx], value, bold=True, size=7, center=True)
    for ridx, label in enumerate(fields, 1):
        _set_cell_text(table.rows[ridx].cells[0], label, bold=True, size=7)
        _set_cell_text(table.rows[ridx].cells[1], _row_value(worst, label), size=7)
        _set_cell_text(table.rows[ridx].cells[2], _row_value(alternative, label), size=7)


def _render_psm(doc: Document, project: Stage2Project) -> None:
    _add_title(doc, "공정안전보고서", PSM_SOURCE, project)
    doc.add_heading("사업개요", level=1)
    _psm_form12(doc, project)

    doc.add_page_break()
    doc.add_heading("공정안전자료", level=1)
    _add_form_table(doc, PSM_FORMS["13"], _psm_form13_rows(project))
    _add_attachment_line(doc, "물질안전보건자료(MSDS)", _doc_value(project, "psm.psi.msds"))
    _add_form_table(doc, PSM_FORMS["14"], _psm_form14_rows(project))
    _add_form_table(doc, PSM_FORMS["15"], _psm_form15_rows(project))
    _add_form_table(doc, PSM_FORMS["16"], _psm_form16_rows(project))
    _add_form_table(doc, PSM_FORMS["17"], _psm_form17_rows(project))

    doc.add_heading("공정개요 및 공정도면", level=2)
    doc.add_paragraph(_text(project, "process.description", default=MISSING))
    _add_attachment_line(doc, "공정흐름도(PFD)", _doc_value(project, "documents.pfd"))
    _add_attachment_line(doc, "공정배관·계장도(P&ID)", _doc_value(project, "documents.pid"))
    _add_attachment_line(doc, "유틸리티 계통도·배관계장도(UFD)", _doc_value(project, "psm.psi.ufd"))

    _add_form_table(
        doc,
        PSM_FORMS["17-2"],
        _generic_form_rows(
            project,
            "psm.psi.interlock_conditions",
            PSM_FORMS["17-2"],
            (("인터록번호",), ("대상설비번호", "대상설비"), ("온도",), ("압력",), ("액위",), ("기타",), ("감지기번호",), ("최종 작동설비번호",), ("가동중지범위",), ("점검주기",), ("비고",)),
        ),
    )
    _add_attachment_line(doc, "공장 전체배치도", _doc_value(project, "documents.site_plan"))
    _add_attachment_line(doc, "설비배치도", _doc_value(project, "psm.psi.equipment_layout"))
    _add_attachment_line(doc, "건물·철구조물 평면도 및 입면도", _doc_value(project, "psm.psi.building_structure"))

    _add_form_table(
        doc, PSM_FORMS["18"],
        _generic_form_rows(project, ("psm.psi.fireproofing_table", "psm.psi.fireproofing"), PSM_FORMS["18"], (("내화설비 또는 지역", "설비", "지역"), ("내화부위",), ("내화시험기준 및 시간", "시험기준", "내화시간"), ("비고",))),
    )
    _add_form_table(
        doc, PSM_FORMS["17-3"],
        _generic_form_rows(project, ("psm.psi.fire_protection_table", "psm.psi.fire_protection"), PSM_FORMS["17-3"], tuple((h,) for h in PSM_FORMS["17-3"].headers)),
    )
    _add_form_table(
        doc, PSM_FORMS["17-4"],
        _generic_form_rows(project, ("psm.psi.fire_detection_table", "psm.psi.fire_detection"), PSM_FORMS["17-4"], tuple((h,) for h in PSM_FORMS["17-4"].headers)),
    )
    _add_form_table(
        doc, PSM_FORMS["17-5"],
        _generic_form_rows(project, ("psm.psi.gas_detection_table", "psm.psi.gas_detection"), PSM_FORMS["17-5"], (
            ("감지기 번호", "감지기번호", "구분기호"), ("검출대상 물질", "감지대상"), ("설치위치", "설치장소"), ("작동시간",),
            ("감지방식", "측정방식"), ("경보 설정값", "경보설정값"), ("경보 위치", "경보기 위치"), ("정밀도",),
            ("경보시 조치내용",), ("유지관리", "점검주기"), ("비고", "관련 도면번호"),
        )),
    )
    _add_attachment_line(doc, "세안·세척시설 및 안전보호장구 설치계획", _text(project, "psm.psi.wash_facility", "psm.psi.ppe", default=MISSING))
    _add_form_table(
        doc, PSM_FORMS["19"],
        _generic_form_rows(project, ("psm.psi.local_exhaust_table", "psm.psi.local_exhaust"), PSM_FORMS["19"], tuple((h,) for h in PSM_FORMS["19"].headers)),
    )
    _add_attachment_line(doc, "폭발위험장소 구분도", _doc_value(project, "psm.psi.hazardous_area"))
    _add_form_table(
        doc, PSM_FORMS["20"],
        _generic_form_rows(project, "psm.psi.ex_equipment", PSM_FORMS["20"], (
            ("설치장소", "공정", "설치장소 또는 공정"), ("전기/계장 기계·기구명", "기기명"), ("0종", "0종장소"), ("1종", "1종장소"), ("2종", "2종장소"),
        )),
    )
    _add_attachment_line(doc, "전기단선도", _doc_value(project, "psm.psi.single_line"))
    _add_attachment_line(doc, "단락용량 계산서", _doc_value(project, "psm.psi.short_circuit"))
    _add_attachment_line(doc, "비상전원 설비용량 자료", _doc_value(project, "psm.psi.emergency_power"))
    _add_attachment_line(doc, "접지계획·배치도", _doc_value(project, "psm.psi.grounding"))
    doc.add_heading("안전설계·제작 및 설치 관련 지침서", level=2)
    doc.add_paragraph(_text(project, "psm.psi.design_installation_guideline", default=MISSING))

    doc.add_page_break()
    doc.add_heading("공정위험성평가서", level=1)
    risk_items = (
        ("위험성 평가의 목적", "psm.risk.purpose"),
        ("공정 위험특성", "psm.risk.characteristics"),
        ("위험성 평가결과에 따른 잠재위험의 종류", "psm.risk.report"),
        ("사고빈도 최소화 및 사고시 피해 최소화 대책", "psm.risk.mitigation"),
        ("위험성 평가 수행 및 평가절차", "psm.risk.procedure"),
    )
    for label, key in risk_items:
        doc.add_heading(label, level=2)
        doc.add_paragraph(_text(project, key, default=MISSING))
    _psm_form19_2(doc, project)
    _add_form_table(
        doc, PSM_FORMS["21"],
        _generic_form_rows(project, "psm.risk.team", PSM_FORMS["21"], (("책임분야",), ("성명",), ("소속회사", "소속"), ("직책",), ("주요경력",))),
    )

    specs = [s for s in selected_requirement_specs(project) if s.system == "PSM"]
    by_key = {s.key: s for s in specs}
    doc.add_page_break()
    doc.add_heading("안전운전계획", level=1)
    for key in (
        "psm.operation.sop", "psm.operation.maintenance", "psm.operation.work_permit", "psm.operation.contractor",
        "psm.operation.training", "psm.operation.prestartup", "psm.operation.moc", "psm.operation.audit", "psm.operation.incident_investigation",
    ):
        spec = by_key.get(key)
        if spec is not None:
            _add_narrative_requirement(doc, spec, project, article_items=PSM_ARTICLE_ITEMS.get(key, ()))

    doc.add_page_break()
    doc.add_heading("비상조치계획", level=1)
    emergency_specs = [s for s in specs if s.section == "비상조치계획"]
    if emergency_specs:
        for spec in emergency_specs:
            items = PSM_ARTICLE_ITEMS.get("psm.emergency.core", ()) if spec.key in {"psm.emergency.core", "psm.emergency.roles_procedures"} else ()
            _add_narrative_requirement(doc, spec, project, article_items=items)
    else:
        p = doc.add_paragraph()
        p.add_run("법정 작성항목: ").bold = True
        p.add_run(" / ".join(PSM_ARTICLE_ITEMS["psm.emergency.core"]))
        doc.add_paragraph(MISSING)


def _render_psm_narrative(doc: Document, project: Stage2Project) -> None:
    """Append PSM content that has no annex-form cell onto an existing document.

    Forms 12~21 (사업개요, 공정안전자료 표, 공정위험성평가서 표, 위험성평가
    참여 전문가 명단) are the regulation-form baseline itself and are filled
    separately by ``psm_baseline_docx.build_psm_baseline_draft``. This only
    adds the surrounding content the statute also requires but that has no
    annex-form counterpart: process description, referenced drawings, and the
    free-text 안전운전계획/비상조치계획 chapters.
    """
    doc.add_page_break()
    caption = doc.add_paragraph("법정서식 기반 검토용 작성본 · 별지서식에 없는 서술형 작성항목")
    caption.runs[0].italic = True

    _add_heading_safe(doc, "공정개요 및 공정도면", level=1)
    doc.add_paragraph(_text(project, "process.description", default=MISSING))
    _add_attachment_line(doc, "공정흐름도(PFD)", _doc_value(project, "documents.pfd"))
    _add_attachment_line(doc, "공정배관·계장도(P&ID)", _doc_value(project, "documents.pid"))
    _add_attachment_line(doc, "유틸리티 계통도·배관계장도(UFD)", _doc_value(project, "psm.psi.ufd"))
    _add_attachment_line(doc, "물질안전보건자료(MSDS)", _doc_value(project, "psm.psi.msds"))
    _add_attachment_line(doc, "공장 전체배치도", _doc_value(project, "documents.site_plan"))
    _add_attachment_line(doc, "설비배치도", _doc_value(project, "psm.psi.equipment_layout"))
    _add_attachment_line(doc, "건물·철구조물 평면도 및 입면도", _doc_value(project, "psm.psi.building_structure"))
    _add_attachment_line(doc, "세안·세척시설 및 안전보호장구 설치계획", _text(project, "psm.psi.wash_facility", "psm.psi.ppe", default=MISSING))
    _add_attachment_line(doc, "폭발위험장소 구분도", _doc_value(project, "psm.psi.hazardous_area"))
    _add_attachment_line(doc, "전기단선도", _doc_value(project, "psm.psi.single_line"))
    _add_attachment_line(doc, "단락용량 계산서", _doc_value(project, "psm.psi.short_circuit"))
    _add_attachment_line(doc, "비상전원 설비용량 자료", _doc_value(project, "psm.psi.emergency_power"))
    _add_attachment_line(doc, "접지계획·배치도", _doc_value(project, "psm.psi.grounding"))

    _add_heading_safe(doc, "안전설계·제작 및 설치 관련 지침서", level=1)
    doc.add_paragraph(_text(project, "psm.psi.design_installation_guideline", default=MISSING))

    _add_heading_safe(doc, "공정위험성평가서", level=1)
    risk_items = (
        ("위험성 평가의 목적", "psm.risk.purpose"),
        ("공정 위험특성", "psm.risk.characteristics"),
        ("위험성 평가결과에 따른 잠재위험의 종류", "psm.risk.report"),
        ("사고빈도 최소화 및 사고시 피해 최소화 대책", "psm.risk.mitigation"),
        ("위험성 평가 수행 및 평가절차", "psm.risk.procedure"),
    )
    for label, key in risk_items:
        _add_heading_safe(doc, label, level=2)
        doc.add_paragraph(_text(project, key, default=MISSING))

    specs = [s for s in selected_requirement_specs(project) if s.system == "PSM"]
    by_key = {s.key: s for s in specs}
    doc.add_page_break()
    _add_heading_safe(doc, "안전운전계획", level=1)
    for key in (
        "psm.operation.sop", "psm.operation.maintenance", "psm.operation.work_permit", "psm.operation.contractor",
        "psm.operation.training", "psm.operation.prestartup", "psm.operation.moc", "psm.operation.audit", "psm.operation.incident_investigation",
    ):
        spec = by_key.get(key)
        if spec is not None:
            _add_narrative_requirement(doc, spec, project, article_items=PSM_ARTICLE_ITEMS.get(key, ()))

    doc.add_page_break()
    _add_heading_safe(doc, "비상조치계획", level=1)
    emergency_specs = [s for s in specs if s.section == "비상조치계획"]
    if emergency_specs:
        for spec in emergency_specs:
            items = PSM_ARTICLE_ITEMS.get("psm.emergency.core", ()) if spec.key in {"psm.emergency.core", "psm.emergency.roles_procedures"} else ()
            _add_narrative_requirement(doc, spec, project, article_items=items)
    else:
        p = doc.add_paragraph()
        p.add_run("법정 작성항목: ").bold = True
        p.add_run(" / ".join(PSM_ARTICLE_ITEMS["psm.emergency.core"]))
        doc.add_paragraph(MISSING)


def _cap_form1(doc: Document, project: Stage2Project) -> None:
    _add_form_heading(doc, "별지 제1호서식", "사업장의 작성수준 구분")
    prepared = build_cap_form1_data(project)

    doc.add_paragraph("1. 단위공장 내 취급시설별 최대보유량 산출").runs[0].bold = True
    spec1 = FormSpec(
        "", "",
        ("단위공장", "유해화학물질", "CAS No.", "함량(%)", "구분기호", "취급시설", "설계용량(m3)", "취급량(ton)"),
    )
    t1_rows = [
        [
            str(row.get("단위공장") or ""),
            str(row.get("유해화학물질") or ""),
            str(row.get("CAS No.") or ""),
            str(row.get("함량(%)") or ""),
            str(row.get("구분기호") or ""),
            str(row.get("취급시설") or ""),
            str(row.get("설계용량(m3)") or ""),
            str(row.get("취급량(ton)") or ""),
        ]
        for row in prepared.facility_rows
    ]
    _add_form_table(doc, spec1, t1_rows or [[MISSING] + [""] * 7])

    p = doc.add_paragraph("2. 유해화학물질별 사업장 내의 최대보유량 산출")
    p.runs[0].bold = True
    spec2 = FormSpec(
        "", "",
        ("물질명", "CAS No.", "물질구분", "사업장 내 최대보유량(ton)", "작성수준", "하위규정수량(ton)", "상위규정수량(ton)"),
    )
    t2_rows = [
        [
            str(row.get("물질명") or ""),
            str(row.get("CAS No.") or ""),
            str(row.get("물질구분") or MISSING),
            str(row.get("사업장 내 최대보유량(ton)") or MISSING),
            str(row.get("작성수준") or project.cap_group or MISSING),
            str(row.get("하위규정수량(ton)") or MISSING),
            str(row.get("상위규정수량(ton)") or MISSING),
        ]
        for row in prepared.chemical_rows
    ]
    _add_form_table(doc, spec2, t2_rows or [[MISSING] + [""] * 6])

    doc.add_paragraph("3. 작성수준 도출: " + (project.cap_group or MISSING))
    if prepared.blockers:
        note = doc.add_paragraph()
        note.add_run("확인 필요: ").bold = True
        note.add_run(" / ".join(prepared.blockers))

def _cap_form3(doc: Document, project: Stage2Project) -> None:
    level = project.cap_group or _text(project, "cap.business.writing_level", default=MISSING)
    rows = (
        ("사업장명", project.company_name or MISSING),
        ("단위공장명", _text(project, "cap.business.unit_plant_name", default=MISSING)),
        ("사업자 등록번호", _text(project, "cap.business.registration_no", default=MISSING)),
        ("대표자", _text(project, "cap.business.representative", default=MISSING)),
        ("우편번호/주소", _text(project, "business.address", default=MISSING)),
        ("산업단지", _text(project, "cap.business.industrial_complex", default=MISSING)),
        ("대표전화", _text(project, "cap.business.contact", default=MISSING)),
        ("제출구분", _text(project, "cap.business.submission_type", default=MISSING)),
        ("작성수준", f"{_checkbox(level == '1군', '1군')}   {_checkbox(level == '2군', '2군')}"),
        ("공동비상대응계획 수립 여부", _text(project, "cap.business.joint_emergency_plan", default=MISSING)),
        ("유사제도 심사결과 활용", _text(project, "cap.business.other_system_review", default=MISSING)),
        ("총괄영향범위내 주민여부", _text(project, "cap.business.residents_in_overall_range", default=MISSING)),
        ("최근 3년간 화학사고 발생 여부", _text(project, "cap.business.recent_accident", default=MISSING)),
        ("화학사고예방관리계획서 작성자", _text(project, "cap.business.writer_info", default=MISSING)),
        ("담당자 연락처", _text(project, "cap.business.writer_contact", default=MISSING)),
        ("담당자 메일주소", _text(project, "cap.business.writer_email", default=MISSING)),
    )
    _add_key_value_form(doc, "별지 제3호서식", "사업장 일반정보", rows)


def _cap_facility_overview(doc: Document, project: Stage2Project, *, detailed: bool) -> None:
    facilities = _facility_rows(project)
    chemicals = _chemical_rows(project)
    chem_lines = []
    for row in chemicals[:12]:
        chem_lines.append(
            f"{_row_value(row, '물질명', '유해화학물질명')} / {_row_value(row, 'CAS 번호', '화학물질식별번호')} / {_row_value(row, '최대보유량', '최대보유량(kg)')}"
        )
    rows = (
        ("단위공장 구성", _text(project, "cap.basic.unit_facility_overview", default=MISSING)),
        ("공정개요", _text(project, "process.description", default=MISSING)),
        ("장치·설비 종류 및 수량", _facility_counts(facilities)),
        ("입·출하 및 운반시설", _text(project, "cap.basic.loading_transport", default=MISSING)),
        ("유해화학물질 및 취급량", "\n".join(chem_lines) if chem_lines else MISSING),
    )
    reference = "별지 제5호서식" if detailed else "별지 제4호서식"
    title = "세부 취급시설 개요" if detailed else "총괄 취급시설 개요"
    _add_key_value_form(doc, reference, title, rows)


def _cap_form6_rows(project: Stage2Project) -> list[list[str]]:
    legal = build_cap_chemical_legal_data(project)
    out = []
    for idx, row in enumerate(legal.rows, 1):
        out.append([
            str(idx),
            _row_value(row, "물질명", "유해화학물질명"),
            _row_value(row, "물질구분"),
            _row_value(row, "CAS 번호", "화학물질식별번호"),
            _row_value(row, "고유번호"),
            _row_value(row, "물리적 상태", "물질상태"),
            _row_value(row, "함량(%)", "함량"),
            _row_value(row, "비중"),
            _row_value(row, "폭발한계 하한", "폭발하한"),
            _row_value(row, "폭발한계 상한", "폭발상한"),
            _row_value(row, "독성구분 항목", "독성구분-항목"),
            _row_value(row, "독성구분", "독성구분-구분"),
            _row_value(row, "위험노출수준", "ERPG", "AEGL", "PAC", "IDLH"),
            _row_value(row, "허용농도값", "TWA", "노출기준"),
            _row_value(row, "증기압", "증기압(20℃, mmHg)"),
            _row_value(row, "부식성", "부식성(유, 무)"),
        ])
    return out


def _cap_form7(doc: Document, project: Stage2Project) -> None:
    source = _value(project, "cap.chemical.hazard_information", default=MISSING)
    items: list[Mapping[str, object]] = []
    if isinstance(source, list):
        items = [dict(v) for v in source if isinstance(v, Mapping)]
    elif isinstance(source, Mapping):
        items = [dict(source)]
    if not items:
        _add_key_value_form(doc, "별지 제7호서식", "유해화학물질의 유해성 정보", (("작성내용", MISSING),))
        return
    for idx, row in enumerate(items, 1):
        if idx > 1:
            doc.add_paragraph()
        _add_key_value_form(doc, "별지 제7호서식", "유해화학물질의 유해성 정보", (
            ("1. 취급물질의 일반정보", _row_value(row, "일반정보", "물질명")),
            ("가. 물질명", _row_value(row, "물질명")),
            ("나. 화학물질식별번호(CAS 번호)", _row_value(row, "CAS 번호", "화학물질식별번호")),
            ("다. 유해화학물질 고유번호", _row_value(row, "고유번호")),
            ("라. 농도(또는 함량 %)", _row_value(row, "농도", "함량")),
            ("마. 최대보유량", _row_value(row, "최대보유량")),
            ("2. 인체유해성", _row_value(row, "인체유해성")),
            ("3. 물리적 위험성", _row_value(row, "물리적 위험성")),
            ("4. 환경유해성", _row_value(row, "환경유해성")),
            ("5. 출처", _row_value(row, "출처")),
            ("6. 선정 사유", _row_value(row, "선정 사유", "선정사유")),
        ))


def _cap_form8(doc: Document, project: Stage2Project) -> None:
    _add_form_heading(doc, "별지 제8호서식", "사업장 주변 환경 정보")
    doc.add_paragraph("1. 사업장 경계선 500m 범위 내의 입지 현황")
    doc.add_paragraph(_text(project, "cap.site.surrounding_environment", default=MISSING))
    spec = FormSpec("", "", ("일련번호", "보호대상 종류", "보호대상 명칭", "거리(m)"))
    rows = _generic_form_rows(project, "cap.site.surrounding_environment", spec, (("일련번호", "연번"), ("보호대상 종류", "종류"), ("보호대상 명칭", "명칭"), ("거리", "거리(m)")))
    _add_form_table(doc, spec, rows)
    _add_attachment_line(doc, "보호대상 위치도", _doc_value(project, "cap.site.surrounding_map"))


def _cap_form9_rows(project: Stage2Project) -> list[list[str]]:
    prepared = build_cap_form9_data(project)
    out: list[list[str]] = []
    for row in prepared.rows:
        out.append([
            str(row.get("연번") or ""),
            str(row.get("구분기호") or ""),
            str(row.get("장치·설비명") or ""),
            str(row.get("취급물질") or ""),
            str(row.get("CAS No.") or ""),
            str(row.get("물질상태") or ""),
            str(row.get("함량(%)") or ""),
            str(row.get("연결구 크기(mm)") or ""),
            str(row.get("압력(MPa)-설계") or ""),
            str(row.get("압력(MPa)-운전") or ""),
            str(row.get("온도(℃)-설계") or ""),
            str(row.get("온도(℃)-운전") or ""),
            str(row.get("설계용량(m3)") or ""),
            str(row.get("취급량(ton)") or ""),
            str(row.get("비고") or ""),
        ])
    return out


def _cap_form10_rows(project: Stage2Project) -> list[list[str]]:
    prepared = build_cap_form10_data(project)
    out: list[list[str]] = []
    for row in prepared.rows:
        out.append([
            str(row.get("연번") or ""),
            str(row.get("설비형태") or ""),
            str(row.get("구분기호") or ""),
            str(row.get("장치·설비명") or ""),
            str(row.get("설계용량") or ""),
            str(row.get("설비종류") or ""),
            str(row.get("필요용량") or ""),
            str(row.get("유효용량") or ""),
            str(row.get("검토결과") or ""),
            str(row.get("비고") or ""),
        ])
    return out


def _cap_form11_rows(project: Stage2Project) -> list[list[str]]:
    prepared = build_cap_form11_data(project)
    out: list[list[str]] = []
    for row in prepared.rows:
        out.append([
            str(row.get("연번") or ""),
            str(row.get("구분기호") or ""),
            str(row.get("감지대상") or ""),
            str(row.get("설치위치") or ""),
            str(row.get("작동시간") or ""),
            str(row.get("측정방식") or ""),
            str(row.get("경보설정값") or ""),
            str(row.get("경보기 설치장소") or ""),
            str(row.get("연동여부") or ""),
            str(row.get("정밀도") or ""),
            str(row.get("유지관리") or ""),
            str(row.get("비고") or ""),
        ])
    return out


def _cap_form12(doc: Document, project: Stage2Project) -> None:
    prepared = build_cap_form12_data(project)
    _add_form_heading(doc, "별지 제12호서식", "사고시나리오 사업장 주변지역 영향 평가")
    if not prepared.rows:
        doc.add_paragraph(MISSING)
    for idx, row in enumerate(prepared.rows, 1):
        if idx > 1:
            doc.add_paragraph()
        _add_key_value_form(doc, "", f"{idx}) {row.get('사고시나리오명') or MISSING}", (
            ("유해화학물질명", str(row.get("유해화학물질명") or MISSING)),
            ("대상 설비번호", str(row.get("대상 설비번호") or MISSING)),
            ("사고유형", str(row.get("사고유형") or MISSING)),
            ("장외거리(m)", str(row.get("장외거리(m)") if row.get("장외거리(m)") not in (None, "") else MISSING)),
            ("영향범위 내 거주민수", str(row.get("거주민수") if row.get("거주민수") not in (None, "") else MISSING)),
            ("영향범위 내 근로자수", str(row.get("근로자수") if row.get("근로자수") not in (None, "") else MISSING)),
            ("갑종 보호대상 수", str(row.get("갑종 보호대상 수") if row.get("갑종 보호대상 수") not in (None, "") else MISSING)),
            ("을종 보호대상 수", str(row.get("을종 보호대상 수") if row.get("을종 보호대상 수") not in (None, "") else MISSING)),
            ("환경수용체 수", str(row.get("환경수용체 수") if row.get("환경수용체 수") not in (None, "") else MISSING)),
            ("사고원점의 좌표", str(row.get("사고원점 좌표") or MISSING)),
            ("KORA/GIS 근거", str(row.get("KORA/GIS 근거") or MISSING)),
        ))
    if prepared.blockers:
        note = doc.add_paragraph()
        note.add_run("확인 필요: ").bold = True
        note.add_run(" / ".join(prepared.blockers))


def _cap_form13(doc: Document, project: Stage2Project) -> None:
    prepared = build_cap_form13_data(project)
    _add_form_heading(doc, "별지 제13호서식", "총괄영향범위 사업장 주변지역 영향 평가")
    summary = prepared.summary or {}
    _add_key_value_form(doc, "", "총괄영향범위 확정정보", (
        ("총괄영향범위 산출방법", str(summary.get("총괄영향범위 산출방법") or MISSING)),
        ("총괄영향범위 결과 요약", str(summary.get("총괄영향범위 결과 요약") or MISSING)),
        ("총괄영향범위 내 거주민수", str(summary.get("총괄영향범위 내 거주민수") if summary.get("총괄영향범위 내 거주민수") not in (None, "") else MISSING)),
        ("총괄영향범위 내 근로자수", str(summary.get("총괄영향범위 내 근로자수") if summary.get("총괄영향범위 내 근로자수") not in (None, "") else MISSING)),
        ("보호대상 없음 여부", str(summary.get("보호대상 없음 여부") or MISSING)),
        ("GIS/KORA 근거", str(summary.get("GIS/KORA 근거") or MISSING)),
    ))
    spec = FormSpec(
        "", "총괄영향범위 내 보호대상",
        ("일련번호", "보호대상 명칭", "보호대상 구분", "보호대상 종류", "주소·위치", "사업장 경계와 거리(m)", "인원수", "GIS 근거"),
    )
    rows = [
        [
            str(row.get("일련번호") or ""),
            str(row.get("보호대상 명칭") or ""),
            str(row.get("보호대상 구분") or ""),
            str(row.get("보호대상 종류") or ""),
            str(row.get("주소·위치") or row.get("좌표") or ""),
            str(row.get("사업장 경계와 거리(m)") if row.get("사업장 경계와 거리(m)") not in (None, "") else ""),
            str(row.get("인원수") if row.get("인원수") not in (None, "") else ""),
            str(row.get("GIS 근거") or ""),
        ]
        for row in prepared.protected_targets
    ]
    if rows:
        _add_form_table(doc, spec, rows)
    elif prepared.no_protected_targets:
        doc.add_paragraph("총괄영향범위 내 보호대상 없음(회사/GIS 확정)")
    else:
        doc.add_paragraph(MISSING)
    _add_attachment_line(doc, "KORA/GIS 총괄영향범위 결과파일", _doc_value(project, "documents.kora_impact_result"))
    if prepared.blockers:
        note = doc.add_paragraph()
        note.add_run("확인 필요: ").bold = True
        note.add_run(" / ".join(prepared.blockers))



def _cap_form14(doc: Document, project: Stage2Project) -> None:
    prepared = build_cap_form14_data(project)
    _add_form_heading(doc, "별지 제14호서식", "사고시나리오별 시설빈도")
    grouped: dict[str, list[dict[str, Any]]] = {}
    for row in prepared.event_rows:
        grouped.setdefault(str(row.get("사고시나리오명") or ""), []).append(dict(row))

    if not grouped:
        doc.add_paragraph(MISSING)
    for scenario, rows in grouped.items():
        p = doc.add_paragraph()
        p.add_run(f"1) {scenario or MISSING}").bold = True
        spec = FormSpec("", "", ("연번", "개시사건", "빈도", "개수", "사고빈도"))
        values = [
            [
                str(idx),
                str(row.get("개시사건") or ""),
                str(row.get("기준빈도(/연)") or ""),
                str(row.get("개수") if row.get("개수") is not None else ""),
                str(row.get("사고빈도(/연)") or ""),
            ]
            for idx, row in enumerate(rows, 1)
        ]
        _add_form_table(doc, spec, values)
        summary = next(
            (item for item in prepared.scenario_rows if str(item.get("사고시나리오명") or "") == scenario),
            {},
        )
        _add_key_value_form(doc, "", "안전성 확보설비", (
            ("시설빈도 합(/연)", str(summary.get("시설빈도(/연)") or MISSING)),
            ("수동적 완화장치", str(summary.get("수동적 완화장치") or MISSING)),
            ("능동적 완화장치", str(summary.get("능동적 완화장치") or MISSING)),
            ("증빙자료", str(summary.get("안전성확보설비 증빙") or MISSING)),
        ))
    if prepared.blockers:
        note = doc.add_paragraph()
        note.add_run("확인 필요: ").bold = True
        note.add_run(" / ".join(prepared.blockers))


def _cap_form15(doc: Document, project: Stage2Project) -> None:
    prepared = build_cap_form15_data(project)
    _add_form_heading(doc, "별지 제15호서식", "위험도 분석")

    if prepared.no_offsite_scenario:
        doc.add_paragraph("장외 사고시나리오 없음: 제24조 및 제25조 작성 생략, 위험도 '다' 적용")
    else:
        spec = FormSpec(
            "", "1. 사업장 내 위험도 판단 요소 선정",
            ("연번", "사고시나리오 명", "사고시나리오 시설빈도", "사고시나리오 거리(장외)", "주민수"),
        )
        rows = [
            [
                str(row.get("연번") or ""),
                str(row.get("사고시나리오 명") or ""),
                str(row.get("사고시나리오 시설빈도") or ""),
                str(row.get("사고시나리오 거리(장외)") or ""),
                str(row.get("위험도 주민수") if row.get("위험도 주민수") not in (None, "") else ""),
            ]
            for row in prepared.scenario_rows
        ]
        _add_form_table(doc, spec, rows or [[MISSING] + [""] * 4])

    doc.add_paragraph("2. 위험도 판단 요소 점수").runs[0].bold = True
    totals = prepared.totals or {}
    scores = prepared.scores or {}
    _add_key_value_form(doc, "", "", (
        ("사고시나리오 총 개수(A)", str(totals.get("사고시나리오 총 개수(A)", MISSING))),
        ("사고시나리오 시설빈도의 합(B)", str(totals.get("사고시나리오 시설빈도의 합(B)", MISSING))),
        ("사고시나리오 거리의 합(C)", str(totals.get("사고시나리오 거리의 합(C)", MISSING))),
        ("주민수 합(D)", str(totals.get("주민수 합(D)", MISSING))),
        ("사고시나리오 개수 구간점수", str(scores.get("사고시나리오 개수 구간점수", MISSING))),
        ("시설빈도 구간점수", str(scores.get("시설빈도 구간점수", MISSING))),
        ("거리 구간점수", str(scores.get("거리 구간점수", MISSING))),
        ("주민수 구간점수", str(scores.get("주민수 구간점수", MISSING))),
        ("사고빈도점수(A+B)", str(scores.get("사고빈도점수(A+B)", MISSING))),
        ("사고영향점수(C+D)", str(scores.get("사고영향점수(C+D)", MISSING))),
        ("위험도 판정표 점수(증감 전)", str(scores.get("위험도 판정표 점수(증감 전)", MISSING))),
        ("증감 전 위험도", str(scores.get("증감 전 위험도", MISSING))),
        ("최종 위험도", str(scores.get("최종 위험도", MISSING))),
    ))
    if prepared.blockers:
        note = doc.add_paragraph()
        note.add_run("확인 필요: ").bold = True
        note.add_run(" / ".join(prepared.blockers))



def _cap_form16(doc: Document, project: Stage2Project) -> None:
    prepared = build_cap_form16_data(project)
    _add_form_heading(doc, "별지 제16호서식", "화학사고예방관리계획서 비상대응분야 요약서")

    business = prepared.business
    _add_key_value_form(doc, "", "1. 사업장 일반정보", (
        ("사업장명", str(business.get("사업장명") or MISSING)),
        ("대표자", str(business.get("대표자") or MISSING)),
        ("우편번호/주소", str(business.get("우편번호/주소") or MISSING)),
        ("사업자 등록번호", str(business.get("사업자 등록번호") or MISSING)),
        ("담당자", str(business.get("담당자") or MISSING)),
        ("담당자 연락처", str(business.get("담당자 연락처") or MISSING)),
        ("담당자 메일주소", str(business.get("담당자 메일주소") or MISSING)),
        ("작성일", str(business.get("작성일") or MISSING)),
    ))

    doc.add_paragraph("2. 사고시나리오 선정 유해화학물질 목록").runs[0].bold = True
    chemical_spec = FormSpec(
        "", "",
        ("연번", "유해화학물질명", "화학물질식별번호(CAS 번호)", "최대함량(%)", "최대보유량(ton)", "사고유형"),
    )
    chemical_rows = [
        [
            str(row.get("연번") or ""),
            str(row.get("유해화학물질명") or ""),
            str(row.get("화학물질식별번호(CAS 번호)") or ""),
            str(row.get("최대함량(%)") or ""),
            str(row.get("최대보유량(ton)") or ""),
            str(row.get("사고유형") or ""),
        ]
        for row in prepared.chemical_rows
    ]
    _add_form_table(doc, chemical_spec, chemical_rows or [[MISSING] + [""] * 5])

    doc.add_paragraph("3. 사고시나리오 및 위험도 핵심정보").runs[0].bold = True
    scenario_spec = FormSpec(
        "", "",
        ("연번", "사고시나리오명", "유해화학물질명", "대상 설비번호", "사고유형", "시설빈도(/연)", "장외거리(m)", "위험도 주민수", "KORA/GIS 근거"),
    )
    scenario_rows = [
        [
            str(row.get("연번") or ""),
            str(row.get("사고시나리오명") or ""),
            str(row.get("유해화학물질명") or ""),
            str(row.get("대상 설비번호") or ""),
            str(row.get("사고유형") or ""),
            str(row.get("시설빈도(/연)") or ""),
            str(row.get("장외거리(m)") if row.get("장외거리(m)") not in (None, "") else ""),
            str(row.get("위험도 주민수") if row.get("위험도 주민수") not in (None, "") else ""),
            str(row.get("KORA/GIS 근거") or ""),
        ]
        for row in prepared.scenario_rows
    ]
    if scenario_rows:
        _add_form_table(doc, scenario_spec, scenario_rows)
    else:
        doc.add_paragraph("장외 사고시나리오 없음 또는 시나리오 영향평가 확인 필요")

    doc.add_paragraph("4. 내부 비상대응 핵심내용").runs[0].bold = True
    _add_key_value_form(
        doc, "", "",
        tuple((label, value or MISSING) for label, value in prepared.internal_summary),
    )

    if project.cap_group == "1군":
        doc.add_paragraph("5. 외부 비상대응 핵심내용").runs[0].bold = True
        _add_key_value_form(
            doc, "", "",
            tuple((label, value or MISSING) for label, value in prepared.external_summary),
        )
    else:
        doc.add_paragraph("※ 2군 사업장은 외부 비상대응 요약항목을 작성대상에서 제외함.")

    if prepared.blockers:
        note = doc.add_paragraph()
        note.add_run("확인 필요: ").bold = True
        note.add_run(" / ".join(prepared.blockers))


def _render_cap(doc: Document, project: Stage2Project) -> None:
    _add_title(doc, "화학사고예방관리계획서", CAP_SOURCE, project, cap_group=project.cap_group)
    doc.add_heading("기본정보", level=1)
    _cap_form1(doc, project)
    _cap_form3(doc, project)
    _cap_facility_overview(doc, project, detailed=False)
    _cap_facility_overview(doc, project, detailed=True)
    _add_form_table(doc, CAP_FORMS["6"], _cap_form6_rows(project))
    _cap_form7(doc, project)
    _cap_form8(doc, project)

    doc.add_page_break()
    doc.add_heading("시설정보", level=1)
    doc.add_heading("공정개요", level=2)
    doc.add_paragraph(_text(project, "process.description", default=MISSING))
    _add_attachment_line(doc, "공정흐름도(PFD)", _doc_value(project, "documents.pfd"))
    _add_attachment_line(doc, "공정배관계장도(P&ID)", _doc_value(project, "documents.pid"))
    _add_form_table(doc, CAP_FORMS["9"], _cap_form9_rows(project))
    doc.add_heading("공정위험성 분석 자료", level=2)
    doc.add_paragraph(_text(project, "cap.facility.process_hazard_analysis", default=MISSING))
    doc.add_heading("운전책임자 및 작업자 현황", level=2)
    doc.add_paragraph(_text(project, "cap.facility.operator_staffing", default=MISSING))
    _add_form_table(doc, CAP_FORMS["10"], _cap_form10_rows(project))
    _add_form_table(doc, CAP_FORMS["11"], _cap_form11_rows(project))
    _add_attachment_line(doc, "안전밸브 및 파열판 명세", _text(project, "cap.safety.relief_device_specs", default=MISSING))
    _add_attachment_line(doc, "배출물질 처리시설 현황", _text(project, "cap.safety.waste_treatment", default=MISSING))

    doc.add_page_break()
    doc.add_heading("장외평가정보", level=1)
    doc.add_heading("예비시나리오 및 사고시나리오 선정", level=2)
    for label, key in (
        ("예비시나리오 대상 설비 선정", "cap.offsite.target_facility_selection"),
        ("대상 설비 취급량 산정", "cap.offsite.target_holding_calculation"),
        ("누출조건", "cap.offsite.release_conditions"),
        ("기상조건", "cap.offsite.weather_conditions"),
        ("영향범위 평가결과", "cap.offsite.impact_range_result"),
    ):
        p = doc.add_paragraph()
        p.add_run(label + ": ").bold = True
        p.add_run(_text(project, key, default=MISSING))
    _cap_form12(doc, project)
    _cap_form13(doc, project)
    _cap_form14(doc, project)
    _cap_form15(doc, project)

    specs = [s for s in selected_requirement_specs(project) if s.system == "CAP"]
    by_key = {s.key: s for s in specs}
    doc.add_page_break()
    doc.add_heading("사전관리방침", level=1)
    for key in (
        "cap.prevention.safety_management", "cap.prevention.training", "cap.prevention.self_inspection",
        "cap.prevention.change_management", "cap.prevention.emergency_system",
    ):
        spec = by_key.get(key)
        if spec is not None:
            _add_narrative_requirement(doc, spec, project, article_items=CAP_ARTICLE_ITEMS.get(key, ()))
    change_log = _rows(project, "cap.prevention.change_log")
    if change_log:
        spec = FormSpec("별지 제2호서식", "화학사고예방관리계획서 변경내역 관리대장", ("번호", "일자", "변경항목", "변경의 종류", "변경 내용(변경전 → 변경후)", "후속조치", "담당자"))
        rows = [[_row_value(row, h) for h in spec.headers] for row in change_log]
        _add_form_table(doc, spec, rows)

    doc.add_page_break()
    doc.add_heading("내부 비상대응계획", level=1)
    for key in ("cap.internal.accident_response", "cap.internal.post_accident"):
        spec = by_key.get(key)
        if spec is not None:
            _add_narrative_requirement(doc, spec, project, article_items=CAP_ARTICLE_ITEMS.get(key, ()))

    if project.cap_group == "1군":
        doc.add_page_break()
        doc.add_heading("외부 비상대응계획", level=1)
        for key in ("cap.external.community_coordination", "cap.external.evacuation", "cap.external.community_notice"):
            spec = by_key.get(key)
            if spec is not None:
                _add_narrative_requirement(doc, spec, project, article_items=CAP_ARTICLE_ITEMS.get(key, ()))

    doc.add_page_break()
    _cap_form16(doc, project)


def _add_review_appendix(doc: Document, project: Stage2Project, system: str, status) -> None:
    doc.add_page_break()
    _add_heading_safe(doc, "검토 참고사항 (법정서식 외)", level=1)
    doc.add_paragraph(
        "아래 내용은 자동작성 상태를 확인하기 위한 내부 검토용 정보이며 법정서식의 일부가 아닙니다. "
        "최종 제출 전 회사 담당자가 [확인 필요], 도면·첨부자료 및 AI 보강문장을 모두 검토해야 합니다."
    )
    table = doc.add_table(rows=4, cols=2)
    _set_table_borders(table)
    info = (
        ("프로젝트 ID", project.project_id),
        ("작성완성도", f"{status.completion_pct:.1f}%"),
        ("작성상태", status.state),
        ("미해결 필드", str(status.unresolved_n)),
    )
    for i, (label, value) in enumerate(info):
        _set_cell_text(table.rows[i].cells[0], label, bold=True, size=8)
        _set_cell_text(table.rows[i].cells[1], value, size=8)


def build_statutory_report_draft(project: Stage2Project, system: str, status) -> bytes:
    system = str(system or "").strip().upper()
    if system == "PSM" and not project.psm_in_scope:
        raise ValueError("공정안전보고서는 현재 작성범위에 포함되어 있지 않습니다.")
    if system == "CAP" and not project.cap_in_scope:
        raise ValueError("화학사고예방관리계획서는 현재 작성범위에 포함되어 있지 않습니다.")
    if system not in {"PSM", "CAP"}:
        raise ValueError(f"지원하지 않는 보고서 종류입니다: {system}")

    doc = Document()
    _configure_doc(doc)
    if system == "PSM":
        _render_psm(doc, project)
    else:
        _render_cap(doc, project)
    _add_review_appendix(doc, project, system, status)
    out = BytesIO()
    doc.save(out)
    return out.getvalue()
