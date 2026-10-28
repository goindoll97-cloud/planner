"""주소·도면·이미지 없이 파일 접수부터 서술형·DOCX까지 연습하는 정상 자료.

기존 오류 탐지용 fixture는 보존합니다. 모든 회사·인물·연락망·수치·조사값은 가상이며
물질명/CAS만 실제입니다. 법적 대상·현장조사·적합통보를 증명하지 않습니다.
"""
from functools import lru_cache
from io import BytesIO, StringIO
from zipfile import ZipFile, ZIP_DEFLATED
import csv
import json

from . import demo_full_company as base
from . import cap_fact_catalog as catalog

COMPANY = "가상 새솔화학(주)"
ZIP_NAME = "화사계_정상작성연습_주소도면제외.zip"

EXTERNAL = {
    "community": {"대상": "가상 인근 주민대표 및 근로자", "방법": "주민대표 간담회와 안내문 배포", "주기": "분기 1회 및 공정 변경 시"},
    "mutual_aid": {"기관": "가상 지역 소방대·새솔 협력사업장(연락반 000-000-0010)", "협력내용": "흡착포 200장·이동식 펌프 1대 상호지원", "합동훈련": "연 1회 누출 대응 합동훈련"},
    "evacuation": {"경보": "주민대표 연락망·확성기 방송", "경로": "바람과 누출 상황을 확인한 후 지정 통로로 이동(연습용 가정)", "장소": "가상 외부 집결지 A(위치·수용성 미검증)"},
    "public_notice": {"대상": "가상 영향범위 내 주민대표", "방법": "연락반의 문자·전화 전달", "내용": "취급물질·사고상황·주민 행동요령·문의 연락반"},
    "investigation": {"책임": "안전환경팀장", "방법": "사고기록·운전기록·관계자 진술 대조", "재발방지": "원인별 개선조치와 담당자·기한 기록 후 교육"},
}


def _csv(headers, rows):
    out = StringIO()
    writer = csv.writer(out)
    writer.writerow(headers)
    writer.writerows(rows)
    return out.getvalue().encode("utf-8-sig")


def _pdf(pairs, headers=("항목", "값")):
    from reportlab.pdfgen import canvas
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.cidfonts import UnicodeCIDFont
    from reportlab.platypus import Table, TableStyle
    from reportlab.lib import colors
    pdfmetrics.registerFont(UnicodeCIDFont("HYSMyeongJo-Medium"))
    out = BytesIO()
    page = canvas.Canvas(out, pagesize=(842, 1191))
    rows = [list(headers), *[list(row) for row in pairs]]
    # 실제 문자 표로 생성: 스캔 이미지와 OCR이 필요하지 않습니다.
    for start in range(0, len(rows), 28):
        table = Table(rows[start:start + 28], colWidths=[190, 570] if len(headers) == 2 else [760 / len(headers)] * len(headers), rowHeights=29)
        table.setStyle(TableStyle([("FONTNAME", (0, 0), (-1, -1), "HYSMyeongJo-Medium"),
                                  ("FONTSIZE", (0, 0), (-1, -1), 9),
                                  ("GRID", (0, 0), (-1, -1), .5, colors.grey)]))
        _, height = table.wrap(760, 1050)
        table.drawOn(page, 40, 1130 - height)
        page.showPage()
    page.save()
    return out.getvalue()


def _pairs(groups):
    return [(catalog.get(f"narrative.{key}.{label}").template_header, value)
            for key, values in groups.items() for label, value in values.items()]


def _xlsx(raw, edit):
    from openpyxl import load_workbook
    book = load_workbook(BytesIO(raw))
    edit(book)
    out = BytesIO()
    book.save(out)
    return out.getvalue()


@lru_cache(maxsize=1)
def _files():
    from openpyxl import load_workbook
    source = base._files()  # 텍스트·표만 생성하는 기존 데이터 생성기를 재사용
    files = {"01_제품목록.xlsx": source["01_제품목록.xlsx"]}
    book = load_workbook(BytesIO(source["02_설비대장.xlsx"]))
    rows = list(book.active.values)
    files["02_설비대장.csv"] = _csv([*rows[0], "단위공장·공정"], [[*r, "제1공장"] for r in rows[1:]])
    c = base.DATA["company"]
    files["03_사업장기본정보.docx"] = base._doc(COMPANY + " 가상 기본정보", [("항목·값", [
        ("사업장명", COMPANY), ("사업자 등록번호", c["registration_no"]), ("대표자", c["representative"]),
        ("산업단지", "가상 산업단지(소재지 미지정)"), ("대표전화", c["contact"]),
        ("공동비상대응계획 수립 여부", "단독제출"), ("유사제도 심사결과 활용", "미해당"),
        ("총괄영향범위내 주민여부", "있음"), ("최근 3년간 화학사고 발생 여부", "없음"),
        ("작성자 소속 부서", "안전환경팀"), ("작성자 성명", c["writer_name"]),
        ("담당자 연락처", c["writer_contact"]), ("담당자 메일주소", c["writer_email"]),
        ("제출구분", "신규제출"), ("제출 사유", "가상 회사 정상작성 연습"), ("단위공장명", c["단위공장명"])])])
    files["04_공정개요.hwpx"] = base._hwpx_from_docx(source["04_공정설명서.docx"])
    sds_names = {p["props"]["SDS 파일명"]: f"{13 + i:02d}_가상SDS_{p['code']}_발췌.pdf"
                 for i, p in enumerate(base.DATA["products"])}
    def sds_links(book):
        for sheet in book:
            for row in sheet:
                for cell in row:
                    if cell.value in sds_names:
                        cell.value = sds_names[cell.value]
    files["05_SDS물성_위험성.xlsx"] = _xlsx(source["05_SDS_물성표.xlsx"], sds_links)
    def clean_specs(book):
        for sheet in book:
            for cell in sheet[1]:
                if cell.value == "보관계획도 최대량":
                    sheet.delete_cols(cell.column)
                    break
            for row in sheet:
                for cell in row:
                    if isinstance(cell.value, str) and "P&ID" in cell.value:
                        cell.value = "가상 설비명세(도면 제외)"
        sheet = book["별지11_가스감지기"]
        column = sheet.max_column + 1
        sheet.cell(1, column, "설치형태")
        for row in range(2, sheet.max_row + 1):
            sheet.cell(row, column, "고정식")
    files["06_설비사양_안전설비.xlsx"] = _xlsx(source["06_설비사양_안전설비.xlsx"], clean_specs)
    def complete_scenarios(book):
        sheet = book["별지12_추가값"]
        first = sheet.max_column + 1
        for col, header in enumerate(("비열비", "시나리오 선정 근거"), first):
            sheet.cell(1, col, header)
        for row in range(2, sheet.max_row + 1):
            sheet.cell(row, first, 1.31 if sheet.cell(row, 1).value == "CY-01" else 1.1)
            sheet.cell(row, first + 1, "가상 검토대상 설비별 대표 사고유형(작성 연습용)")
        freq = book["별지14_개시사건"]
        for row in range(2, freq.max_row + 1):
            for col in range(2, 12):
                if freq.cell(row, col).value in (None, ""):
                    freq.cell(row, col, 0)
    files["07_사고시나리오_검토자료.xlsx"] = _xlsx(source["08_사고시나리오_검토자료.xlsx"], complete_scenarios)
    # 완성된 계획서 대신 연락체계·자원 등 사실만 제공합니다.
    emergency = base.DATA["emergency"]
    files["08_비상연락_대응자원.docx"] = base._doc("가상 비상연락·대응자원", [("항목·값", [
        ("비상연락체계", emergency["contact"]), ("방재 인력·장비·물품", emergency["personnel"]),
        ("취급시설 유형별 응급조치", emergency["first_aid"])])])
    decisions = dict(base.DATA["narrative"])
    recovery = {"recovery": decisions.pop("recovery"), "investigation": EXTERNAL["investigation"]}
    files["09_안전관리_교육점검_기초정보.pdf"] = _pdf(_pairs(decisions))
    external = {k: v for k, v in EXTERNAL.items() if k != "investigation"}
    doc = base._doc("1군 외부대응 단답형 정보(가상)", [("항목·값", _pairs(external))])
    from .hwp5_binary import from_hwpx
    files["10_지역사회_공조대피_기초정보.hwp"] = from_hwpx(base._hwpx_from_docx(doc))
    files["11_사고조사_복구_기초정보.txt"] = ("가상 회사 사실(연습용)\n" + "\n".join(a + ": " + b for a, b in _pairs(recovery))).encode("utf-8")
    # 숫자만으로 집계 연습합니다. 지도 확인·현장 조사 완료로 처리하지 않습니다.
    targets = [{k: v for k, v in row.items() if k != "주소·위치"} for row in base.DATA["surroundings"]]
    for row in targets:
        row["GIS/현장 근거"] = "연습용 가정 수치(지도·현장 미확인)"
    files["12_주변보호대상_가상수치.csv"] = _csv(tuple(targets[0]), [list(r.values()) for r in targets])
    for product in base.DATA["products"]:
        name = sds_names[product["props"]["SDS 파일명"]]
        files[name] = _pdf([(product["code"], product["name"], product["cas"], product["props"]["분자량"])],
                           headers=("제품코드", "제품명", "CAS 번호", "분자량"))
    return files


def files():
    return dict(_files())


README = f"""{COMPANY} — 정상작성 연습 자료 (주소·도면·이미지 제외)
모든 회사·인물·연락처·설비·운전조건·조사값은 가상입니다. 실제 운전이나 법정 제출에 사용하지 마세요.

사용 방법
1. ‘새 사업장 만들기’에서 ‘가상 1군 사업장으로 연습 시작’을 누르세요. 사업장 주소 없이 전용 연습 사업장이 만들어집니다.
2. 회사자료/의 01~15 파일을 한 번에 업로드하고 ‘자료 읽기’를 누르세요. 파일의 종류는 내용으로 인식합니다.
3. ②의 실제 보완 사항을 확인하세요. 이 묶음에는 오류 탐지용 구버전·충돌·스캔본을 넣지 않았습니다.
4. ④ 작성·내려받기에서 작성하고, 연결된 로컬 AI가 있으면 서술형만 다듬으세요.
   1군 외부 소통·공조·대피 문장까지 보려면 1군으로 설정된 연습 사업장이 필요합니다.
   자료에 1군 적합통보나 법적 대상 판정을 조작해 넣지 않았습니다.
5. 규정서식 초안 DOCX를 내려받고 회사 사실이 반영되었는지 확인하세요.
6. 제출·이행 관리에서 제출본을 저장·비교하는 연습도 할 수 있습니다. 실제 접수는 수행하지 않습니다.

제외·제한
- 사업장 주소, 좌표, 부지경계, 도면, 지도, 사진, 스캔 이미지: 제공하지 않습니다.
- 주변보호대상 거리·경계까지 거리·대피장소는 가상 수치와 명칭이며 현장 검증을 뜻하지 않습니다.
- 13~15 PDF는 식별·분자량을 담은 가상 SDS 발췌 표이며 정식 16개 항목의 SDS가 아닙니다.
- SDS 물성은 계산 연습용입니다. 물질/CAS는 실제여도 승인된 SDS나 규정 DB를 대체하지 않습니다.
- 승인 규정 DB·법적 대상 판정·현장 조사·도면·주소가 필요한 항목은 보완으로 남는 것이 정상입니다.
- 따라서 이 자료로 생성되는 것은 작성 연습용 DOCX 초안이며 법정 제출 가능한 최종 완성본이 아닙니다.
- 소통·공조·대피계획 등 완성 문장은 요구하지 않습니다. 09~11의 단답형 사실로 프로그램이 작성합니다.
- 파일을 수정할 때는 같은 파일명으로 다시 올리세요. 수정 사유를 붙일 필요가 없습니다.
"""


@lru_cache(maxsize=1)
def archive():
    out = BytesIO()
    with ZipFile(out, "w", ZIP_DEFLATED) as z:
        for name, content in files().items():
            z.writestr("회사자료/" + name, content)
        z.writestr("읽어주세요.txt", README)
        z.writestr("자료목록.json", json.dumps({"company": COMPANY, "files": list(files()),
                    "intended_group": "1군 서술형 연습(법적 판정 아님)",
                    "excluded": ["사업장 주소", "도면", "이미지", "좌표", "부지경계"],
                    "final_submission_ready": False}, ensure_ascii=False, indent=2))
    return out.getvalue()


if __name__ == "__main__":
    import argparse
    from pathlib import Path
    parser = argparse.ArgumentParser(description="주소·도면·이미지를 제외한 정상 작성 연습자료 ZIP 생성")
    parser.add_argument("--output-dir", default="outputs/authoring_demo")
    args = parser.parse_args()
    directory = Path(args.output_dir)
    directory.mkdir(parents=True, exist_ok=True)
    destination = directory / ZIP_NAME
    destination.write_bytes(archive())
    print(destination)
