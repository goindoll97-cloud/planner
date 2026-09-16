from __future__ import annotations

"""PSM statutory-form writer for the user-supplied replacement DOCX baseline.

This adapter keeps legal-currentness decisions outside the DOCX template.  It
only maps already-confirmed project data into the preserved form layout and
fails closed when the pinned table geometry no longer matches the approved
baseline snapshot.
"""

from io import BytesIO
from typing import Mapping

from docx import Document

from . import psm_baseline_docx as legacy
from .project import Stage2Project


EXPECTED_TABLE_ROWS = (21, 11, 12, 12, 13, 11, 11, 12, 11, 14, 12, 15, 33, 15, 11)


def _validate_replacement_layout(doc: Document) -> None:
    if len(doc.tables) != len(legacy.FORM_TABLE_INDEX):
        raise ValueError("PSM 규정서식 baseline 구조가 예상과 달라 자동작성을 중단했습니다.")
    actual_rows = tuple(len(table.rows) for table in doc.tables)
    if actual_rows != EXPECTED_TABLE_ROWS:
        raise ValueError(
            "PSM 규정서식 baseline의 행 구조가 승인된 교체 서식과 다릅니다: "
            f"{actual_rows}"
        )


def _schedule_parts(project: Stage2Project) -> tuple[str, str, str]:
    value = legacy.base._value(project, "psm.business.schedule", default="")
    if not isinstance(value, Mapping):
        return legacy._clean(value), "", ""

    normalized = {legacy.base._norm(key): val for key, val in value.items()}

    def pick(*keys: str) -> str:
        for key in keys:
            candidate = normalized.get(legacy.base._norm(key))
            if candidate not in (None, "", legacy.MISSING):
                return legacy._clean(candidate)
        return ""

    return (
        pick("총사업기간", "사업기간", "total_period", "project_period"),
        pick("착공예정일", "착공일", "construction_start", "start_date"),
        pick("시운전기간", "시운전", "commissioning_period", "trial_operation_period"),
    )


def _fill_form12(table, project: Stage2Project) -> None:
    """Fill 별지 제12호 using the replacement form's exact semantic rows."""
    chemicals = legacy.base._rows(project, "inventory.chemicals")
    raw_materials = ", ".join(
        value
        for row in chemicals[:8]
        if (
            value := legacy._clean(
                legacy.base._row_value(row, "물질명", "화학물질", "유해화학물질명")
            )
        )
    )
    writer, qualification = legacy._writer_parts(project)

    legacy._append_value(legacy._unique_cells(table.rows[3])[0], project.company_name)
    legacy._write_cell(legacy._unique_cells(table.rows[3])[2], legacy._project_type_options(project))
    legacy._append_value(
        legacy._unique_cells(table.rows[4])[0],
        legacy._field_text(project, "business.registration_no", "cap.business.registration_no"),
    )
    legacy._append_value(
        legacy._unique_cells(table.rows[5])[0],
        legacy._field_text(project, "business.representative", "cap.business.representative"),
    )
    legacy._write_cell(
        legacy._unique_cells(table.rows[5])[2],
        legacy._field_text(project, "psm.business.target_facility"),
    )
    legacy._append_value(
        legacy._unique_cells(table.rows[6])[0],
        legacy._field_text(project, "business.ksic"),
    )
    legacy._append_value(
        legacy._unique_cells(table.rows[7])[0],
        legacy._field_text(project, "business.employee_count"),
    )
    electric = legacy._field_text(project, "business.electric_contract_capacity")
    if electric:
        legacy._write_cell(legacy._unique_cells(table.rows[7])[2], f"{electric} ㎾")

    legacy._write_cell(legacy._unique_cells(table.rows[8])[1], writer)
    legacy._write_cell(legacy._unique_cells(table.rows[8])[3], qualification)
    legacy._write_cell(legacy._unique_cells(table.rows[11])[2], raw_materials)
    legacy._write_cell(
        legacy._unique_cells(table.rows[12])[2],
        legacy._field_text(project, "business.main_products"),
    )
    legacy._write_cell(
        legacy._unique_cells(table.rows[13])[2],
        legacy._field_text(project, "psm.business.overview"),
    )

    # Replacement layout: 위치=14, 부지=15, 주요건물=16.
    legacy._write_cell(
        legacy._unique_cells(table.rows[14])[2],
        legacy._field_text(project, "business.address"),
    )
    legacy._write_cell(
        legacy._unique_cells(table.rows[15])[2],
        legacy._field_text(project, "psm.business.site_area", "business.site_area"),
    )
    legacy._write_cell(
        legacy._unique_cells(table.rows[16])[2],
        legacy._field_text(project, "psm.business.site_building"),
    )

    total_period, construction_start, commissioning_period = _schedule_parts(project)
    legacy._write_cell(legacy._unique_cells(table.rows[17])[2], total_period)
    legacy._write_cell(legacy._unique_cells(table.rows[18])[2], construction_start)
    legacy._write_cell(legacy._unique_cells(table.rows[19])[2], commissioning_period)


def build_psm_baseline_draft(project: Stage2Project) -> bytes:
    if not project.psm_in_scope:
        raise ValueError("공정안전보고서는 현재 작성범위에 포함되어 있지 않습니다.")

    doc = Document(BytesIO(legacy.load_psm_baseline_bytes()))
    _validate_replacement_layout(doc)

    _fill_form12(doc.tables[legacy.FORM_TABLE_INDEX["12"]], project)
    legacy._fill_table_rows(
        doc.tables[legacy.FORM_TABLE_INDEX["13"]],
        legacy._sanitize_rows(legacy.base._psm_form13_rows(project)),
    )
    legacy._fill_table_rows(
        doc.tables[legacy.FORM_TABLE_INDEX["14"]],
        legacy._sanitize_rows(legacy.base._psm_form14_rows(project)),
    )
    legacy._fill_table_rows(
        doc.tables[legacy.FORM_TABLE_INDEX["15"]],
        legacy._sanitize_rows(legacy.base._psm_form15_rows(project)),
    )
    legacy._fill_table_rows(
        doc.tables[legacy.FORM_TABLE_INDEX["16"]],
        legacy._sanitize_rows(legacy.base._psm_form16_rows(project)),
    )
    legacy._fill_table_rows(
        doc.tables[legacy.FORM_TABLE_INDEX["17"]],
        legacy._sanitize_rows(legacy.base._psm_form17_rows(project)),
    )

    for form_key, field_keys in legacy.LATER_FORM_FIELDS.items():
        rows = legacy._structured_rows(project, form_key, field_keys)
        legacy._fill_table_rows(doc.tables[legacy.FORM_TABLE_INDEX[form_key]], rows)

    legacy._fill_form19_2(doc.tables[legacy.FORM_TABLE_INDEX["19-2"]], project)

    out = BytesIO()
    doc.save(out)
    return out.getvalue()


def psm_baseline_filename(project: Stage2Project) -> str:
    return legacy.psm_baseline_filename(project)
