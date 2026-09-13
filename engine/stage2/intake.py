from __future__ import annotations

from dataclasses import asdict, dataclass
from io import BytesIO
import re
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

from engine.legal_terminology import psm_field_label

from .project import CONFIRMED_STATUSES, Stage2Project
from .requirements import (
    COMMON_REQUIREMENTS,
    RequirementSpec,
    cap_field_labels,
    cap_requirement_specs,
    psm_field_labels,
    psm_requirement_specs,
)


COVERAGE_CONFIRMED = "확인"
COVERAGE_PARTIAL = "일부 확인"
COVERAGE_REVIEW = "사람 확인 필요"
COVERAGE_REQUEST = "자료 요청 필요"
COVERAGE_NOT_APPLICABLE = "해당 없음"

SYSTEM_LABELS = {
    "COMMON": "공통자료",
    "PSM": "공정안전보고서",
    "CAP": "화학사고예방관리계획서",
}


@dataclass(frozen=True)
class IntakeRequirement:
    requirement_key: str
    system: str
    system_label: str
    section: str
    label: str
    coverage_status: str
    confirmed_fields: tuple[str, ...]
    received_unconfirmed_fields: tuple[str, ...]
    missing_fields: tuple[str, ...]
    confirmed_labels: tuple[str, ...]
    received_unconfirmed_labels: tuple[str, ...]
    missing_labels: tuple[str, ...]
    suggested_evidence: tuple[str, ...]
    legal_basis: str
    form_references: tuple[str, ...]
    reference_pages: tuple[int, ...]
    reference_label: str
    request_text: str
    input_kind: str
    automation: str
    legal_status: str

    def to_dict(self) -> dict[str, Any]:
        return asdict(self)



def selected_requirement_specs(project: Stage2Project) -> list[RequirementSpec]:
    """Requirements for the user's selected authoring scope only.

    Stage 1 legal applicability is deliberately separate from the Stage 2
    authoring choice. Common requirements are included only after the user has
    selected at least one document to prepare.
    """
    if not project.scope_confirmed:
        return []
    if not project.psm_in_scope and not project.cap_in_scope:
        return []

    specs: list[RequirementSpec] = list(COMMON_REQUIREMENTS)
    if project.psm_in_scope:
        specs.extend(psm_requirement_specs())
    if project.cap_in_scope:
        group = project.cap_group if project.cap_group in {"1군", "2군"} else "1군"
        specs.extend(cap_requirement_specs(group))
    return specs



def _field_labels() -> dict[str, str]:
    labels: dict[str, str] = {
        "business.company_name": "회사명",
        "business.address": "사업장 소재지",
        "inventory.chemicals": "화학물질 목록",
        "inventory.facilities": "시설·설비 목록",
        "process.description": "공정설명서",
        "documents.pfd": "공정흐름도(PFD)",
        "documents.pid": "공정배관·계장도(P&ID)",
        "documents.site_plan": "각종 건물·설비의 배치도",
    }
    labels.update(cap_field_labels())
    psm_labels = psm_field_labels()
    labels.update({key: psm_field_label(key, value) for key, value in psm_labels.items()})
    return labels



def field_label(field_key: str) -> str:
    return _field_labels().get(field_key, field_key)



def extract_form_references(text: str) -> tuple[str, ...]:
    """Extract statutory annex/form references without inventing a form."""
    source = str(text or "")
    found: list[str] = []

    # 별지 제4호·제5호서식 같은 표현을 지원한다.
    for segment in re.findall(r"별지\s*([^;\n]*)", source):
        for number in re.findall(r"제\s*(\d+)호", segment):
            label = f"별지 제{number}호서식"
            if label not in found:
                found.append(label)

    for number in re.findall(r"별표\s*제?\s*(\d+)호?", source):
        label = f"별표 {number}"
        if label not in found:
            found.append(label)

    return tuple(found)



def _reference_label(spec: RequirementSpec) -> str:
    if spec.system == "CAP" and spec.manual_pages:
        return "화학사고예방관리계획서 작성 매뉴얼"
    if spec.system == "PSM" and spec.manual_pages:
        return "공정안전보고서 지원시스템 작성예시집(형식 참고)"
    return ""



def _coverage(project: Stage2Project, spec: RequirementSpec) -> tuple[str, list[str], list[str], list[str]]:
    if not spec.required and spec.legal_status != "VERIFY_CURRENT":
        return COVERAGE_NOT_APPLICABLE, [], [], []

    confirmed: list[str] = []
    received_unconfirmed: list[str] = []
    missing: list[str] = []

    if not spec.field_keys:
        return COVERAGE_REVIEW, [], [], []

    for key in spec.field_keys:
        record = project.get_field(key)
        if record is None:
            missing.append(key)
            continue
        if record.status in CONFIRMED_STATUSES:
            confirmed.append(key)
            continue
        if record.evidence or record.status == "AI_DRAFT":
            received_unconfirmed.append(key)
        else:
            missing.append(key)

    if len(confirmed) == len(spec.field_keys):
        status = COVERAGE_CONFIRMED
    elif confirmed:
        status = COVERAGE_PARTIAL
    elif received_unconfirmed:
        status = COVERAGE_REVIEW
    else:
        status = COVERAGE_REQUEST
    return status, confirmed, received_unconfirmed, missing



def build_intake_catalog(project: Stage2Project) -> list[IntakeRequirement]:
    labels = _field_labels()
    rows: list[IntakeRequirement] = []
    for spec in selected_requirement_specs(project):
        status, confirmed, received_unconfirmed, missing = _coverage(project, spec)
        rows.append(
            IntakeRequirement(
                requirement_key=spec.key,
                system=spec.system,
                system_label=SYSTEM_LABELS.get(spec.system, spec.system),
                section=spec.section,
                label=spec.label,
                coverage_status=status,
                confirmed_fields=tuple(confirmed),
                received_unconfirmed_fields=tuple(received_unconfirmed),
                missing_fields=tuple(missing),
                confirmed_labels=tuple(labels.get(key, key) for key in confirmed),
                received_unconfirmed_labels=tuple(labels.get(key, key) for key in received_unconfirmed),
                missing_labels=tuple(labels.get(key, key) for key in missing),
                suggested_evidence=spec.suggested_evidence,
                legal_basis=spec.legal_basis,
                form_references=extract_form_references(spec.legal_basis),
                reference_pages=spec.manual_pages,
                reference_label=_reference_label(spec),
                request_text=spec.request_text or spec.description,
                input_kind=spec.input_kind,
                automation=spec.automation,
                legal_status=spec.legal_status,
            )
        )
    return rows



def intake_summary(project: Stage2Project) -> dict[str, Any]:
    rows = build_intake_catalog(project)
    counts = {
        COVERAGE_CONFIRMED: 0,
        COVERAGE_PARTIAL: 0,
        COVERAGE_REVIEW: 0,
        COVERAGE_REQUEST: 0,
        COVERAGE_NOT_APPLICABLE: 0,
    }
    for row in rows:
        counts[row.coverage_status] = counts.get(row.coverage_status, 0) + 1
    return {
        "scope_confirmed": project.scope_confirmed,
        "psm_in_scope": project.psm_in_scope,
        "cap_in_scope": project.cap_in_scope,
        "counts": counts,
        "rows": [row.to_dict() for row in rows],
    }



def build_missing_request_text(item: IntakeRequirement) -> str:
    if item.coverage_status == COVERAGE_CONFIRMED:
        return "현재 등록자료에서 필요한 내용이 확인되었습니다."
    if item.coverage_status == COVERAGE_NOT_APPLICABLE:
        return "현재 선택한 작성범위와 적용조건에서 필수 작성항목으로 보지 않습니다."

    parts = [item.request_text or f"{item.label} 관련 자료를 확인해 주세요."]
    if item.received_unconfirmed_labels:
        parts.append(
            "자료는 접수되었으나 내용을 확정하지 못한 항목: "
            + ", ".join(item.received_unconfirmed_labels)
        )
    if item.missing_labels:
        parts.append("현재 확인되지 않은 항목: " + ", ".join(item.missing_labels))
    if item.suggested_evidence:
        parts.append("확인 가능한 자료 예: " + ", ".join(item.suggested_evidence))
    if item.form_references:
        parts.append("관련 법정 서식: " + ", ".join(item.form_references))
    if item.legal_basis:
        parts.append("작성근거: " + item.legal_basis)
    if item.reference_label and item.reference_pages:
        parts.append(
            f"작성 참고자료: {item.reference_label} 관련 쪽 "
            + ", ".join(str(v) for v in item.reference_pages)
        )
    return "\n".join(parts)



def _autosize(ws) -> None:
    for column_cells in ws.columns:
        width = 10
        letter = get_column_letter(column_cells[0].column)
        for cell in column_cells:
            value = "" if cell.value is None else str(cell.value)
            width = max(width, min(len(value) + 2, 60))
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        ws.column_dimensions[letter].width = width



def _append_table(ws, headers: list[str], rows: list[list[Any]]) -> None:
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for row in rows:
        ws.append(row)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    _autosize(ws)



def build_program_input_workbook(project: Stage2Project) -> bytes:
    """Create a convenience workbook, explicitly not a statutory form."""
    catalog = build_intake_catalog(project)
    wb = Workbook()
    guide = wb.active
    guide.title = "안내"
    guide.append(["구분", "내용"])
    guide.append(["문서 성격", "프로그램 입력양식(법정 별지서식이 아님)"])
    guide.append(["목적", "회사가 보유한 자료와 부족한 자료를 작성항목별로 정리하기 위한 보조 입력양식"])
    guide.append(["주의", "법정 별지서식이 있는 항목은 현행 법령·고시의 공식 서식을 별도로 확인해야 합니다."])
    guide.append(["공정안전보고서 작성범위", "선택" if project.psm_in_scope else "미선택"])
    guide.append(["화학사고예방관리계획서 작성범위", "선택" if project.cap_in_scope else "미선택"])
    _autosize(guide)

    ws = wb.create_sheet("요구자료목록")
    rows: list[list[Any]] = []
    for item in catalog:
        rows.append([
            item.system_label,
            item.section,
            item.label,
            item.coverage_status,
            ", ".join(item.missing_labels),
            ", ".join(item.suggested_evidence),
            ", ".join(item.form_references),
            item.legal_basis,
            item.reference_label,
            ", ".join(str(v) for v in item.reference_pages),
            item.request_text,
        ])
    _append_table(
        ws,
        [
            "구분",
            "작성구조",
            "작성항목",
            "자료상태",
            "현재 확인되지 않은 항목",
            "확인 가능한 자료 예",
            "관련 법정 서식",
            "작성근거",
            "작성 참고자료",
            "관련 쪽",
            "요청사항",
        ],
        rows,
    )

    input_ws = wb.create_sheet("입력항목")
    input_rows: list[list[Any]] = []
    for item in catalog:
        spec_fields = item.confirmed_fields + item.received_unconfirmed_fields + item.missing_fields
        if not spec_fields:
            input_rows.append([
                item.system_label,
                item.section,
                item.label,
                "",
                "",
                item.legal_basis,
                ", ".join(item.form_references),
            ])
            continue
        for field_key in spec_fields:
            input_rows.append([
                item.system_label,
                item.section,
                item.label,
                field_label(field_key),
                "",
                item.legal_basis,
                ", ".join(item.form_references),
            ])
    _append_table(
        input_ws,
        ["구분", "작성구조", "작성항목", "확인할 내용", "회사 작성/확인값", "작성근거", "관련 법정 서식"],
        input_rows,
    )

    output = BytesIO()
    wb.save(output)
    return output.getvalue()
