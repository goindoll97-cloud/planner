"""실제 회사 자료처럼 복잡한 가상 시험자료.

‘정답’(TRUTH)을 먼저 정하고, 거기서 회사 파일 7개를 만들면서 실제 현장에서 흔한 문제를 일부러 심습니다.
그래서 프로그램이 자동으로 채운 값이 맞는지, 심어 둔 문제를 사람에게 물어보는지 정답표로 확인할 수 있습니다.

모든 회사명·제품명·설비·운전값은 허구입니다. 물질명과 CAS 번호만 실제이며, 실제 운전·평가에 쓸 수 없습니다.
"""
from __future__ import annotations

import csv
import json
from functools import lru_cache
from io import BytesIO, StringIO
from zipfile import ZIP_DEFLATED, ZipFile

COMPANY = "가상 한별정밀화학"
FILL = 0.85   # 저장탱크 최대보유량 = 용량 × 비중 × 1000 × 충전율 (가상 회사의 내부 기준)

# 코드, 제품명, CAS, 함량(%), 성상, 비중
PRODUCTS = [
    ("P-01", "가상 톨루엔", "108-88-3", 99.0, "액체", 0.87),
    ("P-02", "가상 메탄올", "67-56-1", 99.8, "액체", 0.79),
    ("P-03", "가상 아세톤", "67-64-1", 99.5, "액체", 0.79),
    ("P-04", "가상 자일렌", "1330-20-7", 98.0, "액체", 0.86),
    ("P-05", "가상 염산", "7647-01-0", 35.0, "액체", 1.17),
    ("P-06", "가상 황산", "7664-93-9", 98.0, "액체", 1.84),
    ("P-07", "가상 가성소다 용액", "1310-73-2", 48.0, "액체", 1.51),
    ("P-08", "가상 질산", "7697-37-2", 68.0, "액체", 1.41),
    ("P-09", "가상 암모니아수", "1336-21-6", 25.0, "액체", 0.91),
    ("P-10", "가상 과산화수소", "7722-84-1", 50.0, "액체", 1.20),
    ("P-11", "가상 포르말린", "50-00-0", 37.0, "액체", 1.09),
    ("P-12", "가상 이소프로판올", "67-63-0", 99.9, "액체", 0.79),
    ("P-13", "가상 메틸에틸케톤", "78-93-3", 99.5, "액체", 0.81),
    ("P-14", "가상 초산에틸", "141-78-6", 99.5, "액체", 0.90),
    ("P-15", "가상 벤젠", "71-43-2", 99.9, "액체", 0.88),
    ("P-16", "가상 페놀", "108-95-2", 99.0, "고체", 1.07),
    ("P-17", "가상 아크릴로니트릴", "107-13-1", 99.5, "액체", 0.81),
    ("P-18", "가상 불산", "7664-39-3", 55.0, "액체", 1.20),
    ("P-19", "가상 액화염소", "7782-50-5", 99.5, "기체·고압가스", 1.41),
    ("P-20", "가상 액화암모니아", "7664-41-7", 99.9, "기체·고압가스", 0.68),
    ("P-21", "가상 디메틸포름아미드", "68-12-2", 99.8, "액체", 0.95),
    ("P-22", "가상 노말헥산", "110-54-3", 95.0, "액체", 0.66),
    ("P-23", "가상 빙초산", "64-19-7", 99.8, "액체", 1.05),
    ("P-24", "가상 에탄올", "64-17-5", 95.0, "액체", 0.81),
]
DISCONTINUED = [("P-90", "가상 사염화탄소", "56-23-5", 99.5, "액체", 1.59),
                ("P-91", "가상 클로로포름", "67-66-3", 99.0, "액체", 1.48)]

# 저장탱크: 번호, 이름, 제품, 용량(m³), 운전압력(MPa), 운전온도(℃)
TANKS = [
    ("TK-101", "가상 톨루엔 탱크 A", "P-01", 50, 0.1, 25), ("TK-102", "가상 톨루엔 탱크 B", "P-01", 30, 0.1, 25),
    ("TK-103", "가상 메탄올 탱크 A", "P-02", 45, 0.1, 25), ("TK-104", "가상 자일렌 탱크", "P-04", 30, 0.1, 25),
    ("TK-105", "가상 아세톤 탱크", "P-03", 40, 0.1, 20), ("TK-106", "가상 IPA 탱크", "P-12", 30, 0.1, 25),
    ("TK-107", "가상 MEK 탱크", "P-13", 25, 0.1, 25), ("TK-108", "가상 초산에틸 탱크", "P-14", 20, 0.1, 25),
    ("TK-109", "가상 메탄올 탱크 B", "P-02", 50, 0.1, 25), ("TK-110", "가상 헥산 탱크", "P-22", 30, 0.1, 20),
    ("TK-111", "가상 염산 탱크", "P-05", 20, 0.1, 25), ("TK-112", "가상 황산 탱크", "P-06", 15, 0.1, 25),
    ("TK-113", "가상 가성소다 탱크", "P-07", 25, 0.1, 30), ("TK-114", "가상 질산 탱크", "P-08", 10, 0.1, 25),
    ("TK-115", "가상 과산화수소 탱크", "P-10", 10, 0.1, 20), ("TK-116", "가상 포르말린 탱크", "P-11", 15, 0.1, 40),
    ("TK-117", "가상 빙초산 탱크", "P-23", 20, 0.1, 25), ("TK-118", "가상 에탄올 탱크", "P-24", 30, 0.1, 25),
]
# 반응·혼합: 번호, 이름, 제품, 용량(L), 최대보유량(ton), 운전압력(kPa), 운전온도(℃), 설계압력(MPa), 공정유형
PROCESS = [
    ("RX-201", "가상 중합 반응기", "P-17", 8000, 5.2, 300, 70, 1.0, "반응"),
    ("RX-202", "가상 DMF 반응기", "P-21", 6000, 4.5, 250, 90, 1.0, "반응"),
    ("RX-203", "가상 수지 반응기", "P-11", 5000, 4.4, 300, 80, 1.0, "반응"),
    ("RX-204", "가상 페놀 반응기", "P-16", 4000, 3.6, 200, 85, 1.0, "반응"),
    ("RX-205", "가상 벤젠 반응기", "P-15", 6000, 4.2, 350, 120, 1.0, "반응"),
    ("RX-206", "가상 MEK 반응기", "P-13", 5000, 3.4, 200, 75, 1.0, "반응"),
    ("RX-207", "가상 메탄올 반응기", "P-02", 10000, 6.7, 150, 65, 1.0, "반응"),
    ("RX-208", "가상 불소화 반응기", "P-18", 3000, 2.9, 200, 40, 1.0, "반응"),
    ("MX-301", "가상 IPA 혼합조", "P-12", 4000, 2.6, 101, 25, 0.5, "단순혼합"),
    ("MX-302", "가상 톨루엔 혼합조", "P-01", 5000, 3.5, 101, 30, 0.5, "단순혼합"),
    ("MX-303", "가상 초산에틸 혼합조", "P-14", 3000, 2.2, 101, 30, 0.5, "단순혼합"),
    ("MX-304", "가상 아세톤 혼합조", "P-03", 3000, 2.0, 101, 30, 0.5, "단순혼합"),
    ("MX-305", "가상 암모니아수 혼합조", "P-09", 2000, 1.6, 101, 25, 0.5, "단순혼합"),
    ("MX-306", "가상 빙초산 혼합조", "P-23", 2000, 1.8, 101, 30, 0.5, "단순혼합"),
]
# 보관·가스: 번호, 이름, 제품, 용량(m³), 최대보유량(kg), 운전압력(MPa), 운전온도(℃), 성상
STORES = [
    ("DR-401", "가상 페놀 드럼 창고", "P-16", 2.4, 2000, 0.1, 25, "고체"),
    ("DR-402", "가상 가성소다 IBC 창고", "P-07", 3.0, 3000, 0.1, 25, "액체"),
    ("DR-403", "가상 염산 IBC 창고", "P-05", 2.0, 2400, 0.1, 25, "액체"),
    ("DR-404", "가상 과산화수소 IBC 창고", "P-10", 1.4, 1500, 0.1, 20, "액체"),
    ("CY-501", "가상 액화염소 용기 보관소", "P-19", 0.94, 1000, 0.68, 25, "기체·고압가스"),
    ("CY-502", "가상 액화암모니아 용기 보관소", "P-20", 0.94, 400, 0.86, 25, "기체·고압가스"),
]

README = f"""{COMPANY} — 실제와 비슷한 큰 연습 자료

모든 회사·제품·설비·운전값은 허구입니다(물질명과 CAS 번호만 실제). 실제 운전·평가에 쓸 수 없습니다.

[올리는 방법]
- 01~07을 한 번에 올리세요. 대부분 종류를 알아서 알아봅니다.
- 06(설비 사양서)은 표에 설비번호가 없습니다. ‘번호가 없는 문서’ 칸에 RX-205를 적으세요.
- 07(점검표)은 스캔본입니다. 글자를 읽을 수 없는 것이 정상이며, 종류는 아무거나 골라도 됩니다.
- answer_key.json은 올리지 마세요. 정답과 일부러 넣은 문제 목록입니다.

[파일]
01 제품목록(엑셀): 제목 줄, 숨김 시트(단종 제품), 쓰지 않는 열(제조사·UN번호), 함량 범위 표기, CAS 오타, 성상 표기 흔들림
02 MSDS 요약(워드): 개정된 MSDS라 01과 함량이 다른 제품이 있음
03 설비대장(엑셀): 시트 5개(숨김 1), 시트마다 다른 단위(L·ton·kPa), 병합 셀, 결과 없는 계산식, 합계 줄, 오타
04 2024년 설비목록(CSV, 한글 윈도우 인코딩): 옛 값이 섞여 있음
05 운전조건표(CSV): 03과 같은 값으로 교차 확인되는 값, 다른 값, 03에 빠진 값, 설비번호 없이 이름만 있는 줄
06 설비 사양서(한글 HWPX): 문장으로 적힌 값
07 정기점검표(PDF): 스캔본
"""


# ---------------------------------------------------------------- 정답

def _tank_inventory(capacity, product):
    density = next(p[5] for p in PRODUCTS if p[0] == product)
    return int(round(capacity * density * 1000 * FILL, -1))


def truth():
    """프로그램 단위(m³·MPa·kg·%·℃)로 적은 최종 정답. 제품은 제품코드로 적습니다."""
    out = {}
    for code, name, cas, conc, state, density in PRODUCTS + DISCONTINUED:
        out["material:" + code] = {"name": name, "cas": cas, "concentration": conc, "substance_state": state, "density": density}
    for tid, name, product, cap, pressure, temp in TANKS:
        out["facility:" + tid] = {"name": name, "product": product, "capacity": float(cap), "inventory": float(_tank_inventory(cap, product)),
                                  "pressure": pressure, "temperature": float(temp), "facility_type": "저장탱크", "substance_state": "액체"}
    for fid, name, product, liters, tons, kpa, temp, design, process in PROCESS:
        state = "액체"
        out["facility:" + fid] = {"name": name, "product": product, "capacity": liters / 1000, "inventory": tons * 1000,
                                  "pressure": kpa / 1000, "temperature": float(temp), "design_pressure": design,
                                  "facility_type": "제조·사용시설", "substance_state": state, "process_type": process}
    for sid, name, product, cap, inventory, pressure, temp, state in STORES:
        out["facility:" + sid] = {"name": name, "product": product, "capacity": cap, "inventory": float(inventory),
                                  "pressure": pressure, "temperature": float(temp), "facility_type": "보관시설", "substance_state": state}
    out["facility:TK-190"] = {"name": "가상 폐기예정 탱크", "product": "P-90", "capacity": 10.0, "inventory": 0.0,
                              "facility_type": "취급중단 신고시설", "substance_state": "액체"}
    return out


# 일부러 넣은 문제: 파일, 대상, 항목, 종류, 자료에 적힌 값, 기대하는 프로그램 동작
PLANTED = [
    ("01", "material:P-05", "concentration", "range", "35~37", "범위라서 물어봄"),
    ("01", "material:P-13", "cas", "cas_typo", "78-93-8", "CAS 검산 실패로 물어봄"),
    ("01", "material:P-16", "substance_state", "unknown_word", "결정", "선택지에 없는 말이라 물어봄"),
    ("01", "material:P-19", "density", "text_number", "1.41(액화)", "숫자가 아니라 물어봄"),
    ("01", "material:P-19", "substance_state", "synonym", "가스", "기체·고압가스로 알아서 바꿈"),
    ("01", "material:P-01", "substance_state", "synonym", "액상", "액체로 알아서 바꿈"),
    ("02", "material:P-10", "concentration", "conflict", "50 (01에는 35)", "자료마다 달라 물어봄(개정 MSDS가 정답)"),
    ("03", "facility:TK-101~TK-112", "facility_type", "merged", "병합 셀 ‘저장탱크’", "법적 항목이라 자동으로 채우지 않고 확인을 받음"),
    ("03", "facility:TK-104", "product", "link", "가상자일렌 (띄어쓰기 없음)", "물질 목록과 연결되지 않아 알려 줌"),
    ("03", "facility:TK-109", "inventory", "formula", "=50*0.79*1000*0.85", "계산값 33,575 kg을 제안"),
    ("03", "facility:TK-112", "inventory", "formula_ref", "=E15*1.84*1000*0.85", "셀 참조라 계산하지 않고 물어봄"),
    ("03", "facility:TK-115", "pressure", "unit_typo", "101.3 (MPa 칸에 kPa로 적음)", "허용 범위 밖이라 물어봄"),
    ("03", "facility:TK-117", "facility_type", "unknown_word", "저장조", "선택지에 없는 말이라 물어봄"),
    ("03", "facility:TK-118", "pressure", "missing", "", "자료에 없어서 물어봄"),
    ("03", "facility:TK-118", "temperature", "missing", "", "자료에 없어서 물어봄"),
    ("03", "(합계 줄)", "", "total_row", "합계", "설비로 읽지 않고 건너뜀"),
    ("03", "facility:RX-203", "pressure", "range", "200~400 kPa (05에는 0.3 MPa)", "범위 값은 버리고 05의 값과 비교해 물어봄"),
    ("03", "facility:RX-204", "substance_state", "unknown_word", "용융", "선택지에 없는 말이라 물어봄"),
    ("03", "facility:RX-206", "inventory", "missing", "", "자료에 없어서 물어봄"),
    ("03", "facility:DR-401", "pressure", "text", "상압", "숫자가 아니라 물어봄"),
    ("03", "facility:DR-402", "pressure", "text", "상압", "숫자가 아니라 물어봄"),
    ("03", "facility:MX-301~MX-306", "facility_type", "synonym", "사용시설", "제조·사용시설로 알아서 바꿈"),
    ("03", "facility:TK-190", "", "hidden_sheet", "숨김 시트", "숨김 시트도 읽고 위치에 표시"),
    ("04", "facility:TK-103", "capacity", "conflict", "40 (03에는 45)", "자료마다 달라 물어봄(03이 정답)"),
    ("04", "facility:TK-105", "name", "conflict", "가상 2호 저장탱크 (옛 이름)", "자료마다 달라 물어봄(03이 정답)"),
    ("04", "facility:TK-107", "inventory", "conflict", "15000 (03과 다름)", "자료마다 달라 물어봄(03이 정답)"),
    ("05", "facility:RX-204", "temperature", "conflict", "80 (03에는 85)", "자료마다 달라 물어봄(03이 정답)"),
    ("05", "facility:MX-304", "temperature", "fills_gap", "30 (03에는 빈칸)", "다른 파일 값으로 빈칸을 채움"),
    ("05", "facility:가상 혼합조 7·8", "", "name_only", "설비번호 없음", "이름만으로 찾은 값이라 확인을 받음"),
    ("06", "facility:RX-205", "", "sentence", "‘운전압력(MPa): 0.35’ 같은 문장", "문장에서 찾은 값이라 확인을 받음"),
    ("07", "", "", "scan", "스캔본", "글자를 읽지 못했다고 알려 줌"),
]


# ---------------------------------------------------------------- 파일 만들기

def _product_name(code):
    return next(p[1] for p in PRODUCTS + DISCONTINUED if p[0] == code)


def _materials_xlsx():
    from openpyxl import Workbook
    book = Workbook()
    sheet = book.active
    sheet.title = "제품목록"
    sheet.append([f"{COMPANY} 취급 제품 목록 (2026년 3월 기준)"])
    sheet.merge_cells("A1:I1")
    sheet.append(["※ 함량은 대표 성분 기준. 혼합물 세부 조성은 MSDS 참조"])
    sheet.append([])
    headers = ["제품코드", "제품명", "CAS 번호", "함량(wt%)", "성상", "비중", "제조사", "UN번호", "비고"]
    sheet.append(headers)
    odd = {("P-05", "concentration"): "35~37", ("P-13", "cas"): "78-93-8", ("P-16", "state"): "결정",
           ("P-19", "density"): "1.41(액화)", ("P-19", "state"): "가스", ("P-10", "concentration"): 35}
    for index, (code, name, cas, conc, state, density) in enumerate(PRODUCTS):
        shown_state = odd.get((code, "state"), "액상" if state == "액체" and index % 2 == 0 else state)
        sheet.append([code, name, odd.get((code, "cas"), cas), odd.get((code, "concentration"), conc), shown_state,
                      odd.get((code, "density"), density), "가상 공급사 " + "ABC"[index % 3], f"UN{1000 + index * 37}", ""])
    hidden = book.create_sheet("단종제품")
    hidden.sheet_state = "hidden"
    hidden.append(headers)
    for code, name, cas, conc, state, density in DISCONTINUED:
        hidden.append([code, name, cas, conc, state, density, "가상 공급사 D", "", "2025년 단종"])
    return _book_bytes(book)


def _msds_docx():
    from docx import Document
    doc = Document()
    doc.add_heading(f"{COMPANY} MSDS 요약 (2026년 2월 개정분)", level=1)
    doc.add_paragraph("공급사에서 새로 받은 MSDS의 3번(구성성분)·9번(물리화학적 특성) 항목을 요약했습니다.")
    doc.add_paragraph("과산화수소는 공급 규격이 35%에서 50%로 바뀌었습니다.")
    rows = [("P-07", "가상 가성소다 용액", "1310-73-2", "48", "1.51", "2026-01-15"),
            ("P-10", "가상 과산화수소", "7722-84-1", "50", "1.20", "2026-02-03"),
            ("P-18", "가상 불산", "7664-39-3", "55", "1.20", "2025-11-20")]
    table = doc.add_table(rows=1, cols=6)
    for cell, text in zip(table.rows[0].cells, ["제품코드", "제품명", "CAS 번호", "함량(%)", "비중", "개정일"]):
        cell.text = text
    for row in rows:
        for cell, text in zip(table.add_row().cells, row):
            cell.text = text
    out = BytesIO()
    doc.save(out)
    return out.getvalue()


def _ledger_xlsx():
    from openpyxl import Workbook
    book = Workbook()
    tanks = book.active
    tanks.title = "저장시설"
    tanks.append([f"{COMPANY} 설비대장 — 저장시설"])
    tanks.merge_cells("A1:J1")
    tanks.append(["작성: 가상 안전환경팀 / 기준일 2026-03-31"])
    tanks.append(["설비번호", "설비명", "시설구분", "취급제품", "용량(m3)", "최대보유량(kg)", "운전압력(MPa)", "운전온도(℃)", "물질상태", "비고"])
    first = tanks.max_row + 1
    for tid, name, product, cap, pressure, temp in TANKS:
        inventory = _tank_inventory(cap, product)
        product_text = {"TK-104": "가상자일렌", "TK-106": "P-12", "TK-116": "P-11"}.get(tid, _product_name(product))
        row = [tid, name, "저장탱크", product_text, cap, inventory, pressure, temp, "액상", ""]
        if tid == "TK-109":
            row[5] = "=50*0.79*1000*0.85"
        if tid == "TK-112":
            row[5] = f"=E{tanks.max_row + 1}*1.84*1000*0.85"
        if tid == "TK-115":
            row[6] = 101.3
        if tid == "TK-117":
            row[2] = "저장조"
        if tid == "TK-118":
            row[6], row[7], row[9] = None, None, "운전조건 확인 중"
        tanks.append(row)
    # 실제 대장처럼 같은 구분은 한 칸으로 합쳐 둡니다 (TK-101~TK-112)
    tanks.merge_cells(start_row=first, start_column=3, end_row=first + 11, end_column=3)
    last = tanks.max_row
    tanks.append(["", "합계", "", "", f"=SUM(E{first}:E{last})", f"=SUM(F{first}:F{last})", "", "", "", ""])

    process = book.create_sheet("제조공정")
    process.append([f"{COMPANY} 설비대장 — 제조공정 (단위 주의: L, ton, kPa)"])
    process.append([])
    process.append(["설비번호", "설비명", "시설구분", "공정구분", "취급제품", "용량(L)", "최대보유량(ton)", "운전압력(kPa)",
                    "운전온도(℃)", "설계압력(MPa)", "물질상태"])
    for fid, name, product, liters, tons, kpa, temp, design, kind in PROCESS:
        row = [fid, name, "사용시설" if fid.startswith("MX") else "제조·사용시설", kind, _product_name(product),
               liters, tons, kpa, temp, design, "액체"]
        if fid == "RX-203":
            row[7] = "200~400"
        if fid == "RX-204":
            row[10] = "용융"
        if fid == "RX-206":
            row[6] = None
        if fid == "MX-304":
            row[8] = None
        process.append(row)

    stores = book.create_sheet("보관창고")
    stores.append(["설비번호", "설비명", "시설구분", "취급제품", "용량(m3)", "최대보유량(kg)", "운전압력(MPa)", "운전온도(℃)", "물질상태"])
    for sid, name, product, cap, inventory, pressure, temp, state in STORES[:4]:
        shown = "상압" if sid in {"DR-401", "DR-402"} else pressure
        stores.append([sid, name, "보관시설", _product_name(product), cap, inventory, shown, temp, state])

    gas = book.create_sheet("가스시설")
    gas.append(["설비번호", "설비명", "시설구분", "취급제품", "용량(m3)", "최대보유량(kg)", "운전압력(MPa)", "운전온도(℃)", "물질상태"])
    for sid, name, product, cap, inventory, pressure, temp, state in STORES[4:]:
        gas.append([sid, name, "보관시설", _product_name(product), cap, inventory, pressure, temp, "기체"])

    old = book.create_sheet("폐기예정")
    old.sheet_state = "hidden"
    old.append(["설비번호", "설비명", "시설구분", "취급제품", "용량(m3)", "최대보유량(kg)", "물질상태"])
    old.append(["TK-190", "가상 폐기예정 탱크", "취급중단 신고시설", "가상 사염화탄소", 10, 0, "액체"])
    return _book_bytes(book)


def _old_list_csv():
    rows = [["설비번호", "설비명", "용량(m3)", "최대보유량(kg)", "운전압력(MPa)", "비고"]]
    for tid, name, product, cap, pressure, _ in TANKS[:10]:
        inventory = _tank_inventory(cap, product)
        if tid == "TK-103":
            cap = 40
        if tid == "TK-105":
            name = "가상 2호 저장탱크"
        if tid == "TK-107":
            inventory = 15000
        if tid == "TK-109":
            inventory = ""      # 옛 목록에도 비어 있음
        rows.append([tid, name, cap, inventory, pressure, "2024년 정기검사 기준"])
    return _csv(rows).encode("cp949")


def _operating_csv():
    rows = [["설비번호", "설비명", "운전압력(MPa)", "운전온도(℃)", "설계압력(MPa)"]]
    for fid, name, product, liters, tons, kpa, temp, design, kind in PROCESS:
        shown_temp = 80 if fid == "RX-204" else temp
        rows.append([fid, name, kpa / 1000, shown_temp, design])
    rows.append(["", "가상 혼합조 7", 0.1, 25, 0.5])
    rows.append(["", "가상 혼합조 8", 0.1, 30, 0.5])
    return _csv(rows).encode("utf-8-sig")


def _spec_hwpx():
    from hwpx import HwpxDocument
    doc = HwpxDocument.new()
    for line in ["가상 벤젠 반응기 설비 사양서", "설비명: 가상 벤젠 반응기", "용량(m3): 6", "운전압력(MPa): 0.35",
                 "운전온도(℃): 120", "※ 설계 조건은 아래 표 참조"]:
        doc.add_paragraph(line)
    table = doc.add_table(2, 3)
    for c, (head, value) in enumerate([("설계압력(MPa)", "1.0"), ("운전온도(℃)", "120"), ("공정유형", "반응")]):
        table.set_cell_text(0, c, head)
        table.set_cell_text(1, c, value)
    return doc.to_bytes()


def _scanned_pdf():
    """글자 없이 그림만 있는 PDF(스캔본 흉내)."""
    content = b"q 0.85 g 50 650 500 120 re f Q q 0 g 60 600 m 540 600 l S Q"
    objects = [b"<< /Type /Catalog /Pages 2 0 R >>",
               b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
               b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 595 842] /Contents 4 0 R /Resources << >> >>",
               b"<< /Length %d >>\nstream\n" % len(content) + content + b"\nendstream"]
    out = BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets = []
    for number, body in enumerate(objects, 1):
        offsets.append(out.tell())
        out.write(b"%d 0 obj\n" % number + body + b"\nendobj\n")
    xref = out.tell()
    out.write(b"xref\n0 %d\n0000000000 65535 f \n" % (len(objects) + 1))
    for offset in offsets:
        out.write(b"%010d 00000 n \n" % offset)
    out.write(b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (len(objects) + 1, xref))
    return out.getvalue()


def _book_bytes(book):
    out = BytesIO()
    book.save(out)
    return out.getvalue()


def _csv(rows):
    out = StringIO()
    csv.writer(out, lineterminator="\n").writerows(rows)
    return out.getvalue()


def files():
    """파일 이름 → 내용. 올릴 때 필요한 종류와 번호는 UPLOAD에 있습니다."""
    return {"01_제품목록_2026.xlsx": _materials_xlsx(),
            "02_MSDS_요약_2026개정.docx": _msds_docx(),
            "03_설비대장_2026.xlsx": _ledger_xlsx(),
            "04_설비목록_2024_구버전.csv": _old_list_csv(),
            "05_운전조건표_2026Q1.csv": _operating_csv(),
            "06_설비사양서_RX-205.hwpx": _spec_hwpx(),
            "07_정기점검표_스캔.pdf": _scanned_pdf()}


# 사람이 올릴 때 고르는 값: 종류, 번호 없는 문서의 대상 번호
UPLOAD = {"01_제품목록_2026.xlsx": ("material", ""), "02_MSDS_요약_2026개정.docx": ("material", ""),
          "03_설비대장_2026.xlsx": ("facility", ""), "04_설비목록_2024_구버전.csv": ("facility", ""),
          "05_운전조건표_2026Q1.csv": ("facility", ""), "06_설비사양서_RX-205.hwpx": ("facility", "RX-205"),
          "07_정기점검표_스캔.pdf": ("facility", "")}


def answer_key():
    return {"안내": "정답(truth)은 프로그램 단위(m³·MPa·kg·%·℃)로 적었고 제품은 제품코드입니다. planted는 일부러 넣은 문제입니다.",
            "회사": COMPANY,
            "truth": truth(),
            "planted": [{"파일": f, "대상": e, "항목": fld, "종류": k, "자료에 적힌 값": raw, "기대 동작": expect}
                        for f, e, fld, k, raw, expect in PLANTED]}


@lru_cache(maxsize=1)
def large_demo_zip():   # 화면이 다시 그려질 때마다 파일 7개를 새로 만들지 않도록 보관
    out = BytesIO()
    with ZipFile(out, "w", ZIP_DEFLATED) as archive:
        for name, content in files().items():
            archive.writestr(name, content)
        archive.writestr("README.txt", README.encode("utf-8"))
        archive.writestr("answer_key.json", json.dumps(answer_key(), ensure_ascii=False, indent=2).encode("utf-8"))
    return out.getvalue()
