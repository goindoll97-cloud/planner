from __future__ import annotations

from io import BytesIO
import json

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from .cap_requests import build_cap_data_requests
from .completeness import evaluate_project_completeness
from .intake import build_intake_catalog
from .project import Stage2Project
from .psm_requests import build_psm_data_requests
from .requirements import (
    cap_field_labels,
    cap_manual_source,
    psm_example_source,
    psm_field_labels,
)


HEADER_FILL = PatternFill("solid", fgColor="D9EAF7")
SYSTEM_LABELS = {
    "COMMON": "공통자료",
    "PSM": "공정안전보고서",
    "CAP": "화학사고예방관리계획서",
}
BASE_FIELD_LABELS = {
    "business.company_name": "회사명",
    "business.address": "사업장 소재지",
    "inventory.chemicals": "화학물질 목록",
    "inventory.facilities": "시설·설비 목록",
    "process.description": "공정설명서",
    "documents.pfd": "공정흐름도(PFD)",
    "documents.pid": "공정배관·계장도(P&ID)",
    "documents.site_plan": "각종 건물·설비의 배치도",
    "emergency.internal_plan": "내부 비상대응계획",
    "emergency.external_plan": "외부 비상대응계획",
}


def _header(ws, row: int, values: list[str]) -> None:
    for col, value in enumerate(values, start=1):
        cell = ws.cell(row=row, column=col, value=value)
        cell.font = Font(bold=True)
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(vertical="top", wrap_text=True)


def _field_labels() -> dict[str, str]:
    labels = dict(BASE_FIELD_LABELS)
    try:
        labels.update(psm_field_labels())
    except Exception:
        pass
    try:
        labels.update(cap_field_labels())
    except Exception:
        pass
    return labels


def _field_label(key: str) -> str:
    return _field_labels().get(key, key)


def _field_list(values) -> str:
    return ", ".join(_field_label(value) for value in values)


def build_progress_workbook(project: Stage2Project) -> bytes:
    """Export an audit/progress workbook, not a statutory final submission form."""
    wb = Workbook()
    ws = wb.active
    ws.title = "프로젝트"
    rows = [
        ("프로젝트 ID", project.project_id),
        ("회사명", project.company_name),
        ("사업장명", project.site_name),
        ("공정안전보고서 법적 대상 여부", project.psm_required),
        ("화학사고예방관리계획서 법적 대상 여부", project.cap_required),
        ("화학사고예방관리계획서 작성수준", project.cap_group),
        ("작성범위 선택 완료", project.scope_confirmed),
        ("공정안전보고서 작성 지원 선택", project.psm_in_scope),
        ("화학사고예방관리계획서 작성 지원 선택", project.cap_in_scope),
        ("Stage 1 원본 SHA-256", project.stage1_source_fingerprint),
        ("생성일", project.created_at),
        ("최종수정일", project.updated_at),
        ("스키마", project.schema_version),
    ]
    if project.psm_in_scope:
        sources = psm_example_source()
        source = sources.get("source", {})
        outline = sources.get("outline_source", {})
        legal_ref = sources.get("legal_reference", {})
        rows.extend([
            ("공정안전보고서 작성예시집", source.get("title", "")),
            ("공정안전보고서 작성예시집 PDF SHA-256", source.get("sha256", "")),
            ("공정안전보고서 작성예시집 역할", source.get("source_role", "")),
            ("공정안전보고서 목차 SHA-256", outline.get("sha256", "")),
            ("공정안전보고서 현행구조 확인일", legal_ref.get("checked_as_of", "")),
            ("공정안전보고서 현행 법적 근거", f"{legal_ref.get('statute', '')} / {legal_ref.get('rule', '')}"),
            ("공정안전보고서 현행 세부고시", legal_ref.get("administrative_rule", "")),
        ])
    if project.cap_in_scope:
        source = cap_manual_source()
        rows.extend([
            ("화학사고예방관리계획서 작성 매뉴얼", source.get("title", "")),
            ("화학사고예방관리계획서 작성 매뉴얼 문서번호", source.get("document_code", "")),
            ("화학사고예방관리계획서 작성 매뉴얼 PDF SHA-256", source.get("sha256", "")),
            ("화학사고예방관리계획서 작성 매뉴얼 역할", source.get("source_role", "")),
        ])
    _header(ws, 1, ["항목", "값"])
    for r, (label, value) in enumerate(rows, start=2):
        ws.cell(r, 1, label)
        ws.cell(r, 2, "" if value is None else str(value))
    ws.column_dimensions["A"].width = 46
    ws.column_dimensions["B"].width = 90

    completeness = evaluate_project_completeness(project)
    ws2 = wb.create_sheet("작성현황")
    _header(ws2, 1, ["구분", "작성구조", "작성항목", "상태", "완성도(%)", "미확인 항목", "AI 초안 항목", "검증 보류 항목", "작성근거"])
    for r, item in enumerate(completeness["requirements"], start=2):
        values = [
            SYSTEM_LABELS.get(item["system"], item["system"]),
            item["section"],
            item["label"],
            item["state"],
            item["completion_pct"],
            _field_list(item["missing_fields"]),
            _field_list(item["draft_fields"]),
            _field_list(item["hold_fields"]),
            item["legal_basis"],
        ]
        for c, value in enumerate(values, start=1):
            ws2.cell(r, c, value)
    for col, width in {"A": 28, "B": 26, "C": 38, "D": 18, "E": 14, "F": 50, "G": 42, "H": 42, "I": 72}.items():
        ws2.column_dimensions[col].width = width

    ws_intake = wb.create_sheet("자료준비·접수현황")
    _header(ws_intake, 1, [
        "구분", "작성구조", "작성항목", "자료상태", "현재 확인되지 않은 항목",
        "접수 후 확인이 필요한 항목", "확인 가능한 자료 예", "관련 법정 서식", "작성근거", "작성 참고자료", "관련 쪽"
    ])
    for r, item in enumerate(build_intake_catalog(project), start=2):
        values = [
            item.system_label,
            item.section,
            item.label,
            item.coverage_status,
            ", ".join(item.missing_labels),
            ", ".join(item.received_unconfirmed_labels),
            ", ".join(item.suggested_evidence),
            ", ".join(item.form_references),
            item.legal_basis,
            item.reference_label,
            ", ".join(str(v) for v in item.reference_pages),
        ]
        for c, value in enumerate(values, start=1):
            ws_intake.cell(r, c, value)
    for col, width in {
        "A": 28, "B": 26, "C": 38, "D": 18, "E": 55, "F": 55,
        "G": 70, "H": 34, "I": 78, "J": 50, "K": 18
    }.items():
        ws_intake.column_dimensions[col].width = width

    if project.psm_in_scope:
        ws_psm = wb.create_sheet("공정안전보고서 요청자료")
        _header(ws_psm, 1, [
            "우선순위", "구분", "작성항목", "예시집 PDF 페이지", "미확인 항목",
            "권장 증빙자료", "상호검증 항목", "법적 상태", "법적·작성 근거", "요청문구", "자동화 방식"
        ])
        for r, item in enumerate(build_psm_data_requests(project), start=2):
            values = [
                item.priority,
                item.section,
                item.label,
                ", ".join(str(v) for v in item.example_pages),
                ", ".join(item.missing_labels),
                ", ".join(item.suggested_evidence),
                ", ".join(_field_label(v) for v in item.cross_checks),
                item.legal_status,
                item.legal_basis,
                item.request_text,
                item.automation,
            ]
            for c, value in enumerate(values, start=1):
                ws_psm.cell(r, c, value)
        for col, width in {
            "A": 12, "B": 22, "C": 34, "D": 18, "E": 52, "F": 65,
            "G": 60, "H": 30, "I": 75, "J": 75, "K": 34
        }.items():
            ws_psm.column_dimensions[col].width = width

    if project.cap_in_scope:
        ws_req = wb.create_sheet("화학사고예방관리계획서 요청자료")
        _header(ws_req, 1, ["우선순위", "작성구조", "작성항목", "매뉴얼 페이지", "미확인 항목", "권장 증빙자료", "요청문구", "자동화 방식"])
        for r, item in enumerate(build_cap_data_requests(project), start=2):
            values = [
                item.priority,
                item.section,
                item.label,
                ", ".join(str(v) for v in item.manual_pages),
                ", ".join(item.missing_labels),
                ", ".join(item.suggested_evidence),
                item.request_text,
                item.automation,
            ]
            for c, value in enumerate(values, start=1):
                ws_req.cell(r, c, value)
        for col, width in {"A": 12, "B": 26, "C": 34, "D": 18, "E": 52, "F": 65, "G": 75, "H": 28}.items():
            ws_req.column_dimensions[col].width = width

    ws3 = wb.create_sheet("필드")
    _header(ws3, 1, ["내부 필드 식별자", "표시명", "값", "상태", "비고", "최종수정일", "증빙수"])
    for r, key in enumerate(sorted(project.fields), start=2):
        record = project.fields[key]
        if isinstance(record.value, (list, dict)):
            value = json.dumps(record.value, ensure_ascii=False, default=str)
        else:
            value = "" if record.value is None else str(record.value)
        values = [record.key, record.label, value, record.status, record.note, record.updated_at, len(record.evidence)]
        for c, cell_value in enumerate(values, start=1):
            ws3.cell(r, c, cell_value)
    for col, width in {"A": 38, "B": 38, "C": 80, "D": 18, "E": 55, "F": 28, "G": 10}.items():
        ws3.column_dimensions[col].width = width
    ws3.column_dimensions["A"].hidden = True

    ws4 = wb.create_sheet("증빙목록")
    _header(ws4, 1, ["내부 필드 식별자", "표시명", "근거종류", "근거파일", "SHA-256", "페이지", "저장위치", "비고"])
    r = 2
    for key in sorted(project.fields):
        record = project.fields[key]
        for evidence in record.evidence:
            values = [key, record.label, evidence.source_type, evidence.source_name, evidence.sha256, evidence.page, evidence.location, evidence.note]
            for c, cell_value in enumerate(values, start=1):
                ws4.cell(r, c, cell_value)
            r += 1
    for col, width in {"A": 38, "B": 38, "C": 22, "D": 42, "E": 68, "F": 12, "G": 70, "H": 50}.items():
        ws4.column_dimensions[col].width = width
    ws4.column_dimensions["A"].hidden = True

    for sheet in wb.worksheets:
        sheet.freeze_panes = "A2"
        for row in sheet.iter_rows():
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)

    out = BytesIO()
    wb.save(out)
    return out.getvalue()
