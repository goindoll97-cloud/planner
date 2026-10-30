"""실제 회사 양식의 값 칸에 가상 회사를 채운 업로드·출력 연습 파일."""
from functools import lru_cache
from io import BytesIO

from openpyxl import load_workbook

from . import cap_upload_template as template, demo_authoring_company as demo
from . import evidence_workspace as ew
from .direct_template import create_direct_cap_project

COMPANY = demo.COMPANY
FILENAME = '화사계_가상회사_작성완료_연습용.xlsx'
IDENTIFIER = 'planner-synthetic-company-workbook-v1'
NOTICE = '가상 회사·인물·수치로 작성한 1군 작성 연습용입니다. 실제 법적 대상 판정·현장 확인을 뜻하지 않으며 실제 제출에 사용하지 마세요.'

SITE_VALUES = {
    'site.address': '00000 가상시 가상구 연습산업로 100 (실재하지 않는 연습용 주소)',
    'site.company_operating_roles': '제1공장 운전책임자 박가상, 주간조 작업자 4명, 야간조 2명',
    'site.company_hazard_method': '가상 HAZOP 연습, 2026-10-01, 저장·혼합·반응 공정',
    'site.company_hazard_findings': '가상 TK-01 누출 시 화재 위험: 방류벽·감지기·긴급차단으로 대응하는 연습',
    'narrative.safety_management.policy': '누출 예방과 근로자·주민의 인명 보호 우선 (가상 방침)',
    'narrative.safety_management.technical': '가상 설비명세와 점검절차의 최신본을 안전환경팀에서 관리',
    'narrative.safety_management.execution': '안전환경팀 담당, 2026년 분기별 점검, 연습용 예산 500만원',
    'site.company_response_equipment': '가상 흡착포 200장·화학보호복 4벌·이동식 펌프 1대, 방재창고 보관',
    'site.company_response_use': '가상 대응반장이 방재창고 장비를 동원하고 부족하면 협력사업장에 지원 요청',
    'site.company_medical': '가상 새솔병원, 가상시 연습로 20, 000-000-0020, 119 이송 요청(실제 협의 아님)',
    'site.company_shelter': '가상 집결지 A, 사업장 서측 600m, 가상 수용인원 100명, 000-000-0030',
    'site.company_medical_check': '2026-10-01, 이담당, 의료기관 연락 절차 연습(실제 협의 아님)',
    'site.company_communication_owner': '가상 안전환경팀 이담당, 가상 주민회관',
    'site.company_notice_schedule': '가상 연습일 2026-10-15, 안전환경팀 이담당(법정 고지 주기로 정한 값이 아님)',
    'narrative.emergency_org.총괄': '가상 비상대책본부장 박가상(제1공장장)',
    'narrative.emergency_org.대응반': '가상 소화반·누출차단반·응급처치반 각 2명',
    'narrative.emergency_org.교대': '주간조·야간조 교대, 야간 비상연락망 가동(가상)',
    'narrative.command_center.위치': '가상 사무동 1층 방재센터',
    'narrative.command_center.운영': '24시간 상황실 운영(가상)',
    'narrative.command_center.담당': '가상 안전환경팀 이담당',
    'narrative.shutdown.권한자': '가상 제1공장 운전책임자 박가상',
    'narrative.shutdown.판단기준': '가상 감지기 경보 2단계 이상 또는 누출 확인 시',
    'narrative.shutdown.방법': '가상 긴급차단밸브 작동 후 공급 중단',
    'narrative.communication.수단': '가상 비상방송·무전기·사내 전화',
    'narrative.communication.방법': '상황 발생 즉시 방송 후 대응반에 무전 전파(가상)',
    'narrative.communication.담당': '가상 안전환경팀 이담당',
    'narrative.recovery.책임': '가상 시설팀 박가상',
    'narrative.recovery.순서': '누출 차단 → 잔류물 제거 → 설비 점검 → 재가동 승인(가상)',
    'narrative.recovery.환경조치': '가상 흡착포 회수·폐기물 위탁 처리',
    'narrative.investigation.책임': '가상 안전환경팀장 이담당',
    'narrative.investigation.방법': '사고 후 7일 이내 원인 조사 보고서 작성(가상)',
    'narrative.investigation.재발방지': '개선 과제를 등록하고 분기별 이행 점검(가상)',
    'narrative.community.대상': '가상 주민회관·인근 마을 대표',
    'narrative.community.방법': '연 1회 설명회와 안내문 배포(가상)',
    'narrative.community.주기': '연 1회(가상)',
    'narrative.mutual_aid.기관': '가상 소방서·가상 인근사업장 협의체',
    'narrative.mutual_aid.협력내용': '장비·인력 상호 지원 요청(가상, 실제 협의 아님)',
    'narrative.mutual_aid.합동훈련': '연 1회 합동 훈련 계획(가상)',
    'narrative.evacuation.경보': '가상 사이렌과 비상방송',
    'narrative.evacuation.경로': '사업장 서측 도로를 따라 가상 집결지 A로 이동',
    'narrative.evacuation.장소': '가상 집결지 A',
    'narrative.public_notice.대상': '가상 사업장 인근 주민',
    'narrative.public_notice.방법': '안내문 게시와 문자 안내(가상)',
    'narrative.public_notice.내용': '취급 물질과 사고 시 행동 요령(가상)',
    'site.survey_source': '가상 조사기록 DEMO-01, 2026-10-01(실제 지도·현장 확인 아님)',
    'site.survey_method': '작성 연습을 위해 가정한 거리이며 실제 경계·거리는 미검증',
    'site.survey_map_numbers': 'DEMO-01(실제 지도 아님)',
    'site.survey_scope': '아니오', 'site.survey_no_target': '아니오',
    'site.survey_check_boundary': '아니오', 'site.survey_check_facilities': '아니오',
    'site.survey_check_environment': '아니오', 'site.survey_check_classification': '아니오',
    'site.company_risk_measures': '가상 긴급차단 기능 점검, 시설팀 박가상, 2026-11-01 완료 예정(연습)',
}

# 사고시나리오 검토 시트의 연습값(가상). 경계 거리·사고원점·주변 보호대상(지도)과 사고유형·선정 근거(프로그램 제안)는 넣지 않는다.
_ST = '가상 방류벽(수동)'
SCENARIO_VALUES = {
    'TK-01': {'액위(m)': 6, '누출지점 배관길이(m)': 15, '비열비': 1.06, '연소열(kJ/kg)': 40500, '비점(℃)': 110.6,
              '액체비열(kJ/kg·K)': 1.7, '기화열(kJ/kg)': 365, '증발잠열(kcal/kg)': 87, '액체비열(kcal/kg·℃)': 0.4,
              '증기밀도(kg/m3)': 3.1, '상압 탱크 파열 및 누출': 1, '배관누출': 1, '플랜지 등의 가스켓 파손': 1,
              '입/출하 시설 누출 사고': 1, '외부화재': 1, '개수 산정근거': '가상 P&ID 가상-PID-TK-01의 연결 요소 수 기준(연습)',
              '수동적 완화장치': _ST, '능동적 완화장치': '가상 가스감지기 GD-01 연동 긴급차단', '안전성확보설비 증빙': '가상 설비 점검기록 DEMO-SAFE-01'},
    'TK-02': {'액위(m)': 4, '누출지점 배관길이(m)': 10, '비열비': 1.4, '비점(℃)': 48, '액체비열(kJ/kg·K)': 3.2,
              '기화열(kJ/kg)': 443, '증기밀도(kg/m3)': 1.6, '상압 탱크 파열 및 누출': 1, '배관누출': 1,
              '플랜지 등의 가스켓 파손': 1, '입/출하 시설 누출 사고': 1, '개수 산정근거': '가상 P&ID 연결 요소 수 기준(연습)',
              '수동적 완화장치': _ST, '능동적 완화장치': '가상 누출 경보 및 중화 설비', '안전성확보설비 증빙': '가상 설비 점검기록 DEMO-SAFE-02'},
    'RX-01': {'누출지점 배관길이(m)': 8, '비열비': 1.1, '연소열(kJ/kg)': 38000, '비점(℃)': 100, '액체비열(kJ/kg·K)': 1.7, '기화열(kJ/kg)': 365, '증기밀도(kg/m3)': 2.8,
              '고압용기파열': 1, '배관파열': 1, '배관누출': 1, '플랜지 등의 가스켓 파손': 2, '펌프/컴프레서 누출': 1,
              '안전밸브 오작동 및 조기개방': 1, '냉각수 손실': 1, '개수 산정근거': '가상 P&ID 가상-PID-RX-01 기준(연습)',
              '수동적 완화장치': '가상 방호벽', '능동적 완화장치': '가상 안전밸브 PSV-01·긴급차단', '안전성확보설비 증빙': '가상 설비 점검기록 DEMO-SAFE-03'},
    'CY-01': {'누출지점 배관길이(m)': 3, '비열비': 1.31, '연소열(kJ/kg)': 18600, '비점(℃)': -33.3, '액체비열(kJ/kg·K)': 4.7, '기화열(kJ/kg)': 1370,
              '증기밀도(kg/m3)': 0.7, '고압용기파열': 1, '배관누출': 1, '플랜지 등의 가스켓 파손': 1,
              '입/출하 시설 누출 사고': 1, '외부화재': 1, '개수 산정근거': '가상 용기 수와 연결부 수 기준(연습)',
              '수동적 완화장치': '가상 차양·방호 울타리', '능동적 완화장치': '가상 암모니아 감지기 연동 경보', '안전성확보설비 증빙': '가상 설비 점검기록 DEMO-SAFE-04'},
}


@lru_cache(maxsize=1)
def build():
    # 검증된 가상 회사의 물질·설비·비상대응 사실을 재사용한다.
    project = create_direct_cap_project(COMPANY, cap_group='1군')
    for name, content in demo.files().items():
        ew.ingest(project, name, content, ew.classify_document(name, content))
    entities = {}
    for slot in ew.slots(project):
        if slot['chosen']:
            entities.setdefault(slot['entity'], {})[slot['field']] = slot['chosen']['value']
    book = load_workbook(BytesIO(template.build_company()))
    wanted = template.company_columns()
    for row, item in enumerate(wanted['site'], 2):
        value = SITE_VALUES.get(item.id, entities[ew.SITE_ENTITY].get(item.evidence_field))
        if value is not None:
            book[template.SITE_SHEET].cell(row, 2, value)
    for kind, title in template.ENTITY_SHEETS.items():
        if title not in book.sheetnames:
            continue
        prefix = ('facility' if kind == 'scenario' else kind) + ':'
        ws = book[title]
        columns = {cell.value: cell.column for cell in ws[1]}
        for row, (entity, values) in enumerate(((e, v) for e, v in entities.items() if e.startswith(prefix)), 2):
            ws.cell(row, 1, entity.split(':', 1)[1])
            for item in wanted[kind]:
                value = values.get(item.evidence_field)
                if kind == 'material' and item.id == 'material.product_type':
                    value = '혼합물' if entity.endswith('P-02') else '단일물질'
                if kind == 'facility' and item.id == 'facility.pid_number':
                    value = '가상-PID-' + entity.split(':', 1)[1]
                if kind == 'facility' and item.id == 'facility.process_type':
                    value = '반응' if entity.endswith('RX-01') else '변화없음'
                if kind == 'scenario':
                    value = SCENARIO_VALUES.get(entity.split(':', 1)[1], {}).get(item.template_header)
                if kind == 'target' and item.id == 'target.location':
                    value = '가상 사업장 주변 위치(실제 주소·현장 확인 아님)'
                if value is not None:
                    ws.cell(row, columns[item.template_header], value)
    # 추가 명세도 실제 양식의 입력 열에만 넣는다. 예시 열은 변경하지 않는다.
    tables = {
        'cap-components': [
            {'혼합제품 코드': 'P-02', '혼합제품명': '가상 염산 35%', '구성성분 CAS 번호': '7647-01-0', '구성성분명': '염화수소', '구성성분 함량(%)': 35},
            {'혼합제품 코드': 'P-02', '혼합제품명': '가상 염산 35%', '구성성분 CAS 번호': '7732-18-5', '구성성분명': '물', '구성성분 함량(%)': 65},
        ],
        'cap-waste': [{'시설번호': 'SC-01', '시설명': '가상 스크러버', '처리대상 물질': '암모니아',
                       '처리방식': '습식 세정', '처리용량': '1000 m3/h', '연결 설비·배관': 'CY-01', '최종 배출지점': '가상 배출구'}],
        'cap-relief': [{'안전밸브·파열판 번호': 'PSV-01', '보호대상 설비번호': 'RX-01', '형식': '안전밸브',
                        '설정압력': '0.6 MPa', '배출용량': '1200 kg/h', '정격용량': '1300 kg/h',
                        '노즐크기 입구': '25A', '노즐크기 출구': '40A', '배출물질': '가상 반응기 증기',
                        '배출상태': '기체', '보호기기 운전압력': '0.3 MPa', '보호기기 설계압력': '0.8 MPa',
                        '몸체재질': 'SUS304', 'TRIM 재질': 'SUS316', '정밀도': '±3%',
                        '최종 배출·처리 지점': 'SC-01', '배출원인': '가상 냉각 이상', '관련 P&ID 번호': '가상-PID-RX-01'}],
        'cap-change': [{'일자': '2026-10-01', '변경항목': '가상 감지기 GD-01', '변경의 종류': '점검 기록',
                        '변경 내용(변경전 → 변경후)': '정기 점검 전 → 점검 완료(가상)', '후속조치': '차단 연동 기능 확인 연습', '담당자': '이담당'}],
    }
    for table, rows in tables.items():
        for row, values in enumerate(rows, 2):
            for col, item in enumerate(wanted['tables'][table], 1):
                book[template.sheet_title(table)].cell(row, col, values.get(item.column))
    book.properties.identifier = IDENTIFIER
    book.properties.title = COMPANY + ' · 가상 회사 작성 연습'
    book.properties.description = NOTICE
    out = BytesIO()
    book.save(out)
    return out.getvalue()
