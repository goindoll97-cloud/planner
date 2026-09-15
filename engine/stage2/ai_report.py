from __future__ import annotations

from io import BytesIO
from typing import Mapping
from zipfile import ZIP_DEFLATED, ZipFile

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph
from docx.shared import Pt, RGBColor

from .ai_drafting import ai_draft_field_key, ai_draft_is_current
from .intake import selected_requirement_specs
from .language_policy import assert_public_prose
from .project import Stage2Project
from .report_draft import build_report_draft, draft_filename


def _run_font(run, *, size: float = 9, bold: bool = False, color: str = "1F1F1F") -> None:
    run.font.name = "Malgun Gothic"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Malgun Gothic")
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def _insert_after(paragraph: Paragraph) -> Paragraph:
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    return Paragraph(new_p, paragraph._parent)


def _draft_payload(project: Stage2Project, system: str, requirement_key: str) -> tuple[Mapping[str, object], str] | None:
    record = project.get_field(ai_draft_field_key(system, requirement_key))
    if record is None or not isinstance(record.value, Mapping):
        return None
    # Never export prose generated from an older confirmed-fact snapshot. When
    # company facts change, Stage 5 regenerates the item using the new input
    # fingerprint before it can re-enter an AI-enhanced report.
    if not ai_draft_is_current(project, system, requirement_key):
        return None
    text = str(record.value.get("draft_text") or "").strip()
    if not text:
        return None
    return record.value, record.status


def has_ai_report_prose(project: Stage2Project, system: str) -> bool:
    system = str(system or "").strip().upper()
    for spec in selected_requirement_specs(project):
        if spec.system != system:
            continue
        if _draft_payload(project, system, spec.key):
            return True
    return False


def build_ai_enhanced_report_draft(project: Stage2Project, system: str) -> bytes:
    """Inject grounded AI prose after matching legal-writing headings.

    The deterministic report remains the source-of-truth skeleton. AI prose is
    a separate overlay and never replaces the underlying confirmed field data.
    Before output, every user-facing AI sentence is checked again against the
    legal-language policy so legacy drafts cannot leak developer terminology.
    """
    system = str(system or "").strip().upper()
    base = build_report_draft(project, system)
    doc = Document(BytesIO(base))

    specs = [spec for spec in selected_requirement_specs(project) if spec.system == system]
    by_label = {spec.label: spec for spec in specs}
    inserted = 0
    first_ai_heading: Paragraph | None = None

    for paragraph in list(doc.paragraphs):
        spec = by_label.get(paragraph.text.strip())
        if spec is None:
            continue
        payload_status = _draft_payload(project, system, spec.key)
        if payload_status is None:
            continue
        if first_ai_heading is None:
            first_ai_heading = paragraph
        payload, status = payload_status
        draft_text = assert_public_prose(str(payload.get("draft_text") or "").strip(), system)
        suggestions = payload.get("suggested_additions")

        marker = _insert_after(paragraph)
        marker.alignment = WD_ALIGN_PARAGRAPH.LEFT
        if status == "USER_CONFIRMED":
            run = marker.add_run("AI 보강문장 · 담당자 검토 완료")
            _run_font(run, size=8, bold=True, color="008000")
        else:
            run = marker.add_run("AI 보강 초안 · 사람 검토 필요")
            _run_font(run, size=8, bold=True, color="C00000")

        body = _insert_after(marker)
        for idx, line in enumerate(draft_text.splitlines() or [draft_text]):
            if idx:
                body.add_run("\n")
            run = body.add_run(line)
            _run_font(run, size=9)

        if isinstance(suggestions, list) and suggestions:
            safe_suggestions = [
                assert_public_prose(str(value), system)
                for value in suggestions
                if str(value).strip()
            ]
            if safe_suggestions:
                note = _insert_after(body)
                run = note.add_run("추가 확인하면 좋은 내용: " + " / ".join(safe_suggestions))
                _run_font(run, size=8, color="9C6500")
        inserted += 1

    # Both systems now open on the real regulation-form baseline itself (its
    # own title page, not a synthetic cover this program generated), so the
    # banner must never land ahead of that. Insert it right before the first
    # AI-touched heading instead of unconditionally at paragraph 0.
    if inserted and first_ai_heading is not None:
        target = first_ai_heading
        if target is not None:
            note = OxmlElement("w:p")
            target._p.addprevious(note)
            p = Paragraph(note, target._parent)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run("AI 문장 보강 적용 · 회사 확인자료는 원문 표/필드로 함께 보존")
            _run_font(run, size=8, bold=True, color="C00000")

    out = BytesIO()
    doc.save(out)
    return out.getvalue()


def build_ai_enhanced_draft_bundle(project: Stage2Project) -> bytes:
    systems: list[str] = []
    if project.psm_in_scope:
        systems.append("PSM")
    if project.cap_in_scope:
        systems.append("CAP")
    if not systems:
        raise ValueError("먼저 작성범위를 선택해야 보고서 초안을 생성할 수 있습니다.")

    out = BytesIO()
    with ZipFile(out, "w", compression=ZIP_DEFLATED) as archive:
        manifest = [
            "Stage 2 AI 문장 보강 검토용 보고서 묶음",
            f"프로젝트 ID: {project.project_id}",
            "주의: AI 보강문장은 회사 확인자료와 법정 작성구조를 바탕으로 한 검토용 초안이며 법정 최종본이 아닙니다.",
            "",
        ]
        for system in systems:
            filename = draft_filename(project, system).replace("_검토용_초안.docx", "_AI보강_검토용_초안.docx")
            archive.writestr(filename, build_ai_enhanced_report_draft(project, system))
            manifest.append(f"- {system}: AI 보강문장 {'있음' if has_ai_report_prose(project, system) else '없음'}")
        archive.writestr("AI보강_생성상태.txt", "\n".join(manifest).encode("utf-8"))
    return out.getvalue()
