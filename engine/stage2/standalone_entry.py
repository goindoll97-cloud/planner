from __future__ import annotations

from dataclasses import dataclass
from io import BytesIO
from typing import Any
import re

from openpyxl import load_workbook

from .integrated_workbook import (
    META_SHEET,
    MODE_INPUT,
    SCHEMA_VERSION,
    IntegratedImportResult,
    apply_integrated_authoring_workbook,
)
from .project import EvidenceRef, Stage2Project


DIRECT_ENTRY_MODE = "STAGE2_DIRECT_WORKBOOK"


@dataclass(frozen=True)
class StandaloneWorkbookPreview:
    project_id: str
    company_name: str
    address: str
    psm_selected: bool
    cap_selected: bool
    cap_group: str

    @property
    def scope_labels(self) -> tuple[str, ...]:
        labels: list[str] = []
        if self.psm_selected:
            labels.append("공정안전보고서")
        if self.cap_selected:
            labels.append("화학사고예방관리계획서")
        return tuple(labels)


def _bool_value(value: object) -> bool:
    text = str(value or "").strip().lower()
    return text in {"true", "1", "yes", "y", "예", "선택"}


def _cap_group(value: object) -> str:
    text = re.sub(r"\s+", "", str(value or ""))
    if "1군" in text:
        return "1군"
    if "2군" in text:
        return "2군"
    return ""


def _meta_header(ws) -> dict[str, str]:
    header: dict[str, str] = {}
    for row in range(1, 6):
        key = str(ws.cell(row, 1).value or "").strip()
        value = str(ws.cell(row, 2).value or "").strip()
        if key:
            header[key] = value
    return header


def _meta_records(ws) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    for raw in ws.iter_rows(min_row=8, values_only=True):
        if not any(value not in (None, "") for value in raw):
            continue
        rows.append(
            {
                "kind": str(raw[0] or "").strip(),
                "sheet": str(raw[1] or "").strip(),
                "field_keys": str(raw[2] or "").strip(),
                "row_or_header": int(raw[3]) if raw[3] not in (None, "") else 0,
                "value_column": int(raw[4]) if raw[4] not in (None, "") else 0,
            }
        )
    return rows


def _field_value(wb, records: list[dict[str, Any]], field_key: str) -> Any:
    for item in records:
        keys = [key for key in item["field_keys"].split("|") if key]
        if field_key not in keys or item["kind"] not in {"FORM", "PROTECTED"}:
            continue
        if item["sheet"] not in wb.sheetnames:
            continue
        if not item["row_or_header"] or not item["value_column"]:
            continue
        return wb[item["sheet"]].cell(item["row_or_header"], item["value_column"]).value
    return ""


def _parse_table(ws, header_row: int = 4) -> list[dict[str, Any]]:
    headers = [str(cell.value or "").strip() for cell in ws[header_row]]
    last_col = max((index for index, value in enumerate(headers, 1) if value), default=0)
    if not last_col:
        return []
    headers = headers[:last_col]
    rows: list[dict[str, Any]] = []
    blank_run = 0
    for row_no in range(header_row + 1, ws.max_row + 1):
        values = []
        for col in range(1, last_col + 1):
            value = ws.cell(row_no, col).value
            if isinstance(value, str):
                value = value.strip()
            values.append(value if value is not None else "")
        if not any(value not in (None, "") for value in values):
            blank_run += 1
            if blank_run >= 5:
                break
            continue
        blank_run = 0
        rows.append({header: value for header, value in zip(headers, values) if header})
    return rows


def inspect_standalone_workbook(workbook_bytes: bytes) -> StandaloneWorkbookPreview:
    """Read the program workbook contract needed to start Stage 2 directly.

    This does not determine legal applicability. It only reads the authoring
    scope encoded in the workbook and, for CAP, the visible writing-level value
    confirmed by the company.
    """
    wb = load_workbook(BytesIO(workbook_bytes), data_only=False)
    if META_SHEET not in wb.sheetnames:
        raise ValueError("프로그램 통합 작성자료가 아닙니다. _시스템정보 시트를 찾을 수 없습니다.")

    meta_ws = wb[META_SHEET]
    meta = _meta_header(meta_ws)
    if meta.get("schema_version") != SCHEMA_VERSION:
        raise ValueError("지원하지 않는 통합 작성자료 버전입니다. 최신 프로그램 입력파일을 사용해 주세요.")
    if meta.get("workbook_mode") != MODE_INPUT:
        raise ValueError("작성예시 파일로는 Stage 2를 시작할 수 없습니다. 실제 입력용 통합 작성자료를 사용해 주세요.")

    project_id = str(meta.get("project_id") or "").strip()
    if not re.fullmatch(r"[0-9A-Za-z._-]{3,180}", project_id):
        raise ValueError("통합 작성자료의 프로젝트 식별자가 없거나 올바르지 않습니다.")

    records = _meta_records(meta_ws)
    company_name = str(_field_value(wb, records, "business.company_name") or "").strip()
    address = str(_field_value(wb, records, "business.address") or "").strip()
    writing_level = _field_value(wb, records, "cap.business.writing_level")

    psm_selected = _bool_value(meta.get("psm_selected"))
    cap_selected = _bool_value(meta.get("cap_selected"))
    cap_group = _cap_group(writing_level)

    if not psm_selected and not cap_selected:
        raise ValueError("통합 작성자료에 작성할 보고서가 선택되어 있지 않습니다.")
    if not company_name:
        raise ValueError("01_사업장정보의 회사명을 입력한 뒤 다시 업로드해 주세요.")
    if cap_selected and cap_group not in {"1군", "2군"}:
        raise ValueError("화학사고예방관리계획서를 작성하려면 01_사업장정보의 작성수준을 '1군 사업장' 또는 '2군 사업장'으로 입력해 주세요.")

    return StandaloneWorkbookPreview(
        project_id=project_id,
        company_name=company_name,
        address=address,
        psm_selected=psm_selected,
        cap_selected=cap_selected,
        cap_group=cap_group,
    )


def is_standalone_stage2_project(project: Stage2Project) -> bool:
    return str(project.stage1_snapshot.get("entry_mode") or "") == DIRECT_ENTRY_MODE


def create_project_from_standalone_workbook(
    workbook_bytes: bytes,
    *,
    workbook_evidence: EvidenceRef | None = None,
) -> tuple[Stage2Project, IntegratedImportResult, StandaloneWorkbookPreview]:
    """Create an authoring-only project directly from a completed workbook.

    The resulting project deliberately does not claim that legal applicability
    was decided. The selected report flags enable the Stage 2 authoring,
    validation and draft-report pipeline only. A permanent marker preserves
    that distinction for later review.
    """
    preview = inspect_standalone_workbook(workbook_bytes)
    evidence = [workbook_evidence] if workbook_evidence is not None else []

    project = Stage2Project(
        project_id=preview.project_id,
        company_name=preview.company_name,
        psm_required=preview.psm_selected,
        cap_required=preview.cap_selected,
        cap_group=preview.cap_group if preview.cap_selected else "",
        scope_confirmed=True,
        psm_selected=preview.psm_selected,
        cap_selected=preview.cap_selected,
        stage1_snapshot={
            "entry_mode": DIRECT_ENTRY_MODE,
            "legal_applicability_confirmed": False,
            "source_project_id": preview.project_id,
            "decision": {},
            "business": {},
            "documents": {},
            "chemicals": [],
            "facilities": [],
        },
        notes=[
            "Stage 1 판정진단을 연결하지 않고 통합 작성자료에서 직접 시작한 작성 프로젝트입니다. "
            "작성범위는 법적 대상 판정이 아니라 사용자가 제공한 작성자료의 선택값입니다."
        ],
    )
    project.set_field(
        "business.company_name",
        "회사명",
        preview.company_name,
        "USER_CONFIRMED",
        evidence=evidence,
        note="Stage 2 직접 시작 통합 작성자료에서 회사가 입력한 값",
    )
    if preview.address:
        project.set_field(
            "business.address",
            "사업장 소재지",
            preview.address,
            "USER_CONFIRMED",
            evidence=evidence,
            note="Stage 2 직접 시작 통합 작성자료에서 회사가 입력한 값",
        )

    result = apply_integrated_authoring_workbook(
        project,
        workbook_bytes,
        workbook_evidence=workbook_evidence,
    )

    # A workbook originally produced for group 1 may later be edited to group 2
    # for a test or a changed authoring case. Never retain external emergency
    # response fields when the visible writing level is group 2.
    if preview.cap_selected and preview.cap_group == "2군":
        removed = False
        for key in list(project.fields):
            if key.startswith("cap.external."):
                del project.fields[key]
                removed = True
        if removed:
            project.touch()

    # Stage 1 normally supplies these common inventory facts. In direct Stage 2
    # mode, mirror the completed integrated tables so PSM/CAP share the same
    # company-confirmed facts without requiring a Stage 1 upload.
    wb = load_workbook(BytesIO(workbook_bytes), data_only=False)
    if "02_화학물질정보" in wb.sheetnames:
        chemicals = _parse_table(wb["02_화학물질정보"], 4)
        if chemicals:
            project.set_field(
                "inventory.chemicals",
                "화학물질 목록",
                chemicals,
                "USER_CONFIRMED",
                evidence=evidence,
                note="Stage 2 직접 시작 통합 작성자료의 화학물질 표에서 승계",
            )
            if preview.cap_selected and project.get_field("cap.chemical.details") is None:
                project.set_field(
                    "cap.chemical.details",
                    "유해화학물질 목록 및 취급량",
                    chemicals,
                    "USER_CONFIRMED",
                    evidence=evidence,
                    note="Stage 2 직접 시작 통합 작성자료의 화학물질 표에서 승계",
                )

    if "03_설비정보" in wb.sheetnames:
        facilities = _parse_table(wb["03_설비정보"], 4)
        if facilities:
            project.set_field(
                "inventory.facilities",
                "시설별 최대보유량 자료",
                facilities,
                "USER_CONFIRMED",
                evidence=evidence,
                note="Stage 2 직접 시작 통합 작성자료의 설비 표에서 승계",
            )
            if preview.psm_selected and project.get_field("psm.psi.equipment_specs") is None:
                project.set_field(
                    "psm.psi.equipment_specs",
                    "유해하거나 위험한 설비의 목록 및 사양",
                    facilities,
                    "USER_CONFIRMED",
                    evidence=evidence,
                    note="Stage 2 직접 시작 통합 작성자료의 설비 표에서 승계",
                )
            if preview.cap_selected and project.get_field("cap.facility.equipment_specs") is None:
                project.set_field(
                    "cap.facility.equipment_specs",
                    "장치·설비 목록 및 명세",
                    facilities,
                    "USER_CONFIRMED",
                    evidence=evidence,
                    note="Stage 2 직접 시작 통합 작성자료의 설비 표에서 승계",
                )

    return project, result, preview
