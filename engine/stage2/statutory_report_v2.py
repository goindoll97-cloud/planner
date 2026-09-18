from __future__ import annotations

from io import BytesIO

from docx import Document

from .project import Stage2Project
from . import statutory_report as base


# The base renderer uses small helper tables inside a larger statutory form.
# Those helper tables must not create a fake blank "■" form heading.
_original_add_form_heading = base._add_form_heading


def _safe_add_form_heading(doc: Document, reference: str, title: str) -> None:
    if not str(reference or "").strip() and not str(title or "").strip():
        return
    _original_add_form_heading(doc, reference, title)


base._add_form_heading = _safe_add_form_heading


CAP_ARTICLE_ITEMS = {
    "cap.prevention.safety_management": (
        "사업장의 종합적인 화학사고에 대한 안전관리 방향 및 목표",
        "설정된 목표를 달성하기 위한 구체적인 실행과제",
        "관리적 대책",
        "기술적 대책",
    ),
    "cap.prevention.training": (
        "연간 교육·훈련 계획",
        "화학사고예방관리계획서 전문교육 및 비상대응조직 역할별 교육·훈련",
        "교육·훈련의 평가 방법과 평가 결과에 따른 보완계획",
    ),
    "cap.prevention.self_inspection": (
        "자체점검반 구성, 점검시기, 점검항목 등 자체점검 계획",
        "자체점검 실시 및 자체점검결과 내부 보고체계",
        "자체점검결과를 활용한 안전관리운영계획·계획서 환류",
        "자체점검결과 화학물질안전원 서면 제출계획(1군 사업장만 해당)",
    ),
    "cap.prevention.change_management": (
        "변경내역 확인 담당자",
        "변경 확인 주기",
        "변경 확인 방법 및 변경내역 관리",
    ),
    "cap.prevention.emergency_contact": (
        "사업장 내·외부 사고신고 체계",
        "유관기관 목록 및 사고신고 체계",
        "인근 사업장 공조 연락체계(해당 시)",
        "총괄영향범위 내 다른 시·군 사고신고 체계(해당 시)",
    ),
    "cap.prevention.emergency_org": (
        "비상대응조직별 편성인원 및 임무",
        "협력업체 비상대응조직도 및 임무(해당 시)",
    ),
    "cap.prevention.command_center": (
        "비상통제실 지정",
        "비상통제실 운영에 필요한 계획서·개인보호장구·통신장비 등 비치 또는 신속 설치 방안",
    ),
    "cap.internal.shutdown": (
        "비상상황 발생 시 가동중지 권한",
        "비상 운전정지 판단 및 실행 절차",
    ),
    "cap.internal.resources": (
        "화학사고 초기 대응을 위한 자체 방재 인력 현황",
        "방재장비·물품 및 개인보호장구 보유현황 및 배치",
        "방재장비·물품 등의 관리·유지 및 확충 계획",
        "방재인력 및 장비·물품 운영에 필요한 사항",
    ),
    "cap.internal.communication": (
        "사내 경보시설의 종류 및 경보발령지점",
        "경보전달체계 및 경보전달 담당자",
        "경보시설 유지관리방법",
    ),
    "cap.internal.facility_response": (
        "취급시설 유형별 자동·수동 차단시스템",
        "단계별 내·외부 확산차단 또는 방지대책",
        "2차 오염 방지대책",
        "사내·외 비상대피, 응급의료 및 환자수송 계획",
    ),
    "cap.internal.investigation": (
        "사고조사팀의 구성 및 팀원의 역할",
        "사고조사보고서의 작성항목 및 작성방법",
        "개선대책 및 이행방법",
    ),
    "cap.internal.recovery": (
        "사고복구 조직 및 역할",
        "책임보험 가입계획(해당 시)",
        "폐기물처리 및 토양환경복원 등 환경복원 전문업체 활용계획",
    ),
    "cap.external.communication": (
        "화학사고 발생 시 대외소통 담당 조직 및 임무",
        "정보 제공 방법 또는 절차",
        "정보 제공이 필요한 이해당사자 및 제공정보",
        "평상시 지역사회와의 소통계획",
    ),
    "cap.external.mutual_aid": (
        "화학사고시 비상대응을 위한 협정 내역",
        "자사 보유 자원의 타사 지원 계획",
        "지역 비상대응기관 및 인근사업장과의 합동훈련계획",
    ),
    "cap.external.evacuation": (
        "사고유형에 따른 대피경보 방법",
        "인근 사업장·주민 등 대상별 경보전달 방법",
        "기초지방자치단체 담당부서 및 연락처",
        "주민행동 요령",
        "응급의료 계획",
        "주민대피장소 및 방법",
    ),
    "cap.external.public_notice": (
        "고지대상",
        "고지방법",
        "고지정보",
    ),
}


def _render_cap(doc: Document, project: Stage2Project) -> None:
    # The CAP deliverable should read like the statutory annex set itself.
    # Do not prepend a cover, source declaration, or synthetic category heading;
    # the first visible content is Annex Form 1.
    base._cap_form1(doc, project)
    base._cap_form3(doc, project)
    base._cap_facility_overview(doc, project, detailed=False)
    base._cap_facility_overview(doc, project, detailed=True)
    base._add_form_table(doc, base.CAP_FORMS["6"], base._cap_form6_rows(project))
    base._cap_form7(doc, project)
    base._cap_form8(doc, project)

    doc.add_page_break()
    doc.add_heading("시설정보", level=1)
    doc.add_heading("공정개요", level=2)
    doc.add_paragraph(base._text(project, "process.description", default=base.MISSING))
    base._add_attachment_line(doc, "공정흐름도(PFD)", base._doc_value(project, "documents.pfd"))
    base._add_attachment_line(doc, "공정배관계장도(P&ID)", base._doc_value(project, "documents.pid"))
    base._add_form_table(doc, base.CAP_FORMS["9"], base._cap_form9_rows(project))
    doc.add_heading("공정위험성 분석 자료", level=2)
    doc.add_paragraph(base._text(project, "cap.facility.process_hazard_analysis", default=base.MISSING))
    doc.add_heading("운전책임자 및 작업자 현황", level=2)
    doc.add_paragraph(base._text(project, "cap.facility.operator_staffing", default=base.MISSING))
    base._add_form_table(doc, base.CAP_FORMS["10"], base._cap_form10_rows(project))
    base._add_form_table(doc, base.CAP_FORMS["11"], base._cap_form11_rows(project))
    base._add_attachment_line(doc, "안전밸브 및 파열판 명세", base._text(project, "cap.safety.relief_device_specs", default=base.MISSING))
    base._add_attachment_line(doc, "배출물질 처리시설 현황", base._text(project, "cap.safety.waste_treatment", default=base.MISSING))

    doc.add_page_break()
    doc.add_heading("장외평가정보", level=1)
    doc.add_heading("예비시나리오 및 사고시나리오 선정", level=2)
    for label, key in (
        ("예비시나리오 대상 설비 선정", "cap.offsite.target_facility_selection"),
        ("대상 설비 취급량 산정", "cap.offsite.target_holding_calculation"),
        ("누출조건", "cap.offsite.release_conditions"),
        ("기상조건", "cap.offsite.weather_conditions"),
        ("영향범위 평가결과", "cap.offsite.impact_range_result"),
    ):
        p = doc.add_paragraph()
        p.add_run(label + ": ").bold = True
        p.add_run(base._text(project, key, default=base.MISSING))
    base._cap_form12(doc, project)
    base._cap_form13(doc, project)
    base._add_form_table(doc, base.CAP_FORMS["14"], base._cap_form14_rows(project))
    base._cap_form15(doc, project)

    specs = [s for s in base.selected_requirement_specs(project) if s.system == "CAP"]
    by_key = {s.key: s for s in specs}

    doc.add_page_break()
    doc.add_heading("사전관리방침", level=1)
    for key in (
        "cap.prevention.safety_management",
        "cap.prevention.training",
        "cap.prevention.self_inspection",
        "cap.prevention.change_management",
        "cap.prevention.emergency_contact",
        "cap.prevention.emergency_org",
        "cap.prevention.command_center",
    ):
        spec = by_key.get(key)
        if spec is not None:
            base._add_narrative_requirement(doc, spec, project, article_items=CAP_ARTICLE_ITEMS.get(key, ()))

    change_log = base._rows(project, "cap.prevention.change_log")
    if change_log:
        form = base.FormSpec(
            "별지 제2호서식",
            "화학사고예방관리계획서 변경내역 관리대장",
            ("번호", "일자", "변경항목", "변경의 종류", "변경 내용(변경전 → 변경후)", "후속조치", "담당자"),
        )
        rows = [[base._row_value(row, header) for header in form.headers] for row in change_log]
        base._add_form_table(doc, form, rows)

    doc.add_page_break()
    doc.add_heading("내부 비상대응계획", level=1)
    for key in (
        "cap.internal.shutdown",
        "cap.internal.resources",
        "cap.internal.communication",
        "cap.internal.facility_response",
        "cap.internal.investigation",
        "cap.internal.recovery",
    ):
        spec = by_key.get(key)
        if spec is not None:
            base._add_narrative_requirement(doc, spec, project, article_items=CAP_ARTICLE_ITEMS.get(key, ()))

    if project.cap_group == "1군":
        doc.add_page_break()
        doc.add_heading("외부 비상대응계획", level=1)
        for key in (
            "cap.external.communication",
            "cap.external.mutual_aid",
            "cap.external.evacuation",
            "cap.external.public_notice",
        ):
            spec = by_key.get(key)
            if spec is not None:
                base._add_narrative_requirement(doc, spec, project, article_items=CAP_ARTICLE_ITEMS.get(key, ()))

    doc.add_page_break()
    base._cap_form16(doc, project)


def _render_cap_narrative(doc: Document, project: Stage2Project) -> None:
    """Append CAP content that has no annex-form cell onto an existing document.

    Forms 1~16 (사업장의 작성수준 구분, 사업장 일반정보, 취급시설 개요,
    유해화학물질 목록, 장외평가정보 표, 비상대응분야 요약서 등) are the
    regulation-form baseline itself, filled separately by
    ``cap_baseline_docx.build_cap_baseline_draft``. This only adds the
    surrounding content the statute also requires but that has no annex-form
    counterpart: process/PFD references, 공정위험성 분석·운전책임자 현황,
    예비시나리오 선정 근거, and the free-text 사전관리방침/내부·외부
    비상대응계획 chapters.
    """
    doc.add_page_break()
    caption = doc.add_paragraph("법정서식 기반 검토용 작성본 · 별지서식에 없는 서술형 작성항목")
    caption.runs[0].italic = True

    base._add_heading_safe(doc, "시설정보", level=1)
    base._add_attachment_line(doc, "공정흐름도(PFD)", base._doc_value(project, "documents.pfd"))
    base._add_attachment_line(doc, "공정배관계장도(P&ID)", base._doc_value(project, "documents.pid"))
    base._add_heading_safe(doc, "공정위험성 분석 자료", level=2)
    doc.add_paragraph(base._text(project, "cap.facility.process_hazard_analysis", default=base.MISSING))
    base._add_heading_safe(doc, "운전책임자 및 작업자 현황", level=2)
    doc.add_paragraph(base._text(project, "cap.facility.operator_staffing", default=base.MISSING))
    base._add_attachment_line(doc, "안전밸브 및 파열판 명세", base._text(project, "cap.safety.relief_device_specs", default=base.MISSING))
    base._add_attachment_line(doc, "배출물질 처리시설 현황", base._text(project, "cap.safety.waste_treatment", default=base.MISSING))

    base._add_heading_safe(doc, "장외평가정보", level=1)
    base._add_heading_safe(doc, "예비시나리오 및 사고시나리오 선정", level=2)
    for label, key in (
        ("예비시나리오 대상 설비 선정", "cap.offsite.target_facility_selection"),
        ("대상 설비 취급량 산정", "cap.offsite.target_holding_calculation"),
        ("누출조건", "cap.offsite.release_conditions"),
        ("기상조건", "cap.offsite.weather_conditions"),
        ("영향범위 평가결과", "cap.offsite.impact_range_result"),
    ):
        p = doc.add_paragraph()
        p.add_run(label + ": ").bold = True
        p.add_run(base._text(project, key, default=base.MISSING))

    specs = [s for s in base.selected_requirement_specs(project) if s.system == "CAP"]
    by_key = {s.key: s for s in specs}

    doc.add_page_break()
    base._add_heading_safe(doc, "사전관리방침", level=1)
    for key in (
        "cap.prevention.safety_management",
        "cap.prevention.training",
        "cap.prevention.self_inspection",
        "cap.prevention.change_management",
        "cap.prevention.emergency_contact",
        "cap.prevention.emergency_org",
        "cap.prevention.command_center",
    ):
        spec = by_key.get(key)
        if spec is not None:
            base._add_narrative_requirement(doc, spec, project, article_items=CAP_ARTICLE_ITEMS.get(key, ()))

    doc.add_page_break()
    base._add_heading_safe(doc, "내부 비상대응계획", level=1)
    for key in (
        "cap.internal.shutdown",
        "cap.internal.resources",
        "cap.internal.communication",
        "cap.internal.facility_response",
        "cap.internal.investigation",
        "cap.internal.recovery",
    ):
        spec = by_key.get(key)
        if spec is not None:
            base._add_narrative_requirement(doc, spec, project, article_items=CAP_ARTICLE_ITEMS.get(key, ()))

    if project.cap_group == "1군":
        doc.add_page_break()
        base._add_heading_safe(doc, "외부 비상대응계획", level=1)
        for key in (
            "cap.external.communication",
            "cap.external.mutual_aid",
            "cap.external.evacuation",
            "cap.external.public_notice",
        ):
            spec = by_key.get(key)
            if spec is not None:
                base._add_narrative_requirement(doc, spec, project, article_items=CAP_ARTICLE_ITEMS.get(key, ()))
    else:
        doc.add_paragraph("2군 사업장은 외부 비상대응계획 내용을 생략할 수 있으므로 이 작성본에는 포함하지 않습니다.")


def build_statutory_report_draft(project: Stage2Project, system: str, status) -> bytes:
    system = str(system or "").strip().upper()
    if system == "PSM" and not project.psm_in_scope:
        raise ValueError("공정안전보고서는 현재 작성범위에 포함되어 있지 않습니다.")
    if system == "CAP" and not project.cap_in_scope:
        raise ValueError("화학사고예방관리계획서는 현재 작성범위에 포함되어 있지 않습니다.")
    if system not in {"PSM", "CAP"}:
        raise ValueError(f"지원하지 않는 보고서 종류입니다: {system}")

    if system == "PSM":
        from .psm_baseline_docx import build_psm_baseline_draft

        doc = Document(BytesIO(build_psm_baseline_draft(project)))
        base._render_psm_narrative(doc, project)
    else:
        from .cap_baseline_docx import build_cap_baseline_draft

        doc = Document(BytesIO(build_cap_baseline_draft(project)))
        _render_cap_narrative(doc, project)
    base._add_review_appendix(doc, project, system, status)
    out = BytesIO()
    doc.save(out)
    return out.getvalue()
