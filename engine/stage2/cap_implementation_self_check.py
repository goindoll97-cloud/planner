from __future__ import annotations

"""화학사고예방관리계획서 자체점검(이행규정 별지 1~3호) 작성지원."""

from dataclasses import dataclass
from functools import lru_cache
from io import BytesIO
import json
import zipfile
from pathlib import Path
import re
from typing import Any, Mapping

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

from .ooxml_determinism import canonicalize_docx_zip
from .project import CONFIRMED_STATUSES, Stage2Project
from . import versioning

ROOT = Path(__file__).resolve().parents[2]
CHECKLIST_PATH = ROOT / "data" / "stage2" / "cap_implementation_self_check.json"

CHECKLIST_KEY = "cap.implementation.checklist"
IMPROVEMENT_KEY = "cap.implementation.improvements"
TEAM_KEY = "cap.implementation.team"
CONFIRMER_KEY = "cap.implementation.confirmers"
BASE_VERSION_KEY = "cap.implementation.base_version_id"

CHECK_STATUSES = ("", "확인", "개선필요", "해당없음")
IMPROVEMENT_COLUMNS = (
    "연번", "자체점검결과 개선사항", "조치결과", "조치일자",
    "책임부서 (담당자)", "확인자", "서명",
)
PERSON_COLUMNS = ("소속", "직급", "성명", "서명")


@dataclass(frozen=True)
class SelfCheckReadiness:
    ready: bool
    blockers: tuple[str, ...]
    total: int
    checked: int
    improvements_required: int
    improvements_completed: int


def _clean(value: object) -> str:
    return "" if value is None else str(value).strip()


def _value(project: Stage2Project, *keys: str) -> str:
    for key in keys:
        rec = project.get_field(key)
        if rec is not None and rec.status in CONFIRMED_STATUSES:
            value = _clean(rec.value)
            if value:
                return value
    return ""


@lru_cache(maxsize=1)
def checklist_definition() -> tuple[dict[str, str], ...]:
    raw = json.loads(CHECKLIST_PATH.read_text(encoding="utf-8"))
    if raw.get("schema_version") != "cap-implementation-self-check-v1":
        raise ValueError("지원하지 않는 자체점검표 스키마입니다.")
    return tuple(dict(item) for item in raw["items"])


def _saved_map(project: Stage2Project) -> dict[str, dict[str, str]]:
    rec = project.get_field(CHECKLIST_KEY)
    if rec is None or rec.status not in CONFIRMED_STATUSES or not isinstance(rec.value, list):
        return {}
    result = {}
    for row in rec.value:
        if isinstance(row, Mapping) and _clean(row.get("id")):
            result[_clean(row.get("id"))] = {
                "status": _clean(row.get("status")),
                "note": _clean(row.get("note")),
            }
    return result


def checklist_rows(project: Stage2Project) -> list[dict[str, str]]:
    saved = _saved_map(project)
    rows = []
    for item in checklist_definition():
        state = saved.get(item["id"], {})
        rows.append({
            **item,
            "status": state.get("status", ""),
            "note": state.get("note", ""),
        })
    return rows


def save_checklist(project: Stage2Project, rows: list[Mapping[str, object]]) -> int:
    legal_ids = {item["id"] for item in checklist_definition()}
    cleaned: list[dict[str, str]] = []
    for row in rows:
        item_id = _clean(row.get("id"))
        if item_id not in legal_ids:
            continue
        status = _clean(row.get("status"))
        if status not in CHECK_STATUSES:
            raise ValueError(f"지원하지 않는 자체점검 결과입니다: {status}")
        cleaned.append({
            "id": item_id,
            "status": status,
            "note": _clean(row.get("note")),
        })
    project.set_field(
        CHECKLIST_KEY,
        "화학사고예방관리계획서 자체점검표 결과",
        cleaned,
        "USER_CONFIRMED",
        note="이행 등에 관한 규정 별지 제2호에 대한 회사 자체점검 결과",
    )
    return len(cleaned)


def _list_rows(project: Stage2Project, key: str) -> list[dict[str, Any]]:
    rec = project.get_field(key)
    if rec is None or rec.status not in CONFIRMED_STATUSES or not isinstance(rec.value, list):
        return []
    return [dict(row) for row in rec.value if isinstance(row, Mapping)]


def save_people(project: Stage2Project, key: str, label: str, rows: list[Mapping[str, object]]) -> int:
    cleaned = []
    for row in rows:
        item = {column: _clean(row.get(column)) for column in PERSON_COLUMNS}
        if any(item.values()):
            cleaned.append(item)
    project.set_field(key, label, cleaned, "USER_CONFIRMED")
    return len(cleaned)


def self_check_team(project: Stage2Project) -> list[dict[str, Any]]:
    return _list_rows(project, TEAM_KEY)


def confirmers(project: Stage2Project) -> list[dict[str, Any]]:
    return _list_rows(project, CONFIRMER_KEY)


def improvement_rows(project: Stage2Project) -> list[dict[str, Any]]:
    return _list_rows(project, IMPROVEMENT_KEY)


def proposed_improvements(project: Stage2Project) -> list[dict[str, Any]]:
    current = {str(row.get("_check_id") or ""): dict(row) for row in improvement_rows(project)}
    items = {row["id"]: row for row in checklist_rows(project) if row["status"] == "개선필요"}
    result = []
    for index, (item_id, item) in enumerate(items.items(), start=1):
        previous = current.get(item_id, {})
        result.append({
            "_check_id": item_id,
            "연번": str(index),
            "자체점검결과 개선사항": previous.get("자체점검결과 개선사항") or item["check"],
            "조치결과": previous.get("조치결과", ""),
            "조치일자": previous.get("조치일자", ""),
            "책임부서 (담당자)": previous.get("책임부서 (담당자)", ""),
            "확인자": previous.get("확인자", ""),
            "서명": previous.get("서명", ""),
        })
    return result


def save_improvements(project: Stage2Project, rows: list[Mapping[str, object]]) -> int:
    cleaned = []
    for row in rows:
        item = {column: _clean(row.get(column)) for column in IMPROVEMENT_COLUMNS}
        item["_check_id"] = _clean(row.get("_check_id"))
        if any(item[column] for column in IMPROVEMENT_COLUMNS):
            cleaned.append(item)
    project.set_field(
        IMPROVEMENT_KEY,
        "자체점검 결과 개선사항 조치 내역서",
        cleaned,
        "USER_CONFIRMED",
        note="이행 등에 관한 규정 별지 제3호",
    )
    return len(cleaned)


def save_header(project: Stage2Project, values: Mapping[str, object]) -> None:
    labels = {
        BASE_VERSION_KEY: "자체점검 기준 화학사고예방관리계획서 버전",
        "cap.implementation.target_process": "자체점검 대상 공정",
        "cap.implementation.workplace_registration_no": "자체점검 사업장등록번호",
        "cap.implementation.business_permit_type": "자체점검 영업허가 구분",
        "cap.implementation.period_start": "자체점검 시작일",
        "cap.implementation.period_end": "자체점검 종료일",
        "cap.implementation.report_date": "자체점검 결과서 제출일",
    }
    for key, label in labels.items():
        if key in values and _clean(values[key]):
            project.set_field(key, label, _clean(values[key]), "USER_CONFIRMED")


def header_values(project: Stage2Project) -> dict[str, str]:
    return {
        "상호(명칭)": project.company_name or _value(project, "business.company_name"),
        "대상 공정": _value(project, "cap.implementation.target_process", "cap.business.unit_plant_name"),
        "성명(대표자)": _value(project, "business.representative", "cap.business.representative"),
        "사업장등록번호": _value(project, "cap.implementation.workplace_registration_no"),
        "표준산업분류(업종번호)": _value(project, "business.ksic"),
        "영업허가 구분": _value(project, "cap.implementation.business_permit_type"),
        "주소(사업장)": _value(project, "business.address"),
        "자체점검 시작일": _value(project, "cap.implementation.period_start"),
        "자체점검 종료일": _value(project, "cap.implementation.period_end"),
        "제출일": _value(project, "cap.implementation.report_date"),
        "기준 버전": _value(project, BASE_VERSION_KEY),
    }


def is_major_facility(project: Stage2Project) -> bool:
    # 기존 CAP 엔진에서 1군은 규칙 제19조제8항 주요취급시설로 판정된 값이다.
    return project.cap_group == "1군"


def readiness(project: Stage2Project) -> SelfCheckReadiness:
    blockers: list[str] = []
    header = header_values(project)
    for label in (
        "상호(명칭)", "대상 공정", "성명(대표자)", "사업장등록번호",
        "표준산업분류(업종번호)", "영업허가 구분", "주소(사업장)",
        "자체점검 시작일", "자체점검 종료일", "제출일", "기준 버전",
    ):
        if not header.get(label):
            blockers.append(f"별지 제1호 '{label}' 값이 확인되지 않았습니다.")

    base = header.get("기준 버전", "")
    if base:
        versions = {meta.version_id for meta in versioning.list_versions(project.project_id, "CAP")}
        if base not in versions:
            blockers.append("자체점검 기준 버전을 찾을 수 없습니다.")

    team = self_check_team(project)
    if not team:
        blockers.append("자체점검반이 확인되지 않았습니다.")
    elif any(not _clean(row.get("성명")) for row in team):
        blockers.append("자체점검반 구성원 성명이 비어 있습니다.")

    rows = checklist_rows(project)
    missing = [row for row in rows if row["status"] not in {"확인", "개선필요", "해당없음"}]
    if missing:
        blockers.append(f"별지 제2호 자체점검표에 미확인 항목이 {len(missing)}건 남아 있습니다.")

    required = [row for row in rows if row["status"] == "개선필요"]
    improvements = proposed_improvements(project)
    complete = 0
    for row in improvements:
        if all(_clean(row.get(column)) for column in (
            "자체점검결과 개선사항", "조치결과", "조치일자", "책임부서 (담당자)", "확인자",
        )):
            complete += 1
    if required and complete < len(required):
        blockers.append(f"별지 제3호 개선조치가 {len(required) - complete}건 미완료입니다.")

    if is_major_facility(project):
        change_log = project.get_field("cap.prevention.change_log")
        if change_log is None or change_log.status not in CONFIRMED_STATUSES:
            blockers.append("주요취급시설은 작성 규정 별지 제2호 변경내역 관리대장을 함께 확인해야 합니다.")

    return SelfCheckReadiness(
        ready=not blockers,
        blockers=tuple(blockers),
        total=len(rows),
        checked=len(rows) - len(missing),
        improvements_required=len(required),
        improvements_completed=complete,
    )


def _set_cell(cell, text: object, *, bold: bool = False, center: bool = False, size: int = 8) -> None:
    cell.text = _clean(text)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    for p in cell.paragraphs:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_after = Pt(0)
        for run in p.runs:
            run.font.name = "Malgun Gothic"
            run._element.rPr.rFonts.set(qn("w:eastAsia"), "Malgun Gothic")
            run.font.size = Pt(size)
            run.bold = bold


def _borders(table) -> None:
    tblPr = table._tbl.tblPr
    borders = tblPr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tblPr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = borders.find(qn("w:" + edge))
        if el is None:
            el = OxmlElement("w:" + edge)
            borders.append(el)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "5")
        el.set(qn("w:color"), "777777")


def _doc_landscape() -> Document:
    doc = Document()
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width, sec.page_height = Cm(29.7), Cm(21)
    sec.top_margin = sec.bottom_margin = Cm(1.1)
    sec.left_margin = sec.right_margin = Cm(1.1)
    return doc


def _heading(doc: Document, text: str, size: int = 16) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    run.bold = True
    run.font.name = "Malgun Gothic"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Malgun Gothic")
    run.font.size = Pt(size)


def build_form1_docx(project: Stage2Project) -> bytes:
    state = readiness(project)
    header = header_values(project)
    # 별지 1은 다른 별지의 미완료 때문에 출력을 막지 않고, 자체 표지 필수정보만 fail-closed.
    missing = [label for label in (
        "상호(명칭)", "대상 공정", "성명(대표자)", "사업장등록번호",
        "표준산업분류(업종번호)", "영업허가 구분", "주소(사업장)",
        "자체점검 시작일", "자체점검 종료일", "제출일",
    ) if not header.get(label)]
    if missing:
        raise ValueError("별지 제1호 필수값 미확인: " + ", ".join(missing))

    doc = _doc_landscape()
    _heading(doc, "화학사고예방관리계획서 자체점검 결과서")
    table = doc.add_table(rows=4, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    _borders(table)
    pairs = (
        ("상호(명칭)", header["상호(명칭)"], "대상 공정", header["대상 공정"]),
        ("성명(대표자)", header["성명(대표자)"], "사업장등록번호", header["사업장등록번호"]),
        ("표준산업분류(업종번호)", header["표준산업분류(업종번호)"], "영업허가 구분", header["영업허가 구분"]),
        ("주소(사업장)", header["주소(사업장)"], "", ""),
    )
    for r, pair in enumerate(pairs):
        for c, value in enumerate(pair):
            _set_cell(table.cell(r, c), value, bold=c % 2 == 0, center=c % 2 == 0)

    p = doc.add_paragraph(f"자체점검 일자: {header['자체점검 시작일']} ~ {header['자체점검 종료일']}")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph(
        "「화학물질관리법 시행규칙」 제19조의3제1항제1호 및 「화학사고예방관리계획서 이행 등에 관한 규정」 "
        "제4조제1항에 따른 화학사고예방관리계획서의 자체점검 결과를 제출합니다."
    )
    p = doc.add_paragraph(f"{header['제출일']}\n사업장 대표  {header['성명(대표자)']}  (서명 또는 인)")
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT

    for title, rows in (("자체점검반", self_check_team(project)), ("사업장 확인자", confirmers(project))):
        doc.add_paragraph(title)
        t = doc.add_table(rows=max(2, len(rows) + 1), cols=4)
        _borders(t)
        for c, label in enumerate(PERSON_COLUMNS):
            _set_cell(t.cell(0, c), label, bold=True, center=True)
        for r, row in enumerate(rows, start=1):
            for c, label in enumerate(PERSON_COLUMNS):
                _set_cell(t.cell(r, c), row.get(label, ""), center=True)

    doc.add_paragraph("첨부서류: 화학사고예방관리계획서 자체점검 내용 및 결과 1부")
    out = BytesIO()
    doc.save(out)
    return canonicalize_docx_zip(out.getvalue())


def build_form2_docx(project: Stage2Project) -> bytes:
    rows = checklist_rows(project)
    if any(row["status"] not in {"확인", "개선필요", "해당없음"} for row in rows):
        raise ValueError("별지 제2호에 미확인 항목이 남아 있어 작성을 중단합니다.")
    doc = _doc_landscape()
    _heading(doc, "화학사고예방관리계획서 자체점검표")
    table = doc.add_table(rows=1, cols=6)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    _borders(table)
    headers = ("구분", "항목", "연번", "세부항목", "확인할 항목", "확인")
    for c, label in enumerate(headers):
        _set_cell(table.cell(0, c), label, bold=True, center=True)
    for row in rows:
        cells = table.add_row().cells
        values = (row["section"], row["category"], row["no"], row["detail"], row["check"], row["status"])
        for c, value in enumerate(values):
            _set_cell(cells[c], value, center=c in {2, 5}, size=7)
    doc.add_paragraph("※ 위 자체점검표를 활용하여 자체점검 시 누락된 항목은 없는지 확인하며, 사업장에 해당 없는 내용은 [확인]란에 해당없음으로 표시한다.")
    out = BytesIO()
    doc.save(out)
    return canonicalize_docx_zip(out.getvalue())


def build_form3_docx(project: Stage2Project) -> bytes:
    rows = proposed_improvements(project)
    if not rows:
        rows = [{column: "" for column in IMPROVEMENT_COLUMNS}]
    doc = _doc_landscape()
    _heading(doc, "화학사고예방관리계획서 자체점검 결과 개선사항 조치 내역서")
    table = doc.add_table(rows=1, cols=len(IMPROVEMENT_COLUMNS))
    _borders(table)
    for c, label in enumerate(IMPROVEMENT_COLUMNS):
        _set_cell(table.cell(0, c), label, bold=True, center=True)
    for row in rows:
        cells = table.add_row().cells
        for c, label in enumerate(IMPROVEMENT_COLUMNS):
            _set_cell(cells[c], row.get(label, ""), center=c in {0, 3, 5, 6})
    doc.add_paragraph("※ 별지 제2호서식 자체점검표를 활용하여 자체점검결과 개선사항 조치내역을 작성하고, 개선 전ㆍ후 비교, 사진 증빙 등 추가자료는 별첨으로 제출")
    out = BytesIO()
    doc.save(out)
    return canonicalize_docx_zip(out.getvalue())


def form_filename(project: Stage2Project, no: int) -> str:
    company = re.sub(r"[^0-9A-Za-z가-힣._-]+", "_", project.company_name or project.project_id).strip("._") or "사업장"
    titles = {
        1: "자체점검_결과서",
        2: "자체점검표",
        3: "자체점검_개선사항_조치내역서",
    }
    return f"{company}_이행점검_별지제{no}호_{titles[no]}.docx"


def build_change_log_docx(project: Stage2Project) -> bytes:
    """주요취급시설 연간 제출용: 작성 규정 별지 제2호 변경내역 관리대장."""
    rec = project.get_field("cap.prevention.change_log")
    confirmed = rec is not None and rec.status in CONFIRMED_STATUSES and isinstance(rec.value, list)
    rows = rec.value if confirmed else []
    if is_major_facility(project) and not confirmed:
        raise ValueError("작성 규정 별지 제2호 변경내역 관리대장이 확인되지 않았습니다.")

    doc = _doc_landscape()
    _heading(doc, "화학사고예방관리계획서 변경내역 관리대장")
    header = doc.add_table(rows=1, cols=4)
    _borders(header)
    values = (
        ("사업장명", project.company_name),
        ("단위공장명", _value(project, "cap.business.unit_plant_name") or project.site_name),
    )
    for idx, (label, value) in enumerate(values):
        _set_cell(header.cell(0, idx * 2), label, bold=True, center=True)
        _set_cell(header.cell(0, idx * 2 + 1), value)

    columns = ("일자", "변경항목", "변경의 종류", "변경 내용(변경전 → 변경후)", "후속조치", "담당자")
    table = doc.add_table(rows=1, cols=len(columns))
    _borders(table)
    for i, label in enumerate(columns):
        _set_cell(table.cell(0, i), label, bold=True, center=True)
    for row in rows or [{}]:
        cells = table.add_row().cells
        for i, label in enumerate(columns):
            _set_cell(cells[i], row.get(label, ""), size=7)
    out = BytesIO()
    doc.save(out)
    return canonicalize_docx_zip(out.getvalue())


def build_submission_package(project: Stage2Project) -> bytes:
    state = readiness(project)
    if not state.ready:
        raise ValueError("자체점검 제출 패키지를 만들기 전에 보완할 항목이 있습니다: " + " / ".join(state.blockers))
    output = BytesIO()
    with zipfile.ZipFile(output, "w", compression=zipfile.ZIP_DEFLATED) as zf:
        zf.writestr(form_filename(project, 1), build_form1_docx(project))
        zf.writestr(form_filename(project, 2), build_form2_docx(project))
        zf.writestr(form_filename(project, 3), build_form3_docx(project))
        if is_major_facility(project):
            company = re.sub(r"[^0-9A-Za-z가-힣._-]+", "_", project.company_name or project.project_id).strip("._") or "사업장"
            zf.writestr(f"{company}_작성규정_별지제2호_변경내역_관리대장.docx", build_change_log_docx(project))
    return output.getvalue()


def submission_package_filename(project: Stage2Project) -> str:
    company = re.sub(r"[^0-9A-Za-z가-힣._-]+", "_", project.company_name or project.project_id).strip("._") or "사업장"
    return f"{company}_화학사고예방관리계획서_자체점검_제출패키지.zip"
