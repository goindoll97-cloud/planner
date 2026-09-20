from __future__ import annotations

from dataclasses import dataclass
from io import BytesIO
from pathlib import Path
from typing import Any, Iterable, Mapping
import re

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .project import EvidenceRef, Stage2Project
from .psm_applicability import requirement_explicitly_not_applicable
from .requirements import cap_field_labels, psm_field_labels
from .intake import field_label, selected_requirement_specs


SCHEMA_VERSION = "stage2-integrated-authoring-v1"
META_SHEET = "_시스템정보"
MODE_INPUT = "INPUT"
MODE_EXAMPLE = "EXAMPLE"

PSM_FULL = "공정안전보고서"
CAP_FULL = "화학사고예방관리계획서"

HEADER_FILL = "1F4E78"
SUBHEADER_FILL = "D9EAF7"
INPUT_FILL = "FFF2CC"
EXAMPLE_FILL = "E2F0D9"
NOTE_FILL = "F2F2F2"

PROTECTED_STAGE1_FIELDS = {
    "business.company_name",
    "business.address",
    "inventory.chemicals",
    "inventory.facilities",
}

TABLE_SPECS: tuple[dict[str, Any], ...] = (
    {
        "sheet": "02_화학물질정보",
        "scope": "COMMON",
        "title": "화학물질 정보 및 회사 SDS 확인값",
        "targets": ("inventory.chemicals", "psm.psi.chemical_details", "cap.chemical.details"),
        "headers": (
            "물질명", "CAS 번호", "분자식", "함량(%)", "물리적 상태", "최대보유량", "단위", "일일사용량",
            "사용·저장 공정", "주요 용도",
            "비중", "폭발한계 하한", "폭발한계 상한", "노출기준", "독성치", "인화점", "발화점", "이상반응 유무",
            "독성구분 항목", "독성구분", "위험노출수준", "허용농도값",
            "증기압", "부식성", "SDS 파일명", "SDS 개정일", "비고",
        ),
        "example": (
            (
                "톨루엔", "108-88-3", "C7H8", 99.5, "액체", 15000, "kg", "2500 kg/day",
                "원료 저장·혼합", "원료", "0.87", "1.2 vol%", "7.1 vol%", "TWA 50 ppm",
                "LD50 2600 mg/kg (경구, rat)", "4 ℃", "480 ℃", "해당 없음",
                "급성독성(흡입)", "구분 4", "ERPG-2 300 ppm", "TWA 50 ppm",
                "28.4 mmHg (25℃)", "해당 없음", "toluene_company_SDS.pdf", "2026-06-01",
                "회사 제품 SDS 전사 예시"
            ),
        ),
    },
    {
        "sheet": "03_설비정보",
        "scope": "COMMON",
        "title": "장치·설비 통합정보",
        "targets": ("inventory.facilities", "psm.psi.equipment_specs", "cap.facility.equipment_specs"),
        "headers": (
            "설비번호", "설비명", "설비종류", "단위공장·공정", "취급물질", "용량", "용량단위",
            "설계압력", "설계온도", "운전압력", "운전온도",
            "재질", "부속품재질", "개스킷재질", "용접효율", "계산두께", "부식여유", "사용두께",
            "후열처리 여부", "비파괴검사율", "최대보유량(kg)", "P&ID 번호", "비고", "최대 연결구 크기(mm)",
        ),
        "example": (
            (
                "TK-101", "톨루엔 저장탱크", "저장탱크", "원료저장", "톨루엔", 20, "m3",
                "0.49 MPa", "80 ℃", "0.15 MPa", "30 ℃",
                "SUS304", "SUS304", "PTFE", "1.0", "6.0 mm", "1.0 mm", "8.0 mm",
                "해당 없음", "10%", 15000, "PID-101", "예시값", 50
            ),
        ),
    },
    {
        "sheet": "04_안전밸브_파열판",
        "scope": "COMMON",
        "title": "안전밸브 및 파열판 통합정보",
        "targets": ("psm.psi.relief_device_specs", "cap.safety.relief_device_specs"),
        "headers": (
            "안전밸브·파열판 번호", "보호대상 설비번호", "형식", "설정압력",
            "배출용량", "정격용량", "노즐크기 입구", "노즐크기 출구",
            "배출물질", "배출상태", "보호기기 운전압력", "보호기기 설계압력",
            "몸체재질", "TRIM 재질", "정밀도", "최종 배출·처리 지점", "배출원인",
            "관련 P&ID 번호", "비고",
        ),
        "example": (
            (
                "PSV-101", "TK-101", "안전밸브", "0.45 MPa",
                "1200 kg/h", "1300 kg/h", "25A", "40A",
                "톨루엔 증기", "기체", "0.15 MPa", "0.49 MPa",
                "WCB", "SUS316", "±3%", "스크러버", "화재",
                "PID-101", "예시값"
            ),
        ),
    },
    {
        "sheet": "05_가스누출감지_경보장치",
        "scope": "COMMON",
        "title": "가스누출감지 및 경보장치 통합정보",
        "targets": ("psm.psi.gas_detection_table", "cap.safety.gas_detection"),
        "headers": (
            "감지기 번호", "설치형태", "설치위치", "검출대상 물질",
            "작동시간", "측정방식", "경보 설정값", "경보 위치",
            "연동여부", "연동 설비·조치", "정밀도", "유지관리",
            "비상전원 여부", "관련 도면번호", "비고",
        ),
        "example": (
            ("GD-101", "고정식", "TK-101 방유제 내", "톨루엔",
             "30초 이내", "접촉연소식", "10% LEL", "중앙제어실",
             "예", "HH 경보 시 원료 이송펌프 정지", "±3% F.S.", "월 1회 기능점검·연 1회 교정",
             "예", "GA-101", "예시값"),
        ),
    },
    {
        "sheet": "08_PSM_별지12_사업개요",
        "scope": "PSM",
        "title": "별지 제12호서식 입력자료 · 사업개요",
        "targets": ("psm.business.form12_details",),
        "headers": (
            "사업장명", "제출구분", "사업자등록번호", "대표자", "대상 유해·위험설비",
            "한국표준산업분류", "근로자수", "계약전력(kW)", "작성자 성명", "작성자 자격",
            "주요 원료", "주요 생산품", "사업개요", "사업장 소재지", "전화번호", "전송번호",
            "부지면적", "주요 건물", "총 사업기간", "착공예정일", "시운전기간",
        ),
        "example": (
            (
                "예시화학 제1공장", "변경", "123-45-67890", "홍길동", "반응·저장 공정",
                "C20119", 85, 1500, "김작성", "산업안전기사",
                "톨루엔, 아세톤", "혼합제품 A", "원료 저장·혼합·제품출하 공정",
                "대구광역시 ○○구 산업로 1", "053-000-0000", "053-000-0001",
                "12,500㎡", "생산동 2동 / 연면적 3,200㎡", "2026-10-01 ~ 2027-03-31",
                "2026-10-01", "2027-03-01 ~ 2027-03-31",
            ),
        ),
    },
    {
        "sheet": "10_동력기계",
        "scope": "PSM",
        "title": "동력기계 목록",
        "targets": ("psm.psi.machinery_list",),
        "headers": (
            "기계번호", "기계명", "기계종류", "처리량", "토출압력", "회전수",
            "임펠러반경", "양중하중", "양중높이", "명세",
            "동력", "재질", "방호·보호장치 종류",
            "취급물질", "설치공정", "관련 P&ID 번호", "비고",
        ),
        "example": (
            (
                "P-101", "원료 이송펌프", "원심펌프", "20 m3/h", "0.8 MPa", "1750 rpm",
                "", "", "", "", "7.5 kW", "SUS304",
                "커플링 가드·모터 과부하 보호", "톨루엔", "원료저장", "PID-101", "예시값"
            ),
        ),
    },
    {
        "sheet": "11_배관_개스킷",
        "scope": "PSM",
        "title": "배관 및 개스킷 명세",
        "targets": ("psm.psi.piping_gasket_specs",),
        "headers": (
            "배관번호·Class", "유체명", "배관재질", "호칭경", "설계압력", "설계온도",
            "개스킷 재질", "비파괴검사율", "후열처리여부", "관련 P&ID 번호", "비고",
        ),
        "example": (
            (
                "PCL-150", "톨루엔", "SUS304", "50A", "0.49 MPa", "80 ℃",
                "PTFE", "10%", "해당 없음", "PID-101", "예시값"
            ),
        ),
    },
    {
        "sheet": "09_PSM_조건부서식_적용여부",
        "scope": "PSM",
        "title": "PSM 조건부 별지서식 적용여부 확인",
        "targets": ("psm.psi.form_applicability",),
        "headers": ("서식번호", "서식명", "적용여부", "확인근거"),
        "example": (
            ("17-2", "이상발생시 인터록 작동조건 및 가동중지 범위", "적용", "공정 인터록 목록 및 P&ID 확인"),
            ("17-3", "소화설비 설치계획", "적용", "소화설비 배치도 확인"),
            ("17-4", "화재탐지경보설비 설치계획", "적용", "화재감지·경보설비 현황 확인"),
            ("17-5", "가스누출감지경보기 설치계획", "적용", "가스감지기 목록 확인"),
            ("18", "내화구조 명세", "해당 없음", "내화구조 적용대상 없음 확인"),
            ("19", "국소배기장치 개요", "해당 없음", "국소배기 적용대상 없음 확인"),
            ("20", "방폭전기/계장 기계·기구 선정기준", "적용", "폭발위험장소 구분도 확인"),
        ),
    },
    {
        "sheet": "12_PSM_인터록",
        "scope": "PSM",
        "title": "별지 제17호의2서식 입력자료 · 인터록 작동조건 및 가동중지 범위",
        "targets": ("psm.psi.interlock_conditions",),
        "headers": (
            "인터록번호", "대상설비번호", "설정값-온도(℃)", "설정값-압력(MPa)",
            "설정값-액위(m)", "설정값-기타", "감지기번호", "최종 작동설비번호",
            "가동중지범위", "점검주기", "비고",
        ),
        "example": (
            ("IL-101", "R-101", "120 ℃", "0.8 MPa", "해당 없음", "해당 없음",
             "TI-101/PI-101", "XV-101/P-101", "R-101 원료공급 및 가열 정지", "월 1회", "예시값"),
        ),
    },
    {
        "sheet": "13_PSM_소화설비",
        "scope": "PSM",
        "title": "별지 제17호의3서식 입력자료 · 소화설비 설치계획",
        "targets": ("psm.psi.fire_protection_table",),
        "headers": (
            "설치지역", "소화기", "자동확산소화기", "자동소화장치", "옥내소화전",
            "스프링클러", "물분무소화설비", "포소화설비", "CO2 소화설비",
            "할로겐화합물 소화설비", "청정소화약제 소화설비", "옥외소화전",
        ),
        "example": (
            ("원료저장동", "분말 6기", "해당 없음", "해당 없음", "2개소",
             "전면 설치", "해당 없음", "해당 없음", "해당 없음",
             "해당 없음", "해당 없음", "2개소"),
        ),
    },
    {
        "sheet": "14_PSM_화재탐지",
        "scope": "PSM",
        "title": "별지 제17호의4서식 입력자료 · 화재탐지경보설비 설치계획",
        "targets": ("psm.psi.fire_detection_table",),
        "headers": (
            "설치지역", "단독경보형 감지기", "비상경보설비", "시각경보기",
            "자동화재탐지설비", "비상방송설비", "자동화재속보설비",
            "통합감시시설", "누전경보기",
        ),
        "example": (
            ("반응동", "해당 없음", "1식", "2개소", "연기·열감지기 12개",
             "1식", "1식", "중앙제어실 연동", "1식"),
        ),
    },
    {
        "sheet": "15_PSM_내화구조",
        "scope": "PSM",
        "title": "별지 제18호서식 입력자료 · 내화구조 명세",
        "targets": ("psm.psi.fireproofing_table",),
        "headers": ("내화설비 또는 지역", "내화부위", "내화시험기준 및 시간", "비고"),
        "example": (
            ("R-101 지지철골", "주기둥 및 보", "2시간 내화성능", "내화피복 적용"),
        ),
    },
    {
        "sheet": "16_PSM_국소배기",
        "scope": "PSM",
        "title": "별지 제19호서식 입력자료 · 국소배기장치 개요",
        "targets": ("psm.psi.local_exhaust_table",),
        "headers": (
            "공정 또는 작업장명", "실내외 구분", "발생원", "유해물질 종류",
            "후드형식", "후드 제어풍속(m/s)", "덕트내 반송속도(m/s)",
            "배풍량(m3/min)", "전동기용량(kW)", "배기 및 처리순서", "방폭형식",
        ),
        "example": (
            ("혼합공정", "실내", "원료 투입구", "톨루엔", "포위식", "0.5",
             "12", "80", "7.5", "후드 → 덕트 → 활성탄 흡착기 → 배기구", "Ex d IIB T4"),
        ),
    },
    {
        "sheet": "17_PSM_방폭기기",
        "scope": "PSM",
        "title": "별지 제20호서식 입력자료 · 방폭전기/계장 기계·기구 선정기준",
        "targets": ("psm.psi.ex_equipment",),
        "headers": (
            "설치장소 또는 공정", "전기/계장 기계·기구명",
            "0종장소 선정기준(방폭형식)", "1종장소 선정기준(방폭형식)",
            "2종장소 선정기준(방폭형식)",
        ),
        "example": (
            ("원료저장", "모터·현장계기", "해당 없음", "Ex d IIB T4", "Ex e IIB T4"),
        ),
    },
    {
        "sheet": "18_PSM_위험성평가자",
        "scope": "PSM",
        "title": "별지 제21호서식 입력자료 · 위험성평가 참여 전문가 명단",
        "targets": ("psm.risk.team",),
        "headers": ("책임분야", "성명", "소속회사", "직책", "주요경력"),
        "example": (
            ("공정", "홍길동", "예시화학", "공정팀장", "공정설계 및 운전 15년"),
            ("안전", "김안전", "예시화학", "안전팀장", "공정안전관리 12년"),
        ),
    },
    {
        "sheet": "19_PSM_사고피해예측",
        "scope": "PSM",
        "title": "별지 제19호의2서식 입력자료 · 시나리오 및 피해예측 결과",
        "targets": ("psm.risk.consequence_table",),
        "headers": (
            "단위공장", "사고유형", "시나리오명", "대상 설비번호",
            "시나리오 구분", "풍속(m/s)", "대기안정도(A~F)", "대기온도(℃)", "습도(%)",
            "표면거칠기", "물질명", "물질의 상태", "설비명(또는 배관부위)",
            "운전압력(MPa)", "운전온도(℃)", "누출구의 크기(mm2)", "웅덩이 크기(m2)",
            "누출결과", "직접계산(kg/s or kg)", "웅덩이(kg/s)", "설비/배관(kg/s)",
            "화재-4 kW/m2", "화재-12.5 kW/m2", "화재-37.5 kW/m2",
            "폭발-7 kPa", "폭발-21 kPa", "폭발-70 kPa",
            "인화성-25% LEL", "인화성-LEL", "인화성-UEL",
            "독성-ERPG 1", "독성-ERPG 2", "독성-ERPG 3", "계산모델·결과 근거",
        ),
        "example": (
            (
                "제1공장", "독성 누출", "염소 최악 누출", "V-201",
                "최악의 사고 시나리오", 1.5, "F", 25, 50, "도시", "염소", "기체", "V-201",
                "0.7", "25", "25", "해당 없음", "연속누출", "0.25", "해당 없음", "0.25",
                "해당 없음", "해당 없음", "해당 없음", "해당 없음", "해당 없음", "해당 없음",
                "해당 없음", "해당 없음", "해당 없음", 650, 300, 180, "ALOHA/KORA 계산결과 파일 RISK-01",
            ),
            (
                "제1공장", "독성 누출", "염소 대안 누출 1", "V-201",
                "대안의 사고 시나리오", 3.0, "D", 25, 50, "도시", "염소", "기체", "V-201",
                "0.7", "25", "10", "해당 없음", "연속누출", "0.10", "해당 없음", "0.10",
                "해당 없음", "해당 없음", "해당 없음", "해당 없음", "해당 없음", "해당 없음",
                "해당 없음", "해당 없음", "해당 없음", 320, 150, 90, "ALOHA/KORA 계산결과 파일 RISK-02",
            ),
        ),
    },
    {
        "sheet": "20_배출물질_처리시설",
        "scope": "CAP",
        "title": "배출물질 처리시설 현황",
        "targets": ("cap.safety.waste_treatment",),
        "headers": (
            "시설번호", "시설명", "처리대상 물질", "처리방식", "처리용량", "연결 설비·배관", "최종 배출지점", "비고",
        ),
        "example": (
            ("SC-101", "유기용제 스크러버", "톨루엔 증기", "흡수", "1500 m3/h", "PSV-101", "대기배출구", "예시값"),
        ),
    },
    {
        "sheet": "21_확산방지설비",
        "scope": "CAP",
        "title": "확산방지설비 계산자료",
        "targets": ("cap.safety.dike_calculation",),
        "headers": (
            "적용여부", "대상 설비번호", "설비형태", "확산방지설비 종류",
            "필요용량(m3)", "필요용량 기준·근거",
            "내부 길이(m)", "내부 폭(m)", "유효높이(m)", "내부 차감용적(m3)",
            "직접확인 유효용량(m3)", "비고",
        ),
        "example": (
            ("예", "TK-101", "저장탱크", "방류벽", 22, "적용 시설기준 검토자료의 필요용량", 6, 5, 0.8, 1.2, "", "치수계산 예시"),
        ),
    },
    {
        "sheet": "22_사고시나리오_영향평가",
        "scope": "CAP",
        "title": "사고시나리오 영향평가 확정값",
        "targets": ("cap.offsite.scenario_impact_table",),
        "headers": (
            "사고시나리오명", "유해화학물질명", "대상 설비번호", "사고유형",
            "장외거리(m)", "거주민수", "근로자수",
            "갑종 보호대상 수", "을종 보호대상 수", "환경수용체 수",
            "사고원점 좌표", "KORA/GIS 근거",
        ),
        "example": (
            ("염소 독성누출-1", "염소", "V-201", "독성누출", 180, 25, 10, 1, 0, 1, "35.0,129.0", "KORA 결과파일 KORA-01"),
        ),
    },
    {
        "sheet": "23_사고시나리오_시설빈도",
        "scope": "CAP",
        "title": "사고시나리오별 개시사건 개수",
        "targets": ("cap.offsite.scenario_frequency",),
        "headers": (
            "사고시나리오명",
            "고압용기파열", "배관파열", "배관누출", "상압 탱크 파열 및 누출",
            "플랜지 등의 가스켓 파손", "펌프/컴프레서 누출",
            "안전밸브 오작동 및 조기개방", "냉각수 손실",
            "입/출하 시설 누출 사고", "외부화재",
            "개수 산정근거", "수동적 완화장치", "능동적 완화장치", "안전성확보설비 증빙",
        ),
        "example": (
            ("염소 독성누출-1", 2, 3, 5, 0, 8, 1, 1, 0, 0, 0, "PID-201 및 설비목록 검토", "방류벽", "가스감지기-자동차단밸브 연동", "GA-201/PID-201"),
        ),
    },
    {
        "sheet": "24_장외시나리오_확인",
        "scope": "CAP",
        "title": "장외 사고시나리오 존재 여부 확인",
        "targets": ("cap.offsite.risk_control",),
        "headers": ("장외 사고시나리오 없음 여부", "확인근거", "비고"),
        "example": (("아니오", "KORA 결과 KORA-01", "장외 사고시나리오 존재"),),
    },
    {
        "sheet": "25_총괄영향범위_요약",
        "scope": "CAP",
        "title": "총괄영향범위 GIS·KORA 확정요약",
        "targets": ("cap.offsite.overall_impact_summary",),
        "headers": (
            "총괄영향범위 산출방법", "총괄영향범위 결과 요약", "GIS/KORA 근거",
            "총괄영향범위 내 거주민수", "총괄영향범위 내 근로자수",
            "보호대상 없음 여부", "비고",
        ),
        "example": (
            ("KORA 사고시나리오 영향범위와 GIS 공간중첩", "총괄영향범위는 KORA/GIS 결과도면 기준", "KORA-ALL-01",
             65, 25, "아니오", "형상은 GIS 결과파일 참조"),
        ),
    },
    {
        "sheet": "26_총괄영향범위_보호대상",
        "scope": "CAP",
        "title": "총괄영향범위 내 보호대상 명세",
        "targets": ("cap.offsite.population_and_protected_targets",),
        "headers": (
            "보호대상 명칭", "보호대상 구분", "세부유형", "주소·위치", "좌표",
            "사업장 경계와 거리(m)", "인원수", "GIS 근거", "비고",
        ),
        "example": (
            ("○○초등학교", "갑종", "교육·연구시설", "○○시 ○○로 10", "35.0,129.0", 420, 350, "GIS-PT-01", "예시"),
        ),
    },
    {
        "sheet": "27_사업장주변_500m_보호대상",
        "scope": "CAP",
        "title": "사업장 경계 500m 내 보호대상 현황",
        "targets": ("cap.site.surrounding_environment",),
        "headers": (
            "보호대상 없음 여부", "보호대상 명칭", "보호대상 구분", "세부유형",
            "주소·위치", "좌표", "사업장 경계와 거리(m)", "GIS/현장 근거", "비고",
        ),
        "example": (
            ("아니오", "○○초등학교", "갑종", "교육·연구시설",
             "○○시 ○○로 10", "35.0,129.0", 420, "GIS-SITE-01", "500m 내 보호대상 예시"),
        ),
    },
    {
        "sheet": "28_대표물질_유해성정보",
        "scope": "CAP",
        "title": "별지 제7호 대표물질 유해성 정보",
        "targets": ("cap.chemical.hazard_information",),
        "headers": (
            "물질명", "CAS 번호", "인체유해성", "물리적 위험성", "환경유해성",
            "출처", "선정 사유", "SDS 파일명", "SDS 개정일", "비고",
        ),
        "example": (
            (
                "염소", "7782-50-5",
                "흡입 시 급성 독성 및 호흡기 자극 우려",
                "산화성·가압가스 관련 위험",
                "수생생물에 매우 유독",
                "회사 제품 SDS 제2·11·12항",
                "사고시나리오 대상물질이며 독성영향을 대표하므로 선정",
                "chlorine_company_SDS.pdf", "2026-05-10", "회사 확인자료 예시"
            ),
        ),
    },
    {
        "sheet": "29_변경내역_관리대장",
        "scope": "CAP",
        "title": "별지 제2호 변경내역 관리대장 입력자료",
        "targets": ("cap.prevention.change_log",),
        "headers": (
            "일자", "변경항목", "변경의 종류",
            "변경 내용(변경전 → 변경후)", "후속조치", "담당자",
        ),
        "example": (
            (
                "2026-09-18",
                "장치·설비 목록 및 명세",
                "설비변경",
                "TK-101 저장탱크 → TK-101A 저장탱크",
                "관련 도면·설비명세·계획서 변경사항 갱신",
                "환경안전팀 홍길동",
            ),
        ),
    },
)

ATTACHMENT_KINDS = {
    "DOCUMENT_SET", "DRAWING", "DRAWING_SET", "DRAWING_AND_DATA", "DRAWING_AND_TABLE",
    "ANALYSIS_DOCUMENT", "CALCULATION_AND_DRAWING", "CALCULATION_AND_MODEL",
    "TABLE_AND_DRAWING", "PLAN_CALCULATION_AND_DRAWING", "PLAN_SPEC_CALCULATION_AND_DRAWING",
    "PLAN_TABLE_AND_DRAWING", "PLAN_SPEC_AND_DRAWING",
}

EXAMPLE_VALUES: dict[str, Any] = {
    "cap.business.representative": "홍길동",
    "cap.business.registration_no": "123-45-67890",
    "cap.business.contact": "053-000-0000",
    "cap.business.unit_plant_name": "제1공장",
    "cap.business.industrial_complex": "○○국가산업단지",
    "cap.business.submission_type": "신규제출",
    "cap.business.submission_reason": "최초",
    "cap.business.joint_emergency_plan": "단독제출",
    "cap.business.other_system_review": "미해당",
    "cap.business.residents_in_overall_range": "없음",
    "cap.business.recent_accident": "없음",
    "cap.business.writer_name": "김담당",
    "cap.business.writer_department": "환경안전팀",
    "cap.business.writer_contact": "053-000-0001",
    "cap.business.writer_email": "safety@example.com",
    "psm.business.overview": "본 사업장은 원료 저장, 혼합 및 제품 출하 공정으로 구성되며 주요 공정설비는 저장탱크, 혼합기 및 이송펌프이다.",
    "process.description": "원료는 저장탱크에서 이송펌프로 혼합공정에 공급되고, 정해진 운전조건에서 혼합 후 제품저장설비로 이송된다.",
    "psm.risk.purpose": "공정 내 잠재 유해·위험요인을 체계적으로 확인하고 필요한 개선대책을 도출하기 위함.",
    "psm.risk.characteristics": "인화성 액체 취급과 이송배관 누출에 따른 화재·폭발 위험이 주요 위험특성이다.",
    "psm.risk.mitigation": "누출감지, 긴급차단, 방유제 및 비상대응절차를 운영한다.",
    "psm.operation.sop": "정상운전, 운전개시, 운전정지 및 비상정지 절차를 사내 안전운전지침서에 따라 수행한다.",
    "psm.operation.maintenance": "설비별 점검주기와 정비기준을 정하고 점검·보수 이력을 관리한다.",
    "psm.operation.work_permit": "화기작업, 밀폐공간작업 등 위험작업은 안전작업허가 절차에 따라 승인 후 수행한다.",
    "psm.operation.contractor": "도급작업 전 작업위험과 안전수칙을 공유하고 작업 중 이행상태를 확인한다.",
    "psm.operation.training": "신규·정기·작업변경 시 필요한 안전보건교육을 실시하고 기록을 보관한다.",
    "psm.operation.prestartup": "신설·변경 설비의 가동 전 설계·시공·안전장치·절차 준비상태를 확인한다.",
    "psm.operation.moc": "설비, 원료, 운전조건 및 절차 변경 시 변경요소 관리절차에 따라 사전 검토한다.",
    "psm.operation.audit": "정기적으로 공정안전관리 이행상태를 자체 점검하고 개선조치를 관리한다.",
    "psm.operation.incident_investigation": "공정사고 발생 시 원인조사 후 재발방지대책을 수립하고 이행상태를 확인한다.",
    "psm.emergency.contacts": "비상상황 발생 시 중앙제어실을 통해 비상연락망으로 상황을 전파한다.",
    "psm.emergency.roles_procedures": "비상대응조직별 지휘, 신고, 초기대응, 대피 및 복구 임무를 정하여 운영한다.",
    "psm.emergency.training": "정기 비상조치 교육과 훈련을 실시하고 결과를 기록한다.",
    "cap.prevention.safety_policy": "화학사고 예방과 인명·환경피해 최소화를 안전관리의 기본방침으로 한다.",
    "cap.prevention.training_plan": "취급작업자와 비상대응조직을 대상으로 정기 교육·훈련을 실시한다.",
    "cap.prevention.change_management_plan": "유해화학물질, 취급시설, 운전조건 또는 절차 변경 시 사전 검토 후 변경사항을 관리한다.",
    "cap.internal.shutdown_authority": "공정반장 또는 중앙제어실 책임자",
    "cap.internal.shutdown_procedure": "이상상황 확인 후 원료공급 차단, 관련 설비 정지, 상황전파 순으로 조치한다.",
    "cap.internal.communication_system": "비상방송, 유선전화 및 비상연락망을 이용하여 내부에 상황을 전달한다.",
    "cap.external.communication_plan": "주변 지역사회와 관계기관에 사고위험정보와 비상시 행동요령을 정기적으로 안내한다.",
}


@dataclass(frozen=True)
class IntegratedImportResult:
    updated_fields: int
    table_fields: int
    attachment_declarations: int
    warnings: tuple[str, ...]


def _scope_matches(project: Stage2Project, scope: str) -> bool:
    if scope == "COMMON":
        return project.psm_in_scope or project.cap_in_scope
    if scope == "PSM":
        return project.psm_in_scope
    if scope == "CAP":
        return project.cap_in_scope
    return False


def _safe_sheet_name(name: str) -> str:
    text = re.sub(r"[\\/*?:\[\]]", "_", str(name))
    return text[:31]


def _style_title(ws, title: str, note: str = "") -> None:
    ws["A1"] = title
    ws["A1"].font = Font(size=14, bold=True, color="FFFFFF")
    ws["A1"].fill = PatternFill("solid", fgColor=HEADER_FILL)
    ws["A1"].alignment = Alignment(vertical="center")
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=6)
    if note:
        ws["A2"] = note
        ws["A2"].fill = PatternFill("solid", fgColor=NOTE_FILL)
        ws["A2"].alignment = Alignment(wrap_text=True, vertical="top")
        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=6)


def _style_header(row) -> None:
    for cell in row:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=HEADER_FILL)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _autosize(ws, max_width: int = 34) -> None:
    for column_cells in ws.columns:
        letter = get_column_letter(column_cells[0].column)
        width = 10
        for cell in column_cells:
            value = "" if cell.value is None else str(cell.value)
            width = max(width, min(len(value) + 2, max_width))
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        ws.column_dimensions[letter].width = width


def _normalized_mapping(row: Mapping[str, Any]) -> dict[str, Any]:
    def key(value: object) -> str:
        return re.sub(r"[^0-9a-z가-힣]", "", str(value or "").lower())
    return {key(k): v for k, v in row.items()}


def _pick(row: Mapping[str, Any], *aliases: str) -> Any:
    normalized = _normalized_mapping(row)
    for alias in aliases:
        token = re.sub(r"[^0-9a-z가-힣]", "", alias.lower())
        if token in normalized and normalized[token] not in (None, ""):
            return normalized[token]
    return ""


def _prefill_chemicals(project: Stage2Project) -> list[list[Any]]:
    # Preserve the richer Stage-2 table on subsequent downloads while keeping
    # Stage-1 inventory as the immutable first-pass source. In PSM-only scope,
    # never depend on a CAP-only enrichment field.
    record = None
    preferred_keys = []
    if project.psm_in_scope:
        preferred_keys.append("psm.psi.chemical_details")
    if project.cap_in_scope:
        preferred_keys.append("cap.chemical.details")
    preferred_keys.append("inventory.chemicals")
    for key in preferred_keys:
        candidate = project.get_field(key)
        if candidate is not None and isinstance(candidate.value, list) and candidate.value:
            record = candidate
            break
    rows = record.value if record and isinstance(record.value, list) else []
    out: list[list[Any]] = []
    for row in rows:
        if not isinstance(row, Mapping):
            continue
        out.append([
            _pick(row, "물질명", "화학물질명", "제품명", "substance", "name"),
            _pick(row, "CAS 번호", "CAS", "CAS No", "CAS번호"),
            _pick(row, "분자식", "molecular formula"),
            _pick(row, "함량(%)", "함량", "농도", "content", "purity"),
            _pick(row, "물리적 상태", "상태", "state", "phase"),
            _pick(row, "최대보유량", "최대저장량", "보유량", "quantity", "holding"),
            _pick(row, "단위", "unit"),
            _pick(row, "일일사용량", "취급량", "사용량", "daily use"),
            _pick(row, "공정", "사용공정", "저장공정", "사용·저장 공정", "process"),
            _pick(row, "용도", "주요 용도", "usage"),
            _pick(row, "비중", "밀도/비중"),
            _pick(row, "폭발한계 하한", "폭발하한", "LEL"),
            _pick(row, "폭발한계 상한", "폭발상한", "UEL"),
            _pick(row, "노출기준", "TWA", "허용농도값"),
            _pick(row, "독성치", "toxicity"),
            _pick(row, "인화점", "flash point"),
            _pick(row, "발화점", "ignition point"),
            _pick(row, "이상반응 유무", "이상반응"),
            _pick(row, "독성구분 항목", "독성구분-항목"),
            _pick(row, "독성구분", "독성구분-구분"),
            _pick(row, "위험노출수준", "ERPG", "AEGL", "PAC", "IDLH"),
            _pick(row, "허용농도값", "TWA", "노출기준"),
            _pick(row, "증기압", "증기압(20℃, mmHg)"),
            _pick(row, "부식성", "부식성(유, 무)"),
            _pick(row, "SDS 파일명", "MSDS 파일명"),
            _pick(row, "SDS 개정일", "MSDS 개정일", "SDS 작성·개정일"),
            _pick(row, "비고", "note"),
        ])
    return out


def _prefill_detectors(project: Stage2Project) -> list[list[Any]]:
    rows: list[Mapping[str, Any]] = []
    for key in ("psm.psi.gas_detection_table", "cap.safety.gas_detection", "psm.psi.gas_detection"):
        record = project.get_field(key)
        if record and isinstance(record.value, list):
            rows = [row for row in record.value if isinstance(row, Mapping)]
            if rows:
                break

    out: list[list[Any]] = []
    for row in rows:
        legacy = str(_pick(row, "감지방식") or "").strip()
        legacy_norm = re.sub(r"[^0-9a-z가-힣]", "", legacy.lower())
        installation = _pick(row, "설치형태", "설치 형식")
        measurement = _pick(row, "측정방식", "측정원리", "센서방식")
        if not installation and legacy_norm in {"고정식", "휴대식", "고정식휴대식"}:
            installation = legacy
        if not measurement and legacy and legacy_norm not in {"고정식", "휴대식", "고정식휴대식"}:
            measurement = legacy

        out.append([
            _pick(row, "감지기 번호", "감지기번호", "구분기호"),
            installation,
            _pick(row, "설치위치", "설치장소"),
            _pick(row, "검출대상 물질", "감지대상", "검출대상"),
            _pick(row, "작동시간", "응답시간"),
            measurement,
            _pick(row, "경보 설정값", "경보설정값"),
            _pick(row, "경보 위치", "경보기 설치장소"),
            _pick(row, "연동여부", "인터록 연동여부"),
            _pick(row, "연동 설비·조치", "연동설비", "경보시 조치내용"),
            _pick(row, "정밀도", "정확도"),
            _pick(row, "유지관리", "점검주기", "교정주기"),
            _pick(row, "비상전원 여부"),
            _pick(row, "관련 도면번호", "도면번호"),
            _pick(row, "비고"),
        ])
    return out


def _prefill_facilities(project: Stage2Project) -> list[list[Any]]:
    record = None
    preferred_keys = []
    if project.psm_in_scope:
        preferred_keys.append("psm.psi.equipment_specs")
    if project.cap_in_scope:
        preferred_keys.append("cap.facility.equipment_specs")
    preferred_keys.append("inventory.facilities")
    for key in preferred_keys:
        candidate = project.get_field(key)
        if candidate is not None and isinstance(candidate.value, list) and candidate.value:
            record = candidate
            break
    rows = record.value if record and isinstance(record.value, list) else []
    out: list[list[Any]] = []
    for row in rows:
        if not isinstance(row, Mapping):
            continue
        out.append([
            _pick(row, "설비번호", "시설번호", "장치번호", "tag", "equipment id"),
            _pick(row, "설비명", "시설명", "장치명", "equipment"),
            _pick(row, "설비종류", "시설종류", "type"),
            _pick(row, "단위공장·공정", "단위공장", "공정", "process", "unit"),
            _pick(row, "취급물질", "물질명", "chemical"),
            _pick(row, "용량", "capacity"),
            _pick(row, "용량단위", "단위", "unit"),
            _pick(row, "설계압력", "design pressure"),
            _pick(row, "설계온도", "design temperature"),
            _pick(row, "운전압력", "operating pressure"),
            _pick(row, "운전온도", "operating temperature"),
            _pick(row, "재질", "본체재질", "material"),
            _pick(row, "부속품재질", "부속품"),
            _pick(row, "개스킷재질", "개스킷 재질"),
            _pick(row, "용접효율"),
            _pick(row, "계산두께"),
            _pick(row, "부식여유"),
            _pick(row, "사용두께"),
            _pick(row, "후열처리 여부", "후열처리여부"),
            _pick(row, "비파괴검사율", "비파괴율검사"),
            _pick(row, "최대보유량", "최대보유량(kg)", "holding"),
            _pick(row, "P&ID 번호", "P&ID", "PID"),
            _pick(row, "비고", "note"),
            _pick(row, "최대 연결구 크기(mm)", "연결구 크기(mm)", "호칭경", "connection size"),
        ])
    return out


PSM_CONDITIONAL_FORMS: tuple[tuple[str, str], ...] = (
    ("17-2", "이상발생시 인터록 작동조건 및 가동중지 범위"),
    ("17-3", "소화설비 설치계획"),
    ("17-4", "화재탐지경보설비 설치계획"),
    ("17-5", "가스누출감지경보기 설치계획"),
    ("18", "내화구조 명세"),
    ("19", "국소배기장치 개요"),
    ("20", "방폭전기/계장 기계·기구 선정기준"),
    ("19-2", "시나리오 및 피해예측 결과"),
)


def _prefill_psm_form_applicability(project: Stage2Project) -> list[list[Any]]:
    record = project.get_field("psm.psi.form_applicability")
    saved: dict[str, Mapping[str, Any]] = {}
    if record is not None and isinstance(record.value, list):
        for row in record.value:
            if not isinstance(row, Mapping):
                continue
            form_no = str(_pick(row, "서식번호", "form_no") or "").strip()
            if form_no:
                saved[form_no] = row

    rows: list[list[Any]] = []
    for form_no, form_name in PSM_CONDITIONAL_FORMS:
        row = saved.get(form_no, {})
        rows.append([
            form_no,
            form_name,
            _pick(row, "적용여부", "적용 여부", "applicability"),
            _pick(row, "확인근거", "근거", "basis"),
        ])
    return rows


def _prefill_psm_form12(project: Stage2Project) -> list[list[Any]]:
    saved = project.get_field("psm.business.form12_details")
    if saved is not None and isinstance(saved.value, list) and saved.value:
        row = next((item for item in saved.value if isinstance(item, Mapping)), None)
        if row is not None:
            headers = next(
                spec["headers"] for spec in TABLE_SPECS
                if spec["sheet"] == "08_PSM_별지12_사업개요"
            )
            return [[_pick(row, header) for header in headers]]

    chemicals = project.get_field("psm.psi.chemical_details") or project.get_field("inventory.chemicals")
    raw_materials: list[str] = []
    if chemicals is not None and isinstance(chemicals.value, list):
        for item in chemicals.value[:8]:
            if isinstance(item, Mapping):
                name = _pick(item, "물질명", "화학물질명", "제품명")
                if name:
                    raw_materials.append(str(name))

    def scalar(*keys: str) -> Any:
        for key in keys:
            record = project.get_field(key)
            if record is not None and record.value not in (None, "", [], {}):
                if isinstance(record.value, (str, int, float, bool)):
                    return record.value
        return ""

    writer = project.get_field("psm.business.writer_info")
    writer_name = ""
    writer_qualification = ""
    if writer is not None and isinstance(writer.value, Mapping):
        writer_name = _pick(writer.value, "작성자", "성명", "이름", "name")
        writer_qualification = _pick(writer.value, "작성자 자격", "자격", "qualification")
    elif writer is not None and isinstance(writer.value, str):
        parts = [part.strip() for part in writer.value.split("/", 1)]
        writer_name = parts[0] if parts else ""
        writer_qualification = parts[1] if len(parts) > 1 else ""

    return [[
        project.company_name,
        scalar("psm.business.project_type"),
        scalar("business.registration_no", "cap.business.registration_no"),
        scalar("business.representative", "cap.business.representative"),
        scalar("psm.business.target_facility"),
        scalar("business.ksic"),
        scalar("business.employee_count"),
        scalar("business.electric_contract_capacity"),
        writer_name,
        writer_qualification,
        ", ".join(raw_materials),
        scalar("business.main_products"),
        scalar("psm.business.overview"),
        scalar("business.address"),
        scalar("business.phone", "psm.business.phone"),
        scalar("business.fax", "psm.business.fax"),
        scalar("psm.business.site_area", "business.site_area"),
        scalar("psm.business.main_building", "business.main_building", "psm.business.site_building"),
        scalar("psm.business.total_period", "psm.business.schedule"),
        scalar("psm.business.start_date"),
        scalar("psm.business.commissioning_period"),
    ]]


def _table_rows_for(project: Stage2Project, sheet: str, example: bool, spec: Mapping[str, Any]) -> list[list[Any]]:
    if example:
        return [list(row) for row in spec.get("example", ())]
    if sheet == "08_PSM_별지12_사업개요":
        return _prefill_psm_form12(project)
    if sheet == "09_PSM_조건부서식_적용여부":
        return _prefill_psm_form_applicability(project)
    if sheet == "02_화학물질정보":
        return _prefill_chemicals(project)
    if sheet == "03_설비정보":
        return _prefill_facilities(project)
    if sheet == "05_가스누출감지_경보장치":
        return _prefill_detectors(project)

    # Generic structured-table re-download support. This preserves previously
    # confirmed rows such as Form 7 hazard information, GIS target tables and
    # scenario tables instead of returning a blank sheet on the next download.
    headers = list(spec.get("headers", ()))
    for key in spec.get("targets", ()):
        record = project.get_field(str(key))
        if record is None or not isinstance(record.value, list):
            continue
        source_rows = [row for row in record.value if isinstance(row, Mapping)]
        if not source_rows:
            continue
        return [[_pick(row, header) for header in headers] for row in source_rows]
    return []


def _write_table_sheet(wb: Workbook, project: Stage2Project, spec: Mapping[str, Any], *, example: bool, meta_rows: list[list[Any]]) -> None:
    sheet_name = _safe_sheet_name(str(spec["sheet"]))
    ws = wb.create_sheet(sheet_name)
    _style_title(
        ws,
        str(spec["title"]),
        "노란색 영역에 사업장 정보를 입력하세요. 공정안전보고서와 화학사고예방관리계획서에 공통으로 필요한 정보는 한 번만 작성합니다."
        if not example else
        "작성예시입니다. 아래 값은 실제 사업장 정보가 아니며 입력 방법을 보여주기 위한 예시입니다.",
    )
    header_row = 4
    headers = list(spec["headers"])
    for col, value in enumerate(headers, 1):
        ws.cell(header_row, col, value)
    _style_header(ws[header_row])

    rows = _table_rows_for(project, sheet_name, example, spec)
    first_data = header_row + 1
    for r_idx, row in enumerate(rows, first_data):
        for c_idx, value in enumerate(row, 1):
            ws.cell(r_idx, c_idx, value)
            ws.cell(r_idx, c_idx).fill = PatternFill("solid", fgColor=EXAMPLE_FILL if example else INPUT_FILL)
    reserve_to = max(first_data + 24, first_data + len(rows))
    for row in ws.iter_rows(min_row=first_data, max_row=reserve_to, min_col=1, max_col=len(headers)):
        for cell in row:
            if cell.fill.fill_type is None:
                cell.fill = PatternFill("solid", fgColor=INPUT_FILL if not example else EXAMPLE_FILL)
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    ws.freeze_panes = f"A{first_data}"
    _autosize(ws)
    targets = [key for key in spec.get("targets", ()) if _target_in_scope(project, key)]
    meta_rows.append(["TABLE", sheet_name, "|".join(targets), header_row, "", ""])


def _target_in_scope(project: Stage2Project, field_key: str) -> bool:
    if field_key.startswith("psm."):
        return project.psm_in_scope
    if field_key.startswith("cap."):
        return project.cap_in_scope
    return True


def _question_text(label: str, request_text: str) -> str:
    request = str(request_text or "").strip()
    if request and request != label:
        return request
    return f"{label}을(를) 사업장 실제 기준으로 작성해 주세요."


def _example_value(field_key: str, label: str) -> str:
    if field_key in EXAMPLE_VALUES:
        return str(EXAMPLE_VALUES[field_key])
    if "연락" in label:
        return "예: 담당부서, 담당자, 연락처와 연락 순서를 실제 기준으로 입력"
    if "계획" in label or "절차" in label or "방침" in label:
        return f"예: {label}의 실제 운영방법, 담당자, 주기 및 기록방법을 간단히 입력"
    if "현황" in label or "정보" in label:
        return f"예: {label}에 해당하는 사업장 실제 현황을 입력"
    return f"예: {label}에 해당하는 사업장 실제 내용을 입력"


def _handled_fields(project: Stage2Project) -> set[str]:
    handled = set(PROTECTED_STAGE1_FIELDS)
    for spec in TABLE_SPECS:
        if _scope_matches(project, str(spec["scope"])):
            handled.update(key for key in spec["targets"] if _target_in_scope(project, key))
    handled.update({"process.description"})
    return handled


def _build_question_groups(project: Stage2Project) -> dict[str, list[tuple[str, str, str, str]]]:
    groups: dict[str, list[tuple[str, str, str, str]]] = {}
    handled = _handled_fields(project)
    seen: set[str] = set()
    for spec in selected_requirement_specs(project):
        if spec.input_kind in ATTACHMENT_KINDS or spec.input_kind in {"DOCUMENT_SET", "DRAWING"}:
            continue
        for key in spec.field_keys:
            if key in handled or key in seen or key in PROTECTED_STAGE1_FIELDS:
                continue
            if not _target_in_scope(project, key):
                continue
            seen.add(key)
            if spec.system == "PSM":
                group = f"공정안전보고서_{spec.section}"
            elif spec.system == "CAP":
                section = re.sub(r"^3\.\d+\s*", "", spec.section).strip()
                group = f"화학사고예방관리계획서_{section or '작성정보'}"
            else:
                group = "공통_추가정보"
            label = field_label(key)
            groups.setdefault(group, []).append((key, label, _question_text(label, spec.request_text), spec.legal_basis))
    return groups


def _write_question_sheet(wb: Workbook, sheet_name: str, rows: Iterable[tuple[str, str, str, str]], *, project: Stage2Project, example: bool, meta_rows: list[list[Any]]) -> None:
    ws = wb.create_sheet(_safe_sheet_name(sheet_name))
    _style_title(
        ws,
        sheet_name.replace("_", " "),
        "질문을 읽고 회사 작성값에 실제 사업장 기준으로 입력하세요. 모르면 임의로 작성하지 말고 빈칸으로 두십시오."
        if not example else
        "작성예시입니다. 표현방식을 참고하되 실제 사업장 사실과 다른 내용을 복사하지 마십시오.",
    )
    headers = ["확인할 내용", "회사 작성값", "작성 예", "작성방법"]
    for col, value in enumerate(headers, 1):
        ws.cell(4, col, value)
    _style_header(ws[4])
    current = 5
    for field_key, label, question, basis in rows:
        record = project.get_field(field_key)
        current_value = ""
        if not example and record and record.value not in (None, "", [], {}):
            current_value = record.value if isinstance(record.value, (str, int, float, bool)) else ""
        ws.cell(current, 1, question)
        ws.cell(current, 2, _example_value(field_key, label) if example else current_value)
        ws.cell(current, 3, _example_value(field_key, label))
        ws.cell(current, 4, "확인 가능한 사내 문서나 담당자 확인을 바탕으로 작성. 모르는 내용은 빈칸 유지.")
        ws.cell(current, 2).fill = PatternFill("solid", fgColor=EXAMPLE_FILL if example else INPUT_FILL)
        for col in range(1, 5):
            ws.cell(current, col).alignment = Alignment(vertical="top", wrap_text=True)
        meta_rows.append(["FORM", ws.title, field_key, current, 2, basis])
        current += 1
    ws.freeze_panes = "A5"
    _autosize(ws, max_width=42)
    ws.column_dimensions["A"].width = 44
    ws.column_dimensions["B"].width = 52
    ws.column_dimensions["C"].width = 52
    ws.column_dimensions["D"].width = 36


def _attachment_specs(project: Stage2Project) -> list[tuple[str, str, str]]:
    rows: list[tuple[str, str, str]] = []
    seen: set[str] = set()
    for spec in selected_requirement_specs(project):
        if requirement_explicitly_not_applicable(project, spec.key):
            continue
        if spec.input_kind not in ATTACHMENT_KINDS and spec.input_kind not in {"DOCUMENT_SET", "DRAWING"}:
            continue
        for key in spec.field_keys:
            if key in seen or key in PROTECTED_STAGE1_FIELDS:
                continue
            if key == "cap.site.surrounding_environment":
                # Form 8 surroundings are captured as a structured GIS/field
                # table. Site/layout drawings remain separate attachment keys.
                continue
            if not _target_in_scope(project, key):
                continue
            seen.add(key)
            rows.append((key, field_label(key), spec.request_text or spec.description))
    return rows


def _write_attachment_sheet(wb: Workbook, project: Stage2Project, *, example: bool, meta_rows: list[list[Any]]) -> None:
    ws = wb.create_sheet("07_도면_첨부자료목록")
    _style_title(
        ws,
        "도면·첨부자료 목록",
        "PFD, P&ID, 배치도, 물질안전보건자료(MSDS) 등 파일 자체는 이 Excel에 넣지 말고 파일명과 도면번호만 적은 뒤 프로그램에서 별도로 업로드하세요.",
    )
    headers = ["자료종류", "파일명", "도면·문서번호", "개정번호·일자", "비고", "확인사항"]
    for col, value in enumerate(headers, 1):
        ws.cell(4, col, value)
    _style_header(ws[4])
    current = 5
    for field_key, label, request in _attachment_specs(project):
        record = project.get_field(field_key)
        declared = ""
        if not example and record and isinstance(record.value, str):
            declared = record.value
        ws.cell(current, 1, label)
        ws.cell(current, 2, f"{re.sub(r'[^0-9A-Za-z가-힣]+', '_', label)}.pdf" if example else declared)
        ws.cell(current, 3, "PID-101" if example and "P&ID" in label else "")
        ws.cell(current, 4, "Rev.3 / 2026-08-01" if example else "")
        ws.cell(current, 5, "최신본" if example else "")
        ws.cell(current, 6, request)
        for col in (2, 3, 4, 5):
            ws.cell(current, col).fill = PatternFill("solid", fgColor=EXAMPLE_FILL if example else INPUT_FILL)
        for col in range(1, 7):
            ws.cell(current, col).alignment = Alignment(vertical="top", wrap_text=True)
        meta_rows.append(["ATTACHMENT", ws.title, field_key, current, 2, ""])
        current += 1
    ws.freeze_panes = "A5"
    _autosize(ws, max_width=40)


def _write_guide(wb: Workbook, project: Stage2Project, *, example: bool) -> None:
    ws = wb.active
    ws.title = "00_작성안내"
    title = "통합 작성자료 작성예시" if example else "통합 작성자료"
    _style_title(ws, title, "공정안전보고서와 화학사고예방관리계획서의 중복정보를 한 번만 입력하도록 구성한 프로그램용 작성자료입니다.")
    scope = []
    if project.psm_in_scope:
        scope.append(PSM_FULL)
    if project.cap_in_scope:
        scope.append(CAP_FULL)
    guide_rows = [
        ("작성범위", ", ".join(scope)),
        ("문서 성격", "프로그램 통합 작성자료(법정 별지서식 자체가 아님)"),
        ("작성 원칙", "노란색 칸에 실제 사업장 사실만 입력하고, 모르는 내용은 추측하지 말고 빈칸으로 둡니다."),
        ("중복 입력", "공통 화학물질·설비·안전밸브·감지기 정보는 한 번만 작성하며 선택한 두 보고서에 재사용됩니다."),
        ("도면·이미지", "PFD, P&ID, 배치도, 물질안전보건자료(MSDS) 등 파일은 07_도면_첨부자료목록에 파일명만 적고 프로그램에서 별도 업로드합니다."),
        ("Stage 1 승계", "판정진단에서 확인된 회사명·주소·화학물질·시설자료는 가능한 범위에서 미리 채워집니다. 판정에 사용된 사실을 바꿔야 하면 먼저 판정진단을 다시 수행합니다."),
        ("작성예시", "작성예시 파일은 입력 방법 참고용이며 프로그램에 제출할 수 없습니다."),
        ("법적 근거", "평소 화면에서는 최소화하고, 누락·확인 필요 항목이 발생했을 때 프로그램의 ‘왜 필요한가?’에서 관련 법령·서식·매뉴얼 근거를 확인합니다."),
    ]
    ws["A4"] = "구분"
    ws["B4"] = "내용"
    _style_header(ws[4])
    for idx, (kind, text) in enumerate(guide_rows, 5):
        ws.cell(idx, 1, kind)
        ws.cell(idx, 2, text)
        ws.cell(idx, 2).alignment = Alignment(wrap_text=True, vertical="top")
    ws.column_dimensions["A"].width = 18
    ws.column_dimensions["B"].width = 100


def _write_business_sheet(wb: Workbook, project: Stage2Project, *, example: bool, meta_rows: list[list[Any]]) -> None:
    ws = wb.create_sheet("01_사업장정보")
    _style_title(ws, "사업장 정보", "판정진단에서 확인된 정보는 미리 채워집니다. 내용이 달라졌다면 이 파일에서 고치기 전에 판정진단을 다시 수행하세요.")
    ws.append([])
    headers = ["확인할 내용", "회사 작성값", "작성 예", "작성방법"]
    for col, value in enumerate(headers, 1):
        ws.cell(4, col, value)
    _style_header(ws[4])
    base = [
        ("business.company_name", "회사명", "㈜가상화학", "Stage 1에서 승계된 회사명을 확인"),
        ("business.address", "사업장 소재지", "대구광역시 ○○구 ○○로 00", "Stage 1에서 승계된 사업장 소재지를 확인"),
    ]
    current = 5
    for key, label, example_value, guidance in base:
        record = project.get_field(key)
        value = example_value if example else (record.value if record else "")
        ws.cell(current, 1, label)
        ws.cell(current, 2, value)
        ws.cell(current, 3, example_value)
        ws.cell(current, 4, guidance)
        ws.cell(current, 2).fill = PatternFill("solid", fgColor=EXAMPLE_FILL if example else INPUT_FILL)
        meta_rows.append(["PROTECTED", ws.title, key, current, 2, "Stage 1 판정자료"])
        current += 1

    # CAP의 사업장 일반정보 중 Stage 1에서 받지 않은 값은 여기에서 한 번만 입력한다.
    if project.cap_in_scope:
        labels = cap_field_labels()
        for key in (
            "cap.business.representative", "cap.business.registration_no", "cap.business.contact",
            "cap.business.unit_plant_name", "cap.business.industrial_complex",
            "cap.business.submission_type", "cap.business.submission_reason",
            "cap.business.writing_level",
            "cap.business.joint_emergency_plan", "cap.business.other_system_review",
            "cap.business.residents_in_overall_range", "cap.business.recent_accident",
            "cap.business.writer_name", "cap.business.writer_department",
            "cap.business.writer_contact", "cap.business.writer_email",
        ):
            label = labels.get(key, field_label(key))
            record = project.get_field(key)
            value = _example_value(key, label) if example else (record.value if record and isinstance(record.value, (str, int, float)) else "")
            ws.cell(current, 1, label)
            ws.cell(current, 2, value)
            ws.cell(current, 3, _example_value(key, label))
            ws.cell(current, 4, "사업자등록증, 사내 담당자 확인 등 실제 근거를 바탕으로 입력")
            ws.cell(current, 2).fill = PatternFill("solid", fgColor=EXAMPLE_FILL if example else INPUT_FILL)
            meta_rows.append(["FORM", ws.title, key, current, 2, ""])
            current += 1
    ws.freeze_panes = "A5"
    _autosize(ws, max_width=44)
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 42
    ws.column_dimensions["C"].width = 42
    ws.column_dimensions["D"].width = 44


def _write_process_sheet(wb: Workbook, project: Stage2Project, *, example: bool, meta_rows: list[list[Any]]) -> None:
    ws = wb.create_sheet("06_공정정보")
    _style_title(ws, "공정 정보", "공정 흐름과 정상운전 정보를 회사 담당자가 이해하는 수준으로 적어 주세요. 프로그램이 법정 문서 문장으로 다듬을 때 이 사실정보만 사용합니다.")
    headers = ["확인할 내용", "회사 작성값", "작성 예", "작성방법"]
    for col, value in enumerate(headers, 1):
        ws.cell(4, col, value)
    _style_header(ws[4])
    key = "process.description"
    record = project.get_field(key)
    value = _example_value(key, "공정개요") if example else (record.value if record and isinstance(record.value, str) else "")
    ws.cell(5, 1, "원료 투입부터 제품·부산물 배출까지 주요 공정 흐름과 정상운전 조건을 설명해 주세요.")
    ws.cell(5, 2, value)
    ws.cell(5, 3, _example_value(key, "공정개요"))
    ws.cell(5, 4, "공정단계, 주요 설비, 취급물질, 정상 운전조건을 사실 위주로 입력")
    ws.cell(5, 2).fill = PatternFill("solid", fgColor=EXAMPLE_FILL if example else INPUT_FILL)
    meta_rows.append(["FORM", ws.title, key, 5, 2, ""])
    _autosize(ws, max_width=48)
    ws.column_dimensions["A"].width = 46
    ws.column_dimensions["B"].width = 60
    ws.column_dimensions["C"].width = 60
    ws.column_dimensions["D"].width = 42


def _write_meta(wb: Workbook, project: Stage2Project, *, example: bool, meta_rows: list[list[Any]]) -> None:
    ws = wb.create_sheet(META_SHEET)
    ws.append(["schema_version", SCHEMA_VERSION])
    ws.append(["workbook_mode", MODE_EXAMPLE if example else MODE_INPUT])
    ws.append(["project_id", project.project_id])
    ws.append(["psm_selected", str(project.psm_in_scope)])
    ws.append(["cap_selected", str(project.cap_in_scope)])
    ws.append([])
    ws.append(["kind", "sheet", "field_keys", "row_or_header", "value_column", "basis"])
    for row in meta_rows:
        ws.append(row)
    ws.sheet_state = "hidden"


def build_integrated_authoring_workbook(project: Stage2Project, *, example: bool = False) -> bytes:
    if not project.scope_confirmed or not (project.psm_in_scope or project.cap_in_scope):
        raise ValueError("먼저 작성범위를 선택해야 통합 작성자료를 만들 수 있습니다.")

    wb = Workbook()
    meta_rows: list[list[Any]] = []
    _write_guide(wb, project, example=example)
    _write_business_sheet(wb, project, example=example, meta_rows=meta_rows)

    for spec in TABLE_SPECS:
        if _scope_matches(project, str(spec["scope"])):
            _write_table_sheet(wb, project, spec, example=example, meta_rows=meta_rows)

    _write_process_sheet(wb, project, example=example, meta_rows=meta_rows)
    _write_attachment_sheet(wb, project, example=example, meta_rows=meta_rows)

    groups = _build_question_groups(project)
    for group_name, rows in groups.items():
        if not rows:
            continue
        # 이미 별도 공정정보 시트에서 받는 값은 제외한다.
        filtered = [row for row in rows if row[0] != "process.description"]
        if filtered:
            _write_question_sheet(wb, group_name, filtered, project=project, example=example, meta_rows=meta_rows)

    _write_meta(wb, project, example=example, meta_rows=meta_rows)
    output = BytesIO()
    wb.save(output)
    return output.getvalue()


def _meta_map(wb) -> tuple[dict[str, str], list[dict[str, Any]]]:
    if META_SHEET not in wb.sheetnames:
        raise ValueError("프로그램 통합 작성자료가 아닙니다. _시스템정보 시트를 찾을 수 없습니다.")
    ws = wb[META_SHEET]
    header: dict[str, str] = {}
    for row in range(1, 6):
        key = str(ws.cell(row, 1).value or "").strip()
        value = str(ws.cell(row, 2).value or "").strip()
        header[key] = value
    records: list[dict[str, Any]] = []
    for row in ws.iter_rows(min_row=8, values_only=True):
        if not any(value not in (None, "") for value in row):
            continue
        records.append({
            "kind": str(row[0] or ""),
            "sheet": str(row[1] or ""),
            "field_keys": str(row[2] or ""),
            "row_or_header": int(row[3]) if row[3] not in (None, "") else 0,
            "value_column": int(row[4]) if row[4] not in (None, "") else 0,
            "basis": str(row[5] or ""),
        })
    return header, records


def _clean_value(value: Any) -> Any:
    if value is None:
        return ""
    if isinstance(value, str):
        return value.strip()
    return value


def _parse_table(ws, header_row: int) -> list[dict[str, Any]]:
    headers = [str(cell.value or "").strip() for cell in ws[header_row]]
    last_col = max((idx for idx, value in enumerate(headers, 1) if value), default=0)
    if not last_col:
        return []
    headers = headers[:last_col]
    rows: list[dict[str, Any]] = []
    blank_run = 0
    for row_no in range(header_row + 1, ws.max_row + 1):
        values = [_clean_value(ws.cell(row_no, col).value) for col in range(1, last_col + 1)]
        if not any(value not in (None, "") for value in values):
            blank_run += 1
            if blank_run >= 5:
                break
            continue
        blank_run = 0
        rows.append({header: value for header, value in zip(headers, values) if header})
    return rows


def _norm_compare(value: Any) -> str:
    return re.sub(r"\s+", "", str(value or "")).strip().lower()


def apply_integrated_authoring_workbook(
    project: Stage2Project,
    workbook_bytes: bytes,
    *,
    workbook_evidence: EvidenceRef | None = None,
) -> IntegratedImportResult:
    wb = load_workbook(BytesIO(workbook_bytes), data_only=False)
    meta, records = _meta_map(wb)
    if meta.get("schema_version") != SCHEMA_VERSION:
        raise ValueError("지원하지 않는 통합 작성자료 버전입니다. 프로그램에서 최신 입력파일을 다시 내려받아 작성해 주세요.")
    if meta.get("workbook_mode") != MODE_INPUT:
        raise ValueError("작성예시 파일은 제출할 수 없습니다. '통합 작성자료.xlsx'에 실제 사업장 정보를 작성해 업로드해 주세요.")
    if meta.get("project_id") != project.project_id:
        raise ValueError("다른 작성 프로젝트에서 내려받은 통합 작성자료입니다. 현재 프로젝트의 파일을 다시 내려받아 주세요.")

    warnings: list[str] = []
    updated = 0
    table_fields = 0
    attachment_count = 0
    evidence = [workbook_evidence] if workbook_evidence is not None else []

    for item in records:
        sheet_name = item["sheet"]
        if sheet_name not in wb.sheetnames:
            warnings.append(f"시트 누락: {sheet_name}")
            continue
        ws = wb[sheet_name]
        kind = item["kind"]
        field_keys = [key for key in item["field_keys"].split("|") if key]
        if not field_keys:
            continue

        if kind == "PROTECTED":
            key = field_keys[0]
            current = project.get_field(key)
            cell_value = _clean_value(ws.cell(item["row_or_header"], item["value_column"]).value)
            if current and cell_value not in (None, "") and _norm_compare(cell_value) != _norm_compare(current.value):
                warnings.append(
                    f"{field_label(key)} 값이 Stage 1 판정자료와 다릅니다. 통합 작성자료에서 덮어쓰지 않았으며, 변경이 필요하면 판정진단을 다시 수행해 주세요."
                )
            continue

        if kind == "FORM":
            key = field_keys[0]
            value = _clean_value(ws.cell(item["row_or_header"], item["value_column"]).value)
            if value in (None, ""):
                continue
            project.set_field(
                key,
                field_label(key),
                value,
                "USER_CONFIRMED",
                evidence=evidence,
                note="통합 작성자료 Excel에서 회사가 직접 입력한 내용",
            )
            updated += 1
            continue

        if kind == "TABLE":
            rows = _parse_table(ws, item["row_or_header"])
            if not rows:
                continue
            for key in field_keys:
                if key in PROTECTED_STAGE1_FIELDS:
                    continue
                status = "USER_CONFIRMED"
                note = "통합 작성자료 Excel의 구조화 표에서 회사가 직접 입력한 내용"
                if key == "psm.psi.form_applicability":
                    allowed = {"적용", "해당 없음"}
                    complete = all(
                        str(row.get("적용여부") or "").strip() in allowed
                        and bool(str(row.get("확인근거") or "").strip())
                        for row in rows
                    )
                    if not complete:
                        status = "HOLD"
                        note = (
                            "PSM 조건부 별지서식의 적용여부 또는 확인근거가 일부 비어 있습니다. "
                            "모든 서식에 대해 '적용' 또는 '해당 없음'과 확인근거를 회사가 확인해야 합니다."
                        )
                project.set_field(
                    key,
                    field_label(key),
                    rows,
                    status,
                    evidence=evidence,
                    note=note,
                )
                updated += 1
                table_fields += 1
            continue

        if kind == "ATTACHMENT":
            key = field_keys[0]
            row_no = item["row_or_header"]
            file_name = str(ws.cell(row_no, 2).value or "").strip()
            if not file_name:
                continue
            value = {
                "file_name": file_name,
                "reference_no": str(ws.cell(row_no, 3).value or "").strip(),
                "revision": str(ws.cell(row_no, 4).value or "").strip(),
                "note": str(ws.cell(row_no, 5).value or "").strip(),
            }
            project.set_field(
                key,
                field_label(key),
                value,
                "HOLD",
                evidence=evidence,
                note="통합 작성자료에 파일명이 기재되었으나 실제 도면·첨부파일의 별도 업로드 및 내용 확인이 필요합니다.",
            )
            attachment_count += 1

    return IntegratedImportResult(
        updated_fields=updated,
        table_fields=table_fields,
        attachment_declarations=attachment_count,
        warnings=tuple(warnings),
    )


def declared_attachment_file_names(record_value: Any) -> tuple[str, ...]:
    names: list[str] = []

    def add(value: object) -> None:
        text = str(value or "").strip()
        if text and text not in names:
            names.append(text)

    if isinstance(record_value, Mapping):
        for key in ("file_name", "파일명", "SDS 파일명", "MSDS 파일명", "SDS 원본 파일명"):
            add(record_value.get(key))
    elif isinstance(record_value, list):
        for item in record_value:
            if isinstance(item, Mapping):
                for key in ("file_name", "파일명", "SDS 파일명", "MSDS 파일명", "SDS 원본 파일명"):
                    add(item.get(key))
            elif isinstance(item, str):
                add(item)
    elif isinstance(record_value, str):
        add(record_value)
    return tuple(names)


def declared_attachment_file_name(record_value: Any) -> str:
    names = declared_attachment_file_names(record_value)
    return names[0] if names else ""


def attach_company_file(project: Stage2Project, evidence: EvidenceRef) -> tuple[str, ...]:
    """Attach a separately uploaded drawing/document to fields that declared the same file name."""
    matched: list[str] = []
    target_name = Path(evidence.source_name).name.lower()
    for key, record in list(project.fields.items()):
        declared_names = {
            Path(name).name.lower()
            for name in declared_attachment_file_names(record.value)
            if str(name or "").strip()
        }
        if target_name not in declared_names:
            continue
        refs = list(record.evidence)
        if not any(ref.sha256 == evidence.sha256 for ref in refs):
            refs.append(evidence)
        status = record.status if record.status in {"VERIFIED", "USER_CONFIRMED", "CALCULATED"} else "HOLD"
        note = (
            "통합 작성자료의 회사 확정값에 실제 원본 첨부파일이 증빙으로 연결되었습니다."
            if status != "HOLD"
            else "통합 작성자료에 기재된 파일과 실제 첨부파일이 연결되었습니다. 내용 확인 전까지 사람 확인 필요 상태로 유지합니다."
        )
        project.set_field(
            key,
            record.label,
            record.value,
            status,
            evidence=refs,
            note=note,
        )
        matched.append(key)
    if not matched:
        existing = project.get_field("attachments.unclassified")
        refs = list(existing.evidence) if existing else []
        refs.append(evidence)
        names = list(existing.value) if existing and isinstance(existing.value, list) else []
        names.append(evidence.source_name)
        project.set_field(
            "attachments.unclassified",
            "미분류 첨부자료",
            names,
            "HOLD",
            evidence=refs,
            note="통합 작성자료의 파일명과 자동 연결되지 않은 첨부자료입니다. 사람이 연결대상을 확인해야 합니다.",
        )
    return tuple(matched)
