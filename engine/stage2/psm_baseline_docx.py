from __future__ import annotations

"""Write confirmed PSM company data into the preserved statutory-form DOCX baseline.

The bundled DOCX is a layout baseline derived from the regulation forms supplied
by the user. It is deliberately not treated as the legal-currentness authority;
that remains the existing law.go.kr monitoring/approval path. This writer only
preserves the supplied form geometry and inserts already-confirmed project data.
Unknown values stay blank.
"""

import base64
from copy import deepcopy
from hashlib import sha256
from io import BytesIO
import json
from pathlib import Path
import re
from typing import Any, Mapping, Sequence

from docx import Document
from docx.table import Table

from . import statutory_report as base
from .project import Stage2Project
from .ooxml_determinism import canonicalize_docx_zip

PROJECT_ROOT = Path(__file__).resolve().parents[2]
TEMPLATE_DIR = PROJECT_ROOT / "data" / "templates" / "psm"
METADATA_PATH = TEMPLATE_DIR / "psm_statutory_forms_baseline.json"
BASE64_GLOB = "psm_statutory_forms_baseline.docx.b64.*"
MISSING = base.MISSING

FORM_TABLE_INDEX = {
    "12": 0, "13": 1, "14": 2, "15": 3, "16": 4, "17": 5,
    "17-2": 6, "17-3": 7, "17-4": 8, "17-5": 9,
    "18": 10, "19": 11, "19-2": 12, "20": 13, "21": 14,
}

LATER_FORM_FIELDS: dict[str, tuple[str, ...]] = {
    # Reuse the already-established Stage 2 field used by the existing PSM
    # renderer first; keep the older aliases only as fallback so no interlock
    # data is inferred that the company never confirmed.
    "17-2": ("psm.psi.interlock_conditions", "psm.psi.interlock_specs", "psm.psi.interlocks"),
    "17-3": ("psm.psi.fire_protection_table", "psm.psi.fire_protection"),
    "17-4": ("psm.psi.fire_detection_table", "psm.psi.fire_detection"),
    "17-5": ("psm.psi.gas_detection_table", "psm.psi.gas_detection"),
    "18": ("psm.psi.fireproofing_table", "psm.psi.fireproofing"),
    "19": ("psm.psi.local_exhaust_table", "psm.psi.local_exhaust"),
    "20": ("psm.psi.ex_equipment",),
    "21": ("psm.risk.team",),
}

HEADER_ALIASES: dict[str, tuple[str, ...]] = {
    "인터록번호": ("인터록번호", "인터록 번호", "interlock_no", "interlock"),
    "대상설비번호": ("대상설비번호", "대상설비", "설비번호", "대상 설비"),
    "설정값-온도(℃)": ("설정값-온도(℃)", "설정온도", "온도", "temperature"),
    "설정값-압력(MPa)": ("설정값-압력(MPa)", "설정압력", "압력", "pressure"),
    "설정값-액위(m)": ("설정값-액위(m)", "설정액위", "액위", "level"),
    "설정값-기타": ("설정값-기타", "기타 설정값", "기타"),
    "감지기번호": ("감지기번호", "감지기 번호", "계기번호", "계기 번호"),
    "최종 작동설비번호": ("최종 작동설비번호", "최종작동설비", "최종 작동설비"),
    "가동중지범위": ("가동중지범위", "가동중지 범위", "중지범위"),
    "점검주기": ("점검주기", "점검 주기"),
    "설치지역": ("설치지역", "설치 지역", "지역"),
    "단독경보형 감지기": ("단독경보형 감지기", "단독경보형감지기"),
    "비상경보설비": ("비상경보설비", "비상 경보설비"),
    "시각경보기": ("시각경보기", "시각 경보기"),
    "자동화재탐지설비": ("자동화재탐지설비", "자동 화재탐지설비"),
    "비상방송설비": ("비상방송설비", "비상 방송설비"),
    "자동화재속보설비": ("자동화재속보설비", "자동 화재속보설비"),
    "통합감시시설": ("통합감시시설", "통합 감시시설"),
    "누전경보기": ("누전경보기", "누전 경보기"),
    "감지대상": ("감지대상", "감지 대상", "검출대상 물질", "검출대상", "물질명"),
    "설치장소": ("설치장소", "설치 장소", "설치위치", "설치 위치"),
    "작동시간": ("작동시간", "작동 시간"),
    "측정방식": ("측정방식", "측정 방식"),
    "경보설정값": ("경보설정값", "경보 설정값", "설정값"),
    "경보기 위치": ("경보기 위치", "경보기위치", "경보 위치"),
    "정밀도": ("정밀도", "오차범위", "정밀도(오차범위)"),
    "경보시 조치내용": ("경보시 조치내용", "경보 시 조치내용", "연동 설비·조치", "연동설비", "조치내용"),
    "유지관리": ("유지관리", "유지 관리", "교정주기", "교정 주기"),
    "내화설비 또는 지역": ("내화설비 또는 지역", "내화설비", "지역"),
    "내화부위": ("내화부위", "내화 부위"),
    "내화시험기준 및 시간": ("내화시험기준 및 시간", "내화시험기준", "내화시간"),
    "공정 또는 작업장명": ("공정 또는 작업장명", "공정명", "작업장명"),
    "실내외 구분": ("실내외 구분", "실내외", "구분"),
    "발생원": ("발생원", "발산원"),
    "유해물질 종류": ("유해물질 종류", "유해물질", "물질종류"),
    "후드형식": ("후드형식", "후드 형식"),
    "후드 제어풍속(m/s)": ("후드 제어풍속(m/s)", "제어풍속", "후드의 제어풍속"),
    "덕트내 반송속도(m/s)": ("덕트내 반송속도(m/s)", "반송속도", "덕트내 반송속도"),
    "배풍량(m3/min)": ("배풍량(m3/min)", "배풍량", "배풍량(㎥/min)"),
    "전동기용량(kW)": ("전동기용량(kW)", "전동기용량", "전동기 용량"),
    "배기 및 처리순서": ("배기 및 처리순서", "배기·처리순서", "처리순서"),
    "방폭형식": ("방폭형식", "방폭 형식"),
    "설치장소 또는 공정": ("설치장소 또는 공정", "설치장소", "공정"),
    "전기/계장 기계·기구명": ("전기/계장 기계·기구명", "전기/계장기계 기구명", "기계기구명"),
    "0종장소 선정기준(방폭형식)": ("0종장소 선정기준(방폭형식)", "0종장소", "0종 방폭형식"),
    "1종장소 선정기준(방폭형식)": ("1종장소 선정기준(방폭형식)", "1종장소", "1종 방폭형식"),
    "2종장소 선정기준(방폭형식)": ("2종장소 선정기준(방폭형식)", "2종장소", "2종 방폭형식"),
    "책임분야": ("책임분야", "분야"),
    "성명": ("성명", "이름"),
    "소속회사": ("소속회사", "소속", "회사"),
    "직책": ("직책", "직위"),
    "주요경력": ("주요경력", "경력"),
}

# The uploaded statutory Form 19 places the explosion-proof type before the
# exhaust/treatment sequence. Keep this baseline-specific output ordering here
# so the shared PSM_FORMS schema can remain backward compatible elsewhere.
FORM_OUTPUT_HEADERS: dict[str, tuple[str, ...]] = {
    "19": (
        "공정 또는 작업장명", "실내외 구분", "발생원", "유해물질 종류", "후드형식",
        "후드 제어풍속(m/s)", "덕트내 반송속도(m/s)", "배풍량(m3/min)",
        "전동기용량(kW)", "방폭형식", "배기 및 처리순서",
    ),
}


def _metadata() -> dict[str, Any]:
    with METADATA_PATH.open("r", encoding="utf-8") as handle:
        return json.load(handle)


def load_psm_baseline_bytes() -> bytes:
    meta = _metadata()
    parts = sorted(TEMPLATE_DIR.glob(BASE64_GLOB))
    if not parts:
        raise FileNotFoundError("PSM 규정서식 baseline 파일이 준비되지 않았습니다.")
    decoded_parts: list[bytes] = []
    for part in parts:
        encoded = "".join(part.read_text(encoding="ascii").split())
        try:
            decoded_parts.append(base64.b64decode(encoded, validate=True))
        except Exception as exc:
            raise ValueError("PSM 규정서식 baseline 인코딩을 읽지 못했습니다.") from exc
    data = b"".join(decoded_parts)
    expected = str(meta.get("sha256") or "").lower()
    actual = sha256(data).hexdigest()
    if expected and actual != expected:
        raise ValueError("PSM 규정서식 baseline 해시가 등록값과 일치하지 않습니다.")
    validate_psm_baseline(data)
    return data


def validate_psm_baseline(data: bytes) -> None:
    meta = _metadata()
    try:
        doc = Document(BytesIO(data))
    except Exception as exc:
        raise ValueError("PSM 규정서식 baseline DOCX를 열 수 없습니다.") from exc
    required_count = int(meta.get("required_table_count") or 0)
    if required_count and len(doc.tables) != required_count:
        raise ValueError(
            f"PSM 규정서식 baseline의 표 수가 예상과 다릅니다: {len(doc.tables)}개 / 예상 {required_count}개"
        )
    text = "\n".join(cell.text for table in doc.tables for row in table.rows for cell in _unique_cells(row))
    missing = [str(marker) for marker in meta.get("required_markers") or [] if str(marker) not in text]
    if missing:
        raise ValueError("PSM 규정서식 baseline에서 필수 별지표시를 찾지 못했습니다: " + ", ".join(missing))


def _unique_cells(row) -> list:
    seen: set[int] = set()
    cells = []
    for cell in row.cells:
        marker = id(cell._tc)
        if marker in seen:
            continue
        seen.add(marker)
        cells.append(cell)
    return cells


def _clean(value: object) -> str:
    if value in (None, "", MISSING):
        return ""
    if isinstance(value, Mapping):
        parts = [f"{key}: {val}" for key, val in value.items() if val not in (None, "", MISSING)]
        return " / ".join(parts)
    if isinstance(value, (list, tuple)):
        return ", ".join(_clean(item) for item in value if _clean(item))
    return str(value).strip()


def _write_cell(cell, value: object) -> None:
    text = _clean(value)
    runs = [run for paragraph in cell.paragraphs for run in paragraph.runs]
    if runs:
        runs[0].text = text
        for run in runs[1:]:
            run.text = ""
        return
    paragraph = cell.paragraphs[0] if cell.paragraphs else cell.add_paragraph()
    paragraph.add_run(text)


def _append_value(cell, value: object) -> None:
    text = _clean(value)
    if not text:
        return
    label = cell.text.strip()
    _write_cell(cell, (label + " " + text).strip())


def _row_text(row) -> str:
    return " ".join(cell.text.strip() for cell in _unique_cells(row)).strip()


def _table_data_bounds(table) -> tuple[int, int]:
    note_index = len(table.rows)
    for index, row in enumerate(table.rows[3:], start=3):
        text = _row_text(row)
        if text.startswith("주)") or "210㎜×297㎜" in text:
            note_index = index
            break
    for index in range(3, note_index):
        cells = _unique_cells(table.rows[index])
        if cells and not any(cell.text.strip() for cell in cells):
            return index, note_index
    raise ValueError("PSM 규정서식에서 데이터 입력행을 찾지 못했습니다.")


def _ensure_row_capacity(table, required: int) -> int:
    start, note_index = _table_data_bounds(table)
    capacity = note_index - start
    if required <= capacity:
        return start
    template_tr = table.rows[start]._tr
    note_tr = table.rows[note_index]._tr
    for _ in range(required - capacity):
        note_tr.addprevious(deepcopy(template_tr))
    return start


def _fill_table_rows(table, rows: Sequence[Sequence[object]]) -> None:
    if not rows:
        return
    start = _ensure_row_capacity(table, len(rows))
    for offset, values in enumerate(rows):
        cells = _unique_cells(table.rows[start + offset])
        if len(values) > len(cells):
            raise ValueError(
                f"PSM 규정서식 열 수가 입력자료보다 적습니다: 표 열 {len(cells)}개 / 입력 {len(values)}개"
            )
        for index, cell in enumerate(cells):
            _write_cell(cell, values[index] if index < len(values) else "")


def _field_text(project: Stage2Project, *keys: str) -> str:
    return _clean(base._value(project, *keys, default=""))


def _writer_parts(project: Stage2Project) -> tuple[str, str]:
    value = base._value(project, "psm.business.writer_info", default="")
    if isinstance(value, Mapping):
        name = ""
        qualification = ""
        normalized = {base._norm(key): val for key, val in value.items()}
        for key in ("작성자", "성명", "이름", "name"):
            if normalized.get(base._norm(key)) not in (None, ""):
                name = str(normalized[base._norm(key)]).strip()
                break
        for key in ("작성자 자격", "자격", "qualification"):
            if normalized.get(base._norm(key)) not in (None, ""):
                qualification = str(normalized[base._norm(key)]).strip()
                break
        return name, qualification
    text = _clean(value)
    if " / " in text:
        left, right = text.split(" / ", 1)
        return left.strip(), right.strip()
    return text, ""


def _project_type_options(project: Stage2Project, raw_value: object = "") -> str:
    raw = _clean(raw_value) or _field_text(project, "psm.business.project_type")
    normalized = re.sub(r"\s+", "", raw)
    install = bool(raw and ("설치" in normalized or "이전" in normalized))
    change = bool(raw and "변경" in normalized)
    existing = bool(raw and "기존" in normalized)
    return "\n".join((
        ("☒" if install else "☐") + " 설치·이전",
        ("☒" if change else "☐") + " 변경",
        ("☒" if existing else "☐") + " 기존설비",
    ))


def _mapping_lookup(mapping: Mapping[str, object], *keys: str) -> object:
    normalized = {base._norm(key): value for key, value in mapping.items()}
    for key in keys:
        value = normalized.get(base._norm(key))
        if value not in (None, ""):
            return value
    return ""


def _form12_site_parts(project: Stage2Project) -> tuple[str, str, str, str]:
    address = _field_text(project, "business.address")
    site_area = _field_text(project, "psm.business.site_area", "business.site_area")
    main_building = _field_text(project, "psm.business.main_building", "business.main_building")
    site_building = base._value(project, "psm.business.site_building", default="")
    if isinstance(site_building, Mapping):
        address = address or _clean(_mapping_lookup(site_building, "위치", "주소", "location", "address"))
        site_area = site_area or _clean(_mapping_lookup(site_building, "부지", "부지면적", "site_area", "area"))
        main_building = main_building or _clean(
            _mapping_lookup(site_building, "주요건물", "건물", "main_building", "building")
        )
    else:
        main_building = main_building or _clean(site_building)
    phone = _field_text(project, "business.phone", "psm.business.phone")
    fax = _field_text(project, "business.fax", "psm.business.fax")
    contact = f"전화번호: {phone}\n전송번호: {fax}".rstrip() if phone or fax else ""
    return address, contact, site_area, main_building


def _form12_schedule_parts(project: Stage2Project) -> tuple[str, str, str]:
    schedule = base._value(project, "psm.business.schedule", default="")
    total = _field_text(project, "psm.business.total_period")
    start = _field_text(project, "psm.business.start_date")
    commissioning = _field_text(project, "psm.business.commissioning_period")
    if isinstance(schedule, Mapping):
        total = total or _clean(_mapping_lookup(schedule, "총사업기간", "사업기간", "total_period"))
        start = start or _clean(_mapping_lookup(schedule, "착공예정일", "착공일", "start_date"))
        commissioning = commissioning or _clean(
            _mapping_lookup(schedule, "시운전기간", "시운전", "commissioning_period")
        )
    elif schedule not in (None, "", MISSING):
        total = total or _clean(schedule)
    return total, start, commissioning


def _form12_detail(project: Stage2Project) -> Mapping[str, object]:
    rows = base._rows(project, "psm.business.form12_details")
    return rows[0] if rows and isinstance(rows[0], Mapping) else {}


def _fill_form12(table, project: Stage2Project) -> None:
    detail = _form12_detail(project)

    def value(header: str, *fallback_keys: str) -> str:
        direct = _mapping_lookup(detail, header)
        if direct not in (None, ""):
            return _clean(direct)
        return _field_text(project, *fallback_keys) if fallback_keys else ""

    chemicals = base._rows(project, "psm.psi.chemical_details", "inventory.chemicals")
    inventory_raw_materials = ", ".join(
        item for row in chemicals[:8]
        if (item := _clean(base._row_value(row, "물질명", "화학물질", "유해화학물질명")))
    )
    raw_materials = value("주요 원료") or inventory_raw_materials

    legacy_writer, legacy_qualification = _writer_parts(project)
    writer = value("작성자 성명") or legacy_writer
    qualification = value("작성자 자격") or legacy_qualification

    # Stage 1 identity remains authoritative for protected company/address facts.
    company_name = project.company_name or value("사업장명")
    address = _field_text(project, "business.address") or value("사업장 소재지")

    _append_value(_unique_cells(table.rows[3])[0], company_name)
    _write_cell(_unique_cells(table.rows[3])[2], _project_type_options(project, value("제출구분")))
    _append_value(
        _unique_cells(table.rows[4])[0],
        value("사업자등록번호", "business.registration_no", "cap.business.registration_no"),
    )
    _append_value(
        _unique_cells(table.rows[5])[0],
        value("대표자", "business.representative", "cap.business.representative"),
    )
    _write_cell(
        _unique_cells(table.rows[5])[2],
        value("대상 유해·위험설비", "psm.business.target_facility"),
    )
    _append_value(_unique_cells(table.rows[6])[0], value("한국표준산업분류", "business.ksic"))
    _append_value(_unique_cells(table.rows[7])[0], value("근로자수", "business.employee_count"))
    electric = value("계약전력(kW)", "business.electric_contract_capacity")
    if electric:
        _write_cell(_unique_cells(table.rows[7])[2], f"{electric} ㎾")
    _write_cell(_unique_cells(table.rows[8])[1], writer)
    _write_cell(_unique_cells(table.rows[8])[3], qualification)
    _write_cell(_unique_cells(table.rows[11])[2], raw_materials)
    _write_cell(_unique_cells(table.rows[12])[2], value("주요 생산품", "business.main_products"))
    _write_cell(_unique_cells(table.rows[13])[2], value("사업개요", "psm.business.overview"))

    phone = value("전화번호", "business.phone", "psm.business.phone")
    fax = value("전송번호", "business.fax", "psm.business.fax")
    contact = f"전화번호: {phone}\n전송번호: {fax}".rstrip() if phone or fax else ""
    if address:
        _write_cell(_unique_cells(table.rows[14])[2], address)
    if contact:
        _write_cell(_unique_cells(table.rows[14])[3], contact)

    legacy_address, _, legacy_site_area, legacy_main_building = _form12_site_parts(project)
    site_area = value("부지면적", "psm.business.site_area", "business.site_area") or legacy_site_area
    main_building = value(
        "주요 건물", "psm.business.main_building", "business.main_building", "psm.business.site_building"
    ) or legacy_main_building
    if site_area:
        _write_cell(_unique_cells(table.rows[15])[2], site_area)
    if main_building:
        _write_cell(_unique_cells(table.rows[16])[2], main_building)

    legacy_total, legacy_start, legacy_commissioning = _form12_schedule_parts(project)
    total_period = value("총 사업기간", "psm.business.total_period", "psm.business.schedule") or legacy_total
    start_date = value("착공예정일", "psm.business.start_date") or legacy_start
    commissioning_period = value("시운전기간", "psm.business.commissioning_period") or legacy_commissioning
    if total_period:
        _write_cell(_unique_cells(table.rows[17])[2], total_period)
    if start_date:
        _write_cell(_unique_cells(table.rows[18])[2], start_date)
    if commissioning_period:
        _write_cell(_unique_cells(table.rows[19])[2], commissioning_period)


def _sanitize_rows(rows: Sequence[Sequence[object]]) -> list[list[str]]:
    return [[_clean(value) for value in row] for row in rows]


def _aliases(header: str) -> tuple[str, ...]:
    return HEADER_ALIASES.get(header, (header,))


def _structured_rows(project: Stage2Project, form_key: str, field_keys: Sequence[str]) -> list[list[str]]:
    rows = base._rows(project, *field_keys)
    headers = FORM_OUTPUT_HEADERS.get(form_key, base.PSM_FORMS[form_key].headers)
    return [[_clean(base._row_value(row, *_aliases(header))) for header in headers] for row in rows]


def _option_text(value: object, options: Sequence[tuple[str, Sequence[str]]], *, suffix: str = "") -> str:
    text = _clean(value)
    normalized = base._norm(text)
    selected_index: int | None = None
    for index, (_, aliases) in enumerate(options):
        if any(base._norm(alias) and base._norm(alias) in normalized for alias in aliases):
            selected_index = index
            break
    rendered = [f"{'☒' if selected_index == index else '☐'} {label}" for index, (label, _) in enumerate(options)]
    if selected_index is None and text:
        rendered.append((suffix + text).strip())
    return "  ".join(rendered)


def _threshold_values(case: Mapping[str, object], group: str, thresholds: Sequence[str]) -> list[str]:
    value = _mapping_lookup(case, group)
    nested = value if isinstance(value, Mapping) else {}
    output: list[str] = []
    for threshold in thresholds:
        direct = _mapping_lookup(nested, threshold, threshold.replace(" ", "")) if nested else ""
        if direct in (None, ""):
            direct = _mapping_lookup(case, f"{group}-{threshold}", f"{group} {threshold}", threshold)
        output.append(_clean(direct))
    return output


def _case_from_consequence_row(row: Mapping[str, object]) -> dict[str, object]:
    return {
        "풍속(m/s)": base._row_value(row, "풍속(m/s)"),
        "대기안정도(A~F)": base._row_value(row, "대기안정도(A~F)"),
        "대기온도(℃)": base._row_value(row, "대기온도(℃)"),
        "습도(%)": base._row_value(row, "습도(%)"),
        "표면거칠기(m)": base._row_value(row, "표면거칠기"),
        "물질명": base._row_value(row, "물질명"),
        "물질의 상태": base._row_value(row, "물질의 상태"),
        "설비명(또는 배관부위)": base._row_value(row, "설비명(또는 배관부위)"),
        "운전압력(MPa)": base._row_value(row, "운전압력(MPa)"),
        "운전온도(℃)": base._row_value(row, "운전온도(℃)"),
        "누출구의 크기(mm2)": base._row_value(row, "누출구의 크기(mm2)"),
        "웅덩이 크기(m2)": base._row_value(row, "웅덩이 크기(m2)"),
        "누출결과": base._row_value(row, "누출결과"),
        "직접계산(kg/s or kg)": base._row_value(row, "직접계산(kg/s or kg)"),
        "웅덩이(kg/s)": base._row_value(row, "웅덩이(kg/s)"),
        "설비/배관(kg/s)": base._row_value(row, "설비/배관(kg/s)"),
        "화재-복사열이 미치는 거리": {
            "4 kW/m2": base._row_value(row, "화재-4 kW/m2"),
            "12.5 kW/m2": base._row_value(row, "화재-12.5 kW/m2"),
            "37.5 kW/m2": base._row_value(row, "화재-37.5 kW/m2"),
        },
        "폭발-과압이 미치는 거리": {
            "7 kPa": base._row_value(row, "폭발-7 kPa"),
            "21 kPa": base._row_value(row, "폭발-21 kPa"),
            "70 kPa": base._row_value(row, "폭발-70 kPa"),
        },
        "확산결과-인화성": {
            "25% LEL": base._row_value(row, "인화성-25% LEL"),
            "LEL": base._row_value(row, "인화성-LEL"),
            "UEL": base._row_value(row, "인화성-UEL"),
        },
        "확산결과-독성": {
            "ERPG 1": base._row_value(row, "독성-ERPG 1"),
            "ERPG 2": base._row_value(row, "독성-ERPG 2"),
            "ERPG 3": base._row_value(row, "독성-ERPG 3"),
        },
    }


def _consequence_pairs(project: Stage2Project) -> list[tuple[Mapping[str, object], Mapping[str, object]]]:
    rows = [row for row in base._rows(project, "psm.risk.consequence_table") if isinstance(row, Mapping)]
    groups: dict[tuple[str, str], dict[str, list[Mapping[str, object]]]] = {}
    metadata_present = False
    for row in rows:
        unit = _clean(base._row_value(row, "단위공장", "단위공장·공정"))
        accident = _clean(base._row_value(row, "사고유형"))
        scenario = base._norm(base._row_value(row, "시나리오 구분", "시나리오구분"))
        if unit and unit != MISSING and accident and accident != MISSING:
            metadata_present = True
            bucket = groups.setdefault((base._norm(unit), base._norm(accident)), {"worst": [], "alternative": []})
            if "최악" in scenario:
                bucket["worst"].append(row)
            elif "대안" in scenario:
                bucket["alternative"].append(row)

    if metadata_present and groups:
        pairs: list[tuple[Mapping[str, object], Mapping[str, object]]] = []
        for bucket in groups.values():
            worst = bucket["worst"][0] if bucket["worst"] else {}
            alternatives = bucket["alternative"] or [{}]
            for alternative in alternatives:
                pairs.append((
                    _case_from_consequence_row(worst) if worst else {},
                    _case_from_consequence_row(alternative) if alternative else {},
                ))
        if pairs:
            return pairs

    # Backward-compatible single worst/alternative source.
    output: dict[str, object] = {}
    for row in rows:
        scenario = _clean(base._row_value(row, "시나리오 구분", "시나리오구분"))
        normalized = base._norm(scenario)
        if "최악" in normalized:
            output["최악의 사고 시나리오"] = _case_from_consequence_row(row)
        elif "대안" in normalized and "대안의 사고 시나리오" not in output:
            output["대안의 사고 시나리오"] = _case_from_consequence_row(row)
    if not output:
        legacy = base._value(project, "psm.risk.consequence", default={})
        output = legacy if isinstance(legacy, Mapping) else {}

    worst = _mapping_lookup(output, "최악의 사고 시나리오", "worst_case", "worst")
    alternative = _mapping_lookup(output, "대안의 사고 시나리오", "alternative_case", "alternative")
    worst = worst if isinstance(worst, Mapping) else {}
    alternative = alternative if isinstance(alternative, Mapping) else {}
    return [(worst, alternative)] if worst or alternative else []


def _fill_form19_2_cases(
    table,
    worst: Mapping[str, object],
    alternative: Mapping[str, object],
) -> None:
    direct_rows = {
        4: "풍속(m/s)", 5: "대기안정도(A~F)", 6: "대기온도(℃)", 7: "습도(%)",
        10: "물질명", 12: "설비명(또는 배관부위)", 13: "운전압력(MPa)", 14: "운전온도(℃)",
        15: "누출구의 크기(mm2)", 16: "웅덩이 크기(m2)", 18: "누출결과",
        19: "직접계산(kg/s or kg)", 20: "웅덩이(kg/s)", 21: "설비/배관(kg/s)",
    }
    for row_index, label in direct_rows.items():
        cells = _unique_cells(table.rows[row_index])
        _write_cell(cells[1], _mapping_lookup(worst, label))
        _write_cell(cells[2], _mapping_lookup(alternative, label))

    surface_options = (("시골", ("시골", "rural")), ("도시", ("도시", "urban")), ("물위", ("물위", "water")))
    state_options = (("기체", ("기체", "gas")), ("액체", ("액체", "liquid")), ("2상(액체+기체)", ("2상", "two phase", "two-phase")))
    row8 = _unique_cells(table.rows[8])
    _write_cell(row8[1], _option_text(_mapping_lookup(worst, "표면거칠기(m)", "표면거칠기"), surface_options))
    _write_cell(row8[2], _option_text(_mapping_lookup(alternative, "표면거칠기(m)", "표면거칠기"), surface_options))
    row11 = _unique_cells(table.rows[11])
    _write_cell(row11[1], _option_text(_mapping_lookup(worst, "물질의 상태", "물질상태"), state_options))
    _write_cell(row11[2], _option_text(_mapping_lookup(alternative, "물질의 상태", "물질상태"), state_options))

    groups = (
        (24, "화재-복사열이 미치는 거리", ("4 kW/m2", "12.5 kW/m2", "37.5 kW/m2")),
        (26, "폭발-과압이 미치는 거리", ("7 kPa", "21 kPa", "70 kPa")),
        (28, "확산결과-인화성", ("25% LEL", "LEL", "UEL")),
        (30, "확산결과-독성", ("ERPG 1", "ERPG 2", "ERPG 3")),
    )
    for row_index, group, thresholds in groups:
        cells = _unique_cells(table.rows[row_index])
        worst_values = _threshold_values(worst, group, thresholds)
        alternative_values = _threshold_values(alternative, group, thresholds)
        for index, value in enumerate(worst_values, start=1):
            _write_cell(cells[index], value)
        for index, value in enumerate(alternative_values, start=4):
            _write_cell(cells[index], value)


def _fill_and_repeat_form19_2(doc, project: Stage2Project) -> None:
    pairs = _consequence_pairs(project)
    if not pairs:
        return
    template = doc.tables[FORM_TABLE_INDEX["19-2"]]
    _fill_form19_2_cases(template, *pairs[0])
    anchor = template._tbl
    for worst, alternative in pairs[1:]:
        cloned_xml = deepcopy(template._tbl)
        anchor.addnext(cloned_xml)
        cloned = Table(cloned_xml, template._parent)
        _fill_form19_2_cases(cloned, worst, alternative)
        anchor = cloned_xml

def build_psm_baseline_draft(project: Stage2Project) -> bytes:
    if not project.psm_in_scope:
        raise ValueError("공정안전보고서는 현재 작성범위에 포함되어 있지 않습니다.")
    doc = Document(BytesIO(load_psm_baseline_bytes()))
    if len(doc.tables) != len(FORM_TABLE_INDEX):
        raise ValueError("PSM 규정서식 baseline 구조가 예상과 달라 자동작성을 중단했습니다.")

    _fill_form12(doc.tables[FORM_TABLE_INDEX["12"]], project)
    _fill_table_rows(doc.tables[FORM_TABLE_INDEX["13"]], _sanitize_rows(base._psm_form13_rows(project)))
    _fill_table_rows(doc.tables[FORM_TABLE_INDEX["14"]], _sanitize_rows(base._psm_form14_rows(project)))
    _fill_table_rows(doc.tables[FORM_TABLE_INDEX["15"]], _sanitize_rows(base._psm_form15_rows(project)))
    _fill_table_rows(doc.tables[FORM_TABLE_INDEX["16"]], _sanitize_rows(base._psm_form16_rows(project)))
    _fill_table_rows(doc.tables[FORM_TABLE_INDEX["17"]], _sanitize_rows(base._psm_form17_rows(project)))

    for form_key, field_keys in LATER_FORM_FIELDS.items():
        _fill_table_rows(doc.tables[FORM_TABLE_INDEX[form_key]], _structured_rows(project, form_key, field_keys))

    _fill_and_repeat_form19_2(doc, project)
    out = BytesIO()
    doc.save(out)
    return canonicalize_docx_zip(out.getvalue())


def psm_baseline_filename(project: Stage2Project) -> str:
    company = re.sub(r"[^0-9A-Za-z가-힣._-]+", "_", project.company_name or project.project_id).strip("._")
    company = company or "사업장"
    return f"{company}_공정안전보고서_규정서식_작성본.docx"
