from __future__ import annotations

from io import BytesIO
from typing import Any

from openpyxl import load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from .integrated_workbook import build_integrated_authoring_workbook as _base_build
from .project import Stage2Project


OPTIONS_SHEET = "_선택목록"
INPUT_FILL = "FFF2CC"
EXAMPLE_FILL = "E2F0D9"
NOTE_FILL = "F2F2F2"

OPTION_GROUPS: dict[str, tuple[str, ...]] = {
    "physical_state": ("기체", "액체", "고체", "혼합/기타"),
    "mass_volume_unit": ("kg", "ton", "g", "mg", "L", "mL", "m³", "Nm³", "기타(직접입력)"),
    "capacity_unit": ("L", "m³", "kg", "ton", "Nm³", "kg/h", "m³/h", "Nm³/h", "기타(직접입력)"),
    "process_use": (
        "원료저장", "제품저장", "이송", "혼합", "반응", "분리", "정제", "충전", "포장", "세정", "회수", "폐수처리", "출하", "기타(직접입력)"
    ),
    "material_use": ("원료", "부원료", "중간체", "제품", "촉매", "용매", "세정제", "연료", "냉매", "pH 조정제", "기타(직접입력)"),
    "equipment_type": (
        "저장탱크", "반응기", "혼합조", "압력용기", "열교환기", "증류탑", "흡수탑", "펌프", "압축기", "송풍기", "기타(직접입력)"
    ),
    "relief_type": ("안전밸브", "파열판", "안전밸브+파열판", "기타(직접입력)"),
    "discharge_state": ("기체", "증기", "액체", "2상", "기타(직접입력)"),
    "detector_type": ("고정식", "휴대식", "고정식+휴대식", "기타(직접입력)"),
    "yes_no_na": ("예", "아니오", "해당 없음"),
    "machinery_type": ("원심펌프", "용적식펌프", "왕복동압축기", "원심압축기", "교반기", "송풍기", "팬", "기타(직접입력)"),
    "treatment_type": ("흡수", "흡착", "연소", "응축", "중화", "세정", "회수", "폐수처리", "기타(직접입력)"),
    "submission_type": ("신규", "변경", "재제출", "기타(직접입력)"),
}

TABLE_DROPDOWNS: dict[str, dict[str, str]] = {
    "02_화학물질정보": {
        "물리적 상태": "physical_state",
        "단위": "mass_volume_unit",
        "사용·저장 공정": "process_use",
        "주요 용도": "material_use",
    },
    "03_설비정보": {
        "설비종류": "equipment_type",
        "용량단위": "capacity_unit",
    },
    "04_안전밸브_파열판": {
        "형식": "relief_type",
        "배출상태": "discharge_state",
    },
    "05_가스누출감지_경보장치": {
        "감지방식": "detector_type",
        "비상전원 여부": "yes_no_na",
    },
    "10_동력기계": {
        "형식": "machinery_type",
    },
    "20_배출물질_처리시설": {
        "처리방식": "treatment_type",
    },
}

FIELD_DROPDOWNS = {
    "cap.business.submission_type": "submission_type",
    "cap.business.other_system_review": "yes_no_na",
}

INPUT_HINTS = {
    "process.description": "공정단계, 주요 설비, 취급물질, 정상 운전조건을 사실 위주로 작성",
    "cap.business.submission_type": "신규·변경 등 해당 유형을 선택. 목록에 없으면 직접 입력 가능",
    "cap.business.other_system_review": "해당 여부를 선택",
}

TABLE_EXAMPLES: dict[str, tuple[tuple[Any, ...], ...]] = {
    "02_화학물질정보": (
        ("톨루엔", "108-88-3", 99.5, "액체", 15000, "kg", "원료저장", "원료", "유기용제 원료 예시"),
        ("염소", "7782-50-5", 99.9, "기체", 2.5, "ton", "원료저장", "원료", "가스 저장 예시"),
        ("황산", "7664-93-9", 98, "액체", 8, "m³", "폐수처리", "pH 조정제", "부식성 액체 예시"),
        ("아세톤", "67-64-1", 99, "액체", 2000, "L", "세정", "세정제", "세정·회수 공정 예시"),
    ),
    "03_설비정보": (
        ("TK-101", "톨루엔 저장탱크", "저장탱크", "원료저장", "톨루엔", 20, "m³", "0.49 MPa", "80 ℃", "0.15 MPa", "30 ℃", "SUS304", 15000, "PID-101", "저장설비 예시"),
        ("R-201", "합성 반응기", "반응기", "반응", "원료 A/B", 5, "m³", "1.0 MPa", "150 ℃", "0.45 MPa", "95 ℃", "SUS316L", 3200, "PID-201", "반응설비 예시"),
        ("P-301", "제품 이송펌프", "펌프", "이송", "제품", 25, "m³/h", "0.8 MPa", "60 ℃", "0.4 MPa", "35 ℃", "SUS304", "", "PID-301", "동력기계도 설비목록에 연결 가능"),
        ("E-401", "공정 열교환기", "열교환기", "냉각", "공정유체", 120, "m²", "1.2 MPa", "180 ℃", "0.5 MPa", "90 ℃", "SUS316L", "", "PID-401", "열교환기 예시"),
        ("V-501", "중간제품 압력용기", "압력용기", "분리", "중간제품", 3, "m³", "1.5 MPa", "120 ℃", "0.7 MPa", "70 ℃", "SUS304", 1800, "PID-501", "압력용기 예시"),
    ),
    "04_안전밸브_파열판": (
        ("PSV-101", "TK-101", "안전밸브", "0.45 MPa", "1200 kg/h", "톨루엔 증기", "증기", "스크러버", "PID-101", "저장탱크 보호 예시"),
        ("PSV-201", "R-201", "안전밸브", "0.90 MPa", "850 kg/h", "반응 혼합증기", "2상", "플레어 헤더", "PID-201", "반응기 보호 예시"),
        ("RD-501", "V-501", "파열판", "1.30 MPa", "700 kg/h", "공정유체", "기체", "비상벤트", "PID-501", "파열판 예시"),
    ),
    "05_가스누출감지_경보장치": (
        ("GD-101", "TK-101 방유제 내", "톨루엔", "고정식", "10% LEL", "중앙제어실", "예", "GA-101", "인화성 증기 감지 예시"),
        ("GD-201", "염소 저장실", "염소", "고정식", "0.5 ppm", "중앙제어실·현장 경광등", "예", "GA-201", "독성가스 감지 예시"),
        ("PD-301", "정비팀 비치", "복합가스", "휴대식", "사업장 기준", "휴대기기", "해당 없음", "", "휴대용 감지기 예시"),
    ),
    "10_동력기계": (
        ("P-101", "원료 이송펌프", "원심펌프", "20 m3/h", "7.5 kW", "SUS304", "톨루엔", "원료저장", "PID-101", "펌프 예시"),
        ("C-201", "공정 압축기", "왕복동압축기", "500 Nm3/h", "55 kW", "Carbon Steel", "공정가스", "압축", "PID-202", "압축기 예시"),
        ("AG-301", "혼합조 교반기", "교반기", "120 rpm", "15 kW", "SUS316L", "혼합액", "혼합", "PID-301", "교반기 예시"),
    ),
    "11_배관_개스킷": (
        ("PCL-150", "톨루엔", "SUS304", "50A", "0.49 MPa", "80 ℃", "PTFE", "PID-101", "유기용제 배관 예시"),
        ("PCL-300", "황산", "PTFE Lined CS", "40A", "0.6 MPa", "60 ℃", "PTFE", "PID-302", "부식성 유체 배관 예시"),
        ("PCL-600", "공정가스", "Carbon Steel", "80A", "1.2 MPa", "120 ℃", "Spiral Wound", "PID-601", "가스 배관 예시"),
    ),
    "20_배출물질_처리시설": (
        ("SC-101", "유기용제 스크러버", "톨루엔 증기", "흡수", "1500 m3/h", "PSV-101", "대기배출구", "흡수식 처리 예시"),
        ("AC-201", "활성탄 흡착기", "VOC", "흡착", "2000 m3/h", "공정벤트", "대기배출구", "흡착식 처리 예시"),
        ("NT-301", "산·알칼리 중화조", "산성 폐수", "중화", "10 m3/h", "폐수배관", "폐수처리시설", "수계 처리 예시"),
    ),
}

EXAMPLE_VARIANTS: dict[str, tuple[str, str, str]] = {
    "process.description": (
        "[저장·이송형] 원료는 저장탱크에 입고한 후 이송펌프를 통해 생산설비로 공급하고, 사용 후 잔량은 회수탱크로 이송한다.",
        "[반응공정형] 원료 A와 B를 계량 투입한 뒤 반응기에서 설정 온도·압력으로 반응시키고, 냉각·분리 후 제품저장탱크로 이송한다.",
        "[혼합·충전형] 원료를 혼합조에 순차 투입하여 정해진 시간 동안 혼합한 뒤 품질 확인 후 제품탱크로 이송하고 용기에 충전한다.",
    ),
    "psm.operation.work_permit": (
        "화기작업은 작업허가서 발행, 가스농도 측정, 화재감시자 배치 후 실시한다.",
        "밀폐공간작업은 산소·유해가스 측정, 감시인 배치, 구조장비 준비 후 작업허가 절차에 따라 수행한다.",
        "굴착·고소 등 위험작업은 작업유형별 허가 절차와 현장 안전조치를 확인한 뒤 수행한다.",
    ),
    "psm.operation.moc": (
        "설비 사양 변경 시 공정·기계·전기 담당자가 영향성을 검토하고 승인 후 도면과 절차서를 개정한다.",
        "원료 변경 시 물성·반응성·취급조건 변화를 검토하고 교육 후 변경사항을 적용한다.",
        "운전조건 변경 시 위험성 검토와 관련 인터록·경보 설정값 검토 후 변경이력을 관리한다.",
    ),
    "psm.emergency.contacts": (
        "사고 발견자는 중앙제어실에 즉시 신고하고, 중앙제어실은 비상방송과 비상연락망으로 관련 부서에 상황을 전파한다.",
        "야간에는 당직자가 비상연락망에 따라 공장장·안전담당자·설비담당자에게 순차 연락한다.",
        "외부 지원이 필요한 경우 지정 담당자가 소방서·관계기관에 신고하고 사고물질과 현장상황을 전달한다.",
    ),
    "cap.prevention.safety_policy": (
        "유해화학물질 취급 전 위험요인을 확인하고 예방조치를 우선 적용하여 사고 가능성을 최소화한다.",
        "설비 건전성, 작업자 교육, 변경관리를 핵심 관리항목으로 정하고 정기적으로 이행상태를 점검한다.",
        "사고 발생 시 인명보호와 외부영향 최소화를 최우선 원칙으로 하여 비상대응체계를 운영한다.",
    ),
    "cap.internal.shutdown_procedure": (
        "누출 확인 시 원료공급 차단 → 관련 펌프 정지 → 긴급차단밸브 폐쇄 → 중앙제어실 보고 순으로 조치한다.",
        "반응 이상 시 투입 정지 → 냉각 최대화 → 비상정지 절차 수행 → 현장 접근통제 순으로 조치한다.",
        "가스감지기 고농도 경보 시 해당 구역 출입을 통제하고 원격 차단이 가능한 설비부터 정지한다.",
    ),
}


def _write_option_sheet(wb) -> dict[str, str]:
    if OPTIONS_SHEET in wb.sheetnames:
        del wb[OPTIONS_SHEET]
    ws = wb.create_sheet(OPTIONS_SHEET)
    formulas: dict[str, str] = {}
    for col_idx, (group, values) in enumerate(OPTION_GROUPS.items(), start=1):
        ws.cell(1, col_idx, group)
        for row_idx, value in enumerate(values, start=2):
            ws.cell(row_idx, col_idx, value)
        col_letter = ws.cell(1, col_idx).column_letter
        formulas[group] = f"'{OPTIONS_SHEET}'!${col_letter}$2:${col_letter}${len(values)+1}"
    ws.sheet_state = "hidden"
    return formulas


def _header_map(ws, row: int = 4) -> dict[str, int]:
    return {str(cell.value or "").strip(): cell.column for cell in ws[row] if str(cell.value or "").strip()}


def _add_list_validation(ws, cell_range: str, formula: str, *, prompt: str) -> None:
    dv = DataValidation(type="list", formula1=formula, allow_blank=True)
    dv.showErrorMessage = False
    dv.showInputMessage = True
    dv.promptTitle = "선택 또는 직접입력"
    dv.prompt = prompt
    ws.add_data_validation(dv)
    dv.add(cell_range)


def _apply_table_dropdowns(wb, formulas: dict[str, str]) -> None:
    for sheet_name, mapping in TABLE_DROPDOWNS.items():
        if sheet_name not in wb.sheetnames:
            continue
        ws = wb[sheet_name]
        headers = _header_map(ws)
        for header, group in mapping.items():
            col = headers.get(header)
            if not col:
                continue
            letter = ws.cell(4, col).column_letter
            _add_list_validation(
                ws,
                f"{letter}5:{letter}200",
                formulas[group],
                prompt="목록에서 선택하거나, 목록에 없는 사업장 고유값이면 직접 입력하세요.",
            )

    if "04_안전밸브_파열판" in wb.sheetnames and "03_설비정보" in wb.sheetnames:
        ws = wb["04_안전밸브_파열판"]
        headers = _header_map(ws)
        col = headers.get("보호대상 설비번호")
        if col:
            letter = ws.cell(4, col).column_letter
            _add_list_validation(
                ws,
                f"{letter}5:{letter}200",
                '=INDIRECT("\'03_설비정보\'!$A$5:$A$200")',
                prompt="03_설비정보에 등록한 설비번호(Tag No.)를 선택하세요. 필요한 경우 직접 입력도 가능합니다.",
            )


def _meta_records(wb) -> list[dict[str, Any]]:
    if "_시스템정보" not in wb.sheetnames:
        return []
    ws = wb["_시스템정보"]
    records: list[dict[str, Any]] = []
    for row in ws.iter_rows(min_row=8, values_only=True):
        if not any(v not in (None, "") for v in row):
            continue
        records.append(
            {
                "kind": str(row[0] or ""),
                "sheet": str(row[1] or ""),
                "field_keys": str(row[2] or ""),
                "row": int(row[3]) if row[3] not in (None, "") else 0,
                "col": int(row[4]) if row[4] not in (None, "") else 0,
            }
        )
    return records


def _apply_form_dropdowns(wb, formulas: dict[str, str]) -> None:
    for rec in _meta_records(wb):
        if rec["kind"] != "FORM":
            continue
        field_key = rec["field_keys"].split("|")[0]
        group = FIELD_DROPDOWNS.get(field_key)
        if not group or rec["sheet"] not in wb.sheetnames or not rec["row"] or not rec["col"]:
            continue
        ws = wb[rec["sheet"]]
        coord = ws.cell(rec["row"], rec["col"]).coordinate
        _add_list_validation(
            ws,
            coord,
            formulas[group],
            prompt="정형화 가능한 항목입니다. 목록에서 선택하거나 목록에 없으면 직접 입력하세요.",
        )


def _annotate_identifiers(wb) -> None:
    notes = {
        ("03_설비정보", "설비번호"): "사업장 내 설비를 다른 시트·도면과 연결하기 위한 고유 식별자(Tag No.)입니다. 기존 사업장 Tag를 우선 사용하고, 없으면 중복되지 않는 내부 식별자를 부여하세요.",
        ("04_안전밸브_파열판", "보호대상 설비번호"): "03_설비정보의 설비번호(Tag No.)와 연결합니다. 같은 설비를 다른 표에서 다른 번호로 다시 만들지 마세요.",
        ("05_가스누출감지_경보장치", "감지기 번호"): "감지기별 위치·설정값·도면을 구분하기 위한 식별자입니다. 사업장 관리번호가 있으면 그대로 사용하세요.",
        ("10_동력기계", "기계번호"): "동력기계를 P&ID 및 설비목록과 교차확인하기 위한 식별자입니다.",
        ("20_배출물질_처리시설", "시설번호"): "처리시설별 연결설비·배출지점을 구분하기 위한 식별자입니다. 법정 번호를 새로 만드는 의미는 아닙니다.",
    }
    for (sheet_name, header), text in notes.items():
        if sheet_name not in wb.sheetnames:
            continue
        ws = wb[sheet_name]
        headers = _header_map(ws)
        col = headers.get(header)
        if not col:
            continue
        cell = ws.cell(4, col)
        cell.comment = Comment(text, "PSM/CAP 작성지원")


def _input_hint(field_key: str) -> str:
    return INPUT_HINTS.get(field_key, "사업장 실제 기준으로 작성. 확인되지 않은 내용은 임의로 채우지 말고 빈칸 유지")


def _simplify_input_examples(wb) -> None:
    records = _meta_records(wb)
    for rec in records:
        if rec["kind"] not in {"FORM", "PROTECTED"} or rec["sheet"] not in wb.sheetnames:
            continue
        ws = wb[rec["sheet"]]
        if ws.cell(4, 3).value == "작성 예":
            ws.cell(4, 3, "입력 도움말")
        field_key = rec["field_keys"].split("|")[0]
        if rec["kind"] == "PROTECTED":
            hint = "Stage 1 판정자료에서 자동 입력된 값입니다. 변경이 필요하면 판정진단을 다시 수행하세요."
        else:
            hint = _input_hint(field_key)
        ws.cell(rec["row"], 3, hint)
        ws.cell(rec["row"], 3).fill = PatternFill("solid", fgColor=NOTE_FILL)
        ws.cell(rec["row"], 3).alignment = Alignment(vertical="top", wrap_text=True)


def _generic_variants(question: str) -> tuple[str, str, str]:
    return (
        f"[단순 사업장 예] {question}에 해당하는 실제 담당자·방법·주기를 간단히 작성",
        f"[복수 공정 사업장 예] 공정별 차이가 있으면 공정명과 담당부서를 구분하여 {question} 관련 내용을 작성",
        f"[해당 없음 예] 사업장에 해당하지 않는 경우 '해당 없음'으로 적고 적용되지 않는 이유를 간단히 작성",
    )


def _enrich_example_forms(wb) -> None:
    records = _meta_records(wb)
    for rec in records:
        if rec["kind"] != "FORM" or rec["sheet"] not in wb.sheetnames:
            continue
        ws = wb[rec["sheet"]]
        field_key = rec["field_keys"].split("|")[0]
        row = rec["row"]
        question = str(ws.cell(row, 1).value or "작성항목")
        variants = EXAMPLE_VARIANTS.get(field_key, _generic_variants(question))
        ws.cell(4, 1, "확인할 내용")
        ws.cell(4, 2, "예시 A")
        ws.cell(4, 3, "예시 B")
        ws.cell(4, 4, "예시 C")
        ws.cell(4, 5, "작성 포인트")
        for idx, value in enumerate(variants, start=2):
            ws.cell(row, idx, value)
            ws.cell(row, idx).fill = PatternFill("solid", fgColor=EXAMPLE_FILL)
            ws.cell(row, idx).alignment = Alignment(vertical="top", wrap_text=True)
        ws.cell(row, 5, "예시는 표현방식만 참고하세요. 실제 사업장 사실·설비·조직·수치와 다르면 복사하지 마십시오.")
        ws.cell(row, 5).fill = PatternFill("solid", fgColor=NOTE_FILL)
        ws.cell(row, 5).alignment = Alignment(vertical="top", wrap_text=True)
        ws.column_dimensions["B"].width = 48
        ws.column_dimensions["C"].width = 48
        ws.column_dimensions["D"].width = 48
        ws.column_dimensions["E"].width = 36


def _replace_example_table_rows(wb) -> None:
    for sheet_name, rows in TABLE_EXAMPLES.items():
        if sheet_name not in wb.sheetnames:
            continue
        ws = wb[sheet_name]
        headers = _header_map(ws)
        if not headers:
            continue
        max_col = max(headers.values())
        for row_no in range(5, 35):
            for col_no in range(1, max_col + 1):
                ws.cell(row_no, col_no, None)
        for row_no, values in enumerate(rows, start=5):
            for col_no, value in enumerate(values, start=1):
                ws.cell(row_no, col_no, value)
                ws.cell(row_no, col_no).fill = PatternFill("solid", fgColor=EXAMPLE_FILL)
                ws.cell(row_no, col_no).alignment = Alignment(vertical="top", wrap_text=True)
        ws.cell(2, 1, "여러 사업장 유형을 가정한 작성예시입니다. 실제 사업장 값으로 복사하지 말고 형식과 입력 수준만 참고하세요.")


def _add_example_guide_note(wb) -> None:
    if "00_작성안내" not in wb.sheetnames:
        return
    ws = wb["00_작성안내"]
    target_row = ws.max_row + 2
    ws.cell(target_row, 1, "작성예시 활용방법")
    ws.cell(target_row, 2, "이 파일은 한 회사의 모범답안이 아니라 저장·이송형, 반응공정형, 혼합·충전형 등 여러 상황을 비교해 작성방식을 이해하기 위한 참고자료입니다.")
    ws.cell(target_row + 1, 1, "주의")
    ws.cell(target_row + 1, 2, "예시 문구·수치·Tag No.를 실제 사업장 확인 없이 그대로 복사하지 마십시오.")
    for row in (target_row, target_row + 1):
        ws.cell(row, 1).fill = PatternFill("solid", fgColor=NOTE_FILL)
        ws.cell(row, 2).fill = PatternFill("solid", fgColor=NOTE_FILL)
        ws.cell(row, 1).alignment = Alignment(vertical="top", wrap_text=True)
        ws.cell(row, 2).alignment = Alignment(vertical="top", wrap_text=True)


def build_enhanced_integrated_authoring_workbook(project: Stage2Project, *, example: bool = False) -> bytes:
    raw = _base_build(project, example=example)
    wb = load_workbook(BytesIO(raw), data_only=False)
    formulas = _write_option_sheet(wb)
    _apply_table_dropdowns(wb, formulas)
    _apply_form_dropdowns(wb, formulas)
    _annotate_identifiers(wb)

    if example:
        _replace_example_table_rows(wb)
        _enrich_example_forms(wb)
        _add_example_guide_note(wb)
    else:
        _simplify_input_examples(wb)

    out = BytesIO()
    wb.save(out)
    return out.getvalue()
