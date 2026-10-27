"""회사 자료를 사업장 공통 사실로 바꾸고, 사람이 볼 예외만 남깁니다.

흐름: 자료 읽기 → 사실 후보(출처 포함) → 규칙으로 자동 기입 → 예외(충돌·오류·낮은 신뢰)만 사람에게.

자동 기입 경계는 두 가지로 정합니다.
1. 값이 유효하고 모호하지 않은가 (충돌·단위·범위·CAS 검산·추출 경로).
2. 값이 결과에 미치는 영향이 큰가. 법적 판단에 쓰이는 항목이나 예비 순위를 바꾸는 핵심값은
   더 높은 출처 신뢰도를 요구하고, 자동 기입했더라도 먼저 확인할 목록에 올립니다.
"""
from __future__ import annotations

from collections import OrderedDict
from copy import deepcopy
from hashlib import sha256
from io import BytesIO, StringIO
import csv
import json
import math
import re
from zipfile import BadZipFile, ZIP_DEFLATED, ZipFile
from xml.etree import ElementTree

from . import cap_fact_catalog as catalog
from .project import CONFIRMED_STATUSES, EvidenceRef, Stage2Project

KEY = "workspace.evidence_graph"
LABELS = {"name": "명칭", "product": "취급제품", "cas": "CAS 번호", "capacity": "용량(m³)",
          "pressure": "운전압력(MPa)", "design_pressure": "설계압력(MPa)", "temperature": "운전온도(°C)",
          "inventory": "최대보유량(kg)", "concentration": "함량(%)",
          "facility_type": "시설유형", "substance_state": "물질성상",
          "density": "비중(kg/L)", "process_type": "공정유형"}
# 별지 제1호 계산이 받는 값과 같아야 합니다(tests/test_evidence_workspace.py가 cap_workspace와 일치를 검사).
CHOICES = {"facility_type": ("저장탱크", "제조·사용시설", "보관시설", "기타", "탱크로리·운송차량", "사외배관", "취급중단 신고시설"),
           "substance_state": ("액체", "고체", "기체·고압가스", "복수성상"),
           "process_type": ("변화없음", "단순혼합", "반응")}
CHOICE_SYNONYMS = {"기체": "기체·고압가스", "고압가스": "기체·고압가스", "기체/고압가스": "기체·고압가스", "기체고압가스": "기체·고압가스",
                   "탱크로리": "탱크로리·운송차량", "운송차량": "탱크로리·운송차량", "제조사용시설": "제조·사용시설",
                   "제조시설": "제조·사용시설", "사용시설": "제조·사용시설"}
ALIASES = {"명칭": "name", "물질명": "name", "제품명": "name", "설비명": "name", "시설명": "name",
           "화학물질명": "name", "물질명칭": "name", "설비명칭": "name", "품명": "name",
           "취급제품": "product", "취급물질": "product", "취급제품명": "product",
           "cas": "cas", "casno": "cas", "cas번호": "cas",
           "용량": "capacity", "용량m3": "capacity", "용량l": "capacity", "용량m³": "capacity", "용량㎥": "capacity",
           "운전압력mpa": "pressure", "운전압력kpa": "pressure", "운전압력": "pressure",
           "설계압력mpa": "design_pressure", "설계압력kpa": "design_pressure",
           "운전온도": "temperature", "운전온도c": "temperature", "운전온도°c": "temperature", "운전온도℃": "temperature",
           "최대보유량kg": "inventory", "최대보유량ton": "inventory", "최대보유량": "inventory",
           "함량": "concentration", "함량%": "concentration",
           "시설유형": "facility_type", "설비유형": "facility_type", "시설구분": "facility_type",
           "물질성상": "substance_state", "물질상태": "substance_state", "성상": "substance_state",
           "비중": "density", "밀도": "density", "비중kg/l": "density", "밀도kg/l": "density", "밀도g/cm3": "density",
           "밀도ton/m3": "density", "비중또는밀도": "density", "비중또는밀도kg/l=ton/m3": "density",
           "공정유형": "process_type", "공정구분": "process_type"}
ID_HEADERS = {"설비번호", "시설번호", "설비id", "제품코드", "물질코드", "설비코드", "장치번호", "tagno", "tag번호"}

# ---------------------------------------------------------------- 사실 카탈로그에서 더한 항목
# 위 기본 항목(설비·물질의 법적 계산값) 밖의 회사 사실 — 사업장 일반정보, 공정 설명, 설계온도·연결구, SDS 물성·유해성,
# 확산방지설비·가스감지기 명세 — 은 data/stage2/cap_fact_catalog.json의 머리글 표기로 알아보고 적힌 글자 그대로 둡니다.
# 단위·형식 확인은 그 값을 쓰는 서식 계산 엔진이 다시 합니다.
CORE_FIELDS = frozenset(LABELS)
for _field, _label in catalog.evidence_labels().items():
    LABELS.setdefault(_field, _label)
TEXT_FIELDS = frozenset(LABELS) - CORE_FIELDS
TEXT_OPTIONS = {f.evidence_field: f.options for f in catalog.all_facts()
                if f.evidence_field in TEXT_FIELDS and f.format == "choice" and f.options}
CATALOG_HEADERS = {kind: catalog.header_map(kind) for kind in ("facility", "material", "detector", "site")}
DETECTOR_ID_HEADERS = frozenset(catalog.identifier_headers("detector"))
ID_HEADERS |= set(catalog.identifier_headers("facility")) | set(catalog.identifier_headers("material"))
ENTITY_ID_HEADERS = frozenset(ID_HEADERS)            # 설비·물질 표의 번호 열
ID_HEADERS |= DETECTOR_ID_HEADERS
IGNORED_HEADERS = catalog.ignored_headers()
SITE_ENTITY = "site:사업장"
KINDS = ("facility", "material", "site", "detector")
LABEL_CONFIDENCE = 0.95      # 사업장 자료의 ‘항목: 값’ 문장(사업장은 하나뿐이라 대상이 모호하지 않음)
# 사람 확인 없이 쓸 수 있는 읽기 방식: 표의 칸(table), 사업장 자료의 항목·값(label). 지정 번호 문서의 문장(literal)과 AI 추출은 제외.
READ_METHODS = frozenset({"table", "label"})

NUMERIC_FIELDS = {"capacity", "pressure", "design_pressure", "temperature", "inventory", "concentration", "density"}
BOUNDS = {"capacity": (0, 1e6), "pressure": (0, 100), "design_pressure": (0, 100),
          "temperature": (-273.15, 2000), "inventory": (0, 1e9), "concentration": (0, 100), "density": (0, 30)}
UNIT_SUFFIXES = {"capacity": ("m3", "m³", "㎥", "l"), "pressure": ("mpa", "kpa"), "design_pressure": ("mpa", "kpa"),
                 "inventory": ("kg", "ton"), "temperature": ("c", "°c", "℃"), "concentration": ("%",)}
# 법적 판정(규정수량·방류벽 용량 등)에 직접 쓰이는 항목은 순위와 관계없이 엄격하게 봅니다.
LEGAL_FIELDS = {"inventory", "concentration", "cas", "capacity", "facility_type", "substance_state", "density", "process_type"}
STRICT_CONFIDENCE = 0.98
HINT_CONFIDENCE = 0.88       # 번호 없는 문서의 표: 사람이 지정한 번호로 연결 (문장과 같은 수준)
NUMBER = re.compile(r"[+-]?\d+(?:,\d{3})*(?:\.\d+)?")
RELAXED_CONFIDENCE = 0.90
HEADER_SCAN_ROWS = 20
MAX_ROWS = 10001
MAX_COLUMNS = 200
MAX_UNZIPPED = 80 * 1024 * 1024
DECISION_LABEL = {"AUTO": "자동 기입", "PERSON": "확인 완료", "BULK": "확인 완료"}


def norm(value):
    return re.sub(r"[\s().·_-]", "", str(value)).lower()


# ---------------------------------------------------------------- 저장소

def _read(project):
    """읽기 전용 조회(복사하지 않음). 수정하려면 graph()를 쓰세요."""
    record = project.get_field(KEY)
    return record.value if record else {"documents": [], "facts": [], "decisions": {}}


def graph(project):
    return deepcopy(_read(project))


def _store(project, data):
    data["version"] = data.get("version", 0) + 1
    project.set_field(KEY, "사업장 공통정보와 출처", data, "HOLD", note="추출 후보와 담당자 결정을 분리해 보존")


def _memo(project, name, compute):
    """같은 자료(저장 버전)로 한 번 계산한 결과를 다시 씁니다. 화면 한 번에 같은 계산을 여러 번 하지 않게 합니다.
    보관 위치는 프로젝트 객체 자체라 다른 프로젝트와 섞이지 않고, 저장하면 버전이 바뀌어 자동으로 새로 계산됩니다."""
    record = project.get_field(KEY)
    token = (id(record.value), record.value.get("version")) if record else None
    cache = project.__dict__.setdefault("_evidence_memo", {})
    hit = cache.get(name)
    if hit is not None and token is not None and hit[0] == token:
        return hit[1]
    value = compute()
    cache[name] = (token, value)
    return value


def documents(project):
    """읽기 전용 문서 목록(복사하지 않음)."""
    return _read(project)["documents"]


def register_reference(project, name, content):
    """서식 근거자료를 원본과 연결한다. 표준 사실로 추출하지 않았음을 명시한다."""
    digest = sha256(content).hexdigest()
    data = graph(project)
    if any(d["sha256"] == digest for d in data["documents"]):
        return False
    tables, passages = read_document(name, content)
    has_text = any(p.get("text", "").strip() for p in passages)
    problem = ("글자가 없는 스캔본입니다. OCR 후 원본을 확인해 다시 올려 주세요." if not tables and not has_text else
               "원본은 보관했습니다. 이 문서의 값은 아직 법정 서식에 자동 연결되지 않았습니다.")
    data["documents"].append({"name": name, "sha256": digest, "kind": "reference", "hint": "",
                              "passages": passages, "fact_count": 0, "unmapped_columns": [], "tables": [],
                              "rows_without_id": 0, "needs_reading": True,
                              "needs_ocr": not tables and not has_text, "problem": problem})
    _store(project, data)
    return True


def _invalidate(project):
    for record in project.fields.values():
        if record.note.startswith("공통정보 연결") or record.status == "AI_DRAFT":
            record.status = "HOLD"


# ---------------------------------------------------------------- 값 검증

def has_unit(field, header):
    h = norm(header)
    return any(h.endswith(suffix) for suffix in UNIT_SUFFIXES.get(field, ()))


def cas_valid(text):
    text = str(text)
    if not re.fullmatch(r"\d{2,7}-\d{2}-\d", text):
        return False
    digits = re.sub(r"\D", "", text)
    return sum(int(n) * (i + 1) for i, n in enumerate(digits[-2::-1])) % 10 == int(digits[-1])


def _text_value(field, text):
    """카탈로그 항목은 적힌 그대로 둡니다. 선택지가 있으면 띄어쓰기·기호만 다른 표기를 선택지로 맞춥니다."""
    squeezed = norm(text)
    for option in TEXT_OPTIONS.get(field, ()):
        if squeezed == norm(option):
            return option
    return text


def normalize_value(field, raw, header):
    text = str(raw).strip()
    if field in TEXT_FIELDS:
        return _text_value(field, text)
    if field in CHOICES:
        squeezed = re.sub(r"\s+", "", text)
        for option in CHOICES[field]:
            if squeezed == re.sub(r"\s+", "", option):
                return option
        return CHOICE_SYNONYMS.get(squeezed, text)
    if field in {"name", "product", "cas"}:
        return text
    # 범위나 단위가 불명확한 값을 임의로 숫자로 확정하지 않습니다.
    if not re.fullmatch(r"[+-]?\d+(?:,\d{3})*(?:\.\d+)?", text):
        return text
    number = float(text.replace(",", ""))
    if not math.isfinite(number):
        return text
    h = norm(header)
    if field in {"capacity", "pressure", "design_pressure", "inventory"} and h in {"용량", "운전압력", "설계압력", "최대보유량"}:
        return text
    if field == "capacity" and h.endswith("l"):
        number /= 1000
    if field in {"pressure", "design_pressure"} and h.endswith("kpa"):
        number /= 1000
    if field == "inventory" and h.endswith("ton"):
        number *= 1000
    return number


def value_problem(field, value):
    """승인할 수 없는 값이면 이유를, 괜찮으면 None을 돌려줍니다."""
    if field in {"name", "product"} or field in TEXT_FIELDS:
        return None if str(value).strip() else "빈 값"
    if field in CHOICES:
        return None if value in CHOICES[field] else "허용 값: " + ", ".join(CHOICES[field])
    if field == "cas":
        return None if cas_valid(value) else "CAS 형식 또는 검산값이 맞지 않음"
    if isinstance(value, bool) or not isinstance(value, (int, float)) or not math.isfinite(value):
        return "범위·단위·수식을 확인해야 하는 값"
    low, high = BOUNDS[field]
    if value < low or value > high:
        return f"허용 범위({low:g}~{high:g}) 밖의 값 — 단위를 확인하세요"
    return None


# ---------------------------------------------------------------- 문서 읽기

class Table:
    """문서에서 읽은 표. flags는 (행, 열) → 'merged'|'formula'|'formula_uncached' (0부터)."""

    def __init__(self, location, rows, flags=None):
        self.location = location
        self.rows = rows
        self.flags = flags or {}


def _cell_text(value):
    if value is None:
        return ""
    if isinstance(value, bool):
        return str(value)
    if isinstance(value, float):
        if value == int(value) and abs(value) < 1e15:
            return str(int(value))
        text = repr(value)
        return f"{value:f}".rstrip("0").rstrip(".") if "e" in text.lower() else text
    return str(value)


def _read_csv(content):
    try:
        text = content.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = content.decode("cp949")
    rows = list(csv.reader(StringIO(text)))
    return [Table("CSV", rows)] if rows else []


def _read_xlsx(content):
    from openpyxl import load_workbook
    try:
        values = load_workbook(BytesIO(content), data_only=True)
        formulas = load_workbook(BytesIO(content), data_only=False)
    except (BadZipFile, KeyError, ValueError) as exc:
        raise ValueError("엑셀 파일을 열 수 없습니다.") from exc
    tables = []
    try:
        for sheet in values.worksheets:
            if (sheet.max_row or 0) > MAX_ROWS:
                raise ValueError("한 시트에 10,000행 이하로 나눠 주세요.")
            fsheet = formulas[sheet.title]
            max_row, max_col = sheet.max_row or 0, min(sheet.max_column or 0, MAX_COLUMNS)
            if not max_row or not max_col:
                continue
            grid = [[_cell_text(v) for v in row] for row in
                    sheet.iter_rows(min_row=1, max_row=max_row, max_col=max_col, values_only=True)]
            flags = {}
            for r, row in enumerate(fsheet.iter_rows(min_row=1, max_row=max_row, max_col=max_col, values_only=True)):
                for c, raw in enumerate(row):
                    if isinstance(raw, str) and raw.startswith("="):
                        if grid[r][c] == "":
                            grid[r][c] = raw
                            flags[(r, c)] = "formula_uncached"
                        else:
                            flags[(r, c)] = "formula"
            for span in sheet.merged_cells.ranges:
                if (span.max_row - span.min_row + 1) * (span.max_col - span.min_col + 1) > 10000:
                    continue
                top = grid[span.min_row - 1][span.min_col - 1] if span.min_row <= max_row and span.min_col <= max_col else ""
                for r in range(span.min_row - 1, min(span.max_row, max_row)):
                    for c in range(span.min_col - 1, min(span.max_col, max_col)):
                        if (r, c) != (span.min_row - 1, span.min_col - 1):
                            grid[r][c] = top
                            flags[(r, c)] = "merged"
            while grid and not any(grid[-1]):
                grid.pop()
            label = sheet.title + ("" if sheet.sheet_state == "visible" else " (숨김 시트)")
            if grid:
                tables.append(Table(label, grid, flags))
    finally:
        values.close()
        formulas.close()
    return tables


def _read_docx(content):
    from docx import Document
    doc = Document(BytesIO(content))
    passages = [{"location": f"문단 {i}", "text": p.text} for i, p in enumerate(doc.paragraphs, 1) if p.text.strip()]
    tables = []
    for i, table in enumerate(doc.tables, 1):
        try:
            tables.append(Table(f"표 {i}", [[c.text for c in row.cells] for row in table.rows]))
        except Exception:
            continue
    return tables, passages


def _read_pdf(content):
    import pdfplumber
    tables, passages = [], []
    with pdfplumber.open(BytesIO(content)) as pdf:
        for i, page in enumerate(pdf.pages, 1):
            passages.append({"location": f"p.{i}", "text": page.extract_text() or ""})
            try:
                found = page.extract_tables()
            except Exception:
                found = []
            for j, table in enumerate(found, 1):
                tables.append(Table(f"p.{i} 표 {j}", [[("" if c is None else str(c)) for c in row] for row in table]))
    return tables, passages


def _local(tag):
    return tag.rsplit("}", 1)[-1]


def _read_hwpx(content):
    """HWPX(zip+XML)의 문단과 표를 읽습니다."""
    try:
        archive = ZipFile(BytesIO(content))
    except BadZipFile as exc:
        raise ValueError("HWPX 파일을 열 수 없습니다. 구형 HWP는 HWPX 또는 PDF로 저장해 올려 주세요.") from exc
    with archive:
        if sum(info.file_size for info in archive.infolist()) > MAX_UNZIPPED:
            raise ValueError("압축을 푼 크기가 너무 큽니다.")
        names = sorted(n for n in archive.namelist() if re.fullmatch(r"Contents/section\d+\.xml", n))
        if not names:
            raise ValueError("HWPX 본문을 찾지 못했습니다.")
        tables, passages = [], []

        def text_of(node):
            return "".join(t.text or "" for t in node.iter() if _local(t.tag) == "t")

        def cell_text(cell):
            # 셀 안의 문단은 공백으로 이어 붙여 단어가 붙지 않게 합니다.
            parts = [text_of(p).strip() for p in cell.iter() if _local(p.tag) == "p"]
            return " ".join(part for part in parts if part) or text_of(cell).strip()

        def walk(node, section, state):
            for child in node:
                name = _local(child.tag)
                if name == "tbl":
                    rows = []
                    for tr in (n for n in child.iter() if _local(n.tag) == "tr"):
                        rows.append([cell_text(tc) for tc in tr if _local(tc.tag) == "tc"])
                    state["tables"] += 1
                    if rows:
                        tables.append(Table(f"{section} 표 {state['tables']}", rows))
                elif name == "p":
                    inner_tables = any(_local(n.tag) == "tbl" for n in child.iter())
                    if inner_tables:
                        walk(child, section, state)
                    else:
                        state["paragraphs"] += 1
                        text = text_of(child).strip()
                        if text:
                            passages.append({"location": f"{section} 문단 {state['paragraphs']}", "text": text})
                else:
                    walk(child, section, state)

        for index, name in enumerate(names, 1):
            raw = archive.read(name)
            if b"<!DOCTYPE" in raw[:2048].upper() or b"<!ENTITY" in raw.upper():
                raise ValueError("허용하지 않는 XML 구조입니다.")
            walk(ElementTree.fromstring(raw), f"구역 {index}", {"tables": 0, "paragraphs": 0})
    return tables, passages


def read_document(name, content):
    suffix = name.rsplit(".", 1)[-1].lower()
    if len(content) > 20 * 1024 * 1024:
        raise ValueError("파일은 20MB 이하로 올려 주세요.")
    tables, passages = [], []
    if suffix == "csv":
        tables = _read_csv(content)
    elif suffix == "xlsx":
        tables = _read_xlsx(content)
    elif suffix == "docx":
        tables, passages = _read_docx(content)
    elif suffix == "pdf":
        tables, passages = _read_pdf(content)
    elif suffix == "hwpx":
        tables, passages = _read_hwpx(content)
    elif suffix == "hwp":
        from .hwp5_binary import to_hwpx
        tables, passages = _read_hwpx(to_hwpx(content))
    elif suffix == "txt":
        text = content.decode("utf-8-sig", errors="replace")
        passages = [{"location": f"줄 {i}", "text": line} for i, line in enumerate(text.splitlines(), 1)]
    else:
        raise ValueError("CSV·XLSX·DOCX·HWP 5.0·HWPX·텍스트 PDF·TXT를 지원합니다. 다른 형식은 변환 후 가져오세요.")
    return tables, passages


def _is_known_header(cell):
    key = norm(cell or "")
    return (key in ALIASES or key in ID_HEADERS or key in CATALOG_HEADERS["facility"]
            or key in CATALOG_HEADERS["material"] or key in CATALOG_HEADERS["detector"])


def find_header_row(rows):
    """제목·안내 문구가 위에 있어도 항목명이 가장 많이 모인 줄을 머리글로 고릅니다."""
    best, best_score = None, 0
    for index, row in enumerate(rows[:HEADER_SCAN_ROWS]):
        score = sum(1 for cell in row if _is_known_header(cell))
        if score > best_score:
            best, best_score = index, score
    return best


def _header_score(row):
    return sum(1 for cell in row if _is_known_header(cell))


FACILITY_ID_HEADERS = {"설비번호", "시설번호", "설비id", "설비코드", "장치번호", "tagno", "tag번호"} | set(catalog.identifier_headers("facility"))
MATERIAL_ID_HEADERS = {"제품코드", "물질코드"}
# 단위 표기가 붙어도 인정하도록 접두어로 비교합니다 (예: 용량(m3) → 용량m3).
FACILITY_HINTS = ("설비명", "시설명", "용량", "최대보유량", "운전압력", "운전온도", "설계압력", "시설유형", "설비유형", "공정유형")
MATERIAL_HINTS = ("cas", "제품명", "물질명", "함량", "성상", "비중", "인화점")


def _classify_table(headings):
    """표 하나의 종류. 식별자 열이 한쪽에만 있으면 그쪽으로 정합니다.
    설비대장에 물질명·CAS 열이 함께 있는 것은 흔하므로 항목명은 보조 근거로만 씁니다.
    감지기 번호 열만 있는 표는 가스감지기 명세입니다."""
    facility_id, material_id = headings & FACILITY_ID_HEADERS, headings & MATERIAL_ID_HEADERS
    if facility_id and material_id:
        return "both"          # 설비-제품 대응표처럼 두 대상을 잇는 표: 사람에게 묻습니다
    if facility_id:
        return "facility"
    if material_id:
        return "material"
    if headings & DETECTOR_ID_HEADERS:
        return "detector"
    facility = sum(any(h.startswith(p) for p in FACILITY_HINTS) for h in headings)
    material = sum(any(h.startswith(p) for p in MATERIAL_HINTS) for h in headings)
    if facility >= 2 and facility >= 2 * material:
        return "facility"
    if material >= 2 and material >= 2 * facility:
        return "material"
    return None


def _compact_cells(row):
    """(열 번호, 글자). 병합 셀처럼 바로 옆에 같은 글자가 반복되면 한 번만 둡니다."""
    out = []
    for column, cell in enumerate(row):
        text = str(cell or "").strip()
        if out and out[-1][1] == text:
            continue
        out.append((column, text))
    return out


def site_pairs(rows):
    """사업장 일반정보·공정 설명처럼 ‘항목 | 값’으로 적은 표에서 (행, 열, 항목 표기, 값)을 고릅니다.

    사실 카탈로그의 사업장 항목 이름과 같은 칸만 항목으로 봅니다. 한 줄에 항목·값이 여러 쌍 있어도 읽고,
    항목 이름이 한 줄에 모이고 그 아래 줄에 값이 있는 가로 표도 읽습니다."""
    site = CATALOG_HEADERS["site"]
    pairs = []
    for r, row in enumerate(rows):
        cells = _compact_cells(row)
        i = 0
        while i < len(cells) - 1:
            column, text = cells[i]
            next_column, value = cells[i + 1]
            if norm(text) in site and value and norm(value) not in site:
                pairs.append((r, next_column, text, value))
                i += 2
                continue
            i += 1
    if pairs:
        return pairs
    for r, row in enumerate(rows[:HEADER_SCAN_ROWS]):
        labels = [(c, str(cell).strip()) for c, cell in enumerate(row) if norm(cell or "") in site]
        if len(labels) >= 2 and r + 1 < len(rows):
            below = rows[r + 1]
            return [(r + 1, c, text, str(below[c]).strip()) for c, text in labels
                    if c < len(below) and str(below[c] or "").strip()]
    return []


def _is_site_table(rows, pairs):
    fields = {CATALOG_HEADERS["site"][norm(label)] for _, _, label, _ in pairs}
    if len(fields) >= 2:
        return True
    return len(fields) == 1 and len(rows) <= 2        # 공정개요 한 칸만 적은 작은 표


def _site_lines(passages):
    """문서 문장 중 ‘대표자: 홍길동’처럼 사업장 항목 이름으로 시작하는 줄."""
    site = CATALOG_HEADERS["site"]
    lines = []
    for passage in passages:
        for line in passage.get("text", "").splitlines():
            parts = re.split(r"[:：]", line, maxsplit=1)
            if len(parts) == 2 and norm(parts[0]) in site and parts[1].strip():
                lines.append((passage["location"], parts[0].strip(), parts[1].strip()))
    return lines


def _table_kinds(tables):
    """표마다 종류(facility·material·detector·site·both) 또는 None(머리글은 있으나 모름). 머리글이 없는 표는 빼고 셉니다."""
    kinds = []
    for table in tables:
        pairs = site_pairs(table.rows)
        at = find_header_row(table.rows)
        if pairs and _is_site_table(table.rows, pairs) and (at is None or len(pairs) >= _header_score(table.rows[at])):
            kinds.append("site")
            continue
        if at is None:
            continue
        kinds.append(_classify_table({norm(cell) for cell in table.rows[at] if str(cell).strip()}))
    return kinds


_CLASSIFY_CACHE: OrderedDict = OrderedDict()   # (내용 해시, 확장자) → 종류 또는 None
CLASSIFY_CACHE_LIMIT = 256


def _document_kind(tables, passages):
    kinds = _table_kinds(tables)
    if "both" in kinds or None in kinds:
        return None
    core = set(kinds) - {"site", "detector"}
    if len(core) == 1:
        return core.pop()
    if len(core) > 1:
        return "mixed"          # 시트마다 종류가 분명한 묶음 파일(업로드 양식 등): 표마다 따로 읽습니다
    if set(kinds) == {"detector"}:
        return "detector"
    if "site" in kinds or len({norm(label) for _, label, _ in _site_lines(passages)}) >= 2:
        return "site" if "detector" not in kinds else "mixed"
    return None


def classify_document(name, content):
    """식별자와 항목명이 가리키는 자료 종류를 결정합니다. 모호하면 확인을 요청합니다.

    종류: facility(설비)·material(물질)·site(사업장 일반정보·공정 설명)·detector(가스감지기)·
    mixed(시트마다 종류가 분명한 묶음 파일). 화면은 다시 그려질 때마다 이 함수를 부르므로,
    파일 내용 해시(+확장자)로 결과를 보관해 같은 파일을 다시 읽지 않습니다. 판별 실패도 보관합니다."""
    key = (sha256(content).hexdigest(), str(name).rsplit(".", 1)[-1].lower())
    if key in _CLASSIFY_CACHE:
        _CLASSIFY_CACHE.move_to_end(key)
        kind = _CLASSIFY_CACHE[key]
    else:
        tables, passages = read_document(name, content)
        kind = _document_kind(tables, passages)
        _CLASSIFY_CACHE[key] = kind
        if len(_CLASSIFY_CACHE) > CLASSIFY_CACHE_LIMIT:
            _CLASSIFY_CACHE.popitem(last=False)
    if kind:
        return kind
    raise ValueError("자료 종류를 자동 판별하지 못했습니다. 이 파일의 물질 또는 설비 종류를 선택해 주세요.")


# ---------------------------------------------------------------- 접수

def _rules(data):
    """담당자가 한 번 답한 것을 같은 경우 전체에 적용하는 규칙(이 사업장 자료에만 적용).

    values: 항목별 표기 → 값 (예: 물질성상 ‘액상’ → 액체, 운전압력 ‘상압’ → 0.1)
    units: 파일·열별 단위 (예: 03 파일 ‘운전압력’ 열 → MPa)
    columns: 머리글 → 항목 (예: ‘함량(wt%)’ → 함량)
    """
    rules = data.setdefault("rules", {})
    for key in ("values", "units", "columns"):
        rules.setdefault(key, {})
    return rules


def _is_number(raw):
    return bool(NUMBER.fullmatch(str(raw).strip()))


def _normalize(field, raw, header, rules, sha):
    """(값, 규칙을 썼는지). 규칙은 그대로 쓸 수 없는 값에만 적용합니다."""
    unit = rules["units"].get(f"{sha}|{norm(header)}")
    effective = f"{header}({unit})" if unit and not has_unit(field, header) else header
    value = normalize_value(field, raw, effective)
    used = effective != header
    if value_problem(field, value) and not _is_number(raw):
        learned = rules["values"].get(field, {}).get(norm(raw))
        if learned is not None:
            value, used = learned, True
    return value, used


def _reapply(data):
    """규칙이 바뀌면 이미 읽은 값에 다시 적용합니다. 원문(raw)은 그대로 두고 값만 다시 계산합니다."""
    rules = _rules(data)
    changed = 0
    for fact in data["facts"]:
        if fact.get("method") not in {"table", "literal", "label"}:
            continue
        value, used = _normalize(fact["field"], fact["raw"], fact["header"], rules, fact["sha256"])
        flags = [f for f in fact.get("flags", []) if f != "rule"] + (["rule"] if used else [])
        if value != fact["value"] or flags != fact.get("flags", []):
            fact["value"], fact["flags"] = value, flags
            changed += 1
    return changed


def _field_of(header, rules, kind="facility"):
    """머리글 → 읽을 항목. 설비·물질 표는 기본 항목(ALIASES)을 먼저, 그다음 사실 카탈로그 표기를 봅니다."""
    key = norm(header)
    if kind in ("facility", "material"):
        return ALIASES.get(key) or CATALOG_HEADERS[kind].get(key) or rules["columns"].get(key)
    return CATALOG_HEADERS.get(kind, {}).get(key)


def _table_kind(kind, own):
    """파일 종류(사람이 고르거나 자동 판별)와 표 자체의 종류로 이 표를 어떤 대상으로 읽을지 정합니다.
    감지기 표는 언제나 감지기로, 묶음 파일(mixed)과 사업장 파일은 표 자체의 종류를 따릅니다."""
    if own == "detector":
        return "detector"
    if kind in ("facility", "material"):
        return kind
    return own if own in ("facility", "material") else None


def _link_materials(data):
    """제품코드 없이 이름으로 읽은 물질 줄(SDS 물성표 등)을, CAS 번호가 같은 제품코드 물질 하나에 이어 붙입니다.
    같은 CAS를 가진 제품이 둘 이상이면(혼합물 등) 추측하지 않고 그대로 둡니다. 이어 붙인 수를 돌려줍니다."""
    coded, named = {}, {}
    for fact in data["facts"]:
        if fact.get("kind") != "material" or fact["field"] != "cas" or value_problem("cas", fact["value"]):
            continue
        cas = str(fact["value"]).strip()
        if "by_name" in fact.get("flags", []):
            named.setdefault(fact["entity"], set()).add(cas)
        else:
            coded.setdefault(cas, set()).add(fact["entity"])
    moves = {}
    for entity, numbers in named.items():
        if len(numbers) != 1:
            continue
        targets = coded.get(next(iter(numbers)), set()) - {entity}
        if len(targets) == 1:
            moves[entity] = next(iter(targets))
    for fact in data["facts"]:
        if fact["entity"] in moves and "by_name" in fact.get("flags", []):
            fact["entity"] = moves[fact["entity"]]
            fact["flags"] = [f for f in fact["flags"] if f != "by_name"] + ["cas_link"]
            fact["confidence"] = max(fact.get("confidence", 0), STRICT_CONFIDENCE)
    return len(moves)


def ingest(project, name, content, kind="facility", entity_hint=""):
    if kind not in {"facility", "material", "site", "detector", "mixed"}:
        raise ValueError("자료 종류를 확인하세요.")
    data = graph(project)
    rules = _rules(data)
    digest = sha256(content).hexdigest()
    entity_hint = entity_hint.strip()
    if any(d["sha256"] == digest and d["kind"] == kind and d.get("hint", "") == entity_hint for d in data["documents"]):
        return 0
    tables, passages = read_document(name, content)
    added, unmapped, table_info, skipped_total = [], [], [], 0
    seen_unmapped = set()

    def add(entity_kind, entity, field, raw, location, header, *, method="table", confidence=0.78, flags=()):
        if "merged" in flags or "formula" in flags:
            confidence = min(confidence, RELAXED_CONFIDENCE)
        identifier = sha256(json.dumps([digest, entity_kind, entity, field, location, str(raw)], ensure_ascii=False).encode()).hexdigest()[:24]
        value, used = _normalize(field, raw, header, rules, digest)
        added.append({"id": identifier, "entity": entity_kind + ":" + str(entity).strip(), "kind": entity_kind, "field": field,
                      "value": value, "raw": str(raw), "header": str(header),
                      "document": name, "sha256": digest, "location": location, "method": method,
                      "confidence": confidence, "flags": list(flags) + (["rule"] if used else [])})

    def unknown(header, *, everything=False):
        """읽지 못한 열을 알립니다. 표를 읽었으면 비고·연번 같은 일반 열은 빼고, 표를 통째로 못 읽었으면 모두 알립니다."""
        key = norm(header)
        if header and (everything or key not in IGNORED_HEADERS) and key not in seen_unmapped:
            seen_unmapped.add(key)
            unmapped.append(header)

    site_found = False
    for table in tables:
        rows = table.rows
        if not rows:
            continue
        if len(rows) > MAX_ROWS:
            raise ValueError("한 표에 10,000행 이하로 나눠 주세요.")
        header_at = find_header_row(rows)
        pairs = site_pairs(rows)
        if pairs and _is_site_table(rows, pairs) and (header_at is None or len(pairs) >= _header_score(rows[header_at])):
            site_found = True
            for r, c, label, value in pairs:
                add("site", SITE_ENTITY.split(":", 1)[1], CATALOG_HEADERS["site"][norm(label)], value,
                    f"{table.location} 행 {r + 1} 열 {c + 1}", label, method="label", confidence=STRICT_CONFIDENCE)
            table_info.append({"location": table.location, "header_row": None, "rows": len(rows), "kind": "site",
                               "columns": [{"header": label, "field": CATALOG_HEADERS["site"][norm(label)]} for _, _, label, _ in pairs],
                               "skipped": [], "skipped_count": 0})
            continue
        if header_at is None:
            for h in rows[0]:
                unknown(str(h or "").strip(), everything=True)
            table_info.append({"location": table.location, "header_row": None, "rows": len(rows), "columns": [],
                               "skipped": [], "skipped_count": 0,
                               "problem": f"앞 {HEADER_SCAN_ROWS}줄에서 설비번호·용량·CAS 같은 항목 이름이 있는 줄을 찾지 못함"})
            continue
        headers = [str(h or "").strip() for h in rows[header_at]]
        own = _classify_table({norm(h) for h in headers if h})
        table_kind = _table_kind(kind, own)
        if table_kind is None:
            for h in headers:
                unknown(h, everything=True)
            table_info.append({"location": table.location, "header_row": header_at + 1, "rows": len(rows) - header_at - 1,
                               "columns": [], "skipped": [], "skipped_count": 0,
                               "problem": "설비번호·제품코드·감지기 번호처럼 대상을 가리키는 열을 찾지 못함"})
            continue
        fields = [_field_of(h, rules, table_kind) for h in headers]
        id_headers = DETECTOR_ID_HEADERS if table_kind == "detector" else ENTITY_ID_HEADERS
        id_col = next((i for i, h in enumerate(headers) if norm(h) in id_headers), None)
        columns = [{"header": h, "field": "id" if i == id_col else f} for i, (h, f) in enumerate(zip(headers, fields)) if h]
        skipped = []
        for r in range(header_at + 1, len(rows)):
            row = rows[r]
            if not any(str(c or "").strip() for c in row):
                continue
            id_value = str(row[id_col] or "").strip() if id_col is not None and id_col < len(row) else ""
            # 지정한 번호는 번호 열이 없는 표(사양서 등)에만 씁니다.
            hinted = not id_value and id_col is None and bool(entity_hint) and table_kind != "detector"
            entity = id_value or (entity_hint if hinted else "")
            by_name = False
            if not entity and id_col is None:
                entity = next((str(row[i]).strip() for i, f in enumerate(fields) if f == "name" and i < len(row) and row[i]), "")
                by_name = bool(entity)
            if not entity:
                # 번호 열이 있는 표에서 번호가 빈 줄(합계·메모·번호를 안 적은 설비 등)은 추측하지 않고 건너뛰고 알려 줍니다.
                text = next((str(c).strip() for c in row if str(c or "").strip()), "")
                skipped.append({"row": r + 1, "text": text[:40],
                                "reason": "번호 칸이 비어 있음" if id_col is not None else "번호·이름 칸이 없음"})
                continue
            base = STRICT_CONFIDENCE if id_value else (HINT_CONFIDENCE if hinted else 0.78)
            id_flags = ("merged",) if id_value and table.flags.get((r, id_col)) == "merged" else ()
            if hinted:
                id_flags += ("hint",)
            if by_name:
                id_flags += ("by_name",)
            for i, field in enumerate(fields):
                if i >= len(row) or row[i] is None or not str(row[i]).strip():
                    continue
                if field:
                    flag = table.flags.get((r, i))
                    cell_flags = tuple(f for f in (flag,) if f) + id_flags
                    add(table_kind, entity, field, row[i], f"{table.location} 행 {r + 1} 열 {i + 1}", headers[i],
                        confidence=base, flags=cell_flags)
                elif headers[i] and i != id_col:
                    unknown(headers[i])
        skipped_total += len(skipped)
        table_info.append({"location": table.location, "header_row": header_at + 1, "rows": len(rows) - header_at - 1,
                           "kind": table_kind, "columns": columns, "skipped": skipped[:50], "skipped_count": len(skipped)})
    # 식별 대상을 지정한 명확한 '항목: 값' 문장만 후보로 추출합니다.
    for passage in passages:
        if not entity_hint:
            continue
        for line in passage["text"].splitlines():
            parts = re.split(r"[:：]", line, maxsplit=1)
            hint_kind = kind if kind in ("facility", "material") else "facility"
            if len(parts) == 2 and _field_of(parts[0], rules, hint_kind):
                add(hint_kind, entity_hint, _field_of(parts[0], rules, hint_kind), parts[1].strip(), passage["location"], parts[0],
                    method="literal", confidence=0.88)
    # 사업장 일반정보·공정 설명 문서의 ‘대표자: 홍길동’ 같은 줄(사업장 자료로 볼 근거가 있을 때만)
    lines = _site_lines(passages)
    if lines and (kind == "site" or site_found or len({norm(label) for _, label, _ in lines}) >= 2):
        for location, label, value in lines:
            add("site", SITE_ENTITY.split(":", 1)[1], CATALOG_HEADERS["site"][norm(label)], value, location, label,
                method="label", confidence=LABEL_CONFIDENCE)
    has_text = any(p["text"].strip() for p in passages)
    needs_ocr = bool(passages) and not has_text
    problem = ""
    if not added:
        if needs_ocr:
            problem = "글자가 없는 문서예요. 스캔본이면 글자를 읽을 수 없어요."
        elif not tables and has_text and not entity_hint:
            problem = ("표가 없고 문장만 있어요. 사업장 정보는 ‘대표자: 홍길동’처럼 항목 이름과 값을 한 줄에 적고, "
                       "설비·제품 문서라면 어느 설비·제품의 문서인지 번호를 적으면 ‘항목: 값’ 문장을 읽어요.")
        elif tables and all(t["header_row"] is None for t in table_info):
            problem = "표는 있지만 설비번호·용량·CAS 같은 항목 이름이 있는 줄을 찾지 못했어요."
        else:
            problem = "읽을 수 있는 값을 찾지 못했어요."
    data["documents"].append({"name": name, "sha256": digest, "kind": kind, "hint": entity_hint, "passages": passages,
                              "fact_count": len(added), "unmapped_columns": sorted(set(unmapped)),
                              "tables": table_info, "rows_without_id": skipped_total,
                              "needs_reading": not added, "needs_ocr": needs_ocr, "problem": problem})
    data["facts"].extend(added)
    _link_materials(data)
    _store(project, data)
    _invalidate(project)
    auto_approve_safe(project)
    return len(added)


def remove_document(project, sha):
    """문서와 그 문서에서 읽은 값을 지웁니다(다시 읽기 전에 씁니다). 담당자 직접 입력은 남깁니다."""
    data = graph(project)
    data["documents"] = [d for d in data["documents"] if d["sha256"] != sha]
    data["facts"] = [f for f in data["facts"] if f["sha256"] != sha or f.get("method") == "manual"]
    _store(project, data)
    _invalidate(project)


def reingest(project, sha, content):
    """규칙(머리글 지정 등)이 바뀐 뒤 같은 파일을 다시 읽습니다. 같은 칸의 값은 같은 식별자를 받아 결정이 유지됩니다."""
    doc = next(d for d in _read(project)["documents"] if d["sha256"] == sha)
    name, kind, hint = doc["name"], doc["kind"], doc.get("hint", "")
    remove_document(project, sha)
    return ingest(project, name, content, kind, hint)


def learn_value(project, field, raw, value):
    """담당자가 한 칸에 답한 표기를 같은 항목의 같은 표기 전체에 적용합니다. 바뀐 칸 수를 돌려줍니다."""
    if field not in LABELS or _is_number(raw):
        raise ValueError("숫자로 적힌 값은 칸마다 따로 확인하세요.")
    normalized = normalize_value(field, value, LABELS[field])
    problem = value_problem(field, normalized)
    if problem:
        raise ValueError(problem)
    data = graph(project)
    _rules(data)["values"].setdefault(field, {})[norm(raw)] = normalized
    changed = _reapply(data)
    _store(project, data)
    _invalidate(project)
    auto_approve_safe(project)
    return changed


def learn_unit(project, sha, header, unit):
    """단위가 없는 열에 단위를 정해 그 열의 값 전체에 적용합니다."""
    data = graph(project)
    fields = {f["field"] for f in data["facts"] if f["sha256"] == sha and f["header"] == header}
    if len(fields) != 1 or norm(unit) not in UNIT_SUFFIXES.get(next(iter(fields)), ()):
        raise ValueError("이 열에 쓸 수 있는 단위를 고르세요.")
    _rules(data)["units"][f"{sha}|{norm(header)}"] = unit
    changed = _reapply(data)
    _store(project, data)
    _invalidate(project)
    auto_approve_safe(project)
    return changed


def learn_column(project, header, field):
    """프로그램이 모르는 머리글을 항목으로 지정합니다. 다시 읽어야 하는 문서의 sha 목록을 돌려줍니다."""
    if field not in LABELS:
        raise ValueError("항목을 고르세요.")
    data = graph(project)
    _rules(data)["columns"][norm(header)] = field
    _store(project, data)
    return [d["sha256"] for d in data["documents"] if header in d.get("unmapped_columns", [])]


UNIT_CHOICES = {"capacity": ("m3", "L"), "pressure": ("MPa", "kPa"), "design_pressure": ("MPa", "kPa"),
                "inventory": ("kg", "ton"), "temperature": ("℃",), "concentration": ("%",)}


def problems(project):
    """그대로 쓸 수 없는 값을, 원인을 추측하지 않고 ‘어떤 검사에 걸렸는지’로 묶어 돌려줍니다.

    묶음 종류
    - column: 항목 이름에 단위가 없는 열 → 열마다 한 번 묻기
    - text: 숫자가 아니거나 선택지에 없는 표기 → 같은 항목·같은 표기마다 한 번 묻기(답하면 전체 적용)
    - single: 계산 결과 없는 계산식·허용 범위 밖 숫자·CAS 검산 실패 → 칸마다 묻기
    자료끼리 값이 다른 칸(충돌)은 충돌 질문에서 다루므로 넣지 않습니다.
    """
    def compute():
        rules = _rules(deepcopy({"rules": _read(project).get("rules", {})}))
        groups = {}
        for slot in slots(project):
            if slot["chosen"] or slot["status"] == "값 충돌":
                continue
            for fact in slot["facts"]:
                field = fact["field"]
                problem = value_problem(field, fact["value"])
                if not problem or fact.get("method") == "manual":
                    continue
                raw = str(fact["raw"]).strip()
                low_high = BOUNDS.get(field)
                if "formula_uncached" in fact.get("flags", []) or raw.startswith("="):
                    key, kind, reason = ("single", fact["id"]), "single", "엑셀 계산식만 있고 계산 결과가 파일에 저장돼 있지 않아요"
                elif field == "cas":
                    key, kind, reason = ("single", fact["id"]), "single", "CAS 번호의 형식이나 검산 숫자가 맞지 않아요"
                elif field in UNIT_SUFFIXES and _is_number(raw) and not has_unit(field, fact["header"]) \
                        and f"{fact['sha256']}|{norm(fact['header'])}" not in rules["units"]:
                    key, kind, reason = ("column", fact["sha256"], fact["header"]), "column", "항목 이름에 단위가 없어요"
                elif isinstance(fact["value"], (int, float)) and low_high:
                    key, kind, reason = ("single", fact["id"]), "single", f"허용 범위({low_high[0]:g}~{low_high[1]:g}) 밖의 숫자예요"
                elif field in CHOICES:
                    key, kind, reason = ("text", field, norm(raw)), "text", "선택지에 없는 말이에요"
                else:
                    key, kind, reason = ("text", field, norm(raw)), "text", "숫자가 아니에요"
                group = groups.setdefault(key, {"kind": kind, "field": field, "raw": raw, "reason": reason,
                                                "sha256": fact["sha256"], "header": fact["header"], "facts": [], "keys": []})
                group["facts"].append(fact)
                if slot["key"] not in group["keys"]:
                    group["keys"].append(slot["key"])
        order = {"column": 0, "text": 1, "single": 2}
        return sorted(groups.values(), key=lambda g: (order[g["kind"]], -len(g["facts"]), g["field"], g["raw"]))
    return _memo(project, "problems", compute)


# ---------------------------------------------------------------- 사실 묶음(slot)과 규칙

def _usable_values(facts):
    """그대로 쓸 수 있는 값들(서로 다른 것만). 충돌은 쓸 수 있는 값끼리 다를 때만입니다.
    단위가 없거나 글자로 적힌 값은 ‘다른 값’이 아니라 ‘그대로 쓸 수 없는 값’으로 따로 다룹니다."""
    return {json.dumps(f["value"], ensure_ascii=False, sort_keys=True) for f in facts
            if value_problem(f["field"], f["value"]) is None}


def _slots(data):
    groups = {}
    excluded = set(data.get("excluded", []))
    for fact in data["facts"]:
        if fact["entity"] in excluded:
            continue
        groups.setdefault(fact["entity"] + "|" + fact["field"], []).append(fact)
    result = []
    for key, facts in groups.items():
        values = _usable_values(facts)
        decision = data["decisions"].get(key) or {}
        signature = sha256(json.dumps(sorted(f["id"] for f in facts)).encode()).hexdigest()
        ids = {f["id"] for f in facts}
        confirmed = bool(decision and decision.get("signature") == signature and decision.get("fact_id") in ids)
        chosen = next((f for f in facts if confirmed and f["id"] == decision["fact_id"]), None)
        status = DECISION_LABEL.get(decision.get("method"), "확인 완료") if confirmed else ("값 충돌" if len(values) > 1 else "담당자 확인")
        result.append({"key": key, "entity": facts[0]["entity"], "field": facts[0]["field"], "facts": facts,
                       "signature": signature, "status": status, "chosen": chosen,
                       "method": decision.get("method") if confirmed else None})
    return result


def slots(project):
    """읽기 전용 slot 목록. 같은 저장 버전이면 다시 계산하지 않습니다(돌려받은 목록을 고치지 마세요)."""
    return _memo(project, "slots", lambda: _slots(_read(project)))


def graph_facts(project):
    """읽기 전용 사실 목록(복사하지 않음)."""
    return _read(project)["facts"]


def critical_slot_keys(project):
    """결과 순위를 바꿀 수 있는 값의 slot 키. 순위 계산이 불가능하면 빈 집합."""
    def compute():
        try:
            from .scenario_triage import critical_slot_keys as keys
            return keys(project)
        except Exception:
            return frozenset()
    return _memo(project, "critical", compute)


def review_gate(slot, critical=frozenset()):
    """자동 기입을 막는 이유 목록. 빈 목록이면 자동 기입해도 됩니다.

    모델 신뢰도 자체는 승인 근거가 아니며, 아래 규칙을 모두 통과해야 합니다.
    """
    facts = slot["facts"]
    field = slot["field"]
    if len(_usable_values(facts)) > 1:
        return ["자료 간 값 충돌"]
    for fact in facts:
        problem = value_problem(field, fact["value"])
        if problem:
            return [problem]
    if any(f.get("method") not in READ_METHODS for f in facts):
        return ["문장 또는 AI 추출값"]
    # 담당자가 열의 단위를 정했으면(규칙) 단위가 있는 것으로 봅니다.
    if field in UNIT_SUFFIXES and any(not has_unit(field, f.get("header", "")) and "rule" not in f.get("flags", []) for f in facts):
        return ["항목명에 단위가 없음"]
    strict = field in LEGAL_FIELDS or slot["key"] in critical
    needed = STRICT_CONFIDENCE if strict else RELAXED_CONFIDENCE
    low = [f for f in facts if f.get("confidence", 0) < needed]
    if low:
        causes = []
        for f in low:
            if "hint" in f.get("flags", []):
                causes.append("번호 없는 문서라 지정한 번호로 연결")
            elif f.get("confidence", 0) < RELAXED_CONFIDENCE:
                causes.append("설비번호·제품코드 없이 이름으로 연결")
            elif "formula" in f.get("flags", []):
                causes.append("수식 계산 결과")
            else:
                causes.append("병합 셀에서 가져옴")
        why = "법적 판단에 쓰이는 항목" if field in LEGAL_FIELDS else "예비 순위를 바꿀 수 있는 핵심값"
        prefix = f"{why}이라 엄격 적용: " if strict else ""
        return [prefix + "출처 신뢰도 부족(" + ", ".join(dict.fromkeys(causes)) + ")"]
    return []


def auto_approve_safe(project):
    data = graph(project)
    critical = critical_slot_keys(project)
    count = 0
    for slot in _slots(data):
        if slot["chosen"] or review_gate(slot, critical):
            continue
        data["decisions"][slot["key"]] = {"fact_id": slot["facts"][0]["id"], "signature": slot["signature"], "method": "AUTO"}
        count += 1
    if count:
        _store(project, data)
    return count


def approve(project, key, fact_id):
    slot = next(s for s in slots(project) if s["key"] == key)
    fact = next(f for f in slot["facts"] if f["id"] == fact_id)
    problem = value_problem(slot["field"], fact["value"])
    if problem:
        raise ValueError(problem)
    data = graph(project)
    data["decisions"][key] = {"fact_id": fact_id, "signature": slot["signature"], "method": "PERSON"}
    _invalidate(project)
    _store(project, data)


def approve_bulk(project, keys=None):
    """표에서 읽은, 충돌 없고 유효한 값을 한꺼번에 확인합니다. AI·문장 추출값은 제외합니다."""
    data = graph(project)
    wanted = set(keys) if keys is not None else None
    approved, skipped = 0, []
    for slot in _slots(data):
        if slot["chosen"] or (wanted is not None and slot["key"] not in wanted):
            continue
        facts = slot["facts"]
        reason = None
        if len(_usable_values(facts)) > 1:
            reason = "값 충돌"
        elif any(f.get("method") not in READ_METHODS for f in facts):
            reason = "문장 또는 AI 추출값"
        else:
            reason = next((value_problem(slot["field"], f["value"]) for f in facts if value_problem(slot["field"], f["value"])), None)
        if reason:
            skipped.append((slot["key"], reason))
            continue
        data["decisions"][slot["key"]] = {"fact_id": facts[0]["id"], "signature": slot["signature"], "method": "BULK"}
        approved += 1
    if approved:
        _invalidate(project)
        _store(project, data)
    return {"approved": approved, "skipped": skipped}


def confirm_auto(project, keys):
    """자동 기입한 값을 사람이 확인한 것으로 바꿉니다(핵심값 확인용)."""
    data = graph(project)
    changed = 0
    for key in keys:
        decision = data["decisions"].get(key)
        if decision and decision.get("method") == "AUTO":
            decision["method"] = "PERSON"
            changed += 1
    if changed:
        _store(project, data)
    return changed


def add_manual(project, entity, field, value, reason, confirm=False):
    if not entity.startswith(("facility:", "material:")) or field not in LABELS or not reason.strip():
        raise ValueError("대상 식별번호·항목·확인 근거를 입력하세요.")
    header = LABELS[field]
    normalized = normalize_value(field, value, header)
    if confirm:
        problem = value_problem(field, normalized)
        if problem:
            raise ValueError(problem)
    data = graph(project)
    fact = {"entity": entity, "kind": entity.split(":", 1)[0], "field": field, "value": normalized,
            "raw": str(value), "header": header, "document": "담당자 입력", "sha256": sha256(reason.encode()).hexdigest(),
            "location": reason.strip(), "method": "manual", "confidence": 0.9, "flags": []}
    fact["id"] = sha256(json.dumps(fact, ensure_ascii=False, sort_keys=True).encode()).hexdigest()[:24]
    if fact["id"] not in {f["id"] for f in data["facts"]}:
        data["facts"].append(fact)
    _invalidate(project)
    _store(project, data)
    if confirm:
        slot = next(s for s in slots(project) if s["entity"] == entity and s["field"] == field)
        approve(project, slot["key"], fact["id"])


def confirmed_entities(project):
    entities = {}
    for slot in slots(project):
        if slot["chosen"]:
            entities.setdefault(slot["entity"], {})[slot["field"]] = slot["chosen"]["value"]
    return entities


# ---------------------------------------------------------------- 사람이 볼 목록

def review_queue(project):
    """사람이 처리할 예외 목록: 충돌 → 값 오류 → 확인 필요, 핵심값 먼저."""
    critical = critical_slot_keys(project)
    order = {"충돌": 0, "값 오류": 1, "확인 필요": 2}
    items = []
    for slot in slots(project):
        if slot["chosen"]:
            continue
        reasons = review_gate(slot, critical)
        if slot["status"] == "값 충돌":
            kind = "충돌"
        elif any(value_problem(slot["field"], f["value"]) for f in slot["facts"]):
            kind = "값 오류"
        else:
            kind = "확인 필요"
        items.append({"key": slot["key"], "entity": slot["entity"], "field": slot["field"], "kind": kind,
                      "reasons": reasons, "slot": slot, "critical": slot["key"] in critical})
    items.sort(key=lambda i: (order[i["kind"]], not i["critical"], i["entity"], i["field"]))
    return items


def critical_auto(project):
    """자동 기입했지만 결과를 바꿀 수 있어 먼저 확인할 값."""
    critical = critical_slot_keys(project)
    return [s for s in slots(project) if s["method"] == "AUTO" and s["key"] in critical]


def summary(project):
    all_slots = slots(project)
    queue = review_queue(project)
    return {"documents": len(_read(project)["documents"]), "slots": len(all_slots),
            "auto": sum(s["method"] == "AUTO" for s in all_slots),
            "person": sum(s["method"] in {"PERSON", "BULK"} for s in all_slots),
            "conflicts": sum(i["kind"] == "충돌" for i in queue),
            "errors": sum(i["kind"] == "값 오류" for i in queue),
            "pending": len(queue), "critical_auto": len(critical_auto(project))}


def tasks(project):
    data = _read(project)
    all_slots = _slots(data)
    result = [{"entity": s["entity"], "task": s["status"] + ": " + LABELS[s["field"]], "key": s["key"]}
              for s in all_slots if not s["chosen"]]
    for doc in data["documents"]:
        if doc["needs_reading"]:
            result.append({"entity": doc["name"],
                           "task": "OCR 또는 문서 구조 확인 필요" if doc["needs_ocr"] else
                           "문서를 읽었지만 사실을 연결하지 못했습니다. 표·항목 또는 식별 대상을 확인하세요.", "key": ""})
        elif doc.get("rows_without_id"):
            result.append({"entity": doc["name"],
                           "task": f"번호·이름이 없거나 합계 줄이라 건너뛴 행 {doc['rows_without_id']}개", "key": ""})
    by_entity, confirmed = {}, {}
    for s in all_slots:
        by_entity.setdefault(s["entity"], set()).add(s["field"])
        if s["chosen"]:
            confirmed.setdefault(s["entity"], {})[s["field"]] = s["chosen"]["value"]
    for identifier in sorted(e for e in by_entity if e.startswith("facility:")):
        present = by_entity[identifier]
        for field in ("product", "inventory", "pressure", "temperature", "facility_type", "substance_state"):
            if field not in present:
                result.append({"entity": identifier, "task": "기본 운전조건 미확인: " + LABELS[field], "key": ""})
        product = confirmed.get(identifier, {}).get("product")
        if product:
            matches = [k for k, v in confirmed.items() if k.startswith("material:") and (v.get("name") == product or k.split(":", 1)[1] == product)]
            if len(matches) != 1:
                result.append({"entity": identifier, "task": "취급제품과 물질자료를 하나로 연결할 수 없습니다. 제품코드·이름을 확인하세요.", "key": ""})
    result.extend({"entity": "별지 제1호", "task": need, "key": ""} for need in form1_needs(project))
    return result


def form1_needs(project):
    """연결한 설비 자료로 별지 제1호 최대보유량을 계산할 때 아직 부족한 값(계산 엔진이 알려 줌)."""
    record = project.get_field("inventory.facilities")
    if record is None or record.status not in {"USER_CONFIRMED", "VERIFIED"}:
        return ()
    try:
        from . import cap_workspace
        return tuple(cap_workspace.resolve_form1(project).needs)
    except Exception:
        return ()


# ---------------------------------------------------------------- 연결·보고서

MATERIAL_CORE_COLUMNS = {"name": "물질명", "cas": "CAS 번호", "concentration": "함량(%)", "substance_state": "물질상태", "density": "비중"}
FACILITY_CORE_COLUMNS = {"name": "설비명", "product": "취급물질", "capacity": "용량", "pressure": "운전압력(MPa)",
                         "design_pressure": "설계압력(MPa)", "temperature": "운전온도(℃)", "inventory": "최대보유량(kg)",
                         "facility_type": "시설유형", "substance_state": "물질성상", "density": "비중", "process_type": "공정유형"}


def _catalog_columns(kind, key, core=None):
    """사실 카탈로그에서 이 표(key)에 들어가는 열: 읽은 항목 → 열 이름들. core에 이미 있는 (항목, 열)은 뺍니다."""
    core = core or {}
    out = {}
    for item in catalog.facts(kinds=[kind]):
        if item.identifier or not item.evidence_field:
            continue
        columns = ([item.column] if item.field_key == key and item.column else []) + [c for k, c in item.targets if k == key]
        for column in columns:
            if core.get(item.evidence_field) != column and column not in out.get(item.evidence_field, ()):
                out.setdefault(item.evidence_field, []).append(column)
    return out


MATERIAL_EXTRA_COLUMNS = _catalog_columns("material", "inventory.chemicals", MATERIAL_CORE_COLUMNS)
FACILITY_EXTRA_COLUMNS = _catalog_columns("facility", "inventory.facilities", FACILITY_CORE_COLUMNS)


def publishing_plan(project):
    mapping, facilities = MATERIAL_CORE_COLUMNS, FACILITY_CORE_COLUMNS
    entities = confirmed_entities(project)
    plans = {}
    for kind, target, columns, extra in [("material", "inventory.chemicals", mapping, MATERIAL_EXTRA_COLUMNS),
                                         ("facility", "inventory.facilities", facilities, FACILITY_EXTRA_COLUMNS)]:
        rows = []
        for identifier, values in entities.items():
            if not identifier.startswith(kind + ":"):
                continue
            row = {columns[k]: v for k, v in values.items() if k in columns}
            for field, names in extra.items():       # 카탈로그 항목(설계온도·연결구·SDS 물성 등)과 서식 엔진이 읽는 다른 열 이름
                if field in values:
                    row.update({name: values[field] for name in names})
            row["물질코드" if kind == "material" else "시설번호"] = identifier.split(":", 1)[1]
            if kind == "facility":
                row["설비번호"] = identifier.split(":", 1)[1]
                if "capacity" in values:
                    row["용량단위"] = "m3"
                if "name" in values:
                    row["시설명"] = values["name"]
                matches = [v for k, v in entities.items() if k.startswith("material:") and (v.get("name") == values.get("product") or k.split(":", 1)[1] == values.get("product"))]
                if len(matches) == 1:
                    row["취급물질"] = matches[0].get("name", values.get("product", ""))
                    row["제품명"] = row["취급물질"]
                    if "cas" in matches[0]:
                        row["CAS 번호"] = matches[0]["cas"]
                    if "density" not in values and "density" in matches[0]:
                        row["비중"] = matches[0]["density"]       # 설비대장에 비중이 없으면 취급물질의 비중(별지 제1호 계산에 필요)
            rows.append(row)
        if rows:
            plans[target] = rows
    return plans


def unreviewed_legal(project):
    return [slot for slot in slots(project) if slot["method"] == "AUTO" and slot["field"] in LEGAL_FIELDS]


# ---------------------------------------------------------------- 한 곳의 사실(별지 작성대와 동기화)
#
# 같은 물질·설비 사실이 ‘자료 접수’ 기록과 별지 제1·6호 작성대 표에 따로 있으면 값이 갈라집니다.
# 그래서 자료 접수 기록(이 그래프)을 기준으로 삼고
# - 별지에서 고친 값은 그래프에 ‘직접 입력’ 사실로 되돌려 적고(sync_rows),
# - 그래프에서 보낼 때는 별지 표를 덮어쓰지 않고 같은 번호의 행에 합칩니다(publish).
# 별지에서 행을 지우면 그래프에서도 그 대상을 제외합니다(되살리면 다시 포함).

FACILITY_ID_COLUMNS = ("설비번호", "구분기호", "장치번호", "시설번호")
MATERIAL_ID_COLUMNS = ("CAS 번호", "CAS No.", "CAS No", "CAS")
# 그래프 항목 → 별지 표에서 같은 사실을 담는 열 이름들
FACILITY_ROW_COLUMNS = {"name": ("설비명",), "product": ("취급물질",), "capacity": ("용량",), "facility_type": ("시설유형",),
                        "substance_state": ("물질성상",), "density": ("비중",), "pressure": ("운전압력(MPa)",),
                        "temperature": ("운전온도(℃)",), "process_type": ("공정유형",), "design_pressure": ("설계압력(MPa)",)}
MATERIAL_ROW_COLUMNS = {"name": ("물질명", "화학물질명", "유해화학물질명"), "cas": MATERIAL_ID_COLUMNS,
                        "concentration": ("함량(%)", "함량"), "substance_state": ("물질상태", "물질성상"), "density": ("비중",)}


def _cell(row, names):
    return next((str(row[n]).strip() for n in names if n in row and row[n] is not None and str(row[n]).strip()), "")


def _row_id(kind, row):
    if kind == "facility":
        return _cell(row, FACILITY_ID_COLUMNS)
    return _cell(row, MATERIAL_ID_COLUMNS) or _cell(row, ("물질코드",))     # CAS가 없으면 제품코드로


def _material_entities_by_cas(all_slots):
    by_cas = {}
    for slot in all_slots:
        if slot["field"] == "cas":
            fact = slot["chosen"] or (slot["facts"][0] if len(_usable_values(slot["facts"])) == 1 else None)
            if fact:
                by_cas.setdefault(str(fact["value"]).strip(), slot["entity"])
    return by_cas


def _same_product(current, a, b):
    """취급제품은 제품코드로도 이름으로도 적히므로, 같은 물질을 가리키면 같은 값으로 봅니다."""
    def resolve(text):
        text = str(text).strip()
        hits = {s["entity"] for s in current.values() if s["entity"].startswith("material:") and
                (s["entity"].split(":", 1)[1] == text or (s["field"] == "name" and s["chosen"] and str(s["chosen"]["value"]).strip() == text))}
        return hits.pop() if len(hits) == 1 else text
    return resolve(a) == resolve(b)


def _entity_for(kind, row_id, by_cas):
    if kind == "facility":
        return "facility:" + row_id
    return by_cas.get(row_id, "material:" + row_id)


def sync_rows(project, kind, rows, source, *, removed_ids=()):
    """별지 작성대에서 저장한 표를 그래프에 반영합니다. 그래프와 다른 값만 ‘직접 입력’ 사실로 적고 확인 완료로 둡니다.
    removed_ids: 별지에서 지운 행의 번호(설비번호 또는 CAS). 그래프에 있으면 제외 목록에 넣습니다. 바뀐 값 수를 돌려줍니다.

    별지 작성대가 이미 저장한 값이므로 이미 보낸 작성자료를 다시 보류(HOLD)로 돌리지 않습니다."""
    if kind not in {"facility", "material"}:
        raise ValueError("자료 종류를 확인하세요.")
    columns = FACILITY_ROW_COLUMNS if kind == "facility" else MATERIAL_ROW_COLUMNS
    data = graph(project)
    data.setdefault("excluded", [])
    current = {s["key"]: s for s in _slots(data)}
    by_cas = _material_entities_by_cas(current.values()) if kind == "material" else {}
    known = {s["entity"] for s in current.values()} | set(data["excluded"])
    touched, changed = {}, 0
    stamp = sha256(source.encode()).hexdigest()
    for row in rows:
        row_id = _row_id(kind, row)
        if not row_id:
            continue
        entity = _entity_for(kind, row_id, by_cas)
        if entity in data["excluded"]:
            data["excluded"].remove(entity)
            changed += 1
        for field, names in columns.items():
            raw = _cell(row, names)
            if not raw:
                continue
            header = LABELS[field]
            if field == "capacity":
                unit = _cell(row, ("용량단위",)).lower()
                header = "용량(L)" if unit in {"l", "리터"} else "용량(m3)"
            value = normalize_value(field, raw, header)
            if value_problem(field, value):
                continue        # 별지에서 쓸 수 없는 값은 그래프에 옮기지 않습니다(그래프 값을 망가뜨리지 않음)
            slot = current.get(f"{entity}|{field}")
            if slot and slot["chosen"] and (slot["chosen"]["value"] == value or
                                            (field == "product" and _same_product(current, slot["chosen"]["value"], value))):
                continue
            fact = {"entity": entity, "kind": kind, "field": field, "value": value, "raw": raw, "header": header,
                    "document": source, "sha256": stamp, "location": source, "method": "manual", "confidence": 0.9, "flags": []}
            fact["id"] = sha256(json.dumps(fact, ensure_ascii=False, sort_keys=True).encode()).hexdigest()[:24]
            if fact["id"] not in {f["id"] for f in data["facts"]}:
                data["facts"].append(fact)
            touched[f"{entity}|{field}"] = fact["id"]
            changed += 1
    for row_id in removed_ids:
        entity = _entity_for(kind, str(row_id).strip(), by_cas)
        if entity in known and entity not in data["excluded"]:
            data["excluded"].append(entity)
            changed += 1
    if not changed:
        return 0
    for slot in _slots(data):
        if slot["key"] in touched:
            data["decisions"][slot["key"]] = {"fact_id": touched[slot["key"]], "signature": slot["signature"], "method": "PERSON"}
    _store(project, data)
    return changed


def sync_from_form(project, kind, rows, source, *, removed_ids=()):
    """별지 작성대에서 부르는 쪽. 동기화가 실패해도 별지 저장 자체는 막지 않습니다(바뀐 값 수, 실패하면 0)."""
    try:
        return sync_rows(project, kind, list(rows), source, removed_ids=removed_ids)
    except Exception:
        return 0


def excluded(project):
    return list(_read(project).get("excluded", []))


def restore(project, entity):
    """별지에서 지워 제외한 대상을 다시 포함합니다."""
    data = graph(project)
    if entity in data.get("excluded", []):
        data["excluded"].remove(entity)
        _store(project, data)


CHEMICAL_DETAILS_KEY = "cap.chemical.details"
WORKSPACE_FACILITIES_KEY = "cap.workspace.facilities"
MATERIAL_TABLE_COLUMNS = ("물질명", "CAS 번호", "함량(%)", "물질상태", "비중")
FACILITY_TABLE_COLUMNS = ("설비번호", "설비명", "취급물질", "CAS 번호", "용량", "용량단위", "시설유형", "물질성상", "비중",
                          "운전압력(MPa)", "운전온도(℃)", "공정유형", "설계압력(MPa)")
MATERIAL_MANAGED = MATERIAL_TABLE_COLUMNS + tuple(c for cs in MATERIAL_EXTRA_COLUMNS.values() for c in cs)
FACILITY_MANAGED = FACILITY_TABLE_COLUMNS + tuple(c for cs in FACILITY_EXTRA_COLUMNS.values() for c in cs)
# 확정되지 않은(충돌·확인 전) 값이 있으면 지울 열: 기본 열 + 카탈로그 열
PENDING_COLUMNS = {"facility": {f: tuple(FACILITY_ROW_COLUMNS.get(f, ())) + tuple(FACILITY_EXTRA_COLUMNS.get(f, ()))
                                for f in set(FACILITY_ROW_COLUMNS) | set(FACILITY_EXTRA_COLUMNS)},
                   "material": {f: tuple(MATERIAL_ROW_COLUMNS.get(f, ())) + tuple(MATERIAL_EXTRA_COLUMNS.get(f, ()))
                                for f in set(MATERIAL_ROW_COLUMNS) | set(MATERIAL_EXTRA_COLUMNS)}}
LINK_NOTE = "공통정보 연결(회사 파일에서 자동 기입한 초안)"


def _merge(kind, existing, new, managed, removed, pending=None):
    """같은 번호의 행은 그래프가 아는 열만 바꾸고 나머지 열(별지에서만 적는 값)은 그대로 둡니다.
    새 번호는 뒤에 붙이고, 제외한 대상의 행은 뺍니다. pending(번호 → 열): 그래프에서 아직 확정되지 않은(충돌·확인 전) 값은
    옛 값을 그대로 두지 않고 비웁니다 — 다툼이 있는 값을 조용히 남기지 않기 위해서입니다."""
    return _merge_rows(lambda row: _row_id(kind, row), existing, new, managed, removed, pending,
                       aliases={"CAS 번호": MATERIAL_ID_COLUMNS[1:]})


def _merge_rows(row_id_of, existing, new, managed, removed=(), pending=None, aliases=None):
    pending, aliases = pending or {}, aliases or {}
    by_id = {row_id_of(row): row for row in new if row_id_of(row)}
    out, seen = [], set()
    for row in existing or []:
        row_id = row_id_of(row)
        if row_id and row_id in removed:
            continue
        merged = dict(row)
        for column in pending.get(row_id, ()):
            if column in managed and merged.get(column) not in (None, ""):
                merged[column] = ""
        if row_id in by_id:
            seen.add(row_id)
            for column in managed:
                value = by_id[row_id].get(column)
                if value not in (None, ""):
                    for name in (column,) + tuple(aliases.get(column, ())):
                        if name in merged or name == column:
                            merged[name] = value
        out.append(merged)
    out.extend(row for row_id, row in by_id.items() if row_id not in seen)
    return out


# ---------------------------------------------------------------- 카탈로그 사실을 서식 작성자료로
#
# 사업장 일반정보·공정 설명(한 값), 설비 사양(별지 제9호), 확산방지설비(제10호), 가스감지기(제11호), 대표물질 유해성(제7호)을
# 작성자료 키에 그대로 적습니다. 파일에서 읽어 정해진 값만 보내고(충돌·확인 전 값은 보내지 않음), 사람이 다른 곳에서 적은 값이
# 이미 있으면 덮어쓰지 않고 사실 표에서 ‘저장된 값과 다름’으로 알립니다.

SPEC_KEY = "cap.workspace.equipment_specs"
DIKE_KEY = "cap.safety.dike_calculation"
DIKE_RULE_KEY = "cap.workspace.dike_rule"
GAS_KEY = "cap.safety.gas_detection"
HAZARD_KEY = "cap.chemical.hazard_information"
LORRY_TYPE = "탱크로리·운송차량"
FIXED_DETECTOR = "고정식"


def _chosen_facts(project):
    """대상 → 항목 → 정해진(자동 기입·확인 완료) 사실."""
    out = {}
    for slot in slots(project):
        if slot["chosen"]:
            out.setdefault(slot["entity"], {})[slot["field"]] = slot["chosen"]
    return out


def _refs(facts):
    refs, seen = [], set()
    for fact in facts:
        if fact.get("method") == "manual" or fact["document"] in seen:
            continue
        seen.add(fact["document"])
        refs.append(EvidenceRef(source_type="COMPANY_EVIDENCE", source_name=fact["document"], sha256=fact["sha256"],
                                location=str(fact["location"])))
    return refs


def file_linked(record):
    """회사 파일에서 자동 기입한 작성자료인지(사람이 다른 곳에서 적은 값과 구분)."""
    return record is not None and str(record.note).startswith("공통정보 연결")


def _site_values(project, chosen):
    """사업장 사실 → {작성자료 키: (값, 근거 사실들, 이름)}. 다른 곳에서 적은 값과 다르면 보내지 않습니다."""
    site = chosen.get(SITE_ENTITY, {})
    out = {}
    for item in catalog.facts(kinds=["site", "narrative"], company=True):
        fact = site.get(item.evidence_field)
        if fact is None or not item.field_key or item.field_key.startswith("@") or item.field_key == DIKE_RULE_KEY:
            continue
        value = str(fact["value"]).strip()
        if item.id == "site.loading_facility_count":
            lorries = sum(1 for e, v in chosen.items() if e.startswith("facility:") and v.get("facility_type")
                          and v["facility_type"]["value"] == LORRY_TYPE)
            count = re.sub(r"\s*기$", "", value)
            value = f"입·출하 시설 {count}기" + (f", 보유 탱크로리 {lorries}기" if lorries else "")
        out[item.field_key] = (value, [fact], item.name)
    rule = {name: site.get(field) for name, field in (("비율(%)", "dike_required_ratio"), ("근거", "dike_required_basis"))}
    if any(rule.values()):
        out[DIKE_RULE_KEY] = ({k: (str(f["value"]).strip() if f else "") for k, f in rule.items()},
                              [f for f in rule.values() if f], "확산방지설비 공통 필요용량 기준")
    return out


def _entity_rows(chosen, kind, columns, id_column, extra=None):
    """대상마다 한 행. columns: 항목 → 열 이름. 값이 하나도 없으면 행을 만들지 않습니다."""
    rows, used = [], []
    for entity, values in chosen.items():
        if not entity.startswith(kind + ":"):
            continue
        row = {}
        for field, column in columns.items():
            if field in values:
                row[column] = values[field]["value"]
                used.append(values[field])
        if not row:
            continue
        row[id_column] = entity.split(":", 1)[1]
        if extra:
            row.update(extra(entity, values))
        rows.append(row)
    return rows, used


def _columns_for(kind, key):
    return {item.evidence_field: (item.column if item.field_key == key else next(c for k, c in item.targets if k == key))
            for item in catalog.facts(kinds=[kind])
            if item.evidence_field and not item.identifier and (item.field_key == key or any(k == key for k, _ in item.targets))}


def _table_values(project, chosen):
    """설비 사양·확산방지설비·감지기·대표물질 표 → {작성자료 키: (행 목록, 근거 사실들, 이름)}. 같은 번호의 행에 합칩니다."""
    out = {}
    names = {e: v["name"]["value"] for e, v in chosen.items() if "name" in v}

    spec_rows, used = _entity_rows(chosen, "facility", _columns_for("facility", SPEC_KEY), "설비번호",
                                   lambda e, v: {"설비명": names.get(e, "")})
    if spec_rows:
        record = project.get_field(SPEC_KEY)
        existing = record.value if record and isinstance(record.value, list) else []
        managed = tuple(_columns_for("facility", SPEC_KEY).values()) + ("설비명",)
        out[SPEC_KEY] = (_merge_rows(lambda r: norm(r.get("설비번호") or r.get("설비명") or ""), existing, spec_rows, managed),
                         used, "장치·설비 명세(별지 제9호, 회사 파일)")

    dike_columns = _columns_for("facility", DIKE_KEY)
    dike_rows, used = _entity_rows(chosen, "facility", dike_columns, "구분기호")
    if dike_rows:
        from .cap_form10_workspace import stored_rows

        rule = chosen.get(SITE_ENTITY, {})
        ratio = str(rule["dike_required_ratio"]["value"]) if "dike_required_ratio" in rule else ""
        basis = str(rule["dike_required_basis"]["value"]) if "dike_required_basis" in rule else ""
        for row in dike_rows:
            capacity = chosen.get("facility:" + row["구분기호"], {}).get("capacity")
            row["설계용량(m3)"] = capacity["value"] if capacity else ""
        record = project.get_field(DIKE_KEY)
        existing = record.value if record and isinstance(record.value, list) else []
        new = stored_rows(dike_rows, ratio, basis)
        managed = tuple(sorted({c for r in new for c in r}))
        out[DIKE_KEY] = (_merge_rows(lambda r: norm(r.get("대상 설비번호") or ""), existing, new, managed),
                         used, "확산방지설비 계산자료(별지 제10호, 회사 파일)")

    gas_rows, used = _entity_rows(chosen, "detector", _columns_for("detector", GAS_KEY), "감지기 번호")
    if gas_rows:
        for row in gas_rows:
            row.setdefault("설치형태", FIXED_DETECTOR)
        record = project.get_field(GAS_KEY)
        existing = record.value if record and isinstance(record.value, list) else []
        managed = tuple(_columns_for("detector", GAS_KEY).values())
        out[GAS_KEY] = (_merge_rows(lambda r: norm(r.get("감지기 번호") or ""), existing, gas_rows, managed),
                        used, "고정식 유해감지시설 명세(별지 제11호, 회사 파일)")

    hazard_columns = _columns_for("material", HAZARD_KEY)
    representatives = {e: v for e, v in chosen.items() if e.startswith("material:") and "representative_reason" in v}
    hazard_rows, used = _entity_rows(representatives, "material", hazard_columns, "물질코드",
                                     lambda e, v: {"물질명": v["name"]["value"] if "name" in v else "",
                                                   "CAS 번호": v["cas"]["value"] if "cas" in v else "",
                                                   "SDS 파일명": v["sds_file"]["value"] if "sds_file" in v else "",
                                                   "SDS 개정일": v["sds_revision"]["value"] if "sds_revision" in v else ""})
    if hazard_rows:
        record = project.get_field(HAZARD_KEY)
        existing = record.value if record and isinstance(record.value, list) else []
        managed = tuple(hazard_columns.values()) + ("물질명", "SDS 파일명", "SDS 개정일")
        out[HAZARD_KEY] = (_merge_rows(lambda r: str(r.get("CAS 번호") or r.get("물질명") or "").strip(), existing,
                                       hazard_rows, managed), used, "유해화학물질 유해성 정보(별지 제7호, 회사 파일)")
    return out


def catalog_preview(project):
    """카탈로그 사실로 채울 작성자료: {키: (값, 근거 사실들, 이름)}. 사람이 다른 곳에서 적은 다른 값은 덮어쓰지 않습니다."""
    chosen = _chosen_facts(project)
    out = {}
    for key, (value, facts, label) in {**_site_values(project, chosen), **_table_values(project, chosen)}.items():
        record = project.get_field(key)
        if record is not None and not file_linked(record) and record.value not in (None, "", []) and not isinstance(value, list):
            continue        # 사람이 적은 한 값은 그대로 둡니다. 다르면 사실 표에서 ‘저장된 값과 다름’으로 알립니다.
        out[key] = (value, facts, label)
    return out


def publish_preview(project):
    """보낼 때 각 표가 어떻게 바뀌는지(합친 결과). {필드 키: 행 목록 또는 값}. 화면 미리보기와 publish가 같은 결과를 씁니다."""
    plans = publishing_plan(project)
    data = _read(project)
    removed_material = {_row_id("material", {"CAS 번호": e.split(":", 1)[1]}) for e in data.get("excluded", []) if e.startswith("material:")}
    removed_material |= {str(f["value"]).strip() for f in data["facts"] if f["field"] == "cas" and f["entity"] in data.get("excluded", [])}
    removed_facility = {e.split(":", 1)[1] for e in data.get("excluded", []) if e.startswith("facility:")}
    pending = {"material": {}, "facility": {}}
    plan_ids = {kind: {_row_id(kind, r) for r in plans.get(t, [])} for kind, t in (("material", "inventory.chemicals"), ("facility", "inventory.facilities"))}
    entities = confirmed_entities(project)
    for slot in slots(project):
        kind = slot["entity"].split(":", 1)[0]
        if slot["chosen"] or kind not in pending:
            continue
        values = entities.get(slot["entity"], {})
        row_id = slot["entity"].split(":", 1)[1] if kind == "facility" else (values.get("cas") or slot["entity"].split(":", 1)[1])
        if str(row_id) not in plan_ids[kind]:
            continue
        columns = PENDING_COLUMNS[kind].get(slot["field"], ())
        pending[kind].setdefault(str(row_id), set()).update(columns)
    out = {}
    targets = [("inventory.chemicals", "material", plans.get("inventory.chemicals", []), MATERIAL_MANAGED, removed_material),
               (CHEMICAL_DETAILS_KEY, "material", plans.get("inventory.chemicals", []), MATERIAL_MANAGED, removed_material),
               ("inventory.facilities", "facility", plans.get("inventory.facilities", []), FACILITY_MANAGED + ("최대보유량(kg)", "시설번호", "시설명", "제품명"), removed_facility),
               (WORKSPACE_FACILITIES_KEY, "facility", plans.get("inventory.facilities", []), FACILITY_MANAGED, removed_facility)]
    for key, kind, rows, managed, removed in targets:
        record = project.get_field(key)
        if key in (CHEMICAL_DETAILS_KEY, WORKSPACE_FACILITIES_KEY) and record is None:
            continue        # 별지 작성대 표가 아직 없으면 만들지 않습니다(작성대는 inventory.*를 읽음)
        if not rows and record is None:
            continue
        existing = record.value if record and isinstance(record.value, list) else []
        out[key] = _merge(kind, existing, rows, managed, removed, pending[kind])
    for key, (value, _, _) in catalog_preview(project).items():
        out[key] = value
    return out


def publish(project, *, replace_existing=False):
    """회사 파일에서 읽은 사실을 작성자료(초안)로 보냅니다. 사람의 승인 없이 반영하며, 같은 번호의 행에 합칩니다.

    충돌·값 오류로 정해지지 않은 값은 포함되지 않고, 어느 파일을 고쳐야 하는지 안내로만 남습니다."""
    extended = catalog_preview(project)
    if not publishing_plan(project) and not extended:
        raise ValueError("연결할 확인 완료 사실이 없습니다.")
    preview = publish_preview(project)
    for key, rows in preview.items():
        existing = project.get_field(key)
        if existing and existing.value != rows and not replace_existing:
            raise ValueError("기존 작성자료와 다릅니다. 비교 후 교체 항목을 확인하세요.")
    for key, rows in preview.items():
        if key == WORKSPACE_FACILITIES_KEY:
            from . import cap_workspace
            cap_workspace.save_facility_rows(project, rows)      # 최대보유량 다시 계산(그래프와 이미 같아 되돌려 적을 값 없음)
        elif key == CHEMICAL_DETAILS_KEY:
            old = project.get_field(key)
            project.set_field(key, old.label if old else "화학물질 상세(별지 제6호 작성대)", rows, "USER_CONFIRMED",
                              note="회사 파일에서 자동 기입(초안): 물질명·CAS·함량·성상·비중을 같은 CAS 행에 합침")
        elif key in extended:
            _, facts, label = extended[key]
            old = project.get_field(key)
            project.set_field(key, label, rows, "USER_CONFIRMED", evidence=_refs(facts) or (list(old.evidence) if old else None),
                              note=LINK_NOTE + ": 사실 카탈로그 항목을 파일에서 읽은 그대로 적음")
        else:
            project.set_field(key, "공통정보에서 연결한 " + key, rows, "USER_CONFIRMED",
                              note=LINK_NOTE + ": 출처와 필드별 결정은 workspace.evidence_graph에 보존")
    return list(preview)


def needs_publish(project):
    """보낼 값이 있는데 작성자료가 그와 다르거나 아직 확정 상태가 아니면 True(파일을 다시 올린 뒤 등)."""
    for key, value in publish_preview(project).items():
        record = project.get_field(key)
        if record is None or record.value != value or record.status not in CONFIRMED_STATUSES:
            return True
    return False


def report(project):
    from .scenario_triage import evaluate
    lines = ["# 사업장 공통정보 검토", "", project.company_name, "", "## 확인한 사실"]
    for slot in slots(project):
        if slot["chosen"]:
            f = slot["chosen"]
            how = DECISION_LABEL.get(slot["method"], "확인 완료")
            lines.append(f"- {slot['entity']} / {LABELS[slot['field']]}: {f['value']} [{how}] "
                         f"({f['document']} · {f['location']} · SHA256 {f['sha256']})")
    info = summary(project)
    lines += ["", "## 자동 기입 현황", "",
              f"- 자동 기입 {info['auto']}개, 사람이 확인 {info['person']}개, 남은 확인 {info['pending']}개",
              f"- 자동 기입 중 사람이 아직 보지 않은 핵심값 {info['critical_auto']}개"]
    lines += ["", "## 남은 확인 과제"] + [f"- {t['entity']}: {t['task']}" for t in tasks(project)]
    triage = evaluate(project)
    lines += ["", "## 물질×설비 예비 검토 순서", "", triage["method"]]
    for row in triage["rows"]:
        lines.append(f"- {row['facility']} × {row['material']}: 순위 {row['rank'] if row['rank'] is not None else '취급 관계 미확인'}, "
                     f"지표 {row['score']:g}, 상태 {row['state']}, 가정 {', '.join(row['assumptions']) or '없음'}, "
                     f"핵심값 {', '.join(row['rank_critical']) or '없음'}, 보류 {', '.join(row['blockers']) or '없음'}")
    lines += ["", "## 순위를 바꾸는 결측 질문"] + [f"- {q['facility']} × {q['material']}: {q['question']}" for q in triage["questions"]]
    lines += ["", "이 문서는 자료 검토용입니다. 기본 운전조건 외의 시나리오 선정조건·안전설비·누출조건·물성·계산모델은 기존 계산 화면에서 별도로 검토합니다."]
    return "\n".join(lines)


def scenario_packets(project):
    packets = []
    all_slots = slots(project)
    task_list = tasks(project)
    for entity, values in confirmed_entities(project).items():
        if not entity.startswith("facility:"):
            continue
        packets.append({"설비": entity, "확인한 기본조건": values,
                        "연결된 출처": [{"항목": LABELS[s["field"]], "문서": s["chosen"]["document"], "위치": s["chosen"]["location"]}
                                    for s in all_slots if s["entity"] == entity and s["chosen"]],
                        "자료 확인 과제": [t["task"] for t in task_list if t["entity"] == entity],
                        "별도 시나리오 검토": ["설비·배관 연결관계와 누출 위치", "물성·상태와 적용 모델", "누출 크기·지속시간", "안전설비·차단 조건", "사고유형과 선정 근거"],
                        "상태": "입력 검토용 — 시나리오 선정·계산 미수행"})
    return packets


def demo_zip():
    """가상 회사 시험자료. 모든 값은 허구이며 실제 운전·평가에 쓸 수 없습니다."""
    from openpyxl import Workbook
    files = {"01_materials.csv": "제품코드,제품명,CAS 번호,함량(%)\nP-01,가상 에탄올 제품,64-17-5,99\nP-02,가상 메탄올 제품,67-56-1,100\n",
             "02_equipment.csv": "설비번호,설비명,취급제품,용량(L),최대보유량(kg),운전압력(kPa),운전온도(C)\n"
                                 "TK-01,가상 저장탱크,가상 에탄올 제품,10000,7000,100,25\n"
                                 "TK-02,가상 보조탱크,가상 메탄올 제품,5000,3500,100,25\n",
             "03_old_equipment.csv": "설비번호,용량(m3)\nTK-01,12\n"}
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "설비대장"
    sheet.append(["가상공장 설비대장 (작성 기준일 2026-03)", "", "", "", ""])
    sheet.append(["", "", "", "", ""])
    sheet.append(["설비번호", "설비명", "취급제품", "최대보유량(kg)", "운전압력(MPa)"])
    sheet.append(["RX-01", "가상 반응기", "가상 에탄올 제품", "=1500+500", 0.3])
    sheet.append(["RX-02", "가상 예비 반응기", "가상 에탄올 제품", None, 0.3])
    buffer = BytesIO()
    workbook.save(buffer)
    out = BytesIO()
    with ZipFile(out, "w", ZIP_DEFLATED) as z:
        for name, text in files.items():
            z.writestr(name, text.encode("utf-8-sig"))
        z.writestr("04_ledger_with_title.xlsx", buffer.getvalue())
        z.writestr("README.txt", "모든 회사·설비·운전값은 가상 시험자료입니다. 물질 파일(01)은 제품·물질 자료로, 나머지는 설비 자료로 올리세요. "
                                 "02와 03의 TK-01 용량 10m3/12m3는 의도적인 충돌입니다. 04는 제목 줄과 빈 줄 아래에 머리글이 있고 "
                                 "수식 셀과 빈 칸이 있어, 실제 대장처럼 읽히는지 확인하는 용도입니다.")
    return out.getvalue()
