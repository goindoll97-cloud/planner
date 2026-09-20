from __future__ import annotations

"""PSM 제출·확인 단계 행정서식(별지 제1호·제9호).

보고서 본문 별지 제12~21호와 분리한다.
- 별지 제1호: 최초 심사신청 단계
- 별지 제9호: 심사 후 확인요청 단계

회사에서 이미 확인한 공통 사실은 재사용하고, 단계별로 새로 필요한 사실만
별도 field에 저장한다. 확인되지 않은 값은 추정하거나 현재 날짜로 자동 확정하지 않는다.
"""

from dataclasses import dataclass
from io import BytesIO
import re
from typing import Mapping

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.shared import Pt

from .ooxml_determinism import canonicalize_docx_zip
from .project import Stage2Project


REGULATION_SOURCE = (
    "공정안전보고서의 제출·심사·확인 및 이행상태평가 등에 관한 규정 "
    "[시행 2025.5.30.] [고용노동부고시 제2025-30호]"
)
FORM1_SOURCE = "별지 제1호서식 공정안전보고서 심사신청서 <개정 2020.1.16>"
FORM9_SOURCE = "별지 제9호서식 공정안전보고서확인요청서 <개정 2020.1.16>"

FORM1_KEYS = {
    "사업장명": ("business.site_name", "business.company_name"),
    "사업장관리번호": ("psm.admin.workplace_management_no",),
    "사업자등록번호": ("business.registration_no",),
    "전화번호": ("business.phone",),
    "소재지": ("business.address",),
    "대표자 성명": ("business.representative",),
    "신청일": ("psm.admin.form1.application_date",),
}

FORM9_KEYS = {
    "사업장명": ("business.site_name", "business.company_name"),
    "사업자등록번호": ("business.registration_no",),
    "사업장관리번호": ("psm.admin.workplace_management_no",),
    "전화번호": ("business.phone",),
    "소재지": ("business.address",),
    "대표자 성명": ("business.representative",),
    "담당자 성명": ("psm.admin.contact_name",),
    "담당자 휴대전화번호": ("psm.admin.contact_mobile",),
    "담당자 전자우편 주소": ("psm.admin.contact_email",),
    "확인대상 사업 또는 설비명": ("psm.admin.confirmation_target", "psm.business.target_facility"),
    "공정안전보고서 심사완료일": ("psm.admin.review_completion_date",),
    "공사기간": ("psm.admin.construction_period", "psm.business.total_period"),
    "확인요청일": ("psm.admin.confirmation_request_date",),
    "확인요청 기간 시작": ("psm.admin.confirmation_period_start",),
    "확인요청 기간 종료": ("psm.admin.confirmation_period_end",),
    "신청일": ("psm.admin.form9.application_date",),
}

ADMIN_LABELS = {
    "psm.admin.workplace_management_no": "사업장관리번호",
    "psm.admin.form1.application_date": "심사신청일",
    "psm.admin.contact_name": "확인요청 담당자 성명",
    "psm.admin.contact_mobile": "확인요청 담당자 휴대전화번호",
    "psm.admin.contact_email": "확인요청 담당자 전자우편 주소",
    "psm.admin.confirmation_target": "확인대상 사업 또는 설비명",
    "psm.admin.review_completion_date": "공정안전보고서 심사완료일",
    "psm.admin.construction_period": "공사기간",
    "psm.admin.confirmation_request_date": "확인요청일",
    "psm.admin.confirmation_period_start": "확인요청 기간 시작",
    "psm.admin.confirmation_period_end": "확인요청 기간 종료",
    "psm.admin.form9.application_date": "확인요청서 신청일",
}


@dataclass(frozen=True)
class AdminFormReadiness:
    form_no: str
    title: str
    values: Mapping[str, str]
    blockers: tuple[str, ...]
    manual_items: tuple[str, ...] = ()

    @property
    def ready(self) -> bool:
        return not self.blockers


def _clean(value: object) -> str:
    return "" if value is None else str(value).strip()


def _value(project: Stage2Project, *keys: str) -> str:
    for key in keys:
        record = project.get_field(key)
        if record is not None and record.confirmed:
            value = _clean(record.value)
            if value:
                return value
    return ""


def _current(project: Stage2Project, mapping: Mapping[str, tuple[str, ...]]) -> dict[str, str]:
    return {label: _value(project, *keys) for label, keys in mapping.items()}


def form1_values(project: Stage2Project) -> dict[str, str]:
    return _current(project, FORM1_KEYS)


def form9_values(project: Stage2Project) -> dict[str, str]:
    return _current(project, FORM9_KEYS)


def save_admin_values(project: Stage2Project, values: Mapping[str, object]) -> None:
    """Save only administrative-stage facts; never overwrite shared business facts."""
    for key, label in ADMIN_LABELS.items():
        if key not in values:
            continue
        value = _clean(values.get(key))
        if not value:
            continue
        project.set_field(key, label, value, "USER_CONFIRMED")


def _readiness(
    project: Stage2Project,
    form_no: str,
    title: str,
    mapping: Mapping[str, tuple[str, ...]],
    *,
    manual_items: tuple[str, ...] = (),
) -> AdminFormReadiness:
    values = _current(project, mapping)
    blockers = tuple(
        f"{title}: '{label}' 값이 확인되지 않았습니다."
        for label, value in values.items()
        if not value
    )
    return AdminFormReadiness(
        form_no=form_no,
        title=title,
        values=values,
        blockers=blockers,
        manual_items=manual_items,
    )


def form1_readiness(project: Stage2Project) -> AdminFormReadiness:
    return _readiness(
        project,
        "1",
        "공정안전보고서 심사신청서",
        FORM1_KEYS,
        manual_items=(
            "공정안전보고서 2부를 제출자료로 준비합니다.",
            "수수료는 제출 시점의 고용노동부 고시·공단 안내를 확인합니다.",
            "출력 후 신청인 서명 또는 날인을 확인합니다.",
        ),
    )


def form9_readiness(project: Stage2Project) -> AdminFormReadiness:
    return _readiness(
        project,
        "9",
        "공정안전보고서확인요청서",
        FORM9_KEYS,
        manual_items=(
            "별지 제9호는 공정안전보고서 심사 후 확인을 요청하는 단계에서 사용합니다.",
            "출력 후 신청인 서명 또는 날인을 확인합니다.",
        ),
    )


def _set_font(paragraph, size: int = 10, bold: bool = False) -> None:
    for run in paragraph.runs:
        run.font.name = "Malgun Gothic"
        run.font.size = Pt(size)
        run.bold = bold


def _write(cell, text: object, *, bold: bool = False, center: bool = False) -> None:
    cell.text = _clean(text)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for p in cell.paragraphs:
        if center:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _set_font(p, 9, bold)


def _title(doc: Document, form_label: str, title: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run(form_label)
    r.font.name = "Malgun Gothic"
    r.font.size = Pt(9)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(title)
    r.bold = True
    r.font.name = "Malgun Gothic"
    r.font.size = Pt(16)


def _source_note(doc: Document, source: str) -> None:
    p = doc.add_paragraph()
    r = p.add_run(
        f"작성근거: {REGULATION_SOURCE} / {source}\n"
        "이 파일은 현행 규정서식의 확인된 항목을 바탕으로 프로그램이 작성한 지원본입니다. "
        "제출 전 국가법령정보센터 원문 및 실제 제출기관 안내와 최종 대조하세요."
    )
    r.font.name = "Malgun Gothic"
    r.font.size = Pt(8)


def _new_doc() -> Document:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = section.bottom_margin = Pt(35)
    section.left_margin = section.right_margin = Pt(35)
    return doc


def build_form1_docx(project: Stage2Project) -> bytes:
    readiness = form1_readiness(project)
    if not readiness.ready:
        raise ValueError("별지 제1호 심사신청서 필수값이 확인되지 않아 작성을 중단합니다: " + " / ".join(readiness.blockers))
    v = readiness.values
    doc = _new_doc()
    _title(doc, "[별지 제1호서식] <개정 2020.1.16>", "공정안전보고서 심사신청서")

    receipt = doc.add_table(rows=1, cols=4)
    for cell, value in zip(receipt.rows[0].cells, ("접수번호", "접수일자", "처리일자", "처리기간 30일")):
        _write(cell, value, center=True)

    table = doc.add_table(rows=4, cols=4)
    data = (
        ("신청인", "사업장명", v["사업장명"], f"사업장관리번호\n{v['사업장관리번호']}"),
        ("신청인", "사업자등록번호", v["사업자등록번호"], f"전화번호\n{v['전화번호']}"),
        ("신청인", "소재지", v["소재지"], ""),
        ("신청인", "대표자 성명", v["대표자 성명"], ""),
    )
    for row, values in zip(table.rows, data):
        for cell, value in zip(row.cells, values):
            _write(cell, value)

    p = doc.add_paragraph()
    p.add_run("「산업안전보건법」 제44조제1항에 따라 공정안전보고서 심사를 신청합니다.")
    _set_font(p, 10)
    p = doc.add_paragraph(f"{v['신청일']}\n신청인  {v['대표자 성명']}  (서명 또는 인)\n\n한국산업안전보건공단 이사장 귀하")
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    _set_font(p, 10)

    docs = doc.add_table(rows=1, cols=3)
    _write(docs.cell(0, 0), "신청인 제출서류", bold=True, center=True)
    _write(docs.cell(0, 1), "1. 공정안전보고서 2부")
    _write(docs.cell(0, 2), "수수료\n고용노동부장관이 정하는 수수료 참조", center=True)

    p = doc.add_paragraph("처리절차(안전보건공단, 예방센터)\n신청서 작성 → 접수 → 서류검토 → 심사 → 결과통지")
    _set_font(p, 9)
    _source_note(doc, FORM1_SOURCE)

    out = BytesIO()
    doc.save(out)
    return canonicalize_docx_zip(out.getvalue())


def build_form9_docx(project: Stage2Project) -> bytes:
    readiness = form9_readiness(project)
    if not readiness.ready:
        raise ValueError("별지 제9호 확인요청서 필수값이 확인되지 않아 작성을 중단합니다: " + " / ".join(readiness.blockers))
    v = readiness.values
    doc = _new_doc()
    _title(doc, "[별지 제9호서식] <개정 2020.1.16>", "공정안전보고서확인요청서")

    table = doc.add_table(rows=10, cols=4)
    rows = (
        ("사업장명", v["사업장명"], "사업자등록번호", v["사업자등록번호"]),
        ("사업장관리번호", v["사업장관리번호"], "전화번호", v["전화번호"]),
        ("소재지", v["소재지"], "", ""),
        ("대표자 성명", v["대표자 성명"], "", ""),
        ("담당자", v["담당자 성명"], "휴대전화번호", v["담당자 휴대전화번호"]),
        ("전자우편 주소", v["담당자 전자우편 주소"], "", ""),
        ("확인대상 사업 또는 설비명", v["확인대상 사업 또는 설비명"], "", ""),
        ("공정안전보고서 심사완료일", v["공정안전보고서 심사완료일"], "", ""),
        ("공사기간", v["공사기간"], "확인요청일", v["확인요청일"]),
        ("확인요청 기간", f"{v['확인요청 기간 시작']} ~ {v['확인요청 기간 종료']}", "", ""),
    )
    for row, values in zip(table.rows, rows):
        for cell, value in zip(row.cells, values):
            _write(cell, value)

    p = doc.add_paragraph()
    p.add_run("「산업안전보건법」 제46조제2항 및 같은 법 시행규칙 제53조에 따라 확인을 요청합니다.")
    _set_font(p, 10)
    p = doc.add_paragraph(f"{v['신청일']}\n신청인  {v['대표자 성명']}  (서명 또는 인)\n\n한국산업안전보건공단 이사장 귀하")
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    _set_font(p, 10)
    _source_note(doc, FORM9_SOURCE)

    out = BytesIO()
    doc.save(out)
    return canonicalize_docx_zip(out.getvalue())


def _safe_company(project: Stage2Project) -> str:
    value = re.sub(r"[^0-9A-Za-z가-힣._-]+", "_", project.company_name or project.project_id).strip("._")
    return value or "사업장"


def form1_filename(project: Stage2Project) -> str:
    return f"{_safe_company(project)}_공정안전보고서_별지제1호_심사신청서.docx"


def form9_filename(project: Stage2Project) -> str:
    return f"{_safe_company(project)}_공정안전보고서_별지제9호_확인요청서.docx"
