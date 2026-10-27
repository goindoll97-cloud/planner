from __future__ import annotations

"""점검·내보내기(화학사고예방관리계획서·공정안전보고서 공용). 기존 4단계(점검)·5단계(보고서 작성)와 같은 엔진·기준을 그대로 쓴다.

- 점검: validate_selected_scope 결과를 '지금 보완할 것'과 '최종 제출 전에 담당자가 따로 준비할 것(첨부)'으로 나눈다.
- 내보내기: 승인 절차 없이 언제든 내려받는다. 결과물은 항상 '초안'이다(규정서식 DOCX, 서술형을 포함한 내부 DOCX, 검증정보와 묶음).
  내부 DOCX 끝에는 파일을 보완하면 좋아질 곳을 안내하는 페이지가 붙는다.
"""

from dataclasses import dataclass, field
from typing import Any

from .cross_validation import CrossValidationReport
from .document_output_readiness import DocumentOutputReadiness, evaluate_document_output_readiness
from .intake import selected_requirement_specs
from .output_bundle import build_verified_output_bundle, verified_bundle_filename
from .output_provenance import build_output_provenance, output_provenance_filename, output_provenance_json_bytes
from .project import Stage2Project
from .cap_baseline_docx import build_cap_baseline_draft, cap_baseline_filename
from .psm_baseline_docx import build_psm_baseline_draft, psm_baseline_filename
from .report_draft import build_report_draft, draft_filename
from .scope_validation import validate_selected_scope
from .workflow import ATTACHMENT_MODE_MANUAL, attachment_mode, input_kind_has_attachment

SYSTEMS = {"PSM": "공정안전보고서", "CAP": "화학사고예방관리계획서"}
STATUS_LABELS = {"HOLD": "보완 필요", "REVIEW_REQUIRED": "담당자 확인 필요", "PASS": "확인 완료"}
DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


@dataclass(frozen=True)
class ExportState:
    report: CrossValidationReport
    text_issues: tuple[Any, ...]
    manual_issues: tuple[Any, ...]
    hold: int
    review: int
    passed: int
    readiness: DocumentOutputReadiness | None = None

    @property
    def can_confirm(self) -> bool:
        return self.hold == 0 and self.review == 0

    @property
    def final_ready(self) -> bool:
        return bool(self.readiness and self.readiness.final_ready)


def _attachment_fields(project: Stage2Project) -> set[str]:
    fields = {"documents.pfd", "documents.pid", "documents.site_plan", "psm.psi.msds"}
    for spec in selected_requirement_specs(project):
        if input_kind_has_attachment(spec.input_kind):
            fields.update(spec.field_keys)
    return fields


def _in_scope(project: Stage2Project, system: str) -> bool:
    return project.psm_in_scope if system == "PSM" else project.cap_in_scope


def evaluate(project: Stage2Project, system: str = "PSM") -> ExportState:
    """기존 4단계 화면과 같은 분류. 첨부가 본체인 항목은 '따로 준비'로 분리한다(첨부를 직접 관리하는 모드가 아닐 때)."""
    if system not in SYSTEMS:
        raise ValueError(f"지원하지 않는 보고서 종류입니다: {system}")
    if not _in_scope(project, system):
        raise ValueError(f"현재 작성범위에 {SYSTEMS[system]}이(가) 포함되어 있지 않습니다.")
    report = validate_selected_scope(project)
    manual_mode = attachment_mode(project) == ATTACHMENT_MODE_MANUAL
    attachments = _attachment_fields(project)

    def deferred(issue) -> bool:
        if not manual_mode:
            return False
        keys = set(issue.field_keys)
        if not keys.intersection(attachments):
            return False
        return issue.code == "CROSSCHECK-SOURCE-MISSING" or any(key in attachments for key in keys)

    manual = tuple(i for i in report.issues if deferred(i))
    text = tuple(i for i in report.issues if not deferred(i))
    return ExportState(
        report=report, text_issues=text, manual_issues=manual,
        hold=sum(1 for i in text if i.status == "HOLD"), review=sum(1 for i in text if i.status == "REVIEW_REQUIRED"),
        passed=sum(1 for i in text if i.status == "PASS"),
        readiness=evaluate_document_output_readiness(project, system, report),
    )


@dataclass(frozen=True)
class Outputs:
    regulation_docx: bytes
    regulation_name: str
    review_docx: bytes
    review_name: str
    provenance_json: bytes
    provenance_name: str
    bundle: bytes
    bundle_name: str
    notes: tuple[str, ...] = field(default_factory=tuple)


def ai_review_sheet_name(project: Stage2Project, system: str) -> str:
    return f"{project.project_id}_{SYSTEMS[system]}_AI작성항목_점검표.docx"


def build_ai_review_sheet(project: Stage2Project, system: str) -> bytes:
    """AI가 쓴 항목을 제출 전에 사람이 한 번 더 점검하도록 만든 점검표(규정서식 본문에는 AI 표시를 넣지 않는다)."""
    from io import BytesIO

    from docx import Document

    from .ai_drafting import ai_written_items

    rows = ai_written_items(project, system)
    doc = Document()
    doc.add_heading(f"{SYSTEMS[system]} AI 작성 항목 점검표", level=1)
    doc.add_paragraph(f"프로젝트: {project.company_name} ({project.project_id})")
    doc.add_paragraph("아래 항목의 문장은 로컬 AI가 초안을 쓰고 담당자가 확인·승인한 것입니다. 제출 전에 사실·숫자·설비·직책을 한 번 더 대조하고 점검란에 표시하세요.")
    if not rows:
        doc.add_paragraph("AI가 작성한 항목이 없습니다.")
    else:
        table = doc.add_table(rows=1, cols=6)
        table.style = "Table Grid"
        for cell, text in zip(table.rows[0].cells, ("항목", "상태", "AI 모델", "담당자 수정", "[확인 필요] 표시", "재점검(☐)")):
            cell.text = text
        for row in rows:
            cells = table.add_row().cells
            cells[0].text = f"{row['section']} {row['label']}".strip()
            cells[1].text = row["state"]
            cells[2].text = row["model"] or "-"
            cells[3].text = "수정함" if row["edited_by_reviewer"] else "수정 없음"
            cells[4].text = "있음(채워야 함)" if row["has_check_mark"] else "없음"
            cells[5].text = "☐"
    out = BytesIO()
    doc.save(out)
    return out.getvalue()


def _append_open_items(docx: bytes, project: Stage2Project, system: str) -> bytes:
    """문서 끝에 ‘파일을 보완하면 좋아질 곳’ 안내 페이지를 붙인다(점검할 수 없으면 그대로 둔다)."""
    from io import BytesIO

    from docx import Document

    try:
        state = evaluate(project, system)
    except ValueError:
        return docx
    open_issues = [i for i in state.text_issues if i.status != "PASS"]
    if not open_issues and not state.manual_issues:
        return docx
    doc = Document(BytesIO(docx))
    doc.add_page_break()
    doc.add_heading("프로그램 보완 안내 (초안)", level=1)
    doc.add_paragraph("회사 파일에서 확인되지 않았거나 서로 달라 비워 둔 곳입니다. 해당 파일을 고쳐 다시 올리면 다음 초안에 반영됩니다. "
                      "이 문서를 직접 고쳐 쓰셔도 됩니다.")
    if open_issues:
        table = doc.add_table(rows=1, cols=4)
        try:
            table.style = "Table Grid"
        except KeyError:        # 규정서식 기반 문서에는 이 표 스타일이 없을 수 있다
            pass
        for cell, text in zip(table.rows[0].cells, ("상태", "작성항목", "확인할 내용", "근거")):
            cell.text = text
        for issue in open_issues:
            cells = table.add_row().cells
            cells[0].text = STATUS_LABELS.get(issue.status, issue.status_label)
            cells[1].text = str(issue.legal_item)
            cells[2].text = str(issue.message)
            cells[3].text = str(issue.legal_basis)
    if state.manual_issues:
        doc.add_heading("따로 준비할 자료", level=2)
        for issue in state.manual_issues:
            doc.add_paragraph(f"• {issue.legal_item}: {issue.message}")
    out = BytesIO()
    doc.save(out)
    return out.getvalue()


def build_outputs(project: Stage2Project, system: str = "PSM") -> Outputs:
    """결과물은 항상 초안이다. 규정서식·검증정보는 ‘초안’으로 표기하고, 내부 DOCX 끝에 보완 안내를 붙인다."""
    if system == "CAP":
        regulation, name = build_cap_baseline_draft(project), cap_baseline_filename(project)
    elif system == "PSM":
        regulation, name = build_psm_baseline_draft(project), psm_baseline_filename(project)
    else:
        raise ValueError(f"지원하지 않는 보고서 종류입니다: {system}")
    name = name.replace("_규정서식_작성본.docx", "_규정서식_초안.docx")
    provenance = build_output_provenance(project, system, regulation, file_name=name, final_ready=False)
    review = _append_open_items(build_report_draft(project, system), project, system)
    review_name = draft_filename(project, system).replace("_법정서식_검토용_초안.docx", "_내부_초안.docx").replace(
        "_검토용_초안.docx", "_내부_초안.docx")
    return Outputs(
        regulation_docx=regulation, regulation_name=name, review_docx=review, review_name=review_name,
        provenance_json=output_provenance_json_bytes(provenance), provenance_name=output_provenance_filename(name),
        bundle=build_verified_output_bundle(regulation, provenance), bundle_name=verified_bundle_filename(name),
    )
