from __future__ import annotations

"""CAS-based KOSHA MSDS reference support for Stage 2.

The KOSHA result is deliberately stored as a HOLD reference record. It is not a
company fact and never satisfies the statutory requirement for the actual
supplier/manufacturer/importer product MSDS. Existing company SDS files can be
checked locally for CAS coverage without uploading their contents anywhere.
"""

from dataclasses import dataclass
from io import BytesIO
from pathlib import Path
from typing import Any, Callable, Iterable, Mapping
import re
from zipfile import ZIP_DEFLATED, ZipFile

import pdfplumber
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt

from engine.kosha_msds import KOSHAFullMSDSResult, MSDS_SECTIONS, lookup_full_msds_by_cas

from .project import EvidenceRef, Stage2Project


REFERENCE_PREFIX = "reference.kosha_msds."
SUPPLIER_COMPARISON_KEY = "review.msds_supplier_comparison"
CAS_SCAN_RE = re.compile(r"(?<!\d)(\d{2,7}-\d{2}-\d)(?!\d)")


@dataclass(frozen=True)
class ChemicalCAS:
    cas: str
    chemical_name: str = ""
    product_name: str = ""


@dataclass(frozen=True)
class MSDSRefreshItem:
    cas: str
    chemical_name: str
    status: str
    message: str
    section_count: int


@dataclass(frozen=True)
class SupplierSDSComparison:
    status: str
    expected_cas: tuple[str, ...]
    found_cas: tuple[str, ...]
    matched_cas: tuple[str, ...]
    missing_cas: tuple[str, ...]
    files_checked: tuple[str, ...]
    unreadable_files: tuple[str, ...]
    message: str

    def to_dict(self) -> dict[str, Any]:
        return {
            "status": self.status,
            "expected_cas": list(self.expected_cas),
            "found_cas": list(self.found_cas),
            "matched_cas": list(self.matched_cas),
            "missing_cas": list(self.missing_cas),
            "files_checked": list(self.files_checked),
            "unreadable_files": list(self.unreadable_files),
            "message": self.message,
        }


def _text(value: object) -> str:
    return "" if value is None else str(value).strip()


def _row_value(row: Mapping[str, object], *aliases: str) -> str:
    normalized = {re.sub(r"[^0-9A-Za-z가-힣]", "", str(k)).lower(): v for k, v in row.items()}
    for alias in aliases:
        key = re.sub(r"[^0-9A-Za-z가-힣]", "", alias).lower()
        value = normalized.get(key)
        if value not in (None, ""):
            return _text(value)
    return ""


def extract_cas_numbers(value: object) -> tuple[str, ...]:
    return tuple(dict.fromkeys(match.group(1) for match in CAS_SCAN_RE.finditer(_text(value))))


def inventory_chemicals(project: Stage2Project) -> list[ChemicalCAS]:
    """Return unique CAS identities from company-confirmed chemical tables."""
    rows: list[Mapping[str, object]] = []
    for key in ("inventory.chemicals", "cap.chemical.details"):
        record = project.get_field(key)
        if record is None or not isinstance(record.value, list):
            continue
        rows = [dict(row) for row in record.value if isinstance(row, Mapping)]
        if rows:
            break

    result: list[ChemicalCAS] = []
    seen: set[str] = set()
    for row in rows:
        cas_value = _row_value(row, "CAS 번호", "CAS No.", "CAS No", "CAS", "화학물질식별번호(CAS 번호)", "화학물질식별번호")
        cas_numbers = extract_cas_numbers(cas_value)
        if not cas_numbers:
            continue
        chemical_name = _row_value(row, "물질명", "화학물질명", "유해화학물질명", "성분명")
        product_name = _row_value(row, "제품명", "상품명")
        for cas in cas_numbers:
            if cas in seen:
                continue
            seen.add(cas)
            result.append(ChemicalCAS(cas, chemical_name, product_name))
    return result


def reference_field_key(cas: str) -> str:
    return REFERENCE_PREFIX + re.sub(r"\D", "", cas)


def _reference_evidence(result: KOSHAFullMSDSResult) -> list[EvidenceRef]:
    return [
        EvidenceRef(
            source_type="PUBLIC_REFERENCE",
            source_name=result.source,
            location=f"data.go.kr dataset {result.source_dataset_id} / CAS {result.cas}",
            note="CAS 기반 KOSHA MSDS 작성·검토 참고자료. 회사/제품 법정 MSDS를 대체하지 않음.",
        )
    ]


def refresh_msds_references(
    project: Stage2Project,
    *,
    cas_numbers: Iterable[str] | None = None,
    lookup: Callable[[str], KOSHAFullMSDSResult] = lookup_full_msds_by_cas,
) -> tuple[MSDSRefreshItem, ...]:
    """Fetch and store KOSHA references without promoting them to company facts."""
    available = {item.cas: item for item in inventory_chemicals(project)}
    selected = list(dict.fromkeys(cas_numbers or available.keys()))
    results: list[MSDSRefreshItem] = []

    for cas in selected:
        identity = available.get(cas, ChemicalCAS(cas))
        result = lookup(cas)
        payload = result.to_reference_dict()
        payload["company_input_name"] = identity.chemical_name
        payload["company_product_name"] = identity.product_name
        status = "HOLD"
        note = (
            "KOSHA CAS 기반 MSDS 참고자료입니다. 공급자·제조자·수입자의 실제 제품 MSDS와 회사 제품정보를 확인하기 전에는 "
            "법정 제출자료나 회사 확정사실로 사용하지 않습니다."
        )
        project.set_field(
            reference_field_key(cas),
            f"KOSHA MSDS 참고자료 · {cas}",
            payload,
            status,
            evidence=_reference_evidence(result),
            note=note,
        )
        results.append(
            MSDSRefreshItem(
                cas=cas,
                chemical_name=result.chemical_name or identity.chemical_name,
                status=result.status,
                message=result.message,
                section_count=len(result.sections),
            )
        )
    return tuple(results)


def stored_msds_reference(project: Stage2Project, cas: str) -> Mapping[str, Any] | None:
    record = project.get_field(reference_field_key(cas))
    if record is None or not isinstance(record.value, Mapping):
        return None
    return record.value


def stored_msds_references(project: Stage2Project) -> list[Mapping[str, Any]]:
    rows: list[Mapping[str, Any]] = []
    for item in inventory_chemicals(project):
        value = stored_msds_reference(project, item.cas)
        if value is not None:
            rows.append(value)
    return rows


def _set_run_font(run, size: float = 9, *, bold: bool = False) -> None:
    run.font.name = "Malgun Gothic"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Malgun Gothic")
    run.font.size = Pt(size)
    run.bold = bold


def build_msds_reference_docx(project: Stage2Project, cas: str) -> bytes:
    payload = stored_msds_reference(project, cas)
    if payload is None:
        raise ValueError(f"CAS {cas}의 KOSHA MSDS 참고자료를 먼저 조회해 주세요.")

    doc = Document()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("물질안전보건자료(MSDS) 검토용 참고자료")
    _set_run_font(run, 16, bold=True)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("KOSHA CAS 조회자료 · 법정 제품 MSDS 아님")
    _set_run_font(run, 10, bold=True)

    chemical_name = _text(payload.get("chemical_name")) or _text(payload.get("company_input_name")) or "[확인 필요]"
    for label, value in (
        ("화학물질명", chemical_name),
        ("CAS 번호", cas),
        ("KOSHA 화학물질 식별자", _text(payload.get("chem_id")) or "[확인 필요]"),
        ("출처", _text(payload.get("source"))),
        ("조회시각(UTC)", _text(payload.get("checked_at_utc"))),
    ):
        p = doc.add_paragraph()
        r = p.add_run(label + ": ")
        _set_run_font(r, 9, bold=True)
        r = p.add_run(value)
        _set_run_font(r, 9)

    warning = doc.add_paragraph()
    r = warning.add_run(
        "※ 본 문서는 한국산업안전보건공단의 CAS 기반 참고정보를 자동 정리한 검토용 자료입니다. "
        "제조자·수입자·공급자가 제공하는 실제 제품 MSDS를 대체하지 않으며, 제품명·실제 함량·혼합물 조성·회사 연락처·권고용도 등 "
        "제품 및 회사 고유정보는 반드시 실제 제품자료로 확인해야 합니다."
    )
    _set_run_font(r, 8, bold=True)

    raw_sections = payload.get("sections")
    sections = raw_sections if isinstance(raw_sections, Mapping) else {}
    for number in range(1, 17):
        doc.add_heading(f"{number}. {MSDS_SECTIONS[number]}", level=1)
        section = sections.get(str(number)) or sections.get(number)
        if not isinstance(section, Mapping):
            doc.add_paragraph("[조회 자료 없음 · 공급자 MSDS 확인 필요]")
            continue
        items = section.get("items")
        if isinstance(items, list) and items:
            table = doc.add_table(rows=0, cols=2)
            table.style = "Table Grid"
            for item in items:
                if not isinstance(item, (list, tuple)) or len(item) < 2:
                    continue
                cells = table.add_row().cells
                cells[0].text = _text(item[0]) or "항목"
                cells[1].text = _text(item[1]) or "[자료 없음]"
                for cell in cells:
                    for paragraph in cell.paragraphs:
                        for cell_run in paragraph.runs:
                            _set_run_font(cell_run, 8)
        else:
            doc.add_paragraph(_text(section.get("text")) or "[조회 자료 없음 · 공급자 MSDS 확인 필요]")

    out = BytesIO()
    doc.save(out)
    return out.getvalue()


def build_msds_reference_bundle(project: Stage2Project, cas_numbers: Iterable[str] | None = None) -> bytes:
    available = [item.cas for item in inventory_chemicals(project) if stored_msds_reference(project, item.cas) is not None]
    selected = [cas for cas in dict.fromkeys(cas_numbers or available) if cas in available]
    if not selected:
        raise ValueError("다운로드할 KOSHA MSDS 참고자료가 없습니다.")

    out = BytesIO()
    with ZipFile(out, "w", compression=ZIP_DEFLATED) as archive:
        for cas in selected:
            payload = stored_msds_reference(project, cas) or {}
            name = _text(payload.get("chemical_name")) or _text(payload.get("company_input_name")) or "화학물질"
            safe_name = re.sub(r"[\\/:*?\"<>|]+", "_", name).strip() or "화학물질"
            archive.writestr(f"{safe_name}_{cas}_KOSHA_MSDS_검토용_참고자료.docx", build_msds_reference_docx(project, cas))
        archive.writestr(
            "읽어주세요.txt",
            (
                "KOSHA CAS 기반 MSDS 검토용 참고자료입니다.\n"
                "법정 제품 MSDS가 아니며 제조자·수입자·공급자가 제공하는 실제 제품 MSDS를 대체하지 않습니다.\n"
                "혼합물·UVCB·제품별 조성 및 회사 고유정보는 실제 제품 SDS/MSDS를 우선 확인하십시오.\n"
            ).encode("utf-8"),
        )
    return out.getvalue()


def _candidate_sds_evidence(project: Stage2Project) -> list[EvidenceRef]:
    refs: list[EvidenceRef] = []
    seen: set[tuple[str, str]] = set()
    for key, record in project.fields.items():
        for evidence in record.evidence:
            name = evidence.source_name.lower()
            relevant = key == "psm.psi.msds" or "msds" in name or "sds" in name or "물질안전보건" in name
            if not relevant:
                continue
            identity = (evidence.sha256, evidence.location)
            if identity in seen:
                continue
            seen.add(identity)
            refs.append(evidence)
    return refs


def _read_local_document_text(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        with pdfplumber.open(path) as pdf:
            return "\n".join(page.extract_text() or "" for page in pdf.pages)
    if suffix == ".docx":
        doc = Document(path)
        parts = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                parts.extend(cell.text for cell in row.cells)
        return "\n".join(parts)
    if suffix in {".txt", ".csv", ".tsv"}:
        return path.read_text(encoding="utf-8", errors="ignore")
    raise ValueError("현재 자동검증 지원 형식은 PDF, DOCX, TXT, CSV입니다.")


def compare_supplier_sds(project: Stage2Project) -> SupplierSDSComparison:
    """Compare CAS presence locally; never upload company SDS contents."""
    expected = tuple(item.cas for item in inventory_chemicals(project))
    refs = _candidate_sds_evidence(project)
    if not refs:
        return SupplierSDSComparison(
            "NO_FILES", expected, (), (), expected, (), (),
            "회사/공급자 MSDS 파일이 연결되어 있지 않습니다. KOSHA 참고자료만으로 법정 제품 MSDS를 대체할 수 없습니다.",
        )

    found: list[str] = []
    checked: list[str] = []
    unreadable: list[str] = []
    for ref in refs:
        path = Path(ref.location) if ref.location else None
        if path is None or not path.exists():
            unreadable.append(ref.source_name)
            continue
        try:
            text = _read_local_document_text(path)
        except Exception:
            unreadable.append(ref.source_name)
            continue
        checked.append(ref.source_name)
        found.extend(extract_cas_numbers(text))

    found_unique = tuple(dict.fromkeys(found))
    expected_set = set(expected)
    found_set = set(found_unique)
    matched = tuple(cas for cas in expected if cas in found_set)
    missing = tuple(cas for cas in expected if cas not in found_set)

    if not checked:
        status = "UNREADABLE"
        message = "연결된 MSDS 파일을 자동으로 읽지 못했습니다. 담당자가 실제 파일을 직접 확인해야 합니다."
    elif expected and not missing:
        status = "CAS_COVERED"
        message = "연결된 MSDS 문서에서 현재 화학물질 목록의 CAS 번호를 모두 확인했습니다. 이는 CAS 존재 여부 확인이며 MSDS 내용의 법적 적정성 확정을 의미하지 않습니다."
    elif matched:
        status = "PARTIAL"
        message = "일부 CAS만 연결된 MSDS에서 확인되었습니다. 혼합물의 영업비밀 성분 등 예외 가능성을 포함해 담당자 검토가 필요합니다."
    else:
        status = "NO_CAS_MATCH"
        message = "현재 목록의 CAS 번호가 연결된 MSDS 문서에서 확인되지 않았습니다. 제품·성분 대응관계를 담당자가 확인해야 합니다."

    return SupplierSDSComparison(
        status=status,
        expected_cas=expected,
        found_cas=found_unique,
        matched_cas=matched,
        missing_cas=missing,
        files_checked=tuple(checked),
        unreadable_files=tuple(unreadable),
        message=message,
    )


def store_supplier_sds_comparison(project: Stage2Project) -> SupplierSDSComparison:
    result = compare_supplier_sds(project)
    project.set_field(
        SUPPLIER_COMPARISON_KEY,
        "공급자 MSDS CAS 비교검토",
        result.to_dict(),
        "HOLD",
        note="자동 비교는 MSDS 내 CAS 존재 여부를 확인하는 보조검토입니다. 담당자의 실제 제품 MSDS 내용 확인 후 확정해야 합니다.",
    )
    return result
