from __future__ import annotations

"""CAP 제출·변경 행정서식(시행규칙 별지 제31호·제32호) 작성지원.

사용자가 제공한 2025.8.7 개정 별지 제31호·32호의 항목과 페이지 구성을
기준으로 한다. 회사에서 이미 확인한 사실은 재사용하고, 제출 단계에서만
필요한 값은 별도 field로 저장한다. 확인되지 않은 값은 추정하지 않는다.
"""

from dataclasses import dataclass
from io import BytesIO
import re
from typing import Any, Mapping

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

from .ooxml_determinism import canonicalize_docx_zip
from .project import CONFIRMED_STATUSES, Stage2Project
from . import versioning
from .cap_change_tracking import summarize_changes


FORM31_SOURCE = "화학물질관리법 시행규칙 별지 제31호서식 <개정 2025.8.7>"
FORM32_SOURCE = "화학물질관리법 시행규칙 별지 제32호서식 <개정 2025.8.7>"

SUBMISSION_NEW = "신규제출"
SUBMISSION_RESUBMIT = "재제출"
SUBMISSION_FIVE_YEAR = "5년 재제출"
SUBMISSION_AFTER_UNSUITABLE = "부적합 후 재제출"
SUBMISSION_AFTER_INSPECTION = "이행점검 부적정 재제출"
FORM31_TYPES = (
    SUBMISSION_NEW,
    SUBMISSION_RESUBMIT,
    SUBMISSION_FIVE_YEAR,
    SUBMISSION_AFTER_UNSUITABLE,
    SUBMISSION_AFTER_INSPECTION,
)
FORM32_TYPE = "변경제출"


@dataclass(frozen=True)
class CAPSubmissionFormReadiness:
    form_no: str
    title: str
    values: Mapping[str, Any]
    blockers: tuple[str, ...]

    @property
    def ready(self) -> bool:
        return not self.blockers


def _clean(value: object) -> str:
    return "" if value is None else str(value).strip()


def _value(project: Stage2Project, *keys: str) -> str:
    for key in keys:
        rec = project.get_field(key)
        if rec is not None and rec.status in CONFIRMED_STATUSES:
            value = _clean(rec.value)
            if value:
                return value
        if key == "business.company_name" and _clean(project.company_name):
            return _clean(project.company_name)
        if key == "business.site_name" and _clean(project.site_name):
            return _clean(project.site_name)
    return ""


def _checked(selected: bool, label: str) -> str:
    return f"[{'✓' if selected else ' '}] {label}"


FORM31_KEYS = {
    "상호(명칭)": ("business.company_name", "business.site_name"),
    "사업자등록번호": ("business.registration_no", "cap.business.registration_no"),
    "대표자 성명": ("business.representative", "cap.business.representative"),
    "취급시설 소재지": ("business.address",),
    "담당자 성명": ("cap.business.writer_name", "cap.business.writer_info"),
    "연락처": ("cap.business.writer_contact", "business.phone"),
    "전자우편주소": ("cap.business.writer_email",),
    "관할 유역(지방)환경관서": ("cap.submission.office",),
    "관할 합동방재센터": ("cap.submission.center",),
    "제출방법": ("cap.submission.method",),
    "제출구분": ("cap.business.submission_type",),
    "사업장 구분": ("cap.business.writing_level",),
    "영업허가 대상": ("cap.submission.business_permit",),
    "공동비상대응계획 작성·제출 여부": ("cap.business.joint_emergency_plan",),
    "검토생략 대상": ("cap.submission.review_skip",),
    "신청일": ("cap.submission.form31.application_date",),
}

FORM32_KEYS = {
    "상호(명칭)": ("business.company_name", "business.site_name"),
    "사업자등록번호": ("business.registration_no", "cap.business.registration_no"),
    "대표자 성명": ("business.representative", "cap.business.representative"),
    "취급시설 소재지": ("business.address",),
    "담당자 성명": ("cap.business.writer_name", "cap.business.writer_info"),
    "연락처": ("cap.business.writer_contact", "business.phone"),
    "전자우편 주소": ("cap.business.writer_email",),
    "제출방법": ("cap.submission.method",),
    "(최종) 적합통보 번호": ("cap.submission.final_approval_no",),
    "(최종) 적합통보 일자": ("cap.submission.final_approval_date",),
    "영업허가 대상 여부": ("cap.submission.business_permit",),
    "변경사유 구분": ("cap.submission.change_reason_type",),
    "그 밖의 사유": ("cap.submission.change_other_reason",),
    "변경 전 사업장 구분": ("cap.submission.group_before",),
    "변경 후 사업장 구분": ("cap.submission.group_after", "cap.business.writing_level"),
    "신청일": ("cap.submission.form32.application_date",),
    "비교 기준 버전": ("cap.submission.base_version_id",),
}

ADMIN_KEYS = {
    "cap.submission.office": "관할 유역(지방)환경관서",
    "cap.submission.center": "관할 합동방재센터",
    "cap.submission.method": "제출방법",
    "cap.submission.business_permit": "영업허가 대상 여부",
    "cap.submission.review_skip": "제19조의2제2항에 따른 검토생략 대상",
    "cap.submission.form31.application_date": "별지 제31호 신청일",
    "cap.submission.final_approval_no": "최종 적합통보 번호",
    "cap.submission.final_approval_date": "최종 적합통보 일자",
    "cap.submission.change_reason_type": "변경사유 구분",
    "cap.submission.change_other_reason": "그 밖의 변경사유",
    "cap.submission.group_before": "변경 전 사업장 구분",
    "cap.submission.group_after": "변경 후 사업장 구분",
    "cap.submission.form32.application_date": "별지 제32호 신청일",
    "cap.submission.base_version_id": "변경 비교 기준 버전",
    "cap.submission.oca_details": "장외영향평가서 적합 상세내용",
    "cap.submission.rmp_details": "위해관리계획서 적합 상세내용",
    "cap.submission.change_history": "화학사고예방관리계획서 변경제출 이력",
}


def save_submission_values(project: Stage2Project, values: Mapping[str, object]) -> None:
    for key, label in ADMIN_KEYS.items():
        if key not in values:
            continue
        value = values.get(key)
        if value in (None, "", [], {}):
            continue
        project.set_field(key, label, value, "USER_CONFIRMED")


def form31_values(project: Stage2Project) -> dict[str, str]:
    values = {label: _value(project, *keys) for label, keys in FORM31_KEYS.items()}
    values["사업장 구분"] = values["사업장 구분"] or project.cap_group
    return values


def form32_values(project: Stage2Project) -> dict[str, Any]:
    values: dict[str, Any] = {label: _value(project, *keys) for label, keys in FORM32_KEYS.items()}
    values["변경 후 사업장 구분"] = values["변경 후 사업장 구분"] or project.cap_group
    for key, label in (
        ("cap.submission.oca_details", "장외영향평가서 적합 상세내용"),
        ("cap.submission.rmp_details", "위해관리계획서 적합 상세내용"),
        ("cap.submission.change_history", "화학사고예방관리계획서 변경제출 이력"),
    ):
        rec = project.get_field(key)
        values[label] = rec.value if rec is not None and rec.status in CONFIRMED_STATUSES else ([] if "이력" in label else {})
    base = values.get("비교 기준 버전") or ""
    if base:
        try:
            values["변경사항 상세내용"] = summarize_changes(project, base).form32_details()
        except (FileNotFoundError, ValueError):
            values["변경사항 상세내용"] = {}
    else:
        values["변경사항 상세내용"] = {}
    return values


def _missing(values: Mapping[str, Any], labels: tuple[str, ...]) -> tuple[str, ...]:
    return tuple(label for label in labels if values.get(label) in (None, "", [], {}))


def form31_readiness(project: Stage2Project) -> CAPSubmissionFormReadiness:
    values = form31_values(project)
    blockers: list[str] = []
    submission = values.get("제출구분", "")
    if submission not in FORM31_TYPES:
        blockers.append("별지 제31호는 신규·재제출 계열에서 사용합니다. 제출구분을 확인해 주세요.")
    required = (
        "상호(명칭)", "사업자등록번호", "대표자 성명", "취급시설 소재지",
        "담당자 성명", "연락처", "전자우편주소", "관할 유역(지방)환경관서",
        "관할 합동방재센터", "제출방법", "사업장 구분", "영업허가 대상",
        "공동비상대응계획 작성·제출 여부", "검토생략 대상", "신청일",
    )
    blockers.extend(f"별지 제31호 '{label}' 값이 확인되지 않았습니다." for label in _missing(values, required))
    return CAPSubmissionFormReadiness("31", "화학사고예방관리계획서 검토신청서", values, tuple(blockers))


def form32_readiness(project: Stage2Project) -> CAPSubmissionFormReadiness:
    values = form32_values(project)
    blockers: list[str] = []
    submission = _value(project, "cap.business.submission_type")
    if submission != FORM32_TYPE:
        blockers.append("별지 제32호는 변경제출에서 사용합니다. 제출구분을 '변경제출'로 확인해 주세요.")
    required = (
        "상호(명칭)", "사업자등록번호", "대표자 성명", "취급시설 소재지",
        "담당자 성명", "연락처", "전자우편 주소", "제출방법",
        "(최종) 적합통보 번호", "(최종) 적합통보 일자", "영업허가 대상 여부",
        "변경사유 구분", "신청일", "비교 기준 버전",
    )
    blockers.extend(f"별지 제32호 '{label}' 값이 확인되지 않았습니다." for label in _missing(values, required))
    if values.get("변경사유 구분") == "사업장 구분 변경":
        blockers.extend(
            f"별지 제32호 '{label}' 값이 확인되지 않았습니다."
            for label in _missing(values, ("변경 전 사업장 구분", "변경 후 사업장 구분"))
        )
    if values.get("변경사유 구분") == "그 밖의 사유" and not values.get("그 밖의 사유"):
        blockers.append("별지 제32호 '그 밖의 사유' 내용을 확인해 주세요.")
    base = values.get("비교 기준 버전")
    if base:
        try:
            changes = versioning.diff_versions(project, str(base))
            if not changes:
                blockers.append("비교 기준 버전과 현재 작업본 사이에 변경된 값이 없습니다.")
        except FileNotFoundError:
            blockers.append("선택한 비교 기준 버전 파일을 찾을 수 없습니다.")
    return CAPSubmissionFormReadiness("32", "화학사고예방관리계획서 변경 검토신청서", values, tuple(blockers))


def _shade(cell, fill: str = "D9D9D9") -> None:
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tcPr.append(shd)
    shd.set(qn("w:fill"), fill)


def _cell_margins(cell) -> None:
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in("w:tcMar")
    if tcMar is None:
        tcMar = OxmlElement("w:tcMar")
        tcPr.append(tcMar)
    for name, value in (("top", 55), ("start", 65), ("bottom", 55), ("end", 65)):
        el = tcMar.find(qn("w:" + name))
        if el is None:
            el = OxmlElement("w:" + name)
            tcMar.append(el)
        el.set(qn("w:w"), str(value))
        el.set(qn("w:type"), "dxa")


def _write(cell, text: object, *, size: int = 8, bold: bool = False, center: bool = False, shade: str | None = None) -> None:
    cell.text = _clean(text)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    if shade:
        _shade(cell, shade)
    _cell_margins(cell)
    for p in cell.paragraphs:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        for run in p.runs:
            run.font.name = "Malgun Gothic"
            run._element.rPr.rFonts.set(qn("w:eastAsia"), "Malgun Gothic")
            run.font.size = Pt(size)
            run.bold = bold


def _borders(table) -> None:
    tblPr = table._tbl.tblPr
    borders = tblPr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tblPr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = borders.find(qn("w:" + edge))
        if el is None:
            el = OxmlElement("w:" + edge)
            borders.append(el)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "5")
        el.set(qn("w:color"), "777777")


def _p(doc, text: str = "", *, size: int = 9, bold: bool = False, align=None, before: int = 0, after: int = 0):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    if align is not None:
        p.alignment = align
    run = p.add_run(text)
    run.font.name = "Malgun Gothic"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Malgun Gothic")
    run.font.size = Pt(size)
    run.bold = bold
    return p


def _new_doc() -> Document:
    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.2)
    section.bottom_margin = Cm(1.0)
    section.left_margin = Cm(1.35)
    section.right_margin = Cm(1.35)
    normal = doc.styles["Normal"]
    normal.font.name = "Malgun Gothic"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Malgun Gothic")
    normal.font.size = Pt(9)
    return doc


def _title(doc: Document, no: int, title: str, side: str = "") -> None:
    _p(doc, f"■ 화학물질관리법 시행규칙 [별지 제{no}호서식] <개정 2025. 8. 7.>", size=7, after=5)
    _p(doc, title, size=17, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, before=3, after=4)
    p = _p(doc, "※ 바탕색이 어두운 난은 신청인이 작성하지 않습니다.", size=7)
    if side:
        p.add_run(" " * 42 + side)


def _receipt(doc: Document, *, form32: bool) -> None:
    table = doc.add_table(rows=1, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    labels = ("접수번호", "제출일", "접수일", "처리기간(30일)") if form32 else ("접수번호", "접수일", "발급일", "처리기간 30일")
    for cell, label in zip(table.rows[0].cells, labels):
        _write(cell, label, center=True, shade="C9C9C9")
    _borders(table)


def _applicant(doc: Document, values: Mapping[str, Any], *, form32: bool) -> None:
    table = doc.add_table(rows=6 if form32 else 4, cols=4)
    _borders(table)
    left = table.cell(0, 0)
    for idx in range(1, 4):
        left = left.merge(table.cell(idx, 0))
    _write(table.cell(0, 0), "신\n청\n인", size=10, center=True)
    email_label = "전자우편 주소" if form32 else "전자우편주소"
    rows = (
        ("상호(명칭)", values.get("상호(명칭)"), "성명", values.get("담당자 성명")),
        ("사업자등록번호", values.get("사업자등록번호"), "연락처" if form32 else "연락처(전화번호)", values.get("연락처")),
        ("대표자 성명", values.get("대표자 성명"), email_label, values.get(email_label)),
        ("취급시설 소재지", values.get("취급시설 소재지"), "", ""),
    )
    for idx, (a, b, c, d) in enumerate(rows):
        _write(table.cell(idx, 1), f"{a}\n{_clean(b)}")
        _write(table.cell(idx, 2), c, center=True)
        _write(table.cell(idx, 3), d)
    if form32:
        _write(table.cell(4, 0), "")
        _write(table.cell(4, 1), f"(최종) 적합통보 번호\n{values.get('(최종) 적합통보 번호', '')}")
        _write(table.cell(4, 2), f"(최종) 적합통보 일자\n{values.get('(최종) 적합통보 일자', '')}")
        _write(table.cell(4, 3), "")
        _write(table.cell(5, 0), "")
        _write(table.cell(5, 1), "영업허가 대상 여부")
        permit = values.get("영업허가 대상 여부", "")
        _write(table.cell(5, 2), f"{_checked(permit == '대상', '대상')}   {_checked(permit in {'비대상', '면제대상'}, '비대상')}")
        table.cell(5, 2).merge(table.cell(5, 3))


def build_form31_docx(project: Stage2Project) -> bytes:
    readiness = form31_readiness(project)
    if not readiness.ready:
        raise ValueError("별지 제31호 필수값이 확인되지 않아 작성을 중단합니다: " + " / ".join(readiness.blockers))
    v = readiness.values
    doc = _new_doc()
    _title(doc, 31, "화학사고예방관리계획서 검토신청서")
    _receipt(doc, form32=False)
    _p(doc, f"제출방법   {_checked(v['제출방법'] == '화학물질 종합정보시스템', '화학물질 종합정보시스템')}   {_checked(v['제출방법'] == '서면', '서면')}", size=8, before=2, after=2)
    _applicant(doc, v, form32=False)
    _p(doc, f"관할지방환경관서 :    ( {v['관할 유역(지방)환경관서']} ) 유역(지방)환경관서  /  ( {v['관할 합동방재센터']} )합동방재센터", size=8, before=2, after=2)

    table = doc.add_table(rows=5, cols=2)
    _borders(table)
    sub = v["제출구분"]
    _write(table.cell(0, 0), "제출구분")
    _write(table.cell(0, 1),
           f"{_checked(sub == SUBMISSION_NEW, SUBMISSION_NEW)}\n"
           f"{_checked(sub == SUBMISSION_RESUBMIT, SUBMISSION_RESUBMIT)}      {_checked(sub == SUBMISSION_AFTER_UNSUITABLE, SUBMISSION_AFTER_UNSUITABLE)}\n"
           f"{_checked(sub == SUBMISSION_FIVE_YEAR, SUBMISSION_FIVE_YEAR)}      {_checked(sub == SUBMISSION_AFTER_INSPECTION, SUBMISSION_AFTER_INSPECTION)}")
    group = v["사업장 구분"]
    _write(table.cell(1, 0), "사업장 구분")
    _write(table.cell(1, 1), f"{_checked('1군' in group, '1군 사업장')}      {_checked('2군' in group, '2군 사업장')}")
    permit = v["영업허가 대상"]
    _write(table.cell(2, 0), "영업허가 대상")
    _write(table.cell(2, 1), f"{_checked(permit == '대상', '영업허가 대상')}      {_checked(permit in {'비대상', '면제대상'}, '면제대상')}")
    joint = v["공동비상대응계획 작성·제출 여부"]
    _write(table.cell(3, 0), "공동비상대응계획 작성·제출 여부")
    _write(table.cell(3, 1), f"{_checked(joint in {'Y', '해당'}, '해당')}      {_checked(joint in {'N', '미해당'}, '미해당')}")
    skip = v["검토생략 대상"]
    _write(table.cell(4, 0), "제19조의2제2항에 따른 검토생략 대상")
    _write(table.cell(4, 1), f"{_checked(skip == '안전성향상계획', '안전성향상계획')}   {_checked(skip == '공정안전보고서', '공정안전보고서')}   {_checked(skip == '미해당', '미해당')}")

    _p(doc, "「화학물질관리법」 제23조제1항 및 같은 법 시행규칙 제19조제1항·제4항 또는 제9항에 따라 위와 같이 화학사고예방관리계획서의 검토를 신청합니다.", before=5)
    _p(doc, f"{v['신청일']}\n\n신청인  {v['대표자 성명']}                 (서명 또는 인)", align=WD_ALIGN_PARAGRAPH.RIGHT)
    _p(doc, "화학물질안전원장  귀하", size=12, bold=True, before=2, after=2)
    attach = doc.add_table(rows=1, cols=3)
    _borders(attach)
    _write(attach.cell(0, 0), "첨부\n서류", center=True)
    _write(attach.cell(0, 1), "1. 화학사고예방관리계획서\n2. 별지 제31호의2서식의 공동비상대응계획 작성·제출에 관한 자료(공동으로 비상대응계획을 작성하는 경우만 제출합니다)\n3. 안전성향상계획 또는 공정안전보고서에 대한 심사결과통지서(해당하는 경우만 제출합니다)\n4. 별지 제31호의3서식의 안전성향상계획·공정안전보고 변경사항 부존재 확인서(해당하는 경우만 제출합니다)", size=7)
    _write(attach.cell(0, 2), "수수료\n없음", center=True)
    proc = doc.add_table(rows=1, cols=1)
    _write(proc.cell(0, 0), "처리절차", size=10, center=True, shade="C9C9C9")
    _borders(proc)
    _p(doc, "신청서 작성   ⇒   접수   ⇒   검토   ⇒   결재   ⇒   결과서 통보", size=8, align=WD_ALIGN_PARAGRAPH.CENTER, before=3)
    _p(doc, "신청인                              화학물질안전원                              신청인", size=7, align=WD_ALIGN_PARAGRAPH.CENTER)
    _p(doc, "210㎜×297㎜[백상지 80g/㎡]", size=6, align=WD_ALIGN_PARAGRAPH.RIGHT)
    out = BytesIO()
    doc.save(out)
    return canonicalize_docx_zip(out.getvalue())


def build_form32_docx(project: Stage2Project) -> bytes:
    readiness = form32_readiness(project)
    if not readiness.ready:
        raise ValueError("별지 제32호 필수값이 확인되지 않아 작성을 중단합니다: " + " / ".join(readiness.blockers))
    v = readiness.values
    doc = _new_doc()
    _title(doc, 32, "화학사고예방관리계획서 변경 검토신청서", "(앞쪽)")
    _receipt(doc, form32=True)
    _p(doc, f"제출방법   {_checked(v['제출방법'] == '화학물질 종합정보시스템', '화학물질 종합정보시스템')}   {_checked(v['제출방법'] == '서면', '서면')}", size=8, before=2, after=2)
    _applicant(doc, v, form32=True)

    table = doc.add_table(rows=2, cols=2)
    _borders(table)
    change_type = v["변경사유 구분"]
    _write(table.cell(0, 0), "변경\n사항", center=True)
    _write(table.cell(0, 1), f"{_checked(change_type == '사업장 구분 변경', '사업장 구분 변경')}: 변경 전 ( {v.get('변경 전 사업장 구분', '')} ) 군 사업장, 변경 후 ( {v.get('변경 후 사업장 구분', '')} )군 사업장")
    _write(table.cell(1, 0), "")
    _write(table.cell(1, 1), f"{_checked(change_type == '그 밖의 사유', '그 밖의 사유')}( {v.get('그 밖의 사유', '')} )")
    _p(doc, "「화학물질관리법」 제23조제3항 및 같은 법 시행규칙 제19조제6항에 따라 위와 같이 변경된 화학사고예방관리계획서의 검토를 신청합니다.", before=5)
    _p(doc, f"{v['신청일']}\n\n신청인  {v['대표자 성명']}                 (서명 또는 인)", align=WD_ALIGN_PARAGRAPH.RIGHT)
    _p(doc, "화학물질안전원장  귀하", size=12, bold=True, before=2, after=2)
    attach = doc.add_table(rows=1, cols=3)
    _borders(attach)
    _write(attach.cell(0, 0), "첨부\n서류", center=True)
    _write(attach.cell(0, 1), "1. 변경된 화학사고예방관리계획서\n2. 별지 제31호의2서식의 공동 비상대응계획 작성·제출에 관한 자료(공동으로 비상대응계획을 작성하는 경우만 제출합니다)", size=7)
    _write(attach.cell(0, 2), "수수료\n없음", center=True)
    proc = doc.add_table(rows=1, cols=1)
    _write(proc.cell(0, 0), "처리절차", size=10, center=True, shade="C9C9C9")
    _borders(proc)
    _p(doc, "이 신청서는 아래와 같이 처리됩니다.", size=8, before=2)
    _p(doc, "신청서 작성   ⇒   접수   ⇒   검토   ⇒   결재   ⇒   결과서 통보", size=8, align=WD_ALIGN_PARAGRAPH.CENTER)
    _p(doc, "신청인                              화학물질안전원", size=7, align=WD_ALIGN_PARAGRAPH.CENTER)
    _p(doc, "210㎜×297㎜[백상지 80g/㎡]", size=6, align=WD_ALIGN_PARAGRAPH.RIGHT)

    section = doc.add_section(WD_SECTION.NEW_PAGE)
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.3)
    section.bottom_margin = Cm(1.2)
    section.left_margin = Cm(1.5)
    section.right_margin = Cm(1.5)
    _p(doc, "(뒤쪽)", size=8, align=WD_ALIGN_PARAGRAPH.RIGHT, after=7)
    _p(doc, "<장외영향평가서 및 위해관리계획서 적합 상세내용>", size=10, after=2)

    oca = v.get("장외영향평가서 적합 상세내용") or {}
    table = doc.add_table(rows=2, cols=4)
    _borders(table)
    for cell, label in zip(table.rows[0].cells, ("장외영향평가서 민원번호", "결과번호", "적합날짜", "위험도")):
        _write(cell, label, center=True, shade="C9C9C9")
    for idx, key in enumerate(("민원번호", "결과번호", "적합날짜", "위험도")):
        _write(table.cell(1, idx), oca.get(key, ""), center=True)

    rmp = v.get("위해관리계획서 적합 상세내용") or {}
    table = doc.add_table(rows=2, cols=3)
    _borders(table)
    for cell, label in zip(table.rows[0].cells, ("위해관리계획서 민원번호", "결과번호", "적합날짜")):
        _write(cell, label, center=True, shade="C9C9C9")
    for idx, key in enumerate(("민원번호", "결과번호", "적합날짜")):
        _write(table.cell(1, idx), rmp.get(key, ""), center=True)
    _p(doc, "※ 해당 없을 경우 “해당없음”으로 기재합니다.", size=7, before=2, after=5)

    _p(doc, "<화학사고예방관리계획서 변경제출 이력>", size=10, after=2)
    history = list(v.get("화학사고예방관리계획서 변경제출 이력") or [])
    row_count = max(5, len(history))
    table = doc.add_table(rows=1 + row_count, cols=5)
    _borders(table)
    headers = ("적합 결과번호", "사업장 구분\n(1군/2군)", "위험도", "상세내용", "적합통보일")
    for cell, label in zip(table.rows[0].cells, headers):
        _write(cell, label, center=True, shade="C9C9C9")
    for ridx, row in enumerate(history[:row_count], start=1):
        for cidx, key in enumerate(("적합 결과번호", "사업장 구분", "위험도", "상세내용", "적합통보일")):
            _write(table.cell(ridx, cidx), row.get(key, ""))

    _p(doc, "<화학사고예방관리계획서 변경사항 상세내용>", size=10, before=6, after=2)
    details = v.get("변경사항 상세내용") or {}
    table = doc.add_table(rows=4, cols=3)
    _borders(table)
    for cell, label in zip(table.rows[0].cells, ("변경구분", "변경 전", "변경 후")):
        _write(cell, label, center=True, shade="C9C9C9")
    for ridx, label in enumerate(("유해화학물질추가", "시설추가", "취급저장량 증가"), start=1):
        pair = details.get(label, {})
        _write(table.cell(ridx, 0), label, center=True)
        _write(table.cell(ridx, 1), pair.get("변경 전", ""))
        _write(table.cell(ridx, 2), pair.get("변경 후", ""))

    out = BytesIO()
    doc.save(out)
    return canonicalize_docx_zip(out.getvalue())


def _safe_company(project: Stage2Project) -> str:
    value = re.sub(r"[^0-9A-Za-z가-힣._-]+", "_", project.company_name or project.project_id).strip("._")
    return value or "사업장"


def form31_filename(project: Stage2Project) -> str:
    return f"{_safe_company(project)}_별지제31호_화학사고예방관리계획서_검토신청서.docx"


def form32_filename(project: Stage2Project) -> str:
    return f"{_safe_company(project)}_별지제32호_화학사고예방관리계획서_변경검토신청서.docx"
