from __future__ import annotations

"""혼합제품이 있을 때만 사용하는 두 번째 파일.

기본 물질목록에는 제품 단위 정보만 받고, 이 파일에는 제품 SDS 제3항의
구성성분 CAS No.와 함량(%)만 받는다. 물질명으로 법적 물질을 추정하지 않는다.
"""

from dataclasses import dataclass, field
from io import BytesIO
import hashlib
import re
from typing import Any, Iterable

import pandas as pd

from .cap_chemical_upload import cas_valid


COMPONENT_COLUMNS = ("혼합제품명", "CAS No.", "함량(%)")
ALIASES = {
    "혼합제품명": ("혼합제품명", "제품명", "혼합물명", "품명"),
    "CAS No.": ("CAS No.", "CAS번호", "CAS", "casno", "cas number"),
    "함량(%)": ("함량(%)", "함량", "농도", "함유량", "성분함량"),
}


@dataclass
class ComponentCheck:
    rows: list[dict[str, Any]] = field(default_factory=list)
    errors: list[str] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)
    sha256: str = ""
    file_name: str = ""

    @property
    def ok(self) -> bool:
        return not self.errors and bool(self.rows)


def _clean(value: object) -> str:
    if value is None or (isinstance(value, float) and value != value):
        return ""
    text = str(value).strip()
    return "" if text.lower() == "nan" else text


def _norm(value: object) -> str:
    return re.sub(r"[\s_./·\-()%]", "", _clean(value)).lower()


def _read(data: bytes, file_name: str) -> pd.DataFrame:
    if not data:
        raise ValueError("빈 파일입니다.")
    if file_name.lower().endswith((".xlsx", ".xlsm")):
        raw = pd.read_excel(BytesIO(data), header=None, dtype=object, engine="openpyxl")
    elif file_name.lower().endswith(".csv"):
        raw = None
        for encoding in ("utf-8-sig", "cp949"):
            try:
                raw = pd.read_csv(BytesIO(data), header=None, dtype=object, encoding=encoding, sep=None, engine="python")
                break
            except (UnicodeDecodeError, pd.errors.ParserError):
                continue
        if raw is None:
            raise ValueError("CSV 파일을 읽지 못했습니다.")
    else:
        raise ValueError("엑셀(.xlsx) 또는 CSV 파일만 사용할 수 있습니다.")
    raw = raw.dropna(how="all").reset_index(drop=True)
    if raw.empty:
        raise ValueError("읽을 내용이 없습니다.")
    return raw


def _mapping(columns: list[str]) -> dict[str, str]:
    out: dict[str, str] = {}
    for target, aliases in ALIASES.items():
        wanted = {_norm(v) for v in aliases}
        hit = next((c for c in columns if _norm(c) in wanted), None)
        if hit is not None:
            out[target] = hit
    return out


def blank_template(product_names: Iterable[str]) -> bytes:
    """현재 사업장의 혼합제품명을 미리 넣은 구성성분 입력 양식."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    names = [str(v).strip() for v in product_names if str(v).strip()]
    wb = Workbook()
    ws = wb.active
    ws.title = "혼합물 구성성분"
    ws.append(list(COMPONENT_COLUMNS))
    for cell in ws[1]:
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.font = Font(color="FFFFFF", bold=True)
    ws.column_dimensions["A"].width = 32
    ws.column_dimensions["B"].width = 18
    ws.column_dimensions["C"].width = 14

    # 제품당 5줄을 미리 준비한다. 남는 줄은 비워 둬도 된다.
    for name in names:
        for _ in range(5):
            ws.append([name, "", ""])

    guide = wb.create_sheet("작성 안내")
    lines = (
        ["이 파일은 기본 물질목록에 '혼합물'이 있을 때만 작성합니다."],
        ["제품 SDS 제3항(구성성분의 명칭 및 함유량)을 열어 보고 그대로 옮겨 적으세요."],
        ["혼합제품명: 첫 번째 물질목록에 적은 제품명과 같아야 합니다."],
        ["CAS No.: SDS 제3항에 적힌 각 구성성분의 CAS 번호입니다. 모든 성분 행에 반드시 적습니다."],
        ["함량(%): SDS 제3항의 각 구성성분 함량입니다. 범위로 표시되어 있으면 안전한 판정을 위해 상한값을 적습니다."],
        ["구성성분명은 적지 않습니다. 법적 물질 식별은 CAS No.를 기준으로 합니다."],
        ["제품당 준비된 5줄이 부족하면 행을 추가하고, 남는 줄은 비워 두면 됩니다."],
        ["예: 세척제 A / 67-64-1 / 50"],
        ["예: 세척제 A / 108-88-3 / 30"],
    )
    for line in lines:
        guide.append(line)
    guide.column_dimensions["A"].width = 130

    out = BytesIO()
    wb.save(out)
    return out.getvalue()


def check(data: bytes, file_name: str, allowed_products: Iterable[str]) -> ComponentCheck:
    raw = _read(data, file_name)
    header = 0
    best = -1
    for idx in range(min(10, len(raw))):
        cols = [_clean(v) for v in raw.iloc[idx].tolist()]
        hits = len(_mapping(cols))
        if hits > best:
            best, header = hits, idx
    columns = [_clean(v) or f"열{i+1}" for i, v in enumerate(raw.iloc[header].tolist())]
    mapping = _mapping(columns)
    missing = [col for col in COMPONENT_COLUMNS if col not in mapping]
    result = ComponentCheck(sha256=hashlib.sha256(data).hexdigest(), file_name=file_name)
    if missing:
        result.errors.append("필수 열을 찾지 못했습니다: " + ", ".join(missing))
        return result

    body = raw.iloc[header + 1:].copy()
    body.columns = columns
    allowed = {_norm(name): str(name).strip() for name in allowed_products if str(name).strip()}
    seen: set[tuple[str, str]] = set()
    totals: dict[str, float] = {}

    for excel_row, (_, row) in enumerate(body.iterrows(), start=header + 2):
        product = _clean(row.get(mapping["혼합제품명"]))
        cas = _clean(row.get(mapping["CAS No."]))
        pct_text = _clean(row.get(mapping["함량(%)"]))
        if not product and not cas and not pct_text:
            continue
        if not cas and not pct_text:
            # 템플릿의 미사용 준비행은 제품명이 미리 들어 있으므로 무시한다.
            continue

        canonical = allowed.get(_norm(product))
        if canonical is None:
            result.errors.append(
                f"{excel_row}행: 혼합제품명 '{product or '(빈칸)'}'을 기본 물질목록의 혼합제품에서 찾을 수 없습니다."
            )
            continue
        if not cas_valid(cas):
            result.errors.append(f"{excel_row}행({canonical}): CAS No. '{cas or '(빈칸)'}'를 확인해 주세요.")
            continue
        try:
            pct = float(pct_text.replace(",", ""))
        except ValueError:
            result.errors.append(f"{excel_row}행({canonical}, {cas}): 함량(%)을 숫자로 적어 주세요.")
            continue
        if pct <= 0 or pct > 100:
            result.errors.append(f"{excel_row}행({canonical}, {cas}): 함량(%)은 0 초과 100 이하이어야 합니다.")
            continue
        key = (_norm(canonical), cas)
        if key in seen:
            result.errors.append(f"{excel_row}행({canonical}): CAS {cas}가 중복되었습니다.")
            continue
        seen.add(key)
        totals[canonical] = totals.get(canonical, 0.0) + pct
        result.rows.append({"혼합제품명": canonical, "CAS No.": cas, "함량(%)": pct})

    by_product = {row["혼합제품명"] for row in result.rows}
    for product in allowed.values():
        if product not in by_product:
            result.errors.append(f"{product}: SDS 제3항 구성성분을 한 줄 이상 입력해 주세요.")
    for product, total in totals.items():
        if total > 100.5:
            result.errors.append(f"{product}: 입력한 구성성분 함량 합이 {total:g}%로 100%를 넘습니다.")
        elif total < 99.5:
            result.warnings.append(
                f"{product}: 입력한 구성성분 함량 합이 {total:g}%입니다. 나머지 {100-total:g}%에 규제 대상 성분이 없는지 SDS 제3항을 다시 확인해 주세요."
            )
    return result


def save_to_project(project, rows: list[dict[str, Any]], *, sds_confirmed: bool) -> None:
    """검증된 구성성분을 현재 프로젝트에 저장한다."""
    from . import cap_chemical_workspace as chem
    from . import cap_judgement as judgement

    _, products = chem._rows(project)
    unresolved = set(judgement.composition_rows(project))
    classes = []
    components = []
    for number in sorted(unresolved):
        if not (1 <= number <= len(products)):
            continue
        product = products[number - 1]
        if not judgement._mixture_yes(product.get("혼합물 여부")):
            raise ValueError(
                f"{number}행은 혼합물로 확인되지 않았습니다. 기본 물질목록에서 단일물질/혼합물 구분을 먼저 확인해 주세요."
            )
        name = str(product.get("제품명") or product.get("물질명") or "").strip()
        classes.append({"행": number, "구분": "혼합물", "CAS No.": ""})
        for row in rows:
            if _norm(row.get("혼합제품명")) == _norm(name):
                components.append({
                    "제품목록행번호": number,
                    "CAS No.": row.get("CAS No."),
                    "함량(%)": row.get("함량(%)"),
                    "성분명(선택)": "",
                })
    judgement.save_composition(project, classes, components, sds_confirmed=sds_confirmed)
