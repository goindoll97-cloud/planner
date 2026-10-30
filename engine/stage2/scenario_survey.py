"""사업장 경계(지도에서 그림)와 주변 보호대상 조사표.

화면에서 사람이 직접 하는 일은 지도에서 사업장 경계를 그리는 것 하나입니다.
- 주소점의 필지 경계(브이월드 연속지적도)를 초안으로 불러올 수 있고, 사람이 지도에서 고쳐 그립니다.
- 도면에서 뽑은 경계는 GeoJSON이나 꼭짓점 좌표 CSV(경도·위도)로 올릴 수 있습니다.
- 경계를 저장하면 주변 지도 객체까지의 거리를 그 경계에서 잽니다.
보호대상의 분류·거리·인원은 화면에서 고르지 않습니다. 프로그램이 회사 작성용 엑셀의 ‘주변 보호대상’ 시트에 지도 조회 값을
미리 채워 주면(company_workbook), 사람이 고쳐 자료 접수 ①에 다시 올려 별지 제8호에 들어갑니다. 분류 제안은 이름 규칙과(있으면) 로컬 AI로 하되,
법정 분류 목록에 없는 값은 버립니다.
"""
from __future__ import annotations

import csv
from datetime import date
from io import BytesIO, StringIO
import json
from typing import Any

from . import cap_form8_workspace as f8
from . import cap_surroundings as survey
from . import evidence_workspace as ew
from .project import Stage2Project

TARGET_SHEET, METHOD_SHEET, GUIDE_SHEET = "보호대상", "조사 방법", "안내"
TARGET_SHEET_NAME = "주변 보호대상"          # 회사 작성용 엑셀의 시트 이름
FILENAME = "주변보호대상_조사표.xlsx"
TARGET_COLUMNS = ("보호대상 명칭", "보호대상 구분", "세부유형", "주소·위치", "사업장 경계와 거리(m)", "GIS/현장 근거",
                  "거주민수", "근로자수")
RANGE_M = 500.0
LON_HEADERS = ("경도", "lon", "lng", "longitude", "x")
LAT_HEADERS = ("위도", "lat", "latitude", "y")


def _today() -> str:
    return date.today().isoformat()


# ---------------------------------------------------------------- 경계

def locate(address: str, *, geocoders=None) -> tuple[tuple[float, float] | None, str]:
    """주소 → ((경도, 위도), 어디서 찾았는지). 브이월드 다음 카카오. 못 찾으면 (None, 이유)."""
    from . import cap_site_lookup as kakao, cap_vworld_lookup as vworld

    address = " ".join(str(address or "").split())
    if not address:
        return None, "주소를 적어 주세요."
    geocoders = geocoders or (("브이월드", vworld), ("카카오", kakao))
    tried = []
    for title, provider in geocoders:
        if not provider.api_key():
            continue
        tried.append(title)
        try:
            point = provider.geocode(address)
        except Exception:
            continue
        if point:
            lon, lat = float(point[0]), float(point[1])
            if 124 <= lon <= 132 and 33 <= lat <= 39:
                return (lon, lat), title
    if not tried:
        return None, "브이월드·카카오 키가 없어 주소를 찾을 수 없습니다. 지도를 옮겨 직접 그리거나 경계 파일을 올려 주세요."
    return None, f"{'·'.join(tried)}에서 이 주소를 찾지 못했습니다. 도로명주소나 지번주소로 다시 적어 주세요."


def toggle_parcel(parcels: list[dict], lat: float, lon: float, *, feature=None) -> tuple[list[dict], str]:
    """지도에서 클릭한 점: 이미 고른 필지 안이면 빼고, 아니면 그 점의 필지를 더합니다. (새 목록, 안내)."""
    from shapely.geometry import Point, shape
    from . import cap_vworld_lookup as vworld

    for index, parcel in enumerate(parcels):
        try:
            if shape(parcel["geometry"]).covers(Point(lon, lat)):
                return parcels[:index] + parcels[index + 1:], f"필지 {parcel.get('label', '')}를 뺐습니다."
        except (KeyError, ValueError, TypeError, AttributeError):
            continue
    found, note = (feature or vworld.parcel_feature)(lat, lon)
    if not found:
        return parcels, note
    if any(p.get("id") == found["id"] for p in parcels):
        return parcels, "이미 고른 필지입니다."
    return parcels + [found], f"필지 {found.get('label', '')}를 더했습니다."


def parcels_geometry(parcels: list[dict]) -> dict:
    """고른 필지를 합친 경계(맞닿은 필지는 하나의 면으로)."""
    return survey.boundary_from_drawings([{"type": "Feature", "geometry": p["geometry"], "properties": {}} for p in parcels])


def parcel_draft(project: Stage2Project, *, geocode=None, parcel=None) -> tuple[dict | None, str]:
    """주소점이 든 필지 경계를 초안으로. (도형, 설명) — 실패하면 (None, 이유)."""
    from . import cap_vworld_lookup as vworld

    address = f8.address(project)
    if not address:
        return None, "사업장 주소가 없습니다. 사업장 일반정보 파일에 주소를 적어 올려 주세요."
    geocode = geocode or vworld.geocode
    parcel = parcel or vworld.parcel_boundary
    try:
        point = geocode(address)
    except Exception as exc:
        return None, f"주소 좌표를 찾지 못했습니다({type(exc).__name__}). 지도에 직접 그려 주세요."
    if not point:
        return None, "주소 좌표를 찾지 못했습니다(브이월드 키·주소 확인). 지도에 직접 그려 주세요."
    lon, lat = float(point[0]), float(point[1])
    return parcel(lat, lon)


def confirm(project: Stage2Project, geometry: dict, origin: str, *, collect=None, refresh=None) -> dict[str, Any]:
    """‘이 경계로 확인’: 경계를 저장하고, 경계 기준으로 주변 지도를 조회하고, 조사표 후보 거리를 재고, 시나리오를 다시 계산합니다.
    지도 조회가 안 되어도(키 없음 등) 경계 저장은 그대로 두고 이유를 알려 줍니다."""
    from . import scenario_pipeline as sp

    payload = survey.save_boundary(project, geometry, f"{origin or '화면 지도에서 그린 경계'}, {_today()}", confirmed=True)
    notices: list[str] = []
    try:
        (collect or survey.collect)(project, expand_to_boundary=True)
    except Exception as exc:          # 키 없음·조회 실패: 경계는 저장됨
        notices.append(f"주변 지도 조회를 하지 못했습니다({exc if isinstance(exc, ValueError) else type(exc).__name__}). "
                       "주변 보호대상 조사표에 직접 적어 올려 주세요.")
    data = survey.survey(project)
    notices += [str(n) for n in data.get("notices", []) if "실패" in str(n) or "미완료" in str(n)]
    candidates = candidate_rows(project)
    try:
        (refresh or sp.refresh)(project)
    except Exception as exc:
        notices.append(f"사고시나리오를 다시 계산하지 못했습니다({type(exc).__name__}).")
    return {"boundary": payload, "objects": len(data.get("objects", [])), "candidates": len(candidates), "notices": notices}


def boundary_from_csv(content: bytes) -> dict:
    """꼭짓점 좌표 CSV(머리글 경도·위도, 순서대로) → 면 도형. 처음과 끝이 다르면 닫습니다."""
    text = content.decode("utf-8-sig", errors="replace")
    reader = csv.reader(StringIO(text))
    rows = [r for r in reader if any(c.strip() for c in r)]
    if not rows:
        raise ValueError("꼭짓점 좌표가 없습니다.")
    header = [ew.norm(c) for c in rows[0]]
    lon_at = next((i for i, h in enumerate(header) if h in LON_HEADERS), None)
    lat_at = next((i for i, h in enumerate(header) if h in LAT_HEADERS), None)
    if lon_at is None or lat_at is None:
        raise ValueError("머리글에 ‘경도’와 ‘위도’ 열이 있어야 합니다(EPSG:4326 경도·위도).")
    points = []
    for row in rows[1:]:
        try:
            points.append([float(row[lon_at]), float(row[lat_at])])
        except (ValueError, IndexError):
            raise ValueError("숫자가 아닌 좌표가 있습니다. 경도·위도를 소수로 적어 주세요.") from None
    if len(points) < 3:
        raise ValueError("꼭짓점이 3개 이상 있어야 면 경계가 됩니다.")
    if points[0] != points[-1]:
        points.append(points[0])
    geometry = {"type": "Polygon", "coordinates": [points]}
    survey.checked_geometry(geometry, polygon_only=True)
    return geometry


def boundary_from_geojson(content: bytes) -> dict:
    payload = json.loads(content.decode("utf-8-sig"))
    items = payload.get("features", []) if payload.get("type") == "FeatureCollection" else [payload]
    return survey.boundary_from_drawings(items)


def save_drawn_boundary(project: Stage2Project, drawings: list[dict], origin: str) -> dict:
    """지도에서 그린(또는 불러와 고친) 경계를 저장합니다. 경계를 그린 것 자체가 사람의 확인이라 따로 묻지 않습니다."""
    geometry = survey.boundary_from_drawings(drawings)
    source = f"{origin or '화면 지도에서 그린 경계'}, {_today()}"
    return survey.save_boundary(project, geometry, source, confirmed=True)


# ---------------------------------------------------------------- 조사표

def _ai_categories(names: list[str], client) -> dict[str, tuple[str, str]]:
    """로컬 AI의 갑종·을종·환경수용체 제안. 법정 분류 목록(f8.SUBTYPES)에 없는 답은 버립니다."""
    if client is None or not names:
        return {}
    options = {category: list(subtypes) for category, subtypes in f8.SUBTYPES.items()}
    instructions = ("너는 화학사고예방관리계획서 별지 제8호의 주변 보호대상 분류를 제안하는 도우미다. 시설 이름만 보고 "
                    "categories의 구분과 세부유형 중 하나를 고른다. 이름만으로 알 수 없으면 빈 문자열로 둔다. 새 분류를 만들지 않는다. "
                    "응답은 {\"items\": [{\"name\": \"...\", \"category\": \"...\", \"subtype\": \"...\"}]} JSON 하나다.")
    try:
        answer = client.generate_json(instructions=instructions,
                                      prompt=json.dumps({"categories": options, "names": names}, ensure_ascii=False))
    except Exception:
        return {}
    out = {}
    for item in (answer or {}).get("items") or []:
        if not isinstance(item, dict):
            continue
        name, category, subtype = (str(item.get(k) or "").strip() for k in ("name", "category", "subtype"))
        if name in names and category in f8.SUBTYPES and subtype in f8.SUBTYPES[category]:
            out[name] = (category, subtype)
    return out


def candidate_rows(project: Stage2Project, client=None) -> list[dict[str, Any]]:
    """지도 조회 객체 중 경계 밖 500m 안의 것(이름 있는 것) → 조사표 줄 초안. 경계를 저장하기 전에는 비어 있습니다."""
    if survey.boundary(project) is None or not survey.survey(project) or not survey.is_current(project):
        return []
    best: dict[str, dict[str, Any]] = {}
    for row in survey.annotate(project):
        name = str(row.get("name") or "").strip()
        distance = row.get("boundary_distance_m")
        if not name or name.startswith("명칭 미제공") or row["relation"] in ("INSIDE", "POINT_INSIDE", "NO_GEOMETRY"):
            continue
        if distance is None or distance > RANGE_M:
            continue
        key = f8._norm(name)
        if key not in best or distance < best[key]["_distance"]:
            best[key] = {**row, "_distance": distance}
    suggestions = _ai_categories([r["name"] for r in best.values()], client)
    checked = (survey.survey(project).get("checked_at") or "")[:10] or _today()
    rows = []
    for row in sorted(best.values(), key=lambda r: r["_distance"]):
        category, subtype, _ = f8.classification_suggestion(row["name"])
        how = "이름 규칙 제안"
        if not category and row["name"] in suggestions:
            category, subtype = suggestions[row["name"]]
            how = "AI 제안"
        source = str(row.get("source") or "지도 조회").split(" · ")[0]
        rows.append({"보호대상 명칭": row["name"], "보호대상 구분": category, "세부유형": subtype,
                     "주소·위치": str(row.get("address") or "").strip() or "지도 표시 위치",
                     "사업장 경계와 거리(m)": round(float(row["_distance"])),
                     "GIS/현장 근거": f"{source} 조회 {checked}, 그린 사업장 경계에서 최소거리" + (f" · 분류는 {how}" if category else ""),
                     "거주민수": "", "근로자수": ""})
    return rows


def workbook(project: Stage2Project, client=None) -> bytes:
    """주변 보호대상 조사표. 이미 올린 조사표의 줄 + 지도 조회 후보(새 이름만). 사람이 고쳐 다시 올립니다."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    known = [{c: row.get(c, "") for c in TARGET_COLUMNS} for row in f8.saved_rows(project)]
    names = {f8._norm(r["보호대상 명칭"]) for r in known}
    rows = known + [r for r in candidate_rows(project, client) if f8._norm(r["보호대상 명칭"]) not in names]
    book = Workbook()
    sheet = book.active
    sheet.title = TARGET_SHEET
    bold, fill = Font(bold=True), PatternFill("solid", fgColor="DDEBF7")
    sheet.append(list(TARGET_COLUMNS))
    for cell in sheet[1]:
        cell.font, cell.fill = bold, fill
        sheet.column_dimensions[cell.column_letter].width = 22
    for row in rows:
        sheet.append([row.get(c, "") for c in TARGET_COLUMNS])
    last = max(len(rows) + 1, 2) + 20
    category = DataValidation(type="list", formula1='"갑종,을종,환경수용체"', allow_blank=True, showErrorMessage=False)
    sheet.add_data_validation(category)
    category.add(f"B2:B{last}")
    sheet.freeze_panes = "B2"

    review = f8.review(project)
    boundary = survey.boundary_record(project)
    data = survey.survey(project)
    method = review.get("경계·거리 확인방법") or (
        f"지도에서 그린 사업장 경계({boundary.get('source', '')})에서 각 대상까지 최소거리를 재고 경계 밖 500m 안을 확인"
        if boundary else "")
    source = review.get("사용자료·확인일") or (f"브이월드 공간자료 조회 {(data.get('checked_at') or '')[:10]}" if data else "")
    numbers = review.get("지도 번호") or (",".join(str(i) for i in range(1, len(rows) + 1)) if rows else "")
    checks = review.get("확인항목") or {}
    method_sheet = book.create_sheet(METHOD_SHEET)
    method_sheet.append(["항목", "값"])
    for cell in method_sheet[1]:
        cell.font, cell.fill = bold, fill
    lines = [("확인한 지도·도면과 확인일", source), ("경계·거리 확인 방법", method), ("지도 번호", numbers),
             ("500m 범위 전체 확인", "예" if f8.scope_reviewed(project) else ""),
             ("500m 안 보호대상 없음", "예" if f8.declared_no_target(project) else "")]
    lines += [(label, "예" if checks.get(key) else "") for key, label in f8.REVIEW_ITEMS]
    for line in lines:
        method_sheet.append(list(line))
    yes_no = DataValidation(type="list", formula1='"예,아니오"', allow_blank=True, showErrorMessage=False)
    method_sheet.add_data_validation(yes_no)
    yes_no.add(f"B5:B{len(lines) + 1}")
    method_sheet.column_dimensions["A"].width = 44
    method_sheet.column_dimensions["B"].width = 70

    guide = book.create_sheet(GUIDE_SHEET)
    for text in (
        "주변 보호대상 조사표 — 별지 제8호에 들어갑니다. 고쳐서 ‘사업장 업무 > 자료 올리기’에 그대로 올리세요.",
        "‘보호대상’ 시트: 지도 조회로 찾은 경계 밖 500m 안의 이름 있는 시설을 미리 적었습니다. 보호대상이 아니면 줄을 지우고, "
        "빠진 대상은 줄을 더하세요. 구분·세부유형은 제안값이니 「화학사고예방관리계획서 작성 등에 관한 규정」 별표 4로 확인해 고치세요.",
        "거리는 화면 지도에 그린 사업장 경계에서 잰 최소거리입니다. 거주민수·근로자수는 모르면 비워 두면 0으로 셉니다.",
        "‘조사 방법’ 시트: 지도·현장으로 확인한 뒤 확인한 항목에 ‘예’를 적으세요. 보호대상이 하나도 없으면 ‘500m 안 보호대상 없음’에 ‘예’.",
        "열 이름과 시트 이름은 바꾸지 마세요.",
    ):
        guide.append([text])
        guide.cell(row=guide.max_row, column=1).alignment = Alignment(wrap_text=True)
    guide.column_dimensions["A"].width = 120
    out = BytesIO()
    book.save(out)
    return out.getvalue()


def company_workbook(project: Stage2Project, original: bytes | None = None, client=None) -> bytes:
    """회사 작성용 엑셀(지금 반영한 원본, 없으면 빈 양식)의 ‘주변 보호대상’ 시트에 지도 조회 후보를 더하고,
    ‘사업장 정보’의 조사 근거 칸 중 빈 것만 채웁니다. 사람이 고쳐 ①에 그대로 다시 올리면 별지 제8호가 채워집니다.
    회사가 이미 적은 값은 바꾸지 않고, 확인 여부(예/아니오) 칸은 사람이 적도록 비워 둡니다."""
    from openpyxl import load_workbook
    from . import cap_upload_template as tpl

    book = load_workbook(BytesIO(original or tpl.build_company()))
    sheet = book[TARGET_SHEET_NAME] if TARGET_SHEET_NAME in book.sheetnames else None
    added = 0
    if sheet is not None:
        headers = {str(c.value).strip(): c.column for c in sheet[1] if c.value}
        name_col = headers.get("보호대상 명칭", 1)
        names = {f8._norm(sheet.cell(row=r, column=name_col).value) for r in range(2, sheet.max_row + 1)
                 if sheet.cell(row=r, column=name_col).value}
        row = 2
        for candidate in candidate_rows(project, client):
            if f8._norm(candidate["보호대상 명칭"]) in names:
                continue
            while sheet.cell(row=row, column=name_col).value not in (None, ""):
                row += 1
            for column, value in candidate.items():
                if column in headers and value not in (None, ""):
                    sheet.cell(row=row, column=headers[column]).value = value
            names.add(f8._norm(candidate["보호대상 명칭"]))
            added += 1
        total = sum(1 for r in range(2, sheet.max_row + 1) if sheet.cell(row=r, column=name_col).value)
    else:
        total = 0
    if tpl.SITE_SHEET in book.sheetnames:
        site = book[tpl.SITE_SHEET]
        boundary = survey.boundary_record(project) if survey.boundary(project) is not None else {}
        data = survey.survey(project) if survey.survey(project) and survey.is_current(project) else {}
        defaults = {
            "확인한 지도·도면과 확인일": f"브이월드 공간자료 조회 {(data.get('checked_at') or '')[:10]}" if data else "",
            "경계·거리 확인 방법": (f"지도에서 확인한 사업장 경계({boundary.get('source', '')})에서 각 대상까지 최소거리를 재고 "
                              "경계 밖 500m 안을 확인") if boundary else "",
            "지도 번호": ",".join(str(i) for i in range(1, total + 1)) if added and total else "",
        }
        for r in range(2, site.max_row + 1):
            label = str(site.cell(row=r, column=1).value or "").strip()
            if defaults.get(label) and site.cell(row=r, column=2).value in (None, ""):
                site.cell(row=r, column=2).value = defaults[label]
    out = BytesIO()
    book.save(out)
    return out.getvalue()
