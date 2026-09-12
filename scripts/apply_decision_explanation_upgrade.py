from pathlib import Path

STAGE1 = Path('engine/stage1_workbook.py')
TEMPLATE = Path('engine/template.py')
INIT = Path('engine/__init__.py')
UI = Path('ui/diagnosis_page.py')
TEST = Path('tests/test_decision_explanation_contract.py')

stage1 = STAGE1.read_text(encoding='utf-8')

if 'def _psm_subject_explanation(' not in stage1:
    anchor = '\ndef assess_stage1_from_workbook(intake: IntakeData) -> Stage1WorkbookDecision:\n'
    helper = r'''

def _fmt_num(value: object, digits: int = 3) -> str:
    try:
        number = float(value)
    except (TypeError, ValueError):
        return ""
    text = f"{number:.{digits}f}".rstrip("0").rstrip(".")
    return text or "0"


def _psm_industry_clause(code: str) -> str:
    return {
        "19210": "제1호",
        "19229": "제2호",
        "20111": "제3호",
        "20202": "제3호",
        "20311": "제4호",
        "20312": "제5호",
        "20321": "제6호",
        "20494": "제7호",
    }.get(str(code or "").strip(), "")


def _psm_subject_explanation(psm_result, psm_facts: PSMFollowupFacts) -> str:
    parts = [
        "「산업안전보건법」 제44조제1항은 사업장에 대통령령으로 정하는 유해하거나 위험한 설비가 있는 경우 공정안전보고서를 작성·제출하도록 규정하고 있습니다."
    ]

    if psm_result.industry_trigger:
        code = str(psm_result.base.industry_code or "").strip()
        industry = str(psm_result.base.industry_match or "").strip()
        clause = _psm_industry_clause(code)
        if code == "20202":
            applicable_items = []
            for item_no, label in ((1, "인화성 가스"), (2, "인화성 액체")):
                answer = psm_facts.property_answers.get(item_no, PSMPropertyAnswer())
                if answer.applicable is True:
                    applicable_items.append(f"별표 13 제{item_no}호 {label}")
            condition = ", ".join(applicable_items) or "별표 13 제1호 또는 제2호"
            parts.append(
                f"귀사가 입력한 한국표준산업분류 코드 {code}는 「산업안전보건법 시행령」 제43조제1항{clause}의 '{industry}'에 해당하고, 같은 호 단서에서 요구하는 {condition} 해당 사실도 확인되었습니다."
            )
        else:
            parts.append(
                f"귀사가 입력한 한국표준산업분류 코드 {code}는 「산업안전보건법 시행령」 제43조제1항{clause}의 '{industry}'에 해당합니다."
            )

    if psm_result.quantity_trigger:
        r_text = _fmt_num(psm_result.r_value, 4)
        contributors = sorted(psm_result.ratio_lines, key=lambda row: float(row.controlling_ratio or 0), reverse=True)
        details = []
        for row in contributors[:3]:
            ratio = _fmt_num(row.controlling_ratio, 4)
            details.append(
                f"별표 13 제{row.legal_item_no}호 {row.legal_substance}({row.controlling_basis} 기준 C/T={ratio})"
            )
        detail_text = ", ".join(details)
        sentence = f"또한 「산업안전보건법 시행령」 별표 13 비고 제7호에 따라 산정한 합산한 값(R)이 {r_text}로 1 이상입니다."
        if detail_text:
            sentence += f" 주요 기여 항목은 {detail_text}입니다."
        parts.append(sentence)

    parts.append(
        "회사 입력파일에서 「산업안전보건법 시행령」 제43조제2항의 제외설비에는 해당하지 않는 것으로 확인되었으므로, 위 해당 사유를 종합하여 공정안전보고서 제출 대상으로 판정했습니다."
    )
    return "\n\n".join(parts)


def _cap_row_fact(row: dict[str, Any]) -> str:
    name = str(row.get("product_name") or row.get("legal_substance") or "해당 유해화학물질").strip()
    cas = str(row.get("cas") or "").strip()
    holding = row.get("calculated_max_holding_ton")
    if holding is None:
        holding = row.get("confirmed_max_holding_ton")
    lower = row.get("lower_quantity_ton")
    upper = row.get("upper_quantity_ton")
    identity = f"{name}" + (f"(CAS {cas})" if cas else "")
    values = []
    if holding is not None:
        values.append(f"최대보유량 {_fmt_num(holding)} ton")
    if lower is not None:
        values.append(f"하위 규정수량 {_fmt_num(lower)} ton")
    if upper is not None:
        values.append(f"상위 규정수량 {_fmt_num(upper)} ton")
    return identity + (": " + ", ".join(values) if values else "")


def _cap_subject_explanation(final_cap, quantity_rows: list[dict[str, Any]]) -> str:
    normalized = normalize_quantity_evidence(quantity_rows)
    if final_cap.status not in {"REQUIRED_GROUP_1", "REQUIRED_GROUP_2"}:
        return ""

    parts = [
        "「화학물질관리법」 제23조제1항은 유해화학물질 취급시설을 설치·운영하려는 자에게 화학사고예방관리계획서를 작성하여 제출하도록 규정하고 있습니다."
    ]
    if final_cap.status == "REQUIRED_GROUP_1":
        relevant = [row for row in normalized if row.get("decision_level") == "UPPER"]
        if relevant:
            facts = "; ".join(_cap_row_fact(row) for row in relevant[:3])
            parts.append(
                f"귀사의 물질별 최대보유량을 「유해화학물질의 규정수량에 관한 규정」 제4조 및 별표 4에 따라 확인한 결과, 상위 규정수량 이상인 물질이 확인되었습니다. {facts}."
            )
        parts.append(
            "또한 회사 입력파일에서 해당 물질을 상위 규정수량 이상 취급하는 주요취급시설을 운영하는 것으로 확인되었습니다. 이는 「화학물질관리법 시행규칙」 제19조제8항 및 「화학사고예방관리계획서 작성 등에 관한 규정」 제2조제1항제12의1의 1군 사업장 요건에 해당합니다."
        )
        parts.append(
            "아울러 「화학물질관리법」 제23조제1항 단서의 법정 예외에 해당하지 않는 것으로 확인되어, 화학사고예방관리계획서 작성수준 1군 사업장으로 판정했습니다."
        )
    else:
        relevant = [row for row in normalized if row.get("decision_level") == "LOWER"]
        if relevant:
            facts = "; ".join(_cap_row_fact(row) for row in relevant[:3])
            parts.append(
                f"귀사의 물질별 최대보유량을 「유해화학물질의 규정수량에 관한 규정」 제4조 및 별표 4에 따라 확인한 결과, 하위 규정수량 이상이면서 상위 규정수량 미만인 물질이 확인되었습니다. {facts}."
            )
        parts.append(
            "이는 「화학사고예방관리계획서 작성 등에 관한 규정」 제2조제1항제12의2의 2군 사업장 요건에 해당합니다. 또한 「화학물질관리법」 제23조제1항 단서의 법정 예외에 해당하지 않는 것으로 확인되어, 화학사고예방관리계획서 작성수준 2군 사업장으로 판정했습니다."
        )
    return "\n\n".join(parts)
'''
    if anchor not in stage1:
        raise SystemExit('stage1 helper anchor not found')
    stage1 = stage1.replace(anchor, helper + anchor, 1)

old_psm = 'decision.psm_explanation = "「산업안전보건법 시행령」 제43조제1항의 사업 종류 또는 별표 13 유해·위험물질 규정량 기준에 해당하고, 제43조제2항 제외설비에 해당하지 않는 것으로 회사 입력파일에서 확인되었습니다."'
if old_psm not in stage1:
    raise SystemExit('PSM explanation target not found')
stage1 = stage1.replace(old_psm, 'decision.psm_explanation = _psm_subject_explanation(psm_result, psm_facts)', 1)

old_basis = 'decision.psm_legal_basis = ["「산업안전보건법 시행령」 제43조", "같은 영 별표 13"]'
if old_basis not in stage1:
    raise SystemExit('PSM basis target not found')
stage1 = stage1.replace(old_basis, 'decision.psm_legal_basis = ["「산업안전보건법」 제44조제1항", "「산업안전보건법 시행령」 제43조제1항·제2항", "같은 영 별표 13 및 비고 제7호·제8호(해당 시)"]', 1)

for old in [
    'decision.cap_explanation = "회사 입력파일의 최대보유량, 상위 규정수량, 법적 예외 및 주요취급시설 정보를 기준으로 작성수준을 1군 사업장으로 확인했습니다."',
    'decision.cap_explanation = "회사 입력파일의 최대보유량, 하위·상위 규정수량 및 법적 예외 정보를 기준으로 작성수준을 2군 사업장으로 확인했습니다."',
]:
    if old not in stage1:
        raise SystemExit('CAP explanation target not found')
    stage1 = stage1.replace(old, 'decision.cap_explanation = _cap_subject_explanation(final_cap, quantity_rows)', 1)

STAGE1.write_text(stage1, encoding='utf-8')

ui = UI.read_text(encoding='utf-8')
old_ui = '        if explanation:\n            st.write(explanation)'
if old_ui not in ui:
    raise SystemExit('UI result card target not found')
ui = ui.replace(old_ui, '        if explanation:\n            st.markdown("**판정 근거 설명**")\n            st.write(explanation)', 1)
UI.write_text(ui, encoding='utf-8')

init = INIT.read_text(encoding='utf-8')
init = init.replace('file_name="PSM_CAP_회사입력_작성예시_가이드_v1.0.xlsx"', 'file_name="공정안전보고서_화학사고예방관리계획서_회사입력_작성예시_가이드_v1.0.xlsx"')
init = init.replace('"법령 작성 참고파일 다운로드"', '"공정안전보고서·화학사고예방관리계획서 법령 작성 참고파일 다운로드"')
init = init.replace('file_name="PSM_CAP_법령작성참고_v1.0.xlsx"', 'file_name="공정안전보고서_화학사고예방관리계획서_법령작성참고_v1.0.xlsx"')
INIT.write_text(init, encoding='utf-8')

template = TEMPLATE.read_text(encoding='utf-8')
template = template.replace('"PSM·화학사고예방관리계획서 입력항목별 법령 작성 참고"', '"공정안전보고서·화학사고예방관리계획서 입력항목별 법령 작성 참고"')
template = template.replace('build_reference_sheet("01_PSM_법령참고", psm_rows)', 'build_reference_sheet("01_공정안전보고서_법령참고", psm_rows)')
template = template.replace('build_reference_sheet("02_화사계_법령참고", cap_rows)', 'build_reference_sheet("02_화학사고예방관리계획서_법령참고", cap_rows)')
TEMPLATE.write_text(template, encoding='utf-8')

TEST.write_text(r'''from __future__ import annotations

import unittest
from io import BytesIO
from pathlib import Path

from openpyxl import load_workbook

from engine.template import build_legal_reference_workbook


class DecisionExplanationContractTests(unittest.TestCase):
    def test_stage1_has_company_fact_linked_explanations(self) -> None:
        text = Path('engine/stage1_workbook.py').read_text(encoding='utf-8')
        for expected in [
            '「산업안전보건법」 제44조제1항',
            '한국표준산업분류 코드',
            '합산한 값(R)',
            '「화학물질관리법」 제23조제1항',
            '최대보유량',
            '하위 규정수량',
            '상위 규정수량',
            '주요취급시설',
        ]:
            self.assertIn(expected, text)

    def test_result_card_labels_explanation(self) -> None:
        text = Path('ui/diagnosis_page.py').read_text(encoding='utf-8')
        self.assertIn('판정 근거 설명', text)

    def test_reference_workbook_uses_full_legal_names(self) -> None:
        wb = load_workbook(BytesIO(build_legal_reference_workbook()), data_only=False)
        self.assertEqual(wb.sheetnames, [
            '00_사용안내',
            '01_공정안전보고서_법령참고',
            '02_화학사고예방관리계획서_법령참고',
        ])
        values = [str(cell.value or '') for ws in wb.worksheets for row in ws.iter_rows() for cell in row]
        joined = '\n'.join(values)
        self.assertNotIn('화사계', joined)
        self.assertNotIn('PSM·', joined)

    def test_download_filenames_do_not_use_abbreviations(self) -> None:
        text = Path('engine/__init__.py').read_text(encoding='utf-8')
        self.assertIn('공정안전보고서_화학사고예방관리계획서_법령작성참고_v1.0.xlsx', text)
        self.assertIn('공정안전보고서_화학사고예방관리계획서_회사입력_작성예시_가이드_v1.0.xlsx', text)
        self.assertNotIn('PSM_CAP_법령작성참고', text)


if __name__ == '__main__':
    unittest.main()
''', encoding='utf-8')
